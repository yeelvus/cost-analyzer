#!/opt/anaconda3/bin/python
# -*- coding: utf-8 -*-
"""
Excel 批量包裹替换（图形界面）

用法：
  /opt/anaconda3/bin/python /Users/wyattsun/Desktop/code/excel_wrap_replace.py

流程：
  1. 选择 Excel、工作表、要处理的列
  2. 「查找替换」：输入 查找 xx、替换为 xx，加入规则列表；可一键套用全部规则
  3. 「包裹替换」：粘贴词表 + 自定义前后缀（如 [（人:）xxx]）
  4. 先预览再执行。历史替换会写入「修改记录」并固化，再次处理不重复套改

固化 / 修改记录：
  每次真正替换过的形式会写入 Excel「修改记录」表，并固化到
  同名 .wrap_lock.json。再次处理时，已包裹片段原样保留，不重复套一层。
"""

from __future__ import annotations

import json
import os
import re
import shutil
import sys
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except ImportError:
    print(
        "当前 Python 没有 Tk 图形库。请用：\n"
        "  /opt/anaconda3/bin/python excel_wrap_replace.py"
    )
    sys.exit(1)

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


APP_TITLE = "Excel 批量替换"
DEFAULT_PREFIX = "[（人:）"
DEFAULT_SUFFIX = "]"
LOG_SHEET = "修改记录"
RULES_SHEET = "替换规则"
LOG_HEADERS = (
    "时间",
    "工作表",
    "单元格",
    "列名",
    "原词",
    "固化形式",
    "替换前",
    "替换后",
)
RULES_HEADERS = ("启用", "查找", "替换")
_RULE_ARROW = re.compile(r"\s*(?:→|->|=>|＝|=|\t)\s*")
_RULE_CN = re.compile(
    r"^(?:查找|找|原文)[:：]?\s*(.+?)\s*(?:替换为|替换成|替换|改为|改成|换成)[:：]?\s*(.+)$"
)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def extract_terms(text: str, split_mode: str = "line") -> List[str]:
    """从纯文本抽出待替换词，去空、去重、保序。"""
    raw = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    if split_mode == "punct":
        parts = re.split(r"[,，;；、|\n\t]+", raw)
    else:
        parts = raw.split("\n")

    out: List[str] = []
    seen = set()
    for part in parts:
        s = str(part).strip()
        s = re.sub(r"^[\s\d]+[\.．、\)）]\s*", "", s)
        s = re.sub(r"^[-*•·]+\s*", "", s)
        s = s.strip()
        if not s or s in seen:
            continue
        seen.add(s)
        out.append(s)
    return out


def wrap_term(term: str, prefix: str, suffix: str) -> str:
    return f"{prefix}{term}{suffix}"


@dataclass
class FrozenForm:
    term: str
    prefix: str
    suffix: str
    wrapped: str
    first_seen: str = ""


@dataclass
class ChangeRecord:
    ts: str
    sheet: str
    cell: str
    column: str
    term: str
    wrapped: str
    before: str
    after: str


@dataclass
class ReplaceRule:
    find: str
    replace: str
    enabled: bool = True


def apply_wraps(
    text: str,
    terms: Sequence[str],
    prefix: str,
    suffix: str,
    *,
    skip_existing: bool = True,
    frozen_wrapped: Sequence[str] = (),
) -> Tuple[str, int, List[str]]:
    """
    把 text 里出现的每个 term 换成 前缀+term+后缀。
    已固化 / 已包裹片段先藏起来，再次处理不会重复套一层。
    返回 (新文本, 替换次数, 实际命中的原词列表)。
    """
    if not text:
        return text, 0, []
    work = text
    hidden: List[Tuple[str, str]] = []
    total = 0
    applied: List[str] = []
    seq = 0

    def hide(chunk: str) -> None:
        nonlocal work, seq
        if not chunk or chunk not in work:
            return
        token = f"\u0000H{seq}\u0000"
        seq += 1
        hidden.append((token, chunk))
        work = work.replace(chunk, token)

    frozen = [x for x in frozen_wrapped if x]
    frozen.sort(key=len, reverse=True)
    if skip_existing:
        for chunk in frozen:
            hide(chunk)
        for term in terms:
            wrapped = wrap_term(term, prefix, suffix)
            if wrapped and wrapped != term:
                hide(wrapped)

    for term in terms:
        if not term:
            continue
        wrapped = wrap_term(term, prefix, suffix)
        if wrapped == term:
            continue
        n = work.count(term)
        if n:
            work = work.replace(term, wrapped)
            total += n
            applied.append(term)
        hide(wrapped)

    for token, chunk in hidden:
        work = work.replace(token, chunk)
    return work, total, applied


def parse_replace_line(line: str) -> Optional[Tuple[str, str]]:
    s = str(line or "").strip()
    if not s or s.startswith("#") or s.startswith("//"):
        return None
    m = _RULE_CN.match(s)
    if m:
        a, b = m.group(1).strip(), m.group(2).strip()
        if a:
            return a, b
    parts = _RULE_ARROW.split(s, maxsplit=1)
    if len(parts) == 2 and parts[0].strip():
        return parts[0].strip(), parts[1].strip()
    return None


def parse_replace_rules(text: str) -> List[ReplaceRule]:
    out: List[ReplaceRule] = []
    seen = set()
    for line in (text or "").replace("\r\n", "\n").split("\n"):
        pair = parse_replace_line(line)
        if not pair:
            continue
        find, repl = pair
        if find in seen:
            continue
        seen.add(find)
        out.append(ReplaceRule(find=find, replace=repl, enabled=True))
    return out


def apply_find_replace(
    text: str,
    rules: Sequence[ReplaceRule],
    *,
    skip_existing: bool = True,
    frozen_wrapped: Sequence[str] = (),
) -> Tuple[str, int, List[str]]:
    """按规则顺序把查找词换成替换词。已固化片段先藏起来，避免重复改。"""
    if not text:
        return text, 0, []
    work = text
    hidden: List[Tuple[str, str]] = []
    total = 0
    applied: List[str] = []
    seq = 0

    def hide(chunk: str) -> None:
        nonlocal work, seq
        if not chunk or chunk not in work:
            return
        token = f"\u0000R{seq}\u0000"
        seq += 1
        hidden.append((token, chunk))
        work = work.replace(chunk, token)

    frozen = [x for x in frozen_wrapped if x]
    frozen.sort(key=len, reverse=True)
    if skip_existing:
        for chunk in frozen:
            hide(chunk)
        for rule in rules:
            if rule.enabled and rule.replace and rule.replace != rule.find:
                hide(rule.replace)

    for rule in rules:
        if not rule.enabled or not rule.find:
            continue
        if rule.find == rule.replace:
            continue
        n = work.count(rule.find)
        if n:
            work = work.replace(rule.find, rule.replace)
            total += n
            applied.append(rule.find)
        if rule.replace:
            hide(rule.replace)

    for token, chunk in hidden:
        work = work.replace(token, chunk)
    return work, total, applied


def rules_path_for(excel: Optional[Path]) -> Path:
    if excel:
        return Path(excel).with_name(f"{Path(excel).stem}.replace_rules.json")
    return Path(__file__).resolve().with_name("excel_replace_rules.json")


def load_saved_rules(path: Path) -> List[ReplaceRule]:
    if not path.is_file():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return []
    raw = data.get("rules") if isinstance(data, dict) else data
    out: List[ReplaceRule] = []
    seen = set()
    for item in raw or []:
        if isinstance(item, dict):
            find = str(item.get("find") or item.get("查找") or "").strip()
            repl = str(item.get("replace") or item.get("替换") or "")
            enabled = bool(item.get("enabled", True))
        elif isinstance(item, (list, tuple)) and len(item) >= 2:
            find, repl, enabled = str(item[0]).strip(), str(item[1]), True
        else:
            continue
        if not find or find in seen:
            continue
        seen.add(find)
        out.append(ReplaceRule(find=find, replace=repl, enabled=enabled))
    return out


def save_saved_rules(path: Path, rules: Sequence[ReplaceRule]) -> Path:
    path = Path(path)
    payload = {
        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "rules": [asdict(r) for r in rules],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def _first_row_values(ws) -> List[object]:
    """只读第一行。read_only 工作表禁止用 cell() 逐格扫。"""
    try:
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        return list(row)
    except StopIteration:
        return []


def header_map_from_row(row: Sequence[object]) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for col, raw in enumerate(row, start=1):
        name = _cell_text(raw).strip() or f"第{col}列"
        key = name
        n = 2
        while key in mapping:
            key = f"{name} ({n})"
            n += 1
        mapping[key] = col
    return mapping


def header_map(ws: Worksheet, header_row: int = 1) -> Dict[str, int]:
    if header_row == 1:
        return header_map_from_row(_first_row_values(ws))
    try:
        row = next(
            ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
        )
        return header_map_from_row(row)
    except StopIteration:
        return {}


def rules_from_sheet(ws: Worksheet) -> List[ReplaceRule]:
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    cols = {_cell_text(h).strip(): i for i, h in enumerate(header) if _cell_text(h).strip()}
    i_find = cols.get("查找") if "查找" in cols else cols.get("原词", cols.get("find"))
    i_repl = cols.get("替换") if "替换" in cols else cols.get("固化形式", cols.get("replace"))
    i_on = cols.get("启用")
    if i_find is None:
        return []
    out: List[ReplaceRule] = []
    seen = set()
    for row in rows:
        if i_find >= len(row):
            continue
        find = _cell_text(row[i_find]).strip()
        if not find or find in seen:
            continue
        repl = _cell_text(row[i_repl]) if i_repl is not None and i_repl < len(row) else ""
        on = True
        if i_on is not None and i_on < len(row):
            raw = _cell_text(row[i_on]).strip().lower()
            on = raw not in {"0", "否", "n", "no", "false", "关闭"}
        seen.add(find)
        out.append(ReplaceRule(find=find, replace=repl, enabled=on))
    return out


def ensure_rules_sheet(wb) -> Worksheet:
    if RULES_SHEET in wb.sheetnames:
        ws = wb[RULES_SHEET]
        if _cell_text(ws.cell(1, 1).value) == RULES_HEADERS[0]:
            return ws
    ws = wb[RULES_SHEET] if RULES_SHEET in wb.sheetnames else wb.create_sheet(RULES_SHEET)
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    for i, h in enumerate(RULES_HEADERS, start=1):
        cell = ws.cell(1, i, h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 36
    ws.freeze_panes = "A2"
    return ws


def write_rules_sheet(ws: Worksheet, rules: Sequence[ReplaceRule]) -> None:
    if ws.max_row and ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for i, rule in enumerate(rules, start=2):
        ws.cell(i, 1, "是" if rule.enabled else "否")
        ws.cell(i, 2, rule.find)
        ws.cell(i, 3, rule.replace)
    last = max(2, (ws.max_row or 2))
    ws.auto_filter.ref = f"A1:C{last}"


def merge_rules(*groups: Sequence[ReplaceRule]) -> List[ReplaceRule]:
    out: List[ReplaceRule] = []
    seen = set()
    for group in groups:
        for rule in group:
            find = (rule.find or "").strip()
            if not find or find in seen:
                continue
            seen.add(find)
            out.append(ReplaceRule(find=find, replace=rule.replace, enabled=rule.enabled))
    return out


def lock_path_for(excel: Path) -> Path:
    return Path(excel).with_name(f"{Path(excel).stem}.wrap_lock.json")


def load_lock(excel: Path) -> Tuple[List[FrozenForm], List[ChangeRecord]]:
    path = lock_path_for(excel)
    if not path.is_file():
        return [], []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return [], []
    forms: List[FrozenForm] = []
    seen = set()
    for item in data.get("forms") or []:
        wrapped = str(item.get("wrapped") or "").strip()
        if not wrapped or wrapped in seen:
            continue
        seen.add(wrapped)
        forms.append(
            FrozenForm(
                term=str(item.get("term") or ""),
                prefix=str(item.get("prefix") or ""),
                suffix=str(item.get("suffix") or ""),
                wrapped=wrapped,
                first_seen=str(item.get("first_seen") or ""),
            )
        )
    changes: List[ChangeRecord] = []
    for item in data.get("changes") or []:
        changes.append(
            ChangeRecord(
                ts=str(item.get("ts") or ""),
                sheet=str(item.get("sheet") or ""),
                cell=str(item.get("cell") or ""),
                column=str(item.get("column") or ""),
                term=str(item.get("term") or ""),
                wrapped=str(item.get("wrapped") or ""),
                before=str(item.get("before") or ""),
                after=str(item.get("after") or ""),
            )
        )
    return forms, changes


def save_lock(
    excel: Path,
    forms: Sequence[FrozenForm],
    changes: Sequence[ChangeRecord],
    *,
    keep_changes: int = 5000,
) -> Path:
    path = lock_path_for(excel)
    payload = {
        "excel": Path(excel).name,
        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "forms": [asdict(x) for x in forms],
        "changes": [asdict(x) for x in list(changes)[-keep_changes:]],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def merge_forms(*groups: Sequence[FrozenForm]) -> List[FrozenForm]:
    out: List[FrozenForm] = []
    seen = set()
    for group in groups:
        for item in group:
            key = item.wrapped
            if not key or key in seen:
                continue
            seen.add(key)
            out.append(item)
    out.sort(key=lambda x: -len(x.wrapped))
    return out


def forms_from_log_sheet(ws: Worksheet) -> List[FrozenForm]:
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    cols = {_cell_text(h).strip(): i for i, h in enumerate(header) if _cell_text(h).strip()}
    i_term = cols.get("原词")
    i_wrap = cols.get("固化形式")
    i_time = cols.get("时间")
    if i_wrap is None:
        return []
    forms: List[FrozenForm] = []
    for row in rows:
        if i_wrap >= len(row):
            continue
        wrapped = _cell_text(row[i_wrap]).strip()
        if not wrapped:
            continue
        term = _cell_text(row[i_term]).strip() if i_term is not None and i_term < len(row) else ""
        ts = _cell_text(row[i_time]).strip() if i_time is not None and i_time < len(row) else ""
        prefix, suffix = "", ""
        if term and term in wrapped:
            at = wrapped.find(term)
            prefix = wrapped[:at]
            suffix = wrapped[at + len(term) :]
        forms.append(
            FrozenForm(term=term, prefix=prefix, suffix=suffix, wrapped=wrapped, first_seen=ts)
        )
    return merge_forms(forms)


def ensure_log_sheet(wb) -> Worksheet:
    if LOG_SHEET in wb.sheetnames:
        ws = wb[LOG_SHEET]
        if (ws.max_row or 0) >= 1 and _cell_text(ws.cell(1, 1).value) == LOG_HEADERS[0]:
            return ws
    ws = wb[LOG_SHEET] if LOG_SHEET in wb.sheetnames else wb.create_sheet(LOG_SHEET)
    if _cell_text(ws.cell(1, 1).value) != LOG_HEADERS[0]:
        fill = PatternFill("solid", fgColor="1F4E79")
        font = Font(bold=True, color="FFFFFF")
        for i, h in enumerate(LOG_HEADERS, start=1):
            cell = ws.cell(1, i, h)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center")
        widths = [20, 14, 10, 14, 16, 28, 40, 40]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(LOG_HEADERS))}1"
    return ws


def append_log_sheet(ws: Worksheet, records: Sequence[ChangeRecord]) -> None:
    start = (ws.max_row or 1) + 1
    if start == 2 and not any(_cell_text(ws.cell(2, c).value) for c in range(1, len(LOG_HEADERS) + 1)):
        start = 2
    for i, rec in enumerate(records):
        r = start + i
        ws.cell(r, 1, rec.ts)
        ws.cell(r, 2, rec.sheet)
        ws.cell(r, 3, rec.cell)
        ws.cell(r, 4, rec.column)
        ws.cell(r, 5, rec.term)
        ws.cell(r, 6, rec.wrapped)
        ws.cell(r, 7, rec.before)
        ws.cell(r, 8, rec.after)
    last = ws.max_row or 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LOG_HEADERS))}{last}"


def peek_excel(path: Path) -> Tuple[List[str], Dict[str, Dict[str, int]], List[ReplaceRule]]:
    """只打开一次：工作表名 + 各数据表第一行表头 + 替换规则。不扫明细行。"""
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        names = list(wb.sheetnames)
        headers_by: Dict[str, Dict[str, int]] = {}
        rules: List[ReplaceRule] = []
        for name in names:
            ws = wb[name]
            if name == RULES_SHEET:
                rules = rules_from_sheet(ws)
            elif name == LOG_SHEET:
                continue
            else:
                headers_by[name] = header_map_from_row(_first_row_values(ws))
        return names, headers_by, rules
    finally:
        wb.close()


class ExcelWrapApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("920x740")
        self.root.minsize(820, 640)

        self.excel_path: Optional[Path] = None
        self.wb_sheets: List[str] = []
        self.col_vars: Dict[str, tk.BooleanVar] = {}
        self.col_index: Dict[str, int] = {}
        self.frozen_forms: List[FrozenForm] = []
        self.change_history: List[ChangeRecord] = []
        self.replace_rules: List[ReplaceRule] = []
        self._sheet_headers: Dict[str, Dict[str, int]] = {}

        self._build()

    def _build(self) -> None:
        self.root.geometry("980x820")
        pad = {"padx": 8, "pady": 4}
        main = ttk.Frame(self.root, padding=10)
        main.pack(fill=tk.BOTH, expand=True)

        # —— Excel ——
        box1 = ttk.LabelFrame(main, text="1. Excel 与列", padding=8)
        box1.pack(fill=tk.X)

        row = ttk.Frame(box1)
        row.pack(fill=tk.X)
        ttk.Label(row, text="Excel：", width=8).pack(side=tk.LEFT)
        self.var_excel = tk.StringVar()
        ttk.Entry(row, textvariable=self.var_excel).pack(
            side=tk.LEFT, fill=tk.X, expand=True, **pad
        )
        ttk.Button(row, text="选择文件…", command=self.pick_excel).pack(side=tk.LEFT)

        row = ttk.Frame(box1)
        row.pack(fill=tk.X)
        ttk.Label(row, text="工作表：", width=8).pack(side=tk.LEFT)
        self.var_sheet = tk.StringVar()
        self.cmb_sheet = ttk.Combobox(
            row, textvariable=self.var_sheet, state="readonly", width=36
        )
        self.cmb_sheet.pack(side=tk.LEFT, **pad)
        self.cmb_sheet.bind("<<ComboboxSelected>>", lambda _e: self.reload_columns())
        self.var_skip_header = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            row, text="跳过第 1 行表头", variable=self.var_skip_header
        ).pack(side=tk.LEFT, padx=12)
        ttk.Button(row, text="查看修改记录", command=self.show_change_log).pack(
            side=tk.LEFT, padx=8
        )

        ttk.Label(box1, text="处理列（可多选）：").pack(anchor="w")
        col_wrap = ttk.Frame(box1)
        col_wrap.pack(fill=tk.X)
        self.col_canvas = tk.Canvas(col_wrap, height=88, highlightthickness=0)
        scroll_y = ttk.Scrollbar(col_wrap, orient="vertical", command=self.col_canvas.yview)
        self.col_inner = ttk.Frame(self.col_canvas)
        self.col_inner.bind(
            "<Configure>",
            lambda e: self.col_canvas.configure(scrollregion=self.col_canvas.bbox("all")),
        )
        self.col_canvas.create_window((0, 0), window=self.col_inner, anchor="nw")
        self.col_canvas.configure(yscrollcommand=scroll_y.set)
        self.col_canvas.pack(side=tk.LEFT, fill=tk.X, expand=True)
        scroll_y.pack(side=tk.RIGHT, fill=tk.Y)

        btns = ttk.Frame(box1)
        btns.pack(fill=tk.X)
        ttk.Button(btns, text="全选列", command=lambda: self._set_all_cols(True)).pack(
            side=tk.LEFT
        )
        ttk.Button(btns, text="清空列", command=lambda: self._set_all_cols(False)).pack(
            side=tk.LEFT, padx=6
        )

        nb = ttk.Notebook(main)
        nb.pack(fill=tk.BOTH, expand=True, pady=(8, 0))
        tab_rules = ttk.Frame(nb, padding=6)
        tab_wrap = ttk.Frame(nb, padding=6)
        nb.add(tab_rules, text="查找替换（xx → xx）")
        nb.add(tab_wrap, text="包裹替换")

        # —— 查找替换 ——
        add_row = ttk.Frame(tab_rules)
        add_row.pack(fill=tk.X)
        ttk.Label(add_row, text="查找").pack(side=tk.LEFT)
        self.var_find = tk.StringVar()
        ent_find = ttk.Entry(add_row, textvariable=self.var_find, width=22)
        ent_find.pack(side=tk.LEFT, padx=6)
        ttk.Label(add_row, text="替换为").pack(side=tk.LEFT)
        self.var_repl = tk.StringVar()
        ent_repl = ttk.Entry(add_row, textvariable=self.var_repl, width=22)
        ent_repl.pack(side=tk.LEFT, padx=6)
        ttk.Button(add_row, text="添加规则", command=self.add_replace_rule).pack(
            side=tk.LEFT, padx=4
        )
        ent_find.bind("<Return>", lambda _e: ent_repl.focus_set())
        ent_repl.bind("<Return>", lambda _e: self.add_replace_rule())

        ttk.Label(
            tab_rules,
            text="规则列表（按从上到下顺序依次替换）。也可从 txt 导入：每行 旧 → 新",
        ).pack(anchor="w", pady=(6, 2))

        tree_fr = ttk.Frame(tab_rules)
        tree_fr.pack(fill=tk.BOTH, expand=True)
        cols = ("on", "find", "repl")
        self.tree_rules = ttk.Treeview(tree_fr, columns=cols, show="headings", height=8)
        self.tree_rules.heading("on", text="启用")
        self.tree_rules.heading("find", text="查找")
        self.tree_rules.heading("repl", text="替换为")
        self.tree_rules.column("on", width=50, anchor="center")
        self.tree_rules.column("find", width=260, anchor="w")
        self.tree_rules.column("repl", width=320, anchor="w")
        y_rules = ttk.Scrollbar(tree_fr, orient="vertical", command=self.tree_rules.yview)
        self.tree_rules.configure(yscrollcommand=y_rules.set)
        self.tree_rules.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        y_rules.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree_rules.bind("<Double-1>", lambda _e: self.toggle_rule_enabled())

        rbtns = ttk.Frame(tab_rules)
        rbtns.pack(fill=tk.X, pady=4)
        ttk.Button(rbtns, text="删除选中", command=self.delete_selected_rules).pack(
            side=tk.LEFT
        )
        ttk.Button(rbtns, text="清空规则", command=self.clear_replace_rules).pack(
            side=tk.LEFT, padx=4
        )
        ttk.Button(rbtns, text="从 txt 导入…", command=self.import_rules_txt).pack(
            side=tk.LEFT, padx=4
        )
        ttk.Button(rbtns, text="保存规则", command=self.save_rules_to_disk).pack(
            side=tk.LEFT, padx=4
        )
        ttk.Button(rbtns, text="一键引用全部规则", command=self.cite_all_rules).pack(
            side=tk.LEFT, padx=4
        )
        ttk.Button(rbtns, text="查看全部规则", command=self.show_all_rules).pack(
            side=tk.LEFT, padx=4
        )

        apply_row = ttk.Frame(tab_rules)
        apply_row.pack(fill=tk.X, pady=(6, 0))
        ttk.Button(
            apply_row, text="一键套用全部规则", command=self.do_apply_all_rules
        ).pack(side=tk.LEFT)
        ttk.Button(apply_row, text="预览规则替换", command=self.do_preview_rules).pack(
            side=tk.LEFT, padx=8
        )
        ttk.Button(
            apply_row, text="套用并另存为…", command=self.do_save_as_rules
        ).pack(side=tk.LEFT)

        # —— 词表（包裹） ——
        box2 = ttk.LabelFrame(tab_wrap, text="2. 替换词（纯文本）", padding=8)
        box2.pack(fill=tk.BOTH, expand=True)

        row = ttk.Frame(box2)
        row.pack(fill=tk.X)
        ttk.Button(row, text="从 txt 加载…", command=self.pick_txt).pack(side=tk.LEFT)
        self.var_split = tk.StringVar(value="line")
        ttk.Radiobutton(
            row, text="每行一条", variable=self.var_split, value="line"
        ).pack(side=tk.LEFT, padx=10)
        ttk.Radiobutton(
            row, text="逗号 / 顿号 / 分号拆开", variable=self.var_split, value="punct"
        ).pack(side=tk.LEFT)
        ttk.Label(row, text="（空行忽略，自动去重）").pack(side=tk.LEFT, padx=8)

        self.txt_terms = tk.Text(box2, height=8, wrap="word", font=("PingFang SC", 12))
        self.txt_terms.pack(fill=tk.BOTH, expand=True, pady=4)
        self.txt_terms.insert("1.0", "张三\n李四\n王五")

        # —— 形式 ——
        box3 = ttk.LabelFrame(tab_wrap, text="3. 包裹形式（自定义前后）", padding=8)
        box3.pack(fill=tk.X, pady=(8, 0))

        row = ttk.Frame(box3)
        row.pack(fill=tk.X)
        ttk.Label(row, text="前缀").pack(side=tk.LEFT)
        self.var_prefix = tk.StringVar(value=DEFAULT_PREFIX)
        ttk.Entry(row, textvariable=self.var_prefix, width=22).pack(side=tk.LEFT, padx=6)
        ttk.Label(row, text="　原词　").pack(side=tk.LEFT)
        ttk.Label(row, text="后缀").pack(side=tk.LEFT)
        self.var_suffix = tk.StringVar(value=DEFAULT_SUFFIX)
        ttk.Entry(row, textvariable=self.var_suffix, width=12).pack(side=tk.LEFT, padx=6)

        self.var_preview = tk.StringVar(value="")
        ttk.Label(box3, textvariable=self.var_preview, foreground="#1F4E79").pack(
            anchor="w", pady=2
        )

        row = ttk.Frame(box3)
        row.pack(fill=tk.X)
        self.var_skip_wrapped = tk.BooleanVar(value=True)
        self.var_long_first = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            row,
            text="固化历史形式：已替换内容再次处理时跳过",
            variable=self.var_skip_wrapped,
        ).pack(side=tk.LEFT)
        ttk.Checkbutton(
            row, text="长词优先", variable=self.var_long_first
        ).pack(side=tk.LEFT, padx=12)
        self.var_frozen_info = tk.StringVar(value="已固化 0 种形式")
        ttk.Label(box3, textvariable=self.var_frozen_info, foreground="#666").pack(
            anchor="w", pady=(2, 0)
        )

        self.var_prefix.trace_add("write", lambda *_: self.refresh_preview())
        self.var_suffix.trace_add("write", lambda *_: self.refresh_preview())
        self.txt_terms.bind("<KeyRelease>", lambda _e: self.refresh_preview())

        # —— 操作 ——
        box4 = ttk.Frame(tab_wrap)
        box4.pack(fill=tk.X, pady=8)
        ttk.Button(box4, text="预览（不改文件）", command=self.do_preview).pack(
            side=tk.LEFT
        )
        ttk.Button(box4, text="执行并另存为…", command=self.do_save_as).pack(
            side=tk.LEFT, padx=8
        )
        ttk.Button(box4, text="执行并覆盖原文件（先备份）", command=self.do_overwrite).pack(
            side=tk.LEFT
        )

        # —— 日志 ——
        box5 = ttk.LabelFrame(main, text="日志", padding=6)
        box5.pack(fill=tk.BOTH, expand=True)
        self.txt_log = tk.Text(box5, height=10, wrap="word", font=("Menlo", 11))
        y = ttk.Scrollbar(box5, orient="vertical", command=self.txt_log.yview)
        self.txt_log.configure(yscrollcommand=y.set)
        self.txt_log.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        y.pack(side=tk.RIGHT, fill=tk.Y)

        self.refresh_preview()
        self.refresh_rules_tree()
        self.log("准备就绪。选 Excel 和列后：在「查找替换」里输入 查找/替换，或一键引用全部规则。")

    # ---------- UI helpers ----------
    def log(self, msg: str) -> None:
        ts = datetime.now().strftime("%H:%M:%S")
        self.txt_log.insert(tk.END, f"[{ts}] {msg}\n")
        self.txt_log.see(tk.END)
        self.root.update_idletasks()

    def refresh_preview(self) -> None:
        terms = extract_terms(self.txt_terms.get("1.0", tk.END), self.var_split.get())
        sample = terms[0] if terms else "张三"
        wrapped = wrap_term(sample, self.var_prefix.get(), self.var_suffix.get())
        extra = f"  · 共 {len(terms)} 个词" if terms else ""
        self.var_preview.set(f"效果：{sample}  →  {wrapped}{extra}")

    def _set_all_cols(self, value: bool) -> None:
        for var in self.col_vars.values():
            var.set(value)

    def selected_columns(self) -> List[Tuple[str, int]]:
        out = []
        for name, var in self.col_vars.items():
            if var.get() and name in self.col_index:
                out.append((name, self.col_index[name]))
        return out

    def current_terms(self) -> List[str]:
        terms = extract_terms(self.txt_terms.get("1.0", tk.END), self.var_split.get())
        if self.var_long_first.get():
            terms = sorted(terms, key=lambda s: (-len(s), s))
        return terms

    def frozen_wrapped_list(self) -> List[str]:
        return [x.wrapped for x in self.frozen_forms if x.wrapped]

    def _refresh_frozen_label(self) -> None:
        n = len(self.frozen_forms)
        extra = f"  修改记录 {len(self.change_history)} 条" if self.change_history else ""
        self.var_frozen_info.set(f"已固化 {n} 种形式{extra}（再次处理不会重复包裹）")

    def _load_frozen_from_excel(self, path: Path) -> None:
        forms, changes = load_lock(path)
        self.frozen_forms = merge_forms(forms)
        self.change_history = changes
        self._refresh_frozen_label()

    def _busy(self, on: bool, msg: str = "") -> None:
        try:
            self.root.config(cursor="watch" if on else "")
        except Exception:
            pass
        if msg:
            self.log(msg)
        self.root.update_idletasks()

    # ---------- files ----------
    def pick_excel(self) -> None:
        finance = Path.home() / "Desktop" / "财务数据"
        initial = str(finance) if finance.is_dir() else str(Path.home())
        if self.excel_path:
            initial = str(self.excel_path.parent)
        path = filedialog.askopenfilename(
            title="选择 Excel",
            initialdir=initial,
            filetypes=[
                ("Excel 工作簿", "*.xlsx"),
                ("Excel 宏", "*.xlsm"),
                ("所有文件", "*.*"),
            ],
        )
        if not path:
            return
        self.var_excel.set(path)
        self._busy(True, f"正在读取 {Path(path).name} 的工作表和表头…")
        try:
            self.load_excel(Path(path))
        finally:
            self._busy(False)

    def load_excel(self, path: Path) -> None:
        try:
            names, headers_by, sheet_rules = peek_excel(path)
        except Exception as e:
            messagebox.showerror(APP_TITLE, f"无法打开 Excel：\n{e}")
            return
        self.excel_path = path
        self.wb_sheets = names
        self._sheet_headers = headers_by
        visible = [n for n in names if n not in {LOG_SHEET, RULES_SHEET}] or names
        self.cmb_sheet["values"] = visible
        prefer = "汇总" if "汇总" in visible else visible[0]
        self.var_sheet.set(prefer)
        self._load_frozen_from_excel(path)
        sidecar = load_saved_rules(rules_path_for(path))
        merged = merge_rules(sidecar, sheet_rules) if (sidecar or sheet_rules) else []
        if merged:
            self.replace_rules = merged
            self.refresh_rules_tree()
        self.log(f"已打开：{path.name}  工作表 {len(names)} 个")
        if self.frozen_forms:
            self.log(f"读到已固化形式 {len(self.frozen_forms)} 种，再次处理将跳过这些内容")
        if self.replace_rules:
            self.log(f"已载入替换规则 {len(self.replace_rules)} 条")
        self.reload_columns()

    def reload_columns(self) -> None:
        if not self.excel_path:
            return
        sheet = self.var_sheet.get()
        mapping = self._sheet_headers.get(sheet)
        if mapping is None:
            try:
                wb = load_workbook(
                    self.excel_path, read_only=True, data_only=True, keep_links=False
                )
                try:
                    mapping = header_map_from_row(_first_row_values(wb[sheet]))
                finally:
                    wb.close()
                self._sheet_headers[sheet] = mapping
            except Exception as e:
                messagebox.showerror(APP_TITLE, f"读取表头失败：\n{e}")
                return

        for child in self.col_inner.winfo_children():
            child.destroy()
        self.col_vars.clear()
        self.col_index = mapping

        for i, name in enumerate(mapping):
            var = tk.BooleanVar(value=False)
            self.col_vars[name] = var
            ttk.Checkbutton(self.col_inner, text=name, variable=var).grid(
                row=i // 4, column=i % 4, sticky="w", padx=6, pady=1
            )
        self.log(f"工作表「{sheet}」列：{', '.join(mapping)}")

    def pick_txt(self) -> None:
        path = filedialog.askopenfilename(
            title="选择文本",
            filetypes=[
                ("文本", "*.txt *.csv *.md"),
                ("所有文件", "*.*"),
            ],
        )
        if not path:
            return
        data = None
        for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
            try:
                data = Path(path).read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        if data is None:
            messagebox.showerror(APP_TITLE, "无法识别文本编码。")
            return
        self.txt_terms.delete("1.0", tk.END)
        self.txt_terms.insert("1.0", data)
        self.refresh_preview()
        n = len(extract_terms(data, self.var_split.get()))
        self.log(f"已加载文本 {Path(path).name}，抽出 {n} 个词")

    # ---------- replace core ----------
    def _validate(self) -> Optional[str]:
        if not self.excel_path or not Path(self.excel_path).is_file():
            return "请先选择 Excel 文件。"
        if not self.var_sheet.get():
            return "请选择工作表。"
        if not self.selected_columns():
            return "请至少勾选一列。"
        if not self.current_terms():
            return "词表是空的。请粘贴或加载 txt。"
        if self.var_sheet.get() in {LOG_SHEET, RULES_SHEET}:
            return "请选择数据工作表，不要选「修改记录 / 替换规则」。"
        return None

    def _validate_excel_cols(self) -> Optional[str]:
        if not self.excel_path or not Path(self.excel_path).is_file():
            return "请先选择 Excel 文件。"
        if not self.var_sheet.get():
            return "请选择工作表。"
        if not self.selected_columns():
            return "请至少勾选一列。"
        if self.var_sheet.get() in {LOG_SHEET, RULES_SHEET}:
            return "请选择数据工作表，不要选「修改记录 / 替换规则」。"
        return None

    def enabled_rules(self) -> List[ReplaceRule]:
        return [r for r in self.replace_rules if r.enabled and r.find]

    def refresh_rules_tree(self) -> None:
        tree = self.tree_rules
        tree.delete(*tree.get_children())
        for i, rule in enumerate(self.replace_rules):
            tree.insert(
                "",
                tk.END,
                iid=str(i),
                values=("是" if rule.enabled else "否", rule.find, rule.replace),
            )

    def add_replace_rule(self) -> None:
        find = self.var_find.get().strip()
        repl = self.var_repl.get()
        if not find:
            messagebox.showwarning(APP_TITLE, "请填写「查找」内容。")
            return
        for rule in self.replace_rules:
            if rule.find == find:
                rule.replace = repl
                rule.enabled = True
                self.refresh_rules_tree()
                self.log(f"已更新规则：{find}  →  {repl}")
                self.var_find.set("")
                self.var_repl.set("")
                return
        self.replace_rules.append(ReplaceRule(find=find, replace=repl, enabled=True))
        self.refresh_rules_tree()
        self.log(f"已添加规则：{find}  →  {repl}")
        self.var_find.set("")
        self.var_repl.set("")

    def toggle_rule_enabled(self) -> None:
        sel = self.tree_rules.selection()
        if not sel:
            return
        idx = int(sel[0])
        if 0 <= idx < len(self.replace_rules):
            self.replace_rules[idx].enabled = not self.replace_rules[idx].enabled
            self.refresh_rules_tree()

    def delete_selected_rules(self) -> None:
        sel = sorted((int(i) for i in self.tree_rules.selection()), reverse=True)
        if not sel:
            return
        for idx in sel:
            if 0 <= idx < len(self.replace_rules):
                del self.replace_rules[idx]
        self.refresh_rules_tree()

    def clear_replace_rules(self) -> None:
        if not self.replace_rules:
            return
        if not messagebox.askyesno(APP_TITLE, "清空当前规则列表？"):
            return
        self.replace_rules.clear()
        self.refresh_rules_tree()

    def import_rules_txt(self) -> None:
        path = filedialog.askopenfilename(
            title="导入替换规则（每行：旧 → 新）",
            filetypes=[("文本", "*.txt *.csv *.md *.json"), ("所有文件", "*.*")],
        )
        if not path:
            return
        p = Path(path)
        if p.suffix.lower() == ".json":
            rules = load_saved_rules(p)
        else:
            data = None
            for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
                try:
                    data = p.read_text(encoding=enc)
                    break
                except UnicodeDecodeError:
                    continue
            if data is None:
                messagebox.showerror(APP_TITLE, "无法识别文本编码。")
                return
            rules = parse_replace_rules(data)
        n_before = len(self.replace_rules)
        self.replace_rules = merge_rules(self.replace_rules, rules)
        self.refresh_rules_tree()
        self.log(f"从 {p.name} 导入规则 {len(self.replace_rules) - n_before} 条，合计 {len(self.replace_rules)}")

    def save_rules_to_disk(self) -> None:
        default = rules_path_for(self.excel_path)
        path = filedialog.asksaveasfilename(
            title="保存替换规则",
            defaultextension=".json",
            initialfile=default.name,
            initialdir=str(default.parent),
            filetypes=[("JSON", "*.json"), ("文本", "*.txt")],
        )
        if not path:
            return
        dest = Path(path)
        if dest.suffix.lower() == ".txt":
            lines = [f"{r.find} → {r.replace}" for r in self.replace_rules if r.find]
            dest.write_text("\n".join(lines) + "\n", encoding="utf-8")
        else:
            save_saved_rules(dest, self.replace_rules)
        self.log(f"规则已保存：{dest}")

    def cite_all_rules(self) -> None:
        """一键引用：当前列表 + 磁盘规则 + Excel 规则表 + 修改记录里的固化形式。"""
        groups: List[List[ReplaceRule]] = [list(self.replace_rules)]
        if self.excel_path:
            groups.append(load_saved_rules(rules_path_for(self.excel_path)))
            try:
                wb = load_workbook(
                    self.excel_path, read_only=True, data_only=True, keep_links=False
                )
                try:
                    if RULES_SHEET in wb.sheetnames:
                        groups.append(rules_from_sheet(wb[RULES_SHEET]))
                    if LOG_SHEET in wb.sheetnames:
                        hist: List[ReplaceRule] = []
                        seen = set()
                        for form in forms_from_log_sheet(wb[LOG_SHEET]):
                            if form.term and form.wrapped and form.term not in seen:
                                seen.add(form.term)
                                hist.append(
                                    ReplaceRule(
                                        find=form.term, replace=form.wrapped, enabled=True
                                    )
                                )
                        groups.append(hist)
                finally:
                    wb.close()
            except Exception as e:
                self.log(f"引用 Excel 规则失败：{e}")
        groups.append(load_saved_rules(rules_path_for(None)))
        from_frozen: List[ReplaceRule] = []
        seen_f = set()
        for form in self.frozen_forms:
            if form.term and form.wrapped and form.term not in seen_f:
                seen_f.add(form.term)
                from_frozen.append(
                    ReplaceRule(find=form.term, replace=form.wrapped, enabled=True)
                )
        groups.append(from_frozen)
        merged = merge_rules(*groups)
        self.replace_rules = merged
        self.refresh_rules_tree()
        self.log(f"已引用全部规则 {len(merged)} 条")
        self.show_all_rules()

    def show_all_rules(self) -> None:
        win = tk.Toplevel(self.root)
        win.title("全部替换规则")
        win.geometry("720x460")
        ttk.Label(
            win,
            text=f"共 {len(self.replace_rules)} 条。双击可开关「启用」。套用时只执行已启用规则。",
        ).pack(anchor="w", padx=8, pady=6)
        frm = ttk.Frame(win)
        frm.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 8))
        cols = ("n", "on", "find", "repl")
        tree = ttk.Treeview(frm, columns=cols, show="headings")
        tree.heading("n", text="#")
        tree.heading("on", text="启用")
        tree.heading("find", text="查找")
        tree.heading("repl", text="替换为")
        tree.column("n", width=40, anchor="center")
        tree.column("on", width=50, anchor="center")
        tree.column("find", width=240, anchor="w")
        tree.column("repl", width=320, anchor="w")
        y = ttk.Scrollbar(frm, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=y.set)
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        y.pack(side=tk.RIGHT, fill=tk.Y)
        for i, rule in enumerate(self.replace_rules, start=1):
            tree.insert(
                "",
                tk.END,
                values=(i, "是" if rule.enabled else "否", rule.find, f"→  {rule.replace}"),
            )
        if not self.replace_rules:
            tree.insert("", tk.END, values=("", "", "（暂无规则）", ""))

    def _iter_cell_changes(
        self, ws: Worksheet, *, write: bool
    ) -> Tuple[int, int, List[str], List[ChangeRecord]]:
        """扫描所选列。write=True 时写回单元格。"""
        terms = self.current_terms()
        prefix = self.var_prefix.get()
        suffix = self.var_suffix.get()
        skip_existing = self.var_skip_wrapped.get()
        frozen = self.frozen_wrapped_list()
        start_row = 2 if self.var_skip_header.get() else 1
        cols = self.selected_columns()
        max_row = ws.max_row or 1
        ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        sheet = ws.title

        cell_hits = 0
        total_repl = 0
        samples: List[str] = []
        records: List[ChangeRecord] = []

        for row in range(start_row, max_row + 1):
            for col_name, col in cols:
                cell = ws.cell(row, col)
                if cell.data_type == "f":
                    continue
                original = _cell_text(cell.value)
                if not original:
                    continue
                new, n, applied = apply_wraps(
                    original,
                    terms,
                    prefix,
                    suffix,
                    skip_existing=skip_existing,
                    frozen_wrapped=frozen,
                )
                if new == original:
                    continue
                if write:
                    cell.value = new
                cell_hits += 1
                total_repl += n
                addr = f"{get_column_letter(col)}{row}"
                if len(samples) < 12:
                    samples.append(f"{addr}: {original[:40]}  →  {new[:40]}")
                for term in applied:
                    records.append(
                        ChangeRecord(
                            ts=ts,
                            sheet=sheet,
                            cell=addr,
                            column=col_name,
                            term=term,
                            wrapped=wrap_term(term, prefix, suffix),
                            before=original[:200],
                            after=new[:200],
                        )
                    )
        return cell_hits, total_repl, samples, records

    def _run_replace(self, ws: Worksheet) -> Tuple[int, int, List[str], List[ChangeRecord]]:
        return self._iter_cell_changes(ws, write=True)

    def _iter_rule_changes(
        self, ws: Worksheet, *, write: bool
    ) -> Tuple[int, int, List[str], List[ChangeRecord]]:
        rules = self.enabled_rules()
        skip_existing = self.var_skip_wrapped.get()
        frozen = self.frozen_wrapped_list()
        start_row = 2 if self.var_skip_header.get() else 1
        cols = self.selected_columns()
        max_row = ws.max_row or 1
        ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        sheet = ws.title
        find_to_repl = {r.find: r.replace for r in rules}

        cell_hits = 0
        total_repl = 0
        samples: List[str] = []
        records: List[ChangeRecord] = []

        for row in range(start_row, max_row + 1):
            for col_name, col in cols:
                cell = ws.cell(row, col)
                if cell.data_type == "f":
                    continue
                original = _cell_text(cell.value)
                if not original:
                    continue
                new, n, applied = apply_find_replace(
                    original,
                    rules,
                    skip_existing=skip_existing,
                    frozen_wrapped=frozen,
                )
                if new == original:
                    continue
                if write:
                    cell.value = new
                cell_hits += 1
                total_repl += n
                addr = f"{get_column_letter(col)}{row}"
                if len(samples) < 12:
                    samples.append(f"{addr}: {original[:40]}  →  {new[:40]}")
                for find in applied:
                    records.append(
                        ChangeRecord(
                            ts=ts,
                            sheet=sheet,
                            cell=addr,
                            column=col_name,
                            term=find,
                            wrapped=find_to_repl.get(find, ""),
                            before=original[:200],
                            after=new[:200],
                        )
                    )
        return cell_hits, total_repl, samples, records

    def _scan_preview(self) -> Tuple[int, int, List[str], int, int]:
        path = Path(self.excel_path)
        wb = load_workbook(path, data_only=False)
        try:
            ws = wb[self.var_sheet.get()]
            cells, repl, samples, records = self._iter_cell_changes(ws, write=False)
            skipped = 0
            if self.var_skip_wrapped.get() and self.frozen_forms:
                # 粗略：含固化片段的单元格数（日志提示用）
                frozen = self.frozen_wrapped_list()
                start_row = 2 if self.var_skip_header.get() else 1
                cols = [idx for _n, idx in self.selected_columns()]
                for row in range(start_row, (ws.max_row or 1) + 1):
                    for col in cols:
                        text = _cell_text(ws.cell(row, col).value)
                        if any(f and f in text for f in frozen):
                            skipped += 1
            return cells, repl, samples, len(self.current_terms()), skipped
        finally:
            wb.close()

    def do_preview(self) -> None:
        err = self._validate()
        if err:
            messagebox.showwarning(APP_TITLE, err)
            return
        self.refresh_preview()
        try:
            cells, repl, samples, n_terms, frozen_cells = self._scan_preview()
        except Exception as e:
            messagebox.showerror(APP_TITLE, f"预览失败：\n{e}")
            return
        cols = "、".join(n for n, _ in self.selected_columns())
        self.log(
            f"预览：词 {n_terms} 个 · 列「{cols}」· "
            f"将改 {cells} 个单元格 · 替换 {repl} 处"
        )
        if frozen_cells:
            self.log(f"  其中 {frozen_cells} 个单元格含已固化内容，那些片段会跳过")
        for line in samples:
            self.log("  " + line)
        if not samples:
            self.log("  （所选列里没有新的可替换内容，或都已固化）")
        messagebox.showinfo(
            APP_TITLE,
            f"将改 {cells} 个单元格，共替换 {repl} 处。\n"
            f"已固化形式 {len(self.frozen_forms)} 种会被跳过。\n详见下方日志。",
        )

    def do_preview_rules(self) -> None:
        err = self._validate_excel_cols()
        if err:
            messagebox.showwarning(APP_TITLE, err)
            return
        if not self.enabled_rules():
            messagebox.showwarning(APP_TITLE, "请先添加至少一条启用的替换规则。")
            return
        try:
            path = Path(self.excel_path)
            wb = load_workbook(path, data_only=False)
            try:
                ws = wb[self.var_sheet.get()]
                cells, repl, samples, records = self._iter_rule_changes(ws, write=False)
            finally:
                wb.close()
        except Exception as e:
            messagebox.showerror(APP_TITLE, f"预览失败：\n{e}")
            return
        self.log(
            f"规则预览：{len(self.enabled_rules())} 条 · "
            f"将改 {cells} 个单元格 · 替换 {repl} 处"
        )
        for line in samples:
            self.log("  " + line)
        if not samples:
            self.log("  （所选列没有命中这些规则，或内容已固化）")
        messagebox.showinfo(APP_TITLE, f"将改 {cells} 个单元格，共替换 {repl} 处。")

    def do_apply_all_rules(self) -> None:
        if not self.excel_path:
            messagebox.showwarning(APP_TITLE, "请先选择 Excel。")
            return
        if not messagebox.askyesno(
            APP_TITLE,
            f"将按 {len(self.enabled_rules())} 条启用规则，在勾选列中依次替换。\n"
            "会先备份原文件。确定？",
        ):
            return
        self._execute(Path(self.excel_path), backup_src=True, mode="rules")

    def do_save_as_rules(self) -> None:
        if not self.excel_path:
            messagebox.showwarning(APP_TITLE, "请先选择 Excel。")
            return
        src = Path(self.excel_path)
        path = filedialog.asksaveasfilename(
            title="套用规则并另存为",
            defaultextension=".xlsx",
            initialfile=f"{src.stem}_已替换{src.suffix}",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not path:
            return
        self._execute(Path(path), backup_src=False, mode="rules")

    def _execute(self, dest: Path, *, backup_src: bool, mode: str = "wrap") -> None:
        if mode == "rules":
            err = self._validate_excel_cols()
            if not err and not self.enabled_rules():
                err = "请先添加至少一条启用的替换规则。"
        else:
            err = self._validate()
        if err:
            messagebox.showwarning(APP_TITLE, err)
            return
        src = Path(self.excel_path)
        try:
            if backup_src and dest.resolve() == src.resolve():
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                bak = src.with_name(f"{src.stem}.bak_{stamp}{src.suffix}")
                shutil.copy2(src, bak)
                self.log(f"已备份：{bak.name}")
            wb = load_workbook(src)
            ws = wb[self.var_sheet.get()]
            if mode == "rules":
                cells, repl, samples, records = self._iter_rule_changes(ws, write=True)
                write_rules_sheet(ensure_rules_sheet(wb), self.replace_rules)
                save_saved_rules(rules_path_for(dest), self.replace_rules)
            else:
                cells, repl, samples, records = self._run_replace(ws)
            if records:
                log_ws = ensure_log_sheet(wb)
                append_log_sheet(log_ws, records)
                now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                prefix = self.var_prefix.get()
                suffix = self.var_suffix.get()
                new_forms = [
                    FrozenForm(
                        term=rec.term,
                        prefix=prefix if mode == "wrap" else "",
                        suffix=suffix if mode == "wrap" else "",
                        wrapped=rec.wrapped,
                        first_seen=rec.ts or now,
                    )
                    for rec in records
                ]
                self.frozen_forms = merge_forms(self.frozen_forms, new_forms)
                self.change_history.extend(records)
            tmp = dest.with_name(f".~{dest.stem}.tmp{dest.suffix}")
            wb.save(tmp)
            os.replace(tmp, dest)
            wb.close()
            save_lock(dest, self.frozen_forms, self.change_history)
        except PermissionError:
            messagebox.showerror(
                APP_TITLE,
                "无法保存：请先关闭 Excel 里已打开的这个文件，再试一次。",
            )
            return
        except Exception as e:
            messagebox.showerror(APP_TITLE, f"替换失败：\n{e}")
            return

        self._refresh_frozen_label()
        self.log(f"完成：改 {cells} 个单元格 · 替换 {repl} 处")
        self.log(f"已写入：{dest}")
        self.log(f"修改记录已追加到工作表「{LOG_SHEET}」，形式已固化 {len(self.frozen_forms)} 种")
        for line in samples:
            self.log("  " + line)
        if dest.resolve() != src.resolve():
            self.excel_path = dest
            self.var_excel.set(str(dest))
        messagebox.showinfo(
            APP_TITLE,
            f"完成。改了 {cells} 个单元格，替换 {repl} 处。\n"
            f"已写入「{LOG_SHEET}」，下次遇到这些内容会跳过。",
        )

    def do_save_as(self) -> None:
        if not self.excel_path:
            messagebox.showwarning(APP_TITLE, "请先选择 Excel。")
            return
        src = Path(self.excel_path)
        path = filedialog.asksaveasfilename(
            title="另存为",
            defaultextension=".xlsx",
            initialfile=f"{src.stem}_已包裹{src.suffix}",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not path:
            return
        self._execute(Path(path), backup_src=False)

    def do_overwrite(self) -> None:
        if not messagebox.askyesno(
            APP_TITLE,
            "将直接改原文件（会先做一份 .bak 备份）。确定？",
        ):
            return
        self._execute(Path(self.excel_path), backup_src=True)

    def show_change_log(self) -> None:
        rows = list(self.change_history)
        if self.excel_path:
            try:
                wb = load_workbook(
                    self.excel_path, read_only=True, data_only=True, keep_links=False
                )
                try:
                    if LOG_SHEET in wb.sheetnames:
                        ws = wb[LOG_SHEET]
                        it = ws.iter_rows(values_only=True)
                        header = next(it, None)
                        if header:
                            cols = {
                                _cell_text(h).strip(): i
                                for i, h in enumerate(header)
                                if _cell_text(h).strip()
                            }

                            def col(row, name: str) -> str:
                                i = cols.get(name)
                                if i is None or i >= len(row):
                                    return ""
                                return _cell_text(row[i])

                            sheet_rows: List[ChangeRecord] = []
                            for row in it:
                                if not col(row, "固化形式") and not col(row, "原词"):
                                    continue
                                sheet_rows.append(
                                    ChangeRecord(
                                        ts=col(row, "时间"),
                                        sheet=col(row, "工作表"),
                                        cell=col(row, "单元格"),
                                        column=col(row, "列名"),
                                        term=col(row, "原词"),
                                        wrapped=col(row, "固化形式"),
                                        before=col(row, "替换前"),
                                        after=col(row, "替换后"),
                                    )
                                )
                            if sheet_rows:
                                rows = sheet_rows
                finally:
                    wb.close()
            except Exception as e:
                self.log(f"读取修改记录失败：{e}")

        win = tk.Toplevel(self.root)
        win.title("修改记录 / 已固化形式")
        win.geometry("960x480")
        info = ttk.Label(
            win,
            text=(
                f"已固化 {len(self.frozen_forms)} 种形式。"
                "再次处理时，单元格里已有这些写法的片段会原样保留。"
            ),
        )
        info.pack(anchor="w", padx=8, pady=6)

        frm = ttk.Frame(win)
        frm.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 8))
        cols = ("ts", "sheet", "cell", "column", "term", "wrapped")
        tree = ttk.Treeview(frm, columns=cols, show="headings", height=16)
        titles = {
            "ts": "时间",
            "sheet": "工作表",
            "cell": "单元格",
            "column": "列名",
            "term": "原词",
            "wrapped": "固化形式",
        }
        widths = {"ts": 150, "sheet": 90, "cell": 70, "column": 90, "term": 120, "wrapped": 280}
        for c in cols:
            tree.heading(c, text=titles[c])
            tree.column(c, width=widths[c], anchor="w")
        y = ttk.Scrollbar(frm, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=y.set)
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        y.pack(side=tk.RIGHT, fill=tk.Y)
        for rec in rows[-2000:]:
            tree.insert(
                "",
                tk.END,
                values=(rec.ts, rec.sheet, rec.cell, rec.column, rec.term, rec.wrapped),
            )
        if not rows:
            tree.insert("", tk.END, values=("", "", "", "", "（暂无记录）", ""))


def main() -> int:
    root = tk.Tk()
    try:
        style = ttk.Style()
        if sys.platform == "darwin" and "aqua" in style.theme_names():
            style.theme_use("aqua")
    except Exception:
        pass
    ExcelWrapApp(root)
    root.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
