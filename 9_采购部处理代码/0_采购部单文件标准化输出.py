import os
import re
import importlib.util
from datetime import datetime, date as date_type
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


# ============= 路径配置 =============
RAW_DATA_ROOT = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/9_采购部采购台账/A_采购部采购台账_raw"
OUTPUT_ROOT = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/9_采购部采购台账/B_采购部采购台账_std"
PROJECT_INFO_FILE = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/0_项目基本信息情况/项目基本信息表.xlsx"
PROJECT_MATCH_DB_FILE = "/Users/wyattsun/Desktop/code/0_项目基本信息/project_name_match_db.json"
PROJECT_MATCHER_MODULE_PATH = "/Users/wyattsun/Desktop/code/0_项目基本信息/project_name_matcher.py"


# ============= 标准化流水线配置 =============
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
STANDARDIZATION_PIPELINE_DIR = os.path.join(CURRENT_DIR, "物品名称规格标准化")
PIPELINE_CONTEXT = None


# ============= 列配置 =============
STANDARD_COLUMNS = [
    "材料归类", "物品名称", "规格", "标准单位", "数量",
    "单价", "不含税金额", "7%增值税", "含税金额", "货币种类", "汇率",
    "签订日期", "供货商名称", "备注",
]

CORE_REQUIRED_COLUMNS = ["物品名称", "单位", "数量", "单价", "不含税金额", "含税金额"]

FINAL_OUTPUT_COLUMNS = [
    "项目编码", "项目名称", "材料归类", "物品名称", "规格", "标准单位", "数量",
    "单价", "不含税金额", "7%增值税", "含税金额", "货币种类", "汇率", "年月",
    "标准日期", "供货商名称", "备注", "签订日期",
]

CENTER_ALIGN_COLUMNS = {"标准单位", "数量", "货币种类", "汇率"}
LEFT_ALIGN_COLUMNS = {"签订日期"}
AMOUNT_COLUMNS = {"单价", "不含税金额", "7%增值税", "含税金额"}


COLUMN_MAPPINGS = {
    "物品名称": ["物品名称", "采购物资名称", "物资名称"],
}

UNIT_MAPPING = {
    "m2": "㎡", "M2": "㎡", "平方米": "㎡", "平米": "㎡", "平方": "㎡", "平": "㎡",
    "KG": "kg", "公斤": "kg", "KG-袋": "kg",
    "t": "吨", "dun": "吨",
    "m": "米", "M": "米",
    "立方": "m³",
    "PC": "个", "PCS": "个", "(个)": "个", "个)": "个",
    "(箱)": "箱", "箱)": "箱", "(根)": "根", "根)": "根", "片)": "片",
    "j卷": "卷", "快": "块",
    "桶/7KG": "桶", "桶/20KG": "桶", "捆/10KG": "捆",
}


def _load_module_from_path(module_name, file_path):
    spec = importlib.util.spec_from_file_location(module_name, file_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载模块: {file_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def load_project_matcher_funcs():
    matcher = _load_module_from_path("project_name_matcher", PROJECT_MATCHER_MODULE_PATH)
    return {
        "load_or_init_db": matcher.load_or_init_db,
        "save_db": matcher.save_db,
        "match_project": matcher.match_project,
        "mark_unmatched": matcher.mark_unmatched,
        "add_mapping": matcher.add_mapping,
    }


def load_project_config():
    try:
        df = pd.read_excel(PROJECT_INFO_FILE, sheet_name=0)
    except Exception as exc:
        print(f"警告：无法加载项目配置 {PROJECT_INFO_FILE}: {exc}")
        return []

    required_cols = ["项目编码", "项目名称"]
    if any(col not in df.columns for col in required_cols):
        print("警告：项目配置缺少必要列，后续仅按文件名输出项目名称")
        return []

    projects = []
    for _, row in df.iterrows():
        code = str(row["项目编码"]).strip() if pd.notna(row["项目编码"]) else ""
        name = str(row["项目名称"]).strip() if pd.notna(row["项目名称"]) else ""
        if code or name:
            projects.append({"project_code": code, "project_name": name})
    return projects


def extract_project_name(filename):
    base_name = os.path.splitext(filename)[0]
    if "_" in base_name:
        return base_name.split("_", 1)[0].strip()
    match = re.match(r"^(.+?项目)", base_name)
    if match:
        return match.group(1).strip()
    return base_name.strip()


def find_matching_project(extracted_name, project_config, match_project, match_db):
    hit = match_project(extracted_name, match_db)
    if hit:
        return {
            "project_code": hit.get("project_code", ""),
            "project_name": hit.get("project_name", extracted_name),
        }

    for project in project_config:
        if project["project_name"] == extracted_name:
            return project

    for project in project_config:
        if extracted_name in project["project_name"] or project["project_name"] in extracted_name:
            return project

    return None


def list_excel_files(root_dir):
    excel_paths = []
    for root, _, files in os.walk(root_dir):
        for name in sorted(files):
            if name.startswith("~$"):
                continue
            ext = os.path.splitext(name)[1].lower()
            if ext in {".xlsx", ".xls", ".xlsm"}:
                excel_paths.append(os.path.join(root, name))
    return sorted(excel_paths)


def normalize_headers(headers):
    normalized = []
    for col in headers:
        stripped_col = str(col).strip() if col else ""
        for std_col, variations in COLUMN_MAPPINGS.items():
            if stripped_col in variations:
                normalized.append(std_col)
                break
        else:
            normalized.append(stripped_col)
    return normalized


def find_header_row(sheet):
    for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row and row[0] and str(row[0]).strip() in ["序号", "No.", "编号"]:
            row_values = [str(val).strip() for val in row if val is not None]
            matched = [
                col for col in STANDARD_COLUMNS
                if any(alias in row_values for alias in COLUMN_MAPPINGS.get(col, [col]))
            ]
            if len(matched) >= len(STANDARD_COLUMNS) // 2:
                return row_idx
    return None


def is_summary_row(formula_sheet, row_idx, amount_col):
    amount_cell = formula_sheet.cell(row_idx, amount_col)
    if amount_cell.data_type != "f":
        return False

    formula_str = str(amount_cell.value).upper()
    if "SUM" not in formula_str:
        return False

    col_letter = get_column_letter(amount_col)
    if f"{col_letter}:{col_letter}" in formula_str:
        return True

    range_pattern = f"\\$?{col_letter}\\$?\\d+:\\$?{col_letter}\\$?\\d+"
    return re.search(range_pattern, formula_str) is not None


def standardize_unit(unit_str):
    if not unit_str or str(unit_str).strip() == "":
        return ""
    unit_str = str(unit_str).strip()

    try:
        float(unit_str)
        return ""
    except ValueError:
        pass

    if unit_str in ["上", "想", "和", "干", "寸"]:
        return ""

    return UNIT_MAPPING.get(unit_str, unit_str)


def parse_date(date_str):
    if date_str is None:
        return (None, None, None)

    if isinstance(date_str, datetime):
        return (date_str.year, date_str.month, date_str.day)
    if isinstance(date_str, date_type):
        return (date_str.year, date_str.month, date_str.day)

    s = str(date_str).strip()
    if not s:
        return (None, None, None)

    match = re.match(r"^(\d{4})/(\d{1,2})/(\d{1,2})(?:\s+\d{1,2}:\d{2}:\d{2})?", s)
    if match:
        y, m, d = int(match.group(1)), int(match.group(2)), int(match.group(3))
        if 1 <= m <= 12 and 1 <= d <= 31:
            return (y, m, d)

    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[\s T]\d{1,2}:\d{2}:\d{2})?", s)
    if match:
        y, m, d = int(match.group(1)), int(match.group(2)), int(match.group(3))
        if 1 <= m <= 12 and 1 <= d <= 31:
            return (y, m, d)

    match = re.match(r"^(\d{4})\.(\d{1,2})(?:\.(\d{1,2}))?", s)
    if match:
        y, m = int(match.group(1)), int(match.group(2))
        d = int(match.group(3)) if match.group(3) else None
        if 1 <= m <= 12 and (d is None or 1 <= d <= 31):
            return (y, m, d)

    return (None, None, None)


def infer_date_for_project(project_data):
    result = list(project_data)

    for i in range(1, len(result)):
        row_idx, year, month, day = result[i]
        if year is None:
            for j in range(i - 1, -1, -1):
                _, py, pm, pd_ = result[j]
                if py is not None:
                    result[i] = (row_idx, py, pm, pd_)
                    break

    for i in range(len(result) - 2, -1, -1):
        row_idx, year, month, day = result[i]
        if year is None:
            for j in range(i + 1, len(result)):
                _, ny, nm, nd = result[j]
                if ny is not None:
                    result[i] = (row_idx, ny, nm, nd)
                    break

    return result


def format_date_str(year, month, day):
    if year is None:
        return ""
    if day is None:
        return f"{year}-{month:02d}"
    return f"{year}-{month:02d}-{day:02d}"


def get_pipeline_context():
    """缓存流水线模块与规则模块，避免每个文件重复import。"""
    global PIPELINE_CONTEXT
    if PIPELINE_CONTEXT is not None:
        return PIPELINE_CONTEXT

    pipeline_dir = Path(STANDARDIZATION_PIPELINE_DIR)
    pipeline_script = pipeline_dir / "0_统一标准化流水线.py"
    if not pipeline_script.exists():
        PIPELINE_CONTEXT = {"ok": False, "reason": f"未找到标准化流水线脚本: {pipeline_script}"}
        return PIPELINE_CONTEXT

    spec = importlib.util.spec_from_file_location("_standardization_pipeline", str(pipeline_script))
    if spec is None or spec.loader is None:
        PIPELINE_CONTEXT = {"ok": False, "reason": f"无法加载标准化流水线: {pipeline_script}"}
        return PIPELINE_CONTEXT

    pipeline_module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(pipeline_module)

    rule_names = getattr(pipeline_module, "DEFAULT_RULES", [])
    if not rule_names:
        PIPELINE_CONTEXT = {"ok": False, "reason": "标准化流水线 DEFAULT_RULES 为空"}
        return PIPELINE_CONTEXT

    rule_modules = []
    for rule_name in rule_names:
        rule_file = pipeline_dir / rule_name
        if not rule_file.exists():
            continue
        try:
            rule_module = pipeline_module.load_module_from_file(rule_file)
            rule_modules.append((rule_file.name, rule_module))
        except Exception:
            continue

    PIPELINE_CONTEXT = {
        "ok": True,
        "pipeline_module": pipeline_module,
        "rule_modules": rule_modules,
    }
    return PIPELINE_CONTEXT


def apply_standardization_pipeline(excel_path):
    context = get_pipeline_context()
    if not context.get("ok"):
        print(f"  ⚠️ {context.get('reason', '标准化流水线不可用，已跳过')}")
        return

    workbook = openpyxl.load_workbook(excel_path)
    pipeline_module = context["pipeline_module"]
    for rule_name, rule_module in context["rule_modules"]:
        try:
            stats = pipeline_module.run_rule_on_workbook(rule_module, workbook, rule_name)
            print(
                f"  ✓ {rule_name}: "
                f"命中 {stats['matched']} 行, 更新 {stats['updated']} 行, 未识别 {stats['unresolved']} 行"
            )
        except Exception as exc:
            print(f"  ✗ 规则执行失败 {rule_name}: {exc}")

    workbook.save(excel_path)


def extract_rows_from_single_excel(file_path, project_code, project_name):
    formula_wb = openpyxl.load_workbook(file_path, data_only=False)
    formula_sheet = formula_wb.active
    if formula_sheet is None:
        formula_wb.close()
        print("  ⚠️ 未找到活动工作表，跳过")
        return []

    header_row_idx = find_header_row(formula_sheet)
    if header_row_idx is None:
        formula_wb.close()
        print("  ⚠️ 未识别到表头，跳过")
        return []

    headers = normalize_headers([cell.value for cell in formula_sheet[header_row_idx]])
    header_idx_map = {name: idx for idx, name in enumerate(headers)}

    missing_core = [col for col in CORE_REQUIRED_COLUMNS if col not in headers]
    if missing_core:
        formula_wb.close()
        print(f"  ⚠️ 缺少核心列，跳过: {', '.join(missing_core)}")
        return []

    try:
        amount_col = header_idx_map["含税金额"] + 1
    except ValueError:
        formula_wb.close()
        print("  ⚠️ 缺少含税金额列，跳过")
        return []
    except KeyError:
        formula_wb.close()
        print("  ⚠️ 缺少含税金额列，跳过")
        return []

    summary_rows = set()
    for row_idx in range(header_row_idx + 1, formula_sheet.max_row + 1):
        if is_summary_row(formula_sheet, row_idx, amount_col):
            summary_rows.add(row_idx)

    formula_wb.close()

    value_wb = openpyxl.load_workbook(file_path, data_only=True)
    value_sheet = value_wb.active
    if value_sheet is None:
        value_wb.close()
        print("  ⚠️ 未找到活动工作表，跳过")
        return []

    key_data_cols = ["物品名称", "规格", "单位", "数量", "单价", "不含税金额"]
    key_data_indices = [header_idx_map[col] for col in key_data_cols if col in header_idx_map]

    valid_rows = []
    row_values_map = {}
    for row_idx, row_vals in enumerate(
        value_sheet.iter_rows(min_row=header_row_idx + 1, max_row=value_sheet.max_row, values_only=True),
        start=header_row_idx + 1,
    ):
        if row_idx in summary_rows:
            continue

        if all(v is None or str(v).strip() == "" for v in row_vals):
            continue

        key_data_empty = True
        for idx_ in key_data_indices:
            v = row_vals[idx_] if idx_ < len(row_vals) else None
            if v is not None and str(v).strip() != "":
                key_data_empty = False
                break

        if key_data_empty:
            continue

        valid_rows.append(row_idx)
        row_values_map[row_idx] = row_vals

    rows = []
    for row_idx in valid_rows:
        row_vals = row_values_map[row_idx]
        row_obj = {"项目编码": project_code, "项目名称": project_name}

        for col_name in STANDARD_COLUMNS:
            if col_name == "标准单位":
                unit_idx = header_idx_map.get("单位")
                if unit_idx is not None and unit_idx < len(row_vals):
                    unit_val = row_vals[unit_idx]
                    row_obj[col_name] = standardize_unit(unit_val)
                else:
                    row_obj[col_name] = ""
            elif col_name in header_idx_map:
                source_idx = header_idx_map[col_name]
                row_obj[col_name] = row_vals[source_idx] if source_idx < len(row_vals) else ""
            else:
                row_obj[col_name] = ""

        rows.append(row_obj)

    value_wb.close()
    return rows


def enrich_dates(df):
    if df.empty:
        df["标准日期"] = ""
        df["年月"] = ""
        return df

    parsed = [parse_date(v) for v in df["签订日期"].tolist()]
    project_data = [(idx, y, m, d) for idx, (y, m, d) in enumerate(parsed)]
    inferred = infer_date_for_project(project_data)

    std_dates = [""] * len(df)
    yms = [""] * len(df)
    for idx, y, m, d in inferred:
        std_dates[idx] = format_date_str(y, m, d)
        yms[idx] = f"{y}-{m:02d}" if y and m else ""

    df["标准日期"] = std_dates
    df["年月"] = yms
    return df


def write_output(df, output_path):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    out_df = df.copy()
    for col in FINAL_OUTPUT_COLUMNS:
        if col not in out_df.columns:
            out_df[col] = ""
    out_df = out_df[FINAL_OUTPUT_COLUMNS]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        out_df.to_excel(writer, sheet_name="合并数据", index=False)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["合并数据"]

    ws.row_dimensions[1].height = 40
    ws.freeze_panes = "A2"

    for col_idx, header in enumerate(FINAL_OUTPUT_COLUMNS, start=1):
        cell = ws.cell(1, col_idx)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

        if header in AMOUNT_COLUMNS:
            for row_idx in range(2, ws.max_row + 1):
                c = ws.cell(row_idx, col_idx)
                if isinstance(c.value, (int, float)):
                    c.number_format = "0.00"

        if header in CENTER_ALIGN_COLUMNS:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row_idx, col_idx).alignment = Alignment(horizontal="center", vertical="center")
        elif header in LEFT_ALIGN_COLUMNS:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row_idx, col_idx).alignment = Alignment(horizontal="left", vertical="center")

    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(output_path)


def build_output_path(input_path):
    rel_path = os.path.relpath(input_path, RAW_DATA_ROOT)
    rel_dir = os.path.dirname(rel_path)
    base_name, _ = os.path.splitext(os.path.basename(rel_path))
    output_name = f"{base_name}.xlsx"
    return os.path.join(OUTPUT_ROOT, rel_dir, output_name)


def process_single_file(file_path, project_config, matcher_funcs, match_db):
    file_name = os.path.basename(file_path)
    print(f"\n处理文件: {file_path}")

    extracted_name = extract_project_name(file_name)
    matched = find_matching_project(extracted_name, project_config, matcher_funcs["match_project"], match_db)

    if matched:
        project_code = matched.get("project_code", "")
        project_name = matched.get("project_name", extracted_name)
        matcher_funcs["add_mapping"](match_db, extracted_name, project_name)
    else:
        project_code = ""
        project_name = extracted_name
        matcher_funcs["mark_unmatched"](match_db, extracted_name)

    try:
        rows = extract_rows_from_single_excel(file_path, project_code, project_name)
    except Exception as exc:
        print(f"  ✗ 处理失败: {exc}")
        return False

    if not rows:
        print("  ⚠️ 无有效数据，跳过输出")
        return False

    df = pd.DataFrame(rows)

    # 仅保留金额合计不为0的行。
    amount_df = df[[c for c in ["单价", "不含税金额", "7%增值税", "含税金额"] if c in df.columns]].apply(
        pd.to_numeric, errors="coerce"
    )
    row_sum = amount_df.fillna(0).sum(axis=1)
    df = df[row_sum.ne(0)]
    if df.empty:
        print("  ⚠️ 金额合计均为0，跳过输出")
        return False

    df = enrich_dates(df)

    output_path = build_output_path(file_path)
    write_output(df, output_path)
    apply_standardization_pipeline(output_path)

    print(f"  ✓ 已输出: {output_path}")
    return True


def main():
    if not os.path.isdir(RAW_DATA_ROOT):
        raise FileNotFoundError(f"输入目录不存在: {RAW_DATA_ROOT}")

    os.makedirs(OUTPUT_ROOT, exist_ok=True)

    project_config = load_project_config()
    matcher_funcs = load_project_matcher_funcs()
    match_db = matcher_funcs["load_or_init_db"](PROJECT_MATCH_DB_FILE, project_config)

    excel_files = list_excel_files(RAW_DATA_ROOT)
    if not excel_files:
        print("未发现Excel文件，程序结束")
        return

    success_count = 0
    skip_count = 0

    print(f"扫描到 {len(excel_files)} 个Excel文件，开始逐个标准化输出...")
    for file_path in excel_files:
        ok = process_single_file(file_path, project_config, matcher_funcs, match_db)
        if ok:
            success_count += 1
        else:
            skip_count += 1

    matcher_funcs["save_db"](PROJECT_MATCH_DB_FILE, match_db)

    print("\n处理完成")
    print(f"成功输出: {success_count} 个")
    print(f"跳过/失败: {skip_count} 个")
    print(f"输出根目录: {OUTPUT_ROOT}")


if __name__ == "__main__":
    main()
