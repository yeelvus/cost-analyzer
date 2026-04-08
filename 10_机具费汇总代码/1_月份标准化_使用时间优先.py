from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX as OPENPYXL_COLOR_INDEX
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel


DEFAULT_INPUT_DIR = Path(
    "/Users/wyattsun/Desktop/邯郸中材/01成本统计/10_机具使用费/A_原始数据"
)
DEFAULT_OUTPUT_DIR = Path(
    "/Users/wyattsun/Desktop/邯郸中材/01成本统计/10_机具使用费/B_标准化文件"
)
DEFAULT_OUTPUT_SUFFIX = "_标准化"
DEFAULT_YEAR_WITHOUT_EXPLICIT = 2025
HEADER_SCAN_ROWS = 12
SOFT_SCAN_ROWS = 60
CURRENT_YEAR_MONTH = (date.today().year, date.today().month)
RAW_FILE_MONTH_RE = re.compile(r"(20\d{2})\s*年\s*(1[0-2]|0?[1-9])\s*月")

USAGE_HEADER_KEYWORDS = (
    "使用时间",
    "使用日期",
    "使用期间",
    "起止时间",
    "起止日期",
    "租赁期间",
    "使用起止",
    "日期区间",
    "时间区间",
    "期间",
)

BLOCKED_MONTH_HEADERS = (
    "统计开始月份",
    "开始结算月份",
    "开始月份",
    "结束月份",
    "结束日期",
    "结束月",
    "开始月",
)

MONTH_WORD_TO_NUMBER = {
    "元月": 1,
    "正月": 1,
    "一月": 1,
    "二月": 2,
    "三月": 3,
    "四月": 4,
    "五月": 5,
    "六月": 6,
    "七月": 7,
    "八月": 8,
    "九月": 9,
    "十月": 10,
    "十一月": 11,
    "十二月": 12,
    "腊月": 12,
}

STANDARD_YM_RE = re.compile(r"^\s*(20\d{2})-(0[1-9]|1[0-2])\s*$")
COMPACT_YM_RE = re.compile(r"(?<!\d)(20\d{2})(0[1-9]|1[0-2])(?:月|月份)?(?!\d)")
YEAR_MONTH_RE = re.compile(r"(?<!\d)(20\d{2})\s*[年./-]\s*(1[0-2]|0?[1-9])")
YEAR_MONTH_2DIGIT_RE = re.compile(r"(?<![\d./-])(2[4-6])\s*[年./-]\s*(1[0-2]|0?[1-9])")
MONTH_YEAR_RE = re.compile(r"(?<!\d)(1[0-2]|0?[1-9])\s*[./-]\s*(20\d{2})(?!\d)")
YEAR_4_RE = re.compile(r"(?<!\d)(20\d{2})(?!\d)")
YEAR_2_RE = re.compile(r"(?<![\d./-])(2[4-6])(?=\s*(?:年|[./-]))")
YEAR_2_NIAN_RE = re.compile(r"(?<!\d)(2[4-6])\s*年")  # 允许范围符号'-'前缀，如 -25年
YEAR_2_NIAN_MONTH_RE = re.compile(r"(?<!\d)(2[4-6])\s*年\s*(1[0-2]|0?[1-9])")  # 25年11月
MONTH_SLASH_DAY_RE = re.compile(
    r"(?<![\d/\-~,至、，])(1[0-2]|0?[1-9])\s*/\s*[0-3]?\d(?:\s*[-~,至、，]\s*(?:1[0-2]|0?[1-9])\s*/\s*[0-3]?\d)*"
)
DAY_RANGE_MONTH_END_RE = re.compile(
    r"(?<!\d)(?:\d{1,2}(?:\s*[-~,至]\s*\d{1,2})?)(?:\s*,\s*\d{1,2}(?:\s*[-~,至]\s*\d{1,2})?)*\s*/\s*(1[0-2]|0?[1-9])(?!\d)"
)
MONTH_CN_RE = re.compile(r"(?<!\d)(1[0-2]|0?[1-9])\s*月")
MONTH_DOT_DAY_RE = re.compile(r"(?<!\d)(1[0-2]|0?[1-9])\s*[./]\s*[0-3]?\d")
MONTH_ONLY_RE = re.compile(r"^\s*(1[0-2]|0?[1-9])(?:月|月份)?\s*$")
MONTH_WORD_RE = re.compile(
    "|".join(re.escape(word) for word in sorted(MONTH_WORD_TO_NUMBER, key=len, reverse=True))
)


@dataclass
class HeaderDetection:
    header_row: int
    usage_col: int | None
    month_col: int | None
    usage_source: str
    month_source: str


@dataclass
class RowState:
    row_idx: int
    is_data_row: bool
    month_cell: Any
    usage_value: Any
    month_raw_value: Any
    usage_candidate: tuple[int, int] | None
    month_candidate: tuple[int, int] | None
    yellow_signature: str | None
    starts_yellow_block: bool
    final_ym: tuple[int, int] | None = None
    source: str = ""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="按使用时间优先规则标准化机具台账中的月份列")
    parser.add_argument("input_file", nargs="?", default=None, help="输入 Excel 文件路径（不传则自动选择原始数据目录中最新月份文件）")
    parser.add_argument("-i", "--input", dest="input_override", default=None, help="输入 Excel 文件路径（优先级高于自动选择）")
    parser.add_argument("--suffix", default=DEFAULT_OUTPUT_SUFFIX, help="输出文件名后缀")
    return parser.parse_args()


def extract_year_month_from_name(path: Path) -> tuple[int, int] | None:
    match = RAW_FILE_MONTH_RE.search(path.name)
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def find_latest_raw_input_file(input_dir: Path) -> Path | None:
    if not input_dir.exists():
        return None

    candidates = [
        file_path
        for file_path in input_dir.glob("*.xlsx")
        if file_path.is_file() and not file_path.name.startswith("~$")
    ]
    if not candidates:
        return None

    dated_candidates: list[tuple[tuple[int, int], float, Path]] = []
    undated_candidates: list[tuple[float, Path]] = []
    for file_path in candidates:
        parsed = extract_year_month_from_name(file_path)
        modified_time = file_path.stat().st_mtime
        if parsed is not None:
            dated_candidates.append((parsed, modified_time, file_path))
        else:
            undated_candidates.append((modified_time, file_path))

    if dated_candidates:
        # 优先按文件名中的“年/月”选最新，月份一致时取修改时间更新者。
        dated_candidates.sort(key=lambda item: (item[0], item[1]))
        return dated_candidates[-1][2]

    undated_candidates.sort(key=lambda item: item[0])
    return undated_candidates[-1][1]


def list_all_raw_input_files(input_dir: Path) -> list[Path]:
    if not input_dir.exists():
        return []

    candidates = [
        file_path
        for file_path in input_dir.glob("*.xlsx")
        if file_path.is_file() and not file_path.name.startswith("~$")
    ]

    # 按文件名排序，保证批量执行顺序稳定、可复现。
    return sorted(candidates, key=lambda p: p.name)


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def normalize_header_text(value: object) -> str:
    return re.sub(r"\s+", "", normalize_text(value))


def format_year_month(year_month: tuple[int, int]) -> str:
    return f"{year_month[0]:04d}-{year_month[1]:02d}"


def next_year_month(year_month: tuple[int, int]) -> tuple[int, int]:
    year, month = year_month
    if month == 12:
        return year + 1, 1
    return year, month + 1


def clamp_to_current(year_month: tuple[int, int]) -> tuple[int, int]:
    if year_month > CURRENT_YEAR_MONTH:
        return CURRENT_YEAR_MONTH
    return year_month


def coerce_continuous(
    candidate: tuple[int, int] | None,
    floor: tuple[int, int] | None,
) -> tuple[int, int] | None:
    if candidate is None:
        return floor

    candidate = clamp_to_current(candidate)
    if floor is None:
        return candidate

    if candidate < floor:
        # 年初回看：floor 在 1-2 月时，允许回退到上一年 10-12 月（如 2025-01 前的 2024-10）。
        if floor[1] <= 2 and candidate[0] == floor[0] - 1 and candidate[1] >= 10:
            return candidate
        # 年初局部重排：允许在同一年 1-4 月之间小范围回退（4月后出现1-4月明细）。
        if floor[1] <= 4 and candidate[0] == floor[0] and candidate[1] <= 4:
            return candidate
        return floor
    return candidate


def row_has_any_value(worksheet, row_idx: int) -> bool:
    for col_idx in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row_idx, col_idx).value
        if value is not None and normalize_text(value):
            return True
    return False


def count_nonempty_cells(worksheet, row_idx: int) -> int:
    count = 0
    for col_idx in range(1, worksheet.max_column + 1):
        if normalize_text(worksheet.cell(row_idx, col_idx).value):
            count += 1
    return count


def collect_row_text(worksheet, row_idx: int) -> str:
    values: list[str] = []
    for col_idx in range(1, worksheet.max_column + 1):
        text = normalize_text(worksheet.cell(row_idx, col_idx).value)
        if text:
            values.append(text)
    return " ".join(values)


def score_usage_header(text: str) -> int:
    if not text:
        return 0
    if text in {"使用时间", "使用日期", "使用期间", "起止时间", "起止日期"}:
        return 120
    if text in {"租赁期间", "使用起止", "日期区间", "时间区间"}:
        return 100
    if any(keyword in text for keyword in USAGE_HEADER_KEYWORDS):
        return 80
    return 0


def score_month_header(text: str) -> int:
    if not text:
        return 0
    if any(block in text for block in BLOCKED_MONTH_HEADERS):
        return 0
    if any(block in text for block in ("金额", "租金", "单价", "合计")) and text != "费用月份":
        return 0
    if text == "费用月份":
        return 130
    if text == "月份":
        return 120
    if text == "费用月":
        return 110
    if text == "月":
        return 100
    if text.endswith("月份"):
        return 90
    if text.endswith("月") and len(text) <= 4:
        return 70
    return 0


def detect_header_and_columns(worksheet) -> HeaderDetection:
    best_row: tuple[int, int | None, int | None, int] | None = None
    scan_limit = min(HEADER_SCAN_ROWS, worksheet.max_row)

    for row_idx in range(1, scan_limit + 1):
        usage_col = None
        month_col = None
        usage_score = 0
        month_score = 0

        for col_idx in range(1, worksheet.max_column + 1):
            text = normalize_header_text(worksheet.cell(row_idx, col_idx).value)
            if not text:
                continue

            current_usage_score = score_usage_header(text)
            current_month_score = score_month_header(text)

            if current_usage_score > usage_score:
                usage_score = current_usage_score
                usage_col = col_idx
            if current_month_score > month_score:
                month_score = current_month_score
                month_col = col_idx

        if usage_score or month_score:
            row_score = month_score * 10 + usage_score * 6 + min(count_nonempty_cells(worksheet, row_idx), 10)
            candidate = (row_idx, usage_col, month_col, row_score)
            if best_row is None or candidate[3] > best_row[3] or (candidate[3] == best_row[3] and candidate[0] < best_row[0]):
                best_row = candidate

    if best_row is None:
        header_row = max(range(1, scan_limit + 1), key=lambda idx: count_nonempty_cells(worksheet, idx), default=1)
        usage_col = None
        month_col = None
        usage_source = "none"
        month_source = "none"
    else:
        header_row, usage_col, month_col, _ = best_row
        usage_source = "header" if usage_col is not None else "none"
        month_source = "header" if month_col is not None else "none"

    if usage_col is None:
        usage_col = soft_detect_usage_column(worksheet, header_row, exclude_col=month_col)
        if usage_col is not None:
            usage_source = "soft"

    if month_col is None:
        month_col = soft_detect_month_column(worksheet, header_row, exclude_col=usage_col)
        if month_col is not None:
            month_source = "soft"

    return HeaderDetection(
        header_row=header_row,
        usage_col=usage_col,
        month_col=month_col,
        usage_source=usage_source,
        month_source=month_source,
    )


def coerce_excel_serial_to_date(value: object) -> date | datetime | None:
    if isinstance(value, (int, float)):
        number = float(value)
    elif isinstance(value, str):
        stripped = value.strip()
        if not stripped or not re.fullmatch(r"\d+(?:\.0+)?", stripped):
            return None
        number = float(stripped)
    else:
        return None

    if not (20000 <= number <= 60000):
        return None

    try:
        converted = from_excel(number)
    except Exception:
        return None

    if isinstance(converted, (datetime, date)):
        return converted
    return None


def extract_explicit_year(text: str) -> int | None:
    match = YEAR_4_RE.search(text)
    if match:
        return int(match.group(1))

    match = YEAR_2_RE.search(text)
    if match:
        return 2000 + int(match.group(1))

    # 处理范围分隔符后的两位年份，如 "11月1日-25年11月30日" 中的 "-25年"
    match = YEAR_2_NIAN_RE.search(text)
    if match:
        return 2000 + int(match.group(1))

    return None


def extract_usage_month(text: str) -> int | None:
    compact_text = re.sub(r"\s+", "", text)
    if not compact_text:
        return None

    month_patterns: tuple[tuple[re.Pattern[str], int], ...] = (
        (MONTH_SLASH_DAY_RE, 1),
        (DAY_RANGE_MONTH_END_RE, 1),
        (COMPACT_YM_RE, 2),
        (YEAR_MONTH_RE, 2),
        (YEAR_MONTH_2DIGIT_RE, 2),
        (MONTH_CN_RE, 1),
        (MONTH_DOT_DAY_RE, 1),
    )

    for pattern, group_idx in month_patterns:
        match = pattern.search(compact_text)
        if not match:
            continue
        month = int(match.group(group_idx))
        if 1 <= month <= 12:
            return month

    word_match = MONTH_WORD_RE.search(compact_text)
    if word_match:
        return MONTH_WORD_TO_NUMBER[word_match.group(0)]

    return None


def parse_explicit_year_month_from_text(text: str) -> tuple[int, int] | None:
    compact_text = re.sub(r"\s+", "", text)
    if not compact_text:
        return None

    match = STANDARD_YM_RE.fullmatch(compact_text)
    if match:
        return int(match.group(1)), int(match.group(2))

    match = COMPACT_YM_RE.search(compact_text)
    if match:
        return int(match.group(1)), int(match.group(2))

    match = YEAR_MONTH_RE.search(compact_text)
    if match:
        return int(match.group(1)), int(match.group(2))

    match = YEAR_MONTH_2DIGIT_RE.search(compact_text)
    if match:
        return 2000 + int(match.group(1)), int(match.group(2))

    # 专门处理 "-25年11月" 这类范围分隔符后跟两位年月的情况
    match = YEAR_2_NIAN_MONTH_RE.search(compact_text)
    if match:
        return 2000 + int(match.group(1)), int(match.group(2))

    match = MONTH_YEAR_RE.search(compact_text)
    if match:
        return int(match.group(2)), int(match.group(1))

    year = extract_explicit_year(compact_text)
    if year is None:
        return None

    month = extract_usage_month(compact_text)
    if month is None:
        return None

    return year, month


def parse_usage_year_month(value: object) -> tuple[int, int] | None:
    if value is None:
        return None

    if isinstance(value, (datetime, date)):
        return int(value.year), int(value.month)

    serial_date = coerce_excel_serial_to_date(value)
    if serial_date is not None:
        return int(serial_date.year), int(serial_date.month)

    text = normalize_text(value)
    if not text:
        return None

    explicit_year_month = parse_explicit_year_month_from_text(text)
    if explicit_year_month is not None:
        return explicit_year_month

    month = extract_usage_month(text)
    if month is None:
        return None

    year = extract_explicit_year(text)
    if year is None:
        year = DEFAULT_YEAR_WITHOUT_EXPLICIT
    return year, month


def parse_month_only(text: str) -> int | None:
    compact_text = re.sub(r"\s+", "", text)
    if not compact_text:
        return None

    match = MONTH_ONLY_RE.fullmatch(compact_text)
    if match:
        return int(match.group(1))

    match = MONTH_CN_RE.fullmatch(compact_text)
    if match:
        return int(match.group(1))

    word_match = MONTH_WORD_RE.fullmatch(compact_text)
    if word_match:
        return MONTH_WORD_TO_NUMBER[word_match.group(0)]

    return None


def parse_month_year_month(value: object) -> tuple[int, int] | None:
    if value is None:
        return None

    if isinstance(value, (datetime, date)):
        return int(value.year), int(value.month)

    text = normalize_text(value)
    if not text:
        return None

    explicit_year_month = parse_explicit_year_month_from_text(text)
    if explicit_year_month is not None:
        return explicit_year_month

    month = parse_month_only(text)
    if month is not None:
        return DEFAULT_YEAR_WITHOUT_EXPLICIT, month

    return None


def usage_value_has_explicit_year(value: object) -> bool:
    if value is None:
        return False

    if isinstance(value, (datetime, date)):
        return True

    if coerce_excel_serial_to_date(value) is not None:
        return True

    compact_text = re.sub(r"\s+", "", normalize_text(value))
    if not compact_text:
        return False

    return extract_explicit_year(compact_text) is not None


def month_value_has_explicit_year(value: object) -> bool:
    if value is None:
        return False

    if isinstance(value, (datetime, date)):
        return True

    compact_text = re.sub(r"\s+", "", normalize_text(value))
    if not compact_text:
        return False

    return extract_explicit_year(compact_text) is not None


def infer_implicit_year_month(month: int, floor: tuple[int, int] | None) -> tuple[int, int]:
    month = max(1, min(12, int(month)))
    if floor is None:
        return clamp_to_current((DEFAULT_YEAR_WITHOUT_EXPLICIT, month))

    # 年初出现 11/12 月时，优先判定为上一年，避免被误拉到当年 11/12 再触发后续跨年。
    if floor[1] <= 2 and month >= 11:
        return clamp_to_current((floor[0] - 1, month))

    same_year_candidate = (floor[0], month)
    if same_year_candidate < floor:
        # 仅在年末窗口允许跨年：避免 6月/7月等中段月份被误推到下一年。
        can_rollover = floor[1] >= 11 and month <= 2
        candidate = (floor[0] + 1, month) if can_rollover else same_year_candidate
    else:
        candidate = same_year_candidate

    candidate = clamp_to_current(candidate)
    return candidate


def score_usage_value(value: object) -> int:
    if value is None:
        return 0

    if coerce_excel_serial_to_date(value) is not None:
        return 6

    text = normalize_text(value)
    if not text:
        return 0

    compact_text = re.sub(r"\s+", "", text)
    score = 0
    if extract_explicit_year(compact_text) is not None:
        score += 2
    if DAY_RANGE_MONTH_END_RE.search(compact_text):
        score += 5
    if YEAR_MONTH_RE.search(compact_text) or YEAR_MONTH_2DIGIT_RE.search(compact_text):
        score += 5
    if MONTH_DOT_DAY_RE.search(compact_text):
        score += 4
    if MONTH_CN_RE.search(compact_text) and any(token in compact_text for token in ("日", "号", "-", "~", "至")):
        score += 4
    if compact_text.count("/") >= 2 or compact_text.count(".") >= 2:
        score += 2
    if parse_month_only(compact_text) is not None:
        score -= 2
    return max(score, 0)


def score_month_value(value: object) -> int:
    if value is None:
        return 0

    text = normalize_text(value)
    if not text:
        return 0

    compact_text = re.sub(r"\s+", "", text)
    if STANDARD_YM_RE.fullmatch(compact_text):
        return 10
    if COMPACT_YM_RE.fullmatch(compact_text):
        return 8
    if parse_month_only(compact_text) is not None:
        return 7
    if parse_explicit_year_month_from_text(compact_text) is not None and all(token not in compact_text for token in ("日", "号", "~", "至")):
        return 6
    return 0


def soft_detect_usage_column(worksheet, header_row: int, exclude_col: int | None) -> int | None:
    sample_end = min(worksheet.max_row, header_row + SOFT_SCAN_ROWS)
    best_col = None
    best_score = 0

    for col_idx in range(1, worksheet.max_column + 1):
        if exclude_col is not None and col_idx == exclude_col:
            continue

        total_score = 0
        for row_idx in range(header_row + 1, sample_end + 1):
            total_score += score_usage_value(worksheet.cell(row_idx, col_idx).value)

        if total_score > best_score:
            best_col = col_idx
            best_score = total_score

    if best_score >= 10:
        return best_col
    return None


def soft_detect_month_column(worksheet, header_row: int, exclude_col: int | None) -> int | None:
    sample_end = min(worksheet.max_row, header_row + SOFT_SCAN_ROWS)
    best_col = None
    best_score = 0

    for col_idx in range(1, worksheet.max_column + 1):
        if exclude_col is not None and col_idx == exclude_col:
            continue

        total_score = 0
        for row_idx in range(header_row + 1, sample_end + 1):
            total_score += score_month_value(worksheet.cell(row_idx, col_idx).value)

        if total_score > best_score:
            best_col = col_idx
            best_score = total_score

    if best_score >= 8:
        return best_col
    return None


def rgb_from_color(color) -> str | None:
    if color is None:
        return None

    rgb = None
    if getattr(color, "type", None) == "rgb" and getattr(color, "rgb", None):
        rgb = color.rgb
    elif getattr(color, "type", None) == "indexed" and getattr(color, "indexed", None) is not None:
        idx = color.indexed
        if isinstance(idx, int) and 0 <= idx < len(OPENPYXL_COLOR_INDEX):
            rgb = OPENPYXL_COLOR_INDEX[idx]
    elif getattr(color, "rgb", None):
        rgb = color.rgb

    if not rgb:
        return None

    rgb = str(rgb).replace("#", "").strip()
    if len(rgb) == 8:
        rgb = rgb[2:]
    if len(rgb) != 6:
        return None
    return rgb.upper()


def is_yellow_rgb(rgb: str | None) -> bool:
    if rgb is None:
        return False
    try:
        red = int(rgb[0:2], 16)
        green = int(rgb[2:4], 16)
        blue = int(rgb[4:6], 16)
    except ValueError:
        return False

    return red >= 220 and green >= 170 and blue <= 150


def yellow_signature_of_cell(cell) -> str | None:
    fill = getattr(cell, "fill", None)
    if fill is None:
        return None

    for color in (getattr(fill, "fgColor", None), getattr(fill, "start_color", None), getattr(fill, "end_color", None)):
        rgb = rgb_from_color(color)
        if is_yellow_rgb(rgb):
            return rgb
    return None


def get_row_yellow_signature(worksheet, row_idx: int, usage_col: int | None, month_col: int | None) -> str | None:
    focus_columns: list[int] = []
    if usage_col is not None:
        focus_columns.append(usage_col)
    if month_col is not None and month_col not in focus_columns:
        focus_columns.append(month_col)

    for col_idx in focus_columns:
        signature = yellow_signature_of_cell(worksheet.cell(row_idx, col_idx))
        if signature:
            return signature

    for col_idx in range(1, worksheet.max_column + 1):
        signature = yellow_signature_of_cell(worksheet.cell(row_idx, col_idx))
        if signature:
            return signature
    return None


def process_sheet(worksheet) -> dict[str, Any]:
    detection = detect_header_and_columns(worksheet)

    if detection.month_col is None:
        return {
            "sheet": worksheet.title,
            "status": "skipped",
            "reason": "未找到月份列",
            "header_row": detection.header_row,
            "usage_col": detection.usage_col,
            "month_col": None,
        }

    row_states: list[RowState] = []
    previous_yellow_signature: str | None = None

    for row_idx in range(detection.header_row + 1, worksheet.max_row + 1):
        is_data_row = row_has_any_value(worksheet, row_idx)
        yellow_signature = get_row_yellow_signature(worksheet, row_idx, detection.usage_col, detection.month_col)
        starts_yellow_block = bool(is_data_row and yellow_signature is not None and yellow_signature != previous_yellow_signature)
        month_cell = worksheet.cell(row_idx, detection.month_col)
        usage_value = worksheet.cell(row_idx, detection.usage_col).value if detection.usage_col is not None else None

        _usage_candidate = parse_usage_year_month(usage_value)
        # 显式解析出的未来日期（如序列号 46357=Dec2026、"2026/12/13"）视为数据错误，
        # 不参与推断，回退到上下文填充，避免错误拉高 floor。
        if (_usage_candidate is not None
                and _usage_candidate > CURRENT_YEAR_MONTH
                and usage_value_has_explicit_year(usage_value)):
            _usage_candidate = None

        _month_candidate = parse_month_year_month(month_cell.value)
        if (_month_candidate is not None
                and _month_candidate > CURRENT_YEAR_MONTH
                and month_value_has_explicit_year(month_cell.value)):
            _month_candidate = None

        row_states.append(
            RowState(
                row_idx=row_idx,
                is_data_row=is_data_row,
                month_cell=month_cell,
                usage_value=usage_value,
                month_raw_value=month_cell.value,
                usage_candidate=_usage_candidate,
                month_candidate=_month_candidate,
                yellow_signature=yellow_signature,
                starts_yellow_block=starts_yellow_block,
            )
        )

        previous_yellow_signature = yellow_signature if is_data_row else None

    parsed_by_usage = 0
    parsed_by_month = 0
    filled_by_yellow = 0
    filled_by_context = 0
    updated_cells = 0
    explicit_sources = {"usage_explicit", "month_explicit", "yellow_explicit"}

    last_valid: tuple[int, int] | None = None
    active_yellow_signature: str | None = None
    active_yellow_month: tuple[int, int] | None = None
    active_yellow_is_explicit = False

    for row_state in row_states:
        if not row_state.is_data_row:
            active_yellow_signature = None
            active_yellow_month = None
            active_yellow_is_explicit = False
            continue

        if row_state.yellow_signature is None:
            active_yellow_signature = None
            active_yellow_month = None
            active_yellow_is_explicit = False

        candidate: tuple[int, int] | None = None
        source = ""

        if row_state.usage_candidate is not None:
            candidate = row_state.usage_candidate
            if usage_value_has_explicit_year(row_state.usage_value):
                source = "usage_explicit"
            else:
                candidate = infer_implicit_year_month(candidate[1], last_valid)
                source = "usage_inferred"
            parsed_by_usage += 1
        elif row_state.month_candidate is not None:
            candidate = row_state.month_candidate
            if month_value_has_explicit_year(row_state.month_raw_value):
                source = "month_explicit"
            else:
                candidate = infer_implicit_year_month(candidate[1], last_valid)
                source = "month_inferred"
            parsed_by_month += 1
        else:
            if row_state.yellow_signature is not None:
                if row_state.starts_yellow_block:
                    active_yellow_signature = row_state.yellow_signature
                    row_text_candidate = parse_explicit_year_month_from_text(collect_row_text(worksheet, row_state.row_idx))
                    if (row_text_candidate is not None
                            and row_text_candidate <= CURRENT_YEAR_MONTH):
                        active_yellow_month = row_text_candidate
                        active_yellow_is_explicit = True
                    else:
                        # 黄色行的显式年月也是未来日期时，退化为上下文延续
                        active_yellow_month = last_valid
                        active_yellow_is_explicit = False
                elif active_yellow_signature != row_state.yellow_signature:
                    active_yellow_signature = row_state.yellow_signature
                    active_yellow_month = last_valid
                    active_yellow_is_explicit = False

                if active_yellow_month is not None:
                    candidate = active_yellow_month
                    source = "yellow_explicit" if active_yellow_is_explicit else "yellow"
                    filled_by_yellow += 1

            if candidate is None and last_valid is not None:
                candidate = last_valid
                source = "context_forward"
                filled_by_context += 1

        if candidate is not None:
            if source in explicit_sources:
                row_state.final_ym = clamp_to_current(candidate)
            else:
                row_state.final_ym = coerce_continuous(candidate, last_valid)
            row_state.source = source
            if row_state.final_ym is not None:
                last_valid = row_state.final_ym
                if row_state.yellow_signature is not None:
                    active_yellow_signature = row_state.yellow_signature
                    active_yellow_month = row_state.final_ym
                    active_yellow_is_explicit = source == "yellow_explicit"

    next_valid: tuple[int, int] | None = None
    for row_state in reversed(row_states):
        if not row_state.is_data_row:
            continue
        if row_state.final_ym is not None:
            next_valid = row_state.final_ym
            continue
        if next_valid is not None:
            row_state.final_ym = next_valid
            row_state.source = "context_backward"
            filled_by_context += 1

    first_known = next((row_state.final_ym for row_state in row_states if row_state.is_data_row and row_state.final_ym is not None), None)
    fallback_seed = first_known or (DEFAULT_YEAR_WITHOUT_EXPLICIT, 1)
    for row_state in row_states:
        if row_state.is_data_row and row_state.final_ym is None:
            row_state.final_ym = fallback_seed
            row_state.source = "default_fill"
            filled_by_context += 1

    # 规则：按“连续数据段”处理；每段内首次显式年份出现前的无显式年份推断行，年份不应超过该显式年份。
    segment_rows: list[RowState] = []
    for row_state in row_states + [RowState(-1, False, None, None, None, None, None, None, False)]:
        if row_state.is_data_row:
            segment_rows.append(row_state)
            continue

        if not segment_rows:
            continue

        anchor_row_idx: int | None = None
        anchor_year: int | None = None
        for seg_row in segment_rows:
            if seg_row.final_ym is None:
                continue
            if seg_row.source in explicit_sources:
                anchor_row_idx = seg_row.row_idx
                anchor_year = seg_row.final_ym[0]
                break

        if anchor_row_idx is not None and anchor_year is not None:
            ceiling_year = max(anchor_year, DEFAULT_YEAR_WITHOUT_EXPLICIT)
            for seg_row in segment_rows:
                if seg_row.final_ym is None:
                    continue
                if seg_row.row_idx >= anchor_row_idx:
                    continue
                if seg_row.source in explicit_sources:
                    continue
                if seg_row.final_ym[0] > ceiling_year:
                    seg_row.final_ym = (ceiling_year, seg_row.final_ym[1])

        segment_rows = []

    floor: tuple[int, int] | None = None
    unresolved_rows = 0
    parsed_rows = 0

    for row_state in row_states:
        if not row_state.is_data_row:
            continue

        if row_state.source in explicit_sources and row_state.final_ym is not None:
            row_state.final_ym = clamp_to_current(row_state.final_ym)
        else:
            row_state.final_ym = coerce_continuous(row_state.final_ym, floor)
        if row_state.final_ym is None:
            unresolved_rows += 1
            continue

        floor = row_state.final_ym
        if row_state.source in {
            "usage_explicit",
            "usage_inferred",
            "month_explicit",
            "month_inferred",
            "yellow",
            "yellow_explicit",
        }:
            parsed_rows += 1

        month_cell = row_state.month_cell
        new_value = format_year_month(row_state.final_ym)
        old_value = normalize_text(month_cell.value)
        if old_value != new_value:
            month_cell.value = new_value
            updated_cells += 1

    return {
        "sheet": worksheet.title,
        "status": "ok",
        "header_row": detection.header_row,
        "usage_col": detection.usage_col,
        "month_col": detection.month_col,
        "usage_col_source": detection.usage_source,
        "month_col_source": detection.month_source,
        "parsed_rows": parsed_rows,
        "parsed_by_usage": parsed_by_usage,
        "parsed_by_month": parsed_by_month,
        "filled_by_yellow": filled_by_yellow,
        "filled_by_context": filled_by_context,
        "updated_cells": updated_cells,
        "unresolved_rows": unresolved_rows,
    }


def build_output_path(input_file: Path, suffix: str, output_dir: Path) -> Path:
    normalized_stem = re.sub(
        r"-\d{8}(?:（[^）]*）|\([^)]*\))?$",
        "",
        input_file.stem,
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir / f"{normalized_stem}{suffix}{input_file.suffix}"


def format_column_info(column_index: int | None) -> str:
    if column_index is None:
        return "-"
    return f"{get_column_letter(column_index)}({column_index})"


def process_workbook(input_file: Path, suffix: str) -> dict[str, int | str]:
    output_file = build_output_path(input_file, suffix, DEFAULT_OUTPUT_DIR)

    print(f"读取文件: {input_file}")
    workbook = load_workbook(input_file)

    total_parsed_rows = 0
    total_parsed_usage = 0
    total_parsed_month = 0
    total_filled_yellow = 0
    total_filled_context = 0
    total_updated_cells = 0
    total_unresolved_rows = 0

    for worksheet in workbook.worksheets:
        result = process_sheet(worksheet)
        if result["status"] == "skipped":
            print(
                f"工作表 {result['sheet']}: 表头行={result['header_row']}，"
                f"使用时间列={format_column_info(result['usage_col'])}，"
                f"月份列=未找到月份列，跳过写入"
            )
            continue

        total_parsed_rows += int(result["parsed_rows"])
        total_parsed_usage += int(result["parsed_by_usage"])
        total_parsed_month += int(result["parsed_by_month"])
        total_filled_yellow += int(result["filled_by_yellow"])
        total_filled_context += int(result["filled_by_context"])
        total_updated_cells += int(result["updated_cells"])
        total_unresolved_rows += int(result["unresolved_rows"])

        print(
            f"工作表 {result['sheet']}: 表头行={result['header_row']}，"
            f"使用时间列={format_column_info(result['usage_col'])}[{result['usage_col_source']}]，"
            f"月份列={format_column_info(result['month_col'])}[{result['month_col_source']}]，"
            f"解析行数={result['parsed_rows']}，"
            f"修改单元格数={result['updated_cells']}，"
            f"未解析行数={result['unresolved_rows']}，"
            f"使用时间解析={result['parsed_by_usage']}，"
            f"月份列解析={result['parsed_by_month']}，"
            f"黄色填充={result['filled_by_yellow']}，"
            f"上下文填充={result['filled_by_context']}"
        )

    workbook.save(output_file)

    print("\n处理完成")
    print(f"输出文件: {output_file}")
    print(f"解析行数合计: {total_parsed_rows}")
    print(f"使用时间解析合计: {total_parsed_usage}")
    print(f"月份列解析合计: {total_parsed_month}")
    print(f"黄色填充合计: {total_filled_yellow}")
    print(f"上下文填充合计: {total_filled_context}")
    print(f"修改单元格合计: {total_updated_cells}")
    print(f"未解析行数合计: {total_unresolved_rows}")

    return {
        "input": str(input_file),
        "output": str(output_file),
        "parsed_rows": total_parsed_rows,
        "updated_cells": total_updated_cells,
    }


def main() -> None:
    args = parse_args()
    provided_input = args.input_override or args.input_file

    if provided_input:
        input_file = Path(provided_input).expanduser().resolve()
        input_files = [input_file]
    else:
        input_files = [p.resolve() for p in list_all_raw_input_files(DEFAULT_INPUT_DIR)]
        if not input_files:
            raise FileNotFoundError(
                f"未在原始数据目录找到可用 .xlsx 文件: {DEFAULT_INPUT_DIR}"
            )
        print(f"自动发现原始文件 {len(input_files)} 个，开始批量标准化...")

    batch_parsed_rows = 0
    batch_updated_cells = 0

    for idx, input_file in enumerate(input_files, start=1):
        if not input_file.exists():
            raise FileNotFoundError(f"输入文件不存在: {input_file}")
        if input_file.suffix.lower() != ".xlsx":
            raise ValueError(f"仅支持 .xlsx 文件: {input_file}")

        if len(input_files) > 1:
            print(f"\n[{idx}/{len(input_files)}] 开始处理: {input_file.name}")

        result = process_workbook(input_file, args.suffix)
        batch_parsed_rows += int(result["parsed_rows"])
        batch_updated_cells += int(result["updated_cells"])

    if len(input_files) > 1:
        print("\n批量处理完成")
        print(f"处理文件数: {len(input_files)}")
        print(f"累计解析行数: {batch_parsed_rows}")
        print(f"累计修改单元格: {batch_updated_cells}")


if __name__ == "__main__":
    main()
