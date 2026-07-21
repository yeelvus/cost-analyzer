
# --- path bootstrap (S_数据分析模块) ---
import sys
from pathlib import Path as _PathForBootstrap
try:
    from S_数据分析模块._bootstrap import ensure_code_root, ensure_script_dir
    ensure_code_root(_PathForBootstrap(__file__))
    ensure_script_dir(__file__)
except Exception:
    _p = _PathForBootstrap(__file__).resolve()
    for _c in [_p.parent, *_p.parents]:
        if (_c / "settings.py").exists():
            if str(_c) not in sys.path:
                sys.path.insert(0, str(_c))
            break
    if str(_p.parent) not in sys.path:
        sys.path.insert(0, str(_p.parent))
# --- end bootstrap ---

import re
import shutil
from pathlib import Path
from typing import Dict, List, Tuple

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config_driven_report_builder import create_report_builder


def normalize_project_code(value) -> str:
    if pd.isna(value):
        return ""
    code = str(value).strip().upper()
    if code.lower() in {"nan", "none"}:
        return ""
    return code


def coerce_float(value) -> float:
    if pd.isna(value):
        return 0.0
    try:
        return float(str(value).replace(",", ""))
    except Exception:
        return 0.0


def first_notna(series: pd.Series) -> str:
    non_null = series.dropna()
    if len(non_null) == 0:
        return ""
    return str(non_null.iloc[0])


def ensure_columns(df: pd.DataFrame, columns: List[str]) -> pd.DataFrame:
    df = df.copy()
    for col in columns:
        if col not in df.columns:
            df[col] = None
    return df


def sanitize_filename(name: str) -> str:
    invalid = '\\/:*?"<>|'
    cleaned = "".join("-" if ch in invalid else ch for ch in name)
    return cleaned.strip()


def load_responsibility_cost_mapping(responsibility_dir: Path) -> Dict[str, Path]:
    mapping = {}
    if not responsibility_dir.is_dir():
        return mapping

    pattern = re.compile(r"^(THCC-\d{4}-\d{3})_.*_责任成本\.xlsx$")

    for file_path in responsibility_dir.glob("THCC-*.xlsx"):
        match = pattern.match(file_path.name)
        if match:
            code = match.group(1)
            mapping[code] = file_path

    return mapping


def get_responsibility_cost_data(file_path: Path) -> Dict[str, pd.DataFrame]:
    if not file_path.exists():
        return {}

    try:
        xl = pd.ExcelFile(file_path)
        result = {}
        for sheet_name in xl.sheet_names:
            df_raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)

            if df_raw.empty:
                continue

            header_row_idx = None
            header_keywords = {"序号", "名称", "金额"}

            for idx, row in df_raw.iterrows():
                row_str = " ".join(str(val).strip() for val in row if pd.notna(val))
                matched_keywords = sum(1 for keyword in header_keywords if keyword in row_str)
                if matched_keywords >= 2:
                    header_row_idx = idx
                    break

            if header_row_idx is not None:
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=header_row_idx)
            else:
                try:
                    df = pd.read_excel(file_path, sheet_name=sheet_name, header=2)
                except Exception:
                    df = pd.read_excel(file_path, sheet_name=sheet_name, header=0)

            if df.empty or df.shape[0] == 0:
                continue

            result[sheet_name] = df
        return result
    except Exception as e:
        print(f"警告：读取责任成本文件失败 {file_path}: {e}")
        return {}


def build_provenance_rows(
    contract_dir: Path,
    procurement_path: Path,
    concrete_paths: List[Path],
    equipment_path: Path,
    labor_path: Path,
    finance_path: Path,
    cost_stats_path: Path,
    manager_map_path: Path,
    budget_path: Path,
) -> Tuple[List[Dict], List[Dict], List[Dict], List[Dict], List[Dict], List[Dict], List[Dict], List[Dict]]:
    input_rows: List[Dict] = []
    cost_stats_fields: List[Dict] = []
    procurement_fields: List[Dict] = []
    concrete_fields: List[Dict] = []
    equipment_fields: List[Dict] = []
    labor_fields: List[Dict] = []
    finance_fields: List[Dict] = []
    budget_fields: List[Dict] = []

    def add_input(category: str, item: str, source: Path, sheet: str = "", columns: str = "", note: str = ""):
        input_rows.append({
            "类别": category,
            "项目": item,
            "来源文件": str(source),
            "Sheet": sheet,
            "列": columns,
            "备注": note,
        })

    def add_field(field_list: List[Dict], field_name: str, source_file: str, sheet: str, column: str, note: str = ""):
        field_list.append({
            "字段": field_name,
            "来源文件": source_file,
            "Sheet": sheet,
            "列": column,
            "备注": note,
        })

    add_input("代码", "合同汇总主程序", Path(__file__).resolve(), note="生成项目报告与索引")
    add_input("代码", "成本分析构建器", (Path(__file__).parent / "config_driven_report_builder.py").resolve())
    add_input("配置", "成本分析配置", (Path(__file__).parent / "cost_analysis_report_config.json").resolve())

    add_input("合同", "合同管理根目录", contract_dir, note="总包/分包合同文件夹")
    add_input("采购", "采购明细合并表", procurement_path, sheet="项目汇总(优先)/合并数据", columns="项目编码, 项目名称, 总采购金额")
    for c_path in concrete_paths:
        add_input("商混", "商混明细汇总", c_path, sheet="客户总额汇总", columns="项目编码, 项目名称, 总方量, 总车次, 金额, 金额（万泰铢）")
    add_input("机具", "机具使用费汇总", equipment_path, sheet="项目汇总机具费", columns="项目编码, 项目名称, 费用合计(泰铢)")
    add_input("人工", "人工费汇总", labor_path, sheet="按项目汇总", columns="项目编码, 项目名称, 中方正式工铢, 中方劳务工铢, 子公司属地工铢, 劳务公司铢, 签证费, 总人工费")
    add_input("财务", "分包对账与付款统计", finance_path, sheet="项目汇总（万泰铢）/分类汇总")
    add_input("成本统计", "项目统计数据汇总", cost_stats_path, sheet="默认sheet", columns="成本统计字段集合")
    add_input("项目经理", "项目基本信息表", manager_map_path, sheet="默认sheet（自动识别表头）")
    add_input("预算", "预算总汇总", budget_path, sheet="总汇总")

    cost_stats_columns = [
        "预计总成本", "收入", "毛利", "毛利率",
        "业主预计结算额", "总包合同额(万泰铢)", "签证情况",
        "分包结算额", "分包合同额(万泰铢)", "分包签证情况", "分包已付款",
        "材料采购实际已发生", "材料采购后续预估发生",
        "采购部采购金额", "项目部采购金额",
        "材料租赁费用", "材料租赁后续预估",
        "工资支出", "管理人员工资",
        "机具费支出", "机具费后续预估",
        "人员调遣支出", "人员调遣后续预估",
        "其他费用支出", "其他费用后续预估",
        "临建费用支出", "临建费用后续预估", "实际已发生金额",
    ]
    for col in cost_stats_columns:
        add_field(cost_stats_fields, col, str(cost_stats_path), "默认sheet", col, "引用项目成本统计表（各项目使用各自最新月份数据）")

    procurement_source_cols = [
        ("项目编码", "项目编码", "用于匹配"),
        ("项目名称", "项目名称", "项目名称"),
        ("总采购金额", "总采购金额", "采购金额(泰铢)"),
    ]
    for field, col, note in procurement_source_cols:
        add_field(procurement_fields, field, str(procurement_path), "项目汇总(优先)/合并数据", col, note)

    concrete_source_cols = [
        ("项目编码", "项目编码", "用于匹配"),
        ("项目名称", "项目名称", "项目名称"),
        ("金额", "金额", "商混金额(泰铢)"),
        ("金额（万泰铢）", "金额（万泰铢）", "商混金额(万泰铢)"),
    ]
    for field, col, note in concrete_source_cols:
        add_field(concrete_fields, field, str(concrete_paths[0]) if concrete_paths else "混凝土汇总文件", "客户总额汇总", col, note)

    equipment_source_cols = [
        ("项目编码", "项目编码", "用于匹配"),
        ("项目名称", "项目名称", "项目名称"),
        ("费用合计(泰铢)", "费用合计(泰铢)", "机具费用汇总"),
    ]
    for field, col, note in equipment_source_cols:
        add_field(equipment_fields, field, str(equipment_path), "项目汇总机具费", col, note)

    labor_source_cols = [
        ("项目编码", "项目编码", "用于匹配"),
        ("项目名称", "项目名称", "项目名称"),
        ("中方正式工铢", "中方正式工铢", "按项目汇总分项"),
        ("中方劳务工铢", "中方劳务工铢", "按项目汇总分项"),
        ("子公司属地工铢", "子公司属地工铢", "按项目汇总分项"),
        ("劳务公司铢", "劳务公司铢", "按项目汇总分项"),
        ("签证费", "签证费", "按项目汇总分项"),
        ("总人工费", "总人工费", "按项目汇总总额"),
        ("年", "年", "按年月汇总(若有)"),
        ("月", "月", "按年月汇总(若有)"),
    ]
    for field, col, note in labor_source_cols:
        sheet = "按年月项目汇总" if field in ["年", "月"] else "按项目汇总"
        add_field(labor_fields, field, str(labor_path), sheet, col, note)

    finance_source_cols = [
        ("项目编码", "项目编码", "用于匹配"),
        ("含税实付金额（万泰铢）", "含税实付金额（万泰铢）", "含税实付"),
        ("国内代付金额（人民币）", "国内代付金额（人民币）", "国内代付"),
        ("国内代付人民币金额（万泰铢）", "国内代付人民币金额（万泰铢）", "国内代付换算"),
        ("付款汇总（万泰铢）", "付款汇总（万泰铢）", "总付款"),
    ]
    for field, col, note in finance_source_cols:
        sheet = "分类汇总" if field == "项目编码" else "项目汇总（万泰铢）"
        add_field(finance_fields, field, str(finance_path), sheet, col, note)

    budget_source_cols = [
        ("序号", "序号", "序号"),
        ("项目编码", "项目编码", "用于匹配"),
        ("项目名称", "项目名称", "项目名称"),
        ("币种", "币种", "币种"),
        ("标段人材机汇总", "标段人材机汇总", "标段汇总"),
        ("人_明细", "人_明细", "人工费详细"),
        ("管_明细", "管_明细", "管理人员费用"),
        ("设机_明细", "设机_明细", "设备机具费用"),
        ("材类明细", "材类明细", "材料费用"),
        ("其他类别_明细", "其他类别_明细", "其他费用"),
        ("钢筋", "钢筋", "钢筋"),
        ("混凝土", "混凝土", "混凝土"),
        ("砂浆", "砂浆", "砂浆"),
        ("碎石", "碎石", "碎石"),
        ("砖", "砖", "砖"),
        ("其他材料", "其他材料", "其他材料"),
        ("人管设机材类明细其他类别_合计", "人管设机材类明细其他类别_合计", "总合计1"),
        ("钢筋混凝土砂浆碎石砖其他材料_合计", "钢筋混凝土砂浆碎石砖其他材料_合计", "总合计2"),
    ]
    for field, col, note in budget_source_cols:
        add_field(budget_fields, field, str(budget_path), "总汇总", col, note)

    return input_rows, cost_stats_fields, procurement_fields, concrete_fields, equipment_fields, labor_fields, finance_fields, budget_fields


def export_provenance_excel(
    output_dir: Path,
    input_rows: List[Dict],
    cost_stats_fields: List[Dict],
    procurement_fields: List[Dict],
    concrete_fields: List[Dict],
    equipment_fields: List[Dict],
    labor_fields: List[Dict],
    finance_fields: List[Dict],
    budget_fields: List[Dict],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "数据来源与字段映射.xlsx"

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        if input_rows:
            pd.DataFrame(input_rows).to_excel(writer, index=False, sheet_name="输入数据来源")
        if cost_stats_fields:
            pd.DataFrame(cost_stats_fields).to_excel(writer, index=False, sheet_name="成本统计字段来源")
        if procurement_fields:
            pd.DataFrame(procurement_fields).to_excel(writer, index=False, sheet_name="采购字段来源")
        if concrete_fields:
            pd.DataFrame(concrete_fields).to_excel(writer, index=False, sheet_name="商混字段来源")
        if equipment_fields:
            pd.DataFrame(equipment_fields).to_excel(writer, index=False, sheet_name="机具字段来源")
        if labor_fields:
            pd.DataFrame(labor_fields).to_excel(writer, index=False, sheet_name="人工字段来源")
        if finance_fields:
            pd.DataFrame(finance_fields).to_excel(writer, index=False, sheet_name="财务字段来源")
        if budget_fields:
            pd.DataFrame(budget_fields).to_excel(writer, index=False, sheet_name="预算字段来源")

        ws_input = writer.sheets.get("输入数据来源")
        if ws_input:
            style_sheet(ws_input, amount_headers=[])
            for col_idx in range(1, ws_input.max_column + 1):
                col_letter = get_column_letter(col_idx)
                ws_input.column_dimensions[col_letter].width = 20

        for sheet_name in ["成本统计字段来源", "采购字段来源", "商混字段来源", "机具字段来源", "人工字段来源", "财务字段来源", "预算字段来源"]:
            ws = writer.sheets.get(sheet_name)
            if ws:
                style_sheet(ws, amount_headers=[])
                for col_idx in range(1, ws.max_column + 1):
                    col_letter = get_column_letter(col_idx)
                    ws.column_dimensions[col_letter].width = 20

    writer.book.save(output_path)


def style_sheet(ws, amount_headers: List[str]) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="D3D3D3")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=False, horizontal="center")
        cell.border = border

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    headers = [c.value for c in ws[1]]
    for header in amount_headers:
        if header in headers:
            idx = headers.index(header) + 1
            for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=ws.max_row):
                for item in cell:
                    item.number_format = "#,##0.00"

    def calculate_width(value):
        if value is None:
            return 0
        text = str(value)
        width = 0
        for char in text:
            if "\u4e00" <= char <= "\u9fff":
                width += 2.2
            else:
                width += 1.1
        return width

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        values = [ws.cell(row=r, column=col_idx).value for r in range(1, ws.max_row + 1)]
        max_width = 0
        for value in values:
            width = calculate_width(value)
            if width > max_width:
                max_width = width

        adjusted_width = min(max(max_width + 2, 12), 60)
        ws.column_dimensions[col_letter].width = adjusted_width

    ws.freeze_panes = "A2"


def style_analysis_sheet(ws) -> None:
    report_builder = create_report_builder()
    report_builder.style_report_sheet(ws)


def build_analysis_report_rows(
    project_code: str,
    project_name: str,
    summary_row: Dict,
    cost_stats_row,
    main_rows: pd.DataFrame,
    sub_rows: pd.DataFrame,
    concrete_rows: pd.DataFrame,
    equipment_rows: pd.DataFrame = None,
    finance_category_rows: pd.DataFrame = None,
    procurement_rows: pd.DataFrame = None,
    labor_rows: pd.DataFrame = None,
    extra_proc_rows: pd.DataFrame = None,
    responsibility_cost_data: Dict[str, pd.DataFrame] = None,
    data_source_info: Dict = None,
) -> List[List]:
    report_builder = create_report_builder()

    if extra_proc_rows is not None and not extra_proc_rows.empty:
        if procurement_rows is not None and not procurement_rows.empty:
            procurement_rows = pd.concat([procurement_rows, extra_proc_rows], ignore_index=True)
        else:
            procurement_rows = extra_proc_rows

    return report_builder.build_report_rows(
        project_code=project_code,
        project_name=project_name,
        cost_stats_row=cost_stats_row,
        summary_row=summary_row,
        main_rows=main_rows,
        sub_rows=sub_rows,
        concrete_rows=concrete_rows,
        equipment_rows=equipment_rows,
        labor_rows=labor_rows,
        finance_category_rows=finance_category_rows,
        procurement_rows=procurement_rows,
        responsibility_cost_data=responsibility_cost_data,
        data_source_info=data_source_info,
    )


def build_project_summary(
    main_df: pd.DataFrame,
    sub_df: pd.DataFrame,
    procurement_df: pd.DataFrame,
    material_arrears_df: pd.DataFrame,
    concrete_df: pd.DataFrame,
    equipment_df: pd.DataFrame,
    labor_df: pd.DataFrame,
    labor_month_df: pd.DataFrame,
    finance_summary_df: pd.DataFrame = None,
    finance_category_df: pd.DataFrame = None,
    budget_df: pd.DataFrame = None,
):
    main_df = ensure_columns(main_df, ["项目编码", "项目简称", "项目状态", "金额_THB", "是否采购类合同"])
    sub_df = ensure_columns(sub_df, ["项目编码", "项目简称", "项目状态", "金额_THB", "是否采购类合同"])
    procurement_df = ensure_columns(procurement_df, ["项目编码", "金额_THB"])
    material_arrears_df = ensure_columns(material_arrears_df, ["项目编码", "欠款金额_THB", "标准项目名称"])
    concrete_df = ensure_columns(concrete_df, ["项目编码", "金额_THB", "标准项目名称"])
    equipment_df = ensure_columns(equipment_df, ["项目编码", "金额_THB", "标准项目名称"])
    labor_df = ensure_columns(labor_df, ["项目编码", "人工费汇总（泰铢）", "项目名称", "target_status"])
    labor_month_df = ensure_columns(labor_month_df, ["项目编码", "项目名称", "年", "月"])

    if finance_summary_df is None:
        finance_summary_df = pd.DataFrame()
    if finance_category_df is None:
        finance_category_df = pd.DataFrame()
    if budget_df is None:
        budget_df = pd.DataFrame()

    all_codes = set(main_df.get("项目编码", pd.Series(dtype=str)).unique())
    all_codes |= set(sub_df.get("项目编码", pd.Series(dtype=str)).unique())
    all_codes |= set(procurement_df.get("项目编码", pd.Series(dtype=str)).unique())
    all_codes |= set(material_arrears_df.get("项目编码", pd.Series(dtype=str)).unique())
    all_codes |= set(concrete_df.get("项目编码", pd.Series(dtype=str)))
    all_codes |= set(equipment_df.get("项目编码", pd.Series(dtype=str)))
    all_codes |= set(labor_df.get("项目编码", pd.Series(dtype=str)))
    if not finance_summary_df.empty:
        all_codes |= set(finance_summary_df.get("项目编码", pd.Series(dtype=str)).unique())
    if not budget_df.empty:
        all_codes |= set(budget_df.get("项目编码", pd.Series(dtype=str)).unique())

    proc_sum = procurement_df.groupby("项目编码")["金额_THB"].sum() if not procurement_df.empty else {}
    material_arrears_sum = material_arrears_df.groupby("项目编码")["欠款金额_THB"].sum() if not material_arrears_df.empty else {}
    concrete_sum = concrete_df.groupby("项目编码")["金额_THB"].sum() if not concrete_df.empty else {}
    equipment_sum = equipment_df.groupby("项目编码")["金额_THB"].sum() if not equipment_df.empty else {}
    labor_sum = labor_df.groupby("项目编码")["人工费汇总（泰铢）"].sum() if not labor_df.empty else {}

    budget_sum = {}
    if not budget_df.empty:
        exclude_cols = {"项目编码", "项目名称", "币种"}
        amount_cols = [col for col in budget_df.columns if col not in exclude_cols]

        for code in budget_df["项目编码"].unique():
            code_str = str(code).strip()
            if not code_str:
                continue
            code_rows = budget_df[budget_df["项目编码"] == code]
            budget_sum[code_str] = {}
            for col in amount_cols:
                budget_sum[code_str][col] = code_rows[col].sum() if col in code_rows.columns else 0.0

    finance_sum_dict = {}
    if not finance_summary_df.empty:
        for col in ["含税实付金额（万泰铢）", "国内代付金额（人民币）", "国内代付人民币金额（万泰铢）", "付款汇总（万泰铢）"]:
            if col in finance_summary_df.columns:
                finance_sum_dict[col] = finance_summary_df.groupby("项目编码")[col].sum().astype(float)

    valid_codes: List[str] = []
    for raw_code in all_codes:
        if pd.isna(raw_code):
            continue
        code_str = normalize_project_code(raw_code)
        if not code_str or code_str.lower() in {"nan", "none"} or code_str == "0":
            continue
        valid_codes.append(code_str)

    for code_str in sorted(valid_codes, key=str):
        main_rows = main_df[main_df["项目编码"] == code_str].copy()
        sub_rows = sub_df[sub_df["项目编码"] == code_str].copy()
        project_name = first_notna(main_rows["项目简称"]) or first_notna(sub_rows["项目简称"]) or first_notna(
            procurement_df.loc[procurement_df["项目编码"] == code_str, "标准项目名称"]
        ) or first_notna(
            material_arrears_df.loc[material_arrears_df["项目编码"] == code_str, "标准项目名称"]
        ) or first_notna(
            concrete_df.loc[concrete_df["项目编码"] == code_str, "标准项目名称"]
        ) or first_notna(
            equipment_df.loc[equipment_df["项目编码"] == code_str, "标准项目名称"]
        ) or first_notna(labor_df.loc[labor_df["项目编码"] == code_str, "项目名称"]) or first_notna(
            budget_df.loc[budget_df["项目编码"] == code_str, "项目名称"]
        )

        project_status = first_notna(main_rows["项目状态"]) or first_notna(sub_rows["项目状态"]) or first_notna(
            labor_df.loc[labor_df["项目编码"] == code_str, "target_status"]
            if "target_status" in labor_df.columns
            else pd.Series(dtype=str)
        )

        main_proc = main_rows[main_rows["是否采购类合同"]].copy()
        main_proc["合同来源"] = "总包合同"
        sub_proc = sub_rows[sub_rows["是否采购类合同"]].copy()
        sub_proc["合同来源"] = "分包合同"
        procurement_rows = pd.concat([main_proc, sub_proc], ignore_index=True)

        extra_proc_rows = procurement_df[procurement_df["项目编码"] == code_str].copy()
        concrete_rows = concrete_df[concrete_df["项目编码"] == code_str].copy()
        equipment_rows = equipment_df[equipment_df["项目编码"] == code_str].copy()
        labor_rows = labor_df[labor_df["项目编码"] == code_str].copy()
        labor_month_rows = labor_month_df[labor_month_df["项目编码"] == code_str].copy()

        budget_data = budget_sum.get(code_str, {})

        budget_currency = ""
        if not budget_df.empty:
            budget_project_rows = budget_df[budget_df["项目编码"] == code_str]
            if not budget_project_rows.empty and "币种" in budget_project_rows.columns:
                budget_currency = str(budget_project_rows["币种"].iloc[0]) if pd.notna(budget_project_rows["币种"].iloc[0]) else ""

        summary_row = {
            "项目编码": code_str,
            "项目名称": project_name,
            "项目状态": project_status,
            "总包合同金额": main_rows.get("金额_THB", pd.Series(dtype=float)).sum(),
            "分包合同金额": sub_rows.get("金额_THB", pd.Series(dtype=float)).sum(),
            "材料采购金额": proc_sum.get(code_str, 0.0),
            "材料欠款金额": material_arrears_sum.get(code_str, 0.0),
            "商混站金额": concrete_sum.get(code_str, 0.0),
            "机具使用费": equipment_sum.get(code_str, 0.0),
            "人工费": labor_sum.get(code_str, 0.0),
            "含税实付金额（万泰铢）": (finance_sum_dict.get("含税实付金额（万泰铢）", pd.Series()).get(code_str, 0.0) if "含税实付金额（万泰铢）" in finance_sum_dict else 0.0) * 10000,
            "国内代付金额（人民币）": finance_sum_dict.get("国内代付金额（人民币）", pd.Series()).get(code_str, 0.0) if "国内代付金额（人民币）" in finance_sum_dict else 0.0,
            "国内代付人民币金额（万泰铢）": (finance_sum_dict.get("国内代付人民币金额（万泰铢）", pd.Series()).get(code_str, 0.0) if "国内代付人民币金额（万泰铢）" in finance_sum_dict else 0.0) * 10000,
            "付款汇总（万泰铢）": (finance_sum_dict.get("付款汇总（万泰铢）", pd.Series()).get(code_str, 0.0) if "付款汇总（万泰铢）" in finance_sum_dict else 0.0) * 10000,
            "利润额": 0.0,
            "分包代付款": 0.0,
            "预算币种": budget_currency,
        }

        for col, val in budget_data.items():
            if col not in ["项目编码", "项目名称", "币种"]:
                summary_row[f"预算_{col}"] = val

        finance_category_rows = pd.DataFrame()
        if not finance_category_df.empty:
            finance_category_rows = finance_category_df[finance_category_df["项目编码"] == code_str].copy()

        yield (
            code_str,
            project_name,
            summary_row,
            main_rows,
            sub_rows,
            procurement_rows,
            extra_proc_rows,
            concrete_rows,
            equipment_rows,
            labor_rows,
            labor_month_rows,
            finance_category_rows,
        )


def export_project_report(
    output_dir: Path,
    project_code: str,
    project_name: str,
    summary_row: Dict,
    main_rows: pd.DataFrame,
    sub_rows: pd.DataFrame,
    procurement_rows: pd.DataFrame,
    extra_proc_rows: pd.DataFrame,
    concrete_rows: pd.DataFrame,
    equipment_rows: pd.DataFrame,
    labor_rows: pd.DataFrame,
    labor_month_rows: pd.DataFrame,
    finance_category_rows: pd.DataFrame = None,
    cost_stats_df: pd.DataFrame = None,
    budget_df: pd.DataFrame = None,
    responsibility_cost_mapping: Dict[str, Path] = None,
    data_source_info: Dict = None,
) -> None:
    if finance_category_rows is None:
        finance_category_rows = pd.DataFrame()
    if cost_stats_df is None:
        cost_stats_df = pd.DataFrame()
    if budget_df is None:
        budget_df = pd.DataFrame()
    if responsibility_cost_mapping is None:
        responsibility_cost_mapping = {}
    if data_source_info is None:
        data_source_info = {}

    project_code = str(project_code)
    project_name = str(project_name) if project_name else ""

    output_dir.mkdir(parents=True, exist_ok=True)
    filename = sanitize_filename(f"{project_code}_{project_name}" if project_name else project_code) + ".xlsx"
    output_path = output_dir / filename

    main_columns = [
        "合同子文件夹", "合同类型", "合同金额", "含税金额", "不含税金额", "金额（万泰铢）", "币种", "合同签订日期", "对方公司名称",
    ]
    sub_columns = ["合同子文件夹", "工程范围", "合同类型", "合同金额", "对方公司名称"]

    main_rows = ensure_columns(main_rows, main_columns)
    sub_rows = ensure_columns(sub_rows, sub_columns)

    cost_stats_row = None
    project_stats_month = ""
    if not cost_stats_df.empty:
        if "匹配编码" in cost_stats_df.columns:
            match_rows = cost_stats_df[cost_stats_df["匹配编码"] == project_code]
        else:
            match_rows = cost_stats_df[cost_stats_df.get("项目编码", pd.Series(dtype=str)) == project_code]

        if not match_rows.empty:
            cost_stats_row = match_rows.iloc[0]
            if "月份" in cost_stats_row.index and pd.notna(cost_stats_row["月份"]):
                project_stats_month = str(cost_stats_row["月份"])
                print(f"  → 项目 {project_code} 使用成本统计数据月份: {project_stats_month}")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        workbook = writer.book
        if "Sheet" in workbook.sheetnames:
            del workbook["Sheet"]

        responsibility_cost_data = None
        if project_code in responsibility_cost_mapping:
            responsibility_file = responsibility_cost_mapping[project_code]
            responsibility_cost_data = get_responsibility_cost_data(responsibility_file)

        project_data_source_info = data_source_info.copy()
        if project_stats_month:
            project_data_source_info["项目成本统计汇总表"] = {
                "文件名": "项目统计数据汇总",
                "统计截止日期": project_stats_month,
            }

        ws_report = workbook.create_sheet("成本分析报告", 0)
        writer.sheets["成本分析报告"] = ws_report

        report_rows = build_analysis_report_rows(
            project_code,
            project_name,
            summary_row,
            cost_stats_row,
            main_rows,
            sub_rows,
            concrete_rows,
            equipment_rows,
            finance_category_rows,
            procurement_rows,
            labor_rows,
            extra_proc_rows,
            responsibility_cost_data,
            project_data_source_info,
        )
        for r_idx, row in enumerate(report_rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                ws_report.cell(row=r_idx, column=c_idx, value=val)
            if row[0] and isinstance(row[0], str) and re.match(r"^[一二三四五六七八九]、", row[0]):
                ws_report.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=5)
        ws_report.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
        for r_idx in range(2, ws_report.max_row + 1):
            if ws_report.cell(row=r_idx, column=1).value == "备注":
                ws_report.merge_cells(start_row=r_idx, start_column=2, end_row=r_idx, end_column=5)

        if cost_stats_row is not None:
            cost_stats_columns = [
                "预计总成本", "收入", "毛利", "毛利率", "业主预计结算额", "总包合同额(万泰铢)", "签证情况", "分包结算额", "分包合同额(万泰铢)",
                "分包签证情况", "分包已付款", "材料采购实际已发生", "材料采购后续预估发生", "采购部采购金额", "项目部采购金额", "材料租赁费用",
                "材料租赁后续预估", "工资支出", "管理人员工资", "机具费支出", "机具费后续预估", "人员调遣支出", "人员调遣后续预估", "其他费用支出",
                "其他费用后续预估", "临建费用支出", "临建费用后续预估", "实际已发生金额",
            ]
            dynamic_cost_cols = [
                col for col in cost_stats_row.index if isinstance(col, str) and "付款财务数据" in col and col not in cost_stats_columns
            ]
            cost_stats_columns.extend(dynamic_cost_cols)
            available_cost_cols = {}
            for col in cost_stats_columns:
                if col in cost_stats_row.index and pd.notna(cost_stats_row[col]):
                    available_cost_cols[col] = cost_stats_row[col]
            if available_cost_cols:
                cost_df = pd.DataFrame(list(available_cost_cols.items()), columns=["项目", "金额"])
                cost_df.to_excel(writer, index=False, sheet_name="项目成本统计汇总表")

        main_rows[main_columns].to_excel(writer, index=False, sheet_name="总包合同明细")
        sub_rows[sub_columns].to_excel(writer, index=False, sheet_name="分包合同明细")

        if not extra_proc_rows.empty:
            procurement_columns = [
                "标准项目名称", "物品名称", "规格", "单位", "数量", "单价", "不含税金额", "7%增值税", "含税金额", "货币种类", "汇率", "签订日期", "供货商名称", "采购部门", "备注",
            ]
            available_procurement_cols = [col for col in procurement_columns if col in extra_proc_rows.columns]
            if not available_procurement_cols:
                fallback_cols = ["项目编码", "项目名称", "标准项目名称", "总采购金额", "含税金额", "金额_THB"]
                available_procurement_cols = [col for col in fallback_cols if col in extra_proc_rows.columns]
            extra_proc_rows[available_procurement_cols].to_excel(writer, index=False, sheet_name="材料采购明细")
        if not concrete_rows.empty:
            concrete_cols = [col for col in concrete_rows.columns if col not in {"项目编码", "项目状态"}]
            concrete_rows[concrete_cols].to_excel(writer, index=False, sheet_name="商混站明细")
        if not equipment_rows.empty:
            equipment_output_cols = [
                "sheet名称", "项目编码", "项目名称", "开始月份", "区域", "标准机具种类", "标准型号规格", "序号", "机具种类", "型号规格", "台班费_泰铢", "本次结算台班",
                "加班小时", "加班费", "进出场费运费", "移动费拆装费", "结算费用_泰铢", "增值税7%", "机具租赁预扣税3%", "机具租赁预扣税5%", "本次费用合计", "使用日期", "年月",
                "出租单位", "月份", "备注", "月租金额_泰铢", "备注加班费单价",
            ]
            equipment_export_df = equipment_rows.copy()
            for col in equipment_output_cols:
                if col not in equipment_export_df.columns:
                    equipment_export_df[col] = ""
            equipment_export_df[equipment_output_cols].to_excel(writer, index=False, sheet_name="机具使用费明细")
        if not labor_rows.empty:
            labor_required_cols = ["中方正式工铢", "中方劳务工铢", "子公司属地工铢", "劳务公司铢", "签证费", "总人工费", "人工费汇总（泰铢）"]
            exclude_cols = {"项目编码", "项目状态", "项目名称", "target_status"}
            available_cols = [col for col in labor_required_cols if col in labor_rows.columns]
            other_cols = [col for col in labor_rows.columns if col not in exclude_cols and col not in labor_required_cols]
            labor_cols = available_cols + other_cols
            labor_rows[labor_cols].to_excel(writer, index=False, sheet_name="人工费汇总")
        if not labor_month_rows.empty:
            labor_month_cols = [col for col in labor_month_rows.columns if col not in {"项目编码", "项目状态", "项目名称"}]
            labor_month_rows[labor_month_cols].to_excel(writer, index=False, sheet_name="人工费按年月项目汇总")
        if not finance_category_rows.empty:
            excluded_cols = {
                "项目编码", "项目状态", "扣除借款（泰铢）", "扣除借款（人民币）", "扣除预扣税（泰铢）", "扣除预扣税（人民币）", "公司实际付款（财务数据）（泰铢）", "公司实际付款（财务数据）（人民币）", "汇率",
            }
            finance_cols = [col for col in finance_category_rows.columns if col not in excluded_cols]
            finance_category_rows[finance_cols].to_excel(writer, index=False, sheet_name="财务分包付款明细")

        if not budget_df.empty:
            project_budget_rows = budget_df[budget_df["项目编码"] == project_code].copy()
            if not project_budget_rows.empty:
                budget_cols_needed = [
                    "序号", "项目编码", "项目名称", "币种", "标段人材机汇总", "人_明细", "管_明细", "设机_明细", "材类明细", "其他类别_明细", "钢筋", "混凝土", "砂浆", "碎石", "砖", "其他材料",
                    "人管设机材类明细其他类别_合计", "钢筋混凝土砂浆碎石砖其他材料_合计",
                ]
                available_cols = [col for col in budget_cols_needed if col in project_budget_rows.columns]
                df_project_budget = project_budget_rows[available_cols].copy()

                if not df_project_budget.empty:
                    df_project_budget.to_excel(writer, index=False, sheet_name="预算金额")
                    ws_budget = writer.sheets["预算金额"]
                    budget_amount_cols = [
                        "标段人材机汇总", "人_明细", "管_明细", "设机_明细", "材类明细", "其他类别_明细", "钢筋", "混凝土", "砂浆", "碎石", "砖", "其他材料",
                        "人管设机材类明细其他类别_合计", "钢筋混凝土砂浆碎石砖其他材料_合计",
                    ]
                    style_sheet(ws_budget, amount_headers=budget_amount_cols)

        amount_headers = [
            "金额_THB", "金额_含税优先", "合同金额", "含税金额", "不含税金额", "金额（万泰铢）", "金额（万泰铢）_THB", "计算总金额", "结算费用_泰铢", "本次费用合计", "月租金额_泰铢", "人工费汇总（泰铢）",
            "含税实付金额（万泰铢）", "国内代付金额（人民币）", "国内代付人民币金额（万泰铢）", "付款汇总（万泰铢）", "单价", "数量", "预计总成本", "收入", "毛利", "业主预计结算额", "分包结算额", "分包已付款",
            "材料采购实际已发生", "材料采购后续预估发生", "采购部采购金额", "项目部采购金额", "材料租赁费用", "材料租赁后续预估", "工资支出", "管理人员工资", "机具费支出", "机具费后续预估", "人员调遣支出",
            "人员调遣后续预估", "其他费用支出", "其他费用后续预估", "临建费用支出", "临建费用后续预估",
        ]

        if project_code in responsibility_cost_mapping:
            responsibility_file = responsibility_cost_mapping[project_code]
            responsibility_data = get_responsibility_cost_data(responsibility_file)

            if responsibility_data:
                actual_amounts = {}
                zhongfang_laowu = 0
                zigongsi_didi = 0
                laowu_gongsi = 0

                if not labor_rows.empty:
                    for col_name in ["中方劳务工铢", "求和项:中方劳务工铢"]:
                        if col_name in labor_rows.columns:
                            zhongfang_laowu = coerce_float(labor_rows[col_name].sum())
                            break

                    for col_name in ["子公司属地工铢", "求和项:子公司属地工铢"]:
                        if col_name in labor_rows.columns:
                            zigongsi_didi = coerce_float(labor_rows[col_name].sum())
                            break

                    for col_name in ["劳务公司铢", "求和项:劳务公司铢"]:
                        if col_name in labor_rows.columns:
                            laowu_gongsi = coerce_float(labor_rows[col_name].sum())
                            break

                actual_amounts["自主施工人工费"] = zhongfang_laowu + zigongsi_didi + laowu_gongsi

                if not equipment_rows.empty and "金额_THB" in equipment_rows.columns:
                    actual_amounts["自主施工机械费"] = coerce_float(equipment_rows["金额_THB"].sum())
                else:
                    actual_amounts["自主施工机械费"] = 0.0

                material_actual = coerce_float(cost_stats_row.get("材料采购实际已发生", 0) if cost_stats_row is not None else 0)
                caigou_bu = coerce_float(cost_stats_row.get("采购部采购金额", 0) if cost_stats_row is not None else 0)
                xiangmu_bu = coerce_float(cost_stats_row.get("项目部采购金额", 0) if cost_stats_row is not None else 0)
                zulin_feiyong = coerce_float(cost_stats_row.get("材料租赁费用", 0) if cost_stats_row is not None else 0)
                material_sum = caigou_bu + xiangmu_bu + zulin_feiyong
                actual_amounts["自主施工材料费"] = max(material_actual, material_sum)

                ws_responsibility = workbook.create_sheet("责任成本", 1)
                current_row = 1

                for sheet_idx, (_source_sheet_name, df) in enumerate(responsibility_data.items()):
                    if sheet_idx > 0:
                        current_row += 1

                    amount_col_idx = None
                    amount_col_name = None
                    for col in df.columns:
                        col_text = str(col)
                        col_lower = col_text.lower()
                        if ("金额" in col_text) or ("合价" in col_text) or ("amount" in col_lower) or ("price" in col_lower):
                            amount_col_name = col
                            break

                    if amount_col_name is not None:
                        amount_col_idx = df.columns.get_loc(amount_col_name)
                        df.insert(amount_col_idx + 1, "实际已发生金额", "")

                    for col_idx, col_name in enumerate(df.columns, start=1):
                        cell = ws_responsibility.cell(row=current_row, column=col_idx)
                        cell.value = col_name
                    current_row += 1

                    name_col_idx = None
                    for idx, col in enumerate(df.columns):
                        if "名称" in str(col):
                            name_col_idx = idx
                            break
                    if name_col_idx is None:
                        name_col_idx = 1

                    for row_idx, row_data in enumerate(df.itertuples(index=False, name=None), start=1):
                        row_list = list(row_data)
                        if amount_col_idx is not None:
                            if len(row_list) > name_col_idx:
                                name_value = str(row_list[name_col_idx])
                            else:
                                name_value = ""
                            name_value = re.sub(r"\s+", "", name_value)

                            matched_value = None
                            if "人工费" in name_value or ("人工" in name_value and "费" in name_value):
                                matched_value = actual_amounts.get("自主施工人工费")
                            elif "机械费" in name_value or "机具" in name_value or "设备" in name_value:
                                matched_value = actual_amounts.get("自主施工机械费")
                            elif "材料费" in name_value or "材料" in name_value:
                                matched_value = actual_amounts.get("自主施工材料费")

                            if matched_value is not None:
                                row_list[amount_col_idx + 1] = matched_value

                        for col_idx, value in enumerate(row_list, start=1):
                            cell = ws_responsibility.cell(row=current_row + row_idx - 1, column=col_idx)
                            cell.value = value

                    current_row += len(df)
            else:
                workbook.create_sheet("责任成本", 1)
        else:
            workbook.create_sheet("责任成本", 1)

        for sheet_name in writer.sheets:
            ws = writer.sheets[sheet_name]
            if sheet_name == "成本分析报告":
                style_analysis_sheet(ws)
            else:
                sheet_amount_headers = amount_headers.copy()

                if sheet_name == "项目成本统计汇总表":
                    sheet_amount_headers = ["金额"]
                elif sheet_name == "人工费汇总":
                    sheet_amount_headers = ["中方正式工铢", "中方劳务工铢", "子公司属地工铢", "劳务公司铢", "签证费", "人工费汇总（泰铢）"]
                elif sheet_name == "人工费按年月项目汇总":
                    sheet_amount_headers = [col for col in ws[1] if col.value and "年" not in str(col.value) and "月" not in str(col.value) and "编码" not in str(col.value) and "名称" not in str(col.value)]
                    sheet_amount_headers = [str(col.value) if col.value else "" for col in sheet_amount_headers]
                elif sheet_name == "财务分包付款明细":
                    sheet_amount_headers = [col for col in ws[1] if col.value and ("金额" in str(col.value) or "费" in str(col.value) or "实付" in str(col.value))]
                    sheet_amount_headers = [str(col.value) if col.value else "" for col in sheet_amount_headers]
                elif sheet_name == "责任成本":
                    sheet_amount_headers = []
                    for col in ws[1]:
                        col_text = str(col.value) if col.value else ""
                        if "金额" in col_text or "合价" in col_text or "实际已发生" in col_text:
                            sheet_amount_headers.append(col_text)

                style_sheet(ws, amount_headers=sheet_amount_headers)
        workbook.save(output_path)


def export_index_file(output_dir: Path, summary_rows: List[Dict], cost_stats_df: pd.DataFrame = None, budget_df: pd.DataFrame = None) -> None:
    if not summary_rows:
        return
    output_dir.mkdir(parents=True, exist_ok=True)
    index_path = output_dir / "项目汇总索引.xlsx"
    df = pd.DataFrame(summary_rows)

    if cost_stats_df is not None and not cost_stats_df.empty:
        metadata_cols = {"月份", "项目名称（合同管理）", "项目编码", "标准项目编码"}
        available_cost_cols = [
            col for col in cost_stats_df.columns if col == "匹配编码" or col not in metadata_cols
        ]
        cost_merge_df = cost_stats_df[available_cost_cols].copy()

        df = df.merge(
            cost_merge_df,
            left_on="项目编码",
            right_on="匹配编码",
            how="left",
            suffixes=("", "_成本统计"),
        )

        if "匹配编码" in df.columns:
            df = df.drop(columns=["匹配编码"])

    if "项目状态" in df.columns and "项目编码" in df.columns:
        df = df.sort_values(["项目状态", "项目编码"], ascending=[True, False])
    elif "项目状态" in df.columns:
        df = df.sort_values(["项目状态"])
    elif "项目编码" in df.columns:
        df = df.sort_values(["项目编码"], ascending=False)
    if "采购合同金额" in df.columns:
        df = df.drop(columns=["采购合同金额"])

    df_wan = df.copy()
    amount_columns = [
        "总包合同金额", "分包合同金额", "材料采购金额", "材料欠款金额", "商混站金额", "机具使用费", "人工费", "含税实付金额（万泰铢）", "国内代付人民币金额（万泰铢）", "付款汇总（万泰铢）", "利润额", "分包代付款",
        "预计总成本", "收入", "毛利", "业主预计结算额", "总包合同额(万泰铢)", "签证情况", "分包结算额", "分包合同额(万泰铢)", "分包签证情况", "分包已付款", "材料采购实际已发生", "材料采购后续预估发生",
        "采购部采购金额", "项目部采购金额", "材料租赁费用", "材料租赁后续预估", "工资支出", "管理人员工资", "机具费支出", "机具费后续预估", "人员调遣支出", "人员调遣后续预估", "其他费用支出", "实际已发生金额",
        "其他费用后续预估", "临建费用支出", "临建费用后续预估",
    ]

    for col in df_wan.columns:
        col_text = str(col)
        if ("万泰铢" in col_text or "付款财务数据" in col_text) and col_text not in amount_columns:
            amount_columns.append(col_text)

    for col in df_wan.columns:
        if col.startswith("预算_"):
            amount_columns.append(col)

    for col in amount_columns:
        if col in df_wan.columns:
            df_wan[col] = df_wan[col] / 10000

    with pd.ExcelWriter(index_path, engine="openpyxl") as writer:
        df_overview = df.copy()
        budget_cols = [col for col in df_overview.columns if col.startswith("预算_")]
        if budget_cols:
            df_overview = df_overview.drop(columns=budget_cols)
        if "预算币种" in df_overview.columns:
            df_overview = df_overview.drop(columns=["预算币种"])

        df_overview.to_excel(writer, index=False, sheet_name="项目概览")
        ws = writer.sheets["项目概览"]
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

        headers = [c.value for c in ws[1]]
        total_contract_col = headers.index("总包合同金额") + 1 if "总包合同金额" in headers else None
        sub_contract_col = headers.index("分包合同金额") + 1 if "分包合同金额" in headers else None
        material_col = headers.index("材料采购金额") + 1 if "材料采购金额" in headers else None
        concrete_col = headers.index("商混站金额") + 1 if "商混站金额" in headers else None
        equipment_col = headers.index("机具使用费") + 1 if "机具使用费" in headers else None
        labor_col = headers.index("人工费") + 1 if "人工费" in headers else None
        profit_col = headers.index("利润额") + 1 if "利润额" in headers else None
        sub_payment_col = headers.index("分包代付款") + 1 if "分包代付款" in headers else None
        payment_summary_col = headers.index("付款汇总（万泰铢）") + 1 if "付款汇总（万泰铢）" in headers else None

        for row in range(2, ws.max_row + 1):
            if profit_col and total_contract_col and sub_contract_col and material_col and concrete_col and equipment_col and labor_col:
                profit_formula = f"={get_column_letter(total_contract_col)}{row}-({get_column_letter(sub_contract_col)}{row}+{get_column_letter(material_col)}{row}+{get_column_letter(concrete_col)}{row}+{get_column_letter(equipment_col)}{row}+{get_column_letter(labor_col)}{row})"
                ws[f"{get_column_letter(profit_col)}{row}"] = profit_formula

            if sub_payment_col and sub_contract_col and payment_summary_col:
                sub_payment_formula = f"={get_column_letter(sub_contract_col)}{row}-{get_column_letter(payment_summary_col)}{row}"
                ws[f"{get_column_letter(sub_payment_col)}{row}"] = sub_payment_formula

        style_sheet(ws, amount_headers=[
            "总包合同金额", "分包合同金额", "材料采购金额", "材料欠款金额", "商混站金额", "机具使用费", "人工费", "含税实付金额（万泰铢）", "国内代付金额（人民币）", "国内代付人民币金额（万泰铢）", "付款汇总（万泰铢）", "利润额", "分包代付款",
            "预计总成本", "收入", "毛利", "业主预计结算额", "总包合同额(万泰铢)", "签证情况", "分包结算额", "分包合同额(万泰铢)", "分包签证情况", "分包已付款", "材料采购实际已发生", "材料采购后续预估发生", "采购部采购金额", "项目部采购金额",
            "材料租赁费用", "材料租赁后续预估", "工资支出", "管理人员工资", "机具费支出", "机具费后续预估", "人员调遣支出", "人员调遣后续预估", "其他费用支出", "实际已发生金额", "其他费用后续预估", "临建费用支出", "临建费用后续预估",
        ])

        df_wan.to_excel(writer, index=False, sheet_name="万元汇总")
        ws_wan = writer.sheets["万元汇总"]
        ws_wan.auto_filter.ref = f"A1:{get_column_letter(ws_wan.max_column)}{ws_wan.max_row}"

        headers_wan = [c.value for c in ws_wan[1]]
        total_contract_col_wan = headers_wan.index("总包合同金额") + 1 if "总包合同金额" in headers_wan else None
        sub_contract_col_wan = headers_wan.index("分包合同金额") + 1 if "分包合同金额" in headers_wan else None
        material_col_wan = headers_wan.index("材料采购金额") + 1 if "材料采购金额" in headers_wan else None
        concrete_col_wan = headers_wan.index("商混站金额") + 1 if "商混站金额" in headers_wan else None
        equipment_col_wan = headers_wan.index("机具使用费") + 1 if "机具使用费" in headers_wan else None
        labor_col_wan = headers_wan.index("人工费") + 1 if "人工费" in headers_wan else None
        profit_col_wan = headers_wan.index("利润额") + 1 if "利润额" in headers_wan else None
        sub_payment_col_wan = headers_wan.index("分包代付款") + 1 if "分包代付款" in headers_wan else None
        payment_summary_col_wan = headers_wan.index("付款汇总（万泰铢）") + 1 if "付款汇总（万泰铢）" in headers_wan else None

        for row in range(2, ws_wan.max_row + 1):
            if profit_col_wan and total_contract_col_wan and sub_contract_col_wan and material_col_wan and concrete_col_wan and equipment_col_wan and labor_col_wan:
                profit_formula_wan = f"={get_column_letter(total_contract_col_wan)}{row}-({get_column_letter(sub_contract_col_wan)}{row}+{get_column_letter(material_col_wan)}{row}+{get_column_letter(concrete_col_wan)}{row}+{get_column_letter(equipment_col_wan)}{row}+{get_column_letter(labor_col_wan)}{row})"
                ws_wan[f"{get_column_letter(profit_col_wan)}{row}"] = profit_formula_wan

            if sub_payment_col_wan and sub_contract_col_wan and payment_summary_col_wan:
                sub_payment_formula_wan = f"={get_column_letter(sub_contract_col_wan)}{row}-{get_column_letter(payment_summary_col_wan)}{row}"
                ws_wan[f"{get_column_letter(sub_payment_col_wan)}{row}"] = sub_payment_formula_wan

        style_sheet(ws_wan, amount_headers=[
            "总包合同金额", "分包合同金额", "材料采购金额", "材料欠款金额", "商混站金额", "机具使用费", "人工费", "含税实付金额（万泰铢）", "国内代付金额（人民币）", "国内代付人民币金额（万泰铢）", "付款汇总（万泰铢）", "利润额", "分包代付款",
            "预计总成本", "收入", "毛利", "业主预计结算额", "总包合同额(万泰铢)", "签证情况", "分包结算额", "分包合同额(万泰铢)", "分包签证情况", "分包已付款", "材料采购实际已发生", "材料采购后续预估发生", "采购部采购金额", "项目部采购金额",
            "材料租赁费用", "材料租赁后续预估", "工资支出", "管理人员工资", "机具费支出", "机具费后续预估", "人员调遣支出", "人员调遣后续预估", "其他费用支出", "实际已发生金额", "其他费用后续预估", "临建费用支出", "临建费用后续预估",
        ])

        budget_cols_needed = [
            "序号", "项目编码", "项目名称", "币种", "标段人材机汇总", "人_明细", "管_明细", "设机_明细", "材类明细", "其他类别_明细", "钢筋", "混凝土", "砂浆", "碎石", "砖", "其他材料",
            "人管设机材类明细其他类别_合计", "钢筋混凝土砂浆碎石砖其他材料_合计",
        ]

        if budget_df is not None and not budget_df.empty:
            available_cols = [col for col in budget_cols_needed if col in budget_df.columns]
            df_budget = budget_df[available_cols].copy()
            if not df_budget.empty:
                df_budget.to_excel(writer, index=False, sheet_name="预算金额")
                ws_budget = writer.sheets["预算金额"]
                ws_budget.auto_filter.ref = f"A1:{get_column_letter(ws_budget.max_column)}{ws_budget.max_row}"
                budget_amount_cols = [
                    "标段人材机汇总", "人_明细", "管_明细", "设机_明细", "材类明细", "其他类别_明细", "钢筋", "混凝土", "砂浆", "碎石", "砖", "其他材料",
                    "人管设机材类明细其他类别_合计", "钢筋混凝土砂浆碎石砖其他材料_合计",
                ]
                style_sheet(ws_budget, amount_headers=budget_amount_cols)

        writer.book.save(index_path)


def cleanup_legacy_output_files(
    output_root: Path,
    project_reports_dir: Path,
    summary_rows: List[Dict],
    project_manager_dirname: str,
    project_status_dirname: str,
) -> None:
    root_only_filenames = {
        "项目汇总索引.xlsx",
        "数据来源与字段映射.xlsx",
    }

    # 根目录专属文件：保留在 output_root，并清理项目合同汇总下的历史副本
    for filename in root_only_filenames:
        nested_path = project_reports_dir / filename
        if nested_path.exists() and nested_path.is_file():
            nested_path.unlink()

    managed_filenames = set()

    for row in summary_rows:
        code = str(row.get("项目编码", "")).strip()
        project_name = str(row.get("项目名称", "")).strip()
        if not code:
            continue
        managed_filenames.add(sanitize_filename(f"{code}_{project_name}" if project_name else code) + ".xlsx")

    legacy_dir_mapping = {
        project_manager_dirname: output_root / project_manager_dirname,
        project_status_dirname: output_root / project_status_dirname,
    }

    for dir_name, target_root in legacy_dir_mapping.items():
        legacy_dir = project_reports_dir / dir_name
        if not legacy_dir.is_dir():
            continue
        target_root.mkdir(parents=True, exist_ok=True)
        for item in legacy_dir.iterdir():
            target_path = target_root / item.name
            if item.is_dir():
                shutil.copytree(item, target_path, dirs_exist_ok=True)
                shutil.rmtree(item)
            else:
                if target_path.exists():
                    item.unlink()
                else:
                    shutil.move(str(item), str(target_path))
        shutil.rmtree(legacy_dir, ignore_errors=True)


def replicate_outputs(
    project_reports_dir: Path,
    output_root: Path,
    summary_rows: List[Dict],
    manager_map: Dict[str, str],
    project_manager_dirname: str,
    project_status_dirname: str,
) -> None:
    status_root = output_root / project_status_dirname
    manager_root = output_root / project_manager_dirname
    status_root.mkdir(parents=True, exist_ok=True)
    manager_root.mkdir(parents=True, exist_ok=True)

    for row in summary_rows:
        code = str(row.get("项目编码", "")).strip()
        project_name = str(row.get("项目名称", "")).strip()
        status = str(row.get("项目状态", "其他")).strip() or "其他"
        manager = manager_map.get(code, "其他") or "其他"

        filename = sanitize_filename(f"{code}_{project_name}" if project_name else code) + ".xlsx"
        src = project_reports_dir / filename
        if not src.exists():
            continue

        dst_status_dir = status_root / status
        dst_status_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst_status_dir / filename)

        dst_manager_dir = manager_root / manager
        dst_manager_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst_manager_dir / filename)
