#!/usr/bin/env python3
from __future__ import annotations

import argparse
import html
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


@dataclass
class ProjectMonthlyRecord:
    month: str
    settlement: Optional[float]
    gross_profit: Optional[float]
    gross_margin: Optional[float]
    completion_ratio: Optional[float]
    remark: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="从1_成本报公司月报Excel提取项目关键指标并生成HTML趋势报告"
    )
    parser.add_argument(
        "--input-dir",
        default="/Users/wyattsun/Desktop/邯郸中材/01成本统计/1_成本报公司",
        help="月报Excel目录",
    )
    parser.add_argument(
        "--output-html",
        default="/Users/wyattsun/Desktop/邯郸中材/01成本统计/1_成本报公司/项目月度进度汇总.html",
        help="输出HTML路径",
    )
    return parser.parse_args()


def month_from_filename(file_name: str) -> Optional[str]:
    m = re.search(r"(20\d{2})\s*年\s*(\d{1,2})\s*月", file_name)
    if not m:
        return None
    year = int(m.group(1))
    month = int(m.group(2))
    if 1 <= month <= 12:
        return f"{year:04d}-{month:02d}"
    return None


def as_float(value: object, is_ratio: bool = False) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        num = float(value)
    else:
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        percent_like = text.endswith("%")
        text = text.replace("%", "")
        try:
            num = float(text)
        except ValueError:
            return None
        if percent_like:
            num = num / 100.0

    if is_ratio and num > 1.0:
        num = num / 100.0
    return num


def normalize_header(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", "").replace(" ", "").strip()


def normalize_project_name(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text


def find_required_columns(header_row: List[object]) -> Dict[str, int]:
    keys = {
        "project": "项目名称",
        "settlement": "业主预计结算额",
        "gross_profit": "预估毛利",
        "gross_margin": "预计毛利率",
        "completion_ratio": "完成比例",
    }
    found: Dict[str, int] = {}
    for idx, val in enumerate(header_row, start=1):
        hv = normalize_header(val)
        for key, keyword in keys.items():
            if key in found:
                continue
            if keyword in hv:
                found[key] = idx

    missing = [k for k in keys if k not in found]
    if missing:
        raise ValueError(f"缺少必要列: {missing}")
    return found


def find_optional_column(header_row: List[object], keyword: str) -> Optional[int]:
    for idx, val in enumerate(header_row, start=1):
        hv = normalize_header(val)
        if keyword in hv:
            return idx
    return None


def read_monthly_file(file_path: Path, month: str) -> Dict[str, ProjectMonthlyRecord]:
    wb = load_workbook(file_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    max_col = ws.max_column
    header_row = [ws.cell(2, c).value for c in range(1, max_col + 1)]
    col = find_required_columns(header_row)
    remark_col = find_optional_column(header_row, "备注")
    progress_col = find_optional_column(header_row, "进度概述")

    month_data: Dict[str, ProjectMonthlyRecord] = {}
    for r in range(3, ws.max_row + 1):
        project = normalize_project_name(ws.cell(r, col["project"]).value)
        if not project:
            continue
        if project in {"合计", "总计", "小计", "序号"}:
            continue

        month_data[project] = ProjectMonthlyRecord(
            month=month,
            settlement=as_float(ws.cell(r, col["settlement"]).value),
            gross_profit=as_float(ws.cell(r, col["gross_profit"]).value),
            gross_margin=as_float(ws.cell(r, col["gross_margin"]).value, is_ratio=True),
            completion_ratio=as_float(ws.cell(r, col["completion_ratio"]).value, is_ratio=True),
          remark=(
            "；".join(
              [
                s
                for s in [
                  normalize_text(ws.cell(r, remark_col).value) if remark_col else "",
                  normalize_text(ws.cell(r, progress_col).value) if progress_col else "",
                ]
                if s
              ]
            )
          ),
        )
    wb.close()
    return month_data


def fmt_number(val: Optional[float], digits: int = 2) -> str:
    if val is None:
        return "-"
    return f"{val:,.{digits}f}"


def fmt_percent(val: Optional[float]) -> str:
    if val is None:
        return "-"
    return f"{val * 100:.2f}%"


def fmt_text(val: str) -> str:
  return html.escape(val) if val else "-"


def make_html(project_data: Dict[str, List[ProjectMonthlyRecord]], months: List[str]) -> str:
    sections: List[str] = []
    chart_data: Dict[str, Dict[str, List[Optional[float]]]] = {}

    def missing_count(project: str) -> int:
        records_by_month = {r.month: r for r in project_data[project]}
        miss = 0
        for m in months:
            rec = records_by_month.get(m)
            if rec is None:
                miss += 4
                continue
            miss += int(rec.settlement is None)
            miss += int(rec.gross_profit is None)
            miss += int(rec.gross_margin is None)
            miss += int(rec.completion_ratio is None)
        return miss

    ordered_projects = sorted(project_data.keys(), key=lambda p: (missing_count(p) > 0, missing_count(p), p))

    for idx, project in enumerate(ordered_projects, start=1):
        records_by_month = {r.month: r for r in project_data[project]}

        rows: List[str] = []
        settlement_series: List[Optional[float]] = []
        gross_profit_series: List[Optional[float]] = []
        gross_margin_series: List[Optional[float]] = []
        completion_series: List[Optional[float]] = []
        remark_series: List[str] = []

        for month in months:
            rec = records_by_month.get(month)
            settlement_series.append(None if rec is None else rec.settlement)
            gross_profit_series.append(None if rec is None else rec.gross_profit)
            gross_margin_series.append(None if rec is None else rec.gross_margin)
            completion_series.append(None if rec is None else rec.completion_ratio)
            remark_series.append("" if rec is None else rec.remark)

            rows.append(
                "<tr>"
                f"<td class='metric-cell'>{month}</td>"
                f"<td class='num metric-cell'>{fmt_number(None if rec is None else rec.settlement)}</td>"
                f"<td class='num metric-cell'>{fmt_number(None if rec is None else rec.gross_profit)}</td>"
                f"<td class='num metric-cell'>{fmt_percent(None if rec is None else rec.gross_margin)}</td>"
                f"<td class='num metric-cell'>{fmt_percent(None if rec is None else rec.completion_ratio)}</td>"
              f"<td class='remark-cell'>{fmt_text('' if rec is None else rec.remark)}</td>"
                "</tr>"
            )

        chart_key = f"project_{idx}"
        chart_data[chart_key] = {
            "settlement": settlement_series,
            "gross_profit": gross_profit_series,
            "gross_margin": gross_margin_series,
            "completion_ratio": completion_series,
            "remark": remark_series,
        }

        chart_amount_id = f"chart_amount_{idx}"
        chart_rate_id = f"chart_rate_{idx}"
        card_id = f"card_{idx}"

        section = (
            f"<section class='project-card' id='{card_id}'>"
            f"<div class='card-header'><h2>{html.escape(project)}</h2><button class='export-btn' onclick='exportCardPDF(\"{ card_id}\", \"{html.escape(project)}\")'>📥 导出PDF</button></div>"
            "<div class='chart-grid'>"
            "<div class='chart-wrap'><h3>金额趋势</h3><canvas id='" + chart_amount_id + "'></canvas></div>"
            "<div class='chart-wrap'><h3>比例趋势</h3><canvas id='" + chart_rate_id + "'></canvas></div>"
            "</div>"
            "<table>"
            "<thead><tr>"
            "<th>月份</th><th>业主预计结算额</th><th>预估毛利</th><th>预计毛利率</th><th>完成比例</th><th>备注</th>"
            "</tr></thead>"
            "<tbody>"
            + "".join(rows)
            + "</tbody></table></section>"
        )
        sections.append(section)

    export_all_btn = (
        "<div style='text-align: center; margin: 20px 0;'>"
        "<button class='export-all-btn' onclick='exportAllCardsPDF()'>📦 导出所有项目</button>"
        "</div>"
    )
    all_sections = export_all_btn + ''.join(sections)

    months_json = json.dumps(months, ensure_ascii=False)
    chart_data_json = json.dumps(chart_data, ensure_ascii=False)

    return f"""<!doctype html>
<html lang=\"zh-CN\">
<head>
  <meta charset=\"utf-8\" />
  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1\" />
  <title>项目月度进度汇总</title>
  <script src=\"https://cdn.jsdelivr.net/npm/chart.js@4.4.3/dist/chart.umd.min.js\"></script>
  <script src=\"https://cdn.jsdelivr.net/npm/html2pdf.js@0.10.1/dist/html2pdf.bundle.min.js\"></script>
  <style>
    :root {{
      --bg1: #f3f7f2;
      --bg2: #f9efe4;
      --card: rgba(255,255,255,0.88);
      --line: #dde6dc;
      --txt: #21322a;
      --muted: #5e7167;
      --accent: #c46a2f;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      color: var(--txt);
      font-family: "PingFang SC", "Noto Sans SC", "Microsoft YaHei", sans-serif;
      background:
        radial-gradient(circle at 20% 0%, #e8f2e3 0%, transparent 40%),
        radial-gradient(circle at 80% 0%, #f8e6d6 0%, transparent 35%),
        linear-gradient(160deg, var(--bg1), var(--bg2));
      min-height: 100vh;
    }}
    .container {{
      width: min(1200px, 95vw);
      margin: 28px auto 48px;
    }}
    .hero {{
      background: var(--card);
      border: 1px solid var(--line);
      border-radius: 18px;
      padding: 18px 22px;
      box-shadow: 0 12px 24px rgba(36, 48, 41, 0.08);
      margin-bottom: 16px;
    }}
    h1 {{ margin: 4px 0 8px; font-size: clamp(22px, 3vw, 32px); }}
    .hint {{ color: var(--muted); margin: 0; }}
    .project-card {{
      background: var(--card);
      border: 1px solid var(--line);
      border-radius: 16px;
      padding: 18px;
      margin: 14px 0;
      box-shadow: 0 10px 20px rgba(36, 48, 41, 0.06);
    }}
    h2 {{ margin: 4px 0 10px; font-size: clamp(18px, 2.2vw, 24px); }}
    h3 {{ margin: 0 0 8px; font-size: 14px; color: #2f4338; }}
    .chart-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 12px; margin-bottom: 32px; }}
    .chart-wrap {{
      width: 100%;
      height: 300px;
      margin-bottom: 0;
    }}
    table {{ width: 100%; border-collapse: collapse; font-size: 13px; margin-top: 12px; }}
    th, td {{ border-bottom: 1px solid var(--line); padding: 8px 6px; text-align: center; vertical-align: middle; }}
    th {{ color: #31453a; background: rgba(227, 237, 230, 0.55); }}
    td.num {{ text-align: center; font-variant-numeric: tabular-nums; }}
    td.metric-cell {{ font-size: 12px; }}
    td.remark-cell {{ font-size: 12px; line-height: 1.45; }}
    .card-header {{ display: flex; justify-content: space-between; align-items: center; margin-bottom: 10px; }}
    .card-header h2 {{ margin: 4px 0; flex: 1; }}
    .export-btn, .export-all-btn {{
      padding: 6px 12px;
      font-size: 12px;
      background: #2b8a6e;
      color: white;
      border: none;
      border-radius: 6px;
      cursor: pointer;
      white-space: nowrap;
      transition: background 0.2s;
    }}
    .export-btn:hover, .export-all-btn:hover {{
      background: #1f6348;
    }}
    .export-btn:active, .export-all-btn:active {{
      transform: scale(0.98);
    }}
    @media (max-width: 768px) {{
      .chart-grid {{ grid-template-columns: 1fr; }}
      .chart-wrap {{ height: 260px; }}
      .chart-grid {{ margin-bottom: 24px; }}
      table {{ font-size: 12px; }}
      td.metric-cell {{ font-size: 11px; }}
      td.remark-cell {{ font-size: 11px; }}
      th, td {{ padding: 6px 4px; }}
    }}
  </style>
</head>
<body>
  <div class=\"container\">
    <section class=\"hero\">
      <h1>项目月度进度汇总</h1>
      <p class=\"hint\"></p>
    </section>
    {all_sections}
  </div>

  <script>
    const months = {months_json};
    const allData = {chart_data_json};
    const amountFormatter = (v) => v == null ? '-' : Number(v).toLocaleString('zh-CN', {{maximumFractionDigits: 2}});
    const percentTick = (v) => (v * 100).toFixed(0) + '%';
    const labelFormatter = (datasetLabel, v) => {{
      if (v == null) return '';
      if (datasetLabel.includes('率') || datasetLabel.includes('比例')) {{
        return (v * 100).toFixed(1) + '%';
      }}
      return Number(v).toLocaleString('zh-CN', {{maximumFractionDigits: 0}});
    }};

    const inlineDataLabelPlugin = {{
      id: 'inlineDataLabelPlugin',
      afterDatasetsDraw(chart) {{
        const ctx = chart.ctx;
        chart.data.datasets.forEach((dataset, datasetIndex) => {{
          const meta = chart.getDatasetMeta(datasetIndex);
          if (meta.hidden) return;

          meta.data.forEach((element, index) => {{
            const raw = dataset.data[index];
            if (raw == null) return;

            const text = labelFormatter(dataset.label || '', raw);
            if (!text) return;

            const pos = element.tooltipPosition();
            ctx.save();
            ctx.fillStyle = '#24352d';
            ctx.font = '600 10px "PingFang SC", "Noto Sans SC", "Microsoft YaHei", sans-serif';
            ctx.textAlign = 'center';
            ctx.textBaseline = 'bottom';
            ctx.fillText(text, pos.x, pos.y - 6);
            ctx.restore();
          }});
        }});
      }}
    }};

    Chart.register(inlineDataLabelPlugin);

    const computeRateBounds = (a, b) => {{
      const values = [...a, ...b].filter(v => v != null);
      if (!values.length) return {{ min: 0, max: 1 }};
      let min = Math.min(...values);
      let max = Math.max(...values);
      if (min === max) {{
        min -= 0.05;
        max += 0.05;
      }}
      const pad = Math.max((max - min) * 0.15, 0.03);
      return {{ min: min - pad, max: max + pad }};
    }};

    const computeAmountBounds = (a, b) => {{
      const values = [...a, ...b].filter(v => v != null);
      if (!values.length) return {{ min: 0, max: 1 }};
      const minVal = Math.min(...values);
      const maxVal = Math.max(...values);

      // 全为正值时从0起，保留顶部空间；存在负值时两端都留白。
      if (minVal >= 0) {{
        return {{ min: 0, max: maxVal * 1.15 + 1 }};
      }}

      const span = maxVal - minVal;
      const pad = Math.max(span * 0.12, 1);
      return {{ min: minVal - pad, max: maxVal + pad }};
    }};

    Object.entries(allData).forEach(([chartKey, series]) => {{
      const idx = chartKey.split('_')[1];
      const amountEl = document.getElementById(`chart_amount_${{idx}}`);
      const rateEl = document.getElementById(`chart_rate_${{idx}}`);
      if (!amountEl || !rateEl) return;
      const amountBounds = computeAmountBounds(series.settlement, series.gross_profit);

      new Chart(amountEl, {{
        type: 'bar',
        data: {{
          labels: months,
          datasets: [
            {{
              type: 'bar',
              label: '业主预计结算额',
              data: series.settlement,
              yAxisID: 'yAmount',
              borderColor: '#2b8a6e',
              backgroundColor: 'rgba(43,138,110,0.75)',
              borderWidth: 1,
              order: 3,
              spanGaps: true,
            }},
            {{
              type: 'bar',
              label: '预估毛利',
              data: series.gross_profit,
              yAxisID: 'yAmount',
              borderColor: '#d9782f',
              backgroundColor: 'rgba(217,120,47,0.78)',
              borderWidth: 1,
              order: 3,
              spanGaps: true,
            }}
          ]
        }},
        options: {{
          responsive: true,
          maintainAspectRatio: false,
          interaction: {{ mode: 'index', intersect: false }},
          plugins: {{
            legend: {{ position: 'top' }},
            inlineDataLabelPlugin: {{}},
            tooltip: {{
              callbacks: {{
                label: (ctx) => {{
                  const label = ctx.dataset.label || '';
                  const val = ctx.raw;
                  if (label.includes('率') || label.includes('比例')) {{
                    return `${{label}}: ${{val == null ? '-' : (val * 100).toFixed(2) + '%'}}`;
                  }}
                  return `${{label}}: ${{amountFormatter(val)}}`;
                }}
              }}
            }}
          }},
          scales: {{
            yAmount: {{
              type: 'linear',
              position: 'left',
              min: amountBounds.min,
              max: amountBounds.max,
              ticks: {{ callback: (v) => amountFormatter(v) }}
            }}
          }}
        }}
      }});

      const rateBounds = computeRateBounds(series.gross_margin, series.completion_ratio);
      new Chart(rateEl, {{
        type: 'line',
        data: {{
          labels: months,
          datasets: [
            {{
              type: 'line',
              label: '预计毛利率',
              data: series.gross_margin,
              yAxisID: 'yRate',
              borderColor: '#3f6fcf',
              backgroundColor: 'rgba(63,111,207,0.15)',
              pointRadius: 3,
              pointHoverRadius: 4,
              borderWidth: 2,
              spanGaps: true,
              tension: 0.25,
            }},
            {{
              type: 'line',
              label: '完成比例',
              data: series.completion_ratio,
              yAxisID: 'yRate',
              borderColor: '#8e4db4',
              backgroundColor: 'rgba(142,77,180,0.15)',
              pointRadius: 3,
              pointHoverRadius: 4,
              borderWidth: 2,
              spanGaps: true,
              tension: 0.25,
            }}
          ]
        }},
        options: {{
          responsive: true,
          maintainAspectRatio: false,
          interaction: {{ mode: 'index', intersect: false }},
          plugins: {{
            legend: {{ position: 'top' }},
            inlineDataLabelPlugin: {{}},
            tooltip: {{
              callbacks: {{
                label: (ctx) => {{
                  const label = ctx.dataset.label || '';
                  const val = ctx.raw;
                  return `${{label}}: ${{val == null ? '-' : (val * 100).toFixed(2) + '%'}}`;
                }}
              }}
            }}
          }},
          scales: {{
            yRate: {{
              type: 'linear',
              position: 'left',
              min: rateBounds.min,
              max: rateBounds.max,
              ticks: {{ callback: (v) => percentTick(v) }}
            }}
          }}
        }}
      }});
    }});

    async function exportCardPDF(cardId, projectName) {{
      const element = document.getElementById(cardId);
      if (!element) return;
      
      // 隐藏导出按钮
      const btns = element.querySelectorAll('.export-btn');
      btns.forEach(btn => btn.style.display = 'none');
      
      try {{
        const canvas = await html2canvas(element, {{
          scale: 3,
          useCORS: true,
          logging: false,
          backgroundColor: '#ffffff'
        }});
        
        const imgData = canvas.toDataURL('image/jpeg', 0.98);
        const pdf = new jsPDF({{
          orientation: 'portrait',
          unit: 'mm',
          format: 'a4'
        }});
        
        const pageWidth = pdf.internal.pageSize.getWidth();
        const margin = 10;
        const contentWidth = pageWidth - 2 * margin;
        const contentHeight = (canvas.height * contentWidth) / canvas.width;
        
        pdf.addImage(imgData, 'JPEG', margin, margin, contentWidth, contentHeight);
        pdf.save(projectName + '.pdf');
      }} catch(err) {{
        console.error('PDF导出失败:', err);
        alert('PDF导出失败，请稍后重试');
      }} finally {{
        btns.forEach(btn => btn.style.display = 'inline-block');
      }}
    }}

    async function exportAllCardsPDF() {{
      const cards = document.querySelectorAll('.project-card');
      if (!cards.length) return;
      
      // 隐藏所有导出按钮
      document.querySelectorAll('.export-btn, .export-all-btn').forEach(btn => {{
        btn.style.display = 'none';
      }});
      
      const opt = {{
        margin: 10,
        image: {{ type: 'jpeg', quality: 0.98 }},
        html2canvas: {{ scale: 2, useCORS: true, logging: false }},
        jsPDF: {{ orientation: 'portrait', unit: 'mm', format: 'a4' }}
      }};
      
      const doc = new jsPDF(opt.jsPDF);
      let isFirstPage = true;
      
      for (let i = 0; i < cards.length; i++) {{
        const card = cards[i];
        const canvas = await html2canvas(card, opt.html2canvas);
        const imgData = canvas.toDataURL('image/jpeg', opt.image.quality);
        const imgWidth = 190; // A4宽度-边距
        const imgHeight = (canvas.height * imgWidth) / canvas.width;
        
        if (!isFirstPage) {{
          doc.addPage();
        }}
        doc.addImage(imgData, 'JPEG', 10, 10, imgWidth, imgHeight);
        isFirstPage = false;
      }}
      
      doc.save('\u9879\u76ee\u6708\u5ea6\u8fdb\u5ea6\u6c47\u603b.pdf');
      
      // 恢复按钮显示
      document.querySelectorAll('.export-btn, .export-all-btn').forEach(btn => {{
        btn.style.display = 'inline-block';
      }});
    }}
  </script>
</body>
</html>
"""


def main() -> None:
    args = parse_args()
    input_dir = Path(args.input_dir)
    output_html = Path(args.output_html)

    if not input_dir.exists():
        raise FileNotFoundError(f"输入目录不存在: {input_dir}")

    excel_files: List[Tuple[str, Path]] = []
    for p in sorted(input_dir.glob("*.xls*")):
        month = month_from_filename(p.name)
        if month is None:
            continue
        excel_files.append((month, p))

    if not excel_files:
        raise RuntimeError("未找到可解析月份的Excel文件")

    excel_files.sort(key=lambda x: x[0])
    months = [m for m, _ in excel_files]

    project_data: Dict[str, List[ProjectMonthlyRecord]] = {}
    for month, file_path in excel_files:
        monthly_rows = read_monthly_file(file_path, month)
        for project, record in monthly_rows.items():
            project_data.setdefault(project, []).append(record)

    html_content = make_html(project_data, months)
    output_html.parent.mkdir(parents=True, exist_ok=True)
    output_html.write_text(html_content, encoding="utf-8")

    print(f"已生成HTML: {output_html}")
    print(f"月份数量: {len(months)}")
    print(f"项目数量: {len(project_data)}")


if __name__ == "__main__":
    main()
