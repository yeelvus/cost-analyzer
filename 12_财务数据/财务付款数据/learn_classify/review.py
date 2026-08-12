# -*- coding: utf-8 -*-
"""待复核导出 / 批量重分类 std 文件。"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook

from .engine import ClassifyEngine, ClassifyResult


REVIEW_COLUMNS = [
    "文件",
    "项目名称",
    "科目",
    "摘要",
    "供应商名称",
    "实付金额",
    "一级科目",
    "二级科目",
    "分类ID",
    "命中层",
    "置信度",
    "原因",
    "正确一级科目",
    "正确二级科目",
    "备注",
]


def classify_dataframe(
    df: pd.DataFrame,
    engine: ClassifyEngine,
    *,
    subject_col: str = "科目",
    summary_col: str = "摘要",
    vendor_col: str = "供应商名称",
) -> pd.DataFrame:
    """批量分类 DataFrame；相同 (科目,摘要,供应商) 只算一次。"""
    out = df.copy()
    n = len(out)

    def col_vals(name: str) -> List[Any]:
        if name not in out.columns:
            return [None] * n
        return out[name].tolist()

    triples = list(
        zip(col_vals(subject_col), col_vals(summary_col), col_vals(vendor_col))
    )
    results = engine.classify_many(triples)

    out["一级科目"] = [r.l1_name for r in results]
    out["二级科目"] = [r.l2_name for r in results]
    out["三级科目"] = [getattr(r, "l3_name", "") or "" for r in results]
    out["分类ID"] = [
        (getattr(r, "l3_id", None) or r.l2_id or "") for r in results
    ]
    out["命中层"] = [r.layer for r in results]
    out["置信度"] = [r.confidence for r in results]
    out["原因"] = [r.reason for r in results]
    out["一级科目ID"] = [r.l1_id for r in results]
    return out


def export_review_pool(
    std_dir: Path,
    engine: ClassifyEngine,
    out_path: Path,
    *,
    conf_threshold: float = 0.55,
    only_review: bool = True,
    max_rows: Optional[int] = None,
) -> Path:
    """从 B_财务付款_std 扫描，导出待复核表（含空的正确列供填写）。"""
    std_dir = Path(std_dir)
    rows: List[Dict[str, Any]] = []

    for f in sorted(std_dir.rglob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        try:
            df = pd.read_excel(f, sheet_name=0)
        except Exception:
            continue
        cols = {str(c).strip(): c for c in df.columns}
        c_subj = cols.get("科目")
        c_summ = cols.get("摘要")
        c_vend = cols.get("供应商名称") or cols.get("供应商")
        c_amt = cols.get("实付金额")
        c_proj = cols.get("项目名称")

        for _, row in df.iterrows():
            subj = "" if not c_subj or pd.isna(row.get(c_subj)) else str(row.get(c_subj)).strip()
            summ = "" if not c_summ or pd.isna(row.get(c_summ)) else str(row.get(c_summ)).strip()
            vend = "" if not c_vend or pd.isna(row.get(c_vend)) else str(row.get(c_vend)).strip()
            amt = None
            if c_amt is not None:
                try:
                    from payment_hub.utils import has_actual_payment

                    raw_amt = row.get(c_amt)
                    if not has_actual_payment(raw_amt):
                        continue  # 空实付不进待复核/学习
                    amt = float(raw_amt)
                except Exception:
                    raw_amt = row.get(c_amt) if c_amt is not None else None
                    if raw_amt is None or (isinstance(raw_amt, float) and pd.isna(raw_amt)):
                        continue
                    if str(raw_amt).strip() == "":
                        continue
                    try:
                        amt = float(raw_amt or 0)
                    except Exception:
                        amt = raw_amt
            elif c_amt is None:
                # 无实付列则不导出（口径：只处理有实付金额）
                continue
            r = engine.classify(subj, summ, vend)
            if only_review and not (r.needs_review or r.confidence < conf_threshold or r.layer == "pending"):
                # 也导出「其他」类便于清洗
                if r.l2_name not in {"其他", "待复核", "(空白)"}:
                    continue
            proj = "" if not c_proj or pd.isna(row.get(c_proj)) else str(row.get(c_proj)).strip()
            rows.append(
                {
                    "文件": f.name,
                    "项目名称": proj,
                    "科目": subj,
                    "摘要": summ,
                    "供应商名称": vend,
                    "实付金额": amt,
                    "一级科目": r.l1_name,
                    "二级科目": r.l2_name,
                    "分类ID": r.l2_id,
                    "命中层": r.layer,
                    "置信度": r.confidence,
                    "原因": r.reason,
                    "正确一级科目": "",
                    "正确二级科目": "",
                    "备注": "",
                }
            )
            if max_rows and len(rows) >= max_rows:
                break
        if max_rows and len(rows) >= max_rows:
            break

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows, columns=REVIEW_COLUMNS).to_excel(out_path, index=False)
    return out_path


def _is_valid_xlsx(path: Path) -> bool:
    """xlsx 本质是 zip；损坏/非 xlsx 会触发 File is not a zip file。"""
    path = Path(path)
    if not path.is_file() or path.stat().st_size < 64:
        return False
    try:
        import zipfile

        with zipfile.ZipFile(path, "r") as zf:
            # 至少能列出中央目录
            _ = zf.namelist()
        return True
    except Exception:
        return False


def update_excel_classification_columns(
    excel_path: Path,
    engine: ClassifyEngine,
    sheet_names=None,
    subject_col: str = "科目",
    summary_col: str = "摘要",
    vendor_col: str = "供应商名称",
    level1_col: str = "一级科目",
    level2_col: str = "二级科目",
    level3_col: str = "三级科目",
    extra_cols: bool = True,
) -> int:
    """更新已有 Excel 的分类列（保留其它内容）。"""
    excel_path = Path(excel_path)
    if not _is_valid_xlsx(excel_path):
        raise ValueError(f"不是有效 xlsx（可能损坏）: {excel_path}")

    wb = load_workbook(excel_path)
    target_sheets = sheet_names or wb.sheetnames
    updated = 0

    for sheet_name in target_sheets:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        header_to_col = {}
        for c in range(1, ws.max_column + 1):
            header = ws.cell(row=1, column=c).value
            if header is None:
                continue
            header_to_col[str(header).strip()] = c

        subj_idx = header_to_col.get(subject_col)
        summ_idx = header_to_col.get(summary_col)
        vend_idx = header_to_col.get(vendor_col)
        if subj_idx is None and summ_idx is None:
            continue

        def ensure_col(name: str) -> int:
            if name in header_to_col:
                return header_to_col[name]
            idx = ws.max_column + 1
            ws.cell(row=1, column=idx, value=name)
            header_to_col[name] = idx
            return idx

        lvl1_idx = ensure_col(level1_col)
        lvl2_idx = ensure_col(level2_col)
        lvl3_idx = ensure_col(level3_col)
        id_idx = conf_idx = layer_idx = None
        if extra_cols:
            id_idx = ensure_col("分类ID")
            layer_idx = ensure_col("命中层")
            conf_idx = ensure_col("置信度")

        # 先扫一遍收集待分类行（仅有实付金额），去重后再写回
        amt_idx = header_to_col.get("实付金额")
        pending: List[tuple] = []  # (row_idx, subj, summ, vend)
        for r in range(2, ws.max_row + 1):
            subject_value = ws.cell(row=r, column=subj_idx).value if subj_idx else None
            summary_value = ws.cell(row=r, column=summ_idx).value if summ_idx else None
            vendor_value = ws.cell(row=r, column=vend_idx).value if vend_idx else None
            if subject_value is None and summary_value in (None, ""):
                continue
            # 空实付：不分类，清空科目列标记
            if amt_idx is not None:
                amt_val = ws.cell(row=r, column=amt_idx).value
                try:
                    from payment_hub.utils import has_actual_payment

                    has_pay = has_actual_payment(amt_val)
                except Exception:
                    has_pay = amt_val is not None and str(amt_val).strip() not in {
                        "",
                        "nan",
                        "None",
                    }
                if not has_pay:
                    ws.cell(row=r, column=lvl1_idx, value="")
                    ws.cell(row=r, column=lvl2_idx, value="")
                    ws.cell(row=r, column=lvl3_idx, value="")
                    if id_idx:
                        ws.cell(row=r, column=id_idx, value="")
                    if layer_idx:
                        ws.cell(row=r, column=layer_idx, value="skip_no_payment")
                    if conf_idx:
                        ws.cell(row=r, column=conf_idx, value=None)
                    continue
            pending.append((r, subject_value, summary_value, vendor_value))

        results = engine.classify_many(
            [(s, m, v) for _, s, m, v in pending]
        )
        for (r, _, _, _), result in zip(pending, results):
            ws.cell(row=r, column=lvl1_idx, value=result.l1_name)
            ws.cell(row=r, column=lvl2_idx, value=result.l2_name)
            ws.cell(row=r, column=lvl3_idx, value=getattr(result, "l3_name", "") or "")
            if id_idx:
                leaf = getattr(result, "l3_id", None) or result.l2_id
                ws.cell(row=r, column=id_idx, value=leaf)
            if layer_idx:
                ws.cell(row=r, column=layer_idx, value=result.layer)
            if conf_idx:
                ws.cell(row=r, column=conf_idx, value=result.confidence)
            updated += 1

    # 重分类后刷新一二级下拉（与 taxonomy 同步）
    try:
        hub_root = Path(__file__).resolve().parent.parent
        if str(hub_root) not in sys.path:
            sys.path.insert(0, str(hub_root))
        from payment_hub.category_dropdowns import apply_category_dropdowns

        data_dir = Path(__file__).resolve().parent / "data"
        apply_category_dropdowns(wb, data_dir=data_dir)
    except Exception:
        pass

    # 清理已废弃的分类汇总相关 sheet（若仍存在）
    try:
        hub_root = Path(__file__).resolve().parent.parent
        if str(hub_root) not in sys.path:
            sys.path.insert(0, str(hub_root))
        from payment_hub.category_summary import remove_deprecated_sheets_from_workbook

        remove_deprecated_sheets_from_workbook(wb)
    except Exception:
        pass

    wb.save(excel_path)
    return updated


def reclassify_std_dir(
    std_dir: Path,
    engine: ClassifyEngine,
    *,
    sheet_names=None,
    extra_cols: bool = True,
) -> Dict[str, int]:
    """
    重刷 B_std 下所有项目 xlsx 的分类列。
    遇到损坏/非 zip 的 xlsx 会跳过并记入 skipped，不中断整批。
    """
    std_dir = Path(std_dir)
    stats: Dict[str, Any] = {
        "files": 0,
        "rows": 0,
        "skipped": 0,
        "skipped_files": [],
    }
    for f in sorted(std_dir.rglob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        try:
            rel = str(f.relative_to(std_dir))
        except ValueError:
            rel = str(f)
        if not _is_valid_xlsx(f):
            print(f"  ⚠️ 跳过损坏/非 xlsx: {rel}")
            stats["skipped"] += 1
            stats["skipped_files"].append(rel)
            continue
        try:
            n = update_excel_classification_columns(
                f,
                engine,
                sheet_names=sheet_names or ["明细数据_万泰铢", "明细数据"],
                extra_cols=extra_cols,
            )
        except Exception as e:
            print(f"  ⚠️ reclassify 失败 {rel}: {e}")
            stats["skipped"] += 1
            stats["skipped_files"].append(f"{rel} ({e})")
            continue
        stats["files"] += 1
        stats["rows"] += n
    if stats["skipped"]:
        print(
            f"  reclassify 完成：成功 {stats['files']} 文件 / {stats['rows']} 行，"
            f"跳过 {stats['skipped']} 个坏文件"
        )
    return stats
