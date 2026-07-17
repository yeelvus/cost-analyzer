# -*- coding: utf-8 -*-
"""
科目分类入口（兼容旧接口）。

默认走 learn_classify 学习闭环引擎（稳定 ID + 记忆 + 可迁移）。
若 learn_classify 不可用，回退到同目录 rules.json 的旧关键词逻辑。
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Optional, Tuple

import pandas as pd

CURRENT_DIR = Path(__file__).resolve().parent


def _try_load_learn_mapper(rules_json_path: str):
    try:
        from learn_classify.adapter import SubjectMapper as LearnMapper

        return LearnMapper(rules_json_path=rules_json_path)
    except Exception:
        # 允许以包路径或同级导入
        import sys

        if str(CURRENT_DIR) not in sys.path:
            sys.path.insert(0, str(CURRENT_DIR))
        try:
            from learn_classify.adapter import SubjectMapper as LearnMapper

            return LearnMapper(rules_json_path=rules_json_path)
        except Exception as exc:
            print(f"⚠️ learn_classify 未启用，回退旧规则引擎: {exc}")
            return None


class _LegacySubjectMapper:
    """旧版 JSON 关键词引擎（仅作回退）。"""

    def __init__(self, rules_json_path: str):
        with open(rules_json_path, "r", encoding="utf-8") as f:
            self.rules_dict = json.load(f)
        self.rules = []
        self.subject_to_level1 = {}
        subject_map = self.rules_dict.get("科目映射") if isinstance(self.rules_dict, dict) else None
        keyword_rules = self.rules_dict.get("关键词规则") if isinstance(self.rules_dict, dict) else None
        if isinstance(subject_map, dict) and isinstance(keyword_rules, dict):
            expanded_map = {}
            for k, v in subject_map.items():
                key_text = str(k).strip()
                if isinstance(v, dict) and re.match(r"^\d+_", key_text):
                    items_text = str(v.get("二级科目") or "").strip()
                    for level2 in [x.strip() for x in items_text.split(";") if x.strip()]:
                        expanded_map[level2] = key_text
            self.subject_to_level1 = expanded_map
            for level2, info in keyword_rules.items():
                if not isinstance(info, dict):
                    continue
                self.rules.append(
                    {
                        "二级科目": level2,
                        "一级科目": str(expanded_map.get(level2, "")).strip(),
                        "含有": info.get("含有", ""),
                        "不包含": info.get("不包含", ""),
                    }
                )
            return
        raise ValueError("无法解析旧版 rules.json")

    def _parse_contains(self, contains_str: str):
        if not contains_str or not str(contains_str).strip():
            return []
        parsed = []
        for group in [g.strip() for g in str(contains_str).split(";") if g.strip()]:
            and_keywords = [k.strip() for k in group.split("&") if k.strip()]
            if and_keywords:
                parsed.append(and_keywords)
        return parsed

    def _match_rule(self, summary: str, rule: dict) -> bool:
        if not summary or not isinstance(summary, str):
            return False
        s = summary.lower().strip()
        and_groups = self._parse_contains(rule.get("含有", ""))
        if not and_groups:
            return False
        matched = any(all(kw.lower() in s for kw in group) for group in and_groups)
        if not matched:
            return False
        excludes = rule.get("不包含", "")
        if excludes and str(excludes).strip():
            exclude_list = [k.strip().lower() for k in str(excludes).split(";") if k.strip()]
            if any(kw in s for kw in exclude_list):
                return False
        return True

    def classify(self, subject: str, summary: str) -> Tuple[str, str]:
        if pd.isna(subject) and (pd.isna(summary) or str(summary).strip() == ""):
            return "", ""
        txt = "" if pd.isna(subject) else str(subject).strip()
        if pd.isna(summary) or str(summary).strip() == "":
            return self.subject_to_level1.get(txt, ""), txt
        for rule in self.rules:
            if self._match_rule(str(summary), rule):
                return rule["一级科目"], rule["二级科目"]
        return self.subject_to_level1.get(txt, ""), txt

    def update_excel_subject_columns(
        self,
        excel_path: str,
        sheet_names=None,
        subject_col: str = "科目",
        summary_col: str = "摘要",
        level1_col: str = "一级科目",
        level2_col: str = "二级科目",
    ) -> int:
        from openpyxl import load_workbook

        wb = load_workbook(excel_path)
        target_sheets = sheet_names or wb.sheetnames
        updated_rows = 0
        for sheet_name in target_sheets:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            header_to_col = {}
            for c in range(1, ws.max_column + 1):
                header = ws.cell(row=1, column=c).value
                if header is None:
                    continue
                header_to_col[str(header).strip()] = c
            subj_idx = header_to_col.get(subject_col)
            summ_idx = header_to_col.get(summary_col)
            if subj_idx is None:
                continue
            lvl1_idx = header_to_col.get(level1_col) or (ws.max_column + 1)
            if level1_col not in header_to_col:
                ws.cell(row=1, column=lvl1_idx, value=level1_col)
            lvl2_idx = header_to_col.get(level2_col) or (ws.max_column + 1)
            if level2_col not in header_to_col:
                ws.cell(row=1, column=lvl2_idx, value=level2_col)
            for r in range(2, ws.max_row + 1):
                subject_value = ws.cell(row=r, column=subj_idx).value
                summary_value = ws.cell(row=r, column=summ_idx).value if summ_idx else ""
                if subject_value is None and summary_value in (None, ""):
                    continue
                level1, level2 = self.classify(subject_value, summary_value)
                if not level1:
                    level1 = "6_管理费"
                if not level2:
                    level2 = str(subject_value).strip() if subject_value is not None else "(空白)"
                ws.cell(row=r, column=lvl1_idx, value=level1)
                ws.cell(row=r, column=lvl2_idx, value=level2)
                updated_rows += 1
        wb.save(excel_path)
        return updated_rows


class SubjectMapper:
    """对外统一入口。"""

    def __init__(self, rules_json_path: str):
        # 优先学习闭环
        learn_data = CURRENT_DIR / "learn_classify" / "data" / "taxonomy.json"
        preferred = str(CURRENT_DIR / "learn_classify" / "data" / "rules.json")
        path = preferred if learn_data.exists() else rules_json_path
        self._impl = _try_load_learn_mapper(path)
        if self._impl is None:
            self._impl = _LegacySubjectMapper(rules_json_path)
        self.rules = getattr(self._impl, "rules", [])
        self.subject_to_level1 = getattr(self._impl, "subject_to_level1", {})

    def classify(self, subject, summary) -> Tuple[str, str]:
        return self._impl.classify(subject, summary)

    def process_excel(self, excel_path: str, output_path: Optional[str] = None):
        if hasattr(self._impl, "process_excel"):
            return self._impl.process_excel(excel_path, output_path)
        raise NotImplementedError("当前引擎不支持 process_excel")

    def update_excel_subject_columns(self, *args, **kwargs) -> int:
        return self._impl.update_excel_subject_columns(*args, **kwargs)
