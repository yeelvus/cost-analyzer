# -*- coding: utf-8 -*-
"""
一级 / 二级 / 三级 下拉（正式名，不含别名）。

- 一级：全部一级
- 二级：随一级联动
- 三级：随二级联动（无三级子项时可空）
"""

from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

CODE_ROOT = Path(__file__).resolve().parent.parent
OPTIONS_SHEET = "_分类选项"
DETAIL_SHEETS = ("明细数据", "明细数据_万泰铢", "分类总表")


def _ensure_learn_path() -> None:
    root = str(CODE_ROOT)
    if root not in sys.path:
        sys.path.insert(0, root)


def load_taxonomy_options(data_dir: Optional[Path] = None):
    """
    返回:
      l1_names
      l1_id_to_l2_names
      l1_name_to_id
      l2_name_to_id
      l2_id_to_l3_names
      tax
    下拉仅正式名。
    """
    _ensure_learn_path()
    from learn_classify.taxonomy import Taxonomy

    data_dir = Path(data_dir) if data_dir else CODE_ROOT / "learn_classify" / "data"
    tax = Taxonomy.load(data_dir / "taxonomy.json")

    l1_sorted = sorted(
        [x for x in tax.level1.values() if x.active],
        key=lambda z: (z.order, z.id),
    )
    l1_names = [x.name for x in l1_sorted]
    l1_name_to_id = {x.name: x.id for x in l1_sorted}

    l1_id_to_l2: Dict[str, List[str]] = {x.id: [] for x in l1_sorted}
    l2_name_to_id: Dict[str, str] = {}
    for l2 in sorted(tax.level2.values(), key=lambda z: (z.order, z.name)):
        if not l2.active or l2.parent not in l1_id_to_l2:
            continue
        if l2.name not in l1_id_to_l2[l2.parent]:
            l1_id_to_l2[l2.parent].append(l2.name)
        l2_name_to_id[l2.name] = l2.id

    l2_id_to_l3: Dict[str, List[str]] = {}
    for l3 in sorted(tax.level3.values(), key=lambda z: (z.order, z.name)):
        if not l3.active:
            continue
        if l3.parent not in tax.level2 or not tax.level2[l3.parent].active:
            continue
        l2_id_to_l3.setdefault(l3.parent, [])
        if l3.name not in l2_id_to_l3[l3.parent]:
            l2_id_to_l3[l3.parent].append(l3.name)

    return l1_names, l1_id_to_l2, l1_name_to_id, l2_name_to_id, l2_id_to_l3, tax


def _clear_old_category_validations(ws: Worksheet, col_indices: Sequence[int]) -> None:
    if not col_indices:
        return
    target_cols = set(col_indices)
    keep = []
    for dv in ws.data_validations.dataValidation:
        sq = str(dv.sqref or "")
        drop = False
        for col in target_cols:
            letter = get_column_letter(col)
            if re.search(rf"(^|\s|,)\$?{letter}\$?\d", sq, re.I):
                drop = True
                break
            if re.search(rf"(^|\s|,)\$?{letter}\$?\d*:\$?{letter}\$?\d*", sq, re.I):
                drop = True
                break
        if not drop:
            keep.append(dv)
    ws.data_validations.dataValidation = keep


def _remove_defined_names_with_prefix(wb: Workbook, prefixes: Sequence[str] = ("PAYL2_", "PAYL3_")) -> None:
    to_del = []
    try:
        for name in list(wb.defined_names):
            n = name if isinstance(name, str) else getattr(name, "name", str(name))
            if any(str(n).startswith(p) for p in prefixes) or str(n) in {
                "PAY_L1_LIST",
                "PAY_L1_MAP",
                "PAY_L2_MAP",
            }:
                to_del.append(str(n))
    except Exception:
        try:
            for dn in list(getattr(wb.defined_names, "definedName", [])):
                n = getattr(dn, "name", "")
                if any(n.startswith(p) for p in prefixes) or n in {
                    "PAY_L1_LIST",
                    "PAY_L1_MAP",
                    "PAY_L2_MAP",
                }:
                    to_del.append(n)
        except Exception:
            return
    for n in to_del:
        try:
            del wb.defined_names[n]
        except Exception:
            try:
                wb.defined_names.delete(n)
            except Exception:
                pass


def _write_options_sheet(
    wb: Workbook,
    l1_names: List[str],
    l1_id_to_l2: Dict[str, List[str]],
    l1_name_to_id: Dict[str, str],
    l2_name_to_id: Dict[str, str],
    l2_id_to_l3: Dict[str, List[str]],
) -> None:
    if OPTIONS_SHEET in wb.sheetnames:
        del wb[OPTIONS_SHEET]
    ws = wb.create_sheet(OPTIONS_SHEET)

    # A: 一级
    ws["A1"] = "一级科目"
    for i, name in enumerate(l1_names, start=2):
        ws.cell(i, 1, name)
    l1_end = max(2, 1 + len(l1_names))

    # B-C: 一级名 -> L1_id
    ws["B1"] = "一级科目"
    ws["C1"] = "一级命名键"
    for i, name in enumerate(l1_names, start=2):
        ws.cell(i, 2, name)
        ws.cell(i, 3, l1_name_to_id[name])
    map1_end = max(2, 1 + len(l1_names))

    # D-E: 二级名 -> L2_id（全局，供三级联动）
    ws["D1"] = "二级科目"
    ws["E1"] = "二级命名键"
    r = 2
    for l1n in l1_names:
        lid = l1_name_to_id[l1n]
        for l2n in l1_id_to_l2.get(lid) or []:
            ws.cell(r, 4, l2n)
            ws.cell(r, 5, l2_name_to_id.get(l2n, ""))
            r += 1
    map2_end = max(2, r - 1)

    # 从 G 列起：每个 L1 一列二级列表，表头 L1_id
    col = 7
    range_defs: List[Tuple[str, str]] = []
    for name in l1_names:
        lid = l1_name_to_id[name]
        l2s = l1_id_to_l2.get(lid) or []
        ws.cell(1, col, lid)
        for i, l2 in enumerate(l2s, start=2):
            ws.cell(i, col, l2)
        end_row = max(2, 1 + len(l2s)) if l2s else 2
        if not l2s:
            ws.cell(2, col, "")
        letter = get_column_letter(col)
        range_defs.append((f"PAYL2_{lid}", f"'{OPTIONS_SHEET}'!${letter}$2:${letter}${end_row}"))
        col += 1

    # 继续：每个有三级的 L2 一列，表头 L2_id，内容三级名
    for l2_id, l3s in sorted(l2_id_to_l3.items()):
        if not l3s:
            continue
        ws.cell(1, col, l2_id)
        for i, l3 in enumerate(l3s, start=2):
            ws.cell(i, col, l3)
        end_row = max(2, 1 + len(l3s))
        letter = get_column_letter(col)
        range_defs.append((f"PAYL3_{l2_id}", f"'{OPTIONS_SHEET}'!${letter}$2:${letter}${end_row}"))
        col += 1

    # 对照
    ref = col + 1
    ws.cell(1, ref, "对照_一级")
    ws.cell(1, ref + 1, "对照_二级")
    ws.cell(1, ref + 2, "对照_三级")
    rr = 2
    for l1n in l1_names:
        lid = l1_name_to_id[l1n]
        for l2n in l1_id_to_l2.get(lid) or []:
            l2id = l2_name_to_id.get(l2n, "")
            kids = l2_id_to_l3.get(l2id) or [""]
            for l3n in kids:
                ws.cell(rr, ref, l1n)
                ws.cell(rr, ref + 1, l2n)
                ws.cell(rr, ref + 2, l3n)
                rr += 1

    ws.cell(1, ref + 4, "说明")
    ws.cell(
        2,
        ref + 4,
        "下拉仅业务大类正式名；三级仅挂在有子类的二级下（如办公费→网络通讯/邮寄/办公物品）。别名不进下拉。",
    )

    ws.sheet_state = "hidden"

    _remove_defined_names_with_prefix(wb)

    def _add_name(n: str, attr: str) -> None:
        try:
            wb.defined_names[n] = DefinedName(n, attr_text=attr)
        except Exception:
            try:
                wb.defined_names.add(DefinedName(name=n, attr_text=attr))
            except Exception:
                pass

    _add_name("PAY_L1_LIST", f"'{OPTIONS_SHEET}'!$A$2:$A${l1_end}")
    _add_name("PAY_L1_MAP", f"'{OPTIONS_SHEET}'!$B$2:$C${map1_end}")
    _add_name("PAY_L2_MAP", f"'{OPTIONS_SHEET}'!$D$2:$E${map2_end}")
    for def_name, attr in range_defs:
        _add_name(def_name, attr)


def _header_map(ws: Worksheet) -> Dict[str, int]:
    m = {}
    for c in range(1, (ws.max_column or 1) + 1):
        v = ws.cell(1, c).value
        if v is not None:
            m[str(v).strip()] = c
    return m


def apply_dropdowns_to_worksheet(
    ws: Worksheet,
    *,
    l1_col_name: str = "一级科目",
    l2_col_name: str = "二级科目",
    l3_col_name: str = "三级科目",
    max_row: Optional[int] = None,
) -> bool:
    headers = _header_map(ws)
    c1 = headers.get(l1_col_name)
    c2 = headers.get(l2_col_name)
    c3 = headers.get(l3_col_name)
    if not c1 and not c2:
        return False

    last = max(max_row or ws.max_row or 2, 2) + 80
    cols = [c for c in (c1, c2, c3) if c]
    _clear_old_category_validations(ws, cols)

    if c1:
        letter = get_column_letter(c1)
        dv1 = DataValidation(
            type="list",
            formula1="=PAY_L1_LIST",
            allow_blank=True,
            showDropDown=False,
            showErrorMessage=False,
            showInputMessage=True,
            promptTitle="一级科目",
            prompt="选择一级大类",
        )
        dv1.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv1)

    if c2 and c1:
        l1_letter = get_column_letter(c1)
        l2_letter = get_column_letter(c2)
        formula = f'INDIRECT("PAYL2_"&IFERROR(VLOOKUP({l1_letter}2,PAY_L1_MAP,2,FALSE),""))'
        dv2 = DataValidation(
            type="list",
            formula1=f"={formula}",
            allow_blank=True,
            showDropDown=False,
            showErrorMessage=False,
            showInputMessage=True,
            promptTitle="二级科目",
            prompt="先选一级，再选二级",
        )
        dv2.add(f"{l2_letter}2:{l2_letter}{last}")
        ws.add_data_validation(dv2)

    if c3 and c2:
        l2_letter = get_column_letter(c2)
        l3_letter = get_column_letter(c3)
        # 三级随二级：PAYL3_{l2_id}
        formula = f'INDIRECT("PAYL3_"&IFERROR(VLOOKUP({l2_letter}2,PAY_L2_MAP,2,FALSE),""))'
        dv3 = DataValidation(
            type="list",
            formula1=f"={formula}",
            allow_blank=True,
            showDropDown=False,
            showErrorMessage=False,
            showInputMessage=True,
            promptTitle="三级科目",
            prompt="先选二级；无三级时可留空",
        )
        dv3.add(f"{l3_letter}2:{l3_letter}{last}")
        ws.add_data_validation(dv3)

    return True


def apply_category_dropdowns(
    wb: Workbook,
    data_dir: Optional[Path] = None,
    sheet_names: Optional[Sequence[str]] = None,
) -> int:
    l1_names, l1_id_to_l2, l1_name_to_id, l2_name_to_id, l2_id_to_l3, _tax = (
        load_taxonomy_options(data_dir)
    )
    if not l1_names:
        return 0

    _write_options_sheet(
        wb, l1_names, l1_id_to_l2, l1_name_to_id, l2_name_to_id, l2_id_to_l3
    )

    targets = list(sheet_names) if sheet_names else list(DETAIL_SHEETS)
    applied = 0
    seen = set()
    for name in list(targets) + list(wb.sheetnames):
        if name in seen or name == OPTIONS_SHEET:
            continue
        if name not in wb.sheetnames:
            continue
        seen.add(name)
        ws = wb[name]
        headers = _header_map(ws)
        if "一级科目" not in headers and "二级科目" not in headers:
            continue
        if apply_dropdowns_to_worksheet(ws):
            applied += 1
    return applied


def apply_category_dropdowns_to_file(
    path: Path,
    data_dir: Optional[Path] = None,
    sheet_names: Optional[Sequence[str]] = None,
) -> int:
    path = Path(path)
    wb = load_workbook(path)
    n = apply_category_dropdowns(wb, data_dir=data_dir, sheet_names=sheet_names)
    wb.save(path)
    return n


def apply_category_dropdowns_to_std_dir(
    std_dir: Path,
    data_dir: Optional[Path] = None,
) -> Dict[str, int]:
    std_dir = Path(std_dir)
    stats = {"files": 0, "sheets": 0, "errors": 0}
    for f in sorted(std_dir.rglob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        try:
            n = apply_category_dropdowns_to_file(f, data_dir=data_dir)
            stats["files"] += 1
            stats["sheets"] += n
            print(f"  ✅ 下拉 {f.relative_to(std_dir)} ({n} sheets)")
        except Exception as e:
            stats["errors"] += 1
            print(f"  ⚠️ 失败 {f.name}: {e}")
    return stats
