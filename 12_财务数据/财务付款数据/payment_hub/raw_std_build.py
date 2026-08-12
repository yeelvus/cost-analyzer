# -*- coding: utf-8 -*-
"""
步骤 raw_std：A_财务付款_raw/{组名}/*.xlsx
  → 列标准化 + 标准项目名称匹配 + 组内合并
  → A1_财务付款_raw_std/{组名}/{组名}_财务付款_raw_std.xlsx

标准列（顺序）：
  项目编码, 项目名称, 原始项目名称, 摘要,
  应付金额, 预扣税额, 实付金额, 增值税额, 不含税额, 供应商名称, 科目, 业务分组,
  日期, 年, 月

年 / 月 为数字（Excel 数值格式）；保留原始项目名称便于对比。
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd

from .config import Paths
from .deps import get_excel_layout_fn, get_project_matcher_module, project_match_db_path
from .std_build import (
    AMOUNT_COLUMNS,
    COLUMN_ALIASES,
    TARGET_COLUMNS,
    normalize_date_month_year,
)
from .utils import (
    iter_excel_files,
    list_raw_groups,
    normalize_text,
    to_numeric_series,
)

# 明细导出列顺序：日期/年/月靠后；无年月、无来源列
RAW_STD_ALL_COLUMNS = [
    "项目编码",
    "项目名称",
    "原始项目名称",
    "摘要",
    "应付金额",
    "预扣税额",
    "实付金额",
    "增值税额",
    "不含税额",
    "供应商名称",
    "科目",
    "业务分组",
    "日期",
    "年",
    "月",
]
RAW_STD_CORE_COLUMNS = [
    "项目名称",
    "日期",
    "摘要",
    "应付金额",
    "预扣税额",
    "实付金额",
    "增值税额",
    "不含税额",
    "供应商名称",
]
# 内部追溯用，不写入明细表
_INTERNAL_META = ["来源文件", "来源Sheet"]

# raw_std 专用别名扩展（银行账常见「借方」等）
RAW_STD_ALIASES: Dict[str, List[str]] = {
    **{k: list(v) for k, v in COLUMN_ALIASES.items()},
}
RAW_STD_ALIASES["实付金额"] = list(
    dict.fromkeys(
        RAW_STD_ALIASES.get("实付金额", [])
        + ["借方", "借方金额", "支出", "付款", "发生额"]
    )
)
RAW_STD_ALIASES["应付金额"] = list(
    dict.fromkeys(
        RAW_STD_ALIASES.get("应付金额", []) + ["贷方", "贷方金额", "收入"]
    )
)

# 至少具备：摘要 或 项目名称；以及任一金额列
_AMOUNT_ANY = set(AMOUNT_COLUMNS)


def find_raw_std_column_mapping(df: pd.DataFrame) -> Dict[str, Optional[str]]:
    mapping: Dict[str, Optional[str]] = {}
    normalized_to_original = {normalize_text(col): col for col in df.columns}
    targets = list(dict.fromkeys(TARGET_COLUMNS + RAW_STD_CORE_COLUMNS))
    for target in targets:
        candidates = RAW_STD_ALIASES.get(target, [target])
        found_col = None
        for alias in candidates:
            norm_alias = normalize_text(alias)
            if norm_alias in normalized_to_original:
                found_col = normalized_to_original[norm_alias]
                break
        if found_col is None:
            for alias in candidates:
                norm_alias = normalize_text(alias)
                for norm_col, origin_col in normalized_to_original.items():
                    if norm_alias and (norm_alias in norm_col or norm_col in norm_alias):
                        found_col = origin_col
                        break
                if found_col is not None:
                    break
        mapping[target] = found_col
    return mapping


def process_raw_sheet(
    file_path: Path,
    sheet_name: str,
    *,
    group_name: str,
) -> Optional[pd.DataFrame]:
    try:
        raw_df = pd.read_excel(file_path, sheet_name=sheet_name)
    except Exception as exc:
        print(f"  ⚠️ 读取失败: {file_path.name} | {sheet_name} | {exc}")
        return None
    if raw_df is None or raw_df.empty:
        return None

    # 清理全空列
    raw_df = raw_df.dropna(axis=1, how="all")
    if raw_df.empty:
        return None

    mapping = find_raw_std_column_mapping(raw_df)

    has_name = mapping.get("项目名称") is not None
    has_summary = mapping.get("摘要") is not None
    has_amount = any(mapping.get(c) is not None for c in AMOUNT_COLUMNS)
    if not has_amount:
        print(f"    - 跳过 sheet: {sheet_name}，无金额列")
        return None
    if not (has_name or has_summary):
        print(f"    - 跳过 sheet: {sheet_name}，无项目名称/摘要")
        return None

    out = pd.DataFrame()
    for col in TARGET_COLUMNS:
        src = mapping.get(col)
        out[col] = raw_df[src] if src is not None else pd.NA

    for amount_col in AMOUNT_COLUMNS:
        out[amount_col] = to_numeric_series(out[amount_col])

    # 去掉金额全空/全 0 行
    row_sum = out[AMOUNT_COLUMNS].fillna(0).sum(axis=1)
    out = out[row_sum.ne(0)]
    if out.empty:
        return None

    out = normalize_date_month_year(out)
    out = ensure_numeric_year_month(out)

    # 文本列去空白
    for col in ("项目名称", "摘要", "供应商名称", "科目"):
        if col in out.columns:
            out[col] = (
                out[col]
                .astype(str)
                .str.strip()
                .replace({"nan": "", "None": "", "NaT": ""})
            )

    # 保留原始项目名称（匹配前快照）
    out["原始项目名称"] = out["项目名称"]
    out["业务分组"] = group_name
    out["来源文件"] = file_path.name
    out["来源Sheet"] = str(sheet_name)
    if "项目编码" not in out.columns:
        out["项目编码"] = ""

    return out


def ensure_numeric_year_month(df: pd.DataFrame) -> pd.DataFrame:
    """
    年 / 月 统一为可空整数（Excel 数值）：年=2024, 月=8
    """
    out = df.copy()
    parsed = (
        pd.to_datetime(out["日期"], errors="coerce")
        if "日期" in out.columns
        else pd.Series(pd.NaT, index=out.index)
    )

    year = (
        pd.to_numeric(out["年"], errors="coerce")
        if "年" in out.columns
        else pd.Series(pd.NA, index=out.index)
    )
    month = (
        pd.to_numeric(out["月"], errors="coerce")
        if "月" in out.columns
        else pd.Series(pd.NA, index=out.index)
    )

    year = year.where(year.notna(), parsed.dt.year)
    month = month.where(month.notna(), parsed.dt.month)
    month = month.where((month >= 1) & (month <= 12))

    out["年"] = year.round().astype("Int64")
    out["月"] = month.round().astype("Int64")
    if "年月" in out.columns:
        out = out.drop(columns=["年月"])
    return out


def apply_project_name_match(
    df: pd.DataFrame,
    *,
    prompt: bool = False,
    interactive: bool = True,
    window_title: str = "财务付款 raw_std 项目匹配",
) -> pd.DataFrame:
    """
    标准项目名称匹配（精确别名，无模糊）。
    - 原始项目名称：匹配前名称（已存在则保留）
    - 项目名称：标准名称
    - 项目编码：匹配编码
    """
    if df is None or df.empty:
        return df

    out = df.copy()
    # 以「原始项目名称」为匹配键；若无该列则用当前项目名称
    if "原始项目名称" not in out.columns:
        out["原始项目名称"] = out.get("项目名称", "")
    out["原始项目名称"] = (
        out["原始项目名称"].fillna("").astype(str).str.strip()
        .replace({"nan": "", "None": "", "NaT": ""})
    )
    # 若原始为空则回退当前项目名称
    cur = out.get("项目名称", pd.Series("", index=out.index))
    cur = cur.fillna("").astype(str).str.strip().replace({"nan": "", "None": ""})
    empty_raw = out["原始项目名称"] == ""
    out.loc[empty_raw, "原始项目名称"] = cur.loc[empty_raw]

    mod = get_project_matcher_module()
    ProjectMatcher = getattr(mod, "ProjectMatcher", None)
    db_path = project_match_db_path()

    print(
        f"  项目匹配: 精确别名 only（无模糊）"
        f"{'；弹窗补录=开' if prompt else '；弹窗补录=关（--prompt-project 可开）'}"
    )

    raw_series = out["原始项目名称"]
    unique = sorted(
        {
            n
            for n in raw_series.tolist()
            if n and n.lower() not in {"nan", "none", "nat"}
        }
    )

    if ProjectMatcher is not None:
        m = ProjectMatcher(db_path=db_path, load_info=True)
        mapped = m.match_batch(
            unique,
            prompt=prompt,
            interactive=interactive,
            window_title=window_title,
            fuzzy=False,
        )
        code_map = {k: (v.get("project_code") or "") for k, v in mapped.items()}
        name_map = {k: (v.get("project_name") or k) for k, v in mapped.items()}
        m.save()
    else:
        db = mod.load_or_init_db(db_path, [])
        code_map: Dict[str, str] = {}
        name_map: Dict[str, str] = {}
        pending: List[str] = []
        for raw_name in unique:
            hit = mod.match_project(raw_name, db, fuzzy=False)
            if hit and hit.get("source") not in {"unmatched_keep", "unmatched"}:
                code_map[raw_name] = hit.get("project_code", "") or ""
                name_map[raw_name] = hit.get("project_name", raw_name)
            else:
                pending.append(raw_name)
                code_map[raw_name] = ""
                name_map[raw_name] = raw_name
        if pending and prompt:
            print(f"  ⚠️ 未匹配项目 {len(pending)} 个，弹出手动匹配…")
            mod.prompt_unmatched_project_mappings(
                pending,
                db,
                window_title=window_title,
                db_path=db_path,
                interactive=interactive,
            )
            for raw_name in pending:
                hit = mod.match_project(raw_name, db, fuzzy=False)
                if hit and hit.get("source") not in {"unmatched_keep", "unmatched"}:
                    code_map[raw_name] = hit.get("project_code", "") or ""
                    name_map[raw_name] = hit.get("project_name", raw_name)
                else:
                    mod.mark_unmatched(db, raw_name)
        elif pending:
            print(f"  未匹配保留原名（未弹窗）: {len(pending)} 个")
        mod.save_db(db_path, db)

    out["项目编码"] = raw_series.map(code_map).fillna("").astype(str)
    # 未匹配时标准名仍用原始名，便于对照
    out["项目名称"] = raw_series.map(name_map).fillna(raw_series).astype(str)

    matched_n = int((out["项目编码"].astype(str).str.strip() != "").sum())
    print(
        f"  匹配结果: 有编码 {matched_n}/{len(out)} 行"
        f"（唯一原名 {len(unique)} 个）"
    )
    no_code_names = sorted(
        {n for n in unique if not str(code_map.get(n, "")).strip()}
    )
    if no_code_names:
        print(
            f"  无编码项目名: {len(no_code_names)} 个"
            f"  例: {no_code_names[:8]}{'…' if len(no_code_names) > 8 else ''}"
        )
    return out


def process_raw_file(file_path: Path, *, group_name: str) -> List[pd.DataFrame]:
    parts: List[pd.DataFrame] = []
    try:
        sheets = pd.ExcelFile(file_path).sheet_names
    except Exception as exc:
        print(f"  ⚠️ 打开失败: {file_path.name} | {exc}")
        return parts
    for sheet in sheets:
        part = process_raw_sheet(file_path, sheet, group_name=group_name)
        if part is not None and not part.empty:
            parts.append(part)
            print(f"    ✓ {file_path.name} / {sheet}: {len(part)} 行")
    return parts


def _prepare_excel_export(df: pd.DataFrame) -> pd.DataFrame:
    """导出前：仅明细列；年/月 为 int/None；日期为 datetime。"""
    out = df.copy()
    # 去掉内部列与年月
    drop_cols = [c for c in list(_INTERNAL_META) + ["年月"] if c in out.columns]
    if drop_cols:
        out = out.drop(columns=drop_cols)

    for col in RAW_STD_ALL_COLUMNS:
        if col not in out.columns:
            out[col] = pd.NA
    out = out[RAW_STD_ALL_COLUMNS]

    for col in ("年", "月"):
        if col not in out.columns:
            continue
        nums = pd.to_numeric(out[col], errors="coerce")
        out[col] = [None if pd.isna(v) else int(v) for v in nums]

    if "日期" in out.columns:
        parsed = pd.to_datetime(out["日期"], errors="coerce")
        mixed = []
        for i, raw in enumerate(out["日期"].tolist()):
            if pd.notna(parsed.iloc[i]):
                mixed.append(parsed.iloc[i].to_pydatetime())
            else:
                text = "" if raw is None else str(raw).strip()
                mixed.append(None if text in {"", "nan", "NaT", "None"} else raw)
        out["日期"] = mixed
    return out


def apply_raw_std_format(path: Path) -> None:
    try:
        from openpyxl import load_workbook

        layout_fn = get_excel_layout_fn()
        wb = load_workbook(path)
        for ws in wb.worksheets:
            if ws.title not in {"明细", "项目统计"} and ws.max_row < 1:
                continue
            try:
                layout_fn(
                    ws,
                    numeric_start_col=1,
                    amount_headers=set(AMOUNT_COLUMNS),
                    header_row=1,
                    header_height=35,
                    freeze_panes="A2",
                )
            except Exception:
                pass

            headers = {
                str(c.value).strip(): c.column
                for c in ws[1]
                if c.value is not None
            }
            for col_name in ("年", "月", "日期"):
                col_idx = headers.get(col_name)
                if not col_idx:
                    continue
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=col_idx)
                    if cell.value is None or cell.value == "":
                        continue
                    if col_name == "日期":
                        cell.number_format = "YYYY-MM-DD"
                    else:
                        if isinstance(cell.value, str) and cell.value.strip().isdigit():
                            cell.value = int(cell.value.strip())
                        cell.number_format = "0"

            for amount_name in AMOUNT_COLUMNS:
                col_idx = headers.get(amount_name)
                if not col_idx:
                    continue
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=col_idx)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"

        wb.save(path)
    except Exception as exc:
        print(f"  ⚠️ 样式跳过: {exc}")


def run_raw_std_for_group(
    group_name: str,
    group_raw_dir: Path,
    out_path: Path,
    *,
    prompt_project: bool = False,
) -> Optional[Path]:
    group_raw_dir = Path(group_raw_dir)
    files = iter_excel_files(group_raw_dir)
    if not files:
        print(f"  [{group_name}] 无 Excel，跳过")
        return None

    print(f"  [{group_name}] 源文件 {len(files)} 个 ← {group_raw_dir}")
    parts: List[pd.DataFrame] = []
    for fp in files:
        print(f"  处理: {fp.name}")
        parts.extend(process_raw_file(fp, group_name=group_name))

    if not parts:
        print(f"  [{group_name}] 未提取到有效行")
        return None

    merged = pd.concat(parts, ignore_index=True)
    merged = ensure_numeric_year_month(merged)

    # 标准项目名称匹配（组内合并后统一匹配一次；保留原始项目名称）
    merged = apply_project_name_match(
        merged,
        prompt=prompt_project,
        interactive=True,
        window_title=f"财务付款 raw_std 项目匹配 · {group_name}",
    )

    # 日期排序（无效日期靠后）
    if "日期" in merged.columns:
        merged = merged.copy()
        merged["_sort_date"] = pd.to_datetime(merged["日期"], errors="coerce")
        sort_cols = ["_sort_date"]
        if "来源文件" in merged.columns:
            sort_cols.append("来源文件")
        merged = (
            merged.sort_values(by=sort_cols, kind="mergesort", na_position="last")
            .drop(columns=["_sort_date"])
            .reset_index(drop=True)
        )

    export_df = _prepare_excel_export(merged)
    n_files = (
        int(merged["来源文件"].nunique())
        if "来源文件" in merged.columns
        else len(files)
    )

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name="明细")
        # 项目匹配概况（含原始名对比）
        proj = (
            merged.groupby(["项目编码", "项目名称", "原始项目名称"], dropna=False)
            .size()
            .reset_index(name="行数")
            .sort_values("行数", ascending=False)
        )
        proj.to_excel(writer, index=False, sheet_name="项目统计")

    apply_raw_std_format(out_path)
    print(
        f"  ✅ [{group_name}] 合并 {len(merged)} 行 → {out_path.name}"
        f"（来自 {n_files} 个文件）"
    )
    return out_path


def run_raw_std(paths: Paths, *, prompt_project: bool = False) -> List[str]:
    """
    A_raw 一级文件夹 → 每组合并一份标准化 Excel：
      A1_财务付款_raw_std/{组名}/{组名}_财务付款_raw_std.xlsx
    """
    raw_dir = Path(paths.raw_dir)
    out_root = Path(paths.raw_std_dir)
    if not raw_dir.is_dir():
        raise FileNotFoundError(f"原始目录不存在: {raw_dir}")
    out_root.mkdir(parents=True, exist_ok=True)

    groups = list_raw_groups(raw_dir)
    if not groups:
        print("未发现 A_raw 一级分组，跳过 raw_std")
        return []

    print(f"raw_std 分组: {', '.join(g for g, _ in groups)}")
    print(f"输出目录: {out_root}/{{组名}}/")
    print(
        f"项目匹配弹窗: {'开启' if prompt_project else '关闭'}"
        f"（--prompt-project / --no-prompt-project）"
    )

    written: List[str] = []
    for group_name, group_raw in groups:
        print(f"\n{'=' * 64}")
        print(f"  raw_std · {group_name}")
        print(f"{'=' * 64}")
        out_path = paths.group_raw_std_xlsx(group_name)
        result = run_raw_std_for_group(
            group_name,
            group_raw,
            out_path,
            prompt_project=prompt_project,
        )
        if result is not None:
            written.append(str(result))

    print(f"\n✅ raw_std 完成：{len(written)} 个合并文件 → {out_root}")
    for p in written:
        print(f"   · {p}")
    return written
