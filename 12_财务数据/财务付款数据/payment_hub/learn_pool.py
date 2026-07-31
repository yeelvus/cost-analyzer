# -*- coding: utf-8 -*-
"""
D_财务付款_学习池

- 01_按项目/{编码}_{名称}.xlsx
- 02_按月份/{YYYY-MM}_财务数据.xlsx
- 03_已确认/已确认分类.xlsx

摘要去重；手改或确认后 apply → 记忆 + 剥离 + reclassify B_std。
"""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .category_dropdowns import apply_category_dropdowns
from .config import Paths
from .pkg import ensure_code_root
from .utils import safe_filename, safe_text

POOL_SHEET = "学习池"
STATUS_PENDING = "待确认"
STATUS_CONFIRMED = "确认"
STATUS_VALUES = [STATUS_PENDING, STATUS_CONFIRMED, "已确认", "是", "Y", "y"]

COLUMNS = [
    "确认状态",
    "摘要",
    "一级科目",
    "二级科目",
    "三级科目",
    "供应商名称",
    "出现次数",
]

SNAPSHOT_META = "_meta"


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _is_confirmed(val: Any) -> bool:
    s = safe_text(val)
    if not s:
        return False
    s2 = s.replace(" ", "").replace("　", "")
    return s2 in {
        STATUS_CONFIRMED,
        "已确认",
        "是",
        "Y",
        "y",
        "1",
        "true",
        "True",
        "OK",
        "ok",
        "√",
        "✓",
    } or s2.startswith("确认")


def _norm_cat(v: Any) -> str:
    s = safe_text(v)
    if s in {"(空白)", "-", "—"}:
        return ""
    return s


def _triple(l1: Any, l2: Any, l3: Any) -> Tuple[str, str, str]:
    return _norm_cat(l1), _norm_cat(l2), _norm_cat(l3)


def _month_from_date(val: Any) -> str:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return "未知月份"
    if isinstance(val, datetime):
        return val.strftime("%Y-%m")
    try:
        ts = pd.to_datetime(val, errors="coerce")
        if pd.isna(ts):
            return "未知月份"
        return ts.strftime("%Y-%m")
    except Exception:
        s = safe_text(val)
        m = re.match(r"(20\d{2})[-/]?(\d{1,2})", s)
        if m:
            return f"{m.group(1)}-{int(m.group(2)):02d}"
        return "未知月份"


def _pick_sheet(path: Path) -> str:
    try:
        xls = pd.ExcelFile(path)
        if "明细数据" in xls.sheet_names:
            return "明细数据"
        return xls.sheet_names[0]
    except Exception:
        return "Sheet1"


def load_confirmed_summaries(confirmed_path: Path) -> Dict[str, Dict[str, str]]:
    """摘要 -> {一级,二级,三级,供应商,确认状态,...}"""
    path = Path(confirmed_path)
    if not path.is_file():
        return {}
    try:
        df = pd.read_excel(path, sheet_name=0)
    except Exception:
        return {}
    out: Dict[str, Dict[str, str]] = {}
    for _, row in df.iterrows():
        summ = safe_text(row.get("摘要"))
        if not summ:
            continue
        out[summ] = {
            "确认状态": STATUS_CONFIRMED,
            "摘要": summ,
            "一级科目": _norm_cat(row.get("一级科目")),
            "二级科目": _norm_cat(row.get("二级科目")),
            "三级科目": _norm_cat(row.get("三级科目")),
            "供应商名称": safe_text(row.get("供应商名称")),
            "出现次数": str(row.get("出现次数") or ""),
            "确认时间": safe_text(row.get("确认时间")) or _now(),
            "来源": safe_text(row.get("来源")),
        }
    return out


def load_snapshot(path: Path) -> Dict[str, Dict[str, Dict[str, str]]]:
    """
    返回: { 相对路径或文件名: { 摘要: {一级,二级,三级} } }
    兼容旧版全局 summaries 映射（视为 __global__）。
    """
    p = Path(path)
    if not p.is_file():
        return {}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(data.get("files"), dict):
            out: Dict[str, Dict[str, Dict[str, str]]] = {}
            for fk, mp in data["files"].items():
                if isinstance(mp, dict):
                    out[str(fk)] = {
                        str(sk): sv for sk, sv in mp.items() if isinstance(sv, dict)
                    }
            return out
        # 旧格式
        rows = data.get("summaries") or {}
        if isinstance(rows, dict):
            return {
                "__global__": {
                    str(k): v for k, v in rows.items() if isinstance(v, dict)
                }
            }
        return {}
    except Exception:
        return {}


def save_snapshot(path: Path, files_map: Dict[str, Dict[str, Dict[str, str]]]) -> None:
    """files_map: {文件名: {摘要: {一级,二级,三级}}}"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    n = sum(len(v) for v in files_map.values())
    payload = {
        "exported_at": _now(),
        "file_count": len(files_map),
        "row_count": n,
        "files": files_map,
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def is_thcc_project_code(code: str) -> bool:
    """是否为 THCC 正式项目编码（用于按月份池过滤）。"""
    s = safe_text(code)
    if not s:
        return False
    return bool(re.match(r"^THCC-", s, re.I))


def collect_raw_lines_from_std(std_dir: Path) -> List[Dict[str, Any]]:
    """从 B_std 收集明细行。"""
    std_dir = Path(std_dir)
    lines: List[Dict[str, Any]] = []
    for f in sorted(std_dir.glob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        sheet = _pick_sheet(f)
        try:
            df = pd.read_excel(f, sheet_name=sheet)
        except Exception as e:
            print(f"  ⚠️ 跳过 {f.name}: {e}")
            continue
        if df.empty:
            continue
        # 文件名兜底项目
        stem = f.stem
        # 编码_名称_(日期)
        m = re.match(r"^([^_]+)_(.+?)_\(", stem)
        file_code = m.group(1) if m else ""
        file_name = m.group(2) if m else stem

        for _, row in df.iterrows():
            summ = safe_text(row.get("摘要"))
            if not summ:
                continue
            code = safe_text(row.get("项目编码")) or file_code or "未识别编码"
            pname = safe_text(row.get("项目名称")) or file_name or "未命名项目"
            lines.append(
                {
                    "摘要": summ,
                    "一级科目": _norm_cat(row.get("一级科目")),
                    "二级科目": _norm_cat(row.get("二级科目")),
                    "三级科目": _norm_cat(row.get("三级科目")),
                    "供应商名称": safe_text(row.get("供应商名称")),
                    "项目编码": code,
                    "项目名称": pname,
                    "月份": _month_from_date(row.get("日期")),
                    "实付金额": pd.to_numeric(row.get("实付金额"), errors="coerce") or 0,
                    "is_thcc": is_thcc_project_code(code),
                }
            )
    return lines


def _dedupe_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """按摘要去重：众数科目/供应商，累计出现次数。"""
    buckets: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for r in rows:
        buckets[r["摘要"]].append(r)

    out: List[Dict[str, Any]] = []
    for summ, items in sorted(buckets.items(), key=lambda x: (-len(x[1]), x[0])):
        l1s = Counter(_norm_cat(i["一级科目"]) for i in items)
        l2s = Counter(_norm_cat(i["二级科目"]) for i in items)
        l3s = Counter(_norm_cat(i["三级科目"]) for i in items)
        vendors = Counter(safe_text(i["供应商名称"]) for i in items if safe_text(i["供应商名称"]))
        l1 = l1s.most_common(1)[0][0] if l1s else ""
        l2 = l2s.most_common(1)[0][0] if l2s else ""
        l3 = l3s.most_common(1)[0][0] if l3s else ""
        vend = vendors.most_common(1)[0][0] if vendors else ""
        out.append(
            {
                "确认状态": STATUS_PENDING,
                "摘要": summ,
                "一级科目": l1,
                "二级科目": l2,
                "三级科目": l3,
                "供应商名称": vend,
                "出现次数": len(items),
            }
        )
    # 默认：一级/二级/三级/摘要 归类排序（便于浏览）
    return _sort_by_category(out)


def _sort_by_category(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """一级 → 二级 → 三级 → 摘要 排序。"""
    blank = "\uffff"  # 空科目排最后

    def key(r: Dict[str, Any]):
        return (
            _norm_cat(r.get("一级科目")) or blank,
            _norm_cat(r.get("二级科目")) or blank,
            _norm_cat(r.get("三级科目")) or blank,
            safe_text(r.get("摘要")),
        )

    return sorted(rows, key=key)


def _format_pool_workbook(
    path: Path,
    rows: List[Dict[str, Any]],
    data_dir: Path,
    *,
    title: str = "",
) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = POOL_SHEET

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for c, h in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, row in enumerate(rows, start=2):
        for c, h in enumerate(COLUMNS, start=1):
            val = row.get(h, "")
            if h == "确认状态" and not val:
                val = STATUS_PENDING
            ws.cell(row=i, column=c, value=val if val != "" else None)

    last = max(2, len(rows) + 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    widths = [10, 36, 16, 18, 16, 22, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 确认状态下拉
    dv = DataValidation(
        type="list",
        formula1=f'"{STATUS_PENDING},{STATUS_CONFIRMED}"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=False,
    )
    dv.add(f"A2:A{last + 200}")
    ws.add_data_validation(dv)

    # 一/二/三级下拉
    try:
        apply_category_dropdowns(wb, data_dir=data_dir, sheet_names=[POOL_SHEET])
    except Exception as e:
        print(f"  ⚠️ 下拉写入失败 {path.name}: {e}")

    if title:
        # 说明放最后，打开文件默认先看到「学习池」数据
        note = wb.create_sheet("说明")
        note["A1"] = title
        note["A2"] = "① 在「学习池」表改「一级/二级/三级」，或把「确认状态」改为「确认」"
        note["A3"] = "② 保存 Excel 后，必须运行："
        note["A4"] = "   python 财务付款一体化.py apply-learn-pool"
        note["A5"] = "③ 写入 03_已确认、剥离已确认摘要，并泛化抬升其余未确认行的预分类（相似/去月份同型）"
        note["A6"] = "④ 推荐在 00_汇总/全部摘要_学习池.xlsx 快速改全量（含非THCC项目）"
        note["A7"] = "⑤ 不要只运行 export-learn-pool（重新导出可能覆盖未 apply 的修改）"
        note["A8"] = "⑥ 确认越多，未确认摘要的建议科目会越准（memory → core → similar → 规则）"
        note.column_dimensions["A"].width = 90

    # 保证活动表是学习池
    try:
        wb.active = wb[POOL_SHEET]
    except Exception:
        pass

    wb.save(path)


def export_learn_pool(paths: Paths) -> Dict[str, Any]:
    """从 B_std 导出 01_按项目 / 02_按月份（排除已确认摘要）。"""
    paths.ensure_output_parents()
    confirmed = load_confirmed_summaries(paths.learn_pool_confirmed)
    confirmed_keys = set(confirmed.keys())

    print(f"已确认摘要: {len(confirmed_keys)}")
    print(f"扫描 std: {paths.std_dir}")
    lines = collect_raw_lines_from_std(paths.std_dir)
    print(f"明细行: {len(lines)}")

    # 过滤已确认
    lines = [r for r in lines if r["摘要"] not in confirmed_keys]
    print(f"未确认明细行: {len(lines)}")

    # 按项目：全部项目（含未识别编码 / LSBM / XCY 等）
    # 按月份：仅 THCC- 正式编码项目
    by_project: Dict[Tuple[str, str], List[Dict]] = defaultdict(list)
    by_month: Dict[str, List[Dict]] = defaultdict(list)
    thcc_lines = 0
    non_thcc_lines = 0
    for r in lines:
        by_project[(r["项目编码"], r["项目名称"])].append(r)
        if r.get("is_thcc"):
            by_month[r["月份"]].append(r)
            thcc_lines += 1
        else:
            non_thcc_lines += 1
    print(f"按项目: 全部 {len(lines)} 行 · 按月份仅 THCC: {thcc_lines} 行（排除非THCC {non_thcc_lines} 行）")

    # 清空旧的 00/01/02 文件（避免残留摘要）
    for d in (
        paths.learn_pool_all,
        paths.learn_pool_by_project,
        paths.learn_pool_by_month,
    ):
        d.mkdir(parents=True, exist_ok=True)
        for f in d.glob("*.xlsx"):
            if f.name.startswith("~$"):
                continue
            try:
                f.unlink()
            except OSError:
                pass

    # 按文件快照，避免项目/月份本地众数与全局不一致被误判为手改
    files_snap: Dict[str, Dict[str, Dict[str, str]]] = {}
    project_files = 0
    month_files = 0
    unique_summ: set = set()

    # 00_汇总：全部项目（含非 THCC / 未识别编码），摘要去重，按科目排序
    all_rows = _sort_by_category(_dedupe_rows(lines))
    if all_rows:
        _format_pool_workbook(
            paths.learn_pool_all_xlsx,
            all_rows,
            paths.classify_data,
            title="全量汇总学习池 · 含全部项目与非项目编码（未识别编码等）",
        )
        key_all = f"00_汇总/{paths.learn_pool_all_xlsx.name}"
        files_snap[key_all] = {}
        for r in all_rows:
            unique_summ.add(r["摘要"])
            files_snap[key_all][r["摘要"]] = {
                "一级科目": r["一级科目"],
                "二级科目": r["二级科目"],
                "三级科目": r["三级科目"],
            }
        print(f"  汇总 {paths.learn_pool_all_xlsx.name}: {len(all_rows)} 摘要（全量）")

    for (code, pname), items in sorted(by_project.items(), key=lambda x: (x[0][0], x[0][1])):
        rows = _dedupe_rows(items)
        if not rows:
            continue
        fname = f"{safe_filename(code)}_{safe_filename(pname)}.xlsx"
        out = paths.learn_pool_by_project / fname
        _format_pool_workbook(
            out,
            rows,
            paths.classify_data,
            title=f"项目学习池 · {code} {pname}",
        )
        project_files += 1
        key = f"01_按项目/{fname}"
        files_snap[key] = {}
        for r in rows:
            unique_summ.add(r["摘要"])
            files_snap[key][r["摘要"]] = {
                "一级科目": r["一级科目"],
                "二级科目": r["二级科目"],
                "三级科目": r["三级科目"],
            }
        print(f"  项目 {fname}: {len(rows)} 摘要")

    for month, items in sorted(by_month.items()):
        rows = _dedupe_rows(items)
        if not rows:
            continue
        if month == "未知月份":
            fname = "未知月份_财务数据.xlsx"
        else:
            fname = f"{month}_财务数据.xlsx"
        out = paths.learn_pool_by_month / fname
        _format_pool_workbook(
            out,
            rows,
            paths.classify_data,
            title=f"月份学习池 · {month}",
        )
        month_files += 1
        key = f"02_按月份/{fname}"
        files_snap[key] = {}
        for r in rows:
            unique_summ.add(r["摘要"])
            files_snap[key][r["摘要"]] = {
                "一级科目": r["一级科目"],
                "二级科目": r["二级科目"],
                "三级科目": r["三级科目"],
            }
        print(f"  月份 {fname}: {len(rows)} 摘要")

    save_snapshot(paths.learn_pool_snapshot, files_snap)

    # 确保确认表存在（空表头）
    if not paths.learn_pool_confirmed.is_file():
        _write_confirmed_excel(paths.learn_pool_confirmed, [])

    stats = {
        "unconfirmed_detail_lines": len(lines),
        "all_summary_rows": len(all_rows),
        "project_files": project_files,
        "month_files": month_files,
        "unique_summaries": len(unique_summ),
        "confirmed_excluded": len(confirmed_keys),
        "pool_root": str(paths.learn_pool_root),
        "all_xlsx": str(paths.learn_pool_all_xlsx),
    }
    print(
        f"✅ 导出学习池: 汇总 {len(all_rows)} · 项目文件 {project_files} · "
        f"月份文件 {month_files} · 摘要 {len(unique_summ)}"
    )
    print(f"   目录: {paths.learn_pool_root}")
    print(f"   快速调整: {paths.learn_pool_all_xlsx}")
    return stats


def _write_confirmed_excel(path: Path, rows: List[Dict[str, Any]]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = [
        "确认状态",
        "摘要",
        "一级科目",
        "二级科目",
        "三级科目",
        "供应商名称",
        "出现次数",
        "确认时间",
        "来源",
    ]
    # 归类排序：一级 → 二级 → 三级 → 摘要
    rows = _sort_by_category(list(rows)) if rows else []
    df = pd.DataFrame(rows, columns=cols) if rows else pd.DataFrame(columns=cols)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="已确认", index=False)
        ws = writer.book["已确认"]
        ws.freeze_panes = "A2"
        if len(df):
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"
        for i, w in enumerate([10, 36, 16, 18, 16, 22, 10, 18, 24], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _read_pool_file(path: Path) -> List[Dict[str, Any]]:
    """优先读「学习池」sheet；兼容旧文件把说明放第一页的情况。"""
    path = Path(path)
    df = None
    try:
        xls = pd.ExcelFile(path)
        sheet = POOL_SHEET if POOL_SHEET in xls.sheet_names else None
        if sheet is None:
            # 找含「摘要」列的 sheet
            for sn in xls.sheet_names:
                if sn in {"说明", "_分类选项"}:
                    continue
                tmp = pd.read_excel(xls, sheet_name=sn, nrows=1)
                cols = {str(c).strip() for c in tmp.columns}
                if "摘要" in cols:
                    sheet = sn
                    break
        if sheet is None:
            sheet = xls.sheet_names[0]
        df = pd.read_excel(xls, sheet_name=sheet)
    except Exception as e:
        print(f"  ⚠️ 无法读取 {path.name}: {e}")
        return []

    # 兼容列名空格
    df.columns = [str(c).strip() for c in df.columns]
    rows = []
    for _, row in df.iterrows():
        summ = safe_text(row.get("摘要"))
        if not summ:
            continue
        status = safe_text(row.get("确认状态")) or STATUS_PENDING
        rows.append(
            {
                "确认状态": status,
                "摘要": summ,
                "一级科目": _norm_cat(row.get("一级科目")),
                "二级科目": _norm_cat(row.get("二级科目")),
                "三级科目": _norm_cat(row.get("三级科目")),
                "供应商名称": safe_text(row.get("供应商名称")),
                "出现次数": row.get("出现次数") or "",
                "_file": path.name,
                "_path": str(path),
                "_mtime": path.stat().st_mtime,
            }
        )
    return rows


def _scan_pool_dirs(paths: Paths) -> List[Dict[str, Any]]:
    """扫描 00_汇总 / 01_按项目 / 02_按月份。"""
    rows: List[Dict[str, Any]] = []
    for d in (
        paths.learn_pool_all,
        paths.learn_pool_by_project,
        paths.learn_pool_by_month,
    ):
        if not d.is_dir():
            continue
        for f in sorted(d.glob("*.xlsx")):
            if f.name.startswith("~$"):
                continue
            rows.extend(_read_pool_file(f))
    return rows


def apply_learn_pool(
    paths: Paths,
    *,
    reclassify: bool = True,
) -> Dict[str, Any]:
    """
    检测手改/确认 → 写记忆 → 合并已确认表 → 从 01/02 剥离 → reclassify B_std。
    """
    ensure_code_root()
    from learn_classify.engine import ClassifyEngine
    from learn_classify.learn import learn_one
    from learn_classify.review import reclassify_std_dir
    from learn_classify.store import CorrectionStore, MemoryStore
    from learn_classify.taxonomy import Taxonomy

    paths.ensure_output_parents()
    snapshot = load_snapshot(paths.learn_pool_snapshot)
    pool_rows = _scan_pool_dirs(paths)
    snap_rows = sum(len(v) for v in snapshot.values())
    print(f"扫描学习池行: {len(pool_rows)} · 快照文件 {len(snapshot)} · 快照行 {snap_rows}")

    def _snap_key_for_row(it: Dict[str, Any]) -> str:
        """根据文件所在目录还原导出时的 snapshot key。"""
        p = Path(it.get("_path") or "")
        name = it.get("_file") or p.name
        parent = p.parent.name if p.parent else ""
        if parent == "00_汇总" or "00_汇总" in str(p):
            return f"00_汇总/{name}"
        if parent == "01_按项目" or "01_按项目" in str(p):
            return f"01_按项目/{name}"
        if parent == "02_按月份" or "02_按月份" in str(p):
            return f"02_按月份/{name}"
        for prefix in ("00_汇总", "01_按项目", "02_按月份"):
            k = f"{prefix}/{name}"
            if k in snapshot:
                return k
        return name

    # 按摘要聚合候选锁定
    by_summ: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for r in pool_rows:
        by_summ[r["摘要"]].append(r)

    locked: Dict[str, Dict[str, Any]] = {}
    conflicts: List[Dict[str, Any]] = []

    for summ, items in by_summ.items():
        candidates = []
        for it in items:
            sk = _snap_key_for_row(it)
            file_snap = snapshot.get(sk) or snapshot.get("__global__") or {}
            snap = file_snap.get(summ) or {}
            snap_t = _triple(snap.get("一级科目"), snap.get("二级科目"), snap.get("三级科目"))
            cur_t = _triple(it["一级科目"], it["二级科目"], it["三级科目"])
            confirmed = _is_confirmed(it["确认状态"])
            has_snap = bool(snap)
            changed = has_snap and cur_t != snap_t and (cur_t[0] or cur_t[1] or cur_t[2])
            # 无该文件快照时：仅「确认」才锁定
            if not has_snap and not confirmed:
                continue
            if confirmed or changed:
                if not cur_t[1] and not cur_t[2]:
                    continue
                candidates.append(
                    {
                        **it,
                        "_confirmed": confirmed,
                        "_changed": changed,
                        "_triple": cur_t,
                    }
                )

        if not candidates:
            continue

        # 优先级：确认 > 手改；同级取 mtime 最新
        candidates.sort(
            key=lambda x: (1 if x["_confirmed"] else 0, x["_mtime"]),
            reverse=True,
        )
        best = candidates[0]
        conf_set = {c["_triple"] for c in candidates if c["_confirmed"]}
        if len(conf_set) > 1:
            conflicts.append(
                {
                    "摘要": summ,
                    "冲突说明": "多处确认且科目不一致",
                    "采用": f"{best['_triple'][0]}/{best['_triple'][1]}/{best['_triple'][2]}",
                    "来源": best.get("_file"),
                }
            )
        locked[summ] = best

    print(f"锁定摘要: {len(locked)} · 冲突: {len(conflicts)}")

    # 学习
    tax = Taxonomy.load(paths.classify_data / "taxonomy.json")
    memory = MemoryStore(paths.classify_data / "memory.jsonl")
    corrections = CorrectionStore(paths.classify_data / "corrections.jsonl")
    learned = 0
    learn_errors = 0
    for summ, it in locked.items():
        l1, l2, l3 = it["_triple"]
        try:
            learn_one(
                taxonomy=tax,
                memory=memory,
                corrections=corrections,
                vendor_stats=None,
                subject="",
                summary=summ,
                vendor=it.get("供应商名称") or "",
                new_l1=l1,
                new_l2=l2 or l3,
                new_l3=l3,
                source="learn_pool",
                notes=f"file={it.get('_file')};confirmed={it.get('_confirmed')}",
                also_summary_only=True,
            )
            learned += 1
        except Exception as e:
            learn_errors += 1
            print(f"  ⚠️ 学习失败 {summ[:30]}: {e}")

    print(f"写入记忆: {learned} · 失败: {learn_errors}")

    # 合并已确认表
    confirmed = load_confirmed_summaries(paths.learn_pool_confirmed)
    for summ, it in locked.items():
        l1, l2, l3 = it["_triple"]
        confirmed[summ] = {
            "确认状态": STATUS_CONFIRMED,
            "摘要": summ,
            "一级科目": l1,
            "二级科目": l2,
            "三级科目": l3,
            "供应商名称": it.get("供应商名称") or "",
            "出现次数": it.get("出现次数") or "",
            "确认时间": _now(),
            "来源": it.get("_file") or "",
        }
    conf_rows = list(confirmed.values())
    conf_rows.sort(key=lambda r: r.get("确认时间") or "", reverse=True)
    _write_confirmed_excel(paths.learn_pool_confirmed, conf_rows)
    print(f"已确认表: {paths.learn_pool_confirmed} ({len(conf_rows)} 行)")

    if conflicts:
        try:
            pd.DataFrame(conflicts).to_excel(paths.learn_pool_conflicts, index=False)
            print(f"冲突记录: {paths.learn_pool_conflicts}")
        except Exception as e:
            print(f"  ⚠️ 写冲突表失败: {e}")

    # 从 00/01/02 剥离锁定摘要
    locked_keys = set(locked.keys())
    stripped_files = 0
    for d in (
        paths.learn_pool_all,
        paths.learn_pool_by_project,
        paths.learn_pool_by_month,
    ):
        if not d.is_dir():
            continue
        for f in sorted(d.glob("*.xlsx")):
            if f.name.startswith("~$"):
                continue
            rows = _read_pool_file(f)
            if not rows:
                continue
            kept = [r for r in rows if r["摘要"] not in locked_keys]
            removed = len(rows) - len(kept)
            if removed == 0:
                continue
            # 写回（保持科目归类排序）
            out_rows = _sort_by_category(
                [
                    {
                        "确认状态": r.get("确认状态") or STATUS_PENDING,
                        "摘要": r["摘要"],
                        "一级科目": r["一级科目"],
                        "二级科目": r["二级科目"],
                        "三级科目": r["三级科目"],
                        "供应商名称": r.get("供应商名称") or "",
                        "出现次数": r.get("出现次数") or "",
                    }
                    for r in kept
                ]
            )
            if not out_rows:
                try:
                    f.unlink()
                    print(f"  🗑️ 已空删除 {f.name}")
                except OSError:
                    _format_pool_workbook(f, [], paths.classify_data)
            else:
                _format_pool_workbook(f, out_rows, paths.classify_data)
            stripped_files += 1
            print(f"  剥离 {f.name}: -{removed} 行，剩 {len(out_rows)}")

    # 用新记忆 + 相似匹配刷新未确认池的预分类建议
    pool_refresh = {}
    eng = None
    if learned:
        print("刷新未确认池建议（memory / core / similar）…")
        eng = ClassifyEngine.from_data_dir(paths.classify_data)
        pool_refresh = refresh_pool_suggestions(paths, eng)
        print(f"池建议刷新: {pool_refresh}")

    # 更新按文件快照（剩余行，含刷新后科目）
    remaining = _scan_pool_dirs(paths)
    new_snap: Dict[str, Dict[str, Dict[str, str]]] = defaultdict(dict)
    for r in remaining:
        sk = _snap_key_for_row(r)
        new_snap[sk][r["摘要"]] = {
            "一级科目": r["一级科目"],
            "二级科目": r["二级科目"],
            "三级科目": r["三级科目"],
        }
    save_snapshot(paths.learn_pool_snapshot, dict(new_snap))

    reclass_stats = {}
    if reclassify and learned:
        print("reclassify B_std …")
        if eng is None:
            eng = ClassifyEngine.from_data_dir(paths.classify_data)
        reclass_stats = reclassify_std_dir(paths.std_dir, eng)
        print(f"reclassify: {reclass_stats}")

    stats = {
        "locked": len(locked),
        "learned": learned,
        "learn_errors": learn_errors,
        "conflicts": len(conflicts),
        "confirmed_total": len(conf_rows),
        "stripped_files": stripped_files,
        "pool_refresh": pool_refresh,
        "reclassify": reclass_stats,
    }
    print(f"✅ apply-learn-pool 完成: locked={learned}")
    return stats


# 可接受的引擎层（用于抬升未确认建议）
_REFRESH_LAYERS = {
    "memory",
    "memory_core",
    "similar_memory",
    "rule",
    "rule_hi",
    "pattern",
    "subject_map",
    "vendor",
}

def refresh_pool_suggestions(
    paths: Paths,
    eng: Any,
    *,
    min_confidence: float = 0.55,
) -> Dict[str, Any]:
    """
    用当前引擎重算 00/01/02 未确认行的 L1/L2/L3。

    - 只抬升：空科目 / 待复核类 / 或新建议置信更高且层更强
    - 不改「确认状态」
    - 写回后与快照对齐，避免下次误判为手改
    """
    files_updated = 0
    rows_scanned = 0
    rows_changed = 0
    by_layer: Counter = Counter()

    for d in (
        paths.learn_pool_all,
        paths.learn_pool_by_project,
        paths.learn_pool_by_month,
    ):
        if not d.is_dir():
            continue
        for f in sorted(d.glob("*.xlsx")):
            if f.name.startswith("~$"):
                continue
            rows = _read_pool_file(f)
            if not rows:
                continue
            changed = False
            out_rows: List[Dict[str, Any]] = []
            for r in rows:
                rows_scanned += 1
                status = r.get("确认状态") or STATUS_PENDING
                if _is_confirmed(status):
                    out_rows.append(
                        {
                            "确认状态": status,
                            "摘要": r["摘要"],
                            "一级科目": r["一级科目"],
                            "二级科目": r["二级科目"],
                            "三级科目": r["三级科目"],
                            "供应商名称": r.get("供应商名称") or "",
                            "出现次数": r.get("出现次数") or "",
                        }
                    )
                    continue

                old_t = _triple(r["一级科目"], r["二级科目"], r["三级科目"])
                try:
                    res = eng.classify(
                        "",
                        r["摘要"],
                        r.get("供应商名称") or "",
                    )
                except Exception:
                    res = None

                new_row = {
                    "确认状态": STATUS_PENDING,
                    "摘要": r["摘要"],
                    "一级科目": r["一级科目"],
                    "二级科目": r["二级科目"],
                    "三级科目": r["三级科目"],
                    "供应商名称": r.get("供应商名称") or "",
                    "出现次数": r.get("出现次数") or "",
                }

                if res is not None and res.layer in _REFRESH_LAYERS:
                    conf = float(res.confidence or 0)
                    if conf >= min_confidence:
                        new_t = (
                            _norm_cat(res.l1_name),
                            _norm_cat(res.l2_name),
                            _norm_cat(res.l3_name),
                        )
                        if new_t[1] or new_t[2]:
                            should = False
                            # 原无有效二级/三级 → 直接抬升
                            if not old_t[1] and not old_t[2]:
                                should = True
                            elif new_t != old_t:
                                # 原像待复核 / 空白一级
                                if not old_t[0] or "待" in old_t[0] or "待" in old_t[1]:
                                    should = True
                                # 相似/记忆/规则 且置信足够时覆盖旧建议
                                elif res.layer in {
                                    "memory",
                                    "memory_core",
                                    "similar_memory",
                                    "rule",
                                    "rule_hi",
                                    "pattern",
                                } and conf >= 0.65:
                                    should = True
                            if should and new_t != old_t:
                                new_row["一级科目"] = new_t[0]
                                new_row["二级科目"] = new_t[1]
                                new_row["三级科目"] = new_t[2]
                                rows_changed += 1
                                by_layer[res.layer] += 1
                                changed = True

                out_rows.append(new_row)

            if changed:
                out_rows = _sort_by_category(out_rows)
                _format_pool_workbook(
                    f,
                    out_rows,
                    paths.classify_data,
                    title="学习池 · 已根据确认结果刷新预分类建议",
                )
                files_updated += 1
                print(f"  ↻ 建议更新 {f.name}")

    return {
        "files_updated": files_updated,
        "rows_scanned": rows_scanned,
        "rows_changed": rows_changed,
        "by_layer": dict(by_layer),
    }
