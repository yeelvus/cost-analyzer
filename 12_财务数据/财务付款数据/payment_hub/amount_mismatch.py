# -*- coding: utf-8 -*-
"""
汇总 B_std 中「应付金额 ≠ 预扣税额 + 实付金额」的明细行到一份 Excel。
"""

from __future__ import annotations

from pathlib import Path
from typing import List, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Paths
from .utils import (
    iter_excel_files,
    parse_std_project_filename,
    resolve_business_groups,
    safe_text,
)

AMOUNT_TOLERANCE = 1e-6
SHEET_CANDIDATES = ("明细数据",)
PREFERRED_COLUMNS = [
    "项目编码",
    "项目名称",
    "日期",
    "年",
    "月",
    "科目",
    "摘要",
    "供应商名称",
    "应付金额",
    "预扣税额",
    "实付金额",
    "预扣+实付",
    "差额(应付-预扣实付)",
    "增值税额",
    "不含税额",
    "一级科目",
    "二级科目",
    "三级科目",
    "来源文件",
    "来源相对路径",
]


def _to_num(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").fillna(0.0)


def _read_detail_df(file_path: Path) -> pd.DataFrame:
    try:
        xls = pd.ExcelFile(file_path)
    except Exception as e:
        print(f"  ⚠️ 打开失败 {file_path.name}: {e}")
        return pd.DataFrame()

    sheet = None
    for name in SHEET_CANDIDATES:
        if name in xls.sheet_names:
            sheet = name
            break
    if sheet is None:
        # 不用万泰铢表，避免金额差 10000 倍误判
        for name in xls.sheet_names:
            if "万泰铢" in str(name):
                continue
            sheet = name
            break
    if sheet is None:
        return pd.DataFrame()

    try:
        return pd.read_excel(file_path, sheet_name=sheet)
    except Exception as e:
        print(f"  ⚠️ 读取失败 {file_path.name}: {e}")
        return pd.DataFrame()


def extract_mismatches_from_file(file_path: Path, std_dir: Path) -> pd.DataFrame:
    df = _read_detail_df(file_path)
    if df.empty:
        return pd.DataFrame()

    need = ["应付金额", "预扣税额", "实付金额"]
    missing = [c for c in need if c not in df.columns]
    if missing:
        # 至少要有应付+实付
        if "应付金额" not in df.columns or "实付金额" not in df.columns:
            return pd.DataFrame()
        if "预扣税额" not in df.columns:
            df = df.copy()
            df["预扣税额"] = 0.0

    out = df.copy()
    payable = _to_num(out["应付金额"])
    withhold = _to_num(out["预扣税额"]) if "预扣税额" in out.columns else 0.0
    actual = _to_num(out["实付金额"])
    combo = withhold + actual
    diff = payable - combo
    mask = diff.abs() > AMOUNT_TOLERANCE
    if not mask.any():
        return pd.DataFrame()

    bad = out.loc[mask].copy()
    bad["预扣+实付"] = combo.loc[mask].values
    bad["差额(应付-预扣实付)"] = diff.loc[mask].values

    file_code, file_name = parse_std_project_filename(file_path)
    if "项目编码" not in bad.columns:
        bad["项目编码"] = file_code
    else:
        bad["项目编码"] = bad["项目编码"].map(safe_text).replace("", file_code)
    if "项目名称" not in bad.columns:
        bad["项目名称"] = file_name
    else:
        bad["项目名称"] = bad["项目名称"].map(safe_text).replace("", file_name)

    try:
        rel = file_path.relative_to(std_dir).as_posix()
    except ValueError:
        rel = file_path.name
    bad["来源文件"] = file_path.name
    bad["来源相对路径"] = rel
    return bad


def collect_amount_mismatches(std_dir: Path) -> pd.DataFrame:
    std_dir = Path(std_dir)
    files = iter_excel_files(std_dir)
    print(f"扫描 std 文件数: {len(files)}")
    parts: List[pd.DataFrame] = []
    for fp in files:
        part = extract_mismatches_from_file(fp, std_dir)
        if not part.empty:
            parts.append(part)
            print(f"  ⚠ {fp.relative_to(std_dir)}: {len(part)} 行不平")

    if not parts:
        return pd.DataFrame(columns=PREFERRED_COLUMNS)

    all_df = pd.concat(parts, ignore_index=True)
    # 列顺序：优先列在前，其余附加
    front = [c for c in PREFERRED_COLUMNS if c in all_df.columns]
    rest = [c for c in all_df.columns if c not in front]
    return all_df[front + rest]


def build_project_summary(detail: pd.DataFrame) -> pd.DataFrame:
    if detail is None or detail.empty:
        return pd.DataFrame(
            columns=[
                "项目编码",
                "项目名称",
                "不平笔数",
                "应付金额合计",
                "预扣税额合计",
                "实付金额合计",
                "预扣+实付合计",
                "差额合计",
            ]
        )
    g = detail.copy()
    for col in ["应付金额", "预扣税额", "实付金额", "预扣+实付", "差额(应付-预扣实付)"]:
        if col not in g.columns:
            g[col] = 0.0
        g[col] = _to_num(g[col])

    grouped = g.groupby(["项目编码", "项目名称"], dropna=False, as_index=False).agg(
        {
            "差额(应付-预扣实付)": ["count", "sum"],
            "应付金额": "sum",
            "预扣税额": "sum",
            "实付金额": "sum",
            "预扣+实付": "sum",
        }
    )
    grouped.columns = [
        "项目编码",
        "项目名称",
        "不平笔数",
        "差额合计",
        "应付金额合计",
        "预扣税额合计",
        "实付金额合计",
        "预扣+实付合计",
    ]
    grouped = grouped[
        [
            "项目编码",
            "项目名称",
            "不平笔数",
            "应付金额合计",
            "预扣税额合计",
            "实付金额合计",
            "预扣+实付合计",
            "差额合计",
        ]
    ]
    grouped = grouped.iloc[grouped["差额合计"].abs().argsort()[::-1]].reset_index(drop=True)
    return grouped


def _format_workbook(path: Path) -> None:
    try:
        wb = load_workbook(path)
    except Exception:
        return
    header_fill = PatternFill("solid", fgColor="C65911")
    header_font = Font(bold=True, color="FFFFFF")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    for ws in wb.worksheets:
        if ws.max_row < 1:
            continue
        headers = [c.value for c in ws[1]]
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        amount_like = {
            "应付金额",
            "预扣税额",
            "实付金额",
            "预扣+实付",
            "差额(应付-预扣实付)",
            "增值税额",
            "不含税额",
            "应付金额合计",
            "预扣税额合计",
            "实付金额合计",
            "预扣+实付合计",
            "差额合计",
        }
        for col_idx, h in enumerate(headers, start=1):
            letter = get_column_letter(col_idx)
            width = 14
            if h in {"摘要", "来源相对路径", "来源文件", "项目名称"}:
                width = 28
            elif h in {"项目编码"}:
                width = 16
            ws.column_dimensions[letter].width = width
            if h in amount_like:
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row, column=col_idx)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00"
        # 明细表高亮差额列
        if "差额(应付-预扣实付)" in headers:
            dcol = headers.index("差额(应付-预扣实付)") + 1
            for row in range(2, ws.max_row + 1):
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = warn_fill
                # 差额列再强调
                ws.cell(row=row, column=dcol).font = Font(bold=True, color="C00000")
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def _write_amount_mismatch(
    std_dir: Path,
    out: Path,
    *,
    group_label: str = "",
) -> str:
    std_dir = Path(std_dir)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    label = group_label or std_dir.name

    detail = collect_amount_mismatches(std_dir)
    summary = build_project_summary(detail)

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        if detail.empty:
            pd.DataFrame(
                [{"提示": "未发现应付金额 ≠ 预扣税额+实付金额 的明细行"}]
            ).to_excel(writer, sheet_name="说明", index=False)
        else:
            detail.to_excel(writer, sheet_name="不平明细", index=False)
            summary.to_excel(writer, sheet_name="按项目汇总", index=False)
            overview = pd.DataFrame(
                [
                    {
                        "不平笔数": len(detail),
                        "涉及项目数": summary["项目编码"].nunique()
                        if not summary.empty
                        else 0,
                        "应付金额合计": float(_to_num(detail["应付金额"]).sum())
                        if "应付金额" in detail.columns
                        else 0,
                        "预扣+实付合计": float(_to_num(detail["预扣+实付"]).sum())
                        if "预扣+实付" in detail.columns
                        else 0,
                        "差额合计(应付-预扣实付)": float(
                            _to_num(detail["差额(应付-预扣实付)"]).sum()
                        )
                        if "差额(应付-预扣实付)" in detail.columns
                        else 0,
                        "扫描目录": str(std_dir),
                        "业务分组": label,
                        "口径": "应付金额 ≠ 预扣税额 + 实付金额（容差 1e-6）",
                    }
                ]
            )
            overview.to_excel(writer, sheet_name="总览", index=False)

    _format_workbook(out)
    print(f"  ✅ [{label}] 金额不平汇总: {out}")
    print(f"     不平明细 {len(detail)} 行 · 项目 {len(summary)} 个")
    return str(out)


def run_amount_mismatch(
    paths: Paths,
    *,
    output_path: Optional[Path] = None,
) -> List[str]:
    """
    按 B_std 业务分组分别输出不平明细。
    若显式传入 output_path，则只扫整个 std_dir 写到该路径（兼容旧调用）。
    """
    std_dir = Path(paths.std_dir)
    if not std_dir.is_dir():
        raise FileNotFoundError(f"std 目录不存在: {std_dir}")

    # 兼容：指定单一输出路径时，仍扫全量 std
    if output_path is not None:
        return [
            _write_amount_mismatch(
                std_dir, Path(output_path), group_label="全部"
            )
        ]

    groups = resolve_business_groups(paths.raw_dir, paths.std_dir)
    if not groups:
        print("未发现业务分组，跳过 amount_mismatch")
        return []

    print(f"不平明细按文件夹分组: {', '.join(g for g, _ in groups)}")
    written: List[str] = []
    for group_name, group_std in groups:
        print(f"\n{'=' * 64}")
        print(f"  分组 amount_mismatch · {group_name}  ← {group_std}")
        print(f"{'=' * 64}")
        out = _write_amount_mismatch(
            group_std,
            paths.group_amount_mismatch_xlsx(group_name),
            group_label=group_name,
        )
        written.append(out)

    print(f"\n✅ amount_mismatch 全部完成：{len(written)} 个文件")
    return written
