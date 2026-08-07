# -*- coding: utf-8 -*-
"""
步骤 summary：B_std → 汇总分析 Excel（原 4_）
"""

from __future__ import annotations

import re
from collections import OrderedDict, defaultdict
from pathlib import Path
from typing import Dict, List, Tuple, cast

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import Paths
from .deps import get_excel_layout_fn
from .utils import (
    infer_time_tag_from_path,
    iter_excel_files,
    level1_sort_key,
    safe_text,
    ymd_sort_key,
    ymd_to_month,
)

AMOUNT_COLUMNS = ["应付金额", "预扣税额", "实付金额", "增值税额", "不含税额"]
SUBJECT_WIDE_SHEET = "一级二级科目宽表"
SUBJECT_SUMMARY_ROW = "一级科目汇总"
PROJECT_LEVEL1_SHEET = "项目一级科目汇总"
EXCLUDED_SCAN_DIRS = {"C_财务付款_sum"}


def _l2_detail_header(level1: str, level2: str, used_headers: set) -> str:
    """
    生成二级明细列名，保证不与一级合计列或其它列冲突。
    - 二级与一级同名 →「一级·其他」
    - 仍冲突 → 加「一级·」前缀或序号后缀
    """
    l1 = safe_text(level1) or "未分类一级科目"
    l2 = safe_text(level2) or "未分类二级科目"
    if l2 == l1:
        candidates = [f"{l1}·其他", f"{l1}·{l2}"]
    else:
        candidates = [l2, f"{l1}·{l2}"]
    for base in candidates:
        header = base
        n = 2
        while header in used_headers:
            header = f"{base}_{n}"
            n += 1
        if header not in used_headers:
            return header
    # 兜底
    header = f"{l1}·{l2}"
    n = 2
    while header in used_headers:
        header = f"{l1}·{l2}_{n}"
        n += 1
    return header


def normalize_month_value(value, fallback_year: str = "") -> str:
    text = safe_text(value)
    if not text:
        return ""
    date_match = re.search(r"(20\d{2})[^\d]?(\d{1,2})(?:[^\d]?(\d{1,2}))?", text)
    if date_match:
        year = date_match.group(1)
        month = int(date_match.group(2))
        if 1 <= month <= 12:
            return f"{year}{month:02d}"
    pure = "".join(ch for ch in text if ch.isdigit())
    if len(pure) >= 6:
        y, m = pure[:4], pure[4:6]
        if y.startswith("20") and m.isdigit() and 1 <= int(m) <= 12:
            return f"{y}{int(m):02d}"
    if pure.isdigit() and 1 <= int(pure) <= 12 and fallback_year:
        return f"{fallback_year}{int(pure):02d}"
    return ""


def extract_month_from_dataframe(df: pd.DataFrame, fallback_ymd: str) -> str:
    fb_digits = "".join(ch for ch in safe_text(fallback_ymd) if ch.isdigit())
    fallback_year = fb_digits[:4] if len(fb_digits) >= 4 else ""
    for col in ["年月", "日期"]:
        if col not in df.columns:
            continue
        for value in df[col].dropna().tolist():
            ym = normalize_month_value(value, fallback_year=fallback_year)
            if ym:
                return ym
    if "年" in df.columns and "月" in df.columns:
        for _, row in df[["年", "月"]].dropna(how="all").iterrows():
            year_val = safe_text(row.get("年"))
            year_digits = "".join(ch for ch in year_val if ch.isdigit())
            year = year_digits[:4] if len(year_digits) >= 4 else fallback_year
            ym = normalize_month_value(safe_text(row.get("月")), fallback_year=year)
            if ym:
                return ym
    ym = normalize_month_value(fallback_ymd, fallback_year=fallback_year)
    return ym or fallback_ymd


def extract_month_from_row(row: pd.Series, fallback_ymd: str) -> str:
    fb_digits = "".join(ch for ch in safe_text(fallback_ymd) if ch.isdigit())
    fallback_year = fb_digits[:4] if len(fb_digits) >= 4 else ""
    for key in ("年月", "日期"):
        ym = normalize_month_value(row.get(key), fallback_year=fallback_year)
        if ym:
            return ym
    if "年" in row.index and "月" in row.index:
        year_val = safe_text(row.get("年"))
        year_digits = "".join(ch for ch in year_val if ch.isdigit())
        year = year_digits[:4] if len(year_digits) >= 4 else fallback_year
        ym = normalize_month_value(safe_text(row.get("月")), fallback_year=year)
        if ym:
            return ym
    ym = normalize_month_value(row.get("月"), fallback_year=fallback_year)
    if ym:
        return ym
    return normalize_month_value(fallback_ymd, fallback_year=fallback_year) or fallback_ymd


def parse_project_from_filename(file_path: Path) -> tuple[str, str]:
    stem = safe_text(file_path.stem)
    if not stem:
        return "", ""
    # 去掉尾部 _(日期)
    stem = re.sub(r"_\(\d{4}-\d{2}-\d{2}\)$", "", stem)
    match = re.match(r"^([A-Za-z0-9\-]+)[_\-\s]+(.+)$", stem)
    if not match:
        return "", stem
    return safe_text(match.group(1)), safe_text(match.group(2))


def pick_first_valid_text(series: pd.Series) -> str:
    values = series.dropna().astype(str).str.strip()
    values = values[~values.str.lower().isin(["", "nan", "none"])]
    return safe_text(values.iloc[0]) if not values.empty else ""


def _read_std_df(file_path: Path) -> pd.DataFrame:
    """优先读「明细数据」，否则第一张表。"""
    try:
        xls = pd.ExcelFile(file_path)
        sheet = "明细数据" if "明细数据" in xls.sheet_names else xls.sheet_names[0]
        return pd.read_excel(file_path, sheet_name=sheet)
    except Exception as e:
        print(f"⚠️ 读取失败 {file_path.name}: {e}")
        return pd.DataFrame()


def extract_subject_rows(file_path: Path) -> List[Dict[str, object]]:
    df = _read_std_df(file_path)
    if df.empty:
        return []
    if "实付金额" not in df.columns and "预扣税额" not in df.columns:
        print(f"⚠️ 缺少金额列: {file_path.name}")
        return []

    # 科目汇总口径：预扣税额 + 实付金额
    amount_df = pd.DataFrame()
    for col in ["预扣税额", "实付金额"]:
        amount_df[col] = (
            pd.to_numeric(df[col], errors="coerce").fillna(0.0)
            if col in df.columns
            else 0.0
        )
    df = df.copy()
    df["金额汇总"] = amount_df.sum(axis=1)
    df = df[df["金额汇总"] != 0]
    if df.empty:
        return []

    file_code, file_name = parse_project_from_filename(file_path)
    project_code = pick_first_valid_text(df["项目编码"]) if "项目编码" in df.columns else file_code
    project_name = pick_first_valid_text(df["项目名称"]) if "项目名称" in df.columns else (file_name or file_path.stem)
    if not project_code:
        project_code = file_code
    if not project_name:
        project_name = file_name or file_path.stem

    rows = []
    for _, row in df.iterrows():
        level1 = safe_text(row.get("一级科目")) or safe_text(row.get("科目")) or "未分类一级科目"
        level2 = safe_text(row.get("二级科目")) or safe_text(row.get("科目")) or "未分类二级科目"
        rows.append(
            {
                "项目编码": safe_text(project_code),
                "项目名称": safe_text(project_name),
                "一级科目": level1,
                "二级科目": level2,
                "金额汇总": float(row.get("金额汇总", 0) or 0),
            }
        )
    return rows


def extract_project_monthly_rows(file_path: Path, ymd: str) -> List[Dict]:
    df = _read_std_df(file_path)
    if df.empty or "实付金额" not in df.columns:
        if not df.empty and "实付金额" not in df.columns:
            print(f"⚠️ 缺少实付金额列: {file_path.name}")
        return []

    amt_cols = [c for c in AMOUNT_COLUMNS if c in df.columns]
    if amt_cols:
        numeric = df[amt_cols].apply(pd.to_numeric, errors="coerce").fillna(0)
        df = df[numeric.sum(axis=1) != 0]
    if df.empty:
        return []

    source_ym = extract_month_from_dataframe(df, ymd)
    file_code, file_name = parse_project_from_filename(file_path)
    project_code = pick_first_valid_text(df["项目编码"]) if "项目编码" in df.columns else file_code
    project_name = pick_first_valid_text(df["项目名称"]) if "项目名称" in df.columns else (file_name or file_path.stem)
    if not project_code:
        project_code = file_code
    if not project_name:
        project_name = file_name or file_path.stem

    for col in AMOUNT_COLUMNS:
        if col not in df.columns:
            df[col] = 0
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)

    df = df.copy()
    df["_年月日"] = df.apply(lambda r: extract_month_from_row(r, ymd), axis=1)
    fallback_month = ymd_to_month(source_ym)
    df["_月份"] = df["_年月日"].apply(lambda v: ymd_to_month(safe_text(v)) or fallback_month)
    grouped = df.groupby("_月份", dropna=False, as_index=False)[AMOUNT_COLUMNS].sum()

    result = []
    for _, g_row in grouped.iterrows():
        row = {
            "项目编码": safe_text(project_code),
            "项目名称": safe_text(project_name),
            "年月日": safe_text(g_row.get("_月份")),
        }
        for col in AMOUNT_COLUMNS:
            row[col] = float(g_row.get(col, 0) or 0)
        result.append(row)
    return result


def build_subject_wide_payload(std_dir: Path) -> Dict[str, object]:
    rows: List[Dict[str, object]] = []
    files = iter_excel_files(std_dir, exclude_dir_names=EXCLUDED_SCAN_DIRS)
    print(f"\n检测到 std 文件数: {len(files)}")
    for fp in files:
        rows.extend(extract_subject_rows(fp))
    if not rows:
        return {"project_order": [], "layout": OrderedDict(), "project_map": {}, "grand_map": {}}

    df = pd.DataFrame(rows)
    grouped = df.groupby(["项目名称", "一级科目", "二级科目"], dropna=False, as_index=False)["金额汇总"].sum()
    level1_order = sorted(grouped["一级科目"].dropna().unique().tolist(), key=level1_sort_key)
    layout: OrderedDict[str, List[str]] = OrderedDict()
    for level1 in level1_order:
        l2s = grouped.loc[grouped["一级科目"] == level1, "二级科目"].dropna().unique().tolist()
        layout[level1] = sorted(l2s, key=level1_sort_key)

    project_order = sorted(grouped["项目名称"].dropna().unique().tolist(), key=level1_sort_key)
    project_map: Dict[str, Dict[str, Dict[str, float]]] = defaultdict(lambda: defaultdict(dict))
    grand_map: Dict[str, Dict[str, float]] = defaultdict(dict)
    for _, row in grouped.iterrows():
        pn, l1, l2 = safe_text(row["项目名称"]), safe_text(row["一级科目"]), safe_text(row["二级科目"])
        amount = float(row["金额汇总"] or 0)
        project_map[pn][l1][l2] = amount
        grand_map[l1][l2] = grand_map[l1].get(l2, 0.0) + amount
    return {
        "project_order": project_order,
        "layout": layout,
        "project_map": project_map,
        "grand_map": grand_map,
    }


def build_project_level1_summary_dataframe(
    std_dir: Path,
) -> Tuple[pd.DataFrame, List[List[str]]]:
    """
    项目 × 一级科目宽表，每个一级科目合计后展开其二级类目。

    返回值第二项为各组二级明细列名（每组对应一个一级，列连续），
    供写入 Excel 时做列组合（大纲折叠）。
    """
    rows = []
    for fp in iter_excel_files(std_dir, exclude_dir_names=EXCLUDED_SCAN_DIRS):
        rows.extend(extract_subject_rows(fp))
    if not rows:
        return pd.DataFrame(columns=["项目编码", "项目名称"]), []

    df = pd.DataFrame(rows)
    df["一级科目"] = df["一级科目"].map(safe_text).replace("", "未分类一级科目")
    df["二级科目"] = df["二级科目"].map(safe_text).replace("", "未分类二级科目")

    grouped_l1 = df.groupby(
        ["项目编码", "项目名称", "一级科目"], dropna=False, as_index=False
    )["金额汇总"].sum()
    level1_order = sorted(
        grouped_l1["一级科目"].dropna().unique().tolist(), key=level1_sort_key
    )
    wide = grouped_l1.pivot_table(
        index=["项目编码", "项目名称"],
        columns="一级科目",
        values="金额汇总",
        aggfunc="sum",
        fill_value=0,
    ).reset_index()
    for col in level1_order:
        if col not in wide.columns:
            wide[col] = 0

    # 各一级下的二级明细
    used_headers: set = set(level1_order) | {"项目编码", "项目名称"}
    # l1 -> [(l2_raw, header), ...]
    l1_l2_headers: Dict[str, List[Tuple[str, str]]] = {}
    detail_groups: List[List[str]] = []

    for l1 in level1_order:
        sub = df.loc[df["一级科目"] == l1]
        if sub.empty:
            l1_l2_headers[l1] = []
            continue
        l2_order = sorted(sub["二级科目"].dropna().unique().tolist(), key=level1_sort_key)
        pairs: List[Tuple[str, str]] = []
        for l2 in l2_order:
            header = _l2_detail_header(l1, l2, used_headers)
            used_headers.add(header)
            pairs.append((l2, header))
        l1_l2_headers[l1] = pairs
        if pairs:
            detail_groups.append([h for _, h in pairs])

    # 透视为 项目×(一级,二级)，再按 header 填列
    grouped_l2 = df.groupby(
        ["项目编码", "项目名称", "一级科目", "二级科目"], dropna=False, as_index=False
    )["金额汇总"].sum()
    amount_lookup: Dict[Tuple[str, str, str, str], float] = {}
    for _, row in grouped_l2.iterrows():
        key = (
            safe_text(row["项目编码"]),
            safe_text(row["项目名称"]),
            safe_text(row["一级科目"]),
            safe_text(row["二级科目"]),
        )
        amount_lookup[key] = float(row["金额汇总"] or 0)

    # 列序：一级合计 + 紧跟其二级明细
    col_order = ["项目编码", "项目名称"]
    all_detail_headers: List[str] = []
    for l1 in level1_order:
        col_order.append(l1)
        for _l2, header in l1_l2_headers.get(l1, []):
            col_order.append(header)
            all_detail_headers.append(header)

    # 填充二级明细列
    for header in all_detail_headers:
        wide[header] = 0.0

    # 建立 header -> (l1, l2) 反查
    header_to_l1_l2: Dict[str, Tuple[str, str]] = {}
    for l1, pairs in l1_l2_headers.items():
        for l2, header in pairs:
            header_to_l1_l2[header] = (l1, l2)

    if all_detail_headers and not wide.empty:
        detail_values = {h: [] for h in all_detail_headers}
        for _, prow in wide.iterrows():
            pcode = safe_text(prow["项目编码"])
            pname = safe_text(prow["项目名称"])
            for header in all_detail_headers:
                l1, l2 = header_to_l1_l2[header]
                detail_values[header].append(
                    amount_lookup.get((pcode, pname, l1, l2), 0.0)
                )
        for header in all_detail_headers:
            wide[header] = detail_values[header]

    wide = wide[col_order]
    return (
        wide.sort_values(["项目编码", "项目名称"]).reset_index(drop=True),
        detail_groups,
    )


def apply_level2_column_groups(ws, detail_groups: List[List[str]]) -> None:
    """对各一级科目下的二级明细列启用 Excel 列组合（大纲，汇总列在左侧）。"""
    if not detail_groups:
        return
    headers = [cell.value for cell in ws[1]]
    header_to_col = {h: i for i, h in enumerate(headers, start=1) if h}

    # 合计列在明细左侧
    ws.sheet_properties.outlinePr.summaryRight = False

    detail_fill = PatternFill("solid", fgColor="F3F0E6")
    detail_font = Font(bold=True, color="5C4E2A")

    for group in detail_groups:
        cols = [header_to_col[h] for h in group if h in header_to_col]
        if not cols:
            continue
        start_col, end_col = min(cols), max(cols)
        start_letter = get_column_letter(start_col)
        end_letter = get_column_letter(end_col)
        ws.column_dimensions.group(
            start_letter, end_letter, outline_level=1, hidden=False
        )
        for col_idx in cols:
            cell = ws.cell(1, col_idx)
            cell.fill = detail_fill
            cell.font = detail_font


def build_summary_dataframe(std_dir: Path) -> pd.DataFrame:
    rows = []
    files = iter_excel_files(std_dir, exclude_dir_names=EXCLUDED_SCAN_DIRS)
    print(f"\n检测到 std 文件数: {len(files)}")
    for fp in files:
        ymd = infer_time_tag_from_path(fp, std_dir)
        rows.extend(extract_project_monthly_rows(fp, ymd))
    if not rows:
        return pd.DataFrame(columns=["项目编码", "项目名称"] + AMOUNT_COLUMNS)
    df = pd.DataFrame(rows)
    overall = df.groupby(["项目编码", "项目名称"], dropna=False, as_index=False)[AMOUNT_COLUMNS].sum()
    return overall[["项目编码", "项目名称"] + AMOUNT_COLUMNS]


def build_monthly_wide_dataframe(std_dir: Path) -> pd.DataFrame:
    rows = []
    for fp in iter_excel_files(std_dir, exclude_dir_names=EXCLUDED_SCAN_DIRS):
        ymd = infer_time_tag_from_path(fp, std_dir)
        for row in extract_project_monthly_rows(fp, ymd):
            month_key = ymd_to_month(safe_text(row.get("年月日")))
            rows.append(
                {
                    "项目编码": safe_text(row.get("项目编码")),
                    "项目名称": safe_text(row.get("项目名称")),
                    "月份": month_key,
                    # 按月宽表口径：预扣税额 + 实付金额
                    "月度金额汇总": float(row.get("预扣税额", 0) or 0)
                    + float(row.get("实付金额", 0) or 0),
                }
            )
    if not rows:
        return pd.DataFrame(columns=["项目编码", "项目名称"])
    monthly = pd.DataFrame(rows)
    grouped = monthly.groupby(["项目编码", "项目名称", "月份"], dropna=False, as_index=False)[
        ["月度金额汇总"]
    ].sum()
    month_order = sorted(grouped["月份"].dropna().unique(), key=ymd_sort_key)
    wide = grouped.pivot_table(
        index=["项目编码", "项目名称"],
        columns="月份",
        values="月度金额汇总",
        aggfunc="sum",
        fill_value=0,
    )
    wide = wide.reindex(columns=month_order, fill_value=0)
    wide.columns = [f"金额汇总_{m}" for m in wide.columns]
    return wide.reset_index()


def write_subject_wide_sheet(wb, payload: Dict[str, object]):
    layout = cast(OrderedDict, payload.get("layout", OrderedDict()))
    project_order = cast(List[str], payload.get("project_order", []))
    project_map = cast(Dict, payload.get("project_map", {}))

    ws = wb.create_sheet(title=SUBJECT_WIDE_SHEET)
    ws.freeze_panes = "C2"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr.summaryBelow = False

    title_fill = PatternFill("solid", fgColor="DDEFE3")
    header_fill = PatternFill("solid", fgColor="C9E1D1")
    summary_fill = PatternFill("solid", fgColor="EAF4EC")
    bold_font = Font(bold=True, color="1F3A2E")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="B7C7BC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.cell(1, 1, "一级科目")
    ws.cell(1, 2, "二级科目")
    for c in (1, 2):
        cell = ws.cell(1, c)
        cell.fill = title_fill
        cell.font = bold_font
        cell.alignment = center
        cell.border = border

    start_col = 3
    for idx, project_name in enumerate(project_order):
        cell = ws.cell(1, start_col + idx, project_name)
        cell.fill = title_fill
        cell.font = bold_font
        cell.alignment = center
        cell.border = border

    summary_col = start_col + len(project_order)
    head = ws.cell(1, summary_col, SUBJECT_SUMMARY_ROW)
    head.fill = title_fill
    head.font = bold_font
    head.alignment = center
    head.border = border

    row_idx = 2
    for level1, level2_list in layout.items():
        subtotal_row = row_idx
        for col, val in ((1, level1), (2, "小计")):
            cell = ws.cell(subtotal_row, col, val)
            cell.fill = header_fill
            cell.font = bold_font
            cell.alignment = center
            cell.border = border

        level1_total = 0.0
        for p_idx, project_name in enumerate(project_order):
            pmap = project_map.get(project_name, {}).get(level1, {})
            value = sum(float(pmap.get(l2, 0) or 0) for l2 in level2_list)
            cell = ws.cell(subtotal_row, start_col + p_idx, value)
            cell.number_format = "#,##0.00"
            cell.fill = header_fill
            cell.font = bold_font
            cell.alignment = center
            cell.border = border
            level1_total += value
        total_cell = ws.cell(subtotal_row, summary_col, level1_total)
        total_cell.number_format = "#,##0.00"
        total_cell.fill = header_fill
        total_cell.font = bold_font
        total_cell.alignment = center
        total_cell.border = border
        row_idx += 1

        detail_start = row_idx
        for level2 in level2_list:
            for col, val in ((1, level1), (2, level2)):
                cell = ws.cell(row_idx, col, val)
                cell.alignment = center
                cell.border = border
            l2_total = 0.0
            for p_idx, project_name in enumerate(project_order):
                value = float(
                    project_map.get(project_name, {}).get(level1, {}).get(level2, 0) or 0
                )
                cell = ws.cell(row_idx, start_col + p_idx, value)
                cell.number_format = "#,##0.00"
                cell.alignment = center
                cell.border = border
                l2_total += value
            cell = ws.cell(row_idx, summary_col, l2_total)
            cell.number_format = "#,##0.00"
            cell.alignment = center
            cell.border = border
            row_idx += 1
        detail_end = row_idx - 1
        if detail_end >= detail_start:
            ws.row_dimensions.group(detail_start, detail_end, outline_level=1, hidden=False)

    final_row = row_idx
    for col, val in ((1, "全部汇总"), (2, "-")):
        cell = ws.cell(final_row, col, val)
        cell.fill = summary_fill
        cell.font = bold_font
        cell.alignment = center
        cell.border = border
    grand = 0.0
    for p_idx, project_name in enumerate(project_order):
        project_total = 0.0
        for level1, level2_list in layout.items():
            for level2 in level2_list:
                project_total += float(
                    project_map.get(project_name, {}).get(level1, {}).get(level2, 0) or 0
                )
        cell = ws.cell(final_row, start_col + p_idx, project_total)
        cell.number_format = "#,##0.00"
        cell.fill = summary_fill
        cell.font = bold_font
        cell.alignment = center
        cell.border = border
        grand += project_total
    cell = ws.cell(final_row, summary_col, grand)
    cell.number_format = "#,##0.00"
    cell.fill = summary_fill
    cell.font = bold_font
    cell.alignment = center
    cell.border = border

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 22
    for col_idx in range(start_col, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16
    ws.row_dimensions[1].height = 30
    return ws


def apply_formatting(
    output_path: Path, detail_groups: List[List[str]] | None = None
) -> None:
    layout_fn = get_excel_layout_fn()
    detail_groups = detail_groups or []
    try:
        wb = load_workbook(output_path)
        for sheet in ["总体汇总", "项目按月宽表", PROJECT_LEVEL1_SHEET]:
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            headers = [cell.value for cell in ws[1]]
            if sheet == "项目按月宽表":
                amount_headers = [
                    h for h in headers if isinstance(h, str) and h.startswith("金额汇总_")
                ]
            elif sheet == PROJECT_LEVEL1_SHEET:
                amount_headers = [
                    h for h in headers if isinstance(h, str) and h not in {"项目编码", "项目名称"}
                ]
            else:
                amount_headers = [h for h in headers if h in AMOUNT_COLUMNS]
            layout_fn(
                ws,
                numeric_start_col=ws.max_column + 1,
                amount_headers=amount_headers,
                header_row=1,
                header_height=40,
                freeze_panes="A2",
            )
            if sheet == PROJECT_LEVEL1_SHEET and detail_groups:
                apply_level2_column_groups(ws, detail_groups)
        wb.save(output_path)
        print(f"✅ 样式应用完成: {output_path.name}")
    except Exception as e:
        print(f"⚠️ 样式应用失败: {e}")


def run_summary(paths: Paths) -> List[str]:
    std_dir = Path(paths.std_dir)
    if not std_dir.is_dir():
        raise FileNotFoundError(f"std 目录不存在: {std_dir}")

    summary_df = build_summary_dataframe(std_dir)
    monthly_wide_df = build_monthly_wide_dataframe(std_dir)
    subject_payload = build_subject_wide_payload(std_dir)
    project_level1_df, detail_groups = build_project_level1_summary_dataframe(std_dir)

    written = []
    for path in (Path(paths.sum_xlsx_z), Path(paths.sum_xlsx)):
        path.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            summary_df.to_excel(writer, sheet_name="总体汇总", index=False)
            monthly_wide_df.to_excel(writer, sheet_name="项目按月宽表", index=False)
            project_level1_df.to_excel(writer, sheet_name=PROJECT_LEVEL1_SHEET, index=False)
            write_subject_wide_sheet(writer.book, subject_payload)
        apply_formatting(path, detail_groups=detail_groups)
        written.append(str(path))
        print(f"   → {path}")

    if detail_groups:
        n_detail = sum(len(g) for g in detail_groups)
        print(f"   一级科目展开二级明细: {len(detail_groups)} 组 / 共 {n_detail} 列")
    print(f"✅ 汇总完成：项目 {len(summary_df)} 个")
    return written
