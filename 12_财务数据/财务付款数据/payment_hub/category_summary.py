# -*- coding: utf-8 -*-
"""
已废弃「分类汇总 / 分项查看 / 透视源 / 分类使用说明」。

仅保留：从 B_std 项目 Excel 中删除这些 sheet 的工具。
"""

from __future__ import annotations

from pathlib import Path
from typing import Dict, Iterable, List, Set

from openpyxl import load_workbook

# 不再生成、并应清理的 sheet
DEPRECATED_SHEETS: tuple[str, ...] = (
    "分类使用说明",
    "分类汇总",
    "分项查看",
    "透视源",
)


def remove_deprecated_sheets_from_workbook(wb, sheet_names: Iterable[str] = DEPRECATED_SHEETS) -> List[str]:
    """删除指定 sheet，返回实际删除的名称列表。"""
    removed: List[str] = []
    want: Set[str] = set(sheet_names)
    for name in list(wb.sheetnames):
        if name in want:
            del wb[name]
            removed.append(name)
    return removed


def remove_deprecated_sheets_in_file(
    excel_path: Path,
    sheet_names: Iterable[str] = DEPRECATED_SHEETS,
) -> List[str]:
    path = Path(excel_path)
    wb = load_workbook(path)
    removed = remove_deprecated_sheets_from_workbook(wb, sheet_names)
    if removed:
        wb.save(path)
    else:
        wb.close()
    return removed


def remove_deprecated_sheets_in_std_dir(std_dir: Path) -> Dict[str, int]:
    """批量清理 B_std 下项目 Excel。"""
    std_dir = Path(std_dir)
    stats = {"files": 0, "changed": 0, "removed_sheets": 0, "errors": 0}
    for f in sorted(std_dir.rglob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        stats["files"] += 1
        try:
            removed = remove_deprecated_sheets_in_file(f)
            if removed:
                stats["changed"] += 1
                stats["removed_sheets"] += len(removed)
        except Exception as exc:
            stats["errors"] += 1
            print(f"  ⚠️ 清理失败 {f.name}: {exc}")
    return stats
