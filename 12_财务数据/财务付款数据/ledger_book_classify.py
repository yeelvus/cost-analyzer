#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
账簿汇总：用付款分类学习库按摘要填写一/二/三级科目；
把账簿里已改好的科目写回学习库；再按银行账（科目编码 1002）
做「项目 × 一级科目」四列金额汇总。

分类时只改一/二/三级科目列的值，不改原表其它内容/格式。
摘要前几个字可能带项目名或与学习库略有差异：会同时用原文、去掉项目名前缀、
去掉「名称-」短前缀 三种写法分类，取更可靠的结果。

学习：按摘要去重，把已填且有效的叶子科目写入 memory.jsonl；
叶子与现有记忆相同则跳过；待复核/其他不学。默认不改账簿。

用法：
  python ledger_book_classify.py --learn
  python ledger_book_classify.py
  python ledger_book_classify.py --ledger 账簿.xlsx --out-bank 银行账_财务付款数据汇总.xlsx
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from learn_classify.engine import ClassifyEngine, ClassifyResult
from learn_classify import DEFAULT_DATA_DIR
from learn_classify.learn import learn_one
from learn_classify.normalize import memory_key
from learn_classify.store import CorrectionStore, MemoryStore
from learn_classify.taxonomy import Taxonomy
from payment_hub.utils import level1_sort_key, safe_text

DEFAULT_LEDGER = Path("/Users/wyattsun/Desktop/财务数据/账簿汇总_v7_分类整理2026-08-29_副本.xlsx")
DEFAULT_BANK_OUT = Path("/Users/wyattsun/Desktop/财务数据/银行账_财务付款数据汇总.xlsx")

AMOUNT_COLS = ("借方THB", "借方CNY", "实付金额THB", "实付金额CNY")
BANK_CODES = {"1002", "1002.0"}
PREFERRED_SHEET = "汇总"

_DASH_PREFIX_RE = re.compile(r"^[\s\-－—–﹣−]+")
_ACTION_PREFIX_RE = re.compile(
    r"^(支付|采购|发放|报销|付款|转账|收款|报|借|取现|工资|费用|付)"
)
_PENDING_LAYERS = {"pending", "low_confidence", "conflict", "error"}
_WEAK_L2 = {
    "待复核",
    "其他",
    "(空白)",
    "",
    "待分类",
    "—",
    "-",
    "未分类一级科目",
    "未分类二级科目",
}


def _cell(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in {"nan", "none", "nat"} else s


def strip_project_prefix(project_name: str, summary: str) -> str:
    """摘要以项目名称开头则去掉该前缀及随后的连接符。"""
    text = _cell(summary)
    name = _cell(project_name)
    if name and len(text) >= len(name) and text[: len(name)].casefold() == name.casefold():
        text = _DASH_PREFIX_RE.sub("", text[len(name) :])
    return text


def strip_generic_name_dash(summary: str) -> str:
    """「短名称-后续」形态：去掉连接符前的短前缀（避免把「支付-…」当项目名）。"""
    text = _cell(summary)
    for dash in ("-", "－", "—", "–"):
        if dash not in text:
            continue
        left, right = text.split(dash, 1)
        left, right = left.strip(), right.strip()
        if (
            right
            and 1 <= len(left) <= 16
            and not _ACTION_PREFIX_RE.match(left)
            and not left.isdigit()
        ):
            return right
    return text


def summary_variants(project_name: str, summary: str) -> List[str]:
    """原文 + 去项目名前缀 + 去短名称- 前缀，去重保序。"""
    raw = _cell(summary)
    out: List[str] = []
    for s in (raw, strip_project_prefix(project_name, raw), strip_generic_name_dash(raw)):
        s = _cell(s)
        if s and s not in out:
            out.append(s)
    if not out and raw:
        out.append(raw)
    return out


def _result_rank(r: ClassifyResult) -> Tuple[int, float, int]:
    weak = (
        1
        if (
            r.layer in _PENDING_LAYERS
            or r.confidence < 0.55
            or (r.l2_name or "") in _WEAK_L2
        )
        else 0
    )
    return (weak, -float(r.confidence or 0), 0 if (r.l3_name or "") else 1)


def pick_best_result(results: Iterable[ClassifyResult]) -> ClassifyResult:
    ranked = sorted(results, key=_result_rank)
    return ranked[0]


def _to_float(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        try:
            if v != v:  # NaN
                return 0.0
        except Exception:
            pass
        return float(v)
    s = str(v).strip().replace(",", "").replace("，", "").replace(" ", "")
    if not s or s.lower() in {"nan", "none", "nat", "-"}:
        return 0.0
    try:
        return float(s)
    except (TypeError, ValueError):
        return 0.0


def _is_bank_code(v) -> bool:
    s = _cell(v)
    if s in BANK_CODES:
        return True
    if s.endswith(".0") and s[:-2] == "1002":
        return True
    try:
        return int(float(s)) == 1002
    except (TypeError, ValueError):
        return False


def _header_map(ws, header_row: int = 1) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for col in range(1, (ws.max_column or 1) + 1):
        v = ws.cell(header_row, col).value
        if v is None:
            continue
        key = str(v).strip()
        if key and key not in mapping:
            mapping[key] = col
    return mapping


def _select_sheet(wb, sheet_name: Optional[str] = None):
    """指定表名优先；否则有「汇总」用汇总，否则用活动表。"""
    name = (sheet_name or "").strip()
    if name:
        if name not in wb.sheetnames:
            raise KeyError(f"找不到工作表 {name!r}，现有: {wb.sheetnames}")
        return wb[name]
    if PREFERRED_SHEET in wb.sheetnames:
        return wb[PREFERRED_SHEET]
    return wb.active


def _norm_cat(v) -> str:
    s = _cell(v)
    return "" if s in {"(空白)", "-", "—"} else s


def _triple(l1, l2, l3) -> Tuple[str, str, str]:
    return _norm_cat(l1), _norm_cat(l2), _norm_cat(l3)


def _is_weak_leaf(l2: str, l3: str) -> bool:
    leaf = l3 or l2
    return (not leaf) or leaf in _WEAK_L2


def _resolve_leaf_id(tax: Taxonomy, l1: str, l2: str, l3: str) -> str:
    """优先三级，其次二级列里的三级名/二级名。"""
    if l3 and l3 not in _WEAK_L2:
        lid = tax.resolve_l3_id(l3) or tax.resolve_leaf_id(l3)
        if lid:
            return lid
    if l2 and l2 not in _WEAK_L2:
        lid = tax.resolve_l3_id(l2) or tax.resolve_l2_id(l2) or tax.resolve_leaf_id(l2)
        if lid:
            return lid
    return ""


def _learnable_leaf(tax: Taxonomy, leaf_id: str) -> bool:
    if not leaf_id:
        return False
    if leaf_id == getattr(tax, "pending_l2_id", "") or leaf_id == "pending":
        return False
    _l1, l2, l3 = tax.display_triple(leaf_id)
    return (l3 or l2) not in _WEAK_L2


def _backup_taxonomy(tax_path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = tax_path.with_name(f"taxonomy_backup_{stamp}.json")
    shutil.copy2(tax_path, bak)
    return bak


def _patch_subject_map(data_dir: Path, name_to_leaf: Dict[str, str]) -> int:
    path = Path(data_dir) / "subject_map.json"
    if not name_to_leaf:
        return 0
    data: Dict[str, str] = {}
    if path.is_file():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            data = {}
    n = 0
    for name, lid in name_to_leaf.items():
        if name and lid and data.get(name) != lid:
            data[name] = lid
            n += 1
    if n:
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return n


def classify_ledger_excel(
    excel_path: Path,
    engine: ClassifyEngine,
    *,
    sheet_name: Optional[str] = None,
    header_row: int = 1,
) -> dict:
    """
    按摘要填写一/二/三级科目。只写这三列的值，其它单元格不动。
    每一行都分类（含仅有借方、实付为空的备用金/取现行）。
    """
    path = Path(excel_path)
    print(f"打开账簿: {path}")
    wb = load_workbook(path)
    ws = _select_sheet(wb, sheet_name)
    print(f"  工作表: {ws.title}")
    headers = _header_map(ws, header_row)
    for required in ("摘要", "一级科目", "二级科目", "三级科目"):
        if required not in headers:
            wb.close()
            raise KeyError(f"缺少列 {required!r}，当前表头: {list(headers)}")

    summ_i = headers["摘要"]
    l1_i = headers["一级科目"]
    l2_i = headers["二级科目"]
    l3_i = headers["三级科目"]
    name_i = headers.get("项目名称")
    nature_i = headers.get("支出性质")
    subj_i = headers.get("科目")  # 账簿通常没有费用科目列

    max_row = ws.max_row or header_row
    # 收集唯一 (科目提示, 摘要变体) 再分类
    jobs: List[Tuple[int, str, str, List[str]]] = []
    unique_pairs: List[Tuple[str, str]] = []
    seen_pairs: set = set()
    empty_summary = 0
    for r in range(header_row + 1, max_row + 1):
        summary = ws.cell(r, summ_i).value
        if _cell(summary) == "":
            empty_summary += 1
            jobs.append((r, "", "", []))
            continue
        pname = ws.cell(r, name_i).value if name_i else ""
        subject = ""
        if subj_i:
            subject = _cell(ws.cell(r, subj_i).value)
        if not subject and nature_i:
            subject = _cell(ws.cell(r, nature_i).value)
        variants = summary_variants(_cell(pname), _cell(summary))
        jobs.append((r, subject, _cell(summary), variants))
        for v in variants:
            key = (subject, v)
            if key not in seen_pairs:
                seen_pairs.add(key)
                unique_pairs.append(key)

    print(
        f"  行 {max_row - header_row}，空摘要 {empty_summary}，"
        f"唯一 (科目提示,摘要变体) {len(unique_pairs)}"
    )
    cache: Dict[Tuple[str, str], ClassifyResult] = {}
    for i, (subject, summ) in enumerate(unique_pairs, 1):
        cache[(subject, summ)] = engine.classify(subject, summ, None)
        if i % 2000 == 0 or i == len(unique_pairs):
            print(f"  分类进度 {i}/{len(unique_pairs)}")

    layer_count: Dict[str, int] = defaultdict(int)
    weak = 0
    filled = 0
    for r, subject, _raw, variants in jobs:
        if not variants:
            # 空摘要：三列保持空，不写「待复核」
            continue
        results = [cache[(subject, v)] for v in variants if (subject, v) in cache]
        if not results:
            continue
        best = pick_best_result(results)
        ws.cell(r, l1_i).value = best.l1_name or None
        ws.cell(r, l2_i).value = best.l2_name or None
        ws.cell(r, l3_i).value = (best.l3_name or "") or None
        filled += 1
        layer_count[best.layer or ""] += 1
        if _result_rank(best)[0] == 1:
            weak += 1

    tmp = path.with_name(f".~{path.stem}.tmp{path.suffix}")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except Exception:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
        raise
    finally:
        wb.close()

    stats = {
        "path": str(path),
        "rows": max_row - header_row,
        "filled": filled,
        "empty_summary": empty_summary,
        "weak": weak,
        "layers": dict(layer_count),
    }
    print(
        f"已写回一二三级科目: {path}\n"
        f"  填写 {filled} 行，其中待复核/弱命中 {weak} 行"
    )
    for k, v in sorted(layer_count.items(), key=lambda x: -x[1]):
        print(f"    {k or '(空)'}: {v}")
    return stats


def _iter_ledger_learn_rows(
    excel_path: Path,
    *,
    sheet_name: Optional[str] = None,
    header_row: int = 1,
) -> Tuple[str, List[dict]]:
    """读取账簿用于学习的行：摘要 / 一二三级 / 项目名称。"""
    path = Path(excel_path)
    wb = load_workbook(path, data_only=True)
    try:
        ws = _select_sheet(wb, sheet_name)
        title = ws.title
        headers = _header_map(ws, header_row)
        for required in ("摘要", "一级科目", "二级科目"):
            if required not in headers:
                raise KeyError(f"缺少列 {required!r}，当前表头: {list(headers)}")
        summ_i = headers["摘要"]
        l1_i = headers["一级科目"]
        l2_i = headers["二级科目"]
        l3_i = headers.get("三级科目")
        name_i = headers.get("项目名称")
        vend_i = headers.get("供应商名称") or headers.get("供应商") or headers.get("收款单位")
        rows: List[dict] = []
        for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
            def col(idx: Optional[int]) -> str:
                if not idx:
                    return ""
                return _cell(r[idx - 1] if idx - 1 < len(r) else None)

            summary = col(summ_i)
            if not summary:
                continue
            rows.append(
                {
                    "摘要": summary,
                    "一级科目": col(l1_i),
                    "二级科目": col(l2_i),
                    "三级科目": col(l3_i) if l3_i else "",
                    "项目名称": col(name_i) if name_i else "",
                    "供应商名称": col(vend_i) if vend_i else "",
                }
            )
        return title, rows
    finally:
        wb.close()


def learn_from_ledger(
    excel_path: Path,
    data_dir: Path,
    *,
    sheet_name: Optional[str] = None,
    auto_taxonomy: bool = True,
    dry_run: bool = False,
    header_row: int = 1,
) -> dict:
    """
    把账簿已填的一/二/三级科目写入学习库。

    - 按摘要去重；同一摘要多套科目时取出现次数最多的一套
    - 待复核 / 其他 / 空叶子 不学
    - 叶子与现有记忆相同则跳过
    - 不改账簿文件
    """
    path = Path(excel_path)
    data_dir = Path(data_dir)
    print(f"从账簿学习: {path}")
    sheet_title, rows = _iter_ledger_learn_rows(
        path, sheet_name=sheet_name, header_row=header_row
    )
    print(f"  工作表: {sheet_title} · 有摘要 {len(rows)} 行")

    by_summ: Dict[str, List[dict]] = defaultdict(list)
    for row in rows:
        by_summ[row["摘要"]].append(row)

    locked: Dict[str, Tuple[str, str, str, str, str]] = {}
    conflicts: List[Dict[str, str]] = []
    skipped_weak = 0
    for summ, items in by_summ.items():
        counts: Counter = Counter()
        sample_vend: Dict[Tuple[str, str, str], str] = {}
        sample_proj: Dict[Tuple[str, str, str], str] = {}
        for it in items:
            t = _triple(it["一级科目"], it["二级科目"], it["三级科目"])
            counts[t] += 1
            if it.get("供应商名称") and t not in sample_vend:
                sample_vend[t] = it["供应商名称"]
            if it.get("项目名称") and t not in sample_proj:
                sample_proj[t] = it["项目名称"]
        if not counts:
            continue
        learnable = [(t, n) for t, n in counts.most_common() if not _is_weak_leaf(t[1], t[2])]
        if not learnable:
            skipped_weak += 1
            continue
        best_t, best_n = learnable[0]
        others = [t for t, n in learnable if t != best_t]
        if others:
            conflicts.append(
                {
                    "摘要": summ,
                    "采用": "/".join(x for x in best_t if x),
                    "其它": " | ".join("/".join(x for x in t if x) for t in others[:4]),
                    "采用次数": str(best_n),
                }
            )
        l1, l2, l3 = best_t
        locked[summ] = (
            l1,
            l2,
            l3,
            sample_vend.get(best_t) or items[0].get("供应商名称") or "",
            sample_proj.get(best_t) or items[0].get("项目名称") or "",
        )

    print(
        f"  唯一摘要 {len(by_summ)} · 可学 {len(locked)} · "
        f"弱分类跳过 {skipped_weak} · 冲突 {len(conflicts)}"
    )

    tax_path = data_dir / "taxonomy.json"
    tax = Taxonomy.load(tax_path)
    memory = MemoryStore(data_dir / "memory.jsonl")
    corrections = CorrectionStore(data_dir / "corrections.jsonl")

    learned = 0
    skipped_unchanged = 0
    skipped_invalid = 0
    errors = 0
    tax_created_all: List[str] = []
    tax_backup: Optional[str] = None
    subject_map_hits: Dict[str, str] = {}
    seen_keys: set = set()
    would_register = 0

    total = len(locked)
    for i, (summ, (l1, l2, l3, vend, pname)) in enumerate(locked.items(), 1):
        leaf_id = _resolve_leaf_id(tax, l1, l2, l3)
        if leaf_id and not _learnable_leaf(tax, leaf_id):
            skipped_invalid += 1
            continue
        if not leaf_id and not auto_taxonomy:
            skipped_invalid += 1
            continue
        if not leaf_id:
            would_register += 1

        summaries = [summ]
        stripped = strip_project_prefix(pname, summ)
        if stripped and stripped != summ and len(stripped) >= 4:
            summaries.append(stripped)

        to_write: List[str] = []
        for s in summaries:
            key = memory_key("", s, None)
            if key in seen_keys:
                continue
            if leaf_id:
                old = memory.get(key)
                if old and old.l2_id == leaf_id:
                    seen_keys.add(key)
                    continue
            to_write.append(s)

        if not to_write:
            skipped_unchanged += 1
            if i % 2000 == 0 or i == total:
                print(f"  {'预览' if dry_run else '学习'}进度 {i}/{total}")
            continue

        if dry_run:
            for s in to_write:
                seen_keys.add(memory_key("", s, None))
            learned += 1
            if i % 2000 == 0 or i == total:
                print(f"  预览进度 {i}/{total}")
            continue

        if not leaf_id and tax_backup is None:
            bak = _backup_taxonomy(tax_path)
            tax_backup = str(bak)
            print(f"  📦 taxonomy 已备份: {bak.name}")

        try:
            last_leaf = ""
            for s in to_write:
                key = memory_key("", s, None)
                seen_keys.add(key)
                result = learn_one(
                    taxonomy=tax,
                    memory=memory,
                    corrections=corrections,
                    vendor_stats=None,
                    subject="",
                    summary=s,
                    vendor=vend,
                    new_l1=l1,
                    new_l2=l2 or l3,
                    new_l3=l3 if l2 else "",
                    source="ledger_book",
                    notes=f"sheet={sheet_title};ledger={path.name}",
                    also_summary_only=True,
                    auto_register=auto_taxonomy,
                )
                if len(result) >= 4:
                    l1n, l2n, l3n, clog = result[0], result[1], result[2], result[3]
                else:
                    l1n, l2n, l3n, clog = result[0], result[1], result[2], []
                if clog:
                    tax_created_all.extend(clog)
                    print(f"  ＋ 新科目 {s[:28]}… → {clog}")
                last_leaf = tax.resolve_l3_id(l3n) or tax.resolve_l2_id(l2n) or last_leaf
                if last_leaf and l2n:
                    subject_map_hits[l2n] = last_leaf
            learned += 1
        except Exception as e:
            errors += 1
            print(f"  ⚠️ 学习失败 {summ[:40]}: {e}")
        if i % 2000 == 0 or i == total:
            print(f"  学习进度 {i}/{total}  写入 {learned}  未变 {skipped_unchanged}")

    if (not dry_run) and tax_created_all:
        if "auto" not in (tax.version or ""):
            tax.version = f"{tax.version}+auto"
        tax.description = (
            (tax.description or "")
            + f" | ledger_book {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}: "
            + f"{len(tax_created_all)} nodes"
        ).strip(" |")
        tax.save(tax_path)
        sm_n = _patch_subject_map(data_dir, subject_map_hits)
        print(f"  💾 taxonomy 已保存（新增标记 {len(tax_created_all)} 项）")
        if sm_n:
            print(f"  💾 subject_map 更新 {sm_n} 条")
        try:
            from learn_classify.taxonomy_io import export_taxonomy_excel

            export_taxonomy_excel(data_dir, data_dir / "科目体系.xlsx")
            print("  💾 科目体系.xlsx 已刷新")
        except Exception as e:
            print(f"  ⚠️ 导出科目体系失败: {e}")

    if conflicts and not dry_run:
        out_c = data_dir / "账簿学习冲突.xlsx"
        try:
            from openpyxl import Workbook

            cwb = Workbook()
            cws = cwb.active
            cws.title = "冲突"
            cws.append(["摘要", "采用", "其它", "采用次数"])
            for rec in conflicts:
                cws.append([rec["摘要"], rec["采用"], rec["其它"], rec["采用次数"]])
            cwb.save(out_c)
            print(f"  冲突表: {out_c} ({len(conflicts)} 条)")
        except Exception as e:
            print(f"  ⚠️ 写冲突表失败: {e}")

    stats = {
        "path": str(path),
        "sheet": sheet_title,
        "rows": len(rows),
        "unique": len(by_summ),
        "locked": len(locked),
        "learned": learned,
        "skipped_weak": skipped_weak,
        "skipped_unchanged": skipped_unchanged,
        "skipped_invalid": skipped_invalid,
        "errors": errors,
        "conflicts": len(conflicts),
        "taxonomy_created": tax_created_all,
        "taxonomy_backup": tax_backup,
        "would_register": would_register,
        "dry_run": dry_run,
        "auto_taxonomy": auto_taxonomy,
        "memory_file": str(data_dir / "memory.jsonl"),
    }
    tag = "预览" if dry_run else "写入记忆"
    print(
        f"✅ 账簿学习{tag}: {learned} 条 · 叶子未变 {skipped_unchanged} · "
        f"无效 {skipped_invalid} · 失败 {errors}"
    )
    if would_register:
        action = "需自动登记" if dry_run else "已/将自动登记"
        print(f"   其中体系外、{action}约 {would_register} 条摘要")
    if tax_created_all:
        shown = ", ".join(tax_created_all[:12])
        extra = "…" if len(tax_created_all) > 12 else ""
        print(f"   新科目: {shown}{extra}")
    return stats


def _collect_ledger_rows(excel_path: Path, *, header_row: int = 1, sheet_name=None) -> List[dict]:
    path = Path(excel_path)
    wb = load_workbook(path, data_only=True)
    try:
        ws = _select_sheet(wb, sheet_name)
        headers = _header_map(ws, header_row)
        rows = []
        for r in range(header_row + 1, (ws.max_row or header_row) + 1):
            item = {name: ws.cell(r, idx).value for name, idx in headers.items()}
            rows.append(item)
        return rows
    finally:
        wb.close()


def _display_width(text: str) -> int:
    n = 0
    for ch in text:
        n += 2 if ord(ch) > 127 else 1
    return n


def _style_summary_sheet(ws, *, amount_headers: List[str], freeze: str = "A2") -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 28
    headers = [c.value for c in ws[1]]
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
    amount_set = set(amount_headers)
    for col_idx, h in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        max_len = _display_width(_cell(h))
        is_amt = h in amount_set
        for row in ws.iter_rows(
            min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx
        ):
            cell = row[0]
            if is_amt and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
            cur = _display_width(_cell(cell.value))
            if cur > max_len:
                max_len = cur
        ws.column_dimensions[letter].width = max(10, min(28, max_len + 2))
    if ws.max_row >= 2 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    last = ws[ws.max_row]
    if last and _cell(last[0].value) in {"合计", "全部合计"}:
        fill = PatternFill("solid", fgColor="D6EAF8")
        font = Font(bold=True)
        for cell in last:
            cell.fill = fill
            cell.font = font


def build_bank_project_l1_summary(
    ledger_path: Path,
    out_path: Path,
    *,
    header_row: int = 1,
    sheet_name=None,
) -> dict:
    """
    银行账（科目编码 1002）：项目 × 一级科目，分别汇总四列金额。
    不含「一级二级科目宽表」「项目按月宽表」。
    """
    rows = _collect_ledger_rows(ledger_path, header_row=header_row, sheet_name=sheet_name)
    bank = [r for r in rows if _is_bank_code(r.get("科目编码"))]
    print(f"银行账行数（科目编码 1002）: {len(bank)} / 账簿 {len(rows)}")

    # 总体：项目
    overall: Dict[Tuple[str, str], Dict[str, float]] = defaultdict(
        lambda: {c: 0.0 for c in AMOUNT_COLS}
    )
    # 项目 × 一级
    l1_map: Dict[Tuple[str, str, str], Dict[str, float]] = defaultdict(
        lambda: {c: 0.0 for c in AMOUNT_COLS}
    )
    l1_names: set = set()

    for r in bank:
        code = safe_text(r.get("项目编码"))
        name = safe_text(r.get("项目名称")) or "（无项目名称）"
        l1 = safe_text(r.get("一级科目")) or "未分类一级科目"
        key_p = (code, name)
        key_l = (code, name, l1)
        for col in AMOUNT_COLS:
            val = _to_float(r.get(col))
            overall[key_p][col] += val
            l1_map[key_l][col] += val
        l1_names.add(l1)

    def _project_sort(item):
        code, name = item[0], item[1]
        return (code or "￿", name)

    overall_keys = sorted(overall.keys(), key=_project_sort)
    l1_order = sorted(l1_names, key=level1_sort_key)

    from openpyxl import Workbook

    wb = Workbook()

    # --- 总体汇总 ---
    ws0 = wb.active
    ws0.title = "总体汇总"
    h0 = ["项目编码", "项目名称", *AMOUNT_COLS]
    ws0.append(h0)
    tot0 = {c: 0.0 for c in AMOUNT_COLS}
    for code, name in overall_keys:
        rec = overall[(code, name)]
        ws0.append([code or None, name, *[rec[c] for c in AMOUNT_COLS]])
        for c in AMOUNT_COLS:
            tot0[c] += rec[c]
    ws0.append(["合计", "", *[tot0[c] for c in AMOUNT_COLS]])
    _style_summary_sheet(ws0, amount_headers=list(AMOUNT_COLS))

    # --- 项目一级科目汇总（长表，四列金额分别合计）---
    ws1 = wb.create_sheet("项目一级科目汇总")
    h1 = ["项目编码", "项目名称", "一级科目", *AMOUNT_COLS]
    ws1.append(h1)
    tot1 = {c: 0.0 for c in AMOUNT_COLS}
    l1_keys = sorted(l1_map.keys(), key=lambda k: (_project_sort(k), level1_sort_key(k[2])))
    for code, name, l1 in l1_keys:
        rec = l1_map[(code, name, l1)]
        if all(abs(rec[c]) < 1e-12 for c in AMOUNT_COLS):
            continue
        ws1.append([code or None, name, l1, *[rec[c] for c in AMOUNT_COLS]])
        for c in AMOUNT_COLS:
            tot1[c] += rec[c]
    ws1.append(["合计", "", "", *[tot1[c] for c in AMOUNT_COLS]])
    _style_summary_sheet(ws1, amount_headers=list(AMOUNT_COLS), freeze="A2")

    # --- 每个金额一列一张「项目 × 一级科目」表（一级作列，不是二级宽表）---
    for amt_col in AMOUNT_COLS:
        title = f"项目一级科目_{amt_col}"[:31]
        ws = wb.create_sheet(title)
        headers = ["项目编码", "项目名称", *l1_order, "合计"]
        ws.append(headers)
        col_tot = {l1: 0.0 for l1 in l1_order}
        grand = 0.0
        for code, name in overall_keys:
            vals = []
            row_sum = 0.0
            for l1 in l1_order:
                v = l1_map.get((code, name, l1), {}).get(amt_col, 0.0)
                vals.append(v)
                row_sum += v
                col_tot[l1] += v
            if abs(row_sum) < 1e-12:
                continue
            ws.append([code or None, name, *vals, row_sum])
            grand += row_sum
        ws.append(["合计", "", *[col_tot[l1] for l1 in l1_order], grand])
        _style_summary_sheet(ws, amount_headers=[*l1_order, "合计"])

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_name(f".~{out_path.stem}.tmp{out_path.suffix}")
    try:
        wb.save(tmp)
        os.replace(tmp, out_path)
    except Exception:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
        raise
    print(f"银行账汇总已写入: {out_path}")
    print(
        f"  项目 {len(overall_keys)} 个 · 一级科目 {len(l1_order)} 个 · "
        + " · ".join(f"{c} {tot0[c]:,.2f}" for c in AMOUNT_COLS)
    )
    return {
        "path": str(out_path),
        "projects": len(overall_keys),
        "level1": len(l1_order),
        "amounts": tot0,
        "bank_rows": len(bank),
    }


def main(argv: Optional[List[str]] = None) -> int:
    p = argparse.ArgumentParser(
        description="账簿：学习已填科目 / 按摘要填科目 / 银行账项目一级科目汇总"
    )
    p.add_argument("--ledger", default=str(DEFAULT_LEDGER), help="账簿汇总 Excel")
    p.add_argument("--out-bank", default=str(DEFAULT_BANK_OUT), help="银行账汇总输出")
    p.add_argument("--data-dir", default=str(DEFAULT_DATA_DIR), help="learn_classify/data")
    p.add_argument(
        "--sheet",
        default="",
        help="账簿工作表名；默认有「汇总」用汇总，否则活动表",
    )
    p.add_argument("--skip-classify", action="store_true", help="跳过填科目，只做银行账汇总")
    p.add_argument("--skip-summary", action="store_true", help="只填科目，不写银行账汇总")
    p.add_argument(
        "--learn",
        action="store_true",
        help="把账簿已填一/二/三级写入学习库（按摘要去重，不改账簿）",
    )
    p.add_argument(
        "--no-auto-taxonomy",
        action="store_true",
        help="学习时不自动登记体系里没有的新科目",
    )
    p.add_argument(
        "--dry-run",
        action="store_true",
        help="仅预览学习条数，不写 memory / taxonomy",
    )
    args = p.parse_args(argv)

    ledger = Path(args.ledger).expanduser()
    if not ledger.is_file():
        print(f"❌ 找不到账簿: {ledger}")
        return 1

    data_dir = Path(args.data_dir)
    sheet = args.sheet or None

    if args.learn:
        learn_from_ledger(
            ledger,
            data_dir,
            sheet_name=sheet,
            auto_taxonomy=not args.no_auto_taxonomy,
            dry_run=args.dry_run,
        )
        # 学习默认不覆盖账簿科目、不重出汇总
        return 0

    engine = ClassifyEngine.from_data_dir(data_dir)
    print(
        f"学习库: {data_dir}  记忆 {len(engine.memory.entries) if engine.memory else 0}  "
        f"规则 {len(engine.rules)}"
    )

    if not args.skip_classify:
        classify_ledger_excel(ledger, engine, sheet_name=sheet)
    if not args.skip_summary:
        build_bank_project_l1_summary(
            ledger, Path(args.out_bank).expanduser(), sheet_name=sheet
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
