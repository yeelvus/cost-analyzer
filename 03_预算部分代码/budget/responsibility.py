import re
from pathlib import Path
from datetime import datetime
from typing import cast
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# ============================================================================
# 配置
# ============================================================================
def normalize_text(value):
    """Normalize text values; treat NaN/None/'nan' as empty string."""
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", text)

def parse_amount(value):
    """Parse numeric amounts from cells or strings."""
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_text(value)
    if not text:
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None

def normalize_header_text(value):
    """Normalize header text for fuzzy matching across spaces/newlines/brackets."""
    text = normalize_text(value)
    if not text:
        return ""
    text = text.replace("&#10;", "").replace("&#13;", "")
    text = text.replace("\n", "").replace("\r", "").replace(" ", "")
    text = text.replace("（", "(").replace("）", ")")
    return text.lower()

from budget.config import RESP_RAW_DIR, RESP_STD_DIR
from common.paths import sum_data_path
from budget.xlsx_safe import excel_cell_value, safe_save_workbook

RAW_INPUT_PATH = RESP_RAW_DIR
TARGET_BASE_PATH = RESP_STD_DIR.parent
OUTPUT_FOLDER = RESP_STD_DIR
REPORT_FILE = sum_data_path("责任成本", legacy=OUTPUT_FOLDER) / "责任成本_汇总.xlsx"
LEGACY_REPORT_FILE = OUTPUT_FOLDER / "完整处理报告.xlsx"
LOG_FILE = Path.home() / ".project_cache" / "process_log.txt"
AMOUNT_THRESHOLD = 100  # 金额阈值，≤100 为异常

# ============================================================================
# 日志函数
# ============================================================================
def log_message(msg):
    """写入日志"""
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    log_entry = f"[{timestamp}] {msg}"
    print(log_entry)
    
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(log_entry + "\n")

# ============================================================================
# 第一部分：数据处理和汇总
# ============================================================================
def is_major_category(seq_str):
    """判断序号是否为大类（首级分类）"""
    if not seq_str or not isinstance(seq_str, str):
        return False
    
    seq_str = seq_str.strip()
    
    if re.match(r'^[A-Z]$', seq_str):
        return True
    if re.match(r'^[1-9]$', seq_str):
        return True
    if re.match(r'^[一二三四五六七八九十]$', seq_str):
        return True
    
    return False

def is_category_rollup_row(name):
    """大类滚计行：小计 / 预计…小计 / 预计总成本合计 / 合同金额。

    不含分部工程合计（如「室外工程合计」「机电工程合计」），那些是直接成本明细。
    """
    text = normalize_text(name)
    if not text:
        return False
    if text in {"小计", "合计"}:
        return True
    if "合同金额" in text:
        return True
    rollup_keywords = (
        "预计成本合计",
        "预计收入合计",
        "预计总成本",
        "预计利润",
        "毛利",
    )
    if any(kw in text for kw in rollup_keywords):
        return True
    if text.endswith("小计"):
        return True
    return False


def is_subtotal_row(name):
    """兼容旧名：仅跳过大类滚计，保留分部「xx工程合计」。"""
    return is_category_rollup_row(name)


def sheet_preference_score(sheet_name):
    """多 sheet 时优先：澄清稿 > 责任成本书 > 更高报价点。"""
    name = normalize_text(sheet_name)
    score = 0
    if "澄清" in name:
        score += 100
    if "责任成本" in name:
        score += 50
    if "测算" in name:
        score += 10
    point_match = re.search(r"(\d+)\s*点", name)
    if point_match:
        score += int(point_match.group(1))
    return score

def check_amount_anomaly(category_totals):
    """检查汇总金额是否异常"""
    if not category_totals:
        return True, "无数据"
    
    total_amount = 0
    for cat, total in category_totals.items():
        if total is not None and isinstance(total, (int, float)):
            total_amount += total
            if total < 0:
                return True, f"{cat}类金额为负({total})"
    
    if total_amount <= AMOUNT_THRESHOLD:
        return True, f"汇总金额过小({total_amount})"
    
    return False, "正常"

def read_excel_data(excel_file):
    """读取 Excel 文件，提取序号、名称、金额三列。只读不写回源文件。"""
    wb = None
    try:
        wb = load_workbook(excel_file, data_only=True, keep_links=False)
        
        candidates = []
        for sheet_index, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            
            rows = []
            project_name = None
            
            for row in ws.iter_rows(values_only=False):
                row_data = []
                for cell in row:
                    row_data.append(cell)
                rows.append(row_data)
            
            header_row_idx = None
            for idx, row in enumerate(rows):
                row_values = [cell.value for cell in row]
                has_seq = "序号" in row_values
                has_name = any(name_keyword in str(v) for v in row_values if v 
                              for name_keyword in ["名称", "描述", "子项", "项目"])
                has_amount = any("合价" in str(v) or "金额" in str(v) for v in row_values if v)
                
                if has_seq and has_name and has_amount:
                    header_row_idx = idx
                    break
            
            if header_row_idx is None:
                continue
            
            header_row = [cell.value for cell in rows[header_row_idx]]
            seq_col = None
            name_col = None
            amount_col = None
            amount_cols = []
            material_col = None
            construction_col = None
            
            for col_idx, val in enumerate(header_row):
                normalized_header = normalize_header_text(val)
                if val == "序号":
                    seq_col = col_idx
                elif val and any(kw in str(val) for kw in ["名称", "描述", "子项", "项目"]):
                    if name_col is None:
                        name_col = col_idx
                elif val and ("合价" in str(val) or "金额" in str(val)):
                    amount_cols.append(col_idx)
                if normalized_header and ("材料合价" in normalized_header) and ("thb" in normalized_header or "泰铢" in normalized_header):
                    material_col = col_idx
                if normalized_header and ("施工合价" in normalized_header) and ("thb" in normalized_header or "泰铢" in normalized_header):
                    construction_col = col_idx

            if amount_cols:
                # 多金额列时智能选择：优先 THB/泰铢、综合列，并结合有效数值行数
                best_col = None
                best_score = None
                for candidate_col in amount_cols:
                    header_text = normalize_header_text(header_row[candidate_col])
                    non_null_count = 0
                    non_zero_count = 0
                    for data_row in rows[header_row_idx + 1:]:
                        cell_val = data_row[candidate_col].value if candidate_col < len(data_row) else None
                        num_val = parse_amount(cell_val)
                        if num_val is not None:
                            non_null_count += 1
                            if abs(num_val) > 1e-9:
                                non_zero_count += 1

                    preference_score = 0
                    if "综合" in header_text:
                        preference_score += 50
                    if "thb" in header_text or "泰铢" in header_text or "铢" in header_text:
                        preference_score += 30
                    if "cny" in header_text or "人民币" in header_text:
                        preference_score -= 20

                    score = preference_score + non_zero_count * 2 + non_null_count
                    if best_score is None or score > best_score:
                        best_score = score
                        best_col = candidate_col

                amount_col = best_col
            
            if seq_col is None or name_col is None or amount_col is None:
                continue
            
            if len(rows) > 0 and rows[0]:
                first_cell_value = normalize_text(rows[0][0].value)
                if first_cell_value and first_cell_value != "#REF!":
                    if "项目" in first_cell_value or "工程" in first_cell_value:
                        project_name = first_cell_value.replace("责任成本书", "").strip()
            
            header_info = {
                "seq_header": header_row[seq_col],
                "name_header": header_row[name_col],
                "amount_header": header_row[amount_col],
                "material_header": header_row[material_col] if material_col is not None else "",
                "construction_header": header_row[construction_col] if construction_col is not None else "",
                "sheet_name": sheet_name,
            }
            
            data_rows = []
            for row_idx in range(header_row_idx + 1, len(rows)):
                row = rows[row_idx]
                seq = row[seq_col].value if seq_col < len(row) else None
                name = row[name_col].value if name_col < len(row) else None
                amount = row[amount_col].value if amount_col < len(row) else None
                material_amount = row[material_col].value if material_col is not None and material_col < len(row) else None
                construction_amount = row[construction_col].value if construction_col is not None and construction_col < len(row) else None
                
                seq_str = str(seq).strip() if seq else None
                
                if name and isinstance(name, str):
                    stop_keywords = ["预计收入", "收入", "利润", "毛利率"]
                    if any(kw in name for kw in stop_keywords):
                        if seq_str and is_major_category(seq_str):
                            if seq_str not in ['A', 'B', 'D', '1', '2', '一', '二']:
                                break
                
                if seq_str is None and name is None and amount is None:
                    continue
                
                data_rows.append({
                    "序号": seq_str,
                    "名称": name,
                    "金额": amount,
                    "材料合价（THB）": material_amount,
                    "施工合价（THB）": construction_amount,
                    "row_idx": row_idx
                })
            
            if data_rows:
                candidates.append((
                    sheet_preference_score(sheet_name),
                    -sheet_index,
                    sheet_name,
                    project_name,
                    data_rows,
                    header_info,
                ))

        if candidates:
            candidates.sort(reverse=True)
            _, _, sheet_name, project_name, data_rows, header_info = candidates[0]
            log_message(f"✓ 在工作表 '{sheet_name}' 中找到有效数据")
            if len(candidates) > 1:
                others = ", ".join(item[2] for item in candidates[1:])
                log_message(f"  已忽略其他候选表: {others}")
            return project_name, data_rows, header_info
        
        log_message(f"⚠ 未找到表头行，跳过 {excel_file.name}")
        return None, [], None
    
    except Exception as e:
        log_message(f"✗ 读取文件失败 ({excel_file.name}): {e}")
        return None, [], None
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass

def _flush_category_total(category_totals, category, rollup_amount, children_sum, header_amount, has_child_amount):
    """大类金额：优先小计/滚计，其次子项合计，再次大类行自身金额。"""
    if not category:
        return
    if rollup_amount is not None:
        category_totals[category] = rollup_amount
    elif has_child_amount:
        category_totals[category] = children_sum
    elif header_amount is not None:
        category_totals[category] = header_amount
    else:
        category_totals[category] = children_sum


def process_and_summarize(data_rows):
    """处理数据行，提取所有成本大类及其子项（遇到收入/利润类大类时停止）。

    兼容两种源表：
    - 责任成本书：A / A.1 / A.2
    - 投标成本测算统计表：A 下为 0101 等分部编码、B.4 机械费、室外工程合计
    """
    result_rows = []
    category_totals = {}

    current_major_category = None
    current_category_sum = 0.0
    current_header_amount = None
    current_rollup_amount = None
    has_child_amount = False

    INCOME_STOP_KEYWORDS = ["收入", "利润", "毛利"]

    def flush_current():
        _flush_category_total(
            category_totals,
            current_major_category,
            current_rollup_amount,
            current_category_sum,
            current_header_amount,
            has_child_amount,
        )

    for row in data_rows:
        seq = row["序号"]
        name = row["名称"]
        amount = parse_amount(row["金额"])
        name_str = normalize_text(name)
        seq_str = normalize_text(seq)

        if seq is None and name is None and row.get("金额") is None:
            continue

        if is_category_rollup_row(name_str):
            if current_major_category and amount is not None:
                # 已有「小计」时，后面的光杆「合计」是 A+B+C 总计，不能覆盖大类
                if name_str == "合计" and current_rollup_amount is not None:
                    pass
                elif "预计总成本" in name_str or "预计成本合计" in name_str:
                    pass
                elif "合同金额" in name_str:
                    pass
                else:
                    current_rollup_amount = amount
            continue

        if seq_str and is_major_category(seq_str):
            flush_current()
            if any(kw in name_str for kw in INCOME_STOP_KEYWORDS):
                current_major_category = None
                break

            current_major_category = seq_str
            current_category_sum = 0.0
            current_header_amount = amount
            current_rollup_amount = None
            has_child_amount = False
            result_rows.append(row)
            continue

        if current_major_category:
            if not name_str and amount is None:
                continue
            result_rows.append(row)
            if amount is not None:
                current_category_sum += amount
                has_child_amount = True

    flush_current()
    return result_rows, category_totals

def extract_contract_amount(data_rows):
    """Extract contract amount: prefer income-category subtotal/total, else contract amount line."""
    contract_amount = None
    income_total = None

    for row in data_rows:
        name = normalize_text(row.get("名称"))
        seq = normalize_text(row.get("序号"))
        amount = parse_amount(row.get("金额"))
        if amount is not None and ((name and "合同金额" in name) or (seq and "合同金额" in seq)):
            contract_amount = amount
            break

    for idx, row in enumerate(data_rows):
        seq = normalize_text(row.get("序号"))
        name = normalize_text(row.get("名称"))
        if seq and is_major_category(str(seq)) and name and ("收入" in name):
            for j in range(idx + 1, len(data_rows)):
                next_row = data_rows[j]
                next_seq = normalize_text(next_row.get("序号"))
                if next_seq and is_major_category(str(next_seq)):
                    break
                next_name = normalize_text(next_row.get("名称"))
                label_hits = [next_seq, next_name]
                if any(lbl and ("合计" in lbl or "小计" in lbl) for lbl in label_hits):
                    income_total = parse_amount(next_row.get("金额"))
                    if income_total is not None:
                        break
            if income_total is not None:
                break

    return income_total if income_total is not None else contract_amount

def resolve_existing_report_file():
    """优先读新文件名，兼容旧版「完整处理报告.xlsx」。"""
    if REPORT_FILE.exists():
        return REPORT_FILE
    if LEGACY_REPORT_FILE.exists():
        return LEGACY_REPORT_FILE
    return REPORT_FILE


def read_old_report(report_file):
    """读取之前生成的责任成本汇总"""
    if not report_file.exists():
        return None
    try:
        wb = load_workbook(report_file, data_only=True, keep_links=False)
        try:
            ws = wb.active
            old_data = []
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                if row_idx == 1:
                    # 表头行
                    headers = row
                    continue
                # 跳过空行
                if not row[0]:
                    continue
                row_dict = {headers[col_idx]: row[col_idx] for col_idx in range(len(headers)) if col_idx < len(row)}
                old_data.append(row_dict)
            return old_data
        finally:
            wb.close()
    except Exception as e:
        log_message(f"⚠ 读取旧报告失败: {e}")
        return None

def compare_reports(old_data, new_data):
    """对比新旧报告数据，找出变化"""
    changes = []
    
    if not old_data:
        # 首次生成，所有数据都是新增
        for item in new_data:
            changes.append({
                "变化类型": "新增项目",
                "项目编码": item.get("项目编码", ""),
                "项目名称": item.get("项目名称", ""),
                "旧利润额": "N/A",
                "新利润额": item.get("利润额", ""),
                "变化幅度": "N/A",
                "异常标记": item.get("异常标记", "")
            })
        return changes
    
    # 构建旧数据的映射
    old_dict = {item.get("项目编码"): item for item in old_data if item.get("项目编码")}
    new_dict = {item.get("项目编码"): item for item in new_data if item.get("项目编码")}
    
    # 检查变化
    for code, new_item in new_dict.items():
        if code not in old_dict:
            # 新增项目
            changes.append({
                "变化类型": "新增项目",
                "项目编码": code,
                "项目名称": new_item.get("项目名称", ""),
                "旧利润额": "",
                "新利润额": new_item.get("利润额", ""),
                "变化幅度": "",
                "异常标记": new_item.get("异常标记", "")
            })
        else:
            # 对比利润额
            old_profit = pd.to_numeric(old_dict[code].get("利润额"), errors='coerce')
            new_profit = pd.to_numeric(new_item.get("利润额"), errors='coerce')
            
            if pd.isna(old_profit) or pd.isna(new_profit):
                continue
            
            if abs(old_profit - new_profit) > 0.01:  # 容差0.01
                delta = new_profit - old_profit
                changes.append({
                    "变化类型": "利润额变化",
                    "项目编码": code,
                    "项目名称": new_item.get("项目名称", ""),
                    "旧利润额": old_profit,
                    "新利润额": new_profit,
                    "变化幅度": delta,
                    "异常标记": new_item.get("异常标记", "")
                })
    
    # 检查删除的项目
    for code, old_item in old_dict.items():
        if code not in new_dict:
            changes.append({
                "变化类型": "删除项目",
                "项目编码": code,
                "项目名称": old_item.get("项目名称", ""),
                "旧利润额": old_item.get("利润额", ""),
                "新利润额": "",
                "变化幅度": "",
                "异常标记": ""
            })
    
    return changes

def write_changes_tracking(wb, changes):
    """在工作簿中追加数据变化追踪sheet"""
    try:
        # 创建新sheet
        if "数据变化追踪" in wb.sheetnames:
            ws = wb["数据变化追踪"]
        else:
            ws = wb.create_sheet("数据变化追踪")
        
        # 清空数据（保留表头）
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.value = None
        
        # 写入表头
        headers = ["变化类型", "项目编码", "项目名称", "旧利润额", "新利润额", "变化幅度", "异常标记", "记录时间"]
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx).value = header
        
        # 写入数据
        for row_idx, change in enumerate(changes, 2):
            ws.cell(row=row_idx, column=1).value = change.get("变化类型", "")
            ws.cell(row=row_idx, column=2).value = change.get("项目编码", "")
            ws.cell(row=row_idx, column=3).value = change.get("项目名称", "")
            ws.cell(row=row_idx, column=4).value = change.get("旧利润额", "")
            ws.cell(row=row_idx, column=5).value = change.get("新利润额", "")
            ws.cell(row=row_idx, column=6).value = change.get("变化幅度", "")
            ws.cell(row=row_idx, column=7).value = change.get("异常标记", "")
            ws.cell(row=row_idx, column=8).value = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # 设置格式
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_font = Font(bold=True)
        for cell in ws[1]:
            cell.alignment = header_alignment
            cell.font = header_font
        
        # 设置列宽
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 15
        ws.column_dimensions["G"].width = 12
        ws.column_dimensions["H"].width = 18
        
        # 数值格式
        for row_idx in range(2, len(changes) + 2):
            ws.cell(row=row_idx, column=4).number_format = "#,##0.00"
            ws.cell(row=row_idx, column=5).number_format = "#,##0.00"
            ws.cell(row=row_idx, column=6).number_format = "#,##0.00"
        
        return True
    except Exception as e:
        log_message(f"✗ 写入数据变化追踪失败: {e}")
        return False

def write_excel_output(data_rows, project_name, output_file, project_code, header_info):
    """将处理后的数据写入新 Excel 文件"""
    try:
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        
        def display_header(value, default=""):
            text = normalize_text(value) or default
            return text.replace("&#10;", "\n").replace("&#13;", "\n")

        row_num = 1
        # 首行直接写表头
        ws.cell(row=row_num, column=1).value = display_header(header_info.get("seq_header"), "序号")
        ws.cell(row=row_num, column=2).value = display_header(header_info.get("name_header"), "名称")
        ws.cell(row=row_num, column=3).value = display_header(
            header_info.get("amount_header"), "金额（人民币）"
        )
        ws.cell(row=row_num, column=4).value = "材料合价（THB）"
        ws.cell(row=row_num, column=5).value = "施工合价（THB）"
        row_num += 1
        
        for row_data in data_rows:
            ws.cell(row=row_num, column=1).value = excel_cell_value(row_data.get("序号"))
            ws.cell(row=row_num, column=2).value = excel_cell_value(row_data.get("名称"))

            amount_raw = row_data.get("金额")
            amount_num = parse_amount(amount_raw)
            if amount_num is not None:
                ws.cell(row=row_num, column=3).value = excel_cell_value(amount_num)
            else:
                ws.cell(row=row_num, column=3).value = excel_cell_value(amount_raw)

            material_raw = row_data.get("材料合价（THB）")
            material_num = parse_amount(material_raw)
            if material_num is not None:
                ws.cell(row=row_num, column=4).value = excel_cell_value(material_num)
            else:
                ws.cell(row=row_num, column=4).value = excel_cell_value(material_raw)

            construction_raw = row_data.get("施工合价（THB）")
            construction_num = parse_amount(construction_raw)
            if construction_num is not None:
                ws.cell(row=row_num, column=5).value = excel_cell_value(construction_num)
            else:
                ws.cell(row=row_num, column=5).value = excel_cell_value(construction_raw)
            row_num += 1

        # 金额列统一设置为带千分位逗号格式
        for r in range(2, row_num):
            ws.cell(row=r, column=3).number_format = "#,##0.00"
            ws.cell(row=r, column=4).number_format = "#,##0.00"
            ws.cell(row=r, column=5).number_format = "#,##0.00"
        
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 18
        
        safe_save_workbook(wb, output_file)
        log_message(f"✓ 生成汇总文件: {output_file.name}")
        return True
    
    except Exception as e:
        log_message(f"✗ 写入文件失败 ({output_file.name}): {e}")
        return False

# ============================================================================
# 第三部分：生成责任成本汇总
# ============================================================================
def write_complete_report(report_data, report_file):
    """
    生成责任成本汇总 Excel 文件
    包含：项目信息、Excel提取情况、处理结果
    """
    try:
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = "责任成本汇总"
        
        # 写入表头
        headers = [
            "项目编码", "项目名称",
            "A类总计", "B类总计", "合计金额", "合同金额汇总", "利润额", "毛利率", "税金补差",
            "利润额（扣除税金补差）", "扣除税金补差利润率",
            "材料合价（THB）汇总", "施工合价（THB）汇总",
            "源文件夹存在", "Excel文件数量", "主文件名", "副文件数量", "复制状态", "异常标记", "输出状态",
            "原始名称", "汇总名称"
        ]
        header_to_col = {header: idx + 1 for idx, header in enumerate(headers)}
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx).value = header  # type: ignore
        
        # 写入数据行
        for row_idx, data in enumerate(report_data, 2):
            for header, value in data.items():
                if header in header_to_col:
                    ws.cell(row=row_idx, column=header_to_col[header]).value = excel_cell_value(value)

            is_normal = data.get("异常标记", "") == "正常"
            if is_normal:
                C = get_column_letter(header_to_col['合同金额汇总'])
                T = get_column_letter(header_to_col['合计金额'])
                X = get_column_letter(header_to_col['税金补差'])
                P = get_column_letter(header_to_col['利润额（扣除税金补差）'])
                formulas = {
                    "毛利率": f"=IFERROR(({C}{row_idx}-{T}{row_idx})/{C}{row_idx},\"\")",
                    "税金补差": f"=IFERROR({C}{row_idx}/1.07*0.07-{T}{row_idx}/1.07*0.07,\"\")",
                    "利润额": f"=IFERROR({C}{row_idx}-{T}{row_idx},\"\")",
                    "利润额（扣除税金补差）": f"=IFERROR({C}{row_idx}-{T}{row_idx}-{X}{row_idx},\"\")",
                    "扣除税金补差利润率": f"=IFERROR({P}{row_idx}/{C}{row_idx},\"\")"
                }
                for fname, formula in formulas.items():
                    ws.cell(row=row_idx, column=header_to_col[fname]).value = formula  # type: ignore
        
        # 冻结首行、设置行高、筛选、自动换行
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 40
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_font = Font(bold=True)
        for cell in ws[1]:  # type: ignore
            cell.alignment = header_alignment  # type: ignore
            cell.font = header_font  # type: ignore

        # 数值格式：按业务字段类型判断金额或百分比
        amount_format = "#,##0.00"
        percent_format = "0.00%"
        percent_headers = {"毛利率", "扣除税金补差利润率"}
        money_headers = {
            "A类总计", "B类总计", "合计金额", "合同金额汇总",
            "税金补差", "利润额", "利润额（扣除税金补差）",
            "材料合价（THB）汇总", "施工合价（THB）汇总"
        }
        percent_cols = [header_to_col[h] for h in percent_headers if h in header_to_col]
        money_cols = [header_to_col[h] for h in money_headers if h in header_to_col]
        for row_idx in range(2, ws.max_row + 1):
            for col_idx in money_cols:
                ws.cell(row=row_idx, column=col_idx).number_format = amount_format
            for col_idx in percent_cols:
                ws.cell(row=row_idx, column=col_idx).number_format = percent_format

        # 列宽：指定关键列为 20.5，其余尽量小
        wide_headers = {
            "A类总计", "B类总计", "合计金额", "合同金额汇总", "毛利率",
            "税金补差", "利润额", "利润额（扣除税金补差）", "扣除税金补差利润率",
            "材料合价（THB）汇总", "施工合价（THB）汇总"
        }
        for header, col_idx in header_to_col.items():
            width = 20.5 if header in wide_headers else 12
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        
        # 保存文件
        safe_save_workbook(wb, report_file)
        log_message(f"✓ 生成责任成本汇总: {report_file.name}")
        return True
    
    except Exception as e:
        log_message(f"✗ 生成责任成本汇总失败: {e}")
        return False

# ============================================================================
# 主逻辑
# ============================================================================
def main():
    log_message("=" * 70)
    log_message("开始处理责任成本文件")
    log_message(f"步骤1: 从 raw 目录读取 Excel 文件")
    log_message(f"步骤2: 处理并生成汇总版本")
    log_message(f"步骤3: 生成责任成本汇总")
    log_message("=" * 70)
    
    # 创建目标文件夹
    TARGET_BASE_PATH.mkdir(parents=True, exist_ok=True)
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    
    # ========================================================================
    # 步骤1: 从 raw 目录读取 Excel 文件
    # ========================================================================
    log_message("\n" + "=" * 70)
    log_message("步骤1: 从 raw 目录读取 Excel 文件")
    log_message("=" * 70)
    
    report_data = []
    if not RAW_INPUT_PATH.exists():
        log_message(f"✗ raw 输入目录不存在: {RAW_INPUT_PATH}")
        return
    
    # ========================================================================
    # 步骤2: 处理 raw 目录中的 Excel 文件
    # ========================================================================
    log_message("\n" + "=" * 70)
    log_message("步骤2: 处理 raw 目录中的 Excel 文件")
    log_message("=" * 70)
    
    # 扫描 raw 目录中的所有 Excel 文件
    excel_files = list(RAW_INPUT_PATH.glob("*.xlsx"))
    excel_files = [f for f in excel_files if not f.name.startswith("~") and "_汇总版" not in f.name and "_责任成本.xlsx" not in f.name]
    
    log_message(f"找到 {len(excel_files)} 个待处理的 Excel 文件")
    
    # 创建项目编码到报告数据的映射
    report_dict = {}
    for data in report_data:
        report_dict[data["项目编码"]] = data
    
    processed_count = 0
    failed_count = 0
    anomaly_count = 0
    
    for excel_file in excel_files:
        log_message("-" * 70)
        log_message(f"处理文件: {excel_file.name}")
        
        # 提取项目信息
        project_code = None
        project_name_from_file = None
        if "_" in excel_file.name:
            parts = excel_file.name.replace(".xlsx", "").split("_", 1)
            if len(parts) == 2:
                project_code = parts[0]
                project_name_from_file = normalize_text(parts[1])
        
        # 读取 Excel 数据
        excel_project_name, data_rows, header_info = read_excel_data(excel_file)
        
        if not data_rows:
            log_message(f"⚠ 无有效数据，跳过")
            failed_count += 1
            
            # 如果项目在报告中不存在，添加新记录
            if project_code and project_code not in report_dict:
                report_dict[project_code] = {
                    "项目编码": project_code, "项目名称": project_name_from_file or "未知",
                    "A类总计": "", "B类总计": "", "合计金额": "", "合同金额汇总": "",
                    "源文件夹存在": "是", "Excel文件数量": 1, "主文件名": excel_file.name,
                    "副文件数量": 0, "复制状态": "raw直读", "异常标记": "无数据", "输出状态": "跳过",
                    "原始名称": "", "汇总名称": ""
                }
            elif project_code and project_code in report_dict:
                report_dict[project_code]["异常标记"] = "无数据"
                report_dict[project_code]["输出状态"] = "跳过"
            
            continue
        
        # 使用文件名中的项目名称
        if project_name_from_file:
            final_project_name = project_name_from_file
        elif normalize_text(excel_project_name):
            excel_project_name = normalize_text(excel_project_name)
            final_project_name = excel_project_name
        else:
            final_project_name = "项目"
        
        # 收集数据统计
        original_names = [row.get("名称", "") for row in data_rows if row.get("名称")]
        original_rows = len(data_rows)
        log_message(f"✓ 读取数据行数: {original_rows}")
        
        # 处理数据
        processed_rows, category_totals = process_and_summarize(data_rows)
        contract_amount = extract_contract_amount(data_rows)
        material_values = [
            parse_amount(row.get("材料合价（THB）"))
            for row in processed_rows
        ]
        construction_values = [
            parse_amount(row.get("施工合价（THB）"))
            for row in processed_rows
        ]
        material_total = sum(v for v in material_values if v is not None)
        construction_total = sum(v for v in construction_values if v is not None)
        material_has_numeric = any(v is not None for v in material_values)
        construction_has_numeric = any(v is not None for v in construction_values)
        has_material_col = bool(normalize_text(header_info.get("material_header", ""))) if header_info else False
        has_construction_col = bool(normalize_text(header_info.get("construction_header", ""))) if header_info else False
        
        summary_names = [row.get("名称", "") for row in processed_rows if row.get("名称")]
        deleted_names = [name for name in original_names if name not in summary_names]
        
        deleted_rows = original_rows - len(processed_rows)
        log_message(f"✓ 汇总后数据行数: {len(processed_rows)}")
        log_message(f"✓ 删除行数: {deleted_rows}")
        
        for cat, total in category_totals.items():
            if str(cat).upper() == "C":
                continue
            log_message(f"  {cat} 类总计: {total}")
        
        # 计算合计金额
        total_amount = sum(category_totals.values())
        
        # 检查金额异常
        is_anomaly, anomaly_msg = check_amount_anomaly(category_totals)
        
        # 生成输出文件名
        if project_code:
            output_filename = f"{project_code}_{final_project_name}_责任成本.xlsx"
        else:
            output_filename = f"{excel_file.stem}_责任成本.xlsx"
        
        output_file = OUTPUT_FOLDER / output_filename
        
        # 判断是否输出
        if is_anomaly:
            log_message(f"⚠ 异常：{anomaly_msg}，不输出文件")
            anomaly_count += 1
            output_status = "异常-未输出"
            failed_count += 1
        else:
            if write_excel_output(processed_rows, final_project_name, output_file, project_code, header_info):
                processed_count += 1
                output_status = "成功输出"
            else:
                output_status = "输出失败"
                failed_count += 1
        
        # 更新报告数据
        if project_code and project_code not in report_dict:
            report_dict[project_code] = {
                "项目编码": project_code, "项目名称": final_project_name,
                "A类总计": "", "B类总计": "", "合计金额": "", "合同金额汇总": "",
                    "材料合价（THB）汇总": "", "施工合价（THB）汇总": "",
                "源文件夹存在": "是", "Excel文件数量": 1, "主文件名": excel_file.name,
                "副文件数量": 0, "复制状态": "raw直读", "异常标记": "", "输出状态": "",
                "原始名称": "", "汇总名称": ""
            }
        
        if project_code and project_code in report_dict:
            report_dict[project_code].update({
                "原始名称": ", ".join(original_names),
                "汇总名称": ", ".join(summary_names),
                "A类总计": category_totals.get("A", ""),
                "B类总计": category_totals.get("B", ""),
                "合计金额": total_amount,
                "合同金额汇总": contract_amount if contract_amount is not None else "",
                "材料合价（THB）汇总": material_total if (has_material_col and material_has_numeric) else "",
                "施工合价（THB）汇总": construction_total if (has_construction_col and construction_has_numeric) else "",
                "异常标记": anomaly_msg if is_anomaly else "正常",
                "输出状态": output_status
            })
    
    # ========================================================================
    # 步骤3: 生成责任成本汇总
    # ========================================================================
    log_message("\n" + "=" * 70)
    log_message("步骤3: 生成责任成本汇总")
    log_message("=" * 70)
    
    # 读取旧报告数据（用于对比）
    old_report_data = read_old_report(resolve_existing_report_file())

    # 合并旧报告中的项目，避免本次汇总缺失历史项目
    if old_report_data:
        for old_item in old_report_data:
            code = old_item.get("项目编码")
            if code and code not in report_dict:
                report_dict[code] = old_item
                if not report_dict[code].get("输出状态"):
                    report_dict[code]["输出状态"] = "未更新"
    
    # 转换字典为列表
    final_report_data = list(report_dict.values())
    
    # 排序：先按项目编码倒序，再按利润额顺序
    # 利润额 = 合同金额汇总 - 合计金额
    # 使用稳定排序：先按利润额升序，再按项目编码降序
    final_report_data.sort(
        key=lambda x: (pd.to_numeric(x.get("合同金额汇总"), errors='coerce') or 0) - 
                      (pd.to_numeric(x.get("合计金额"), errors='coerce') or 0)
    )
    final_report_data.sort(key=lambda x: x.get("项目编码", ""), reverse=True)
    
    # 生成报告
    write_complete_report(final_report_data, REPORT_FILE)
    
    # 对比新旧数据，追加变化追踪sheet
    changes = compare_reports(old_report_data, final_report_data)
    if changes:
        log_message(f"✓ 检测到 {len(changes)} 个数据变化")
        # 打开报告文件，追加变化追踪sheet
        try:
            wb = load_workbook(REPORT_FILE, keep_links=False)
            try:
                write_changes_tracking(wb, changes)
                safe_save_workbook(wb, REPORT_FILE)
                log_message("✓ 已追加数据变化追踪sheet")
            finally:
                wb.close()
        except Exception as e:
            log_message(f"✗ 追加变化追踪失败: {e}")
    else:
        log_message("✓ 无数据变化")
    
    # 总结
    log_message("\n" + "=" * 70)
    log_message("处理完成总结:")
    log_message(f"  处理项目数: {len(report_dict)}")
    log_message(f"  有Excel文件: {sum(1 for d in final_report_data if d.get('Excel文件数量', 0) > 0)}")
    log_message(f"  无Excel文件: {sum(1 for d in final_report_data if d.get('Excel文件数量', 0) == 0)}")
    log_message(f"  成功输出汇总: {processed_count} 个文件")
    log_message(f"  异常未输出: {anomaly_count} 个文件")
    log_message(f"  处理失败: {failed_count} 个文件")
    if LEGACY_REPORT_FILE.exists() and LEGACY_REPORT_FILE != REPORT_FILE:
        try:
            LEGACY_REPORT_FILE.unlink()
            log_message(f"✓ 已移除旧文件: {LEGACY_REPORT_FILE.name}")
        except OSError as e:
            log_message(f"⚠ 未能移除旧文件 {LEGACY_REPORT_FILE.name}: {e}")

    log_message(f"  责任成本汇总: {REPORT_FILE}")
    log_message("=" * 70)

if __name__ == "__main__":
    main()
