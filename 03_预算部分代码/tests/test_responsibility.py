# -*- coding: utf-8 -*-
"""责任成本汇总：责任成本书 A.x 与测算统计表 0101/分部合计。"""

from __future__ import annotations

import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CODE = ROOT.parent
for _p in (str(CODE), str(ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from budget.responsibility import (  # noqa: E402
    extract_contract_amount,
    is_category_rollup_row,
    parse_amount,
    process_and_summarize,
    read_excel_data,
    sheet_preference_score,
    write_excel_output,
)

from common.paths import raw_data_path  # noqa: E402

RESP_RAW = raw_data_path(
    "责任成本",
    legacy=Path.home() / "Desktop/邯郸中材/01成本统计/3.1_责任成本/1_责任成本_raw",
    create=False,
)
JINGSAI = RESP_RAW / "THCC-2026-005_晶赛科技.xlsx"
TIANDUN = RESP_RAW / "THCC-2026-004_天顿科技.xlsx"


def test_rollup_vs_section_total():
    assert is_category_rollup_row("小计")
    assert is_category_rollup_row("合计")
    assert is_category_rollup_row("预计工程直接成本小计")
    assert is_category_rollup_row("预计总成本合计")
    assert is_category_rollup_row("合同金额")
    assert not is_category_rollup_row("室外工程合计")
    assert not is_category_rollup_row("机电工程合计")
    assert not is_category_rollup_row("钢结构工程合计")
    print("rollup vs section total OK")


def test_sheet_score():
    assert sheet_preference_score("投标成本测算统计表_18点 (澄清)") > sheet_preference_score(
        "投标成本测算统计表_15点"
    )
    print("sheet score OK")


def test_output_sheet_name(tmp_path):
    output_file = tmp_path / "责任成本.xlsx"
    assert write_excel_output([], "项目", output_file, "", {})

    workbook = load_workbook(output_file, read_only=True)
    try:
        assert workbook.sheetnames == ["责任成本"]
    finally:
        workbook.close()


def test_wbs_format_synthetic():
    rows = [
        {"序号": None, "名称": "合同金额", "金额": 1480.0},
        {"序号": "A", "名称": "预计工程直接成本", "金额": None},
        {"序号": "0101", "名称": "土石方工程", "金额": 100.0},
        {"序号": "B.4", "名称": "机械费用", "金额": 50.0},
        {"序号": None, "名称": "室外工程合计", "金额": 200.0},
        {"序号": None, "名称": "预计工程直接成本小计", "金额": 350.0},
        {"序号": "B", "名称": "预计间接成本", "金额": None},
        {"序号": "B.1", "名称": "管理费", "金额": 30.0},
        {"序号": None, "名称": "预计间接成本小计", "金额": 30.0},
        {"序号": "C", "名称": "风险准备金", "金额": 10.0},
        {"序号": None, "名称": "风险准备金小计", "金额": 10.0},
        {"序号": None, "名称": "预计总成本合计", "金额": 390.0},
        {"序号": "D", "名称": "预计收入(不含税)", "金额": None},
        {"序号": "D.1", "名称": "合同收入", "金额": 500.0},
        {"序号": None, "名称": "预计收入合计", "金额": 500.0},
    ]
    processed, totals = process_and_summarize(rows)
    names = [r["名称"] for r in processed]
    assert "土石方工程" in names
    assert "机械费用" in names
    assert "室外工程合计" in names
    assert "预计工程直接成本小计" not in names
    assert abs(totals["A"] - 350.0) < 1e-6
    assert abs(totals["B"] - 30.0) < 1e-6
    assert abs(totals["C"] - 10.0) < 1e-6
    assert "D" not in totals
    assert abs(extract_contract_amount(rows) - 500.0) < 1e-6
    print("synthetic WBS format OK")


def test_responsibility_book_ax_items():
    rows = [
        {"序号": "A", "名称": "预计直接成本", "金额": None},
        {"序号": "A.1", "名称": "主材", "金额": 100.0},
        {"序号": "A.2", "名称": "机械费", "金额": 20.0},
        {"序号": None, "名称": "小计", "金额": 120.0},
        {"序号": "B", "名称": "预计间接成本", "金额": None},
        {"序号": "B.1", "名称": "管理费", "金额": 10.0},
        {"序号": None, "名称": "小计", "金额": 10.0},
        {"序号": "C", "名称": "税金及附加", "金额": None},
        {"序号": "C.1", "名称": "印花税", "金额": 5.0},
        {"序号": None, "名称": "小计", "金额": 5.0},
        {"序号": None, "名称": "合计", "金额": 135.0},
        {"序号": "D", "名称": "预计收入", "金额": None},
        {"序号": "D.1", "名称": "合同收入", "金额": 200.0},
        {"序号": None, "名称": "合计", "金额": 200.0},
    ]
    processed, totals = process_and_summarize(rows)
    assert abs(totals["A"] - 120.0) < 1e-6
    assert abs(totals["B"] - 10.0) < 1e-6
    assert abs(totals["C"] - 5.0) < 1e-6
    assert "D" not in totals
    names = [r["名称"] for r in processed]
    assert "主材" in names and "小计" not in names
    print("responsibility book A.x OK")


def test_jingsai_real_file():
    if not JINGSAI.exists():
        print(f"skip real 晶赛: missing {JINGSAI}")
        return
    project_name, data_rows, header_info = read_excel_data(JINGSAI)
    assert data_rows, "晶赛未读到数据"
    assert "澄清" in str(header_info.get("sheet_name", ""))
    processed, totals = process_and_summarize(data_rows)
    contract = extract_contract_amount(data_rows)
    a_total = totals.get("A")
    b_total = totals.get("B")
    c_total = totals.get("C")
    assert a_total and abs(a_total - 110516705.76839058) < 1.0
    assert b_total and abs(b_total - 3346454.55) < 1.0
    assert c_total and abs(c_total - 1138631.6031839058) < 1.0
    assert contract and abs(contract - 138347155.6816541) < 1.0
    names = [r.get("名称") for r in processed]
    assert "土石方工程" in names
    assert "室外工程合计" in names
    assert "机电工程合计" in names
    material = sum(
        v for v in (parse_amount(r.get("材料合价（THB）")) for r in processed) if v
    )
    construction = sum(
        v for v in (parse_amount(r.get("施工合价（THB）")) for r in processed) if v
    )
    assert abs(material - 75233917.29820558) < 1.0
    assert abs(construction - 32632788.470184997) < 1.0
    print(
        f"晶赛 OK  A={a_total:.2f} B={b_total:.2f} C={c_total:.2f} "
        f"合同={contract:.2f} sheet={header_info.get('sheet_name')}"
    )


def test_tiandun_regression():
    if not TIANDUN.exists():
        print(f"skip real 天顿: missing {TIANDUN}")
        return
    _, data_rows, _ = read_excel_data(TIANDUN)
    processed, totals = process_and_summarize(data_rows)
    contract = extract_contract_amount(data_rows)
    assert abs(totals["A"] - 325917975.252829) < 1.0
    assert abs(totals["B"] - 10339583.233) < 1.0
    grand = sum(totals.values())
    assert abs(grand - 361871576.6503795) < 1.0
    assert contract and abs(contract - 382757539.67) < 1.0
    names = [r.get("名称") for r in processed]
    assert "自主施工主材供应" in names
    print(f"天顿回归 OK  A={totals['A']:.2f} 合计={grand:.2f}")


if __name__ == "__main__":
    test_rollup_vs_section_total()
    test_sheet_score()
    test_wbs_format_synthetic()
    test_responsibility_book_ax_items()
    test_jingsai_real_file()
    test_tiandun_regression()
    print("all responsibility tests OK")
