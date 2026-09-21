# -*- coding: utf-8 -*-
"""Mac Excel 兼容：openpyxl 写出后应无 inlineStr / 空 workbookProtection。"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402

from budget.xlsx_safe import safe_save_workbook, sanitize_xlsx  # noqa: E402


def test_safe_save_uses_shared_strings(tmp_path: Path | None = None):
    out = (tmp_path or Path("/tmp")) / "xlsx_safe_test.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "责任成本汇总"
    ws["A1"] = "项目编码"
    ws["B1"] = "晶赛科技"
    ws["C1"] = 110.5
    ws["D1"] = '=IFERROR(C1-1,"")'
    ws.auto_filter.ref = "A1:D2"
    safe_save_workbook(wb, out)

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "xl/sharedStrings.xml" in names
        workbook = z.read("xl/workbook.xml").decode("utf-8")
        assert "workbookProtection" not in workbook
        assert "_xlnm._FilterDatabase" not in workbook
        assert 'fullCalcOnLoad="1"' not in workbook
        sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert "inlineStr" not in sheet
        sst = z.read("xl/sharedStrings.xml").decode("utf-8")
        assert "项目编码" in sst
        assert "晶赛科技" in sst
    print("safe_save sharedStrings OK", out)


def test_sanitize_idempotent(tmp_path: Path | None = None):
    out = (tmp_path or Path("/tmp")) / "xlsx_safe_test2.xlsx"
    wb = Workbook()
    wb.active["A1"] = "hello"
    safe_save_workbook(wb, out)
    sanitize_xlsx(out)
    with zipfile.ZipFile(out) as z:
        assert "inlineStr" not in z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    print("sanitize idempotent OK")


if __name__ == "__main__":
    test_safe_save_uses_shared_strings()
    test_sanitize_idempotent()
    print("all xlsx_safe tests OK")
