import os
import pandas as pd
from pathlib import Path
import importlib.util
import openpyxl
from openpyxl.styles import numbers, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import Cell
import re
from collections import defaultdict
from typing import cast

CONTRACT_COL_NAME = "总包合同额"
OWNER_VISA_COL_NAME = "业主签证"
SUBCONTRACT_CONTRACT_COL_NAME = "分包合同额"
SUBCONTRACT_VISA_COL_NAME = "分包签证"
PROCUREMENT_SUMMARY_COL_NAME = "采购数据汇总"
CONCRETE_SUMMARY_COL_NAME = "商混站混凝土汇总"
TOOL_SUMMARY_COL_NAME = "机具费汇总"
LABOR_SUMMARY_COL_NAME = "人工费汇总"
FINANCE_COL_KEYWORDS = (
    "财务分包已付款",
    "分包已付款财务数据（万泰铢）",
    "分包预付款财务数据（万泰铢）",
)
SUBCONTRACT_PAID_COL_NAME = "工程部记录分包已付款"

# ==================== 自动化指令编辑区 ====================
SCRIPT_DIR = Path(__file__).resolve().parent
CODE_ROOT_DIR = SCRIPT_DIR.parent
DEFAULT_MATCHER_DIR = CODE_ROOT_DIR / "0_项目基本信息"

CONFIG = {
"base_dir": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/2_成本统计收集/B_成本统计_std",
"output_file": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/2_成本统计收集/B_成本统计_sum/项目统计数据汇总_全月份.xlsx",
"profit_rate_chart_file": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/0_html/项目毛利率变化图.html",
"project_matcher_dir": str(DEFAULT_MATCHER_DIR),
"project_match_db_file": "",
"project_info_file": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/0_项目基本信息情况/项目基本信息表.xlsx",
"finance_payment_dir": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/12-财务数据💰",
"contract_management_dir": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/5_合同管理📑",
"contract_management_sheet": "1_总包分包汇总(全部)",
"procurement_summary_dir": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/9_采购台账📦/E_采购部采购台账_sum",
"tool_summary_dir": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/10_机具费🚜/B_机具费_sum",
"labor_summary_dir": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/11_人工费🙎🏻‍♂️/A_1_人工费台账/B_人工费台账_std",
"concrete_summary_dir": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/9.1_商混站🚛/B_商混站_std",
"subcontract_summary_dir": r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/5.1_分包合同/2_分包合同台账_std",
"finance_payment_date": "2026-03-14",
"rmb_exchange_rate": 5,
"items": {
"预计总成本": ("概况统计", "D3", "0.00"),
"毛利": ("概况统计", "D5", "0.00"),
"毛利率": ("概况统计", "D6", "%"),
"业主预计结算额": ("业主预计结算额", "C3", "0.00"),
"分包结算额": ("1分包款", "C", "0.00"),
"材料采购实际已发生": ("2材料汇总", "G5", "0.00"),
"材料采购后续预估发生": ("2材料汇总", "J5", "0.00"),
"采购部采购金额": ("2-1采购部采购", "sum(J)/10000", "0.00"),
"项目部采购金额": ("2-2项目部采购", "sum(J)/10000", "0.00"),
"材料租赁费用": ("3材料租赁", "G", "0.00"),
"材料租赁后续预估": ("3材料租赁", "I5", "0.00"),
"工资支出": ("4，5工资", "D4 + D7", "0.00"),
"机具费支出": ("6机具费", "E5", "0.00"),
"机具费后续预估": ("6机具费", "H5", "0.00"),
"人员调遣支出": ("7人员调遣", "C", "0.00"),
"人员调遣后续预估": ("7人员调遣", "D", "0.00"),
"其他费用支出": ("8其他", "C", "0.00"),
"其他费用后续预估": ("8其他", "D", "0.00"),
"临建费用支出": ("9临建费用", "C", "0.00"),
"临建费用后续预估": ("9临建费用", "D", "0.00"),
    }
}
# =========================================================

def col_letter_to_index(letter):
    index = 0
    for char in letter.upper():
        index = index * 26 + (ord(char) - ord('A') + 1)
    return index - 1

def safe_to_float(val):
    if val is None or pd.isna(val):
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0

def normalize_project_code(value):
    if value is None or pd.isna(value):
        return None
    code = str(value).strip()
    if not code:
        return None
    if re.fullmatch(r"\d+\.0+", code):
        code = code.split('.', 1)[0]
    return code

class ProjectProcessor:
    def __init__(self, config):
        self.config = config
        self.wb_cache = {}
        self.sheet_name_cache = defaultdict(dict)

    def get_workbook(self, file_path):
        path_str = str(file_path)
        if path_str not in self.wb_cache:
            try:
                self.wb_cache[path_str] = openpyxl.load_workbook(file_path, data_only=True)
            except Exception as e:
                print(f"⚠️ 加载文件失败 {file_path.name}: {e}")
                self.wb_cache[path_str] = None
        return self.wb_cache[path_str]

    def get_sheet_by_key(self, file_path, key):
        path_str = str(file_path)
        if key not in self.sheet_name_cache[path_str]:
            wb = self.get_workbook(file_path)
            if not wb:
                self.sheet_name_cache[path_str][key] = None
                return None
            for sheet in wb.worksheets:
                if key in sheet.title:
                    self.sheet_name_cache[path_str][key] = sheet.title
                    break
            else:
                self.sheet_name_cache[path_str][key] = None
        return self.sheet_name_cache[path_str][key]

    def evaluate_expression(self, file_path, sheet_key, expr):
        wb = self.get_workbook(file_path)
        if not wb:
            return 0.0
        sheet_name = self.get_sheet_by_key(file_path, sheet_key)
        if not sheet_name or sheet_name not in wb:
            return 0.0
        ws = wb[sheet_name]

        divisor = 1.0
        if '/' in expr:
            main_part, div_part = expr.rsplit('/', 1)
            try:
                divisor = float(div_part.strip())
            except:
                divisor = 1.0
        else:
            main_part = expr

        total = 0.0
        if main_part.startswith("sum("):
            internal = main_part[4:-1].strip()
            if '+' in internal:
                cells = [c.strip() for c in internal.split('+')]
                for cell in cells:
                    col_letter = ''.join([ch for ch in cell if ch.isalpha()])
                    row_num = int(''.join([ch for ch in cell if ch.isdigit()]))
                    col_idx = col_letter_to_index(col_letter) + 1
                    try:
                        total += safe_to_float(ws.cell(row=row_num, column=col_idx).value)
                    except:
                        pass
            else:
                col_letter = internal.strip()
                col_idx = col_letter_to_index(col_letter) + 1
                for row in range(2, ws.max_row + 1):
                    total += safe_to_float(ws.cell(row=row, column=col_idx).value)
        elif '+' in main_part:
            cells = [c.strip() for c in main_part.split('+')]
            for cell in cells:
                col_letter = ''.join([ch for ch in cell if ch.isalpha()])
                row_num = int(''.join([ch for ch in cell if ch.isdigit()]))
                col_idx = col_letter_to_index(col_letter) + 1
                try:
                    total += safe_to_float(ws.cell(row=row_num, column=col_idx).value)
                except:
                    pass
        else:
            col_letter = ''.join([ch for ch in main_part if ch.isalpha()])
            row_part = ''.join([ch for ch in main_part if ch.isdigit()])
            col_idx = col_letter_to_index(col_letter) + 1
            if row_part:
                row_num = int(row_part)
                try:
                    total = safe_to_float(ws.cell(row=row_num, column=col_idx).value)
                except:
                    total = 0.0
            else:
                for row in range(2, ws.max_row + 1):
                    total += safe_to_float(ws.cell(row=row, column=col_idx).value)
        return total / divisor if divisor != 0 else 0.0

def find_latest_finance_file(directory):
    base_dir_path = Path(directory)
    preferred_subdir = base_dir_path / "B_输出汇总文件"
    dir_path = preferred_subdir if preferred_subdir.exists() else base_dir_path
    if not dir_path.exists():
        print(f"⚠️ 财务数据目录不存在: {directory}")
        return None
    files = [f for f in dir_path.rglob("分包台账汇总_*.xlsx") if not f.name.startswith("~$")]
    if not files:
        files = [f for f in dir_path.rglob("*.xlsx") if not f.name.startswith("~$")]
    if not files:
        print(f"⚠️ 未找到财务汇总Excel: {dir_path}")
        return None
    ts_pattern = re.compile(r"分包台账汇总_(\d{8})_(\d{6})\.xlsx$")
    def sort_key(path_obj):
        filename = path_obj.name
        match = ts_pattern.match(filename)
        if match:
            return (1, match.group(1) + match.group(2))
        return (0, f"{os.path.getmtime(path_obj):020.6f}")
    latest_file = max(files, key=sort_key)
    print(f"✓ 自动选择最新的分包台账文件: {latest_file.name}")
    print(f" 文件路径: {latest_file}")
    return str(latest_file)

def find_latest_contract_file(directory):
    dir_path = Path(directory)
    if not dir_path.exists():
        print(f"⚠️ 合同管理目录不存在: {directory}")
        return None
    files = list(dir_path.rglob("合同管理_总包分包汇总_*.xlsx"))
    if not files:
        print(f"⚠️ 未找到合同管理汇总文件: {directory}")
        return None
    ts_pattern = re.compile(r"合同管理_总包分包汇总_(\d{8})_(\d{6})\.xlsx$")
    def sort_key(path_obj):
        filename = path_obj.name
        match = ts_pattern.match(filename)
        if match:
            return (1, match.group(1) + match.group(2))
        return (0, f"{os.path.getmtime(path_obj):020.6f}")
    latest_file = max(files, key=sort_key)
    print(f"✓ 自动选择最新的合同管理文件: {latest_file.name}")
    print(f" 文件路径: {latest_file}")
    return str(latest_file)

def find_latest_excel_file(directory):
    dir_path = Path(directory)
    if not dir_path.exists():
        print(f"⚠️ 目录不存在: {directory}")
        return None
    files = [f for f in dir_path.rglob("*.xls*") if f.is_file() and not f.name.startswith("~$")]
    if not files:
        print(f"⚠️ 未找到Excel文件: {directory}")
        return None

    ts_full = re.compile(r"(\d{8})_(\d{6})")
    ts_ymd = re.compile(r"(\d{8})")
    ts_ym = re.compile(r"(20\d{2}(0[1-9]|1[0-2]))")

    def sort_key(path_obj):
        name = path_obj.name
        m_full = ts_full.search(name)
        if m_full:
            return (3, m_full.group(1) + m_full.group(2))
        m_ymd = ts_ymd.search(name)
        if m_ymd:
            return (2, m_ymd.group(1) + "000000")
        m_ym = ts_ym.search(name)
        if m_ym:
            return (1, m_ym.group(1) + "01000000")
        return (0, f"{os.path.getmtime(path_obj):020.6f}")

    latest_file = max(files, key=sort_key)
    print(f"✓ 自动选择最新汇总文件: {latest_file.name}")
    print(f" 文件路径: {latest_file}")
    return str(latest_file)

def find_sheet_name(sheet_names, sheet_candidates):
    for candidate in sheet_candidates:
        if candidate in sheet_names:
            return candidate
    for candidate in sheet_candidates:
        for sheet_name in sheet_names:
            if candidate in str(sheet_name):
                return sheet_name
    return None

def load_project_amount_map(directory, sheet_candidates, code_col_candidates, amount_col_candidates, divisor=1.0):
    amount_map = {}
    latest_file = find_latest_excel_file(directory)
    if not latest_file:
        return amount_map
    try:
        xls = pd.ExcelFile(latest_file)
        sheet_name = find_sheet_name(xls.sheet_names, sheet_candidates)
        if not sheet_name:
            raise KeyError(f"未找到sheet，候选: {sheet_candidates}，实际: {xls.sheet_names}")
        df = pd.read_excel(latest_file, sheet_name=sheet_name)
        df.columns = [str(c).strip() for c in df.columns]
        code_col = pick_first_existing_column(df.columns, code_col_candidates)
        amount_col = pick_first_existing_column(df.columns, amount_col_candidates)
        if not code_col or not amount_col:
            raise KeyError(f"sheet缺少必要列，当前列: {list(df.columns)}")
        for _, row in df.iterrows():
            code = normalize_project_code(row.get(code_col))
            if not code:
                continue
            amount = safe_to_float(row.get(amount_col))
            if divisor:
                amount = amount / divisor
            amount_map[code] = amount
        print(f"✓ 已加载 {len(amount_map)} 条项目金额映射: {Path(latest_file).name} / {sheet_name}")
    except Exception as e:
        print(f"⚠️ 加载项目金额映射失败: {directory} -> {e}")
    return amount_map

def insert_columns_after_anchor(columns, anchor, new_columns):
    if anchor in columns:
        idx = columns.index(anchor) + 1
        for new_col in new_columns:
            if new_col in columns:
                columns.remove(new_col)
            columns.insert(idx, new_col)
            idx += 1
    else:
        for new_col in new_columns:
            if new_col not in columns:
                columns.append(new_col)

def pick_first_existing_column(columns, candidates):
    for candidate in candidates:
        if candidate in columns:
            return candidate
    return None

def find_data_root_from_base_dir(base_dir):
    base_path = Path(base_dir).expanduser().resolve()
    for candidate in [base_path, *base_path.parents]:
        if (candidate / "0_项目基本信息情况").exists():
            return candidate
    return None

def resolve_matcher_paths(config):
    matcher_dir_raw = config.get("project_matcher_dir")
    matcher_dir = Path(matcher_dir_raw).expanduser().resolve() if matcher_dir_raw else DEFAULT_MATCHER_DIR

    db_file_raw = config.get("project_match_db_file")
    if db_file_raw:
        db_path = Path(db_file_raw).expanduser().resolve()
    else:
        db_path = (matcher_dir / "project_name_match_db.json").resolve()

    # If an old/invalid DB path is configured, prefer the DB next to matcher module.
    local_db_path = (matcher_dir / "project_name_match_db.json").resolve()
    if not db_path.exists() and local_db_path.exists():
        db_path = local_db_path

    project_info_raw = config.get("project_info_file")
    if project_info_raw:
        project_info_path = Path(project_info_raw).expanduser().resolve()
    else:
        data_root = find_data_root_from_base_dir(config.get("base_dir", ""))
        project_info_path = (data_root / "0_项目基本信息情况" / "项目基本信息表.xlsx").resolve() if data_root else None

    return matcher_dir, db_path, project_info_path

def init_project_matcher(project_code_file=None, match_db_file=None, matcher_dir=None):
    if matcher_dir:
        matcher_dir = Path(matcher_dir).expanduser().resolve()
    elif match_db_file:
        matcher_dir = Path(match_db_file).expanduser().resolve().parent
    elif project_code_file:
        matcher_dir = DEFAULT_MATCHER_DIR
    else:
        matcher_dir = DEFAULT_MATCHER_DIR

    if match_db_file:
        match_db_path = Path(match_db_file).expanduser().resolve()
    else:
        match_db_path = (matcher_dir / "project_name_match_db.json").resolve()

    matcher_file = matcher_dir / "project_name_matcher.py"
    if not matcher_file.exists():
        raise FileNotFoundError(f"未找到项目匹配模块: {matcher_file}")
    module_name = "project_name_matcher_shared"
    module_spec = importlib.util.spec_from_file_location(module_name, str(matcher_file))
    if module_spec is None or module_spec.loader is None:
        raise ImportError(f"无法加载项目匹配模块: {matcher_file}")
    matcher_module = importlib.util.module_from_spec(module_spec)
    module_spec.loader.exec_module(matcher_module)

    build_standard_projects_from_dataframe = matcher_module.build_standard_projects_from_dataframe
    load_or_init_db = matcher_module.load_or_init_db
    match_project = matcher_module.match_project
    mark_unmatched = matcher_module.mark_unmatched
    save_db = matcher_module.save_db

    standard_projects = []
    project_code_path = Path(project_code_file).expanduser().resolve() if project_code_file else None
    if not match_db_path.exists() and project_code_path and project_code_path.exists():
        try:
            project_df = pd.read_excel(project_code_path)
            standard_projects = build_standard_projects_from_dataframe(project_df)
            print(f"✓ JSON匹配库不存在，已从项目基本信息初始化标准项目: {len(standard_projects)} 条")
        except Exception as e:
            print(f"⚠️ 使用项目基本信息初始化匹配库失败: {e}")

    match_db = load_or_init_db(str(match_db_path), standard_projects)

    contract_name_by_code = {}
    contract_name_by_standard_name = {}
    for standard_name, meta in match_db.get('projects', {}).items():
        std_name = str(standard_name).strip()
        if not std_name:
            continue
        code = normalize_project_code(meta.get('project_code'))
        if code:
            contract_name_by_code[code] = std_name
        contract_name_by_standard_name[std_name] = std_name

    return {
        'db': match_db,
        'db_path': str(match_db_path),
        'match_project': match_project,
        'mark_unmatched': mark_unmatched,
        'save_db': save_db,
        'contract_name_by_code': contract_name_by_code,
        'contract_name_by_standard_name': contract_name_by_standard_name,
    }

def normalize_margin_value(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return None
        is_percent_string = stripped.endswith('%')
        stripped = stripped.replace('%', '').replace(',', '')
        try:
            numeric_value = float(stripped)
        except ValueError:
            return None
        if is_percent_string:
            return numeric_value / 100.0
        return numeric_value / 100.0 if 1 < abs(numeric_value) <= 100 else numeric_value
    numeric_value = safe_to_float(value)
    return numeric_value / 100.0 if 1 < abs(numeric_value) <= 100 else numeric_value

def parse_month_to_datetime(month_str):
    match = re.match(r'^(\d{4})年(\d{1,2})月$', str(month_str).strip())
    if not match:
        return pd.NaT
    year = int(match.group(1))
    month = int(match.group(2))
    if month < 1 or month > 12:
        return pd.NaT
    return pd.Timestamp(year=year, month=month, day=1)

def parse_month_sort_key(month_text):
    match = re.match(r'^(\d{4})年(\d{1,2})月$', str(month_text).strip())
    if not match:
        return (0, 0)
    return (int(match.group(1)), int(match.group(2)))


COST_FILE_RE = re.compile(
    r"^(?P<project>.+?)_成本统计_(?P<yyyymm>\d{6})(?:\(\d+\))?\s*\.(?P<ext>xlsx|xls|xlsm)$",
    re.IGNORECASE,
)
COST_FILE_RE_ALT = re.compile(
    r"^(?P<year>\d{4})年(?P<month>\d{1,2})月_成本统计_(?P<project>.+?)\s*\.(?P<ext>xlsx|xls|xlsm)$",
    re.IGNORECASE,
)
MONTH_FOLDER_RE = re.compile(r"^(?:\d+(?:\.\d+)?[_-])*(\d{4}年\d{1,2}月)$")


def yyyymm_to_month_label(yyyymm: str) -> str:
    return f"{int(yyyymm[:4])}年{int(yyyymm[4:6]):02d}月"


def parse_cost_stat_filename(name: str):
    """从成本统计文件名解析 (月份标签, 项目名)。无法识别则返回 None。"""
    if name.startswith("~$") or name.startswith(".") or "汇总" in name:
        return None
    match = COST_FILE_RE.match(name)
    if match:
        return yyyymm_to_month_label(match.group("yyyymm")), match.group("project").strip()
    match = COST_FILE_RE_ALT.match(name)
    if match:
        return (
            f"{int(match.group('year'))}年{int(match.group('month')):02d}月",
            match.group("project").strip(),
        )
    return None


def collect_cost_stat_files(base_path: Path):
    """收集成本统计 Excel。

    两种布局都支持：
    - 按月子目录（A_成本统计_raw）
    - 扁平目录，月份写在文件名里（B_成本统计_std）
    """
    grouped = defaultdict(list)
    skipped = []
    month_folders = [
        item for item in base_path.iterdir()
        if item.is_dir() and MONTH_FOLDER_RE.match(item.name)
    ]
    search_roots = month_folders if month_folders else [base_path]
    use_folder_month = bool(month_folders)

    for root in search_roots:
        for file in sorted(root.glob("*.xls*"), key=lambda p: p.name):
            if not file.is_file():
                continue
            parsed = parse_cost_stat_filename(file.name)
            if not parsed:
                if not file.name.startswith("~$") and not file.name.startswith("."):
                    skipped.append(file)
                continue
            month_label, project_name = parsed
            if use_folder_month:
                m = MONTH_FOLDER_RE.match(root.name)
                month_label = m.group(1) if m else root.name
            grouped[month_label].append((file, project_name))
    return grouped, skipped

def header_col_index(headers, header_name):
    return headers.index(header_name) + 1 if header_name in headers else None

def find_finance_payment_col(headers):
    return next(
        (
            idx for idx, col_name in enumerate(headers, 1)
            if col_name and any(keyword in str(col_name) for keyword in FINANCE_COL_KEYWORDS)
        ),
        None,
    )

def write_ws_cell(ws, row, col, value=None, number_format=None):
    cell = cast(Cell, ws.cell(row=row, column=col))
    if value is not None:
        cell.value = value
    if number_format is not None:
        cell.number_format = number_format
    return cell


def write_dataframe_to_sheet(ws, df):
    if df is None or df.empty:
        ws.cell(row=1, column=1, value="提示")
        ws.cell(row=2, column=1, value="无可用数据")
        return

    for col_idx, column_name in enumerate(df.columns, start=1):
        ws.cell(row=1, column=col_idx, value=column_name)

    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for col_idx, column_name in enumerate(df.columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row[column_name])


def apply_sheet_default_style(ws, df):
    if df is None or df.empty:
        return

    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 50
    for cell in ws[1]:
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

    num_format = numbers.FORMAT_NUMBER_00
    pct_format = numbers.FORMAT_PERCENTAGE_00
    numeric_cols = set(CONFIG["items"].keys()) | {
        "财务分包已付款",
        "总包合同额",
        "业主签证",
        "分包合同额",
        "分包签证",
        "工程部记录分包已付款",
        "采购数据汇总",
        "商混站混凝土汇总",
        "机具费汇总",
        "人工费汇总",
        "分包结算比例",
        "毛利率",
        "实际已发生金额",
        "计算总成本",
        "部门记录汇总金额",
        "后续预估汇总",
    }
    percent_cols = {"毛利率", "分包结算比例"}
    monthly_excel_cols = set(CONFIG["items"].keys()) | {"计算总成本", "部门记录汇总金额", "后续预估汇总"}
    pink_summary_cols = {"计算总成本", "部门记录汇总金额", "后续预估汇总"}
    no_fill_cols = {"项目名称", "标准项目名称"}

    headers = [cell.value for cell in ws[1]]
    for col_idx, header in enumerate(headers, start=1):
        if not header:
            continue
        fmt = pct_format if header in percent_cols else (num_format if header in numeric_cols else None)
        if fmt is not None:
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=col_idx).number_format = fmt

    yellow_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    blue_fill = PatternFill(fill_type="solid", fgColor="DDEBF7")
    pink_fill = PatternFill(fill_type="solid", fgColor="F4CCCC")
    for col_idx, header in enumerate(headers, start=1):
        if not header or header in no_fill_cols:
            continue
        fill = pink_fill if header in pink_summary_cols else (yellow_fill if header in monthly_excel_cols else blue_fill)
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col_idx).fill = fill

    for col_idx in range(1, len(df.columns) + 1):
        column_letter = get_column_letter(col_idx)
        max_length = 0
        for cell in ws[column_letter]:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 35)


def generate_profit_rate_chart(summary_df, output_html):
    required_cols = ["月份", "项目名称", "毛利率"]
    missing_cols = [col for col in required_cols if col not in summary_df.columns]
    if missing_cols:
        print(f"⚠️ 无法生成毛利率图表，缺少列: {', '.join(missing_cols)}")
        return
    chart_df = summary_df[required_cols].copy()
    chart_df["毛利率"] = chart_df["毛利率"].apply(normalize_margin_value)
    chart_df["月份日期"] = chart_df["月份"].apply(parse_month_to_datetime)
    chart_df = chart_df.dropna(subset=["项目名称", "毛利率", "月份日期"])
    if chart_df.empty:
        print("⚠️ 无有效毛利率数据，未生成图表（可能所有毛利率为空或未提取到）")
        return
    chart_df = chart_df.sort_values(["项目名称", "月份日期"])
    chart_df["月份"] = chart_df["月份日期"].apply(lambda d: f"{d.year}年{d.month}月")
    month_order = [f"{d.year}年{d.month}月" for d in sorted(chart_df["月份日期"].unique())]
    try:
        import plotly.express as px
    except ImportError:
        print("⚠️ 未安装 plotly，无法生成图表。请先安装: pip install plotly")
        return
    fig = px.line(
        chart_df,
        x="月份",
        y="毛利率",
        color="项目名称",
        markers=True,
        category_orders={"月份": month_order},
        title="各项目毛利率变化图",
    )
    fig.update_layout(
        template="plotly_white",
        xaxis_title="月份",
        yaxis_title="毛利率",
        legend_title_text="项目名称",
        hovermode="closest",
    )
    fig.update_traces(
        hovertemplate="项目名称: %{fullData.name}<br>月份: %{x}<br>毛利率: %{y:.2%}<extra></extra>"
    )
    fig.update_yaxes(tickformat=".2%")
    output_path = Path(output_html)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.write_html(str(output_path), full_html=True, include_plotlyjs=True)
    print(f"✓ 毛利率变化图已生成: {output_path}")

def build_margin_change_wide_df(summary_df):
    required_cols = ["月份", "项目名称", "毛利率"]
    missing_cols = [col for col in required_cols if col not in summary_df.columns]
    if missing_cols:
        print(f"⚠️ 无法生成毛利率环比宽表，缺少列: {', '.join(missing_cols)}")
        return pd.DataFrame()

    margin_df = summary_df[required_cols].copy()
    margin_df["毛利率"] = margin_df["毛利率"].apply(normalize_margin_value)
    margin_df["月份日期"] = margin_df["月份"].apply(parse_month_to_datetime)
    margin_df = margin_df.dropna(subset=["项目名称", "月份日期", "毛利率"])
    if margin_df.empty:
        return pd.DataFrame()

    margin_df = margin_df.sort_values(["项目名称", "月份日期"])
    margin_df["毛利率环比变化"] = margin_df.groupby("项目名称")["毛利率"].diff()

    ordered_month_dates = sorted(margin_df["月份日期"].unique())
    ordered_month_labels = [f"{d.year}年{d.month}月" for d in ordered_month_dates]
    margin_df["月份"] = margin_df["月份日期"].apply(lambda d: f"{d.year}年{d.month}月")

    wide_df = margin_df.pivot_table(
        index="项目名称",
        columns="月份",
        values="毛利率环比变化",
        aggfunc="last",
    )
    wide_df = wide_df.reindex(columns=ordered_month_labels)
    wide_df.index.name = "项目名称"
    return wide_df

def write_margin_change_sheet(wb, wide_df, sheet_name="毛利率环比变化宽表"):
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(title=sheet_name)

    if wide_df is None or wide_df.empty:
        ws.cell(row=1, column=1, value="项目名称")
        ws.cell(row=1, column=2, value="提示")
        ws.cell(row=2, column=2, value="暂无可用毛利率环比变化数据")
        return

    month_cols = list(wide_df.columns)
    ws.cell(row=1, column=1, value="项目名称")
    for col_idx, month_label in enumerate(month_cols, start=2):
        ws.cell(row=1, column=col_idx, value=month_label)

    for row_idx, project_name in enumerate(wide_df.index.tolist(), start=2):
        ws.cell(row=row_idx, column=1, value=project_name)
        for col_idx, month_label in enumerate(month_cols, start=2):
            value = wide_df.at[project_name, month_label]
            if pd.isna(value):
                ws.cell(row=row_idx, column=col_idx, value=None)
            else:
                ws.cell(row=row_idx, column=col_idx, value=float(value))
                ws.cell(row=row_idx, column=col_idx).number_format = numbers.FORMAT_PERCENTAGE_00

    ws.row_dimensions[1].height = 38
    for cell in ws[1]:
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

    for col in ws.columns:
        max_length = 0
        column_idx = col[0].column
        if column_idx is None:
            continue
        column = get_column_letter(column_idx)
        for cell in col:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column].width = min(max_length + 2, 35)

    ws.auto_filter.ref = ws.dimensions

def process_all_projects():
    base_path = Path(CONFIG["base_dir"])
    finance_col_name = "财务分包已付款"
    processor = ProjectProcessor(CONFIG)

    print("正在加载项目匹配库(JSON)...")

    matcher_dir, match_db_path, project_info_path = resolve_matcher_paths(CONFIG)
    print(f"✓ 项目匹配模块目录: {matcher_dir}")
    print(f"✓ 项目匹配库路径: {match_db_path}")
    if project_info_path:
        print(f"✓ 项目基本信息文件: {project_info_path}")

    matcher = init_project_matcher(
        project_code_file=str(project_info_path) if project_info_path else None,
        match_db_file=str(match_db_path),
        matcher_dir=str(matcher_dir),
    )
    match_db = matcher['db']
    match_project_fn = matcher['match_project']
    mark_unmatched_fn = matcher['mark_unmatched']
    save_db_fn = matcher['save_db']
    print(
        f"✓ 项目匹配库已加载：{matcher['db_path']}，标准项目 {len(match_db.get('projects', {}))} 个，"
        f"别名 {len(match_db.get('alias_to_standard', {}))} 个"
    )

    print("正在加载财务部分包付款数据...")
    finance_dict_by_code = {}
    finance_dict_by_name = {}
    try:
        finance_file = find_latest_finance_file(CONFIG["finance_payment_dir"])
        if not finance_file:
            raise FileNotFoundError("未找到财务数据文件")
        finance_df = pd.read_excel(finance_file, sheet_name='项目汇总（万泰铢）')
        required_finance_cols = {'项目编码', '付款汇总（万泰铢）'}
        missing_finance_cols = required_finance_cols - set(finance_df.columns)
        if missing_finance_cols:
            raise KeyError(f"财务sheet缺少必要列: {', '.join(missing_finance_cols)}")
        print(f"已加载 {len(finance_df)} 个项目的财务付款数据")
        for _, row in finance_df.iterrows():
            payment_amount = safe_to_float(row.get('付款汇总（万泰铢）', 0))
            normalized_code = normalize_project_code(row.get('项目编码'))
            if normalized_code:
                finance_dict_by_code[normalized_code] = payment_amount
            standard_name = str(row.get('标准项目名称', '')).strip()
            if standard_name:
                finance_dict_by_name[standard_name] = payment_amount
    except Exception as e:
        print(f"⚠️ 加载财务数据失败: {e}")
        finance_dict_by_code = {}
        finance_dict_by_name = {}

    print("正在加载合同管理总包合同额数据...")
    contract_dict_by_code = {}
    subcontract_contract_dict_by_code = {}
    try:
        contract_file = find_latest_contract_file(CONFIG["contract_management_dir"])
        if not contract_file:
            raise FileNotFoundError("未找到合同管理汇总文件")
        contract_df = pd.read_excel(contract_file, sheet_name=CONFIG["contract_management_sheet"])
        print(f"已加载 {len(contract_df)} 条合同管理记录")
        code_col_candidates = ["项目编码", "标准项目编码"]
        amount_col_candidates = [
            "总包合同额(泰铢)",
            "总包合同额（泰铢）",
            "总包合同额(万泰铢)",
            "总包合同额（万泰铢）",
            "总包合同额",
        ]
        code_col = next((c for c in code_col_candidates if c in contract_df.columns), None)
        amount_col = next((c for c in amount_col_candidates if c in contract_df.columns), None)
        if not code_col or not amount_col:
            raise KeyError(f"合同管理sheet缺少必要列，找到列: {list(contract_df.columns)}")
        # 成本汇总口径为万泰铢：列名含「泰铢」且不含「万」时 ÷10000
        amount_is_thb = ("泰铢" in str(amount_col)) and ("万" not in str(amount_col))
        for _, row in contract_df.iterrows():
            code = normalize_project_code(row.get(code_col))
            if not code:
                continue
            contract_amount = safe_to_float(row.get(amount_col, 0))
            if amount_is_thb:
                contract_amount = contract_amount / 10000.0
            contract_dict_by_code[code] = contract_dict_by_code.get(code, 0.0) + contract_amount
        print(f"已构建 {len(contract_dict_by_code)} 个项目编码的总包合同额索引\n")
    except Exception as e:
        print(f"⚠️ 加载合同管理数据失败: {e}")

    print("正在加载工程部分包合同台账数据（分包合同额 / 分包已付款）...")
    subcontract_paid_dict_by_code = {}
    try:
        sc_dir = CONFIG.get("subcontract_summary_dir", "")
        sc_file = find_latest_excel_file(sc_dir) if sc_dir else None
        if not sc_file:
            raise FileNotFoundError(f"未找到分包合同台账汇总文件: {sc_dir}")
        sc_df = pd.read_excel(sc_file)
        sc_code_col = next((c for c in ["项目编码", "标准项目编码"] if c in sc_df.columns), None)
        sc_amount_col = next((c for c in ["合同额（泰铢）", "合同额(泰铢）", "合同额(泰铢)"] if c in sc_df.columns), None)
        sc_paid_col = next((c for c in ["分包已付款", "公司实付含税"] if c in sc_df.columns), None)
        if not sc_code_col or not sc_amount_col:
            raise KeyError(f"分包台账缺少必要列，找到列: {list(sc_df.columns)}")
        for _, row in sc_df.iterrows():
            code = normalize_project_code(row.get(sc_code_col))
            if not code:
                continue
            amt = safe_to_float(row.get(sc_amount_col, 0))
            subcontract_contract_dict_by_code[code] = subcontract_contract_dict_by_code.get(code, 0.0) + (amt / 10000.0)
            if sc_paid_col:
                paid = safe_to_float(row.get(sc_paid_col, 0))
                subcontract_paid_dict_by_code[code] = subcontract_paid_dict_by_code.get(code, 0.0) + (paid / 10000.0)
        print(f"已构建 {len(subcontract_contract_dict_by_code)} 个项目编码的分包合同额索引\n")
    except Exception as e:
        print(f"⚠️ 加载工程部分包合同台账失败: {e}")
        subcontract_contract_dict_by_code = {}
        subcontract_paid_dict_by_code = {}

    print("正在加载采购数据汇总...")
    procurement_summary_by_code = load_project_amount_map(
        CONFIG["procurement_summary_dir"],
        ["项目汇总"],
        ["项目编码", "标准项目编码", "项目代码"],
        ["总采购金额"],
        divisor=10000.0,
    )

    print("正在加载商混站混凝土汇总...")
    concrete_summary_by_code = load_project_amount_map(
        CONFIG["concrete_summary_dir"],
        ["客户总额汇总"],
        ["项目编码", "标准项目编码", "项目代码"],
        ["金额（万泰铢）", "金额(万泰铢)", "金额"],
        divisor=1.0,
    )

    print("正在加载机具费汇总...")
    tool_summary_by_code = load_project_amount_map(
        CONFIG["tool_summary_dir"],
        ["项目汇总机具费"],
        ["项目编码", "标准项目编码", "项目代码"],
        ["费用合计(泰铢)", "费用合计（泰铢）", "费用合计"],
        divisor=10000.0,
    )

    print("正在加载人工费汇总...")
    labor_summary_by_code = load_project_amount_map(
        CONFIG["labor_summary_dir"],
        ["项目求和"],
        ["项目编码", "标准项目编码", "项目代码"],
        ["人工费汇总金额"],
        divisor=10000.0,
    )

    print(f"✓ 成本统计源目录: {base_path}")
    files_by_month, skipped_files = collect_cost_stat_files(base_path)
    month_labels = sorted(files_by_month.keys(), key=parse_month_sort_key, reverse=True)
    if not month_labels:
        print("⚠️ 未找到成本统计文件（需要「项目名_成本统计_YYYYMM.xlsx/.xlsm」）")
        return

    print(f"✓ 找到 {len(month_labels)} 个月份、{sum(len(v) for v in files_by_month.values())} 个项目文件:\n")
    for month_label in month_labels:
        print(f" - {month_label}  ({len(files_by_month[month_label])} 个)")
    for skipped in skipped_files:
        print(f" ✗ 跳过不符合格式的文件: {skipped.name}")
    print()

    summary_list = []

    for month_label in month_labels:
        print(f"\n{'='*60}")
        print(f"处理月份: {month_label}")
        print(f"{'='*60}")
        files = files_by_month[month_label]
        if not files:
            print(" ⚠️ 该月份未找到符合格式的文件\n")
            continue

        for file, project_name in files:
            print(f"\n  正在分析项目文件: {file.name}")
            row_data = {
"月份": month_label,
"项目名称": project_name,
"完整文件名": file.name
            }

            print(f" 🔍 正在用项目匹配库匹配项目名: {project_name}")
            hit = match_project_fn(project_name, match_db)

            if hit:
                project_code = normalize_project_code(hit.get('project_code'))
                standard_name = str(hit.get('project_name', '')).strip() or project_name
                contract_name = ''
                if project_code:
                    contract_name = matcher['contract_name_by_code'].get(project_code, '')
                if not contract_name:
                    contract_name = matcher['contract_name_by_standard_name'].get(standard_name, '')

                row_data["项目编码"] = project_code or "需手动核对"
                row_data["项目名称"] = standard_name
                row_data["匹配度"] = "100%"
                finance_payment = 0.0
                if project_code and project_code in finance_dict_by_code:
                    finance_payment = finance_dict_by_code[project_code]
                elif standard_name in finance_dict_by_name:
                    finance_payment = finance_dict_by_name[standard_name]
                contract_amount = contract_dict_by_code.get(project_code, 0.0) if project_code else 0.0
                subcontract_contract_amount = subcontract_contract_dict_by_code.get(project_code, 0.0) if project_code else 0.0

                row_data[finance_col_name] = finance_payment
                row_data[CONTRACT_COL_NAME] = contract_amount
                row_data[SUBCONTRACT_CONTRACT_COL_NAME] = subcontract_contract_amount
                row_data[SUBCONTRACT_PAID_COL_NAME] = subcontract_paid_dict_by_code.get(project_code, 0.0) if project_code else 0.0
                row_data[PROCUREMENT_SUMMARY_COL_NAME] = procurement_summary_by_code.get(project_code, 0.0) if project_code else 0.0
                row_data[CONCRETE_SUMMARY_COL_NAME] = concrete_summary_by_code.get(project_code, 0.0) if project_code else 0.0
                row_data[TOOL_SUMMARY_COL_NAME] = tool_summary_by_code.get(project_code, 0.0) if project_code else 0.0
                row_data[LABOR_SUMMARY_COL_NAME] = labor_summary_by_code.get(project_code, 0.0) if project_code else 0.0
            else:
                row_data["项目编码"] = "未匹配"
                row_data["项目名称"] = project_name
                row_data["匹配度"] = "0%"
                row_data[finance_col_name] = 0.0
                row_data[CONTRACT_COL_NAME] = 0.0
                row_data[SUBCONTRACT_CONTRACT_COL_NAME] = 0.0
                row_data[SUBCONTRACT_PAID_COL_NAME] = 0.0
                row_data[PROCUREMENT_SUMMARY_COL_NAME] = 0.0
                row_data[CONCRETE_SUMMARY_COL_NAME] = 0.0
                row_data[TOOL_SUMMARY_COL_NAME] = 0.0
                row_data[LABOR_SUMMARY_COL_NAME] = 0.0
                mark_unmatched_fn(match_db, project_name)

            try:
                for col, (s_key, expr, fmt) in CONFIG["items"].items():
                    value = processor.evaluate_expression(file, s_key, expr)
                    row_data[col] = value
            except Exception as e:
                print(f" 警告: 处理文件 {file.name} 时出错: {e}")

            summary_list.append(row_data)

    if summary_list:
        df = pd.DataFrame(summary_list)
        df["_month_sort"] = df["月份"].apply(parse_month_sort_key)
        df = df.sort_values(by=["项目名称", "_month_sort"], ascending=[True, False])

        latest_month_df = df.drop_duplicates(subset=["项目名称"], keep="first").sort_values(by="_month_sort", ascending=False)
        latest_month_df = latest_month_df.drop(columns=["_month_sort"])

        df = df.sort_values(by="_month_sort", ascending=False).drop(columns=["_month_sort"])
        chart_source_df = df.copy()

        contract_col_name = CONTRACT_COL_NAME
        visa_col_name = OWNER_VISA_COL_NAME
        subcontract_contract_col_name = SUBCONTRACT_CONTRACT_COL_NAME
        subcontract_visa_col_name = SUBCONTRACT_VISA_COL_NAME

        if "完整文件名" in df.columns:
            df = df.drop(columns=["完整文件名"])
        if "匹配度" in df.columns:
            df = df.drop(columns=["匹配度"])
        for col_to_remove in ["标准项目编码", "标准项目名称", "项目名称（合同管理）"]:
            if col_to_remove in df.columns:
                df = df.drop(columns=[col_to_remove])

        if contract_col_name not in df.columns:
            df[contract_col_name] = 0.0
        if subcontract_contract_col_name not in df.columns:
            df[subcontract_contract_col_name] = 0.0

        if "业主预计结算额" in df.columns:
            df[visa_col_name] = (
                df["业主预计结算额"].apply(safe_to_float)
                - df[contract_col_name].apply(safe_to_float)
            )
        else:
            df[visa_col_name] = 0.0

        if "分包结算额" in df.columns:
            df[subcontract_visa_col_name] = (
                df["分包结算额"].apply(safe_to_float)
                - df[subcontract_contract_col_name].apply(safe_to_float)
            )
        else:
            df[subcontract_visa_col_name] = 0.0

        for calc_col in ["计算总成本", "部门记录汇总金额", "后续预估汇总"]:
            if calc_col not in df.columns:
                df[calc_col] = 0.0

        ordered_columns = [
            "月份",
            "项目编码",
            "项目名称",
            "预计总成本",
            "计算总成本",
            "毛利",
            "毛利率",
            "业主预计结算额",
            contract_col_name,
            visa_col_name,
            "分包结算额",
            subcontract_contract_col_name,
            subcontract_visa_col_name,
            finance_col_name,
            SUBCONTRACT_PAID_COL_NAME,
            "材料采购实际已发生",
            "采购部采购金额",
            PROCUREMENT_SUMMARY_COL_NAME,
            "项目部采购金额",
            CONCRETE_SUMMARY_COL_NAME,
            "材料租赁费用",
            "工资支出",
            "人员调遣支出",
            LABOR_SUMMARY_COL_NAME,
            "机具费支出",
            TOOL_SUMMARY_COL_NAME,
            "部门记录汇总金额",
            "其他费用支出",
            "临建费用支出",
            "后续预估汇总",
            "材料采购后续预估发生",
            "材料租赁后续预估",
            "机具费后续预估",
            "人员调遣后续预估",
            "其他费用后续预估",
            "临建费用后续预估",
        ]
        existing_ordered_cols = [c for c in ordered_columns if c in df.columns]
        remaining_cols = [c for c in df.columns if c not in existing_ordered_cols]
        df = df[existing_ordered_cols + remaining_cols]

        df.to_excel(CONFIG["output_file"], index=False, engine='openpyxl')

        wb = openpyxl.load_workbook(CONFIG["output_file"])
        ws = wb.active
        if ws is None:
            raise ValueError("输出工作簿无可用活动工作表")

        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 50
        for cell in ws[1]:
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

        for col in ws.columns:
            max_length = 0
            column_idx = col[0].column
            if column_idx is None:
                continue
            column = get_column_letter(column_idx)
            for cell in col:
                try:
                    if cell.value:
                        if cell.row == 1:
                            max_length = min(len(str(cell.value)), 30)
                        else:
                            max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 35)
            ws.column_dimensions[column].width = adjusted_width

        headers = [cell.value for cell in ws[1]]
        col_total_cost = header_col_index(headers, "预计总成本")
        col_income = header_col_index(headers, "业主预计结算额")
        col_gross_profit = header_col_index(headers, "毛利")
        col_gross_margin = header_col_index(headers, "毛利率")
        col_subcontract_settlement = header_col_index(headers, "分包结算额")
        col_finance_payment = find_finance_payment_col(headers)

        if col_gross_margin is None and col_income and col_gross_profit:
            ws.cell(row=1, column=ws.max_column + 1, value="毛利率")
            col_gross_margin = ws.max_column

        col_settlement_ratio = None
        if col_finance_payment:
            ws.insert_cols(col_finance_payment + 1)
            col_settlement_ratio = col_finance_payment + 1
            ws.cell(row=1, column=col_settlement_ratio, value="分包结算比例")
            ws.cell(row=1, column=col_settlement_ratio).alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            if col_gross_margin and col_gross_margin > col_finance_payment:
                col_gross_margin += 1

        col_actual_spending = None
        if col_gross_margin:
            ws.insert_cols(col_gross_margin + 1)
            col_actual_spending = col_gross_margin + 1
            ws.cell(row=1, column=col_actual_spending, value="实际已发生金额")
            ws.cell(row=1, column=col_actual_spending).alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')

        headers_after_insert = [cell.value for cell in ws[1]]
        col_total_cost = header_col_index(headers_after_insert, "预计总成本")
        col_income = header_col_index(headers_after_insert, "业主预计结算额")
        col_gross_profit = header_col_index(headers_after_insert, "毛利")
        col_gross_margin = header_col_index(headers_after_insert, "毛利率")
        col_subcontract_settlement = header_col_index(headers_after_insert, "分包结算额")
        col_subcontract_contract = header_col_index(headers_after_insert, SUBCONTRACT_CONTRACT_COL_NAME)
        col_subcontract_visa = header_col_index(headers_after_insert, SUBCONTRACT_VISA_COL_NAME)
        col_settlement_ratio = header_col_index(headers_after_insert, "分包结算比例")
        col_actual_spending = header_col_index(headers_after_insert, "实际已发生金额")
        col_finance_payment = find_finance_payment_col(headers_after_insert)
        col_calc_total_cost = header_col_index(headers_after_insert, "计算总成本")
        col_dept_summary_total = header_col_index(headers_after_insert, "部门记录汇总金额")
        col_future_estimate_total = header_col_index(headers_after_insert, "后续预估汇总")

        calc_total_cost_terms = [
            "分包结算额",
            "材料采购实际已发生",
            "材料租赁费用",
            "工资支出",
            "人员调遣支出",
            "机具费支出",
            "其他费用支出",
            "临建费用支出",
            "材料采购后续预估发生",
            "材料租赁后续预估",
            "机具费后续预估",
            "人员调遣后续预估",
            "其他费用后续预估",
            "临建费用后续预估",
        ]
        dept_summary_terms = [
            PROCUREMENT_SUMMARY_COL_NAME,
            CONCRETE_SUMMARY_COL_NAME,
            LABOR_SUMMARY_COL_NAME,
            TOOL_SUMMARY_COL_NAME,
        ]
        future_estimate_terms = [
            "材料采购后续预估发生",
            "材料租赁后续预估",
            "机具费后续预估",
            "人员调遣后续预估",
            "其他费用后续预估",
            "临建费用后续预估",
        ]

        calc_total_cost_term_cols = [header_col_index(headers_after_insert, name) for name in calc_total_cost_terms]
        calc_total_cost_term_cols = [idx for idx in calc_total_cost_term_cols if idx]
        dept_summary_term_cols = [header_col_index(headers_after_insert, name) for name in dept_summary_terms]
        dept_summary_term_cols = [idx for idx in dept_summary_term_cols if idx]
        future_estimate_term_cols = [header_col_index(headers_after_insert, name) for name in future_estimate_terms]
        future_estimate_term_cols = [idx for idx in future_estimate_term_cols if idx]

        if col_calc_total_cost and calc_total_cost_term_cols:
            for row in range(2, ws.max_row + 1):
                terms_expr = "+".join(f"{get_column_letter(col)}{row}" for col in calc_total_cost_term_cols)
                write_ws_cell(ws, row, col_calc_total_cost, value=f"={terms_expr}")

        if col_dept_summary_total and dept_summary_term_cols:
            for row in range(2, ws.max_row + 1):
                terms_expr = "+".join(f"{get_column_letter(col)}{row}" for col in dept_summary_term_cols)
                write_ws_cell(ws, row, col_dept_summary_total, value=f"={terms_expr}")

        if col_future_estimate_total and future_estimate_term_cols:
            for row in range(2, ws.max_row + 1):
                terms_expr = "+".join(f"{get_column_letter(col)}{row}" for col in future_estimate_term_cols)
                write_ws_cell(ws, row, col_future_estimate_total, value=f"={terms_expr}")

        if col_subcontract_contract and col_subcontract_settlement:
            for row in range(2, ws.max_row + 1):
                raw_value = safe_to_float(ws.cell(row=row, column=col_subcontract_contract).value)
                write_ws_cell(ws, row, col_subcontract_contract, value=f"={raw_value}")
        if col_subcontract_visa and col_subcontract_settlement and col_subcontract_contract:
            subcontract_settlement_col_letter = get_column_letter(col_subcontract_settlement)
            subcontract_contract_col_letter = get_column_letter(col_subcontract_contract)
            for row in range(2, ws.max_row + 1):
                write_ws_cell(
                    ws,
                    row,
                    col_subcontract_visa,
                    value=f"={subcontract_settlement_col_letter}{row}-{subcontract_contract_col_letter}{row}",
                )

        if col_settlement_ratio and col_subcontract_settlement and col_finance_payment:
            for row in range(2, ws.max_row + 1):
                settlement_col_letter = get_column_letter(col_subcontract_settlement)
                payment_col_letter = get_column_letter(col_finance_payment)
                write_ws_cell(
                    ws,
                    row,
                    col_settlement_ratio,
                    value=f'=IF({settlement_col_letter}{row}=0, 0, {payment_col_letter}{row}/{settlement_col_letter}{row})',
                    number_format=numbers.FORMAT_PERCENTAGE_00,
                )

        if col_gross_margin and col_income and col_gross_profit:
            for row in range(2, ws.max_row + 1):
                income_letter = get_column_letter(col_income)
                profit_letter = get_column_letter(col_gross_profit)
                write_ws_cell(
                    ws,
                    row,
                    col_gross_margin,
                    value=f'=IF({income_letter}{row}=0, 0, {profit_letter}{row}/{income_letter}{row})',
                )

        estimated_cols = []
        for idx, col_name in enumerate(headers_after_insert, 1):
            if col_name and "后续预估" in str(col_name):
                estimated_cols.append(idx)

        if col_actual_spending and col_total_cost and col_subcontract_settlement and len(estimated_cols) > 0 and col_finance_payment:
            for row in range(2, ws.max_row + 1):
                total_cost_col_letter = get_column_letter(col_total_cost)
                settlement_col_letter = get_column_letter(col_subcontract_settlement)
                finance_col_letter = get_column_letter(col_finance_payment)
                estimated_sum = " + ".join([get_column_letter(col) + str(row) for col in estimated_cols])
                formula = f'={total_cost_col_letter}{row} - {settlement_col_letter}{row} - ({estimated_sum}) + {finance_col_letter}{row}'
                write_ws_cell(ws, row, col_actual_spending, value=formula)

        num_format = numbers.FORMAT_NUMBER_00
        pct_format = numbers.FORMAT_PERCENTAGE_00
        no_stat_cols = {"项目名称", "标准项目名称"}
        numeric_cols = set(CONFIG["items"].keys()) | {
            finance_col_name,
            contract_col_name,
            visa_col_name,
            subcontract_contract_col_name,
            subcontract_visa_col_name,
            SUBCONTRACT_PAID_COL_NAME,
            PROCUREMENT_SUMMARY_COL_NAME,
            CONCRETE_SUMMARY_COL_NAME,
            TOOL_SUMMARY_COL_NAME,
            LABOR_SUMMARY_COL_NAME,
            "分包结算比例",
            "毛利率",
            "实际已发生金额",
            "计算总成本",
            "部门记录汇总金额",
            "后续预估汇总",
        }
        for col in range(1, ws.max_column + 1):
            col_name = ws.cell(1, col).value
            if not col_name or col_name in no_stat_cols or col_name not in numeric_cols:
                continue
            if col_name in CONFIG["items"]:
                fmt = CONFIG["items"][col_name][2]
                cell_fmt = pct_format if fmt == '%' else num_format
            elif col_name in ("分包结算比例", "毛利率"):
                cell_fmt = pct_format
            else:
                cell_fmt = num_format
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col)
                cell.number_format = cell_fmt

        yellow_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        blue_fill = PatternFill(fill_type="solid", fgColor="DDEBF7")
        pink_fill = PatternFill(fill_type="solid", fgColor="F4CCCC")
        monthly_excel_cols = set(CONFIG["items"].keys())
        monthly_excel_cols.update({"计算总成本", "部门记录汇总金额", "后续预估汇总"})
        pink_summary_cols = {"计算总成本", "部门记录汇总金额", "后续预估汇总"}
        no_fill_cols = {"项目名称", "标准项目名称"}
        for col in range(1, ws.max_column + 1):
            col_name = ws.cell(1, col).value
            if not col_name or col_name in no_fill_cols:
                continue
            if col_name in pink_summary_cols:
                fill = pink_fill
            else:
                fill = yellow_fill if col_name in monthly_excel_cols else blue_fill
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=col).fill = fill

        margin_change_wide_df = build_margin_change_wide_df(chart_source_df)
        write_margin_change_sheet(wb, margin_change_wide_df)

        latest_month_sheet_name = "最新月份项目"
        if latest_month_sheet_name in wb.sheetnames:
            del wb[latest_month_sheet_name]
        latest_ws = wb.create_sheet(title=latest_month_sheet_name)
        write_dataframe_to_sheet(latest_ws, latest_month_df)
        apply_sheet_default_style(latest_ws, latest_month_df)

        wb.save(CONFIG["output_file"])
        print(f"\n✓ 成功！所有项目数据已汇总至: {CONFIG['output_file']}")

        chart_output = CONFIG.get("profit_rate_chart_file")
        if not chart_output:
            chart_output = str(Path("/Users/wyattsun/Desktop/邯郸中材/01成本统计/0_html") / "项目毛利率变化图.html")
        generate_profit_rate_chart(chart_source_df, chart_output)
        save_db_fn(matcher['db_path'], match_db)
        print(f"✓ 匹配库已保存: {matcher['db_path']}")
    else:
        print("\n未找到符合格式的文件，未生成汇总文件。")

if __name__ == "__main__":
    process_all_projects()