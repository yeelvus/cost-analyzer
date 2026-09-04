#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""读取各项目文件夹中的 项目编码.json，汇总成 Excel。

默认输出到业主结算额根目录：业主结算汇总.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from config import DEFAULT_SETTLEMENT_ROOT, EXCEL_NAME, JSON_NAME_TMPL, PROJECT_CODE_RE
from schema import JSON_FIELDS, MONEY_FIELDS, effective_rate, to_thb

CODE_RE = re.compile(PROJECT_CODE_RE)

HEADER_FILL = "1B365D"
HEADER_FONT = "FFFFFF"
ALT_FILL = "F7F9FC"


def parse_folder_name(name: str) -> tuple[str, str]:
    if "_" in name:
        code, title = name.split("_", 1)
        return code.strip(), title.strip()
    return name.strip(), ""


def load_json(path: Path) -> dict:
    raw = path.read_bytes()
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
            continue

        repaired_lines = []
        for line in text.splitlines():
            stripped = line.strip()
            if stripped.startswith('"  ""'):
                if stripped.endswith('",'):
                    stripped = stripped[1:-2] + ","
                elif stripped.endswith('"'):
                    stripped = stripped[1:-1]
                line = stripped.replace('""', '"')
            repaired_lines.append(line)
        repaired = "\n".join(repaired_lines)

        for candidate in (text, repaired):
            try:
                data = json.loads(candidate)
            except json.JSONDecodeError as exc:
                last_error = exc
                continue
            return data if isinstance(data, dict) else {}

    raise ValueError(f"无法读取 JSON 文件: {path}") from last_error


def collect_rows(root: Path) -> list[dict]:
    rows: list[dict] = []
    for folder in sorted(p for p in root.iterdir() if p.is_dir() and not p.name.startswith(".")):
        code, name = parse_folder_name(folder.name)
        if not CODE_RE.match(code):
            continue
        json_path = folder / JSON_NAME_TMPL.format(code=code)
        if not json_path.exists():
            rows.append(
                {
                    "项目编码": code,
                    "项目名称": name,
                    "备注": "尚未创建 JSON",
                    "_json路径": "",
                }
            )
            continue
        data = load_json(json_path)
        row = {k: data.get(k, "") for k in JSON_FIELDS}
        if not row.get("项目编码"):
            row["项目编码"] = code
        if not row.get("项目名称"):
            row["项目名称"] = name
        currency = row.get("币种") or row.get("结算金额单位") or "THB"
        rate = effective_rate(currency, row.get("汇率"))
        row["有效汇率"] = rate
        for field in MONEY_FIELDS:
            row[f"{field}_泰铢"] = to_thb(row.get(field), currency, row.get("汇率"))
        row["_json路径"] = str(json_path)
        rows.append(row)
    return rows


def write_excel(rows: list[dict], out_path: Path) -> None:
    extra = ["有效汇率"] + [f"{f}_泰铢" for f in MONEY_FIELDS]
    headers = JSON_FIELDS + extra + ["JSON路径"]
    wb = Workbook()
    ws = wb.active
    ws.title = "业主结算"

    thin = Border(
        left=Side(style="thin", color="D5D8DC"),
        right=Side(style="thin", color="D5D8DC"),
        top=Side(style="thin", color="D5D8DC"),
        bottom=Side(style="thin", color="D5D8DC"),
    )
    head_font = Font(name="PingFang SC", bold=True, color=HEADER_FONT, size=11)
    cell_font = Font(name="PingFang SC", size=10)
    fill_head = PatternFill("solid", fgColor=HEADER_FILL)
    fill_alt = PatternFill("solid", fgColor=ALT_FILL)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.font = head_font
        cell.fill = fill_head
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 2)}"

    for i, row in enumerate(rows, 1):
        values = [row.get(k, "") for k in JSON_FIELDS]
        values.append(row.get("有效汇率", ""))
        values.extend(row.get(f"{f}_泰铢") for f in MONEY_FIELDS)
        values.append(row.get("_json路径", ""))
        for col, val in enumerate(values, 1):
            cell = ws.cell(i + 1, col, val if val is not None else "")
            cell.font = cell_font
            cell.border = thin
            cell.alignment = center if col <= 2 else left
            header = headers[col - 1]
            if header in MONEY_FIELDS or header.endswith("_泰铢") or header in ("汇率", "有效汇率"):
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            if i % 2 == 0:
                cell.fill = fill_alt
        ws.row_dimensions[i + 1].height = 20

    widths = {
        "项目编码": 16,
        "项目名称": 36,
        "结算金额": 16,
        "结算金额单位": 12,
        "结算时间": 14,
        "合同金额": 16,
        "申报金额": 16,
        "已付款金额": 16,
        "余款": 14,
        "质保金": 14,
        "币种": 10,
        "汇率": 10,
        "备注": 24,
        "有效汇率": 12,
        "结算金额_泰铢": 16,
        "合同金额_泰铢": 16,
        "申报金额_泰铢": 16,
        "已付款金额_泰铢": 16,
        "余款_泰铢": 14,
        "质保金_泰铢": 14,
        "JSON路径": 50,
    }
    for col, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(h, 14)

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="汇总业主结算 JSON 为 Excel")
    parser.add_argument("--root", default=str(DEFAULT_SETTLEMENT_ROOT))
    parser.add_argument(
        "--out",
        default="",
        help="输出 xlsx 路径，默认写在根目录下 业主结算汇总.xlsx",
    )
    args = parser.parse_args()

    root = Path(args.root)
    if not root.exists():
        print(f"目录不存在（网络盘未挂载？）: {root}")
        return 1

    rows = collect_rows(root)
    out = Path(args.out) if args.out else root / EXCEL_NAME
    write_excel(rows, out)
    print(f"项目 {len(rows)} 条 → {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
