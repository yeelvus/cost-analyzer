import os
import re
import sys
import json
import html
import importlib.util
from pathlib import Path
from dataclasses import dataclass
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
_CODE_ROOT = SCRIPT_DIR.parent
TEMPLATE_PATH = _CODE_ROOT / "templates" / "人工费汇总_template.html"

if str(_CODE_ROOT) not in sys.path:
    sys.path.insert(0, str(_CODE_ROOT))

from common.project import (
    build_standard_projects_from_dataframe,
    is_unmatched_alias,
    load_or_init_db,
    match_project,
    save_db,
)
from common.excel import apply_common_excel_layout
from settings import MATCH_DB_PATH as _SETTINGS_MATCH_DB_PATH
from common.paths import html_data_path, raw_data_path, std_data_path, sum_data_path

# ====================== 配置区 ======================
class Config:
    # 输入
    INPUT_SOURCE = str(
        raw_data_path(
            "人工费",
            "人工费台账",
            legacy=Path('/Users/wyattsun/Desktop/邯郸中材/01成本统计/11_人工费🙎🏻\u200d♂️/A_1_人工费台账/A_人工费台账_raw'),
        )
    )
    INPUT_GLOB = "泰国子公司项目人工费统计台账*.xlsx"
    
    # 项目匹配
    MAPPING_EXCEL_PATH = r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/0_项目基本信息情况/项目基本信息表.xlsx"
    MATCH_DB_PATH = str(_SETTINGS_MATCH_DB_PATH)
    
    # 输出
    OUTPUT_STD = str(
        std_data_path(
            "人工费",
            "B_人工费台账_std",
            legacy=Path('/Users/wyattsun/Desktop/邯郸中材/01成本统计/11_人工费🙎🏻\u200d♂️/A_1_人工费台账/B_人工费台账_std'),
        )
    )
    OUTPUT_Z_REPORT = r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/Z_report"
    
    MONEY_COLUMNS = ['中方正式工铢', '中方劳务工铢', '子公司属地工铢', '劳务公司铢', '签证费']



# ====================== 辅助函数 ======================
def _parse_date_key_from_name(file_name: str):
    stem = Path(file_name).stem
    m = re.search(r'(20\d{2})[.年\-/](\d{1,2})(?:[.月\-/](\d{1,2}))?', stem)
    if not m:
        return None
    y = int(m.group(1))
    mm = int(m.group(2))
    dd = int(m.group(3)) if m.group(3) else 1
    if mm < 1 or mm > 12 or dd < 1 or dd > 31:
        return None
    return y, mm, dd


def resolve_input_file() -> Path:
    p = Path(Config.INPUT_SOURCE).expanduser()
    if p.is_file():
        return p.resolve()
    if p.is_dir():
        candidates = [x for x in p.glob(Config.INPUT_GLOB) if x.is_file()]
        if not candidates:
            raise FileNotFoundError(f"未找到匹配文件：{Config.INPUT_GLOB}")

        def sort_key(fp: Path):
            parsed = _parse_date_key_from_name(fp.name)
            return (1, parsed or (0, 0, 0), fp.stat().st_mtime)

        return sorted(candidates, key=sort_key)[-1].resolve()
    raise FileNotFoundError(f"输入路径不存在：{p}")


def parse_money(series: pd.Series) -> pd.Series:
    cleaned = (
        series.astype(str)
        .str.replace('\u3000', '', regex=False)
        .str.replace(',', '', regex=False)
        .str.replace('，', '', regex=False)
        .str.replace('（', '(', regex=False)
        .str.replace('）', ')', regex=False)
        .str.replace('－', '-', regex=False)
        .str.replace('–', '-', regex=False)
        .str.replace('—', '-', regex=False)
        .str.replace('THB', '', regex=False)
        .str.replace('thb', '', regex=False)
    )

    def to_num(val):
        if pd.isna(val):
            return 0
        text = str(val).strip()
        if not text:
            return 0
        is_negative = text.startswith('(') and text.endswith(')')
        text = text.strip('()')
        text = re.sub(r'[^0-9.-]', '', text)
        if text.count('-') > 1:
            text = text.replace('-', '', text.count('-') - 1)
        try:
            num = float(text)
        except:
            num = 0
        return -num if is_negative else num

    return cleaned.apply(to_num)


def _sanitize_filename(name: str) -> str:
    """移除文件名中的非法字符。"""
    return re.sub(r'[\\/*?:"<>|\r\n\t]', '_', str(name)).strip()


@dataclass
class LaborMonthlyRecord:
        month: str
        zh_formal: Optional[float]
        zh_labor: Optional[float]
        local_labor: Optional[float]
        outsourcing_labor: Optional[float]
        visa_fee: Optional[float]
        total_labor: Optional[float]
        status: str


def _fmt_number(val: Optional[float], digits: int = 2) -> str:
        if val is None:
                return "-"
        return f"{val:,.{digits}f}"


def _fmt_text(val: str) -> str:
        return html.escape(val) if val else "-"


def make_labor_html(project_data: Dict[str, List[LaborMonthlyRecord]], months: List[str]) -> str:
        sections: List[str] = []
        toc_items: List[str] = []
        chart_data: Dict[str, Dict[str, List[Optional[float]]]] = {}

        def missing_count(project: str) -> int:
                records_by_month = {r.month: r for r in project_data[project]}
                miss = 0
                for m in months:
                        rec = records_by_month.get(m)
                        if rec is None:
                                miss += 6
                                continue
                        miss += int(rec.zh_formal is None)
                        miss += int(rec.zh_labor is None)
                        miss += int(rec.local_labor is None)
                        miss += int(rec.outsourcing_labor is None)
                        miss += int(rec.visa_fee is None)
                        miss += int(rec.total_labor is None)
                return miss

        ordered_projects = sorted(project_data.keys(), key=lambda p: (missing_count(p) > 0, missing_count(p), p))

        for idx, project in enumerate(ordered_projects, start=1):
                records_by_month = {r.month: r for r in project_data[project]}

                rows: List[str] = []
                zh_formal_series: List[Optional[float]] = []
                zh_labor_series: List[Optional[float]] = []
                local_labor_series: List[Optional[float]] = []
                outsourcing_labor_series: List[Optional[float]] = []
                visa_fee_series: List[Optional[float]] = []
                total_labor_series: List[Optional[float]] = []
                status_text = ""

                for month in months:
                        rec = records_by_month.get(month)
                        zh_formal_series.append(None if rec is None else rec.zh_formal)
                        zh_labor_series.append(None if rec is None else rec.zh_labor)
                        local_labor_series.append(None if rec is None else rec.local_labor)
                        outsourcing_labor_series.append(None if rec is None else rec.outsourcing_labor)
                        visa_fee_series.append(None if rec is None else rec.visa_fee)
                        total_labor_series.append(None if rec is None else rec.total_labor)

                        if rec and rec.status:
                                status_text = rec.status

                        rows.append(
                                "<tr>"
                                f"<td class='metric-cell'>{month}</td>"
                                f"<td class='num metric-cell'>{_fmt_number(None if rec is None else rec.zh_formal)}</td>"
                                f"<td class='num metric-cell'>{_fmt_number(None if rec is None else rec.zh_labor)}</td>"
                                f"<td class='num metric-cell'>{_fmt_number(None if rec is None else rec.local_labor)}</td>"
                                f"<td class='num metric-cell'>{_fmt_number(None if rec is None else rec.outsourcing_labor)}</td>"
                                f"<td class='num metric-cell'>{_fmt_number(None if rec is None else rec.visa_fee)}</td>"
                                f"<td class='num metric-cell strong-cell'>{_fmt_number(None if rec is None else rec.total_labor)}</td>"
                                f"<td class='remark-cell'>{_fmt_text('' if rec is None else rec.status)}</td>"
                                "</tr>"
                        )

                chart_key = f"project_{idx}"
                chart_data[chart_key] = {
                        "zh_formal": zh_formal_series,
                        "zh_labor": zh_labor_series,
                        "local_labor": local_labor_series,
                        "outsourcing_labor": outsourcing_labor_series,
                        "visa_fee": visa_fee_series,
                        "total_labor": total_labor_series,
                }

                chart_amount_id = f"chart_amount_{idx}"
                chart_detail_id = f"chart_detail_{idx}"
                card_id = f"card_{idx}"
                toc_items.append(f"<a class='toc-link' href='#{card_id}'>{html.escape(project)}</a>")

                section = (
                        f"<section class='project-card' id='{card_id}'>"
                        f"<div class='card-header'><h2>{html.escape(project)}</h2><span class='status-tag'>{_fmt_text(status_text)}</span>"
                        f"<button class='export-btn' onclick='exportCardPDF(\"{card_id}\", \"{html.escape(project)}\")'>导出PDF</button></div>"
                        "<div class='chart-grid'>"
                        "<div class='chart-wrap'><h3>累计趋势</h3><canvas id='" + chart_amount_id + "'></canvas></div>"
                        "<div class='chart-wrap'><h3>构成趋势</h3><canvas id='" + chart_detail_id + "'></canvas></div>"
                        "</div>"
                        "<table>"
                        "<thead><tr>"
                        "<th>月份</th><th>中方正式工(万泰铢)</th><th>中方劳务工(万泰铢)</th><th>子公司属地工(万泰铢)</th><th>劳务公司(万泰铢)</th><th>签证费(万泰铢)</th><th>总人工费(万泰铢)</th><th>项目状态</th>"
                        "</tr></thead>"
                        "<tbody>"
                        + "".join(rows)
                        + "</tbody></table></section>"
                )
                sections.append(section)

        export_all_btn = (
                "<div style='text-align: center; margin: 20px 0;'>"
                "<button class='export-all-btn' onclick='exportAllCardsPDF()'>导出所有项目</button>"
                "</div>"
        )
        all_sections = export_all_btn + ''.join(sections)
        toc_html = "".join(toc_items)

        months_json = json.dumps(months, ensure_ascii=False)
        chart_data_json = json.dumps(chart_data, ensure_ascii=False)

        template = TEMPLATE_PATH.read_text(encoding="utf-8")
        return (
            template
            .replace("__TOC_HTML__", toc_html)
            .replace("__ALL_SECTIONS__", all_sections)
            .replace("__MONTHS_JSON__", months_json)
            .replace("__CHART_DATA_JSON__", chart_data_json)
        )

# ====================== 按项目拆分导出 ======================
def export_per_project_excel(
    grouped: pd.DataFrame,
    df_raw: pd.DataFrame,
    money_columns: list,
    output_dir: Path,
) -> None:
    """将人工费数据按项目拆分，每个项目输出一个独立 Excel 文件。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    metrics = money_columns + ['总人工费']
    non_numeric_headers = {'项目名称', '地区', '年', '月', '年-月', '项目编码', '项目状态', '原始项目名称'}

    projects = grouped[['项目编码', '项目名称']].drop_duplicates()
    print(f"\n📂 开始按项目拆分输出（共 {len(projects)} 个项目）...")

    for _, proj_row in projects.iterrows():
        code = str(proj_row['项目编码']).strip()
        name = str(proj_row['项目名称']).strip()

        prefix = f"{code}_{name}" if code and code != 'nan' else name
        filename = _sanitize_filename(prefix) + "_人工费.xlsx"
        out_path = output_dir / filename

        mask_grouped = (grouped['项目名称'] == proj_row['项目名称']) & (grouped['项目编码'] == proj_row['项目编码'])
        mask_raw = (df_raw['项目名称'] == proj_row['项目名称']) & (df_raw['项目编码'] == proj_row['项目编码'])
        proj_grouped = grouped[mask_grouped].copy()
        proj_raw = df_raw[mask_raw].copy()

        if proj_grouped.empty:
            continue

        with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
            # Sheet1: 月度汇总
            monthly_cols = ['年-月'] + money_columns + ['总人工费']
            available = [c for c in monthly_cols if c in proj_grouped.columns]
            monthly_df = proj_grouped[available].sort_values('年-月').reset_index(drop=True)
            
            # 添加汇总行
            summary_row = pd.DataFrame([{col: None for col in monthly_df.columns}])
            summary_row['年-月'] = '合计'
            for col in money_columns + ['总人工费']:
                if col in monthly_df.columns:
                    summary_row[col] = monthly_df[col].sum()
            monthly_df = pd.concat([monthly_df, summary_row], ignore_index=True)
            
            monthly_df.to_excel(writer, sheet_name='月度汇总', index=False)

            # Sheet2: 原始数据
            drop_cols = [c for c in ['项目状态', '原始项目名称'] if c in proj_raw.columns]
            proj_raw.drop(columns=drop_cols).to_excel(writer, sheet_name='原始数据', index=False)

            # Sheet3+: 各金额指标按月透视
            for metric in metrics:
                if metric not in proj_grouped.columns:
                    continue
                pivot = proj_grouped.pivot_table(
                    index=['项目编码', '项目名称'],
                    columns='年-月',
                    values=metric,
                    aggfunc='sum',
                    fill_value=0,
                )
                cols_sorted = sorted(
                    pivot.columns,
                    key=lambda s: (int(str(s).split('-')[0]), int(str(s).split('-')[1])),
                )
                pivot = pivot.reindex(columns=cols_sorted)
                sheet_name = metric[:31]
                pivot.to_excel(writer, sheet_name=sheet_name)

        wb = load_workbook(out_path)
        for ws in wb.worksheets:
            headers = [cell.value for cell in ws[1]]
            numeric_start = next(
                (i + 1 for i, h in enumerate(headers) if h and str(h) not in non_numeric_headers),
                3,
            )
            apply_common_excel_layout(
                ws,
                numeric_start_col=numeric_start,
                freeze_panes='B2',
                thousand_number_format='#,##0.00',
            )
        wb.save(out_path)
        print(f"  ✅ {filename}")

    print(f"✅ 按项目拆分 Excel 已输出至：{output_dir}")


# ====================== 主程序 ======================
def main():
    print("🚀 开始处理人工费台账...")

    excel_file = resolve_input_file()
    base_name = excel_file.stem
    print(f"✅ 使用文件：{excel_file.name}")

    df = pd.read_excel(excel_file, sheet_name="基础数据")
    df.dropna(how='all', inplace=True)
    df.dropna(axis=1, how='all', inplace=True)
    df.columns = df.columns.str.strip()

    # 列名映射
    column_mapping = {
        '项目名称': ['项目名称'],
        '地区': ['地区'],
        '中方正式工铢': ['中方正式工铢'],
        '中方劳务工铢': ['中方劳务工铢', '中方劳务工株'],
        '子公司属地工铢': ['子公司属地工铢'],
        '劳务公司铢': ['劳务公司铢'],
        '签证费': ['签证费'],
        '年': ['年'],
        '月': ['月']
    }

    found_columns = {}
    for standard, variants in column_mapping.items():
        for col in df.columns:
            if any(v in col for v in variants):
                found_columns[standard] = col
                break

    required = ['项目名称', '中方正式工铢', '中方劳务工铢', '子公司属地工铢', '劳务公司铢', '年', '月']
    missing = [col for col in required if col not in found_columns]
    if missing:
        print(f"\n错误：以下必要列未找到：{missing}")
        sys.exit(1)

    df = df.rename(columns={v: k for k, v in found_columns.items()})

    # 金额转换
    money_columns = Config.MONEY_COLUMNS
    for col in money_columns:
        if col in df.columns:
            df[col] = parse_money(df[col])
        else:
            df[col] = 0.0

    # 数据清洗
    INVALID_PROJECT_NAMES = {"", "nan", "none", "null", "n/a", "na", "-", "--", "/"}
    df = df[df['项目名称'].apply(lambda x: str(x).strip().lower() not in INVALID_PROJECT_NAMES)]
    df = df[~df['项目名称'].astype(str).str.contains('合计|总计|小计|Grand', case=False, na=False)]
    df.reset_index(drop=True, inplace=True)

    # ========== 项目名称标准化 ==========
    print("\n🔄 开始项目名称标准化...")
    if os.path.exists(Config.MAPPING_EXCEL_PATH):
        mapping_xls = pd.ExcelFile(Config.MAPPING_EXCEL_PATH)
        mapping_df = pd.read_excel(mapping_xls, sheet_name=mapping_xls.sheet_names[0])
        mapping_df.columns = [str(c).strip() for c in mapping_df.columns]

        # 列选择逻辑（保留原始完整逻辑）
        def pick_first_existing_column(columns, candidates):
            for name in candidates:
                if name in columns:
                    return name
            return None

        code_col = pick_first_existing_column(mapping_df.columns, ['项目编码', '项目代码', '项目编号', '编码', 'project_code'])
        name_col = pick_first_existing_column(mapping_df.columns, [
            '项目名称（合同管理）', '项目名称(合同管理)', '合同管理项目名称',
            '项目名称（成本梳理）', '项目名称(成本梳理)', '成本梳理项目名称',
            '项目名称', '标准项目名称', 'project_name'
        ])
        status_col = pick_first_existing_column(mapping_df.columns, ['备注', '项目状态', '状态', 'target_status'])

        matcher_info_df = pd.DataFrame({
            '项目名称': mapping_df[name_col],
            '项目编码': mapping_df[code_col] if code_col else '',
            '备注': mapping_df[status_col] if status_col else '需手动核对',
        })

        standard_projects = build_standard_projects_from_dataframe(matcher_info_df)
        match_db = load_or_init_db(Config.MATCH_DB_PATH, standard_projects)

        def standardize_project_name(raw_name):
            if str(raw_name).strip().lower() in INVALID_PROJECT_NAMES:
                return ("", str(raw_name), "")
            raw_name = str(raw_name).strip()
            hit = match_project(raw_name, match_db) if match_db else None
            if hit:
                return (
                    hit.get('project_code', '') or '',
                    hit.get('project_name', raw_name) or raw_name,
                    hit.get('target_status', '') or '',
                )
            return ("", raw_name, "")

        standardized = df['项目名称'].apply(standardize_project_name)
        df['项目编码'] = standardized.apply(lambda x: x[0])
        df['项目名称'] = standardized.apply(lambda x: x[1])
        df['项目状态'] = standardized.apply(lambda x: x[2])

        save_db(Config.MATCH_DB_PATH, match_db)
        print("✅ 项目名称标准化完成（共享匹配库）")
    else:
        print("⚠️ 映射表不存在，使用原始项目名称")
        df['项目编码'] = ""
        df['项目状态'] = ""

    # ========== 聚合 ==========
    grouped = df.groupby(['年', '月', '项目名称', '项目编码'])[money_columns].sum().reset_index()
    grouped['年'] = grouped['年'].astype(int)
    grouped['月'] = grouped['月'].astype(int)
    grouped['年-月'] = grouped.apply(lambda r: f"{int(r['年'])}-{int(r['月']):02d}", axis=1)
    grouped['总人工费'] = grouped[money_columns].sum(axis=1)

    summary_df = df.groupby(['项目编码', '项目名称'], as_index=False)[money_columns].sum()
    summary_df['总人工费'] = summary_df[money_columns].sum(axis=1)
    summary_df = summary_df[['项目编码', '项目名称'] + money_columns + ['总人工费']]

    summary_wan_df = summary_df.copy()
    summary_wan_df[money_columns + ['总人工费']] /= 10000.0
    summary_wan_df = summary_wan_df.rename(columns={c: f"{c}(万泰铢)" for c in money_columns + ['总人工费']})

    # 项目求和
    sum_components = money_columns[:]
    proj_summary = df.groupby(['项目编码', '项目名称'], as_index=False)[sum_components].sum()
    proj_summary['人工费汇总金额'] = proj_summary[sum_components].sum(axis=1)
    proj_summary = proj_summary[['项目编码', '项目名称', '人工费汇总金额']]

    # ====================== 输出 ======================
    output_files = [
        sum_data_path("人工费") / "人工费_汇总.xlsx",
    ]
    metrics = money_columns + ['总人工费']

    for out_path in output_files:
        out_path.parent.mkdir(parents=True, exist_ok=True)

        with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
            proj_summary.to_excel(writer, sheet_name='项目求和', index=False)
            summary_df.to_excel(writer, sheet_name='按项目汇总', index=False)
            summary_wan_df.to_excel(writer, sheet_name='按项目汇总(万泰铢)', index=False)

            raw_output_df = df.drop(columns=['项目状态', '原始项目名称'], errors='ignore')
            raw_output_df.to_excel(writer, sheet_name='原始数据', index=False)

            # 各指标按月透视表
            for metric in metrics:
                pivot = grouped.pivot_table(
                    index=['项目编码', '项目名称'],
                    columns='年-月',
                    values=metric,
                    aggfunc='sum',
                    fill_value=0
                )
                cols_sorted = sorted(pivot.columns, key=lambda s: (int(s.split('-')[0]), int(s.split('-')[1])))
                pivot = pivot.reindex(columns=cols_sorted)
                sheet_name = metric if len(metric) <= 31 else metric[:31]
                pivot.to_excel(writer, sheet_name=sheet_name)

        # 应用样式
        wb = load_workbook(out_path)
        skip_headers = {'项目名称', '地区', '年', '月', '年-月', '项目编码'}
        for ws in wb.worksheets:
            headers = [cell.value for cell in ws[1]]
            numeric_start = next((i+1 for i, h in enumerate(headers) if h and str(h) not in skip_headers), 3)
            apply_common_excel_layout(
                ws,
                numeric_start_col=numeric_start,
                freeze_panes='B2',
                thousand_number_format='#,##0.00'
            )
        wb.save(out_path)
        print(f"✅ 已输出：{out_path}")

    # 按项目拆分导出，直接输出到标准目录，不再创建子目录
    export_per_project_excel(
        grouped=grouped,
        df_raw=df,
        money_columns=money_columns,
        output_dir=Path(Config.OUTPUT_STD),
    )

    # ====================== HTML 报告 ======================
    print("\n📊 生成人工费项目月度 HTML 报告...")
    months_sorted = sorted(grouped['年-月'].unique(), key=lambda s: (int(s.split('-')[0]), int(s.split('-')[1])))
    project_data: Dict[str, List[LaborMonthlyRecord]] = {}
    grouped_for_html = grouped.copy()
    grouped_for_html['项目状态'] = grouped_for_html['项目名称'].map(
        df.groupby('项目名称')['项目状态'].first().to_dict()
    ).fillna('未知')

    for _, row in grouped_for_html.iterrows():
        project = str(row.get('项目名称', '')).strip()
        if not project:
            continue
        project_data.setdefault(project, []).append(
            LaborMonthlyRecord(
                month=str(row['年-月']),
                zh_formal=(float(row['中方正式工铢']) / 10000.0) if pd.notna(row['中方正式工铢']) else None,
                zh_labor=(float(row['中方劳务工铢']) / 10000.0) if pd.notna(row['中方劳务工铢']) else None,
                local_labor=(float(row['子公司属地工铢']) / 10000.0) if pd.notna(row['子公司属地工铢']) else None,
                outsourcing_labor=(float(row['劳务公司铢']) / 10000.0) if pd.notna(row['劳务公司铢']) else None,
                visa_fee=(float(row['签证费']) / 10000.0) if pd.notna(row['签证费']) else None,
                total_labor=(float(row['总人工费']) / 10000.0) if pd.notna(row['总人工费']) else None,
                status=str(row.get('项目状态', '') or ''),
            )
        )

    for project in project_data:
        project_data[project].sort(key=lambda x: x.month)

    html_file = html_data_path() / f'人工费汇总分析_{base_name}.html'
    html_content = make_labor_html(project_data, months_sorted)
    html_file.write_text(html_content, encoding='utf-8')
    print(f"✅ HTML 图表生成完成：{html_file}")

    print("\n🎉 处理完成！")
    print(f"共记录项目：{len(summary_df)} 个")
    print(f"Excel 输出位置：按项目拆分 → {Config.OUTPUT_STD}；汇总 → 0_Sum/11_人工费")


if __name__ == "__main__":
    main()