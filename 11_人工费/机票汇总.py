import html
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import List, Optional

import pandas as pd
from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
_CODE_ROOT = SCRIPT_DIR.parent

if str(_CODE_ROOT) not in sys.path:
    sys.path.insert(0, str(_CODE_ROOT))

from common.project import (
    build_standard_projects_from_dataframe,
    load_or_init_db,
    match_project,
    prompt_unmatched_project_mappings,
    save_db,
)
from common.excel import apply_common_excel_layout
from settings import MATCH_DB_PATH as _SETTINGS_MATCH_DB_PATH, REPORT_ROOT
from common.paths import html_data_path, raw_data_path


# ====================== 配置区 ======================
class Config:
    INPUT_SOURCE = str(
        raw_data_path(
            "人工费",
            "机票",
            legacy=Path('/Users/wyattsun/Desktop/邯郸中材/01成本统计/11_人工费🙎🏻\u200d♂️/A_2_飞签机票/A_飞签台账_raw'),
        )
    )
    INPUT_GLOB = "*.xlsx"

    MAPPING_EXCEL_PATH = r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/0_项目基本信息情况/项目基本信息表.xlsx"
    MATCH_DB_PATH = str(_SETTINGS_MATCH_DB_PATH)
    OUTPUT_Z_REPORT = str(REPORT_ROOT)

    OUTPUT_STD = r"/Users/wyattsun/Desktop/邯郸中材/01成本统计/11_人工费🙎🏻‍♂️/A_2_飞签机票/B_飞签机票_std"

    MONEY_COL = "本次机票价格（泰铢）"

    # 目标保留列（顺序即输出顺序）
    TARGET_COLUMNS = [
        "所属单位",
        "姓名",
        "人次",
        "订票时间",
        "年-月",
        "去程机票日期",
        "回程机票日期",
        "出发地-到达地",
        "本次机票航班号",
        "舱位",
        "本次机票价格（泰铢）",
        "人均机票价格（泰铢）",
        "机票类型",
        "票代",
        "护照",
        "项目编码",
        "项目名称",
        "原始项目名称",
        "_来源文件",
    ]


# ====================== 辅助函数 ======================
def _sanitize_filename(name: str) -> str:
    return re.sub(r'[\\/*?:"<>|\r\n\t]', "_", str(name)).strip()


def _parse_money(series: pd.Series) -> pd.Series:
    def to_num(val):
        if pd.isna(val):
            return 0.0
        text = str(val).strip().replace(",", "").replace("，", "").replace("THB", "").replace("thb", "")
        is_neg = text.startswith("(") and text.endswith(")")
        text = text.strip("()")
        text = re.sub(r"[^0-9.\-]", "", text)
        try:
            num = float(text)
        except ValueError:
            num = 0.0
        return -num if is_neg else num

    return series.apply(to_num)


def _parse_person_count(series: pd.Series) -> pd.Series:
    """根据姓名列解析人次：按常见分隔符拆分后计数。"""
    _sep = re.compile(r"[，,、/\\\s\u3000]+")

    def count_names(val):
        if pd.isna(val):
            return 1
        parts = [p.strip() for p in _sep.split(str(val)) if p.strip()]
        return max(len(parts), 1)

    return series.apply(count_names)


def _parse_year_month(series: pd.Series) -> pd.Series:
    """从日期列解析出 'YYYY-MM' 字符串。"""
    converted = pd.to_datetime(series, errors="coerce", dayfirst=False)
    return converted.dt.strftime("%Y-%m").where(converted.notna(), other="")


def _format_date_only(series: pd.Series) -> pd.Series:
    """将日期序列统一为 YYYY/M/D，去掉时分秒。"""
    converted = pd.to_datetime(series, errors="coerce", dayfirst=False)

    def _to_ymd(v):
        if pd.isna(v):
            return ""
        return f"{v.year}/{v.month}/{v.day}"

    return converted.apply(_to_ymd)


def _pick_col(columns, candidates: List[str]) -> Optional[str]:
    """返回第一个在 columns 中能模糊匹配到的候选列名。"""
    for cand in candidates:
        for col in columns:
            if cand in str(col):
                return col
    return None


INVALID_NAMES = {"", "nan", "none", "null", "n/a", "na", "-", "--", "/"}


# ====================== 读取所有原始 Excel ======================
def load_all_raw(input_dir: Path) -> pd.DataFrame:
    files = sorted(input_dir.glob(Config.INPUT_GLOB))
    if not files:
        raise FileNotFoundError(f"未在 {input_dir} 找到任何 .xlsx 文件")

    frames: List[pd.DataFrame] = []
    for f in files:
        print(f"  📄 读取：{f.name}")
        try:
            df_s = pd.read_excel(f, sheet_name=0)
            df_s.dropna(how="all", inplace=True)
            df_s.dropna(axis=1, how="all", inplace=True)
            df_s.columns = [str(c).strip() for c in df_s.columns]
            df_s["_来源文件"] = f.name
            frames.append(df_s)
        except Exception as e:
            print(f"  ⚠️ 跳过（读取失败）：{f.name} — {e}")

    if not frames:
        raise ValueError("所有文件均无有效数据")

    combined = pd.concat(frames, ignore_index=True, sort=False)
    combined.columns = [str(c).strip() for c in combined.columns]
    return combined


# ====================== 列标准化 ======================
def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    col_map = {
        "所属单位":       ["所属单位", "单位", "部门"],
        "姓名":           ["姓名", "员工姓名", "人员"],
        "订票时间":       ["订票时间", "订票日期", "购票时间", "购票日期"],
        "去程机票日期":   ["去程机票日期", "去程日期", "出发日期"],
        "回程机票日期":   ["回程机票日期", "回程日期", "返回日期"],
        "出发地":         ["出发地", "起飞地", "始发地", "出发城市"],
        "到达地":         ["到达地", "目的地", "抵达地", "到达城市"],
        "本次机票航班号": ["本次机票航班号", "航班号", "航班"],
        "舱位":           ["舱位", "座位等级"],
        "本次机票价格（泰铢）": ["本次机票价格", "机票价格", "票价", "金额"],
        "机票类型":       ["机票类型", "类型"],
        "票代":           ["票代", "票务代理"],
        "护照":           ["护照", "护照号", "证件"],
    }

    rename_dict: dict = {}
    for standard, candidates in col_map.items():
        found = _pick_col(df.columns, candidates)
        if found and found not in rename_dict:
            rename_dict[found] = standard

    # "项目"列使用精确匹配，防止误命中"项目名称"等子串列
    if "项目" not in rename_dict.values():
        for exact_name in ["项目", "所属项目"]:
            if exact_name in df.columns and exact_name not in rename_dict:
                rename_dict[exact_name] = "项目"
                break

    df = df.rename(columns=rename_dict)

    # 补全缺失列
    for col in list(col_map.keys()):
        if col not in df.columns:
            df[col] = ""

    dep = df["出发地"].fillna("").astype(str).str.strip()
    arr = df["到达地"].fillna("").astype(str).str.strip()
    df["出发地-到达地"] = dep.where(dep != "", "") + "-" + arr.where(arr != "", "")
    df.loc[(dep == "") & (arr != ""), "出发地-到达地"] = arr[arr != ""]
    df.loc[(dep != "") & (arr == ""), "出发地-到达地"] = dep[dep != ""]
    df.loc[(dep == "") & (arr == ""), "出发地-到达地"] = ""

    return df


# ====================== 项目名称标准化 ======================
def standardize_projects(df: pd.DataFrame) -> pd.DataFrame:
    df["原始项目名称"] = df["项目"].astype(str).str.strip()

    if not os.path.exists(Config.MAPPING_EXCEL_PATH):
        print("⚠️ 映射表不存在，跳过项目标准化")
        df["项目编码"] = ""
        df["项目名称"] = df["原始项目名称"]
        return df

    mapping_df = pd.read_excel(Config.MAPPING_EXCEL_PATH)
    mapping_df.columns = [str(c).strip() for c in mapping_df.columns]

    def pick(cols, candidates):
        for c in candidates:
            if c in cols:
                return c
        return None

    code_col   = pick(mapping_df.columns, ["项目编码", "项目代码", "项目编号", "编码"])
    name_col   = pick(mapping_df.columns, [
        "项目名称（合同管理）", "项目名称(合同管理)", "合同管理项目名称",
        "项目名称（成本梳理）", "项目名称(成本梳理)", "成本梳理项目名称",
        "项目名称", "标准项目名称",
    ])
    status_col = pick(mapping_df.columns, ["备注", "项目状态", "状态"])

    matcher_df = pd.DataFrame({
        "项目名称": mapping_df[name_col],
        "项目编码": mapping_df[code_col] if code_col else "",
        "备注":     mapping_df[status_col] if status_col else "需手动核对",
    })

    standard_projects = build_standard_projects_from_dataframe(matcher_df)
    match_db = load_or_init_db(Config.MATCH_DB_PATH, standard_projects)

    def _match(raw_name: str):
        raw = str(raw_name).strip()
        if raw.lower() in INVALID_NAMES:
            return ("", raw)
        hit = match_project(raw, match_db) if match_db else None
        if hit:
            return (
                hit.get("project_code", "") or "",
                hit.get("project_name", raw) or raw,
            )
        return ("", raw)

    results = df["原始项目名称"].apply(_match)
    df["项目编码"] = results.apply(lambda x: x[0])
    df["项目名称"] = results.apply(lambda x: x[1])

    # 收集未命中项目（编码为空且名称等于原始名称）
    unmatched_mask = (df["项目编码"] == "") & (df["项目名称"] == df["原始项目名称"])
    unmatched_raw_names = df.loc[unmatched_mask, "原始项目名称"].unique().tolist()
    unmatched_raw_names = [n for n in unmatched_raw_names if n.lower() not in INVALID_NAMES]

    if unmatched_raw_names:
        print(f"⚠️ 有 {len(unmatched_raw_names)} 个项目未自动匹配，弹窗手动选择...")
        manual_map = prompt_unmatched_project_mappings(
            unmatched_raw_names, match_db, db_path=Config.MATCH_DB_PATH
        )
        # 将手动匹配结果回写到 df
        for raw_name, std_name in manual_map.items():
            if not std_name:
                continue
            row_mask = df["原始项目名称"] == raw_name
            # 从 db 重新查一次以获取编码
            hit = match_project(raw_name, match_db)
            code = hit.get("project_code", "") or "" if hit else ""
            name = hit.get("project_name", std_name) or std_name if hit else std_name
            df.loc[row_mask, "项目编码"] = code
            df.loc[row_mask, "项目名称"] = name

    save_db(Config.MATCH_DB_PATH, match_db)
    print("✅ 项目名称标准化完成")
    return df


# ====================== 按项目拆分导出 ======================
def export_per_project_excel(df: pd.DataFrame, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    non_numeric = {"所属单位", "姓名", "订票时间", "年-月", "去程机票日期", "回程机票日期",
                   "出发地-到达地", "本次机票航班号", "舱位", "机票类型", "票代", "护照", "项目编码", "项目名称",
                   "原始项目名称", "_来源文件"}

    projects = df[["项目编码", "项目名称"]].drop_duplicates()
    print(f"\n📂 按项目拆分输出（共 {len(projects)} 个项目）...")

    for _, row in projects.iterrows():
        code = str(row["项目编码"]).strip()
        name = str(row["项目名称"]).strip()

        prefix = f"{code}_{name}" if code and code != "nan" else name
        filename = _sanitize_filename(prefix) + "_机票.xlsx"
        out_path = output_dir / filename

        mask = (df["项目名称"] == row["项目名称"]) & (df["项目编码"] == row["项目编码"])
        proj_df = df[mask].copy()

        if proj_df.empty:
            continue

        display_cols = [c for c in Config.TARGET_COLUMNS if c in proj_df.columns]

        # 月度汇总
        monthly = (
            proj_df.groupby("年-月")[Config.MONEY_COL]
            .sum()
            .reset_index()
            .sort_values("年-月")
            .rename(columns={Config.MONEY_COL: "机票费用合计（泰铢）"})
        )

        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            proj_df[display_cols].to_excel(writer, sheet_name="明细数据", index=False)
            monthly.to_excel(writer, sheet_name="月度汇总", index=False)

        wb = load_workbook(out_path)
        for ws in wb.worksheets:
            headers = [cell.value for cell in ws[1]]
            numeric_start = next(
                (i + 1 for i, h in enumerate(headers) if h and str(h) not in non_numeric),
                3,
            )
            apply_common_excel_layout(
                ws,
                numeric_start_col=numeric_start,
                freeze_panes="B2",
                thousand_number_format="#,##0.00",
            )
        wb.save(out_path)
        print(f"  ✅ {filename}")

    print(f"✅ 按项目拆分输出完成：{output_dir}")


# ====================== HTML 报告生成 ======================
def build_html_report(df: pd.DataFrame, out_path: Path) -> None:
    """生成机票费用分析 HTML 报告（Chart.js 风格，对齐人工费汇总样式）。"""
    MONEY = Config.MONEY_COL
    AVG_COL = "人均机票价格（泰铢）"

    months_all = sorted(
        [m for m in df["年-月"].dropna().unique() if "-" in str(m)],
        key=lambda s: (int(s.split("-")[0]), int(s.split("-")[1])),
    )
    ticket_types = sorted(t for t in df["机票类型"].dropna().astype(str).unique() if t and t != "nan")

    def _month_series(sub: pd.DataFrame, col: str, agg: str) -> list:
        s = sub.groupby("年-月")[col].mean() if agg == "mean" else sub.groupby("年-月")[col].sum()
        return [round(float(s[m]), 2) if m in s.index else None for m in months_all]

    type_chart_data: dict = {}
    for t in ticket_types:
        sub = df[df["机票类型"].astype(str) == t]
        type_chart_data[t] = {
            "avg_price": _month_series(sub, AVG_COL, "mean"),
            "total":     _month_series(sub, MONEY, "sum"),
            "trips":     _month_series(sub, "人次", "sum"),
        }
    overview = {t: type_chart_data[t]["avg_price"] for t in ticket_types}

    # KPIs
    total_cost      = df[MONEY].sum()
    total_trips     = int(df["人次"].sum()) if "人次" in df.columns else len(df)
    avg_price_all   = df[AVG_COL].mean() if AVG_COL in df.columns and len(df) > 0 else 0
    project_count   = df["项目名称"].nunique() if "项目名称" in df.columns else 0
    date_range      = f"{months_all[0]} 至 {months_all[-1]}" if months_all else "-"

    COLORS = ["#2c7fb8", "#3fae72", "#d08b2f", "#8f5ccf", "#245f4b", "#e05c5c", "#e8a838", "#3ea3a3", "#c46a2f"]

    toc_items: list = ["<a class='toc-link' href='#card_overview'>📊 总览对比</a>"]
    sections:  list = []

    # ── 总览卡片 ──────────────────────────────────────────────
    overview_card = (
        "<section class='project-card' id='card_overview'>"
        "<div class='card-header'><h2>各机票类型人均价格对比</h2>"
        "<span class='status-tag'>全类型汇总</span></div>"
        "<div class='chart-wrap' style='height:360px'>"
        "<h3>各类型每月人均价格趋势（泰铢）</h3>"
        "<canvas id='chart_overview'></canvas></div>"
        "</section>"
    )
    sections.append(overview_card)

    # ── 各类型卡片 ────────────────────────────────────────────
    for idx, t in enumerate(ticket_types):
        card_id       = f"card_{idx}"
        chart_avg_id  = f"chart_avg_{idx}"
        chart_tot_id  = f"chart_tot_{idx}"
        sub           = df[df["机票类型"].astype(str) == t]
        type_total    = sub[MONEY].sum()
        type_trips    = int(sub["人次"].sum()) if "人次" in sub.columns else len(sub)
        type_avg      = sub[AVG_COL].mean() if AVG_COL in sub.columns and len(sub) > 0 else 0

        toc_items.append(f"<a class='toc-link' href='#{card_id}'>{html.escape(t)}</a>")

        rows = []
        for m in months_all:
            msub = sub[sub["年-月"] == m]
            if msub.empty:
                continue
            m_trips = int(msub["人次"].sum()) if "人次" in msub.columns else len(msub)
            m_avg   = msub[AVG_COL].mean() if AVG_COL in msub.columns else 0
            m_total = msub[MONEY].sum()
            rows.append(
                f"<tr>"
                f"<td class='metric-cell'>{m}</td>"
                f"<td class='num metric-cell'>{m_trips}</td>"
                f"<td class='num metric-cell'>{m_avg:,.2f}</td>"
                f"<td class='num metric-cell strong-cell'>{m_total:,.2f}</td>"
                f"</tr>"
            )

        sections.append(
            f"<section class='project-card' id='{card_id}'>"
            f"<div class='card-header'><h2>{html.escape(t)}</h2>"
            f"<span class='status-tag'>累计 {type_total:,.0f} 铢 · {type_trips} 人次 · 均价 {type_avg:,.0f}</span></div>"
            "<div class='chart-grid'>"
            f"<div class='chart-wrap'><h3>每月人均价格趋势（泰铢）</h3><canvas id='{chart_avg_id}'></canvas></div>"
            f"<div class='chart-wrap'><h3>每月总费用趋势（泰铢）</h3><canvas id='{chart_tot_id}'></canvas></div>"
            "</div>"
            "<table>"
            "<thead><tr><th>月份</th><th>人次</th><th>人均价格（泰铢）</th><th>总费用（泰铢）</th></tr></thead>"
            "<tbody>" + "".join(rows) + "</tbody></table>"
            "</section>"
        )

    toc_html    = "".join(toc_items)
    all_sections = "".join(sections)

    months_json    = json.dumps(months_all,    ensure_ascii=False)
    type_data_json = json.dumps(type_chart_data, ensure_ascii=False)
    overview_json  = json.dumps(overview,      ensure_ascii=False)
    colors_json    = json.dumps(COLORS)

    html_content = f"""<!doctype html>
<html lang="zh-CN">
<head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <title>机票费用分析报告</title>
    <script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.3/dist/chart.umd.min.js"></script>
    <style>
        :root {{
            --bg1: #f3f7f2;
            --bg2: #f9efe4;
            --card: rgba(255,255,255,0.9);
            --line: #dde6dc;
            --txt: #21322a;
            --muted: #5e7167;
            --accent: #2f7d61;
            --accent2: #c46a2f;
        }}
        * {{ box-sizing: border-box; }}
        body {{
            margin: 0;
            color: var(--txt);
            font-family: "PingFang SC", "Noto Sans SC", "Microsoft YaHei", sans-serif;
            background:
                radial-gradient(circle at 20% 0%, #e8f2e3 0%, transparent 40%),
                radial-gradient(circle at 80% 0%, #f8e6d6 0%, transparent 35%),
                linear-gradient(160deg, var(--bg1), var(--bg2));
            min-height: 100vh;
            scroll-behavior: smooth;
        }}
        .container {{ width: min(1400px, 97vw); margin: 28px auto 48px; }}
        .layout {{ display: grid; grid-template-columns: 260px minmax(0, 1fr); gap: 16px; align-items: start; }}
        .toc {{
            position: sticky; top: 14px;
            max-height: calc(100vh - 28px); overflow: auto;
            background: var(--card); border: 1px solid var(--line);
            border-radius: 14px; box-shadow: 0 10px 20px rgba(36,48,41,.06); padding: 12px;
        }}
        .toc-title {{ margin: 0 0 8px; font-size: 14px; color: #2f4338; font-weight: 700; }}
        .toc-links {{ display: grid; gap: 6px; }}
        .toc-link {{
            display: block; padding: 6px 8px; border-radius: 8px;
            color: #2a473c; text-decoration: none; font-size: 12px; line-height: 1.35;
            background: rgba(227,237,230,0.4);
        }}
        .toc-link:hover {{ background: rgba(196,106,47,0.18); }}
        .hero {{
            background: var(--card); border: 1px solid var(--line);
            border-radius: 18px; padding: 18px 22px;
            box-shadow: 0 12px 24px rgba(36,48,41,.08); margin-bottom: 16px;
        }}
        h1 {{ margin: 4px 0 8px; font-size: clamp(22px, 3vw, 32px); }}
        .hint {{ color: var(--muted); margin: 0; }}
        .kpi-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(140px, 1fr)); gap: 10px; margin-top: 14px; }}
        .kpi {{ background: rgba(227,237,230,0.5); border-radius: 10px; padding: 12px 14px; text-align: center; }}
        .kpi-val {{ font-size: 22px; font-weight: 700; color: var(--accent); }}
        .kpi-lbl {{ font-size: 11px; color: var(--muted); margin-top: 3px; }}
        .project-card {{
            background: var(--card); border: 1px solid var(--line);
            border-radius: 16px; padding: 18px; margin: 14px 0;
            box-shadow: 0 10px 20px rgba(36,48,41,.06);
        }}
        .card-header {{ display: flex; justify-content: space-between; align-items: center; gap: 10px; margin-bottom: 10px; }}
        .card-header h2 {{ margin: 4px 0; flex: 1; font-size: clamp(18px, 2.2vw, 24px); }}
        .status-tag {{
            padding: 4px 10px; font-size: 12px; border-radius: 999px;
            background: rgba(47,125,97,0.12); color: #285846; white-space: nowrap;
        }}
        h3 {{ margin: 0 0 8px; font-size: 14px; color: #2f4338; }}
        .chart-grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 12px; margin-bottom: 28px; }}
        .chart-wrap {{ width: 100%; height: 320px; margin-bottom: 0; }}
        table {{ width: 100%; border-collapse: collapse; font-size: 13px; margin-top: 12px; }}
        th, td {{ border-bottom: 1px solid var(--line); padding: 8px 6px; text-align: center; vertical-align: middle; }}
        th {{ color: #31453a; background: rgba(227,237,230,0.55); }}
        td.num {{ text-align: center; font-variant-numeric: tabular-nums; }}
        td.metric-cell {{ font-size: 12px; }}
        td.strong-cell {{ font-weight: 700; color: #244d3e; }}
        @media (max-width: 768px) {{
            .layout {{ grid-template-columns: 1fr; }}
            .toc {{ position: static; max-height: none; }}
            .chart-grid {{ grid-template-columns: 1fr; margin-bottom: 20px; }}
            .chart-wrap {{ height: 260px; }}
            table {{ min-width: 480px; }}
        }}
        @media print {{
            body {{ background: white; }}
            .hero, .project-card {{ box-shadow: none; border: 1px solid #ccc; }}
        }}
    </style>
</head>
<body>
    <div class="container">
        <div class="layout">
            <aside class="toc">
                <h3 class="toc-title">机票类型目录</h3>
                <div class="toc-links">{toc_html}</div>
            </aside>
            <main>
                <section class="hero">
                    <h1>✈️ 机票费用分析报告</h1>
                    <p class="hint">数据区间：{date_range}（单位：泰铢）</p>
                    <div class="kpi-grid">
                        <div class="kpi"><div class="kpi-val">{total_cost:,.0f}</div><div class="kpi-lbl">总机票费用（泰铢）</div></div>
                        <div class="kpi"><div class="kpi-val">{total_trips:,}</div><div class="kpi-lbl">总人次</div></div>
                        <div class="kpi"><div class="kpi-val">{avg_price_all:,.0f}</div><div class="kpi-lbl">平均人均票价（铢）</div></div>
                        <div class="kpi"><div class="kpi-val">{len(ticket_types)}</div><div class="kpi-lbl">机票类型数</div></div>
                        <div class="kpi"><div class="kpi-val">{project_count}</div><div class="kpi-lbl">涉及项目数</div></div>
                        <div class="kpi"><div class="kpi-val">{len(months_all)}</div><div class="kpi-lbl">统计月份数</div></div>
                    </div>
                </section>
                {all_sections}
            </main>
        </div>
    </div>

    <script>
        const months        = {months_json};
        const allTypeData   = {type_data_json};
        const overviewData  = {overview_json};
        const COLORS        = {colors_json};

        const amtFmt = (v) => v == null ? '-' : Number(v).toLocaleString('zh-CN', {{maximumFractionDigits: 0}});
        const lblFmt = (v) => v == null ? '' : Number(v).toLocaleString('zh-CN', {{maximumFractionDigits: 0}});

        const inlineLabelPlugin = {{
            id: 'inlineLabelPlugin',
            afterDatasetsDraw(chart) {{
                const ctx = chart.ctx;
                chart.data.datasets.forEach((ds, di) => {{
                    const meta = chart.getDatasetMeta(di);
                    if (meta.hidden) return;
                    meta.data.forEach((el, i) => {{
                        const raw = ds.data[i];
                        if (raw == null) return;
                        const txt = lblFmt(raw);
                        if (!txt) return;
                        const pos = el.tooltipPosition();
                        ctx.save();
                        ctx.fillStyle = '#24352d';
                        ctx.font = '600 10px "PingFang SC","Noto Sans SC","Microsoft YaHei",sans-serif';
                        ctx.textAlign = 'center';
                        ctx.textBaseline = 'bottom';
                        ctx.fillText(txt, pos.x, pos.y - 6);
                        ctx.restore();
                    }});
                }});
            }}
        }};
        Chart.register(inlineLabelPlugin);

        const bounds = (...lists) => {{
            const vals = lists.flat().filter(v => v != null && !isNaN(v));
            if (!vals.length) return {{ min: 0, max: 1 }};
            const lo = Math.min(...vals), hi = Math.max(...vals);
            return lo >= 0 ? {{ min: 0, max: hi * 1.2 + 1 }} : {{ min: lo - (hi-lo)*0.12 - 1, max: hi + (hi-lo)*0.12 + 1 }};
        }};

        const mkLine = (id, datasets, b) => {{
            const el = document.getElementById(id);
            if (!el) return;
            new Chart(el, {{
                type: 'line',
                data: {{ labels: months, datasets }},
                options: {{
                    responsive: true, maintainAspectRatio: false,
                    interaction: {{ mode: 'index', intersect: false }},
                    plugins: {{
                        legend: {{ position: 'top' }},
                        inlineLabelPlugin: {{}},
                        tooltip: {{ callbacks: {{ label: (c) => `${{c.dataset.label}}: ${{amtFmt(c.raw)}} 铢` }} }}
                    }},
                    scales: {{ y: {{ min: b.min, max: b.max, ticks: {{ callback: amtFmt }} }} }}
                }}
            }});
        }};

        // 总览：各类型人均价格对比（多线）
        (() => {{
            const types = Object.keys(overviewData);
            const datasets = types.map((t, i) => ({{
                label: t,
                data: overviewData[t],
                borderColor: COLORS[i % COLORS.length],
                backgroundColor: COLORS[i % COLORS.length] + '28',
                pointRadius: 4, borderWidth: 2, tension: 0.25, spanGaps: true,
            }}));
            mkLine('chart_overview', datasets, bounds(...types.map(t => overviewData[t])));
        }})();

        // 各类型单独卡片
        Object.entries(allTypeData).forEach(([t, s], idx) => {{
            const c = COLORS[idx % COLORS.length];
            const ds = (label, data) => ({{
                label, data,
                borderColor: c, backgroundColor: c + '28',
                pointRadius: 4, borderWidth: 2, tension: 0.25, spanGaps: true,
            }});
            mkLine(`chart_avg_${{idx}}`, [ds('人均价格（铢）', s.avg_price)], bounds(s.avg_price));
            mkLine(`chart_tot_${{idx}}`, [ds('总费用（铢）',   s.total)],     bounds(s.total));
        }});
    </script>
</body>
</html>"""

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html_content, encoding="utf-8")


# ====================== 主程序 ======================
def main():
    print("🚀 开始处理机票台账...")

    input_dir = Path(Config.INPUT_SOURCE)
    if not input_dir.exists():
        raise FileNotFoundError(f"输入目录不存在：{input_dir}")

    # 1. 读取所有原始文件
    df = load_all_raw(input_dir)
    print(f"✅ 原始数据共 {len(df)} 行")

    # 2. 列标准化
    df = normalize_columns(df)
    df = df.loc[:, ~df.columns.duplicated()]  # 去除重命名后可能产生的重复列

    # 3. 金额转换
    df[Config.MONEY_COL] = _parse_money(df[Config.MONEY_COL])

    # 3.1 人次 & 人均价格
    df["人次"] = _parse_person_count(df["姓名"])
    df["人均机票价格（泰铢）"] = (df[Config.MONEY_COL] / df["人次"]).round(2)

    # 3.2 日期列去掉时分秒，仅保留 YYYY/M/D
    for date_col in ["订票时间", "去程机票日期", "回程机票日期"]:
        if date_col in df.columns:
            df[date_col] = _format_date_only(df[date_col])

    # 4. 年-月列
    df["年-月"] = _parse_year_month(df["订票时间"])

    # 5. 清洗无效项目行
    df = df[df["项目"].apply(lambda x: str(x).strip().lower() not in INVALID_NAMES).values]
    df = df[~df["项目"].astype(str).str.contains("合计|总计|小计|Grand", case=False, na=False).values]
    df.reset_index(drop=True, inplace=True)

    # 6. 项目名称标准化
    df = standardize_projects(df)

    # ========== 聚合 ==========
    proj_summary = (
        df.groupby(["项目编码", "项目名称"], as_index=False)[Config.MONEY_COL]
        .sum()
        .rename(columns={Config.MONEY_COL: "机票费用合计（泰铢）"})
        .sort_values("机票费用合计（泰铢）", ascending=False)
    )

    monthly_pivot = df.pivot_table(
        index=["项目编码", "项目名称"],
        columns="年-月",
        values=Config.MONEY_COL,
        aggfunc="sum",
        fill_value=0,
    )
    if not monthly_pivot.empty:
        cols_sorted = sorted(
            monthly_pivot.columns,
            key=lambda s: (int(str(s).split("-")[0]), int(str(s).split("-")[1])) if "-" in str(s) else (0, 0),
        )
        monthly_pivot = monthly_pivot.reindex(columns=cols_sorted)

    # ========== 输出主汇总 Excel ==========
    out_dir = Path(Config.OUTPUT_STD)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "机票汇总.xlsx"

    display_cols = [c for c in Config.TARGET_COLUMNS if c in df.columns]

    non_numeric_headers = {"所属单位", "姓名", "订票时间", "年-月", "去程机票日期", "回程机票日期",
                           "出发地-到达地", "本次机票航班号", "舱位", "机票类型", "票代", "护照", "项目编码", "项目名称",
                           "原始项目名称", "_来源文件"}

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # Sheet1: 基础汇总（所有行）
        df[display_cols].to_excel(writer, sheet_name="基础汇总", index=False)
        # Sheet2: 按项目汇总成本
        proj_summary.to_excel(writer, sheet_name="按项目汇总", index=False)
        # Sheet3: 按项目×月度透视
        if not monthly_pivot.empty:
            monthly_pivot.to_excel(writer, sheet_name="月度透视")

    wb = load_workbook(out_path)
    for ws in wb.worksheets:
        headers = [cell.value for cell in ws[1]]
        numeric_start = next(
            (i + 1 for i, h in enumerate(headers) if h and str(h) not in non_numeric_headers),
            3,
        )
        apply_common_excel_layout(
            ws,
            numeric_start_col=numeric_start,
            freeze_panes="B2",
            thousand_number_format="#,##0.00",
        )
    wb.save(out_path)
    print(f"✅ 主汇总已输出：{out_path}")

    # ========== 按项目拆分 Excel ==========
    export_per_project_excel(
        df=df,
        output_dir=out_dir / "按项目拆分",
    )

    print("\n🎉 机票台账处理完成！")

    # ========== 输出 HTML 分析报告 ==========
    html_path = html_data_path() / "机票分析报告.html"
    build_html_report(df, out_path=html_path)
    print(f"✅ HTML 报告已输出：{html_path}")


if __name__ == "__main__":
    main()

