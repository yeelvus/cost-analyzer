# 广联达预算人材机全流程处理工具 - 样式完整保留版（修复版）
# 作者: Grok (基于用户反馈修复)
# 日期: 2025-10-31
# 修复内容:
#   1. 修复 .xls 填充样式复制错误：`XFBackground` 无 `pattern` 属性 → 使用 `fill_pattern`
#   2. 修复字体颜色类型错误：`openpyxl` 要求 `color` 为字符串 → 统一转为 hex
#   3. 修复索引越界：主表样式复制时严格限制行列范围
#   4. 增加异常捕获，防止单文件失败影响整体
# 功能: 
#   1. 直接处理 .xls/.xlsx，完整保留原始样式
#   2. 输出 _processed.xlsx 包含：原始所有工作表（带样式） + 统计结果（主表保留原始列样式）
#   3. 总汇总文件带时间戳
# 使用:
#   交互模式: python this_script.py (弹窗选择输入/输出文件夹和列)
#   自动模式: 设置 AUTO_MODE = True 并填写 INPUT_DIR / OUTPUT_DIR / SELECTED_COLUMNS 后运行
# 依赖: pandas, openpyxl, xlrd (<2.0), tkinter

import os
import json
import pandas as pd
import xlrd
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill, numbers
from openpyxl.styles.colors import Color
try:
    from tkinter import Tk, filedialog, messagebox, Listbox, MULTIPLE, Button, Scrollbar, END, RIGHT, Y
except Exception:
    Tk = None
    filedialog = None
    messagebox = None
import re
import logging
from datetime import datetime
import difflib

# ==================== 日志配置 ====================
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler("run.log", encoding="utf-8")
    ]
)
logger = logging.getLogger(__name__)

# ==================== 配置区 ====================
PROJECT_CODE_MAPPING_PATH = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/0_项目基本信息情况/项目编码名称汇总.xlsx"
# 自动模式控制:
#   AUTO_MODE = True 时不弹窗, 使用下面的 INPUT_DIR / OUTPUT_DIR / SELECTED_COLUMNS
#   SELECTED_COLUMNS 为列名列表，None 或空列表表示全部保留。
AUTO_MODE = True  # 修改为 False 可恢复原交互模式
INPUT_DIR = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/6_预算部分/5.1_主材汇总（广联达软件导出）/1_主材汇总-广联达导出文件"  # 请替换为实际输入目录
OUTPUT_DIR = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/6_预算部分/5.1_主材汇总（广联达软件导出）/2_数据汇总分析"      # 请替换为实际输出目录
SELECTED_COLUMNS = []  # 空列表表示保留全部；可填: ["编码", "类别", "名称", "单位", "含税市场价合计"]
KEYWORD_IN_SHEETNAME = "标段人材机汇总"
CODE_COL = "编码"
CATEGORY_COL = "类别"
NAME_COL = "名称"
UNIT_COL = "单位"
TARGET_COL_NAME = "不含税市场价合计"
FALLBACK_TARGET_COL_NAME = "含税市场价合计"
NEW_COL = "不含税市场价合计（万）"
MATERIAL_CATS = {"材", "商浆", "商砼", "主"}
PEOPLE_CAT = "人"
PIPE_CAT = "管"
MECH_CATS = {"设", "机"}
CONFIG_FILE = "selected_columns_cache.json"

MATERIAL_SUB_FILTERS = [
    ("钢筋", ["钢筋"], [], ["t", "kg"]),
    ("混凝土", ["混凝土"], [], ["m³", "m3", "立方米"]),
    ("砂浆", ["砂浆"], ["踢脚线"], []),
    ("碎石", ["碎石"], [], []),
    ("砖", ["砖", "砌块"], ["釉面砖", "花岗岩地砖"], []),
]

# 材料单位映射
MATERIAL_UNITS = {
    "钢筋": "kg",
    "混凝土": "m³",
    "砂浆": "m³",
    "碎石": "m³",
    "砖": "块",
    "其他材料": ""
}

# 单位转换规则（转换为标准单位）
UNIT_CONVERSION = {
    "t": 1000,      # 吨转kg
    "kg": 1,        # kg保持不变
    "m³": 1,        # 立方米保持不变
    "m3": 1,        # m3 = m³
    "立方米": 1,    # 立方米保持不变
}

SUMMARY_SHEET_NAME = "汇总_不含税市场价"
AMOUNT_COL = "金额"
TOTAL_SUMMARY_PREFIX = "【总汇总】所有标段_"

BASE_COLUMNS = [
    "标段人材机汇总", "人_明细", "管_明细", "设机_明细",
    "材类明细", "其他类别_明细",
    "钢筋", "混凝土", "砂浆", "碎石", "砖", "其他材料"
]

GROUP1 = ["人_明细", "管_明细", "设机_明细", "材类明细", "其他类别_明细"]
GROUP2 = ["钢筋", "混凝土", "砂浆", "碎石", "砖", "其他材料"]

FINAL_COLUMNS = (
    ["序号", "项目编码", "项目名称", "币种"] +
    BASE_COLUMNS
)

MAIN_KEYWORDS = ["标段人材机汇总", "人材机汇总", "汇总表"]
QUANTITY_COL = "数量"  # 数量列名
PRICE_COL_CANDIDATES = [
    "不含税市场价",
    "含税市场价",
    "市场价",
    "单价",
    "材料单价",
    "不含税市场价(元)",
    "含税市场价(元)",
    TARGET_COL_NAME,
    FALLBACK_TARGET_COL_NAME,
]

# ==================== 工具函数 ====================
def choose_folder(prompt="选择文件夹"):
    if AUTO_MODE:
        return None
    if Tk is None or filedialog is None:
        return None
    root = Tk(); root.withdraw()
    folder = filedialog.askdirectory(title=prompt)
    root.destroy()
    return folder

def ensure_numeric_series(s):
    return pd.to_numeric(s.astype(str).str.replace(",", "").str.strip(), errors='coerce')

def extract_project_code_and_name(filename):
    """从文件名中提取项目编码和项目名称
    
    Returns:
        (code, name) 元组
    """
    name = Path(filename).stem
    if "_processed" in name:
        name = name.replace("_processed", "")
    
    # 尝试匹配项目编码格式 THCC-YYYY-XXX_
    code_match = re.match(r"^(THCC-\d{4}-\d{3})_(.*)", name)
    if code_match:
        code = code_match.group(1)
        rest = code_match.group(2)
        # 提取项目名称（去除后缀）
        project_name = re.split(r"_标段|_单位工程|_人材机", rest)[0].strip()
        return code, project_name
    
    # 兼容旧格式（数字开头）
    match = re.match(r"^\d+[_ -]?(.*)", name)
    if match:
        project = match.group(1)
        project = re.sub(r"[_ -]?标段.*|0$", "", project, flags=re.IGNORECASE).strip()
        return None, project if project else name
    
    for suffix in ["标段人材机汇总", "标段", "汇总", "0"]:
        if suffix in name:
            project = name.split(suffix)[0].strip()
            project = re.sub(r"[_ -]+$", "", project)
            if project:
                return None, project
    
    return None, re.sub(r"^\d+[_ -]*", "", name).strip()

def load_project_code_mapping():
    """加载项目编码映射表"""
    try:
        df = pd.read_excel(PROJECT_CODE_MAPPING_PATH, sheet_name=0)
        mapping = {}
        
        # 读取所有可能的名称列
        name_cols = ["项目长名称", "项目名称（合同管理）", "项目名称（成本梳理）", "最近匹配文件", "项目名称", "项目名"]
        code_col = "项目编码"
        
        if code_col not in df.columns:
            logger.warning("项目编码映射表中未找到'项目编码'列")
            return {}
        
        for idx, row in df.iterrows():
            code = row.get(code_col)
            if pd.isna(code):
                continue
            code = str(code).strip()
            
            # 从所有名称列中提取
            for col in name_cols:
                if col in df.columns:
                    name = row.get(col)
                    if pd.notna(name):
                        name_str = str(name).strip()
                        if name_str and name_str not in mapping:
                            mapping[name_str] = code
        
        return mapping
    except Exception as e:
        logger.warning(f"加载项目编码映射表失败: {e}")
        return {}

def match_project_code(project_name, mapping, cutoff=0.75):
    """匹配项目编码"""
    if not mapping:
        return None
    
    # 精确匹配
    if project_name in mapping:
        return mapping[project_name]
    
    # 部分匹配
    for key, code in mapping.items():
        if project_name in key or key in project_name:
            return code
    
    # 模糊匹配
    matches = difflib.get_close_matches(project_name, mapping.keys(), n=1, cutoff=cutoff)
    if matches:
        return mapping[matches[0]]
    
    return None

def detect_currency(df):
    """根据混凝土单价判断币种
    
    Args:
        df: 包含材料数据的DataFrame，需要有'名称'、目标金额列(不含税优先)和'数量'列
    
    Returns:
        str: 'CNY'(人民币) 或 'THB'(泰铢) 或 'Unknown'
    """
    try:
        # 查找混凝土相关数据
        if NAME_COL not in df.columns:
            return "Unknown"
        
        name_series = df[NAME_COL].astype(str).fillna("")
        concrete_mask = name_series.str.contains("混凝土", case=False, na=False)
        df_concrete = df[concrete_mask].copy()
        
        if df_concrete.empty:
            logger.debug("  未找到混凝土数据，无法判断币种")
            return "Unknown"
        
        # 计算单价 = 目标金额列 / 数量
        if TARGET_COL_NAME in df_concrete.columns:
            df_concrete[TARGET_COL_NAME] = ensure_numeric_series(df_concrete[TARGET_COL_NAME])
            
            # 尝试找到数量列
            quantity_col = None
            for col in ["数量", "工程量", "quantity"]:
                if col in df_concrete.columns:
                    quantity_col = col
                    break
            
            if quantity_col:
                df_concrete[quantity_col] = ensure_numeric_series(df_concrete[quantity_col])
                # 过滤掉数量为0或空的行
                valid_mask = (df_concrete[quantity_col] > 0) & (df_concrete[TARGET_COL_NAME] > 0)
                df_valid = df_concrete[valid_mask]
                
                if not df_valid.empty:
                    # 计算单价
                    unit_prices = df_valid[TARGET_COL_NAME] / df_valid[quantity_col]
                    avg_unit_price = unit_prices.mean()
                    
                    logger.info(f"  混凝土平均单价: {avg_unit_price:.2f}")
                    
                    # 判断币种：400左右为人民币，1700左右为泰铢
                    # 设定阈值：< 1000 为人民币，>= 1000 为泰铢
                    if avg_unit_price < 1000:
                        return "CNY"
                    else:
                        return "THB"
        
        return "Unknown"
    except Exception as e:
        logger.warning(f"  币种判断失败: {e}")
        return "Unknown"

def choose_columns(columns):
    # 自动模式：直接使用 SELECTED_COLUMNS 或全部
    if AUTO_MODE:
        if isinstance(SELECTED_COLUMNS, (list, tuple)) and SELECTED_COLUMNS:
            return [c for c in SELECTED_COLUMNS if c in columns]
        return columns
    # 交互模式
    if Tk is None:
        return columns
    try:
        if os.path.exists(CONFIG_FILE):
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                cache = json.load(f)
            if cache and messagebox and Tk:
                root_tmp = Tk(); root_tmp.withdraw()
                use_prev = messagebox.askyesno("列选择", "检测到上次的列选择，是否沿用？")
                root_tmp.destroy()
                if use_prev:
                    return cache
    except Exception:
        pass
    root = Tk(); root.title("请选择需要保留的列（可多选，Ctrl/Shift多选）")
    lb = Listbox(root, selectmode=MULTIPLE, width=60, height=20)
    sb = Scrollbar(root, orient="vertical")
    sb.pack(side=RIGHT, fill=Y)
    lb.config(yscrollcommand=sb.set); sb.config(command=lb.yview)
    lb.pack()
    for col in columns: lb.insert(END, col)
    selected_cols = []
    def confirm():
        for i in lb.curselection(): selected_cols.append(lb.get(i))
        root.destroy()
    Button(root, text="确定", command=confirm).pack(); root.mainloop()
    try:
        if selected_cols:
            with open(CONFIG_FILE, "w", encoding="utf-8") as f:
                json.dump(selected_cols, f, ensure_ascii=False, indent=2)
    except Exception:
        pass
    return selected_cols or columns

def safe_color(color_obj, color_map):
    """将 xlrd 颜色索引转为 hex 字符串"""
    if color_obj is None:
        return None
    if isinstance(color_obj, str) and len(color_obj) == 6:
        return color_obj.upper()
    if isinstance(color_obj, int):
        rgb = color_map.get(color_obj)
        if isinstance(rgb, tuple) and len(rgb) == 3:
            return ''.join(f'{x:02X}' for x in rgb)
        return None
    if isinstance(color_obj, tuple) and len(color_obj) == 3:
        return ''.join(f'{x:02X}' for x in color_obj)
    return None

def calculate_material_quantity(df_sub, sheet_name):
    """计算材料数量（含单位转换）
    
    Args:
        df_sub: 材料子集DataFrame
        sheet_name: 材料名称（如"钢筋"、"混凝土"等）
    
    Returns:
        (total_quantity, unit) 元组
    """
    try:
        if df_sub.empty or QUANTITY_COL not in df_sub.columns or UNIT_COL not in df_sub.columns:
            return 0.0, MATERIAL_UNITS.get(sheet_name, "")
        
        df_sub = df_sub.copy()
        df_sub[QUANTITY_COL] = ensure_numeric_series(df_sub[QUANTITY_COL])
        df_sub[UNIT_COL] = df_sub[UNIT_COL].astype(str).fillna("").str.strip()
        
        total = 0.0
        for idx, row in df_sub.iterrows():
            qty = row[QUANTITY_COL]
            unit = row[UNIT_COL]
            if pd.isna(qty) or qty == 0:
                continue
            
            # 单位转换
            conversion_factor = UNIT_CONVERSION.get(unit, 1)
            total += qty * conversion_factor
        
        return round(total, 2), MATERIAL_UNITS.get(sheet_name, "")
    except Exception as e:
        logger.debug(f"计算数量失败 {sheet_name}: {e}")
        return 0.0, MATERIAL_UNITS.get(sheet_name, "")

# ==================== 样式复制核心函数 ====================
def copy_sheet_with_style(source_book, source_sheet_name, target_wb, target_name):
    """从 xlrd 或 openpyxl 复制 sheet 到 openpyxl，带完整样式"""
    if isinstance(source_book, xlrd.Book):
        # --- 处理 .xls ---
        xls_sheet = source_book.sheet_by_name(source_sheet_name)
        ws = target_wb.create_sheet(title=target_name)
        color_map = source_book.colour_map or {}

        for r in range(xls_sheet.nrows):
            # 行高
            if r in xls_sheet.rowinfo_map:
                rowinfo = xls_sheet.rowinfo_map[r]
                if rowinfo.height > 0:
                    ws.row_dimensions[r + 1].height = rowinfo.height / 20.0

            for c in range(xls_sheet.ncols):
                value = xls_sheet.cell_value(r, c)
                cell_xf_index = xls_sheet.cell_xf_index(r, c)
                if cell_xf_index >= len(source_book.xf_list):
                    continue
                xf = source_book.xf_list[cell_xf_index]
                tgt_cell = ws.cell(row=r + 1, column=c + 1, value=value)

                # 字体
                try:
                    font = source_book.font_list[xf.font_index]
                    color_hex = safe_color(font.colour_index, color_map)
                    tgt_cell.font = Font(
                        name=font.name or 'Calibri',
                        size=font.height / 20.0,
                        bold=bool(font.bold),
                        italic=bool(font.italic),
                        strike=bool(font.struck_out),
                        underline='single' if font.underline_type else None,
                        color=color_hex or '000000'
                    )
                except Exception as e:
                    logger.debug(f"字体复制失败: {e}")

                # 对齐
                try:
                    align = xf.alignment
                    horz = {0: 'general', 1: 'left', 2: 'center', 3: 'right'}.get(align.hor_align, 'general')
                    vert = {0: 'bottom', 1: 'top', 2: 'center', 3: 'bottom'}.get(align.vert_align, 'bottom')
                    tgt_cell.alignment = Alignment(
                        horizontal=horz,
                        vertical=vert,
                        wrap_text=bool(align.text_wrapped)
                    )
                except Exception as e:
                    logger.debug(f"对齐复制失败: {e}")

                # 边框
                try:
                    border = xf.border
                    sides = {}
                    for side_name, style in [
                        ('left', border.left_line_style),
                        ('right', border.right_line_style),
                        ('top', border.top_line_style),
                        ('bottom', border.bottom_line_style)
                    ]:
                        if style and style > 0:
                            sides[side_name] = Side(style='thin')
                    if sides:
                        tgt_cell.border = Border(**sides)
                except Exception as e:
                    logger.debug(f"边框复制失败: {e}")

                # 填充 - 修复：使用 fill_pattern
                try:
                    fill = xf.background
                    if fill.fill_pattern != 0:  # xlrd 中属性为 fill_pattern
                        fg_color = safe_color(fill.pattern_colour_index, color_map)
                        if fg_color:
                            tgt_cell.fill = PatternFill(start_color=fg_color, end_color=fg_color, fill_type="solid")
                except Exception as e:
                    logger.debug(f"填充复制失败: {e}")

                # 数字格式
                try:
                    fmt = source_book.format_map[xf.format_index].format_str
                    if fmt and fmt != 'General':
                        tgt_cell.number_format = fmt
                except Exception as e:
                    logger.debug(f"数字格式复制失败: {e}")

        # 列宽
        try:
            for col_idx, colinfo in xls_sheet.colinfo_map.items():
                width = colinfo.width / 256.0
                if width > 0:
                    ws.column_dimensions[get_column_letter(col_idx + 1)].width = width
        except Exception as e:
            logger.debug(f"列宽复制失败: {e}")

        # 合并单元格
        try:
            for rlo, rhi, clo, chi in xls_sheet.merged_cells:
                if rhi > rlo and chi > clo:
                    ws.merge_cells(start_row=rlo + 1, start_column=clo + 1, end_row=rhi, end_column=chi)
        except Exception as e:
            logger.debug(f"合并单元格复制失败: {e}")

        # 筛选 + 冻结
        if xls_sheet.nrows > 1:
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions

    else:
        # --- 处理 .xlsx ---
        src_ws = source_book[source_sheet_name]
        ws = target_wb.create_sheet(title=target_name)

        for row in src_ws.iter_rows():
            for cell in row:
                new_cell = ws.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    new_cell.font = cell.font.copy()
                    new_cell.border = cell.border.copy()
                    new_cell.fill = cell.fill.copy()
                    new_cell.number_format = cell.number_format
                    new_cell.alignment = cell.alignment.copy()

        # 列宽、行高
        for col, dim in src_ws.column_dimensions.items():
            ws.column_dimensions[col].width = dim.width
        for row, dim in src_ws.row_dimensions.items():
            ws.row_dimensions[row].height = dim.height

        # 合并单元格
        for m in src_ws.merged_cells.ranges:
            ws.merge_cells(str(m))

        ws.freeze_panes = src_ws.freeze_panes
        ws.auto_filter.ref = src_ws.auto_filter.ref

    return ws

# ==================== 单文件处理函数（带样式） ====================
def process_file(in_path, out_dir, selected_cols_global=None):
    p = Path(in_path)
    logger.info(f"\n处理文件: {p.name}")

    # 1. 打开原始文件
    is_xls = p.suffix.lower() == ".xls"
    try:
        if is_xls:
            source_book = xlrd.open_workbook(str(p), formatting_info=True)
        else:
            source_book = load_workbook(str(p), data_only=False)
    except Exception as e:
        logger.error(f"  打开文件失败: {e}")
        return

    # 2. 读取 sheet 名称
    sheet_names = [s.name for s in source_book.sheets()] if is_xls else source_book.sheetnames

    # 4. 读取数据
    try:
        all_dfs = pd.read_excel(str(p), sheet_name=None, dtype=object)
    except Exception as e:
        logger.error(f"  读取数据失败: {e}")
        return

    # 3/4. 在所有工作表中查找含关键列的目标表（不再校验表名，只校验列）
    required = {CODE_COL, CATEGORY_COL, TARGET_COL_NAME}
    target_sheet = None
    for sname, df0 in all_dfs.items():
        cols = [str(c).strip() for c in df0.columns]
        if required.issubset(set(cols)):
            target_sheet = sname
            break
    if not target_sheet:
        logger.warning("  跳过：未找到包含关键列的工作表")
        return

    df = all_dfs[target_sheet].copy()
    df.columns = [str(c).strip() for c in df.columns]

    # 5. 列选择
    if selected_cols_global is not None:
        selected_cols = selected_cols_global
    else:
        selected_cols = choose_columns(df.columns.tolist())

    keep_cols = [c for c in selected_cols if c in df.columns] if selected_cols else df.columns.tolist()
    missing = [c for c in selected_cols if c not in df.columns] if selected_cols else []
    if missing:
        logger.info(f"  忽略缺失列：{missing}")

    # 6. 数据处理
    if TARGET_COL_NAME not in df.columns and FALLBACK_TARGET_COL_NAME in df.columns:
        logger.warning(f"  未找到列 '{TARGET_COL_NAME}'，将使用 '{FALLBACK_TARGET_COL_NAME}' 进行汇总")
        df[TARGET_COL_NAME] = df[FALLBACK_TARGET_COL_NAME]
    if TARGET_COL_NAME not in df.columns:
        logger.warning(f"  跳过：缺少目标金额列 '{TARGET_COL_NAME}'")
        return
    df[TARGET_COL_NAME] = ensure_numeric_series(df[TARGET_COL_NAME])

    final_keep = [c for c in keep_cols if c in df.columns]
    if not final_keep:
        final_keep = df.columns.tolist()
    df = df[final_keep].copy()

    # 7. 创建输出工作簿
    out_path = os.path.join(out_dir, f"{p.stem}_processed.xlsx")
    wb = Workbook()
    wb.remove(wb.active)

    # --- 复制原始 sheet ---
    logger.info("  正在复制原始工作表（保留样式）...")
    for sheet_name in sheet_names:
        safe_name = sheet_name[:31]
        counter = 1
        temp_name = safe_name
        while temp_name in wb.sheetnames:
            temp_name = f"{safe_name[:28]}({counter})"
            counter += 1
        try:
            copy_sheet_with_style(source_book, sheet_name, wb, temp_name)
        except Exception as e:
            logger.error(f"  复制 sheet 失败 {sheet_name}: {e}")

    # --- 主表（带样式）---
    main_name = f"{target_sheet[:28]}_统计"
    main_ws = wb.create_sheet(main_name)
    for r in dataframe_to_rows(df, index=False, header=True):
        main_ws.append(r)

    # 应用原始样式到主表（仅数据行）
    try:
        src_ws = source_book.sheet_by_name(target_sheet) if is_xls else source_book[target_sheet]
        max_r = min(main_ws.max_row, src_ws.nrows if is_xls else src_ws.max_row)
        max_c = min(main_ws.max_column, src_ws.ncols if is_xls else src_ws.max_column)
        color_map = source_book.colour_map if is_xls else {}

        for r in range(2, max_r + 1):  # 跳过表头
            for c in range(1, max_c + 1):
                tgt_cell = main_ws.cell(r, c)
                src_r = r - 1 if is_xls else r
                src_c = c - 1 if is_xls else c

                if is_xls:
                    xf_index = src_ws.cell_xf_index(src_r - 1, src_c - 1)
                    xf = source_book.xf_list[xf_index]
                    # 字体
                    font = source_book.font_list[xf.font_index]
                    color_hex = safe_color(font.colour_index, color_map)
                    tgt_cell.font = Font(
                        name=font.name or 'Calibri',
                        size=font.height / 20.0,
                        bold=bool(font.bold),
                        italic=bool(font.italic),
                        color=color_hex or '000000'
                    )
                    # 填充
                    fill = xf.background
                    if fill.fill_pattern != 0:
                        fg_color = safe_color(fill.pattern_colour_index, color_map)
                        if fg_color:
                            tgt_cell.fill = PatternFill(start_color=fg_color, end_color=fg_color, fill_type="solid")
                    # 边框、对齐、数字格式同上
                    border = xf.border
                    sides = {}
                    for side, style in [('left', border.left_line_style), ('right', border.right_line_style),
                                       ('top', border.top_line_style), ('bottom', border.bottom_line_style)]:
                        if style > 0:
                            sides[side] = Side(style='thin')
                    tgt_cell.border = Border(**sides)
                    align = xf.alignment
                    tgt_cell.alignment = Alignment(
                        horizontal={1: 'left', 2: 'center', 3: 'right'}.get(align.hor_align, 'general'),
                        vertical={1: 'top', 2: 'center', 3: 'bottom'}.get(align.vert_align, 'bottom')
                    )
                    fmt = source_book.format_map[xf.format_index].format_str
                    tgt_cell.number_format = fmt if fmt != 'General' else numbers.FORMAT_GENERAL
                else:
                    src_cell = src_ws.cell(r, c)
                    if src_cell.has_style:
                        tgt_cell.font = src_cell.font.copy()
                        tgt_cell.fill = src_cell.fill.copy()
                        tgt_cell.border = src_cell.border.copy()
                        tgt_cell.alignment = src_cell.alignment.copy()
                        tgt_cell.number_format = src_cell.number_format
    except Exception as e:
        logger.error(f"  主表样式复制失败: {e}")

    main_ws.freeze_panes = 'A2'
    main_ws.auto_filter.ref = main_ws.dimensions

    # --- 明细表和汇总表 ---
    def append_df(df_ex, sheetname):
        if df_ex is None or df_ex.empty:
            return None, 0.0
        safe_name = sheetname[:31]
        counter = 1
        temp_name = safe_name
        while temp_name in wb.sheetnames:
            temp_name = f"{safe_name[:28]}({counter})"
            counter += 1
        wsx = wb.create_sheet(temp_name)
        for r in dataframe_to_rows(df_ex, index=False, header=True):
            wsx.append(r)
        format_sheet(wsx, [TARGET_COL_NAME])
        total = df_ex[TARGET_COL_NAME].sum() if TARGET_COL_NAME in df_ex.columns else 0.0
        return wsx, total

    # 判断币种
    currency = detect_currency(df)
    logger.info(f"  检测到币种: {currency}")
    
    summary_rows = []
    main_total = df[TARGET_COL_NAME].sum() if TARGET_COL_NAME in df.columns else 0.0
    summary_rows.append({"类别": f"{target_sheet}", "金额": main_total})
    # 添加币种信息行
    summary_rows.append({"类别": "币种", "金额": currency})

    if CATEGORY_COL in df.columns:
        cat_series = df[CATEGORY_COL].astype(str).str.strip()

        df_people = df[cat_series == PEOPLE_CAT]
        _, people_total = append_df(df_people, "人_明细")
        summary_rows.append({"类别": "人_明细", "金额": people_total})

        df_pipe = df[cat_series == PIPE_CAT]
        _, pipe_total = append_df(df_pipe, "管_明细")
        summary_rows.append({"类别": "管_明细", "金额": pipe_total})

        df_mech = df[cat_series.isin(MECH_CATS)]
        _, mech_total = append_df(df_mech, "设机_明细")
        summary_rows.append({"类别": "设机_明细", "金额": mech_total})

        df_mat = df[cat_series.isin(MATERIAL_CATS)]
        _, mat_total = append_df(df_mat, "材类明细")
        summary_rows.append({"类别": "材类明细", "金额": mat_total})

        sub_totals = {}
        if NAME_COL in df_mat.columns and not df_mat.empty:
            name_series = df_mat[NAME_COL].astype(str).fillna("")
            unit_series = df_mat[UNIT_COL].astype(str).fillna("") if UNIT_COL in df_mat.columns else pd.Series("", index=df_mat.index)

            for sheet_name, include_kw, exclude_kw, unit_kw in MATERIAL_SUB_FILTERS:
                mask = pd.Series(False, index=df_mat.index)
                for kw in include_kw:
                    mask |= name_series.str.contains(kw, case=False, na=False)
                for kw in exclude_kw:
                    mask &= ~name_series.str.contains(kw, case=False, na=False)
                if unit_kw:
                    unit_mask = pd.Series(False, index=df_mat.index)
                    for u in unit_kw:
                        unit_mask |= unit_series.str.contains(u, case=False, na=False)
                    mask &= unit_mask

                df_sub = df_mat[mask].copy()
                if not df_sub.empty:
                    _, total = append_df(df_sub, sheet_name)
                    sub_totals[sheet_name] = total
                    summary_rows.append({"类别": sheet_name, "金额": total})

        sub_sum = sum(sub_totals.values())
        other_mat = mat_total - sub_sum
        summary_rows.append({"类别": "其他材料", "金额": other_mat})

        known_cats = MATERIAL_CATS | {PEOPLE_CAT, PIPE_CAT} | MECH_CATS
        df_other = df[~cat_series.isin(known_cats)]
        if not df_other.empty:
            _, other_cat_total = append_df(df_other, "其他类别_明细")
            summary_rows.append({"类别": "其他类别_明细", "金额": other_cat_total})

    summary_df = pd.DataFrame(summary_rows)
    ws_summary = wb.create_sheet(SUMMARY_SHEET_NAME)
    for r in dataframe_to_rows(summary_df, index=False, header=True):
        ws_summary.append(r)
    format_sheet(ws_summary, ["金额"])

    # 保存
    try:
        wb.save(out_path)
        logger.info(f"  已生成：{out_path}")
    except Exception as e:
        logger.error(f"  保存失败：{e}")
    finally:
        if not is_xls:
            source_book.close()

def format_sheet(ws, numeric_cols):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []
    for col_name in numeric_cols:
        if col_name in header:
            col_idx = header.index(col_name) + 1
            col_letter = get_column_letter(col_idx)
            for cell in ws[col_letter][1:]:
                cell.number_format = '#,##0'

def cleanup_old_total_summary_excels(folder, current_output_path):
    """删除旧的总汇总文件，仅保留当前新生成文件。"""
    current_name = Path(current_output_path).name
    excel_exts = {".xls", ".xlsx", ".xlsm"}
    deleted, failed = [], []

    for fname in os.listdir(folder):
        fpath = os.path.join(folder, fname)
        suffix = Path(fname).suffix.lower()
        if not os.path.isfile(fpath):
            continue
        if not fname.startswith(TOTAL_SUMMARY_PREFIX):
            continue
        if suffix not in excel_exts:
            continue
        if fname == current_name:
            continue
        try:
            os.remove(fpath)
            deleted.append(fname)
        except Exception as e:
            failed.append((fname, str(e)))

    if deleted:
        logger.info(f"已删除 {len(deleted)} 个旧总汇总文件")
        for name in deleted:
            logger.info(f"  删除: {name}")
    else:
        logger.info("未发现需要删除的旧总汇总文件")

    if failed:
        logger.warning(f"有 {len(failed)} 个旧总汇总文件删除失败")
        for name, err in failed:
            logger.warning(f"  删除失败: {name} ({err})")

# ==================== 总汇总函数（不变） ====================
def generate_summary_report(folder):
    logger.info("开始总汇总...")
    files = [f for f in os.listdir(folder) if f.endswith("_processed.xlsx")]
    if not files:
        logger.error("未找到 _processed.xlsx 文件！")
        print("\n[错误] 未找到 _processed.xlsx 文件！\n")
        return

    def sort_key(f):
        match = re.search(r"(\d+)", f)
        return int(match.group(1)) if match else 9999
    files.sort(key=sort_key)

    # 加载项目编码映射表
    logger.info("加载项目编码映射表...")
    code_mapping = load_project_code_mapping()
    logger.info(f"已加载 {len(code_mapping)} 个项目映射")

    all_data = []
    file_paths = {}

    for idx, file in enumerate(files, start=1):
        file_path = os.path.join(folder, file)
        file_paths[idx] = file_path
        
        # 提取项目编码和名称
        project_code, project_name = extract_project_code_and_name(file)
        
        # 如果文件名中没有编码，尝试从映射表匹配
        if not project_code:
            project_code = match_project_code(project_name, code_mapping)
            if not project_code:
                logger.warning(f"  未能匹配项目编码 - {project_name}")

        try:
            df_summary = pd.read_excel(file_path, sheet_name=SUMMARY_SHEET_NAME)
        except Exception as e:
            logger.error(f"  跳过 {file}：读取失败 ({e})")
            continue

        # 读取币种信息
        currency = "Unknown"
        if CATEGORY_COL in df_summary.columns:
            currency_row = df_summary[df_summary[CATEGORY_COL].astype(str) == "币种"]
            if not currency_row.empty and AMOUNT_COL in df_summary.columns:
                currency_val = currency_row[AMOUNT_COL].iloc[0]
                if pd.notna(currency_val):
                    currency = str(currency_val).strip()

        row = {"序号": idx, "项目编码": project_code or "", "项目名称": project_name, "币种": currency}

        # 保留原有口径：标段人材机汇总优先从主关键字匹配，不做强制回算
        row["标段人材机汇总"] = 0.0
        if CATEGORY_COL in df_summary.columns and AMOUNT_COL in df_summary.columns:
            cat_series = df_summary[CATEGORY_COL].astype(str)
            amt_series = pd.to_numeric(df_summary[AMOUNT_COL], errors='coerce').fillna(0)
            mask = pd.Series(False, index=df_summary.index)
            for kw in MAIN_KEYWORDS:
                mask |= cat_series.str.contains(kw, case=False, na=False)
            row["标段人材机汇总"] = round(float(amt_series[mask].iloc[0]), 2) if mask.any() else round(float(amt_series.iloc[0]), 2)

        if CATEGORY_COL in df_summary.columns and AMOUNT_COL in df_summary.columns:
            cat_to_amt = dict(zip(
                df_summary[CATEGORY_COL].astype(str),
                pd.to_numeric(df_summary[AMOUNT_COL], errors='coerce').fillna(0)
            ))
            for col in BASE_COLUMNS[1:]:
                row[col] = round(float(cat_to_amt.get(col, 0.0)), 2)
        else:
            for col in BASE_COLUMNS[1:]:
                row[col] = 0.0

        all_data.append(row)
        code_display = f"({project_code})" if project_code else "(无编码)"
        logger.info(f"  [{idx}] {code_display} {project_name}")

    if not all_data:
        logger.error("无有效数据！")
        print("\n[错误] 无有效数据！\n")
        return

    result_df = pd.DataFrame(all_data)
    result_df = result_df[FINAL_COLUMNS]

    total_row = {"序号": "合计", "项目编码": "", "项目名称": "", "币种": ""}
    for col in BASE_COLUMNS:
        total_row[col] = result_df[col].sum()
    result_df = pd.concat([result_df, pd.DataFrame([total_row])], ignore_index=True)

    # 添加时间戳
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(folder, f"{TOTAL_SUMMARY_PREFIX}{timestamp}.xlsx")

    try:
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            result_df.to_excel(writer, index=False, sheet_name="总汇总")
            wb = writer.book
            ws = writer.sheets["总汇总"]

            for i, col in enumerate(result_df.columns, 1):
                max_len = max(result_df[col].astype(str).map(len).max(), len(str(col)))
                ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 50)

            header_font = Font(bold=True, color="FFFFFF")
            meta_header_fill = PatternFill("solid", fgColor="366092")
            total_header_fill = PatternFill("solid", fgColor="0F766E")
            detail_header_fill = PatternFill("solid", fgColor="D97706")
            sum_font = Font(bold=True, color="D00000")
            link_font = Font(color="0000FF", underline="single")

            total_header_cols = {"标段人材机汇总", "材类明细"}
            detail_header_cols = set(BASE_COLUMNS) - total_header_cols

            for cell in ws[1]:
                cell.font = header_font
                col_name = str(cell.value).strip() if cell.value is not None else ""
                if col_name in total_header_cols:
                    cell.fill = total_header_fill
                elif col_name in detail_header_cols:
                    cell.fill = detail_header_fill
                else:
                    cell.fill = meta_header_fill
                cell.alignment = Alignment(horizontal="center")

            # 为金额列应用千分位格式
            amount_cols = BASE_COLUMNS
            for col_name in amount_cols:
                if col_name in result_df.columns:
                    col_idx = list(result_df.columns).index(col_name) + 1
                    col_letter = get_column_letter(col_idx)
                    for cell in ws[col_letter][1:]:
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##0'

            # 超链接(第3列是项目名称)
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row - 1, min_col=3, max_col=3):
                cell = row[0]
                seq = cell.row - 1
                if seq in file_paths:
                    file_url = str(Path(file_paths[seq]).resolve()).replace("\\", "/")
                    cell.value = result_df.iloc[seq - 1, 2]  # 项目名称在第3列(索引2)
                    cell.hyperlink = f"file:///{file_url}"
                    cell.font = link_font

            for cell in ws[ws.max_row]:
                cell.font = sum_font

            ws.freeze_panes = "A2"

            # ============ 新增：材料数量统计表 ============
            material_names = ["钢筋", "混凝土", "砂浆", "碎石", "砖", "其他材料"]
            material_summary_rows = []
            
            # 数据行
            for idx, file in enumerate(files, start=1):
                file_path = os.path.join(folder, file)
                project_code, project_name = extract_project_code_and_name(file)
                if not project_code:
                    project_code = match_project_code(project_name, code_mapping)
                
                mat_row = {"项目编码": project_code or "", "项目名称": project_name}
                
                # 检测币种
                currency = "Unknown"
                try:
                    df_for_currency = pd.read_excel(file_path, sheet_name=0)
                    currency = detect_currency(df_for_currency)
                except Exception as e:
                    logger.debug(f"检测币种失败 {file}: {e}")
                mat_row["币种"] = currency
                
                try:
                    df_summary = pd.read_excel(file_path, sheet_name=SUMMARY_SHEET_NAME)
                    
                    for mat in material_names:
                        # 获取金额
                        mat_amount = 0.0
                        if CATEGORY_COL in df_summary.columns and AMOUNT_COL in df_summary.columns:
                            cat_to_amt = dict(zip(
                                df_summary[CATEGORY_COL].astype(str),
                                pd.to_numeric(df_summary[AMOUNT_COL], errors='coerce').fillna(0)
                            ))
                            mat_amount = round(float(cat_to_amt.get(mat, 0.0)), 2)
                        
                        # 从对应的材料Sheet读取数量
                        mat_quantity = 0.0
                        try:
                            # 尝试读取材料对应的Sheet
                            df_mat_sheet = pd.read_excel(file_path, sheet_name=mat)
                            
                            if QUANTITY_COL in df_mat_sheet.columns and UNIT_COL in df_mat_sheet.columns:
                                # 计算该材料的总数量
                                qty_series = ensure_numeric_series(df_mat_sheet[QUANTITY_COL])
                                unit_series = df_mat_sheet[UNIT_COL].astype(str).fillna("").str.strip()
                                price_col = next((col for col in PRICE_COL_CANDIDATES if col in df_mat_sheet.columns), None)
                                if price_col:
                                    price_series = ensure_numeric_series(df_mat_sheet[price_col]).fillna(0)
                                else:
                                    # 未找到价格列时保持原口径
                                    price_series = pd.Series(1, index=df_mat_sheet.index, dtype=float)
                                
                                for q, u, price in zip(qty_series, unit_series, price_series):
                                    # 排除价格为0(或无效)的数量
                                    if pd.notna(q) and q > 0 and pd.notna(price) and price > 0:
                                        conversion_factor = UNIT_CONVERSION.get(u, 1)
                                        mat_quantity += q * conversion_factor
                                
                                mat_quantity = round(mat_quantity, 2)
                        except Exception as e:
                            logger.debug(f"读取材料Sheet失败 {mat} ({file}): {e}")
                            mat_quantity = 0.0
                        
                        # 计算平均单价
                        avg_price = 0.0
                        if mat_quantity > 0:
                            avg_price = round(mat_amount / mat_quantity, 2)
                        
                        # 顺序：数量、单位、金额、平均单价
                        mat_row[f"{mat}_数量"] = mat_quantity
                        mat_row[f"{mat}_单位"] = MATERIAL_UNITS.get(mat, "")
                        mat_row[f"{mat}_金额"] = mat_amount
                        mat_row[f"{mat}_平均单价"] = avg_price
                    
                except Exception as e:
                    logger.warning(f"读取文件失败 {file}: {e}")
                    for mat in material_names:
                        mat_row[f"{mat}_数量"] = 0.0
                        mat_row[f"{mat}_单位"] = MATERIAL_UNITS.get(mat, "")
                        mat_row[f"{mat}_金额"] = 0.0
                        mat_row[f"{mat}_平均单价"] = 0.0
                
                material_summary_rows.append(mat_row)
            
            # 合计行
            total_mat_row = {"项目编码": "", "项目名称": "合计", "币种": "合计"}
            for mat in material_names:
                total_quantity = round(sum(
                    r.get(f"{mat}_数量", 0) for r in material_summary_rows 
                    if isinstance(r.get(f"{mat}_数量", 0), (int, float))
                ), 2)
                total_amount = round(sum(
                    r.get(f"{mat}_金额", 0) for r in material_summary_rows 
                    if isinstance(r.get(f"{mat}_金额", 0), (int, float))
                ), 2)
                total_avg_price = 0.0
                if total_quantity > 0:
                    total_avg_price = round(total_amount / total_quantity, 2)
                
                total_mat_row[f"{mat}_数量"] = total_quantity
                total_mat_row[f"{mat}_单位"] = MATERIAL_UNITS.get(mat, "")
                total_mat_row[f"{mat}_金额"] = total_amount
                total_mat_row[f"{mat}_平均单价"] = total_avg_price
            material_summary_rows.append(total_mat_row)
            
            # 写入新Sheet（按顺序：项目编码、项目名称、币种、钢筋_数量、钢筋_单位、钢筋_金额、钢筋_平均单价...）
            ordered_cols = ["项目编码", "项目名称", "币种"]
            for mat in material_names:
                ordered_cols.extend([f"{mat}_数量", f"{mat}_单位", f"{mat}_金额", f"{mat}_平均单价"])
            
            material_df = pd.DataFrame(material_summary_rows)[ordered_cols]
            material_df.to_excel(writer, index=False, sheet_name="材料汇总")
            ws_mat = writer.sheets["材料汇总"]

            # 每个材料分组（数量/单位/金额/平均单价）使用同一表头色，并区分不同材料
            material_group_colors = {
                "钢筋": PatternFill("solid", fgColor="1D4E89"),
                "混凝土": PatternFill("solid", fgColor="0F766E"),
                "砂浆": PatternFill("solid", fgColor="B45309"),
                "碎石": PatternFill("solid", fgColor="BE123C"),
                "砖": PatternFill("solid", fgColor="374151"),
                "其他材料": PatternFill("solid", fgColor="166534"),
            }
            
            # 格式化材料汇总表
            header_row = ws_mat[1]
            for cell in header_row:
                cell.font = header_font
                col_name = str(cell.value).strip() if cell.value is not None else ""
                mat_name = col_name.split("_", 1)[0] if "_" in col_name else ""
                cell.fill = material_group_colors.get(mat_name, meta_header_fill)
                cell.alignment = Alignment(horizontal="center")
            
            # 设置列宽
            for i, col in enumerate(material_df.columns, 1):
                max_len = max(material_df[col].astype(str).map(len).max() if col in material_df else 0, len(str(col)))
                ws_mat.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 50)
            
            # 金额、数量、平均单价列格式化
            for col_name in material_df.columns:
                if "_金额" in col_name or "_数量" in col_name or "_平均单价" in col_name:
                    col_idx = list(material_df.columns).index(col_name) + 1
                    col_letter = get_column_letter(col_idx)
                    for cell in ws_mat[col_letter][1:]:
                        if isinstance(cell.value, (int, float)):
                            if "_数量" in col_name:
                                # 数量仅保留2位小数，不使用金额格式
                                cell.number_format = '0.00'
                            else:
                                cell.number_format = '#,##0.00'

            # 单位列（如: 钢筋_单位）统一居中
            for col_name in material_df.columns:
                if col_name.endswith("_单位"):
                    col_idx = list(material_df.columns).index(col_name) + 1
                    col_letter = get_column_letter(col_idx)
                    for cell in ws_mat[col_letter][1:]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # 合计行加粗
            last_row_idx = ws_mat.max_row
            for cell in ws_mat[last_row_idx]:
                cell.font = sum_font
            
            ws_mat.freeze_panes = "A2"

        # 删除旧的“所有标段总汇总”Excel，仅保留当前新文件
        cleanup_old_total_summary_excels(folder, output_path)

        logger.info(f"\n汇总完成！共 {len(all_data)} 个标段")
        logger.info(f"文件：{output_path}")
        print(f"\n{'='*60}")
        print(f"汇总完成！")
        print(f"共 {len(all_data)} 个标段")
        print(f"文件：{output_path}")
        print(f"{'='*60}\n")

    except Exception as e:
        logger.error(f"保存失败：{e}")
        print(f"\n[错误] 保存失败：{e}\n")

# ==================== 主入口 ====================
def main():
    if AUTO_MODE:
        # 校验目录
        if not INPUT_DIR or not os.path.isdir(INPUT_DIR):
            print(f"[ERROR] INPUT_DIR 不存在: {INPUT_DIR}")
            return
        if not OUTPUT_DIR:
            print("[ERROR] 未配置 OUTPUT_DIR")
            return
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        exts = {".xls", ".xlsx", ".xlsm"}
        files = [os.path.join(INPUT_DIR, f) for f in os.listdir(INPUT_DIR) if Path(f).suffix.lower() in exts]
        if not files:
            print("[WARN] 输入目录中未找到 Excel 文件")
            return
        selected_cols_global = None
        for idx, f in enumerate(files):
            if idx == 0:
                try:
                    cols = pd.read_excel(f, nrows=0).columns.tolist()
                except Exception:
                    try:
                        cols = pd.read_excel(f, sheet_name=0, nrows=0).columns.tolist()
                    except Exception:
                        cols = []
                selected_cols_global = choose_columns(cols if cols else [])
            process_file(f, OUTPUT_DIR, selected_cols_global)
        generate_summary_report(OUTPUT_DIR)
        print(f"[OK] 自动模式完成，结果目录: {OUTPUT_DIR}")
        return
    # 交互模式
    if Tk is None:
        print("[WARN] 无法进入交互模式（Tk 不可用），请启用 AUTO_MODE")
        return
    root = Tk(); root.withdraw()
    input_dir = choose_folder("请选择输入文件夹（包含 Excel 文件）")
    if not input_dir:
        return
    output_dir = choose_folder("请选择输出文件夹（保存结果）")
    if not output_dir:
        return
    exts = {".xls", ".xlsx", ".xlsm"}
    files = [os.path.join(input_dir, f) for f in os.listdir(input_dir) if Path(f).suffix.lower() in exts]
    if not files:
        if messagebox:
            messagebox.showwarning("提示", "输入文件夹中未找到 Excel 文件！")
        else:
            print("[WARN] 输入文件夹中未找到 Excel 文件！")
        return
    selected_cols_global = None
    for idx, f in enumerate(files):
        if idx == 0:
            try:
                cols = pd.read_excel(f, nrows=0).columns.tolist()
            except Exception:
                try:
                    cols = pd.read_excel(f, sheet_name=0, nrows=0).columns.tolist()
                except Exception:
                    cols = []
            selected_cols_global = choose_columns(cols if cols else [])
        process_file(f, output_dir, selected_cols_global)
    generate_summary_report(output_dir)
    print(f"\n{'='*60}")
    print(f"全流程完成！")
    print(f"结果保存在：{output_dir}")
    print(f"{'='*60}\n")

if __name__ == "__main__":
    main()