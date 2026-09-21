#!/usr/bin/env python3
"""
步骤2：扫描 raw 各月台账，按项目对比上月，剔除未更新文件；
每个项目只保留 1 份最新有效文件 → EXTRACT_DIR
（B_采购部采购台账_extract）。

判定优先级：
  文件哈希相同 / 内容哈希相同 → 复用上月（不体现）
  金额+行数一致 → 建议复核后复用
  仅金额一致 → 建议人工复核
  金额变化 → 保留
"""
from __future__ import annotations

import hashlib
import re
import shutil
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter
import sys

_PKG_ROOT = Path(__file__).resolve().parents[1]
_CODE_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(_PKG_ROOT), str(_CODE_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from _采购路径 import (
    EXTRACT_DIR,
    RAW_DIR,
    extract_month_iso,
    extract_project_name,
    is_excel,
    month_sort_key,
)

AMOUNT_CANDIDATES = ("含税金额", "总金额", "价税合计", "含税总金额", "金额")
SKIP_DECISION = "复用上月，当前月不体现"
ANALYSIS_DIR_NAME = "去重分析输出"


def _norm(v: object) -> str:
    if v is None:
        return ""
    return str(v).replace("\n", " ").replace("\r", " ").strip()


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _best_amount_col(columns: list[str]) -> str | None:
    cols = [_norm(c) for c in columns]
    for name in AMOUNT_CANDIDATES:
        if name in cols:
            return name
    for c in cols:
        if "金额" in c:
            return c
    return None


def _summary_row_indexes(ws, columns: list[str], amount_col: str | None) -> set[int]:
    """识别金额列上的 SUM 汇总行（DataFrame 0-based 索引）。"""
    if ws is None or not amount_col:
        return set()
    try:
        col_idx = [_norm(c) for c in columns].index(amount_col) + 1
    except ValueError:
        return set()

    letter = get_column_letter(col_idx)
    found: set[int] = set()
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, col_idx)
        if cell.data_type != "f":
            continue
        f = str(cell.value).upper().replace(" ", "")
        if "SUM(" not in f:
            continue
        if f"{letter}:{letter}" in f or re.search(
            rf"\$?{letter}\$?\d+:\$?{letter}\$?\d+", f
        ):
            found.add(r - 2)
    return found


def excel_metrics(path: Path) -> tuple[float, int, str]:
    """返回 (含税金额合计, 数据行数, 内容哈希)。"""
    formula_wb = openpyxl.load_workbook(path, data_only=False)
    total_sum = 0.0
    total_rows = 0
    hash_acc = hashlib.sha256()

    try:
        xls = pd.ExcelFile(path)
        for sheet in xls.sheet_names:
            try:
                df = pd.read_excel(path, sheet_name=sheet)
            except Exception:
                continue
            if df.empty:
                continue

            df.columns = [_norm(c) for c in df.columns]
            amount_col = _best_amount_col(list(df.columns))
            fws = formula_wb[sheet] if sheet in formula_wb.sheetnames else None
            drop_idx = _summary_row_indexes(fws, list(df.columns), amount_col)
            if drop_idx:
                df = df.drop(index=[i for i in drop_idx if 0 <= i < len(df)]).reset_index(drop=True)

            for col in df.columns:
                if pd.api.types.is_numeric_dtype(df[col]):
                    df[col] = pd.to_numeric(df[col], errors="coerce").round(6)
                else:
                    df[col] = df[col].map(_norm)

            if amount_col and amount_col in df.columns:
                total_sum += float(pd.to_numeric(df[amount_col], errors="coerce").fillna(0).sum())
            total_rows += len(df)

            sig = df.copy()
            if amount_col and amount_col in sig.columns:
                sig[amount_col] = pd.to_numeric(sig[amount_col], errors="coerce").round(2)
            sig = sig.fillna("")
            payload = "|".join(sig.columns) + "\n"
            for row in sig.itertuples(index=False, name=None):
                payload += "\t".join(_norm(v) for v in row) + "\n"
            hash_acc.update(f"{sheet}:{hashlib.sha256(payload.encode()).hexdigest()}".encode())
    finally:
        formula_wb.close()

    return round(total_sum, 2), total_rows, hash_acc.hexdigest()


def scan_raw(raw_root: Path) -> pd.DataFrame:
    rows = []
    for path in sorted(raw_root.rglob("*")):
        if not is_excel(path):
            continue
        month = extract_month_iso(str(path))
        if not month:
            continue
        amount, n_rows, content_hash = excel_metrics(path)
        rows.append(
            {
                "月份": month,
                "项目名称": extract_project_name(path.name),
                "文件名": path.name,
                "文件路径": str(path),
                "含税金额合计": amount,
                "行数": n_rows,
                "文件哈希": _sha256_file(path),
                "内容哈希": content_hash,
            }
        )
    return pd.DataFrame(rows)


def decide(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(
            columns=list(df.columns) + ["对比结论", "处理建议", "是否在最新口径体现"]
        )

    out = []
    for _, g in df.groupby("项目名称", sort=False):
        g = g.sort_values("月份", key=lambda s: s.map(month_sort_key))
        prev = None
        for _, row in g.iterrows():
            reason, decision = "首月保留", "保留"
            if prev is not None:
                amt_diff = abs(float(row["含税金额合计"]) - float(prev["含税金额合计"]))
                row_diff = int(row["行数"]) - int(prev["行数"])
                if row["文件哈希"] == prev["文件哈希"]:
                    reason, decision = "与上月文件字节完全一致", SKIP_DECISION
                elif row["内容哈希"] == prev["内容哈希"]:
                    reason, decision = "与上月内容值一致（忽略文件元数据）", SKIP_DECISION
                elif amt_diff <= 0.01 and row_diff == 0:
                    reason, decision = "含税金额合计一致且行数一致（疑似未更新）", "建议复核后复用上月"
                elif amt_diff <= 0.01:
                    reason, decision = "含税金额合计一致但行数变化", "建议人工复核"
                else:
                    reason, decision = f"金额变化 {amt_diff:,.2f}", "保留"
            out.append(
                {
                    **row.to_dict(),
                    "对比结论": reason,
                    "处理建议": decision,
                    "是否在最新口径体现": "否" if decision == SKIP_DECISION else "是",
                }
            )
            prev = row
    return pd.DataFrame(out)


def export_latest_per_project(decision_df: pd.DataFrame, out_dir: Path) -> tuple[int, int]:
    """每项目保留最新非明确重复文件；分析文件放 去重分析输出/。"""
    out_dir.mkdir(parents=True, exist_ok=True)
    analysis = out_dir / ANALYSIS_DIR_NAME
    analysis.mkdir(exist_ok=True)

    # 清空输出根目录（保留分析子目录）
    for item in out_dir.iterdir():
        if item.name == ANALYSIS_DIR_NAME:
            continue
        if item.is_dir():
            shutil.rmtree(item)
        else:
            item.unlink()

    if decision_df.empty:
        return 0, 0

    cand = decision_df[decision_df["处理建议"] != SKIP_DECISION].copy()
    if cand.empty:
        return 0, 0

    priority = {"保留": 3, "建议复核后复用上月": 2, "建议人工复核": 1}
    cand["_mk"] = cand["月份"].map(month_sort_key)
    cand["_pri"] = cand["处理建议"].map(priority).fillna(0)
    selected = (
        cand.sort_values(["_mk", "_pri"], ascending=[False, False])
        .groupby("项目名称", sort=False, as_index=False)
        .first()
    )

    copied, conflicts = [], []
    used: set[str] = set()
    for _, row in selected.iterrows():
        src = Path(str(row["文件路径"]))
        if not src.is_file():
            continue
        dst = out_dir / src.name
        if src.name in used or dst.exists():
            conflicts.append({"项目名称": row["项目名称"], "文件名": src.name, "路径": str(src)})
            continue
        shutil.copy2(src, dst)
        used.add(src.name)
        copied.append(
            {
                "项目名称": row["项目名称"],
                "选中月份": row["月份"],
                "处理建议": row["处理建议"],
                "对比结论": row["对比结论"],
                "文件名": src.name,
            }
        )

    pd.DataFrame(copied).to_excel(analysis / "A_拷贝清单.xlsx", index=False)
    if conflicts:
        pd.DataFrame(conflicts).to_excel(analysis / "同名冲突清单.xlsx", index=False)
    return len(copied), len(conflicts)


def run(raw_root: Path = RAW_DIR, extract_dir: Path = EXTRACT_DIR) -> None:
    print(f"源目录: {raw_root}")
    print(f"输出目录: {extract_dir}")

    df = scan_raw(raw_root)
    if df.empty:
        print("未扫描到 Excel，结束")
        decision = pd.DataFrame()
        summary = pd.DataFrame()
    else:
        df = df.sort_values(["项目名称", "月份"]).reset_index(drop=True)
        decision = decide(df)
        summary = (
            decision.groupby(["月份", "处理建议"], as_index=False)
            .agg(文件数=("文件名", "count"), 项目数=("项目名称", "nunique"))
            .sort_values(["月份", "处理建议"])
        )

    analysis = extract_dir / ANALYSIS_DIR_NAME
    analysis.mkdir(parents=True, exist_ok=True)
    report = analysis / "采购台账_月度去重分析.xlsx"
    with pd.ExcelWriter(report, engine="openpyxl") as w:
        decision.to_excel(w, index=False, sheet_name="逐文件判定")
        summary.to_excel(w, index=False, sheet_name="月份汇总")

    n_copy, n_conflict = export_latest_per_project(decision, extract_dir)
    print(f"分析报告: {report}")
    print(f"判定行数: {len(decision)} | 拷贝: {n_copy} | 冲突: {n_conflict}")
    print(f"提取目录: {extract_dir}")


def main() -> None:
    run()


if __name__ == "__main__":
    main()
