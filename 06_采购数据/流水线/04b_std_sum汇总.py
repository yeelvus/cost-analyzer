#!/us,/bin/env python3
"""
步骤4（默认日常）：只输出「对比结构」标准化汇总。

数据源：D_raw_sum/采购部采购台账_raw_sum.xlsx（需先跑 4a）
输出唯一主文件：
  E_sum/采购部采购台账_std_sum.xlsx
  （含：原始名称/规格/单位 + 标准化后列 + 转换单位/计算数量
   + 项目汇总 + 采购总金额）

纯合并请手动运行：
  流水线/04a_raw_sum纯合并.py → D_raw_sum/*_raw_sum.xlsx
"""
from __future__ import annotations

import argparse
import importlib.util
import re
import shutil
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_PKG_ROOT = Path(__file__).resolve().parents[1]
_CODE_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(_PKG_ROOT), str(_CODE_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from common.excel import apply_common_excel_layout

from _采购路径 import (
    CONTENT_STD_UTILS,
    FX_RATE_FILE,
    RAW_SUM_DIR,
    RAW_SUM_FILE_NAME,
    STD_DIR,
    STD_SUM_FILE_NAME,
    SUM_DIR,
    UNIT_PIPELINE,
    is_excel,
    raw_sum_xlsx,
)

# 正式输出列（merged / std_merged 不变）
FINAL_OUTPUT_COLUMNS = [
    "项目编码", "项目名称", "材料归类", "物品名称", "规格", "单位", "数量",
    "单价", "不含税金额", "7%增值税", "含税金额", "货币种类", "汇率",
    "标准日期", "供货商名称", "备注", "年月", "签订日期",
]

# 对比版：序号 + 原始列 + 转换列 + 标准化后列
COMPARE_OUTPUT_COLUMNS = [
    "序号",
    "项目编码",
    "项目名称",
    "原始物品名称",
    "原始规格",
    "原始单位",
    "转换单位",
    "计算数量",
    "材料归类",
    "物品名称",
    "规格",
    "单位",
    "数量",
    "单价",
    "不含税金额",
    "7%增值税",
    "含税金额",
    "货币种类",
    "汇率",
    "标准日期",
    "供货商名称",
    "备注",
    "年月",
    "签订日期",
]

OUTPUT_STD_SUM = STD_SUM_FILE_NAME
UNIT_CONV_PIPELINE = UNIT_PIPELINE

AMOUNT_FIELD_NAMES = {"单价", "不含税金额", "7%增值税", "含税金额", "数量", "计算数量", "汇率"}
# Excel 错误字面量 / 非数值，写入或汇总时当作空
_EXCEL_ERROR_TOKENS = {
    "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!",
    "#GETTING_DATA", "#SPILL!", "#CALC!", "#ERROR!",
}

# 非泰铢且含税金额空、不含税有值时：按月查汇率换算，填入单元格用绿色高亮
_FX_GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_FX_GREEN_FONT = Font(color="006100")
_FX_FILL_COLS = ("汇率", "含税金额")  # 自动填入并标绿的列

_YM_RE = re.compile(r"((?:19|20)\d{2})[-./年](0?[1-9]|1[0-2])")

# 货币识别（归一到 THB / CNY / USD）
_CURRENCY_THB = {
    "thb", "tha", "฿", "baht", "泰铢", "泰幣", "泰币", "铢", "銖",
}
_CURRENCY_CNY = {
    "cny", "rmb", "cnh", "￥", "¥", "元", "人民币", "人民幣", "人名币",
}
_CURRENCY_USD = {
    "usd", "us$", "usa", "$", "美元", "美金", "美圆",
}

_STD_UTILS = None
_UNIT_CONV = None
_FX_RATE_CACHE: dict[str, dict[str, float]] | None = None


def _is_excel_error_token(value: object) -> bool:
    if value is None:
        return False
    s = str(value).strip().upper()
    if not s:
        return False
    if s in {t.upper() for t in _EXCEL_ERROR_TOKENS}:
        return True
    return s.startswith("#") and s.endswith("!")


def coerce_amount(value: object) -> float | None:
    """将金额/数量安全转为 float；#VALUE!/空/无法解析 → None。"""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if pd.isna(value):
            return None
        return float(value)
    if isinstance(value, str):
        s = value.strip().replace(",", "").replace("，", "")
        if not s or _is_excel_error_token(s):
            return None
        try:
            return float(s)
        except ValueError:
            return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


# ── 非泰铢汇率换算（不含税×月度汇率 → 含税金额/泰铢）────────────────

def normalize_currency(value: object) -> str | None:
    """
    归一货币种类：
      THB / CNY / USD / 其它大写代码；空 → None
    """
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    s = str(value).strip()
    if not s or _is_excel_error_token(s):
        return None
    # 去掉括号说明，如 "人民币(CNY)"
    s_clean = re.sub(r"[（(].*?[）)]", "", s).strip()
    key = s_clean.lower().replace(" ", "").replace("　", "")
    if key in _CURRENCY_THB or any(a in key for a in ("泰铢", "泰幣", "baht")):
        return "THB"
    if key in _CURRENCY_CNY or "人民币" in key or "人民幣" in key or key in ("元",):
        return "CNY"
    if key in _CURRENCY_USD or "美元" in key or "美金" in key:
        return "USD"
    # 常见字母代码
    letters = re.sub(r"[^a-zA-Z]", "", s_clean).upper()
    if letters in ("THB", "CNY", "RMB", "USD"):
        return "CNY" if letters == "RMB" else letters
    if letters:
        return letters
    return s_clean.upper() if s_clean else None


def _extract_year_month(*values: object) -> str | None:
    """从 年月 / 标准日期 / 签订日期 等提取 YYYY-MM。"""
    for value in values:
        if value is None:
            continue
        try:
            if pd.isna(value):
                continue
        except Exception:
            pass
        if hasattr(value, "year") and hasattr(value, "month"):
            try:
                return f"{int(value.year):04d}-{int(value.month):02d}"
            except Exception:
                pass
        s = str(value).strip()
        if not s:
            continue
        m = _YM_RE.search(s)
        if m:
            return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}"
    return None


def load_fx_rate_table(path: Path | None = None) -> dict[str, dict[str, float]]:
    """
    读取汇率三表.xlsx → { 'YYYY-MM': {'CNY': 人民币对泰铢, 'USD': 美元对泰铢}, ... }

    表头：年月 | 人民币对泰铢 | 美元对泰铢 | …
    """
    global _FX_RATE_CACHE
    rate_path = Path(path) if path else FX_RATE_FILE
    # 同路径才用缓存
    cache_key = str(rate_path.resolve()) if rate_path.exists() else str(rate_path)
    if _FX_RATE_CACHE is not None and getattr(load_fx_rate_table, "_path", None) == cache_key:
        return _FX_RATE_CACHE

    table: dict[str, dict[str, float]] = {}
    if not rate_path.is_file():
        print(f"⚠️ 汇率表不存在: {rate_path}")
        _FX_RATE_CACHE = table
        load_fx_rate_table._path = cache_key  # type: ignore[attr-defined]
        return table

    try:
        df = pd.read_excel(rate_path, sheet_name=0)
    except Exception as exc:
        print(f"⚠️ 读取汇率表失败: {exc}")
        _FX_RATE_CACHE = table
        load_fx_rate_table._path = cache_key  # type: ignore[attr-defined]
        return table

    if df.empty:
        _FX_RATE_CACHE = table
        load_fx_rate_table._path = cache_key  # type: ignore[attr-defined]
        return table

    # 列名可能含换行
    col_map: dict[str, str] = {}
    for c in df.columns:
        cs = str(c).replace("\n", "").replace(" ", "").strip()
        if "年月" in cs or cs.lower() in ("date", "month", "ym"):
            col_map["ym"] = c
        elif "人民币" in cs and "泰铢" in cs:
            col_map["CNY"] = c
        elif "美元" in cs and "泰铢" in cs:
            col_map["USD"] = c

    if "ym" not in col_map:
        # 回退首列
        col_map["ym"] = df.columns[0]
    if "CNY" not in col_map and len(df.columns) >= 2:
        col_map["CNY"] = df.columns[1]
    if "USD" not in col_map and len(df.columns) >= 3:
        col_map["USD"] = df.columns[2]

    for _, row in df.iterrows():
        ym = _extract_year_month(row.get(col_map["ym"]))
        if not ym:
            continue
        entry: dict[str, float] = {}
        for code in ("CNY", "USD"):
            if code not in col_map:
                continue
            rate = coerce_amount(row.get(col_map[code]))
            if rate is not None and rate > 0:
                entry[code] = float(rate)
        if entry:
            table[ym] = entry

    _FX_RATE_CACHE = table
    load_fx_rate_table._path = cache_key  # type: ignore[attr-defined]
    print(f"  · 汇率表: {rate_path.name} 共 {len(table)} 个月份")
    return table


def apply_non_thb_fx_conversion(
    df: pd.DataFrame,
    *,
    rate_path: Path | None = None,
) -> tuple[pd.DataFrame, set[int]]:
    """
    非泰铢 + 含税金额为空 + 不含税金额有值：
      1) 按「年月」查 人民币对泰铢 / 美元对泰铢
      2) 填入「汇率」
      3) 含税金额 = 不含税金额 × 汇率（折合泰铢）
    返回 (修改后的 df, 发生换算的行索引集合)。
    """
    filled: set[int] = set()
    if df is None or df.empty:
        return df, filled

    work = df.copy()
    for col in ("货币种类", "不含税金额", "含税金额", "汇率", "年月", "标准日期", "签订日期"):
        if col not in work.columns:
            work[col] = ""

    rates = load_fx_rate_table(rate_path)
    if not rates:
        print("  ⚠️ 汇率表为空，跳过非泰铢换算")
        return work, filled

    stats = defaultdict(int)
    miss_ym: set[str] = set()
    miss_ccy: set[str] = set()

    for idx, row in work.iterrows():
        ccy = normalize_currency(row.get("货币种类"))
        if ccy is None or ccy == "THB":
            continue

        tax_incl = coerce_amount(row.get("含税金额"))
        if tax_incl is not None:
            continue  # 已有含税金额，不覆盖

        tax_excl = coerce_amount(row.get("不含税金额"))
        if tax_excl is None:
            continue  # 不含税也空，无法换算

        if ccy not in ("CNY", "USD"):
            miss_ccy.add(ccy)
            stats["skip_ccy"] += 1
            continue

        ym = _extract_year_month(
            row.get("年月"), row.get("标准日期"), row.get("签订日期")
        )
        if not ym:
            stats["skip_no_ym"] += 1
            continue

        month_rates = rates.get(ym)
        if not month_rates or ccy not in month_rates:
            miss_ym.add(f"{ym}/{ccy}")
            stats["skip_no_rate"] += 1
            continue

        rate = float(month_rates[ccy])
        converted = round(tax_excl * rate, 2)
        work.at[idx, "汇率"] = rate
        work.at[idx, "含税金额"] = converted
        filled.add(idx)
        stats[ccy] += 1

    if filled:
        parts = [f"{k}×{v}" for k, v in stats.items() if k in ("CNY", "USD") and v]
        print(
            f"  · 非泰铢换算: 成功 {len(filled)} 行"
            + (f"（{', '.join(parts)}）" if parts else "")
            + f"  含税金额=不含税×月度汇率 → 标绿"
        )
    skips = []
    if stats.get("skip_no_ym"):
        skips.append(f"无年月 {stats['skip_no_ym']}")
    if stats.get("skip_no_rate"):
        skips.append(f"无汇率 {stats['skip_no_rate']}")
    if stats.get("skip_ccy"):
        skips.append(f"不支持币种 {stats['skip_ccy']}")
    if skips:
        print(f"  · 非泰铢换算跳过: {'; '.join(skips)}")
    if miss_ym:
        sample = ", ".join(sorted(miss_ym)[:8])
        more = f" 等{len(miss_ym)}项" if len(miss_ym) > 8 else ""
        print(f"    缺月份汇率: {sample}{more}")
    if miss_ccy:
        print(f"    未支持币种: {', '.join(sorted(miss_ccy))}")

    return work, filled


def _apply_fx_green_fill(
    ws,
    columns: list[str],
    excel_rows: list[int],
    fill_col_names: tuple[str, ...] = _FX_FILL_COLS,
) -> int:
    """对指定 Excel 行号的 汇率/含税金额 单元格标绿。返回着色单元格数。"""
    col_idx = {name: i + 1 for i, name in enumerate(columns)}
    targets = [col_idx[c] for c in fill_col_names if c in col_idx]
    if not targets or not excel_rows:
        return 0
    n = 0
    for r in excel_rows:
        for c in targets:
            cell = ws.cell(r, c)
            cell.fill = _FX_GREEN_FILL
            cell.font = _FX_GREEN_FONT
            if c == col_idx.get("含税金额") or c == col_idx.get("汇率"):
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00" if c == col_idx.get("含税金额") else "0.0000"
            n += 1
    return n


def _load_procurement_standardization_utils():
    """加载采购内容标准化工具（采购内容标准化_utils.py）。"""
    global _STD_UTILS
    if _STD_UTILS is not None:
        return _STD_UTILS if _STD_UTILS is not False else None

    if not CONTENT_STD_UTILS.exists():
        print(f"⚠️ 未找到内容标准化工具: {CONTENT_STD_UTILS}")
        _STD_UTILS = False
        return None

    spec = importlib.util.spec_from_file_location(
        "procurement_content_std_utils", str(CONTENT_STD_UTILS)
    )
    if spec is None or spec.loader is None:
        print(f"⚠️ 无法加载内容标准化工具: {CONTENT_STD_UTILS}")
        _STD_UTILS = False
        return None

    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    if not hasattr(mod, "apply_standardization_pipeline"):
        print("⚠️ 内容标准化工具缺少 apply_standardization_pipeline")
        _STD_UTILS = False
        return None

    _STD_UTILS = mod
    return mod


def _load_unit_conversion_pipeline():
    """加载单位数量转换流水线（钢筋网片 m2→kg、砖块→m3 等）。"""
    global _UNIT_CONV
    if _UNIT_CONV is not None:
        return _UNIT_CONV if _UNIT_CONV is not False else None

    if not UNIT_CONV_PIPELINE.exists():
        print(f"⚠️ 未找到单位转换流水线: {UNIT_CONV_PIPELINE}")
        _UNIT_CONV = False
        return None

    spec = importlib.util.spec_from_file_location(
        "unit_qty_conversion_pipeline", str(UNIT_CONV_PIPELINE)
    )
    if spec is None or spec.loader is None:
        _UNIT_CONV = False
        return None
    mod = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(mod)
    except Exception as exc:
        print(f"⚠️ 单位转换流水线加载失败: {exc}")
        _UNIT_CONV = False
        return None
    _UNIT_CONV = mod
    return mod


# 正文左对齐的文本列（表头仍居中）
_LEFT_ALIGN_HEADER_KEYS = (
    "物品名称",
    "规格",
    "原始物品名称",
    "原始规格",
    "项目名称",
    "材料归类",
    "供货商名称",
    "备注",
    "供应商",
)


def _apply_left_align_text_columns(ws, header_row: int = 1) -> None:
    """物品名称、规格等文本列正文左对齐。"""
    if not ws.max_row or ws.max_row < header_row:
        return
    headers = [cell.value for cell in ws[header_row]]
    left_cols = [
        idx + 1
        for idx, h in enumerate(headers)
        if isinstance(h, str) and any(k in h for k in _LEFT_ALIGN_HEADER_KEYS)
    ]
    if not left_cols:
        return
    left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    for col_idx in left_cols:
        for row_idx in range(header_row + 1, (ws.max_row or header_row) + 1):
            ws.cell(row=row_idx, column=col_idx).alignment = left


def apply_unified_workbook_format(output_wb) -> None:
    amount_keywords = ("金额", "单价", "税")
    for sheet_name in (
        "合并数据",
        "项目金额汇总",
        "项目汇总",
        "采购总金额",
    ):
        if sheet_name not in output_wb.sheetnames:
            continue
        ws = output_wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        amount_headers = [
            h for h in headers if isinstance(h, str) and any(k in h for k in amount_keywords)
        ]
        if sheet_name == "采购总金额":
            apply_common_excel_layout(
                ws,
                numeric_start_col=3,
                amount_headers=[h for h in headers[2:] if isinstance(h, str)],
                header_row=1,
                header_height=40,
                freeze_panes="A2",
            )
        else:
            apply_common_excel_layout(
                ws,
                numeric_start_col=ws.max_column + 1,
                amount_headers=amount_headers,
                header_row=1,
                header_height=40,
                freeze_panes="A2",
            )
        _apply_left_align_text_columns(ws, header_row=1)


def _period_key(period_text: str) -> tuple[int, int]:
    y, m = str(period_text).split("-")
    return int(y), int(m)


def build_procurement_total_sheet(merged_df: pd.DataFrame) -> pd.DataFrame:
    required = {"项目编码", "项目名称", "年月", "含税金额"}
    if not required.issubset(set(merged_df.columns)):
        return pd.DataFrame(columns=["项目编码", "项目名称"])

    work = merged_df[["项目编码", "项目名称", "年月", "含税金额"]].copy()
    work["项目编码"] = work["项目编码"].fillna("").astype(str).str.strip()
    work["项目名称"] = work["项目名称"].fillna("").astype(str).str.strip()
    work["年月"] = work["年月"].fillna("").astype(str).str.strip()
    work["含税金额"] = work["含税金额"].map(coerce_amount)
    work = work[(work["项目编码"] != "") | (work["项目名称"] != "")]
    work = work[work["年月"].str.fullmatch(r"(?:19|20)\d{2}-(?:0[1-9]|1[0-2])", na=False)]
    work = work[work["含税金额"].notna()]
    if work.empty:
        return pd.DataFrame(columns=["项目编码", "项目名称"])

    grouped = work.groupby(["项目编码", "项目名称", "年月"], as_index=False).agg(
        含税金额=("含税金额", "sum")
    )
    periods = sorted(grouped["年月"].unique().tolist(), key=_period_key)
    label_map = {p: f"{int(p.split('-')[0])}年{int(p.split('-')[1])}月" for p in periods}
    pivot = (
        grouped.pivot_table(
            index=["项目编码", "项目名称"],
            columns="年月",
            values="含税金额",
            aggfunc="sum",
            fill_value=0.0,
        )
        .reindex(columns=periods, fill_value=0.0)
    )
    amount_wide = pivot.rename(columns=label_map).reset_index()
    sort_cols = [label_map[p] for p in reversed(periods)]
    amount_wide["__编码为空"] = amount_wide["项目编码"].astype(str).str.strip().eq("")
    amount_wide = amount_wide.sort_values(
        ["__编码为空", *sort_cols],
        ascending=[True] + [False] * len(sort_cols),
    ).drop(columns=["__编码为空"]).reset_index(drop=True)
    return amount_wide


def list_std_excels(folder: Path) -> list[Path]:
    files = []
    for p in folder.rglob("*"):
        if not is_excel(p):
            continue
        name = p.name
        if name.startswith("merged") or name.endswith("_std_merged.xlsx"):
            continue
        if "对比" in name:
            continue
        if name.startswith("采购部采购台账_"):
            continue
        files.append(p)
    return sorted(files)


def _new_merge_workbook(columns: list[str]) -> tuple[openpyxl.Workbook, openpyxl.worksheet.worksheet.Worksheet]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "合并数据"
    ws.append(columns)
    return wb, ws


def _append_data_row(ws, row_idx: int, columns: list[str], row_data: dict) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        value = row_data.get(col_name, "")
        try:
            if pd.isna(value):
                value = ""
        except Exception:
            pass

        # 金额/数量列：清洗 #VALUE! 等，只写真正的数字或空
        if col_name in AMOUNT_FIELD_NAMES or col_name.endswith("金额"):
            num = coerce_amount(value)
            cell = ws.cell(row_idx, col_idx, num if num is not None else None)
            if num is not None and col_name in {
                "单价", "不含税金额", "7%增值税", "含税金额", "计算数量",
            }:
                cell.number_format = "#,##0.00"
            continue

        if _is_excel_error_token(value):
            value = ""
        cell = ws.cell(row_idx, col_idx, value)


def _finalize_workbook(
    output_wb: openpyxl.Workbook,
    columns: list[str],
    project_totals: dict[tuple[str, str], float],
    data_last_row: int,
) -> openpyxl.Workbook | None:
    if data_last_row < 2:
        return None

    merged_ws = output_wb["合并数据"]
    merged_ws.row_dimensions[1].height = 40
    for cell in merged_ws[1]:
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    merged_ws.freeze_panes = "A2"
    last_col = get_column_letter(len(columns))
    merged_ws.auto_filter.ref = f"A1:{last_col}{data_last_row}"

    # 简单金额合计（与「分类项目汇总」区分；分类宽表 sheet 名仍为「项目汇总」）
    summary_ws = output_wb.create_sheet("项目金额汇总")
    summary_ws.append(["项目编码", "项目名称", "总采购金额"])
    grand_total = 0.0
    for row_idx, ((code, name), total) in enumerate(sorted(project_totals.items()), start=2):
        amt = float(total) if total is not None else 0.0
        summary_ws.append([code, name, amt])
        summary_ws.cell(row_idx, 3).number_format = "#,##0.00"
        grand_total += amt
    last = summary_ws.max_row + 1
    summary_ws.cell(last, 1, "")
    summary_ws.cell(last, 2, "【合计】")
    summary_ws.cell(last, 3, round(grand_total, 2)).number_format = "#,##0.00"

    merged_df = pd.DataFrame(merged_ws.values)
    merged_df.columns = merged_df.iloc[0]
    merged_df = merged_df.iloc[1:].reset_index(drop=True)
    # 金额列再清洗一遍，供宽表聚合
    if "含税金额" in merged_df.columns:
        merged_df["含税金额"] = merged_df["含税金额"].map(coerce_amount)

    total_df = build_procurement_total_sheet(merged_df)
    total_ws = output_wb.create_sheet("采购总金额")
    total_ws.append(total_df.columns.tolist())
    for row in total_df.itertuples(index=False, name=None):
        cleaned = []
        for v in row:
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cleaned.append(float(v))
            else:
                # 非金额文本列（项目编码/名称）原样
                cleaned.append("" if _is_excel_error_token(v) else v)
        total_ws.append(cleaned)

    # 只保留业务 sheet（分类「项目汇总」由后续注入）
    keep = {"合并数据", "项目金额汇总", "采购总金额", "项目汇总"}
    for sn in list(output_wb.sheetnames):
        if sn not in keep:
            del output_wb[sn]

    return output_wb


def merge_standardized_excels(
    folder: Path,
    *,
    use_content_standardization: bool = False,
    include_compare: bool = False,
) -> tuple[openpyxl.Workbook | None, openpyxl.Workbook | None]:
    """
    从 folder（B_std）合并正式汇总工作簿。

    注意：对比版不再由此函数从 B_std 生成，请用
    build_compare_workbook_from_raw_sum()（数据源 D_raw_sum）。
    include_compare 参数保留兼容，恒忽略。
    """
    del include_compare  # 对比版改由 raw_sum 生成
    excel_paths = list_std_excels(folder)
    std_utils = None

    if use_content_standardization:
        std_utils = _load_procurement_standardization_utils()
        if std_utils is None:
            print("⚠️ 内容标准化不可用，本轮回退为纯汇总")
            use_content_standardization = False
        else:
            try:
                unit_ref = std_utils.build_global_unit_ref(folder)
                std_utils.set_global_unit_ref(unit_ref)
                print(f"全局单位参考库: {len(unit_ref)} 个物品")
            except Exception as exc:
                print(f"⚠️ 全局单位库构建失败（将仅用文件内补全）: {exc}")

    mode = "启用采购内容标准化" if use_content_standardization else "纯汇总"
    print("\n" + "=" * 72)
    print(f"【{mode}】目录: {folder}")
    print(f"【{mode}】文件数: {len(excel_paths)}")
    print("=" * 72 + "\n")

    if not excel_paths:
        print("⚠️ 未找到任何 Excel 文件")
        return None, None

    formal_wb, formal_ws = _new_merge_workbook(FINAL_OUTPUT_COLUMNS)
    formal_row = 2
    project_totals: dict[tuple[str, str], float] = defaultdict(float)

    for idx, file_path in enumerate(excel_paths, start=1):
        print(f"[{idx}/{len(excel_paths)}] {file_path.name}")
        read_path = file_path
        temp_dir = None

        if use_content_standardization and std_utils is not None:
            try:
                temp_dir = tempfile.mkdtemp(prefix="procurement_std_")
                temp_file = Path(temp_dir) / file_path.name
                shutil.copy2(file_path, temp_file)
                std_utils.apply_standardization_pipeline(
                    str(temp_file),
                    auto_prompt_units=False,
                    resolve_pending=False,
                )
                read_path = temp_file
            except Exception as exc:
                print(f"  ⚠️ 标准化失败，用原文件: {exc}")
                read_path = file_path

        try:
            df = pd.read_excel(read_path)
        except Exception as exc:
            print(f"  ✗ 读取失败: {exc}")
            continue
        finally:
            if temp_dir and Path(temp_dir).is_dir():
                shutil.rmtree(temp_dir, ignore_errors=True)

        if df.empty:
            continue

        for col in FINAL_OUTPUT_COLUMNS:
            if col not in df.columns:
                df[col] = ""

        for col in list(df.columns):
            if col in AMOUNT_FIELD_NAMES or str(col).endswith("金额"):
                df[col] = df[col].map(coerce_amount)

        # 纯合并（04a）不改金额；非泰铢换算仅在 04b compare 路径做
        for _, row in df.iterrows():
            formal_data = {c: row[c] for c in FINAL_OUTPUT_COLUMNS}
            _append_data_row(formal_ws, formal_row, FINAL_OUTPUT_COLUMNS, formal_data)

            code = "" if pd.isna(row["项目编码"]) else str(row["项目编码"]).strip()
            name = "" if pd.isna(row["项目名称"]) else str(row["项目名称"]).strip()
            amt = coerce_amount(row["含税金额"])
            if amt is not None:
                project_totals[(code, name)] += float(amt)
            formal_row += 1

    formal_out = _finalize_workbook(formal_wb, FINAL_OUTPUT_COLUMNS, project_totals, formal_row - 1)
    if formal_out is None:
        print("⚠️ 没有有效数据")
    return formal_out, None


def build_compare_workbook_from_raw_sum(
    raw_sum_path: Path | None = None,
) -> openpyxl.Workbook | None:
    """
    对比版数据源：D_采购部采购台账_raw_sum/采购部采购台账_raw_sum.xlsx

    流程：
      1) 读取 raw_sum「合并数据」→ 作为原始列
      2) 对副本做内容标准化 → 标准化后列
      3) 单位转换 → 转换单位/计算数量
      4) 写出对比版结构
    """
    raw_sum_path = Path(raw_sum_path) if raw_sum_path else raw_sum_xlsx()
    if not raw_sum_path.is_file():
        print(
            f"⚠️ 对比版跳过：未找到 raw_sum 源文件\n"
            f"   期望: {raw_sum_path}\n"
            f"   请先运行: 流水线/04a_raw_sum纯合并.py"
        )
        return None

    print("\n" + "=" * 72)
    print(f"【对比版】数据源 D_raw_sum: {raw_sum_path}")
    print("=" * 72 + "\n")

    try:
        try:
            orig_df = pd.read_excel(raw_sum_path, sheet_name="合并数据")
        except Exception:
            orig_df = pd.read_excel(raw_sum_path, sheet_name=0)
    except Exception as exc:
        print(f"⚠️ 读取 raw_sum 失败: {exc}")
        return None

    if orig_df.empty:
        print("⚠️ raw_sum 无数据，跳过对比版")
        return None

    for col in FINAL_OUTPUT_COLUMNS:
        if col not in orig_df.columns:
            orig_df[col] = ""

    # 原始三列（标准化前）
    for c in ("物品名称", "规格", "单位"):
        if c not in orig_df.columns:
            orig_df[c] = ""
    orig_snapshot = orig_df[["物品名称", "规格", "单位"]].copy()
    orig_snapshot.columns = ["原始物品名称", "原始规格", "原始单位"]
    orig_snapshot = orig_snapshot.reset_index(drop=True)

    std_utils = _load_procurement_standardization_utils()
    unit_conv = _load_unit_conversion_pipeline()
    if unit_conv is None:
        print("⚠️ 单位转换不可用：对比版「转换单位/计算数量」将为空")

    temp_dir = tempfile.mkdtemp(prefix="procurement_compare_rawsum_")
    try:
        temp_file = Path(temp_dir) / "from_raw_sum.xlsx"
        # 写一张可被规格流水线处理的表
        with pd.ExcelWriter(temp_file, engine="openpyxl") as writer:
            orig_df.to_excel(writer, sheet_name="合并数据", index=False)

        if std_utils is not None:
            try:
                # 单位参考：用 raw_sum 所在目录扫不到多文件时，用文件内逻辑
                try:
                    unit_ref = std_utils.build_global_unit_ref(raw_sum_path.parent)
                    std_utils.set_global_unit_ref(unit_ref)
                except Exception:
                    pass
                std_utils.apply_standardization_pipeline(
                    str(temp_file),
                    auto_prompt_units=False,
                    resolve_pending=False,
                )
            except Exception as exc:
                print(f"⚠️ 对 raw_sum 做内容标准化失败，对比版标准化列将接近原始: {exc}")

        try:
            std_df = pd.read_excel(temp_file, sheet_name="合并数据")
        except Exception:
            std_df = pd.read_excel(temp_file, sheet_name=0)
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

    for col in FINAL_OUTPUT_COLUMNS:
        if col not in std_df.columns:
            std_df[col] = ""

    for col in list(std_df.columns):
        if col in AMOUNT_FIELD_NAMES or str(col).endswith("金额"):
            std_df[col] = std_df[col].map(coerce_amount)

    std_df = std_df.reset_index(drop=True)
    # 行数对齐
    n = min(len(orig_snapshot), len(std_df))
    if len(orig_snapshot) != len(std_df):
        print(
            f"  ⚠️ raw 与标准化后行数不一致: raw={len(orig_snapshot)} std={len(std_df)}，按 {n} 行对齐"
        )
    orig_snapshot = orig_snapshot.iloc[:n].reset_index(drop=True)
    std_df = std_df.iloc[:n].reset_index(drop=True)

    # 非泰铢：含税空 + 不含税有值 → 按月汇率折算泰铢并标绿
    std_df, fx_filled_idx = apply_non_thb_fx_conversion(std_df)

    conv_units: list = []
    conv_qtys: list = []
    for _, r in std_df.iterrows():
        if unit_conv is not None and hasattr(unit_conv, "convert_row"):
            cr = unit_conv.convert_row(
                item_name=r.get("物品名称", ""),
                spec=r.get("规格", ""),
                unit=r.get("单位", ""),
                qty=r.get("数量", ""),
            )
            conv_units.append(cr.get("转换单位", ""))
            conv_qtys.append(cr.get("计算数量", ""))
        else:
            conv_units.append("")
            conv_qtys.append("")
    std_df["转换单位"] = conv_units
    std_df["计算数量"] = conv_qtys

    compare_df = pd.concat([orig_snapshot, std_df], axis=1)

    compare_wb, compare_ws = _new_merge_workbook(COMPARE_OUTPUT_COLUMNS)
    compare_row = 2
    compare_totals: dict[tuple[str, str], float] = defaultdict(float)
    seq = 1
    fx_excel_rows: list[int] = []

    for df_idx, row in compare_df.iterrows():
        compare_data = {
            c: row[c] if c in row.index else "" for c in COMPARE_OUTPUT_COLUMNS
        }
        compare_data["序号"] = seq
        seq += 1
        _append_data_row(compare_ws, compare_row, COMPARE_OUTPUT_COLUMNS, compare_data)
        if df_idx in fx_filled_idx:
            fx_excel_rows.append(compare_row)
        code = "" if pd.isna(row.get("项目编码", "")) else str(row.get("项目编码", "")).strip()
        name = "" if pd.isna(row.get("项目名称", "")) else str(row.get("项目名称", "")).strip()
        amt = coerce_amount(row.get("含税金额"))
        if amt is not None:
            compare_totals[(code, name)] += float(amt)
        compare_row += 1

    # 汇率换算填入的 汇率 / 含税金额 → 绿色
    if fx_excel_rows:
        n_cells = _apply_fx_green_fill(compare_ws, COMPARE_OUTPUT_COLUMNS, fx_excel_rows)
        print(f"  · 绿色高亮: {len(fx_excel_rows)} 行 / {n_cells} 格（汇率、含税金额）")

    # 序号列居中
    if "序号" in COMPARE_OUTPUT_COLUMNS:
        c_idx = COMPARE_OUTPUT_COLUMNS.index("序号") + 1
        for r in range(2, compare_row):
            compare_ws.cell(r, c_idx).alignment = Alignment(
                horizontal="center", vertical="center"
            )

    return _finalize_workbook(
        compare_wb, COMPARE_OUTPUT_COLUMNS, compare_totals, compare_row - 1
    )


def _load_step5_category_module():
    """加载 5_项目采购分类汇总（build_summary / 样式）。"""
    path = Path(__file__).resolve().parent / "05_分类汇总刷新.py"
    spec = importlib.util.spec_from_file_location("proc_step5_cat", str(path))
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载: {path}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def inject_category_summary_into_workbook(wb: openpyxl.Workbook) -> bool:
    """
    把分类「项目汇总」宽表写入 std_sum（不再生成 分类合计/分类明细）。
    数据取自本簿「合并数据」的标准化列。
    """
    if "合并数据" not in wb.sheetnames:
        print("  ⚠️ 无「合并数据」，跳过分类项目汇总")
        return False

    ws_m = wb["合并数据"]
    rows = list(ws_m.iter_rows(values_only=True))
    if len(rows) < 2:
        print("  ⚠️ 合并数据为空，跳过分类项目汇总")
        return False
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    detail_src = pd.DataFrame(rows[1:], columns=headers)

    try:
        cat = _load_step5_category_module()
    except Exception as exc:
        print(f"  ⚠️ 加载分类汇总模块失败: {exc}")
        return False

    unit_conv = None
    try:
        unit_conv = cat._load_unit_conv()
    except Exception:
        unit_conv = None

    try:
        wide, _class_sum, _detail = cat.build_summary(detail_src, unit_conv)
    except Exception as exc:
        print(f"  ⚠️ 分类汇总计算失败: {exc}")
        return False

    if wide is None or wide.empty:
        print("  ⚠️ 分类项目汇总无有效数据")
        return False

    for sn in ("项目汇总", "分类合计", "分类明细"):
        if sn in wb.sheetnames:
            del wb[sn]

    qty_headers = {c[1] for c in cat.CLASS_SPECS} | {
        "明细行数",
        "汇总数量",
        "数量",
        "行数",
    }
    amt_headers = {c[2] for c in cat.CLASS_SPECS} | {
        "合计金额",
        "分类金额合计",
        "金额差异",
        "含税金额",
    }

    # 项目汇总：首列加序号
    wide = wide.reset_index(drop=True)
    wide.insert(0, "序号", range(1, len(wide) + 1))

    ws1 = wb.create_sheet("项目汇总")
    ws1.append(list(wide.columns))
    for row in wide.itertuples(index=False, name=None):
        ws1.append(list(row))
    if len(wide) > 0:
        total_row = []
        for col in wide.columns:
            if col == "序号":
                total_row.append("")
            elif col == "项目编码":
                total_row.append("")
            elif col == "项目名称":
                total_row.append("【全部项目合计】")
            elif col == "明细行数":
                total_row.append(int(wide[col].sum()))
            elif col in qty_headers:
                total_row.append(round(float(wide[col].sum()), 0))
            else:
                try:
                    total_row.append(float(wide[col].sum()))
                except Exception:
                    total_row.append("")
        ws1.append(total_row)
        from openpyxl.styles import Font as _Font

        for cell in ws1[ws1.max_row]:
            cell.font = _Font(bold=True)
    cat._style_sheet(ws1, qty_headers=qty_headers, amount_headers=amt_headers)
    # 序号居中
    if "序号" in list(wide.columns):
        c_idx = list(wide.columns).index("序号") + 1
        for r in range(2, (ws1.max_row or 1) + 1):
            cell = ws1.cell(r, c_idx)
            if cell.value not in ("", None):
                cell.alignment = Alignment(horizontal="center", vertical="center")

    preferred = [
        "合并数据",
        "项目汇总",
        "项目金额汇总",
        "采购总金额",
    ]
    for i, name in enumerate(preferred):
        if name in wb.sheetnames:
            wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

    print("  ✓ 已写入分类「项目汇总」（不含分类合计/分类明细）")
    return True


def main() -> None:
    parser = argparse.ArgumentParser(
        description="D_raw_sum → E_sum/采购部采购台账_std_sum.xlsx（含分类项目汇总，单文件）"
    )
    parser.add_argument(
        "--output-dir",
        default="",
        help="覆盖输出目录，默认 SUM_DIR（E_sum）",
    )
    parser.add_argument(
        "--raw-sum",
        default="",
        help=f"数据源 raw_sum 文件，默认 {RAW_SUM_DIR / RAW_SUM_FILE_NAME}",
    )
    parser.add_argument(
        "--output-name",
        default=OUTPUT_STD_SUM,
        help=f"输出文件名，默认 {OUTPUT_STD_SUM}",
    )
    args = parser.parse_args()

    out_dir = Path(args.output_dir).expanduser() if args.output_dir else SUM_DIR
    raw_sum_path = (
        Path(args.raw_sum).expanduser()
        if args.raw_sum
        else raw_sum_xlsx()
    )
    out_dir.mkdir(parents=True, exist_ok=True)

    print("标准化汇总 std_sum（对比结构 + 分类项目汇总，单文件）")
    print(f"- 数据源 D_raw_sum: {raw_sum_path}")
    print(f"- 输出 E_sum:       {out_dir / args.output_name}")
    print(f"- 内容标准化:       {CONTENT_STD_UTILS}")
    print(f"- 汇率表:           {FX_RATE_FILE}")
    print("  （非泰铢 + 含税空 + 不含税有值 → 月度汇率换算，含税/汇率标绿）")
    print("- 若无 raw_sum，请先: 流水线/04a_raw_sum纯合并.py")
    print("- 不再输出: 项目采购分类汇总.xlsx（已并入本文件「项目汇总」等 sheet）")

    wb = build_compare_workbook_from_raw_sum(raw_sum_path)
    if wb is None:
        print(f"未产生有效数据: {args.output_name}")
        return

    # 分类「项目汇总」等并入同一工作簿
    inject_category_summary_into_workbook(wb)

    out_path = out_dir / args.output_name
    apply_unified_workbook_format(wb)
    wb.save(out_path)
    print(f"已输出: {out_path}")
    print(f"  sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
