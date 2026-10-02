# -*- coding: utf-8 -*-
"""
项目名称匹配（全仓库共用，位于 common/）

【推荐用法 — 三行搞定】
    from common.project import ProjectMatcher

    m = ProjectMatcher()                    # 自动加载默认 JSON + 可选主数据表
    hit = m.match("永强")                   # → dict 或 None
    name, code = m.resolve("永强")          # → ("永强车业项目", "THCC-2025-005")
    m.save()

【图形界面】
    python -m common.project_name_matcher
    python -m common.project_name_matcher gui
    python -m common.project_name_matcher excel 台账.xlsx --process-summary

【批量 + 弹窗补录】
    mapped = m.match_batch(names, prompt=True)

【旧函数式 API 全部保留】（load_or_init_db / match_project / save_db …）
兼容已有脚本，无需立刻改写。

默认路径：
  01成本统计/config/common/project_name_match_db.json
  邯郸中材/…/0_项目基本信息情况/项目基本信息表.xlsx（若存在则合并标准名）
"""

from __future__ import annotations

import json
import os
import re
import sys
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple, Union

import pandas as pd

try:
    import tkinter as tk
    from tkinter import ttk
except Exception:  # pragma: no cover
    tk = None
    ttk = None

# ---------------------------------------------------------------------------
# 路径
# ---------------------------------------------------------------------------
MODULE_DIR = Path(__file__).resolve().parent
_CODE_ROOT = MODULE_DIR.parent
if str(_CODE_ROOT) not in sys.path:
    sys.path.insert(0, str(_CODE_ROOT))
from settings import MATCH_DB_PATH  # noqa: E402

DEFAULT_DB_PATH = Path(MATCH_DB_PATH)

# 项目基本信息表候选（按优先级）
DEFAULT_INFO_EXCEL_CANDIDATES: Tuple[Path, ...] = (
    Path(
        "/Users/wyattsun/Desktop/邯郸中材/01成本统计/0_项目基本信息情况/项目基本信息表.xlsx"
    ),
    MODULE_DIR.parent.parent
    / "邯郸中材"
    / "01成本统计"
    / "0_项目基本信息情况"
    / "项目基本信息表.xlsx",
)

# 运行期缓存
_ALIAS_CACHE: dict[int, dict] = {}
_FUZZY_CACHE: dict[tuple[int, str], dict | None] = {}

PathLike = Union[str, Path]


def ensure_on_sys_path() -> Path:
    """把 code 根目录加入 sys.path，便于 `from common.project import ...`。"""
    p = str(_CODE_ROOT)
    if p not in sys.path:
        sys.path.insert(0, p)
    return _CODE_ROOT


def find_info_excel(explicit: PathLike | None = None) -> Optional[Path]:
    """定位「项目基本信息表.xlsx」。"""
    if explicit:
        p = Path(explicit).expanduser()
        return p if p.is_file() else None
    for c in DEFAULT_INFO_EXCEL_CANDIDATES:
        if c.is_file():
            return c
    return None


def default_db_path() -> Path:
    return DEFAULT_DB_PATH


# ---------------------------------------------------------------------------
# 文本工具
# ---------------------------------------------------------------------------
def _clean_name(text) -> str:
    if text is None:
        return ""
    s = str(text).strip()
    if not s:
        return ""
    s = re.sub(r"（.*?）|\(.*?\)", "", s)
    s = s.replace("泰国", "").replace("项目", "").replace("工程", "")
    s = re.sub(r"\s+", "", s)
    return s.strip().lower()


def _normalize_key(text) -> str:
    if text is None:
        return ""
    return re.sub(r"\s+", "", str(text)).strip().lower()


def _clean_field(value, default: str = "") -> str:
    if value is None:
        return default
    s = str(value).strip()
    if not s or s.lower() in {"nan", "none", "nat"}:
        return default
    return s


# 摘要前缀清理：去掉与项目名称相同的开头，以及随后的连接符
_SUMMARY_DASH_PREFIX_RE = re.compile(r"^[\s\-－—–﹣−]+")


def strip_project_prefix_from_summary(project_name, summary):
    """
    若摘要以项目名称开头，则删去该前缀；删完后若以「-」等连接符开头，一并去掉。

    摘要本身不含此前缀时原样返回。summary 为 None 时返回 None。
    """
    if summary is None:
        return None
    text = str(summary)
    name = "" if project_name is None else str(project_name).strip()
    if name and len(text) >= len(name) and text[: len(name)].casefold() == name.casefold():
        text = _SUMMARY_DASH_PREFIX_RE.sub("", text[len(name) :])
    return text


def _invalidate_db_caches(db=None) -> None:
    global _ALIAS_CACHE, _FUZZY_CACHE
    if db is None:
        _ALIAS_CACHE.clear()
        _FUZZY_CACHE.clear()
        return
    _ALIAS_CACHE.pop(id(db), None)
    kid = id(db)
    for k in [k for k in _FUZZY_CACHE if k[0] == kid]:
        _FUZZY_CACHE.pop(k, None)


def _cached_alias_map(db) -> dict:
    key = id(db)
    m = _ALIAS_CACHE.get(key)
    if m is None:
        m = _build_alias_map(db)
        _ALIAS_CACHE[key] = m
    return m


def _build_alias_map(db) -> dict:
    """别名 → 标准名；显式 aliases 优先，再标准名自身。"""
    alias_map: dict = {}
    projects = db.get("projects", {}) or {}
    for std_name, meta in projects.items():
        std_name = str(std_name).strip()
        if not std_name:
            continue
        for alias in meta.get("aliases", []) or []:
            alias = str(alias).strip()
            if alias:
                alias_map.setdefault(alias, std_name)
    for std_name in projects.keys():
        std_name = str(std_name).strip()
        if std_name:
            alias_map.setdefault(std_name, std_name)
    return alias_map


def _pick_status_column(columns) -> str | None:
    cols = [str(c).strip() for c in columns]
    for name in ("项目状态", "备注", "状态", "target_status"):
        if name in cols:
            return name
    return None


def _pick_column(columns, *candidates: str) -> str | None:
    cols = {str(c).strip(): c for c in columns}
    for name in candidates:
        if name in cols:
            return str(cols[name]) if not isinstance(cols[name], str) else name
        # 宽松：去空格后比
    norm = {_normalize_key(c): str(c).strip() for c in columns}
    for name in candidates:
        if _normalize_key(name) in norm:
            return norm[_normalize_key(name)]
    return None


# ---------------------------------------------------------------------------
# DB 读写（函数式 API — 保持兼容）
# ---------------------------------------------------------------------------
def build_standard_projects_from_dataframe(info_df: pd.DataFrame) -> List[dict]:
    """从项目基本信息表 DataFrame 构建标准项目列表。"""
    status_col = _pick_status_column(info_df.columns)
    name_col = _pick_column(info_df.columns, "项目名称", "标准项目名称", "工程名称") or "项目名称"
    code_col = _pick_column(info_df.columns, "项目编码", "工程编码", "编码") or "项目编码"

    rows = []
    for _, row in info_df.iterrows():
        name = _clean_field(row.get(name_col, ""))
        if not name:
            continue
        code = _clean_field(row.get(code_col, ""), default="需手动核对")
        if status_col:
            status = _clean_field(row.get(status_col, ""), default="需手动核对")
        else:
            status = "需手动核对"
        rows.append(
            {
                "project_name": name,
                "project_code": code,
                "target_status": status,
            }
        )
    return rows


def load_or_init_db(db_path, standard_projects=None) -> dict:
    """
    加载匹配库，并可选合并标准项目列表。

    合并规则：
    - 新项目：写入编码/状态
    - 已有项目：仅当传入了有意义的编码/状态时才覆盖
    """
    db = {
        "version": 1,
        "projects": {},
        "unmatched_aliases": [],
    }
    path = str(db_path)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                loaded = json.load(f)
            if isinstance(loaded, dict):
                db.update(
                    {
                        "version": loaded.get("version", 1),
                        "projects": loaded.get("projects", {}) or {},
                        "unmatched_aliases": loaded.get("unmatched_aliases", []) or [],
                    }
                )
        except Exception:
            pass

    db["unmatched_aliases"] = [
        str(name).strip()
        for name in db.get("unmatched_aliases", [])
        if str(name).strip()
    ]

    for item in standard_projects or []:
        std_name = _clean_field(item.get("project_name", ""))
        if not std_name:
            continue
        has_code_key = "project_code" in item
        has_status_key = "target_status" in item
        code = _clean_field(item.get("project_code", "")) if has_code_key else ""
        status = _clean_field(item.get("target_status", "")) if has_status_key else ""

        if std_name not in db["projects"]:
            db["projects"][std_name] = {
                "project_code": code or "需手动核对",
                "target_status": status or "需手动核对",
                "aliases": [],
            }
        else:
            meta = db["projects"][std_name]
            if has_code_key and code and code != "需手动核对":
                meta["project_code"] = code
            elif not _clean_field(meta.get("project_code", "")):
                meta["project_code"] = code or "需手动核对"
            if has_status_key and status:
                meta["target_status"] = status
            elif not _clean_field(meta.get("target_status", "")):
                meta["target_status"] = "需手动核对"
            meta.setdefault("aliases", [])

        if std_name in db["unmatched_aliases"]:
            db["unmatched_aliases"].remove(std_name)

    return db


def save_db(db_path, db) -> None:
    db = dict(db)
    db.pop("updated_at", None)
    path = Path(db_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)


def get_project_options(db) -> List[str]:
    return sorted(db.get("projects", {}).keys())


def get_project_lookup(db) -> dict:
    lookup = {}
    for name, meta in db.get("projects", {}).items():
        lookup[name] = (
            meta.get("project_code", "需手动核对"),
            meta.get("target_status", "需手动核对"),
        )
    return lookup


def add_mapping(db, match_name, standard_name) -> None:
    if not match_name or not standard_name:
        return
    match_name = str(match_name).strip()
    standard_name = str(standard_name).strip()
    if not match_name or not standard_name:
        return
    if standard_name not in db["projects"]:
        db["projects"][standard_name] = {
            "project_code": "需手动核对",
            "target_status": "需手动核对",
            "aliases": [],
        }
    aliases = db["projects"][standard_name].setdefault("aliases", [])
    if match_name not in aliases and match_name != standard_name:
        aliases.append(match_name)
    clear_unmatched_aliases(db, match_name)
    _invalidate_db_caches(db)


def mark_unmatched(db, raw_name) -> None:
    if raw_name is None:
        return
    raw_name = str(raw_name).strip()
    if not raw_name:
        return
    std_name = _cached_alias_map(db).get(raw_name)
    if std_name and std_name in db.get("projects", {}):
        return
    unmatched_aliases = db.setdefault("unmatched_aliases", [])
    raw_key = _normalize_key(raw_name)
    for name in unmatched_aliases:
        if _normalize_key(name) == raw_key:
            return
    unmatched_aliases.append(raw_name)
    _invalidate_db_caches(db)


def clear_unmatched_aliases(db, raw_name=None) -> None:
    unmatched_aliases = db.setdefault("unmatched_aliases", [])
    if raw_name is None:
        unmatched_aliases.clear()
        return
    raw_key = _normalize_key(raw_name)
    db["unmatched_aliases"] = [
        name for name in unmatched_aliases if _normalize_key(name) != raw_key
    ]


def is_unmatched_alias(db, raw_name) -> bool:
    if raw_name is None:
        return False
    raw_key = _normalize_key(raw_name)
    if not raw_key:
        return False
    for name in db.get("unmatched_aliases", []):
        if _normalize_key(name) == raw_key:
            return True
    return False


def _result(raw_name, std_name, meta, source) -> dict:
    return {
        "raw_name": raw_name,
        "project_name": std_name,
        "project_code": meta.get("project_code", "需手动核对"),
        "target_status": meta.get("target_status", "需手动核对"),
        "source": source,
    }


def _fuzzy_match_project(raw_name: str, db, *, min_ratio: float = 0.82):
    cache_key = (id(db), _normalize_key(raw_name))
    if cache_key in _FUZZY_CACHE:
        return _FUZZY_CACHE[cache_key]

    projects = db.get("projects", {}) or {}
    if not projects:
        _FUZZY_CACHE[cache_key] = None
        return None

    raw_clean = _clean_name(raw_name)
    raw_key = _normalize_key(raw_name)
    if not raw_clean and not raw_key:
        _FUZZY_CACHE[cache_key] = None
        return None

    alias_map = _cached_alias_map(db)
    clean_hits = []
    for alias, std in alias_map.items():
        if std not in projects:
            continue
        if _clean_name(alias) == raw_clean or _normalize_key(alias) == raw_key:
            clean_hits.append(std)
    clean_hits = list(dict.fromkeys(clean_hits))
    if len(clean_hits) == 1:
        std = clean_hits[0]
        res = _result(raw_name, std, projects[std], "fuzzy_clean_exact")
        _FUZZY_CACHE[cache_key] = res
        return res

    if len(raw_clean) >= 4:
        contain = []
        for std in projects.keys():
            sc = _clean_name(std)
            if not sc:
                continue
            if raw_clean in sc or sc in raw_clean:
                shorter, longer = (
                    (raw_clean, sc) if len(raw_clean) <= len(sc) else (sc, raw_clean)
                )
                if len(shorter) >= 4 and len(shorter) / max(len(longer), 1) >= 0.5:
                    contain.append(std)
        contain = list(dict.fromkeys(contain))
        if len(contain) == 1:
            std = contain[0]
            res = _result(raw_name, std, projects[std], "fuzzy_contain")
            _FUZZY_CACHE[cache_key] = res
            return res

    best_std, best_r, second_r = None, 0.0, 0.0
    for std in projects.keys():
        sc = _clean_name(std) or _normalize_key(std)
        if not sc:
            continue
        r = SequenceMatcher(None, raw_clean or raw_key, sc).ratio()
        if r > best_r:
            second_r = best_r
            best_r = r
            best_std = std
        elif r > second_r:
            second_r = r
    if best_std and best_r >= min_ratio and (best_r - second_r) >= 0.05:
        res = _result(
            raw_name, best_std, projects[best_std], f"fuzzy_ratio_{best_r:.2f}"
        )
        _FUZZY_CACHE[cache_key] = res
        return res

    _FUZZY_CACHE[cache_key] = None
    return None


def match_project(
    raw_name,
    db,
    *,
    fuzzy: bool = True,
    skip_marked_unmatched: bool = True,
) -> Optional[dict]:
    """
    匹配项目名 → dict 或 None。

    返回字段: raw_name, project_name, project_code, target_status, source

    匹配顺序：
    1. 精确 / 规范化别名（已确认映射始终优先）
    2. 若已在 unmatched_aliases 中标记过 → 直接保留原名，不再模糊匹配
       （自动批处理不再弹窗；手动运行 prompt=True 时仍会弹出补录对话框）
    3. 模糊匹配（仅对从未标记过的新名称）
    """
    if raw_name is None:
        return None
    raw_name = str(raw_name).strip()
    if not raw_name:
        return None

    projects = db.get("projects", {})
    alias_map = _cached_alias_map(db)

    std_name = alias_map.get(raw_name)
    if std_name and std_name in projects:
        return _result(raw_name, std_name, projects[std_name], "alias")

    raw_key = _normalize_key(raw_name)
    for alias, std in alias_map.items():
        if _normalize_key(alias) == raw_key and std in projects:
            meta = projects[std]
            add_mapping(db, raw_name, std)
            return _result(raw_name, std, meta, "normalized_alias")

    # 二次及以后运行：曾标记为未匹配的名称，跳过模糊匹配与再次弹窗
    if skip_marked_unmatched and is_unmatched_alias(db, raw_name):
        return {
            "raw_name": raw_name,
            "project_name": raw_name,
            "project_code": "",
            "target_status": "未匹配保留",
            "source": "unmatched_keep",
        }

    if fuzzy:
        hit = _fuzzy_match_project(raw_name, db)
        if hit:
            add_mapping(db, raw_name, hit["project_name"])
            return hit

    return None


def collect_names_needing_prompt(
    raw_names,
    db,
    *,
    fuzzy: bool = True,
    include_marked_unmatched: bool = False,
) -> list[str]:
    need = []
    seen = set()
    for raw in raw_names:
        name = str(raw or "").strip()
        if not name:
            continue
        key = _normalize_key(name)
        if key in seen:
            continue
        seen.add(key)
        hit = match_project(name, db, fuzzy=fuzzy, skip_marked_unmatched=True)
        if hit and hit.get("source") not in {"unmatched_keep", "unmatched"}:
            continue
        if (
            hit
            and hit.get("source") == "unmatched_keep"
            and not include_marked_unmatched
        ):
            continue
        need.append(name)
    return sorted(need)


def prompt_unmatched_project_mappings(
    unmatched_names,
    db,
    window_title="手动项目匹配",
    db_path=None,
    *,
    interactive: bool = True,
    parent=None,
):
    """弹窗补录未匹配项目。interactive=False 时全部 mark_unmatched。

    parent 传入已有 Tk 根窗口时复用，避免图形界面里再创建一个 Tk。
    """
    unmatched_names = [str(n).strip() for n in unmatched_names if str(n).strip()]
    unmatched_names = collect_names_needing_prompt(
        unmatched_names, db, fuzzy=True, include_marked_unmatched=True
    )
    if not unmatched_names:
        return {}

    if not interactive or tk is None or ttk is None:
        for name in unmatched_names:
            mark_unmatched(db, name)
        if db_path:
            parent = os.path.dirname(str(db_path)) or "."
            if os.path.exists(parent):
                save_db(db_path, db)
        return {name: "" for name in unmatched_names}

    options = get_project_options(db)
    if not options:
        for name in unmatched_names:
            mark_unmatched(db, name)
        return {name: "" for name in unmatched_names}

    suggestions: dict[str, str] = {}
    for raw in unmatched_names:
        hit = _fuzzy_match_project(raw, db, min_ratio=0.72)
        if hit:
            suggestions[raw] = hit["project_name"]

    created_root = False
    root = parent if parent is not None else getattr(tk, "_default_root", None)
    if root is None:
        root = tk.Tk()
        root.withdraw()
        created_root = True

    win = tk.Toplevel(root)
    win.title(f"{window_title}（{len(unmatched_names)}项）")
    win.geometry("980x660")
    if not created_root:
        win.transient(root)
    win.lift()
    win.attributes("-topmost", True)
    win.after_idle(win.attributes, "-topmost", False)
    try:
        win.grab_set()
    except tk.TclError:
        pass

    tip = (
        "以下为未匹配项目（含此前保留原名称的项）。下拉可改；建议已预填。"
        "选「保持原名称」后，自动批处理会跳过；手动运行仍会弹出。"
    )
    tk.Label(
        win, text=tip, font=("Arial", 11, "bold"), pady=8, wraplength=900, justify="left"
    ).pack()

    canvas = tk.Canvas(win)
    scrollbar = tk.Scrollbar(win, orient="vertical", command=canvas.yview)
    frame = tk.Frame(canvas)
    frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
    canvas.create_window((0, 0), window=frame, anchor="nw")
    canvas.configure(yscrollcommand=scrollbar.set)
    canvas.pack(side="left", fill="both", expand=True, padx=10)
    scrollbar.pack(side="right", fill="y")

    blank_option = "（保持原名称，不匹配）"
    display_options = [blank_option] + list(options)
    option_to_standard = {s: s for s in options}

    var_map = {}
    for raw_name in unmatched_names:
        row = tk.Frame(frame, relief=tk.RIDGE, borderwidth=1)
        row.pack(fill="x", padx=8, pady=4)
        tk.Label(row, text=raw_name, font=("Arial", 11, "bold"), width=28, anchor="w").pack(
            side="left", padx=8, pady=6
        )
        tk.Label(row, text="→", font=("Arial", 12)).pack(side="left")
        suggest = suggestions.get(raw_name, "")
        default = suggest if suggest in option_to_standard else blank_option
        var = tk.StringVar(value=default)
        combo = ttk.Combobox(
            row, textvariable=var, values=display_options, width=68, state="readonly"
        )
        combo.pack(side="left", padx=8, pady=6, fill="x", expand=True)
        if suggest:
            tk.Label(
                row, text=f"建议:{suggest[:20]}", fg="#666", font=("Arial", 9)
            ).pack(side="left", padx=4)
        var_map[raw_name] = var

    result = {"ok": False}

    def on_confirm():
        result["ok"] = True
        win.destroy()

    def on_skip_all():
        result["ok"] = False
        win.destroy()

    btns = tk.Frame(win)
    btns.pack(pady=10)
    tk.Button(
        btns, text="确定并应用", command=on_confirm, width=16, bg="#4CAF50", fg="white"
    ).pack(side="left", padx=10)
    tk.Button(btns, text="全部保留原名称", command=on_skip_all, width=16).pack(
        side="left", padx=10
    )
    win.protocol("WM_DELETE_WINDOW", on_skip_all)
    root.wait_window(win)
    if created_root:
        root.destroy()

    selected_map = {}
    for raw_name, var in var_map.items():
        selected = var.get().strip()
        if not result["ok"] or selected == blank_option or not selected:
            mark_unmatched(db, raw_name)
            selected_map[raw_name] = ""
            continue
        std_name = option_to_standard.get(selected, "")
        selected_map[raw_name] = std_name
        if std_name:
            add_mapping(db, raw_name, std_name)

    if db_path and selected_map:
        parent = os.path.dirname(str(db_path))
        if parent and os.path.exists(parent):
            save_db(db_path, db)
            print(f"✅ 手动匹配结果已保存: {db_path}")

    return selected_map


def update_db_from_excel(
    db_path,
    excel_path,
    status_column=None,
    name_column=None,
    code_column=None,
):
    """
    从项目基本信息 Excel 更新匹配库（以项目编码对齐）。
    列名可省略，会自动识别「项目名称/项目编码/项目状态」。
    """
    if not os.path.exists(excel_path):
        print(f"❌ Excel 文件不存在: {excel_path}")
        return None

    print(f"正在读取 Excel: {excel_path}")
    df = pd.read_excel(excel_path)
    df = df.dropna(how="all")

    name_column = name_column or _pick_column(
        df.columns, "项目名称", "标准项目名称", "工程名称"
    )
    code_column = code_column or _pick_column(
        df.columns, "项目编码", "工程编码", "编码"
    )
    status_column = status_column or _pick_status_column(df.columns)

    if not name_column or not code_column:
        print(f"⚠️ 缺少名称/编码列，当前列: {list(df.columns)}")
        return None

    db = load_or_init_db(db_path)
    updated_count = 0
    new_count = 0

    for _, row in df.iterrows():
        code = _clean_field(row.get(code_column, ""))
        if not code:
            continue
        name = _clean_field(row.get(name_column, ""))
        status = (
            _clean_field(row.get(status_column, "")) if status_column else ""
        )

        found = False
        for std_name, meta in list(db["projects"].items()):
            if str(meta.get("project_code", "")).strip() == code:
                if status and str(meta.get("target_status", "")).strip() != status:
                    meta["target_status"] = status
                    updated_count += 1
                if name and name != std_name and name not in meta.get("aliases", []):
                    meta.setdefault("aliases", []).append(name)
                found = True
                break

        if not found:
            if not name:
                name = f"未知项目_{code}"
            db["projects"][name] = {
                "project_code": code,
                "target_status": status or "需手动核对",
                "aliases": [],
            }
            new_count += 1
            print(f"✅ 新增项目: {name} (编码: {code})")

    save_db(db_path, db)
    _invalidate_db_caches(db)
    print("\n🎉 更新完成！")
    print(f"   • 更新状态: {updated_count} 条")
    print(f"   • 新增项目: {new_count} 个")
    print(f"   • 当前标准项目总数: {len(db.get('projects', {}))}")
    return db


# ---------------------------------------------------------------------------
# 面向对象门面 — 推荐新代码使用
# ---------------------------------------------------------------------------
class ProjectMatcher:
    """
    项目匹配器（状态：内存中的 db + 落盘路径）。

    示例:
        m = ProjectMatcher()                 # 默认 JSON，有主数据表则合并
        hit = m.match("永强")
        name, code = m.resolve("永强")
        out = m.enrich_dataframe(df, col="项目名称")
        m.save()
    """

    def __init__(
        self,
        db_path: PathLike | None = None,
        info_excel: PathLike | None = None,
        *,
        load_info: bool = True,
        standard_projects: Optional[List[dict]] = None,
    ):
        self.db_path = Path(db_path) if db_path else DEFAULT_DB_PATH
        self.info_excel = find_info_excel(info_excel) if load_info or info_excel else None

        std = list(standard_projects or [])
        if load_info and not std:
            info = self.info_excel or find_info_excel()
            if info and info.is_file():
                try:
                    df = pd.read_excel(info)
                    std = build_standard_projects_from_dataframe(df)
                    self.info_excel = info
                except Exception as e:
                    print(f"⚠️ 读取项目基本信息失败（忽略，仅用 JSON）: {e}")
        self.db = load_or_init_db(str(self.db_path), std)

    # --- 单条 ---
    def match(
        self,
        raw_name,
        *,
        fuzzy: bool = True,
        skip_marked_unmatched: bool = True,
    ) -> Optional[dict]:
        return match_project(
            raw_name,
            self.db,
            fuzzy=fuzzy,
            skip_marked_unmatched=skip_marked_unmatched,
        )

    def resolve(self, raw_name) -> Tuple[str, str]:
        """返回 (标准项目名称, 项目编码)；未匹配则 (原名, "")。"""
        hit = self.match(raw_name)
        if not hit:
            n = str(raw_name or "").strip()
            return n, ""
        return (
            str(hit.get("project_name") or raw_name or "").strip(),
            str(hit.get("project_code") or "").strip(),
        )

    def code(self, raw_name) -> str:
        return self.resolve(raw_name)[1]

    def name(self, raw_name) -> str:
        return self.resolve(raw_name)[0]

    def status(self, raw_name) -> str:
        hit = self.match(raw_name)
        return str(hit.get("target_status") or "") if hit else ""

    # --- 批量 ---
    def match_batch(
        self,
        raw_names: Iterable,
        *,
        prompt: bool = False,
        interactive: bool = True,
        window_title: str = "手动项目匹配",
        fuzzy: bool = True,
        parent=None,
    ) -> Dict[str, dict]:
        """
        批量匹配。prompt=True 时对未命中弹窗补录后再匹配一次
        （含此前「保持原名称」记入 unmatched_aliases 的项）。
        返回 {原始名: hit_dict}。

        注意：去空白/大小写后相同的名称（如「KB项目」「kb项目」「KB 项目」）
        只匹配一次，结果会回填到所有写法，避免只命中其中一种写法。
        """
        # 规范化键 → 该键下出现过的全部原始写法
        groups: Dict[str, List[str]] = {}
        for n in raw_names:
            s = str(n or "").strip()
            if not s or s.lower() == "nan":
                continue
            k = _normalize_key(s)
            if not k:
                continue
            bucket = groups.setdefault(k, [])
            if s not in bucket:
                bucket.append(s)

        # 每个规范化键取一个代表名做匹配
        representatives = [variants[0] for variants in groups.values()]

        def _broadcast(rep: str, hit: dict) -> None:
            """把命中结果写到同一规范化键下的全部原始写法。"""
            k = _normalize_key(rep)
            for original in groups.get(k, [rep]):
                h = dict(hit)
                h["raw_name"] = original
                out[original] = h

        out: Dict[str, dict] = {}
        pending: List[str] = []
        skipped_keep = 0
        for n in representatives:
            hit = self.match(n, fuzzy=fuzzy)
            if hit and hit.get("source") != "unmatched_keep":
                _broadcast(n, hit)
            elif hit and hit.get("source") == "unmatched_keep":
                # 手动弹窗：已保留原名的项也要再确认；自动批处理才跳过
                if prompt:
                    pending.append(n)
                else:
                    _broadcast(n, hit)
                    skipped_keep += 1
            else:
                pending.append(n)

        if skipped_keep:
            print(
                f"  跳过已标记未匹配（未弹窗）: {skipped_keep} 个"
            )

        if pending and prompt:
            print(f"  未匹配项目需确认: {len(pending)} 个")
            prompt_unmatched_project_mappings(
                pending,
                self.db,
                window_title=window_title,
                db_path=str(self.db_path),
                interactive=interactive,
                parent=parent,
            )
            for n in pending:
                # 弹窗后：有别名则用别名；已点跳过/关闭的标记未匹配
                hit = self.match(n, fuzzy=False)
                if hit and hit.get("source") not in {"unmatched_keep", "unmatched"}:
                    _broadcast(n, hit)
                else:
                    mark_unmatched(self.db, n)
                    _broadcast(
                        n,
                        {
                            "raw_name": n,
                            "project_name": n,
                            "project_code": "",
                            "target_status": "未匹配",
                            "source": "unmatched",
                        },
                    )
                    # 同键其它写法也记入未匹配，避免下次再弹
                    for original in groups.get(_normalize_key(n), []):
                        if original != n:
                            mark_unmatched(self.db, original)
        elif pending:
            # 不弹窗：保留原名，不写入 unmatched_aliases，便于下次主动 --prompt-project
            print(f"  未匹配保留原名（未弹窗、未标记）: {len(pending)} 个")
            for n in pending:
                _broadcast(
                    n,
                    {
                        "raw_name": n,
                        "project_name": n,
                        "project_code": "",
                        "target_status": "未匹配",
                        "source": "unmatched",
                    },
                )
        return out

    def enrich_dataframe(
        self,
        df: pd.DataFrame,
        col: str = "项目名称",
        *,
        code_col: str = "项目编码",
        name_col: Optional[str] = None,
        prompt: bool = False,
        interactive: bool = True,
        window_title: str = "项目匹配",
        parent=None,
    ) -> pd.DataFrame:
        """
        给 DataFrame 写入标准化项目名称与编码。
        name_col 默认覆盖原 col；若想保留原列可设 name_col="标准项目名称"。
        """
        if col not in df.columns:
            raise KeyError(f"DataFrame 无列: {col}")
        out = df.copy()
        series = out[col].fillna("").astype(str).str.strip()
        mapped = self.match_batch(
            series.tolist(),
            prompt=prompt,
            interactive=interactive,
            window_title=window_title,
            parent=parent,
        )
        codes = series.map(lambda x: mapped.get(x, {}).get("project_code", "") if x else "")
        names = series.map(
            lambda x: mapped.get(x, {}).get("project_name", x) if x else x
        )
        out[code_col] = codes
        target_name = name_col if name_col is not None else col
        out[target_name] = names
        return out

    def fill_excel(self, excel_path: PathLike, **kwargs) -> dict:
        """按列填写 Excel 标准名称/编码（只改指定单元格，保留原格式）。"""
        kwargs.setdefault("matcher", self)
        return fill_excel_project_columns(excel_path, **kwargs)

    # --- 维护 ---
    def add_alias(self, raw_name: str, standard_name: str) -> None:
        add_mapping(self.db, raw_name, standard_name)

    def save(self) -> None:
        save_db(str(self.db_path), self.db)

    def reload(self, *, load_info: bool = False) -> None:
        std = None
        if load_info:
            info = self.info_excel or find_info_excel()
            if info and info.is_file():
                std = build_standard_projects_from_dataframe(pd.read_excel(info))
        self.db = load_or_init_db(str(self.db_path), std)
        _invalidate_db_caches(self.db)

    def refresh_from_excel(self, excel_path: PathLike | None = None) -> Optional[dict]:
        path = Path(excel_path) if excel_path else (self.info_excel or find_info_excel())
        if not path or not Path(path).is_file():
            print("⚠️ 未找到项目基本信息表")
            return None
        db = update_db_from_excel(str(self.db_path), str(path))
        if db is not None:
            self.db = db
            self.info_excel = Path(path)
        return db

    @property
    def project_count(self) -> int:
        return len(self.db.get("projects", {}) or {})

    def __repr__(self) -> str:
        return (
            f"ProjectMatcher(db={self.db_path.name}, "
            f"projects={self.project_count})"
        )


def load_matcher(
    db_path: PathLike | None = None,
    info_excel: PathLike | None = None,
    *,
    load_info: bool = True,
) -> ProjectMatcher:
    """创建匹配器（最常用工厂函数）。"""
    ensure_on_sys_path()
    return ProjectMatcher(db_path=db_path, info_excel=info_excel, load_info=load_info)


@lru_cache(maxsize=4)
def get_matcher(
    db_path: str = "",
    info_excel: str = "",
    load_info: bool = True,
) -> ProjectMatcher:
    """
    缓存的默认匹配器（路径用字符串；空串 = 默认）。
    注意：写别名后需 m.save()；多进程勿共享。
    """
    return ProjectMatcher(
        db_path=db_path or None,
        info_excel=info_excel or None,
        load_info=load_info,
    )


def clear_matcher_cache() -> None:
    get_matcher.cache_clear()
    _invalidate_db_caches()


# ---------------------------------------------------------------------------
# Excel 填写（保留原格式）
# ---------------------------------------------------------------------------
def peek_excel_sheets_and_headers(
    excel_path: PathLike, header_row: int = 1
) -> Tuple[List[str], Dict[str, List[str]]]:
    """只读扫描工作表名与表头，供图形界面下拉框使用。"""
    from openpyxl import load_workbook

    path = Path(excel_path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = list(wb.sheetnames)
        headers: Dict[str, List[str]] = {}
        for name in sheets:
            ws = wb[name]
            row_vals: List[str] = []
            for i, row in enumerate(
                ws.iter_rows(min_row=1, max_row=header_row, values_only=True), 1
            ):
                if i != header_row:
                    continue
                row_vals = [str(c).strip() if c is not None else "" for c in row]
            while row_vals and not row_vals[-1]:
                row_vals.pop()
            headers[name] = row_vals
        return sheets, headers
    finally:
        wb.close()


def _read_header_map(ws, header_row: int = 1) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    last = ws.max_column or 1
    for col in range(1, last + 1):
        value = ws.cell(header_row, col).value
        if value is None:
            continue
        key = str(value).strip()
        if key and key not in mapping:
            mapping[key] = col
    return mapping


def _resolve_header_col(header_map: Dict[str, int], name: str, *fallbacks: str) -> Optional[int]:
    if not name:
        return None
    candidates = (name,) + fallbacks
    for cand in candidates:
        if cand in header_map:
            return header_map[cand]
        want = _normalize_key(cand)
        for key, idx in header_map.items():
            if _normalize_key(key) == want:
                return idx
    return None


def _last_header_col(ws, header_row: int = 1) -> int:
    last = 1
    for col in range(1, (ws.max_column or 1) + 1):
        value = ws.cell(header_row, col).value
        if value is not None and str(value).strip():
            last = col
    return last


def _ensure_header_column(ws, header_row: int, col_name: str, header_map: Dict[str, int]) -> int:
    existing = _resolve_header_col(header_map, col_name)
    if existing:
        return existing
    new_idx = _last_header_col(ws, header_row) + 1
    cell = ws.cell(header_row, new_idx, col_name)
    src = ws.cell(header_row, new_idx - 1)
    try:
        if src.has_style:
            cell.font = src.font.copy()
            cell.fill = src.fill.copy()
            cell.alignment = src.alignment.copy()
            cell.border = src.border.copy()
    except Exception:
        pass
    header_map[col_name] = new_idx
    return new_idx


def _cell_is_empty(value) -> bool:
    if value is None:
        return True
    s = str(value).strip()
    return (not s) or s.lower() in {"nan", "none", "nat"}


def _is_unmatched_hit(hit: Optional[dict]) -> bool:
    if not hit:
        return True
    if hit.get("source") in {"unmatched", "unmatched_keep"}:
        return True
    return False


def fill_excel_project_columns(
    excel_path: PathLike,
    *,
    sheet_name: Optional[str] = None,
    match_col: str = "项目名称",
    name_out_col: str = "项目标准名称",
    code_out_col: str = "项目编码",
    summary_col: Optional[str] = None,
    summary_out_col: Optional[str] = None,
    process_summary: bool = False,
    header_row: int = 1,
    matcher: Optional[ProjectMatcher] = None,
    prompt: bool = False,
    interactive: bool = True,
    parent=None,
    only_empty: bool = False,
    output_path: PathLike | None = None,
    create_missing_columns: bool = True,
    log=None,
) -> dict:
    """
    读取 Excel，按「待匹配列」匹配标准项目名称与编码，写回指定列。

    只修改目标列的单元格值，不改列宽、样式、冻结、筛选、其它列内容。
    process_summary=True 时，按待匹配列的项目名称清理摘要前缀，写入摘要输出列。
    """
    from openpyxl import load_workbook

    def _log(msg: str) -> None:
        print(msg)
        if log:
            log(msg)

    path = Path(excel_path).expanduser()
    if not path.is_file():
        raise FileNotFoundError(f"Excel 文件不存在: {path}")

    out_path = Path(output_path).expanduser() if output_path else path
    m = matcher or load_matcher()

    _log(f"打开: {path}")
    wb = load_workbook(path)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise KeyError(f"工作表不存在: {sheet_name}；现有: {wb.sheetnames}")
        ws = wb[sheet_name]
    else:
        ws = wb.active

    header_map = _read_header_map(ws, header_row)
    match_idx = _resolve_header_col(header_map, match_col, "项目名称", "工程名称")
    if not match_idx:
        wb.close()
        raise KeyError(f"找不到待匹配列 {match_col!r}，当前表头: {list(header_map)}")

    def _out_col(wanted: str, *fallback: str) -> int:
        idx = _resolve_header_col(header_map, wanted, *fallback)
        if idx:
            return idx
        if not create_missing_columns:
            raise KeyError(f"找不到输出列 {wanted!r}，当前表头: {list(header_map)}")
        _log(f"  新增列: {wanted}")
        return _ensure_header_column(ws, header_row, wanted, header_map)

    name_idx = _out_col(name_out_col, "项目标准名称", "标准项目名称")
    code_idx = _out_col(code_out_col, "项目编码", "工程编码")

    summary_idx = None
    summary_out_idx = None
    if process_summary:
        summary_name = summary_col or "摘要"
        summary_out_name = summary_out_col or "摘要处理"
        summary_idx = _resolve_header_col(header_map, summary_name, "摘要")
        if not summary_idx:
            wb.close()
            raise KeyError(f"找不到摘要列 {summary_name!r}，当前表头: {list(header_map)}")
        summary_out_idx = _out_col(summary_out_name, "摘要处理")

    max_row = ws.max_row or header_row
    raw_names: List[str] = []
    for row_idx in range(header_row + 1, max_row + 1):
        value = ws.cell(row_idx, match_idx).value
        if _cell_is_empty(value):
            continue
        raw_names.append(str(value).strip())

    _log(
        f"工作表「{ws.title}」共 {max_row - header_row} 行，"
        f"待匹配列「{ws.cell(header_row, match_idx).value}」非空 {len(raw_names)}，"
        f"唯一名称 {len({_normalize_key(n) for n in raw_names})}"
    )
    mapped = m.match_batch(
        raw_names,
        prompt=prompt,
        interactive=interactive,
        window_title="Excel 项目匹配",
        parent=parent,
    )
    if prompt:
        m.save()

    matched = unmatched = empty_name = 0
    summary_stripped = summary_processed = 0
    unmatched_names: List[str] = []
    seen_unmatched: set[str] = set()

    for row_idx in range(header_row + 1, max_row + 1):
        raw_val = ws.cell(row_idx, match_idx).value
        if _cell_is_empty(raw_val):
            empty_name += 1
            raw_s = ""
            std_name = ""
            code = ""
            hit = None
        else:
            raw_s = str(raw_val).strip()
            hit = mapped.get(raw_s)
            if _is_unmatched_hit(hit):
                unmatched += 1
                std_name = raw_s
                code = ""
                key = _normalize_key(raw_s)
                if key not in seen_unmatched:
                    seen_unmatched.add(key)
                    unmatched_names.append(raw_s)
            else:
                matched += 1
                std_name = str(hit.get("project_name") or raw_s).strip()
                code = str(hit.get("project_code") or "").strip()

        name_cell = ws.cell(row_idx, name_idx)
        code_cell = ws.cell(row_idx, code_idx)
        if not (only_empty and not _cell_is_empty(name_cell.value)):
            name_cell.value = std_name or None
        if not (only_empty and not _cell_is_empty(code_cell.value)):
            code_cell.value = code or None

        if process_summary and summary_idx and summary_out_idx:
            summary_val = ws.cell(row_idx, summary_idx).value
            processed = strip_project_prefix_from_summary(raw_s, summary_val)
            out_cell = ws.cell(row_idx, summary_out_idx)
            if not (only_empty and not _cell_is_empty(out_cell.value)):
                out_cell.value = processed
                summary_processed += 1
                orig = "" if summary_val is None else str(summary_val)
                new = "" if processed is None else str(processed)
                if new != orig:
                    summary_stripped += 1

    sheet_title = ws.title
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = out_path.with_name(f".~{out_path.stem}.tmp{out_path.suffix}")
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, out_path)
    except Exception:
        if tmp_path.exists():
            try:
                tmp_path.unlink()
            except OSError:
                pass
        raise
    finally:
        wb.close()

    stats = {
        "excel_path": str(path),
        "output_path": str(out_path),
        "sheet": sheet_title,
        "rows": max_row - header_row,
        "matched": matched,
        "unmatched": unmatched,
        "empty_name": empty_name,
        "summary_stripped": summary_stripped,
        "summary_processed": summary_processed,
        "unique_names": len({_normalize_key(n) for n in raw_names}),
        "unmatched_names": unmatched_names,
    }
    _log(
        f"已保存: {out_path}\n"
        f"  匹配 {matched} 行，未匹配 {unmatched} 行"
        + (f"，摘要去掉项目名前缀 {summary_stripped} 行" if process_summary else "")
    )
    if unmatched_names:
        preview = "、".join(unmatched_names[:20])
        extra = f" 等 {len(unmatched_names)} 个" if len(unmatched_names) > 20 else ""
        _log(f"  未匹配名称: {preview}{extra}")
    return stats


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def _cli(argv: Optional[List[str]] = None) -> int:
    import argparse

    p = argparse.ArgumentParser(
        description="项目名称匹配工具",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  python -m common.project_name_matcher              # 打开图形界面
  python -m common.project_name_matcher gui
  python -m common.project_name_matcher match 永强
  python -m common.project_name_matcher excel 台账.xlsx --process-summary
  python -m common.project_name_matcher refresh
  python -m common.project_name_matcher stats
        """,
    )
    sub = p.add_subparsers(dest="cmd")
    sub.add_parser("gui", help="打开图形界面（选 Excel / 列 / 填写）")

    c = sub.add_parser("match", help="匹配一个名称")
    c.add_argument("name")
    c.add_argument("--db", default="")
    c.add_argument("--no-info", action="store_true")

    c = sub.add_parser("refresh", help="从项目基本信息表刷新 JSON")
    c.add_argument("--excel", default="")
    c.add_argument("--db", default="")

    sub.add_parser("stats", help="显示库统计")

    c = sub.add_parser("demo", help="演示 resolve 几个名称")
    c.add_argument("names", nargs="*", default=["永强", "樱花", "不存在的项目XYZ"])

    c = sub.add_parser("excel", help="填写 Excel 标准名称/编码（保留原格式）")
    c.add_argument("path", help="Excel 路径")
    c.add_argument("--sheet", default="")
    c.add_argument("--match", default="项目名称", help="待匹配列")
    c.add_argument("--name-out", default="项目标准名称", help="标准名称输出列")
    c.add_argument("--code-out", default="项目编码", help="项目编码输出列")
    c.add_argument("--summary", default="", help="摘要列（处理摘要时用）")
    c.add_argument("--summary-out", default="", help="摘要处理输出列")
    c.add_argument("--process-summary", action="store_true", help="清理摘要前的项目名称")
    c.add_argument("--prompt", action="store_true", help="未匹配时弹窗补录")
    c.add_argument("--only-empty", action="store_true", help="只填空白单元格")
    c.add_argument("--out", default="", help="另存路径，默认覆盖原文件")
    c.add_argument("--db", default="")

    args = p.parse_args(argv)
    if not args.cmd or args.cmd == "gui":
        from common.project_name_matcher_gui import main as gui_main

        return gui_main()

    if args.cmd == "match":
        m = load_matcher(
            db_path=args.db or None,
            load_info=not args.no_info,
        )
        hit = m.match(args.name)
        print(json.dumps(hit, ensure_ascii=False, indent=2) if hit else "未匹配")
        m.save()
        return 0

    if args.cmd == "refresh":
        m = load_matcher(db_path=args.db or None, load_info=False)
        m.refresh_from_excel(args.excel or None)
        return 0

    if args.cmd == "stats":
        m = load_matcher()
        print(m)
        print(f"db: {m.db_path}")
        print(f"info: {m.info_excel}")
        print(f"unmatched_aliases: {len(m.db.get('unmatched_aliases') or [])}")
        return 0

    if args.cmd == "demo":
        m = load_matcher()
        for n in args.names:
            name, code = m.resolve(n)
            print(f"  {n!r:20} → {name} | {code}")
        return 0

    if args.cmd == "excel":
        m = load_matcher(db_path=args.db or None)
        process_summary = bool(args.process_summary or args.summary or args.summary_out)
        fill_excel_project_columns(
            args.path,
            sheet_name=args.sheet or None,
            match_col=args.match,
            name_out_col=args.name_out,
            code_out_col=args.code_out,
            summary_col=args.summary or None,
            summary_out_col=args.summary_out or None,
            process_summary=process_summary,
            matcher=m,
            prompt=args.prompt,
            only_empty=args.only_empty,
            output_path=args.out or None,
        )
        if args.prompt:
            m.save()
        return 0

    return 0


if __name__ == "__main__":
    raise SystemExit(_cli())
