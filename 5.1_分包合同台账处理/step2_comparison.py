# -*- coding: utf-8 -*-
"""
步骤2：逐月对比 + 项目合同额透视 / 增量 / 协作体统计。
"""

from __future__ import annotations

import json
import os
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from config import AMOUNT_TOLERANCE, COMPARISON_ONLY_NAME, CODE_ROOT, STD_DIR, SUM_DIR

# 格式化工具
import sys

if str(CODE_ROOT) not in sys.path:
    sys.path.insert(0, str(CODE_ROOT))
from common.excel import apply_common_excel_layout  # noqa: E402


def extract_year_month(filename: str) -> str | None:
    match1 = re.search(r"(\d{4})[-_]?(\d{2})", filename)
    if match1:
        return match1.group(1) + match1.group(2)
    match2 = re.search(r"(\d{4})-(\d{1,2})", filename)
    if match2:
        return match2.group(1) + match2.group(2).zfill(2)
    # 兼容 2026-07 文件名中的 2026-07
    match3 = re.search(r"(20\d{2})[-_年](\d{1,2})", filename)
    if match3:
        return match3.group(1) + f"{int(match3.group(2)):02d}"
    return None


def read_contract_file(filepath: str) -> pd.DataFrame | None:
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext in [".xlsx", ".xls"]:
            df = None
            with pd.ExcelFile(filepath) as xls:
                for sheet_name in ["Sheet1", "分包合同", "总包合同", xls.sheet_names[0]]:
                    if sheet_name in xls.sheet_names:
                        df = pd.read_excel(xls, sheet_name=sheet_name)
                        break
            if df is None:
                df = pd.read_excel(filepath)
            return df.dropna(axis=1, how="all")
        if ext == ".json":
            with open(filepath, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                for key in ["分包合同", "总包合同", "data", "contracts"]:
                    if key in data:
                        data = data[key]
                        break
            return pd.DataFrame(data) if isinstance(data, list) else pd.DataFrame([data])
    except Exception as e:
        print(f"  警告：读取文件失败 {filepath}: {e}")
        return None
    return None


def scan_contract_files(data_dir: str | Path | None = None) -> list:
    data_dir = Path(data_dir or STD_DIR)
    files_data = []
    print(f"正在扫描目录: {data_dir}")
    if not data_dir.exists():
        print(f"错误：目录不存在 {data_dir}")
        return files_data

    for filename in os.listdir(data_dir):
        filepath = data_dir / filename
        if not filepath.is_file() or filename.startswith(".") or filename.startswith("~$"):
            continue
        if "钢结构" in filename:
            continue
        ext = filepath.suffix.lower()
        if ext not in [".xlsx", ".xls", ".json"]:
            continue
        year_month = extract_year_month(filename)
        if not year_month:
            print(f"  跳过（无法提取年月）: {filename}")
            continue
        print(f"  发现文件: {filename} -> {year_month}")
        df = read_contract_file(str(filepath))
        if df is not None and len(df) > 0:
            files_data.append((year_month, filename, str(filepath), df))
        else:
            print(f"  警告：文件为空或读取失败: {filename}")

    files_data.sort(key=lambda x: x[0])
    unique_data = {}
    for year_month, filename, filepath, df in files_data:
        if year_month not in unique_data:
            unique_data[year_month] = (year_month, filename, filepath, df)
        else:
            current_ext = os.path.splitext(filepath)[1].lower()
            existing_ext = os.path.splitext(unique_data[year_month][2])[1].lower()
            if current_ext in [".xlsx", ".xls"] and existing_ext not in [".xlsx", ".xls"]:
                unique_data[year_month] = (year_month, filename, filepath, df)

    result = sorted(unique_data.values(), key=lambda x: x[0])
    print(f"\n共找到 {len(result)} 个有效月份的数据")
    for ym, fn, _, _ in result:
        print(f"  {ym}: {fn}")
    return result


def compare_two_months(
    df_prev,
    df_curr,
    prev_month,
    curr_month,
    amount_col="合同额（泰铢）",
    contract_id_col="合同编号",
):
    if contract_id_col not in df_prev.columns or contract_id_col not in df_curr.columns:
        print(f"  警告：找不到合同编号列 '{contract_id_col}'")
        return {"new": pd.DataFrame(), "deleted": pd.DataFrame(), "changed": pd.DataFrame()}

    prev_contracts = set(df_prev[contract_id_col].astype(str).dropna().unique())
    curr_contracts = set(df_curr[contract_id_col].astype(str).dropna().unique())
    new_contract_ids = curr_contracts - prev_contracts
    deleted_contract_ids = prev_contracts - curr_contracts

    new_df = df_curr[df_curr[contract_id_col].astype(str).isin(new_contract_ids)].copy()
    deleted_df = df_prev[df_prev[contract_id_col].astype(str).isin(deleted_contract_ids)].copy()

    common_contracts = prev_contracts & curr_contracts
    changed_list = []
    possible_amount_cols = [
        amount_col,
        "合同金额",
        "合同额",
        "金额（万泰铢）",
        "合同额（万泰铢）",
    ]
    amount_col_prev = next((c for c in possible_amount_cols if c in df_prev.columns), None)
    amount_col_curr = next((c for c in possible_amount_cols if c in df_curr.columns), None)

    if amount_col_prev and amount_col_curr:
        for contract_id in common_contracts:
            prev_row = df_prev[df_prev[contract_id_col].astype(str) == contract_id]
            curr_row = df_curr[df_curr[contract_id_col].astype(str) == contract_id]
            if len(prev_row) == 0 or len(curr_row) == 0:
                continue
            prev_amount = pd.to_numeric(prev_row[amount_col_prev], errors="coerce").sum()
            curr_amount = pd.to_numeric(curr_row[amount_col_curr], errors="coerce").sum()
            if abs(curr_amount - prev_amount) >= AMOUNT_TOLERANCE:
                project_name = ""
                for col in ["标准项目名称(JSON)", "项目名称", "标准项目名称"]:
                    if col in curr_row.columns:
                        project_name = curr_row.iloc[0][col]
                        break
                changed_list.append(
                    {
                        "合同编号": contract_id,
                        "项目名称": project_name,
                        f"{prev_month}合同额": prev_amount,
                        f"{curr_month}合同额": curr_amount,
                        "变化额": curr_amount - prev_amount,
                        "变化率(%)": (
                            (curr_amount - prev_amount) / prev_amount * 100 if prev_amount != 0 else 0
                        ),
                    }
                )
    return {
        "new": new_df,
        "deleted": deleted_df,
        "changed": pd.DataFrame(changed_list),
    }


def format_month_display(year_month: str) -> str:
    if len(year_month) == 6 and year_month.isdigit():
        return f"{year_month[:4]}年{year_month[4:]}月"
    return year_month


def _first_nonempty(row, candidates):
    for col in candidates:
        if col in row.index:
            val = row.get(col)
            if pd.notna(val):
                text = str(val).strip()
                if text and text.lower() != "nan":
                    return val
    return ""


def _extract_amount_from_row(row):
    for col in [
        "合同额（泰铢）",
        "分包合同额（泰铢）",
        "合同金额",
        "合同额",
        "金额（万泰铢）",
        "合同额（万泰铢）",
    ]:
        if col in row.index:
            val = pd.to_numeric(pd.Series([row.get(col)]), errors="coerce").iloc[0]
            if pd.notna(val):
                return float(val)
    return 0.0


def build_simple_change_rows(comparison, prev_month, curr_month):
    rows = []
    display_period = f"{format_month_display(prev_month)}→{format_month_display(curr_month)}"
    new_df = comparison.get("new", pd.DataFrame())
    if new_df is not None and len(new_df) > 0:
        for _, row in new_df.iterrows():
            rows.append(
                {
                    "对比月份": display_period,
                    "变化类型": "新增合同额",
                    "合同编号": _first_nonempty(row, ["合同编号"]),
                    "项目名称": _first_nonempty(
                        row, ["项目名称", "标准项目名称(JSON)", "标准项目名称"]
                    ),
                    "上月合同额": 0.0,
                    "本月合同额": _extract_amount_from_row(row),
                    "变化额": _extract_amount_from_row(row),
                }
            )
    changed_df = comparison.get("changed", pd.DataFrame())
    if changed_df is not None and len(changed_df) > 0:
        prev_col = f"{prev_month}合同额"
        curr_col = f"{curr_month}合同额"
        for _, row in changed_df.iterrows():
            prev_amount = float(
                pd.to_numeric(pd.Series([row.get(prev_col, 0)]), errors="coerce").fillna(0).iloc[0]
            )
            curr_amount = float(
                pd.to_numeric(pd.Series([row.get(curr_col, 0)]), errors="coerce").fillna(0).iloc[0]
            )
            diff_amount = float(
                pd.to_numeric(
                    pd.Series([row.get("变化额", curr_amount - prev_amount)]), errors="coerce"
                )
                .fillna(curr_amount - prev_amount)
                .iloc[0]
            )
            rows.append(
                {
                    "对比月份": display_period,
                    "变化类型": "合同额变化",
                    "合同编号": row.get("合同编号", ""),
                    "项目名称": row.get("项目名称", ""),
                    "上月合同额": prev_amount,
                    "本月合同额": curr_amount,
                    "变化额": diff_amount,
                }
            )
    if not rows:
        return pd.DataFrame(
            columns=["对比月份", "变化类型", "合同编号", "项目名称", "上月合同额", "本月合同额", "变化额"]
        )
    return pd.DataFrame(rows)


def build_project_subcontractor_summary(files_data):
    records = []
    for year_month, _, _, df in files_data:
        if df is None or len(df) == 0:
            continue
        project_col = next(
            (c for c in ["项目名称", "标准项目名称(JSON)", "标准项目名称"] if c in df.columns), None
        )
        unit_col = next(
            (c for c in ["协作体名称", "对方公司名称", "对方单位公司名称"] if c in df.columns), None
        )
        amount_col = next(
            (
                c
                for c in [
                    "合同额（泰铢）",
                    "分包合同额（泰铢）",
                    "合同金额",
                    "合同额",
                    "金额（万泰铢）",
                    "合同额（万泰铢）",
                ]
                if c in df.columns
            ),
            None,
        )
        if not project_col or not unit_col or not amount_col:
            continue
        tmp = df[[project_col, unit_col, amount_col]].copy()
        tmp = tmp.dropna(subset=[project_col, unit_col])
        tmp[project_col] = tmp[project_col].astype(str).str.strip()
        tmp[unit_col] = tmp[unit_col].astype(str).str.strip()
        tmp = tmp[(tmp[project_col] != "") & (tmp[project_col].str.lower() != "nan")]
        tmp = tmp[(tmp[unit_col] != "") & (tmp[unit_col].str.lower() != "nan")]
        tmp[amount_col] = pd.to_numeric(tmp[amount_col], errors="coerce").fillna(0)
        tmp = tmp[tmp[amount_col] != 0]
        for _, row in tmp.iterrows():
            records.append({"项目名称": row[project_col], "分包单位": row[unit_col]})
    if not records:
        return pd.DataFrame(columns=["项目名称", "分包数量"])
    result_df = pd.DataFrame(records)
    summary = result_df.drop_duplicates(subset=["项目名称", "分包单位"])
    summary = summary.groupby("项目名称")["分包单位"].nunique().reset_index(name="分包数量")
    return summary.sort_values(by=["分包数量", "项目名称"], ascending=[False, True])


def build_project_month_amount_summary(files_data):
    records = []
    for year_month, _, _, df in files_data:
        if df is None or len(df) == 0:
            continue
        project_col = next(
            (c for c in ["项目名称", "标准项目名称(JSON)", "标准项目名称"] if c in df.columns), None
        )
        amount_col = next(
            (
                c
                for c in [
                    "合同额（泰铢）",
                    "分包合同额（泰铢）",
                    "合同金额",
                    "合同额",
                    "金额（万泰铢）",
                    "合同额（万泰铢）",
                ]
                if c in df.columns
            ),
            None,
        )
        code_col = next((c for c in ["项目编码", "项目编码(匹配)"] if c in df.columns), None)
        if project_col is None or amount_col is None:
            continue
        tmp = df[[project_col, amount_col]].copy()
        if code_col:
            tmp[code_col] = df[code_col]
        tmp[project_col] = tmp[project_col].astype(str).str.strip()
        tmp = tmp[(tmp[project_col] != "") & (tmp[project_col].str.lower() != "nan")]
        tmp[amount_col] = pd.to_numeric(tmp[amount_col], errors="coerce").fillna(0)
        tmp["年月"] = year_month
        tmp = tmp.rename(columns={project_col: "项目名称", amount_col: "分包签订合同额汇总"})
        if code_col:
            tmp = tmp.rename(columns={code_col: "项目编码"})
            records.append(tmp[["项目编码", "项目名称", "年月", "分包签订合同额汇总"]])
        else:
            records.append(tmp[["项目名称", "年月", "分包签订合同额汇总"]])
    if not records:
        return pd.DataFrame(columns=["项目编码", "项目名称"])
    all_df = pd.concat(records, ignore_index=True)
    if "项目编码" in all_df.columns:
        grouped = all_df.groupby(["项目编码", "项目名称", "年月"], as_index=False)[
            "分包签订合同额汇总"
        ].sum()
        pivot = grouped.pivot_table(
            index=["项目编码", "项目名称"], columns="年月", values="分包签订合同额汇总", fill_value=0
        )
    else:
        grouped = all_df.groupby(["项目名称", "年月"], as_index=False)["分包签订合同额汇总"].sum()
        pivot = grouped.pivot(index="项目名称", columns="年月", values="分包签订合同额汇总").fillna(0)
    sorted_cols = sorted(list(pivot.columns))
    pivot.columns = [
        format_month_display(col) if len(str(col)) == 6 and str(col).isdigit() else col
        for col in sorted_cols
    ]
    return pivot.reset_index()


def build_project_month_increment_summary(files_data):
    records = []
    for year_month, _, _, df in files_data:
        if df is None or len(df) == 0:
            continue
        project_col = next(
            (c for c in ["项目名称", "标准项目名称(JSON)", "标准项目名称"] if c in df.columns), None
        )
        amount_col = next(
            (
                c
                for c in [
                    "合同额（泰铢）",
                    "分包合同额（泰铢）",
                    "合同金额",
                    "合同额",
                    "金额（万泰铢）",
                    "合同额（万泰铢）",
                ]
                if c in df.columns
            ),
            None,
        )
        code_col = next((c for c in ["项目编码", "项目编码(匹配)"] if c in df.columns), None)
        if project_col is None or amount_col is None:
            continue
        tmp = df[[project_col, amount_col]].copy()
        if code_col:
            tmp[code_col] = df[code_col]
        tmp[project_col] = tmp[project_col].astype(str).str.strip()
        tmp = tmp[(tmp[project_col] != "") & (tmp[project_col].str.lower() != "nan")]
        tmp[amount_col] = pd.to_numeric(tmp[amount_col], errors="coerce").fillna(0)
        tmp["年月"] = year_month
        tmp = tmp.rename(columns={project_col: "项目名称", amount_col: "分包签订合同额汇总"})
        if code_col:
            tmp = tmp.rename(columns={code_col: "项目编码"})
            records.append(tmp[["项目编码", "项目名称", "年月", "分包签订合同额汇总"]])
        else:
            records.append(tmp[["项目名称", "年月", "分包签订合同额汇总"]])
    if not records:
        return pd.DataFrame(columns=["项目编码", "项目名称"])
    all_df = pd.concat(records, ignore_index=True)
    if "项目编码" in all_df.columns:
        grouped = all_df.groupby(["项目编码", "项目名称", "年月"], as_index=False)[
            "分包签订合同额汇总"
        ].sum()
        pivot = grouped.pivot_table(
            index=["项目编码", "项目名称"], columns="年月", values="分包签订合同额汇总", fill_value=0
        )
    else:
        grouped = all_df.groupby(["项目名称", "年月"], as_index=False)["分包签订合同额汇总"].sum()
        pivot = grouped.pivot(index="项目名称", columns="年月", values="分包签订合同额汇总").fillna(0)
    sorted_cols = sorted(list(pivot.columns))
    increment = pd.DataFrame(index=pivot.index)
    for i, col in enumerate(sorted_cols):
        display_col = format_month_display(col) if len(str(col)) == 6 and str(col).isdigit() else col
        if i == 0:
            increment[display_col] = None
        else:
            increment[display_col] = pivot[col].values - pivot[sorted_cols[i - 1]].values
    return increment.reset_index()


def build_project_contract_total_summary(files_data):
    if not files_data:
        return pd.DataFrame(columns=["项目编码", "项目名称", "分包合同汇总金额"])
    latest_month, _, _, df = files_data[-1]
    if df is None or len(df) == 0:
        return pd.DataFrame(columns=["项目编码", "项目名称", "分包合同汇总金额"])
    project_col = next(
        (c for c in ["项目名称", "标准项目名称(JSON)", "标准项目名称"] if c in df.columns), None
    )
    amount_col = next(
        (
            c
            for c in [
                "分包合同额（泰铢）",
                "合同额（泰铢）",
                "合同金额",
                "合同额",
                "金额（万泰铢）",
                "合同额（万泰铢）",
            ]
            if c in df.columns
        ),
        None,
    )
    code_col = next((c for c in ["项目编码", "项目编码(匹配)"] if c in df.columns), None)
    if project_col is None or amount_col is None:
        return pd.DataFrame(columns=["项目编码", "项目名称", "分包合同汇总金额"])
    use_cols = [project_col, amount_col] + ([code_col] if code_col else [])
    tmp = df[use_cols].copy()
    tmp[project_col] = tmp[project_col].astype(str).str.strip()
    tmp = tmp[(tmp[project_col] != "") & (tmp[project_col].str.lower() != "nan")]
    tmp[amount_col] = pd.to_numeric(tmp[amount_col], errors="coerce").fillna(0)
    if code_col:
        grouped = tmp.groupby([code_col, project_col], as_index=False)[amount_col].sum()
        grouped = grouped.rename(
            columns={code_col: "项目编码", project_col: "项目名称", amount_col: "分包合同汇总金额"}
        )
    else:
        grouped = tmp.groupby([project_col], as_index=False)[amount_col].sum()
        grouped = grouped.rename(columns={project_col: "项目名称", amount_col: "分包合同汇总金额"})
        grouped.insert(0, "项目编码", "")
    return grouped.sort_values(by="分包合同汇总金额", ascending=False).reset_index(drop=True)


def build_comparison_sheets(files_data: list | None = None, data_dir=None) -> dict[str, pd.DataFrame]:
    """构建对比相关全部 Sheet（全量重算合同变化，避免追加脏数据）。"""
    if files_data is None:
        files_data = scan_contract_files(data_dir)

    empty_change = pd.DataFrame(
        columns=["对比月份", "变化类型", "合同编号", "项目名称", "上月合同额", "本月合同额", "变化额"]
    )
    if len(files_data) < 1:
        return {
            "项目分包合同汇总": pd.DataFrame(columns=["项目编码", "项目名称", "分包合同汇总金额"]),
            "项目年月合同额汇总": pd.DataFrame(columns=["项目编码", "项目名称"]),
            "签订金额增量对比": pd.DataFrame(columns=["项目编码", "项目名称"]),
            "项目分包单位统计": pd.DataFrame(columns=["项目名称", "分包数量"]),
            "合同变化汇总": empty_change,
        }

    all_change_parts = []
    if len(files_data) >= 2:
        print("\n开始逐月对比...")
        for i in range(len(files_data) - 1):
            prev_month, _, _, df_prev = files_data[i]
            curr_month, _, _, df_curr = files_data[i + 1]
            print(f"  对比: {format_month_display(prev_month)} vs {format_month_display(curr_month)}")
            comparison = compare_two_months(df_prev, df_curr, prev_month, curr_month)
            simple_change_df = build_simple_change_rows(comparison, prev_month, curr_month)
            if len(simple_change_df) > 0:
                all_change_parts.append(simple_change_df)
                print(f"    ✓ 变化记录: {len(simple_change_df)} 条")

    print("\n生成对比汇总表...")
    sheets = {
        "项目分包合同汇总": build_project_contract_total_summary(files_data),
        "项目年月合同额汇总": build_project_month_amount_summary(files_data),
        "签订金额增量对比": build_project_month_increment_summary(files_data),
        "项目分包单位统计": build_project_subcontractor_summary(files_data),
        "合同变化汇总": (
            pd.concat(all_change_parts, ignore_index=True) if all_change_parts else empty_change
        ),
    }
    return sheets


def format_comparison_sheets(workbook, sheet_names: list[str] | None = None) -> None:
    targets = sheet_names or list(workbook.sheetnames)
    for name in targets:
        if name not in workbook.sheetnames:
            continue
        # 付款汇总 sheet 由 step4 自己格式化
        if name.startswith("项目付款汇总"):
            continue
        apply_common_excel_layout(workbook[name], numeric_start_col=3)


def export_comparison_only(data_dir=None, output_file=None) -> Path:
    SUM_DIR.mkdir(parents=True, exist_ok=True)
    output_file = Path(output_file or (SUM_DIR / COMPARISON_ONLY_NAME))
    sheets = build_comparison_sheets(data_dir=data_dir)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
    wb = load_workbook(output_file)
    format_comparison_sheets(wb)
    wb.save(output_file)
    wb.close()
    print(f"\n✓ 对比汇总完成: {output_file}")
    return output_file


def main() -> None:
    print("=" * 80)
    print("合同汇总自动对比系统")
    print("=" * 80)
    export_comparison_only()


if __name__ == "__main__":
    main()
