# -*- coding: utf-8 -*-
"""
步骤4：按项目汇总分包合同额与已付款。
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from config import STD_DIR, SUM_DIR


def pick_column(columns, candidates):
    col_map = {str(c).strip().replace(" ", ""): c for c in columns}
    for candidate in candidates:
        key = str(candidate).strip().replace(" ", "")
        if key in col_map:
            return col_map[key]
    return None


def find_latest_excel(input_dir: Path | None = None) -> Path:
    input_dir = Path(input_dir or STD_DIR)
    excel_files = [
        p
        for p in input_dir.glob("*.xlsx")
        if not p.name.startswith("~$") and "钢结构" not in p.name
    ]
    if not excel_files:
        raise FileNotFoundError(f"未在 {input_dir} 下找到 Excel 文件")
    return max(excel_files, key=lambda p: p.stat().st_mtime)


def build_paid_summary_frames(input_file: Path | str) -> dict[str, pd.DataFrame]:
    """返回 {'项目付款汇总_泰铢': df, '项目付款汇总_万泰铢': df}。"""
    input_file = Path(input_file)
    df = pd.read_excel(input_file)

    project_code_col = pick_column(df.columns, ["项目编码", "项目代码", "Project Code"])
    project_name_col = pick_column(df.columns, ["项目名称", "Project Name"])
    amount_col = pick_column(
        df.columns,
        [
            "合同额（泰铢）",
            "合同额(泰铢)",
            "合同额",
            "分包合同额（泰铢）",
            "分包合同额(泰铢)",
            "合同金额",
        ],
    )
    paid_col = pick_column(df.columns, ["分包已付款", "公司实付含税", "已付款", "实付含税"])

    if not project_code_col or not project_name_col or not amount_col or not paid_col:
        raise ValueError(
            f"无法识别所需列。项目编码={project_code_col}, 项目名称={project_name_col}, "
            f"合同额={amount_col}, 分包已付款={paid_col}"
        )

    work_df = df[[project_code_col, project_name_col, amount_col, paid_col]].copy()
    work_df.columns = ["项目编码", "项目名称", "合同额（泰铢）", "分包已付款（泰铢）"]

    for col in ["合同额（泰铢）", "分包已付款（泰铢）"]:
        work_df[col] = pd.to_numeric(work_df[col], errors="coerce").fillna(0)

    work_df["项目编码"] = work_df["项目编码"].fillna("").astype(str).str.strip()
    work_df["项目名称"] = work_df["项目名称"].fillna("").astype(str).str.strip()
    work_df = work_df[(work_df["项目编码"] != "") | (work_df["项目名称"] != "")]

    summary = (
        work_df.groupby(["项目编码", "项目名称"], as_index=False)
        .agg({"合同额（泰铢）": "sum", "分包已付款（泰铢）": "sum"})
        .sort_values(by="合同额（泰铢）", ascending=False)
        .reset_index(drop=True)
    )

    summary["分包付款比例"] = np.where(
        summary["合同额（泰铢）"] > 0,
        summary["分包已付款（泰铢）"] / summary["合同额（泰铢）"],
        0,
    )
    summary["分包付款比例"] = summary["分包付款比例"].fillna(0)

    summary_thb = summary[
        ["项目编码", "项目名称", "合同额（泰铢）", "分包已付款（泰铢）", "分包付款比例"]
    ].copy()

    summary_wan = summary_thb.copy()
    summary_wan = summary_wan.rename(
        columns={
            "合同额（泰铢）": "合同额（万泰铢）",
            "分包已付款（泰铢）": "分包已付款（万泰铢）",
        }
    )
    summary_wan["合同额（万泰铢）"] = summary_wan["合同额（万泰铢）"] / 10000
    summary_wan["分包已付款（万泰铢）"] = summary_wan["分包已付款（万泰铢）"] / 10000

    return {
        "项目付款汇总_泰铢": summary_thb,
        "项目付款汇总_万泰铢": summary_wan,
    }


def format_paid_sheets(workbook, sheet_names: list[str] | None = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    targets = sheet_names or [
        s for s in workbook.sheetnames if s.startswith("项目付款汇总")
    ]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for name in targets:
        if name not in workbook.sheetnames:
            continue
        ws = workbook[name]
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for column in ws.columns:
            col_cells = list(column)
            max_length = 0
            for cell in col_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            letter = col_cells[0].column_letter
            ws.column_dimensions[letter].width = min(max_length + 2, 40)

        # C/D 金额，E 比例
        for row in range(2, ws.max_row + 1):
            ws[f"C{row}"].number_format = "#,##0.00"
            ws[f"D{row}"].number_format = "#,##0.00"
            ws[f"E{row}"].number_format = "0.00%"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions


def export_paid_only(
    input_file: Path | str | None = None,
    output_file: Path | str | None = None,
) -> Path:
    input_file = Path(input_file) if input_file else find_latest_excel(STD_DIR)
    SUM_DIR.mkdir(parents=True, exist_ok=True)
    if output_file is None:
        output_file = SUM_DIR / f"{input_file.stem}{Path('_项目汇总.xlsx').name.replace('_项目汇总.xlsx', '')}_项目汇总.xlsx"
        # clearer
        output_file = SUM_DIR / f"{input_file.stem}_项目汇总.xlsx"
    else:
        output_file = Path(output_file)

    sheets = build_paid_summary_frames(input_file)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for name, df in sheets.items():
            # 单独导出时用短名
            short = "泰铢" if "泰铢" in name and "万" not in name else "万泰铢"
            if name.endswith("万泰铢"):
                short = "万泰铢"
            elif name.endswith("泰铢"):
                short = "泰铢"
            df.to_excel(writer, sheet_name=short, index=False)

    from openpyxl import load_workbook

    wb = load_workbook(output_file)
    # 映射短名格式化
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = __import__("openpyxl.styles", fromlist=["Font"]).Font(bold=True)
        for row in range(2, ws.max_row + 1):
            ws[f"C{row}"].number_format = "#,##0.00"
            ws[f"D{row}"].number_format = "#,##0.00"
            ws[f"E{row}"].number_format = "0.00%"
        for column in ws.columns:
            col_cells = list(column)
            max_length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_length + 2, 40)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(output_file)
    wb.close()
    print(f"已生成汇总文件: {output_file}")
    return output_file


def main() -> None:
    export_paid_only()


if __name__ == "__main__":
    main()
