from __future__ import annotations

import argparse
import math
import re
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
SUMMARY_INPUT_ROOT = SCRIPT_DIR / "B_汇总数据"
ANALYSIS_OUTPUT_ROOT = SCRIPT_DIR / "C_数据分析"
INPUT_GLOB = "*_材料采购_merged.xlsx"
OUTPUT_SUFFIX = "_钢筋规格标准化"

TARGET_KEYWORDS = ("钢筋网片", "螺纹钢", "钢筋")
MESH_KEYWORDS = ("钢筋网片", "网片")
REBAR_KEYWORDS = ("螺纹钢",)
# 名称含「钢筋」但不是棒材本身的物品，跳过规格标准化
SKIP_NAME_KEYWORDS = (
    "套筒", "弯曲机", "切断机", "调直", "加工棚", "工棚",
    "检查井", "方桩", "管桩", "扎钩", "网片",
)

PHI_PATTERN = re.compile(r"[φΦØø]\s*(\d+(?:\.\d+)?)")
# D16 / d16 直径写法（避免把 SD40 / DB40 里的 D 误吃掉，要求 D 前非字母数字或在边界）
D_DIA_PATTERN = re.compile(r"(?<![A-Za-z0-9])[Dd]\s*(\d{1,2}(?:\.\d+)?)(?!\d)")
HRB_DIA_PATTERN = re.compile(
    r"(hrb|hpb|rrb)\s*[- ]?\s*(\d{3})([a-zA-Z]*)\s*[- ]?\s*(\d{1,2}(?:\.\d+)?)",
    re.IGNORECASE,
)
GRADE_PATTERN = re.compile(r"(hrb|hpb|rrb)\s*[- ]?\s*(\d{3})([a-zA-Z]*)", re.IGNORECASE)
# 泰国螺纹钢 / 圆钢牌号
#   SD40 / SD40T / SD50
#   DB40T / DB40 / DB50T  （Deformed Bar，台账常见 DB40T10*12M）
#   SR24 / SR30           （Smooth Round，台账常见 SR24 6*10M）
# 强度位固定两位（40/50/24…），避免 SR246x10m 把 SR246 整段当成牌号
SD_GRADE_PATTERN = re.compile(r"(SD\s*\d{2}\s*T?)", re.IGNORECASE)
DB_GRADE_PATTERN = re.compile(r"(DB\s*\d{2}\s*T?)", re.IGNORECASE)
SR_GRADE_PATTERN = re.compile(r"(SR\s*\d{2})", re.IGNORECASE)
# 紧凑写法（牌号+直径+定尺）优先整体解析，避免把牌号里的 40 当成直径
#   DB40T10*12M → DB40T / 10 / 12
#   SD40T16x12  → SD40T / 16 / 12
#   SR24 6*10M → SR24 / 6 / 10  （去空格后 SR246x10m）
COMPACT_DB_SD_PATTERN = re.compile(
    r"(?P<grade>(?:DB|SD)\s*\d{2}\s*T?)"
    r"(?P<dia>\d{1,2}(?:\.\d+)?)"
    r"\s*[xX*×＊]\s*"
    r"(?P<len>\d+(?:\.\d+)?)\s*m?",
    re.IGNORECASE,
)
COMPACT_SR_PATTERN = re.compile(
    r"(?P<grade>SR\s*\d{2})"
    r"\s*"
    r"(?P<dia>\d{1,2}(?:\.\d+)?)"
    r"\s*[xX*×＊]\s*"
    r"(?P<len>\d+(?:\.\d+)?)\s*m?",
    re.IGNORECASE,
)
# 仅牌号+直径（无长度）：DB40T16、SR246（SR24+Φ6）
GRADE_DIA_ONLY_PATTERN = re.compile(
    r"(?P<grade>(?:DB|SD)\s*\d{2}\s*T?|SR\s*\d{2})"
    r"(?P<dia>\d{1,2}(?:\.\d+)?)(?!\d)",
    re.IGNORECASE,
)
MM_DIA_PATTERN = re.compile(r"(?<!\d)(\d{1,2}(?:\.\d+)?)\s*(?:mm|毫米)\b", re.IGNORECASE)
RAW_DIA_PATTERN = re.compile(r"(?<!\d)(\d{1,2}(?:\.\d+)?)(?!\d)")
# 显式带 m/米 的长度
LENGTH_M_PATTERN = re.compile(r"(?<!\d)(\d+(?:\.\d+)?)\s*m\b", re.IGNORECASE)
# Φ12x12 / φ16*12 / 16x12 （第二段为定尺长度，可无单位）
PHI_LEN_PATTERN = re.compile(
    r"[φΦØø]\s*(\d+(?:\.\d+)?)\s*[xX*]\s*(\d+(?:\.\d+)?)(?!\d)",
    re.IGNORECASE,
)
DIA_LEN_PATTERN = re.compile(
    r"(?<![φΦØøA-Za-z])(\d{1,2}(?:\.\d+)?)\s*(?:mm)?\s*[xX*]\s*(\d+(?:\.\d+)?)(?!\d)",
    re.IGNORECASE,
)
SPACING_PATTERN = re.compile(r"(?<!\d)(\d{2,3})\s*[xX×*＊/]\s*(\d{2,3})(?!\d)")
SHEET_SIZE_PATTERN = re.compile(r"(?<!\d)(\d+(?:\.\d+)?)\s*[xX×*＊]\s*(\d+(?:\.\d+)?)\s*(?:m|米)", re.IGNORECASE)

# 常见定尺（米）；也允许 6~18 的整数米
COMMON_REBAR_LENGTHS_M = {6, 8, 9, 10, 11, 12, 15}
# 常见棒材直径（mm），用于紧凑牌号解析校验
COMMON_REBAR_DIA_MM = {
    6, 8, 9, 10, 12, 13, 14, 16, 18, 20, 22, 25, 28, 32, 36, 40, 50,
}

def format_number(value: float) -> str:
    rounded = round(value)
    if math.isclose(value, rounded, abs_tol=1e-9):
        return str(int(rounded))
    return f"{value:.6f}".rstrip("0").rstrip(".")


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""

    text = text.replace("（", "(").replace("）", ")")
    text = text.replace("，", ",")
    text = text.replace("：", ":")
    text = text.replace("＋", "+")
    text = text.replace("Ｘ", "x").replace("×", "x").replace("＊", "x").replace("*", "x")
    text = text.replace("Ｍ", "m").replace("ｍ", "m")
    text = text.replace("毫米", "mm").replace("米", "m")
    # 数字后的 M/m 视为米单位（DB40T10*12M、SR24 6*10M）
    text = re.sub(r"(?<=\d)[Mm]\b", "m", text)
    text = re.sub(r"(?<=\d)[Mm](?=[xX*×＊\d]|$)", "m", text)
    text = re.sub(r"\s+", "", text)
    return text


def _is_rebar_length_m(val: float) -> bool:
    """判断是否像钢筋定尺长度（米）。"""
    if val in COMMON_REBAR_LENGTHS_M:
        return True
    # 允许 6~18 的整数米（如 7m、14m 偶发）
    if 6 <= val <= 18 and math.isclose(val, round(val), abs_tol=1e-9):
        return True
    return False


def _norm_thai_grade(raw: str) -> str:
    """统一泰国牌号空格与大小写：DB40T / SD40T / SR24。"""
    g = re.sub(r"\s+", "", (raw or "").upper())
    # DB40 / SD40 可选 T
    m = re.fullmatch(r"(DB|SD)(\d{2})(T?)", g)
    if m:
        return f"{m.group(1)}{m.group(2)}{m.group(3)}"
    m = re.fullmatch(r"(SR)(\d{2})", g)
    if m:
        return f"{m.group(1)}{m.group(2)}"
    return g


def parse_compact_thai_rebar(text: str) -> dict[str, float | str] | None:
    """
    解析泰国紧凑规格：
      DB40T10*12M  → grade=DB40T, dia=10, len=12
      DB40T16*12M  → grade=DB40T, dia=16, len=12
      SR24 6*10M   → grade=SR24,  dia=6,  len=10
      SD40T16x12   → grade=SD40T, dia=16, len=12
    """
    if not text:
        return None

    for pat in (COMPACT_DB_SD_PATTERN, COMPACT_SR_PATTERN):
        m = pat.search(text)
        if not m:
            continue
        grade = _norm_thai_grade(m.group("grade"))
        dia = float(m.group("dia"))
        length = float(m.group("len"))
        # 直径合理且定尺像米
        if not (4 <= dia <= 50):
            continue
        if dia not in COMMON_REBAR_DIA_MM and not (4 <= dia <= 50):
            continue
        if not _is_rebar_length_m(length) and not (1 <= length <= 18):
            # 仍允许 1~18m 显式米
            if not (1 <= length <= 18):
                continue
        return {"grade": grade, "diameter": dia, "length_m": length}

    # 无长度：DB40T16 / SR246（SR24+Φ6）
    m = GRADE_DIA_ONLY_PATTERN.search(text)
    if m:
        grade = _norm_thai_grade(m.group("grade"))
        dia = float(m.group("dia"))
        if 4 <= dia <= 50 and (dia in COMMON_REBAR_DIA_MM or 6 <= dia <= 40):
            # 避免把 DB40 的 40 当成直径：牌号已含数字时 dia 不能等于牌号强度位
            strength = re.search(r"(\d{2,3})", grade)
            if strength and math.isclose(dia, float(strength.group(1)), abs_tol=1e-6):
                return None
            return {"grade": grade, "diameter": dia, "length_m": None}
    return None


def find_diameter(text: str, *, prefer_mesh: bool = False) -> float | None:
    if not text:
        return None

    compact = parse_compact_thai_rebar(text)
    if compact is not None:
        return float(compact["diameter"])  # type: ignore[arg-type]

    match = PHI_PATTERN.search(text)
    if match:
        return float(match.group(1))

    match = D_DIA_PATTERN.search(text)
    if match:
        val = float(match.group(1))
        low, high = (2, 20) if prefer_mesh else (4, 50)
        if low <= val <= high:
            return val

    match = HRB_DIA_PATTERN.search(text)
    if match:
        return float(match.group(4))

    match = MM_DIA_PATTERN.search(text)
    if match:
        return float(match.group(1))

    # 去掉牌号数字后再扫裸数字，避免 SD40T16x12 / DB40T16x12 把 40 当成直径
    stripped = SD_GRADE_PATTERN.sub(" ", text)
    stripped = DB_GRADE_PATTERN.sub(" ", stripped)
    stripped = SR_GRADE_PATTERN.sub(" ", stripped)
    stripped = GRADE_PATTERN.sub(" ", stripped)

    # 钢筋直径通常在 4~50mm，网片钢丝常见 3~12mm
    low, high = (2, 20) if prefer_mesh else (4, 50)
    for m in RAW_DIA_PATTERN.finditer(stripped):
        val = float(m.group(1))
        if low <= val <= high:
            return val
    return None


def find_grade(text: str) -> str | None:
    compact = parse_compact_thai_rebar(text)
    if compact is not None and compact.get("grade"):
        return str(compact["grade"])

    match = GRADE_PATTERN.search(text)
    if match:
        prefix = match.group(1).upper()
        level = match.group(2)
        suffix = match.group(3).upper() if match.group(3) else ""
        return f"{prefix}{level}{suffix}"

    for pat in (DB_GRADE_PATTERN, SD_GRADE_PATTERN, SR_GRADE_PATTERN):
        match = pat.search(text)
        if match:
            return _norm_thai_grade(match.group(1))
    return None


def find_length_m(text: str, diameter: float | None = None) -> float | None:
    """
    解析钢筋长度（米）。

    支持：
      - 显式单位：φ12x12m、12米、12m、12M
      - 台账常见无单位写法：φ12*12、φ16x12、16*12（第二段为定尺米数）
      - 泰国紧凑：DB40T10*12M、SR24 6*10M
    """
    if not text:
        return None

    compact = parse_compact_thai_rebar(text)
    if compact is not None and compact.get("length_m") is not None:
        return float(compact["length_m"])  # type: ignore[arg-type]

    # 1) 显式 m（normalize 后 米→m）
    for m in LENGTH_M_PATTERN.finditer(text):
        val = float(m.group(1))
        if 1 <= val <= 18:
            return val

    # 2) Φd x L（无单位）：φ12x12 → 长度 12m
    m = PHI_LEN_PATTERN.search(text)
    if m:
        length = float(m.group(2))
        if _is_rebar_length_m(length):
            return length

    # 3) d x L（无 Φ，可带 mm）：16x12、16mmx12
    m = DIA_LEN_PATTERN.search(text)
    if m:
        dia = float(m.group(1))
        length = float(m.group(2))
        if diameter is not None and not math.isclose(dia, diameter, abs_tol=1e-6):
            # 与已识别直径不一致时仍可接受常见定尺
            pass
        if 4 <= dia <= 50 and _is_rebar_length_m(length):
            return length

    return None


def standardize_rebar_or_steel(spec_value: object) -> str | None:
    text = normalize_text(spec_value)
    if not text:
        return None

    # 优先泰国紧凑牌号（DB40T10*12M / SR24 6*10M）
    compact = parse_compact_thai_rebar(text)
    if compact is not None:
        diameter = float(compact["diameter"])  # type: ignore[arg-type]
        grade = str(compact["grade"]) if compact.get("grade") else None
        length = (
            float(compact["length_m"])  # type: ignore[arg-type]
            if compact.get("length_m") is not None
            else None
        )
    else:
        diameter = find_diameter(text, prefer_mesh=False)
        if diameter is None:
            return None
        grade = find_grade(text)
        length = find_length_m(text, diameter=diameter)

    base = f"Φ{format_number(diameter)}"
    if length is not None:
        base = f"{base}x{format_number(length)}m"
    if grade:
        return f"{grade}-{base}"
    return base


def find_spacing(text: str) -> tuple[int, int] | None:
    for m in SPACING_PATTERN.finditer(text):
        a = int(m.group(1))
        b = int(m.group(2))
        # 网片间距一般是 50~400mm
        if 50 <= a <= 400 and 50 <= b <= 400:
            return a, b
    return None


def find_sheet_size(text: str) -> tuple[float, float] | None:
    for m in SHEET_SIZE_PATTERN.finditer(text):
        a = float(m.group(1))
        b = float(m.group(2))
        # 网片尺寸一般在 0.5~12m
        if 0.5 <= a <= 12 and 0.5 <= b <= 12:
            return a, b
    return None


def standardize_mesh(spec_value: object) -> str | None:
    """网片规格委托给专用模块，输出「间距 型号 大小」，如 200x200 Φ8 3x6m。"""
    try:
        from pathlib import Path
        import importlib.util

        mesh_path = Path(__file__).resolve().parent / "钢筋网片规格标准化.py"
        if mesh_path.exists():
            spec = importlib.util.spec_from_file_location("_mesh_spec_std", str(mesh_path))
            if spec is not None and spec.loader is not None:
                mod = importlib.util.module_from_spec(spec)
                spec.loader.exec_module(mod)
                return mod.standardize_mesh(spec_value)
    except Exception:
        pass

    # 回退：旧逻辑（尽量给出直径）
    text = normalize_text(spec_value)
    if not text:
        return None
    diameter = find_diameter(text, prefer_mesh=True)
    if diameter is None:
        return None
    return f"Φ{format_number(diameter)}"


def looks_like_thai_rebar_code(text: object) -> bool:
    """规格/名称是否像泰国钢筋编码（DB40T… / SR24… / SD40T…）。"""
    t = normalize_text(text)
    if not t:
        return False
    if parse_compact_thai_rebar(t) is not None:
        return True
    if DB_GRADE_PATTERN.search(t) or SR_GRADE_PATTERN.search(t) or SD_GRADE_PATTERN.search(t):
        return True
    return False


def detect_category(item_name: object, spec_value: object = None) -> str | None:
    if item_name is None and spec_value is None:
        return None

    item_text = str(item_name).strip() if item_name is not None else ""

    # 网片单独处理
    if item_text and any(keyword in item_text for keyword in MESH_KEYWORDS):
        return "mesh"

    # 排除套筒/设备/桩等，避免误改规格
    if item_text and any(k in item_text for k in SKIP_NAME_KEYWORDS):
        return None

    if item_text and any(keyword in item_text for keyword in REBAR_KEYWORDS):
        return "rebar"
    if item_text and "钢筋" in item_text:
        return "steel"
    # 物品名称仅为编码或规格列是 DB40T/SR24 等
    if looks_like_thai_rebar_code(item_text) or looks_like_thai_rebar_code(spec_value):
        return "rebar"
    return None


def get_header_map(worksheet):
    return {
        str(cell.value).strip(): index
        for index, cell in enumerate(worksheet[1], start=1)
        if cell.value is not None and str(cell.value).strip()
    }


def process_worksheet(worksheet):
    header_map = get_header_map(worksheet)
    if "物品名称" not in header_map or "规格" not in header_map:
        return {
            "sheet_name": worksheet.title,
            "matched_rows": 0,
            "updated_rows": 0,
            "unresolved_rows": [],
            "skipped": True,
        }

    item_col = header_map["物品名称"]
    spec_col = header_map["规格"]

    matched_rows = 0
    updated_rows = 0
    unresolved_rows = []

    for row_idx in range(2, worksheet.max_row + 1):
        item_name = worksheet.cell(row_idx, item_col).value
        spec_cell = worksheet.cell(row_idx, spec_col)
        category = detect_category(item_name, spec_cell.value)
        if category is None:
            continue

        matched_rows += 1

        if category == "mesh":
            standardized = standardize_mesh(spec_cell.value)
        else:
            # 规格为空时，尝试用物品名称里的泰国编码
            raw_spec = spec_cell.value
            standardized = standardize_rebar_or_steel(raw_spec)
            if standardized is None and looks_like_thai_rebar_code(item_name):
                standardized = standardize_rebar_or_steel(item_name)

        if standardized is None:
            unresolved_rows.append(
                (
                    row_idx,
                    str(item_name).strip(),
                    "" if spec_cell.value is None else str(spec_cell.value).strip(),
                    category,
                )
            )
            continue

        old_text = "" if spec_cell.value is None else str(spec_cell.value).strip()
        if old_text != standardized:
            spec_cell.value = standardized
            updated_rows += 1

    return {
        "sheet_name": worksheet.title,
        "matched_rows": matched_rows,
        "updated_rows": updated_rows,
        "unresolved_rows": unresolved_rows,
        "skipped": False,
    }


def build_output_path(input_file: Path, suffix: str) -> Path:
    ANALYSIS_OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    return ANALYSIS_OUTPUT_ROOT / f"{input_file.stem}{suffix}{input_file.suffix}"


def extract_year_month_from_name(path: Path) -> tuple[int, int] | None:
    m = re.search(r"(?<!\d)((?:19|20)\d{2})年(0?[1-9]|1[0-2])月", path.stem)
    if not m:
        return None
    return int(m.group(1)), int(m.group(2))


def find_latest_summary_file() -> Path | None:
    if not SUMMARY_INPUT_ROOT.is_dir():
        return None
    candidates = [p for p in SUMMARY_INPUT_ROOT.glob(INPUT_GLOB) if p.is_file()]
    if not candidates:
        return None

    def sort_key(p: Path):
        ym = extract_year_month_from_name(p)
        if ym is not None:
            return (1, ym[0], ym[1], p.stat().st_mtime)
        return (0, 0, 0, p.stat().st_mtime)

    return max(candidates, key=sort_key)


def parse_args():
    parser = argparse.ArgumentParser(description="标准化采购台账中钢筋/螺纹钢/钢筋网片的规格格式")
    parser.add_argument(
        "input_file",
        nargs="?",
        default="",
        help="待处理的 Excel 文件路径；不传时自动读取 B_汇总数据 下最新 *_材料采购_merged.xlsx",
    )
    parser.add_argument(
        "--suffix",
        default=OUTPUT_SUFFIX,
        help="输出文件名后缀，默认追加到原文件名后",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    if args.input_file:
        input_file = Path(args.input_file).expanduser().resolve()
    else:
        latest = find_latest_summary_file()
        if latest is None:
            raise FileNotFoundError(f"未在 {SUMMARY_INPUT_ROOT} 找到 {INPUT_GLOB}")
        input_file = latest.resolve()

    output_file = build_output_path(input_file, args.suffix)

    if not input_file.exists():
        raise FileNotFoundError(f"未找到输入文件: {input_file}")

    print(f"正在处理文件: {input_file}")
    workbook = load_workbook(input_file)

    all_results = []
    total_matched = 0
    total_updated = 0
    total_unresolved = 0

    for worksheet in workbook.worksheets:
        result = process_worksheet(worksheet)
        all_results.append(result)

        if result["skipped"]:
            print(f"工作表 {worksheet.title}: 未找到“物品名称”或“规格”列，跳过")
            continue

        total_matched += result["matched_rows"]
        total_updated += result["updated_rows"]
        total_unresolved += len(result["unresolved_rows"])

        print(
            f"工作表 {worksheet.title}: 命中 {result['matched_rows']} 行，"
            f"成功标准化 {result['updated_rows']} 行，"
            f"未识别 {len(result['unresolved_rows'])} 行"
        )

    workbook.save(output_file)

    print("\n处理完成")
    print(f"输出文件: {output_file}")
    print(f"命中总行数: {total_matched}")
    print(f"成功标准化: {total_updated}")
    print(f"未识别规格: {total_unresolved}")

    unresolved_preview = []
    for result in all_results:
        unresolved_preview.extend(
            (result["sheet_name"], row_idx, item_name, spec_value, category)
            for row_idx, item_name, spec_value, category in result["unresolved_rows"]
        )

    if unresolved_preview:
        print("\n以下是前 20 条未识别规格样本:")
        for sheet_name, row_idx, item_name, spec_value, category in unresolved_preview[:20]:
            print(
                f"  - 工作表={sheet_name}, 行号={row_idx}, 类别={category}, 物品名称={item_name}, 原规格={spec_value}"
            )


if __name__ == "__main__":
    main()
