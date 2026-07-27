#!/usr/bin/env python3
"""
各项目采购分类汇总（独立程序）。

按项目汇总：
  - 钢筋：数量(kg) + 金额
  - 砂浆/水泥：数量(kg) + 金额
  - 砖/砌块：数量(m³) + 金额
  - 模板（含木方）：数量(m²，张/根按规格换算) + 金额
  - 碎石：数量(t) + 金额
  - 其他：数量 + 金额

数量列使用整数数字格式（无小数位）；金额保留两位小数。

数据源优先：对比版 → std_merged → merged
输出：SUM/项目采购分类汇总.xlsx
"""
from __future__ import annotations

import argparse
import importlib.util
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_PKG_ROOT = Path(__file__).resolve().parents[1]
_CODE_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(_PKG_ROOT), str(_CODE_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from _采购路径 import SUM_DIR, UNIT_PIPELINE

SCRIPT_DIR = Path(__file__).resolve().parent
UNIT_CONV_PIPELINE = UNIT_PIPELINE

PREFERRED_INPUTS = [
    "采购部采购台账_std_merged_对比.xlsx",
    "采购部采购台账_std_merged.xlsx",
    "采购部采购台账_merged.xlsx",
]
DEFAULT_OUTPUT = "项目采购分类汇总.xlsx"

# 分类定义：(分类名, 数量列名, 金额列名, 数量单位说明)
CLASS_SPECS = [
    ("钢筋", "钢筋_kg", "钢筋_金额", "kg"),
    ("砂浆水泥", "砂浆水泥_kg", "砂浆水泥_金额", "kg"),
    ("砖", "砖_m3", "砖_金额", "m3"),
    ("模板", "模板_m2", "模板_金额", "m2"),  # 含木方，张/根→m2
    ("碎石", "碎石_t", "碎石_金额", "t"),
    ("其他", "其他_数量", "其他_金额", "原单位"),
]

REBAR_NAME = ("钢筋", "螺纹钢", "网片")
REBAR_NAME_EXCLUDE = ("套筒", "弯曲机", "加工棚", "工棚", "扎钩", "运费")
MORTAR_NAME = ("砂浆", "沙浆", "砂桨", "水泥", "干混")
MORTAR_NAME_EXCLUDE = ("砖", "砌块")
BRICK_NAME = (
    "空心砖", "实心砖", "加气块", "加气砖", "水泥砖", "红砖", "砌块",
    "多孔砖", "页岩砖", "标砖", "通风砖", "植草砖", "面包砖", "墙砖",
)
TEMPLATE_NAME = ("模板", "模版", "木模板", "红模板", "小模板")
WOOD_NAME = ("木方", "方木", "杂木方", "木方条")
WOOD_NAME_EXCLUDE = ("模板", "模版", "机", "开孔", "轮", "尾款", "五金")
GRAVEL_NAME = ("碎石", "石子", "卵石", "瓜子片", "石屑", "级配碎石", "毛石")


def _load_unit_conv():
    if not UNIT_CONV_PIPELINE.exists():
        return None
    spec = importlib.util.spec_from_file_location("_unit_conv_cls", str(UNIT_CONV_PIPELINE))
    if spec is None or spec.loader is None:
        return None
    mod = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(mod)
    except Exception as exc:
        print(f"⚠️ 单位转换加载失败: {exc}")
        return None
    return mod


def _to_num(v) -> float:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    try:
        s = str(v).strip().replace(",", "")
        if not s or s.lower() in {"nan", "none"} or (s.startswith("#") and s.endswith("!")):
            return 0.0
        return float(s)
    except ValueError:
        return 0.0


def _s(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def classify_row(name: str, category: str) -> str:
    """返回分类名。"""
    cat = category.strip()
    n = name.strip()

    # 材料归类优先
    if cat == "钢筋":
        return "钢筋"
    if cat == "砖":
        return "砖"
    if cat in {"模板", "木方"}:  # 木方归入模板
        return "模板"
    if cat in {"碎石", "砂石"}:
        return "碎石"
    if cat == "砂浆":
        return "砂浆水泥"

    if n and not any(k in n for k in REBAR_NAME_EXCLUDE):
        if any(k in n for k in REBAR_NAME):
            return "钢筋"

    if n and not any(k in n for k in MORTAR_NAME_EXCLUDE):
        if any(k in n for k in MORTAR_NAME):
            return "砂浆水泥"

    if n and any(k in n for k in BRICK_NAME):
        return "砖"
    if n and "砖" in n and "砂浆" not in n and "胶泥" not in n and "砖刀" not in n:
        return "砖"

    # 模板 + 木方 统一为「模板」
    if n and any(k in n for k in TEMPLATE_NAME):
        return "模板"
    if n and any(k in n for k in WOOD_NAME) and not any(k in n for k in WOOD_NAME_EXCLUDE):
        return "模板"
    if n and "木" in n and "方" in n:
        return "模板"

    if n and any(k in n for k in GRAVEL_NAME):
        return "碎石"

    return "其他"


def resolve_qty_for_class(cls: str, row: pd.Series, unit_conv) -> tuple[float, str]:
    name = _s(row.get("物品名称", ""))
    spec = _s(row.get("规格", ""))
    unit = _s(row.get("单位", ""))
    qty = _to_num(row.get("数量", 0))
    conv_unit = _s(row.get("转换单位", ""))
    conv_qty = row.get("计算数量", "")

    def _from_conv(want_units: set[str]) -> float | None:
        if conv_unit and str(conv_qty).strip() not in {"", "nan", "None"}:
            if conv_unit.lower() in want_units or conv_unit in want_units:
                return _to_num(conv_qty)
        return None

    # 表内转换列
    if cls == "钢筋":
        v = _from_conv({"kg", "公斤", "千克"})
        if v is not None:
            return v, "kg"
    if cls == "砂浆水泥":
        v = _from_conv({"kg", "公斤", "千克"})
        if v is not None:
            return v, "kg"
    if cls == "砖":
        v = _from_conv({"m3", "m³", "立方", "立方米"})
        if v is not None:
            return v, "m3"
    if cls == "模板":
        v = _from_conv({"m2", "m²", "㎡", "平方米", "平米"})
        if v is not None:
            return v, "m2"

    # 现场转换
    if unit_conv is not None and cls in {"钢筋", "砂浆水泥", "砖", "模板"}:
        try:
            r = unit_conv.convert_row(
                item_name=name, spec=spec, unit=unit, qty=qty if qty else row.get("数量", "")
            )
            cu = _s(r.get("转换单位", "")).lower()
            cq = r.get("计算数量", "")
            if cu and str(cq).strip() not in {"", "nan"}:
                if cls == "钢筋" and cu == "kg":
                    return _to_num(cq), "kg"
                if cls == "砂浆水泥" and cu == "kg":
                    return _to_num(cq), "kg"
                if cls == "砖" and cu in {"m3", "m³"}:
                    return _to_num(cq), "m3"
                if cls == "模板" and cu in {"m2", "m²", "㎡"}:
                    return _to_num(cq), "m2"
        except Exception:
            pass

    u = unit.lower()
    if cls == "钢筋" and u in {"kg", "公斤", "千克"}:
        return qty, "kg"
    if cls == "砂浆水泥" and u in {"kg", "公斤", "千克"}:
        return qty, "kg"
    if cls == "砂浆水泥" and unit in {"袋", "包"}:
        return qty * 50.0, "kg"
    if cls == "砖" and u in {"m3", "m³", "立方米", "立方"}:
        return qty, "m3"
    if cls == "模板" and (u in {"m2", "m²", "㎡"} or "平方" in unit):
        return qty, "m2"

    # 碎石：t / 吨
    if cls == "碎石":
        if u in {"t", "ton", "吨", "dun"}:
            return qty, "t"
        return qty, unit or "t"

    if cls == "其他":
        return qty, unit or "原单位"

    # 模板等未能换算：数量记 0，金额仍累计
    return 0.0, "未换算"


def find_input_file(sum_dir: Path, explicit: str = "") -> Path:
    if explicit:
        p = Path(explicit).expanduser().resolve()
        if not p.exists():
            raise FileNotFoundError(f"输入文件不存在: {p}")
        return p
    for name in PREFERRED_INPUTS:
        p = sum_dir / name
        if p.exists():
            return p
    raise FileNotFoundError(
        f"在 {sum_dir} 未找到汇总源文件，请先运行 4_采购部数据汇总.py\n"
        f"期望之一: {PREFERRED_INPUTS}"
    )


def load_detail(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=0)
    for c in ("项目编码", "项目名称", "材料归类", "物品名称", "规格", "单位", "数量", "含税金额"):
        if c not in df.columns:
            df[c] = ""
    if "转换单位" not in df.columns:
        df["转换单位"] = ""
    if "计算数量" not in df.columns:
        df["计算数量"] = ""
    return df


def build_summary(df: pd.DataFrame, unit_conv) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    rows = []
    for _, r in df.iterrows():
        name = _s(r.get("物品名称", ""))
        cat = _s(r.get("材料归类", ""))
        cls = classify_row(name, cat)
        qty, qty_unit = resolve_qty_for_class(cls, r, unit_conv)
        amt = _to_num(r.get("含税金额", 0))
        rows.append(
            {
                "项目编码": _s(r.get("项目编码", "")),
                "项目名称": _s(r.get("项目名称", "")),
                "分类": cls,
                "物品名称": name,
                "材料归类": cat,
                "规格": _s(r.get("规格", "")),
                "单位": _s(r.get("单位", "")),
                "数量": _to_num(r.get("数量", 0)),
                "汇总数量": qty,
                "汇总数量单位": qty_unit,
                "含税金额": amt,
            }
        )
    detail = pd.DataFrame(rows)
    if detail.empty:
        empty = pd.DataFrame()
        return empty, empty, empty

    g = (
        detail.groupby(["项目编码", "项目名称", "分类"], as_index=False)
        .agg(汇总数量=("汇总数量", "sum"), 含税金额=("含税金额", "sum"), 行数=("分类", "count"))
    )

    projects = (
        detail.groupby(["项目编码", "项目名称"], as_index=False)
        .agg(明细行数=("分类", "count"), 合计金额=("含税金额", "sum"))
    )

    wide = projects.copy()
    for cls, qty_col, amt_col, _u in CLASS_SPECS:
        q = g[g["分类"] == cls][["项目编码", "项目名称", "汇总数量"]].rename(
            columns={"汇总数量": qty_col}
        )
        a = g[g["分类"] == cls][["项目编码", "项目名称", "含税金额"]].rename(
            columns={"含税金额": amt_col}
        )
        wide = wide.merge(q, on=["项目编码", "项目名称"], how="left")
        wide = wide.merge(a, on=["项目编码", "项目名称"], how="left")

    qty_cols = [c[1] for c in CLASS_SPECS]
    amt_cols = [c[2] for c in CLASS_SPECS]
    for c in qty_cols + amt_cols:
        if c not in wide.columns:
            wide[c] = 0.0
        wide[c] = pd.to_numeric(wide[c], errors="coerce").fillna(0.0)

    # 数量取整（汇总展示用整数）
    for c in qty_cols:
        wide[c] = wide[c].round(0)

    wide["分类金额合计"] = wide[amt_cols].sum(axis=1)
    wide["金额差异"] = (wide["合计金额"] - wide["分类金额合计"]).round(2)
    wide = wide.sort_values(["合计金额", "项目编码"], ascending=[False, True]).reset_index(drop=True)

    ordered = ["项目编码", "项目名称"]
    for _, q, a, _ in CLASS_SPECS:
        ordered += [q, a]
    ordered += ["合计金额", "分类金额合计", "金额差异", "明细行数"]
    wide = wide[ordered]

    class_sum = (
        detail.groupby("分类", as_index=False)
        .agg(汇总数量=("汇总数量", "sum"), 含税金额=("含税金额", "sum"), 行数=("分类", "count"))
    )
    class_sum["汇总数量"] = class_sum["汇总数量"].round(0)
    # 固定分类顺序
    order_map = {c[0]: i for i, c in enumerate(CLASS_SPECS)}
    class_sum["_ord"] = class_sum["分类"].map(lambda x: order_map.get(x, 99))
    class_sum = class_sum.sort_values(["_ord", "含税金额"], ascending=[True, False]).drop(
        columns=["_ord"]
    )

    return wide, class_sum, detail


def _style_sheet(ws, *, qty_headers: set[str] | None = None, amount_headers: set[str] | None = None) -> None:
    """数量列整数格式，金额列两位小数。"""
    qty_headers = qty_headers or set()
    amount_headers = amount_headers or set()
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"
    headers = [c.value for c in ws[1]]
    for cell in ws[1]:
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.font = Font(bold=True)

    for col_idx, h in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = min(max(len(str(h or "")) + 2, 10), 18)
        h_str = str(h or "")
        is_qty = h in qty_headers or h_str.endswith("_kg") or h_str.endswith("_m3") or h_str.endswith("_m2") or h_str.endswith("_t") or h_str.endswith("_数量") or h_str in {"汇总数量", "数量", "行数", "明细行数"}
        is_amt = h in amount_headers or ("金额" in h_str)
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, col_idx)
            if not isinstance(cell.value, (int, float)):
                continue
            if is_qty and not is_amt:
                cell.number_format = "#,##0"  # 整数，无小数
            elif is_amt:
                cell.number_format = "#,##0.00"

    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def write_output(
    out_path: Path,
    wide: pd.DataFrame,
    class_sum: pd.DataFrame,
    detail: pd.DataFrame,
    source: Path,
) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    qty_headers = {c[1] for c in CLASS_SPECS} | {"明细行数", "汇总数量", "数量", "行数"}
    amt_headers = {c[2] for c in CLASS_SPECS} | {"合计金额", "分类金额合计", "金额差异", "含税金额"}

    # 项目汇总
    ws1 = wb.active
    ws1.title = "项目汇总"
    ws1.append(list(wide.columns))
    for row in wide.itertuples(index=False, name=None):
        ws1.append(list(row))
    if not wide.empty:
        total_row = ["", "【全部项目合计】"]
        for col in wide.columns[2:]:
            if col == "明细行数":
                total_row.append(int(wide[col].sum()))
            elif col in qty_headers:
                total_row.append(round(float(wide[col].sum()), 0))
            else:
                total_row.append(float(wide[col].sum()))
        ws1.append(total_row)
        for cell in ws1[ws1.max_row]:
            cell.font = Font(bold=True)
    _style_sheet(ws1, qty_headers=qty_headers, amount_headers=amt_headers)

    # 分类合计
    ws2 = wb.create_sheet("分类合计")
    ws2.append(["说明", f"数据源: {source.name}"])
    ws2.append(
        [
            "分类口径",
            "钢筋kg / 砂浆水泥kg / 砖m3 / 模板m2(含木方) / 碎石t / 其他",
        ]
    )
    ws2.append([])
    ws2.append(["分类", "汇总数量", "数量单位", "含税金额", "行数"])
    unit_map = {c[0]: c[3] for c in CLASS_SPECS}
    if not class_sum.empty:
        for _, r in class_sum.iterrows():
            cls = r["分类"]
            ws2.append(
                [
                    cls,
                    round(float(r["汇总数量"]), 0),
                    unit_map.get(cls, ""),
                    float(r["含税金额"]),
                    int(r["行数"]),
                ]
            )
    _style_sheet(
        ws2,
        qty_headers={"汇总数量", "行数"},
        amount_headers={"含税金额"},
    )

    # 分类明细
    ws3 = wb.create_sheet("分类明细")
    cols = [
        "项目编码", "项目名称", "分类", "物品名称", "材料归类", "规格",
        "单位", "数量", "汇总数量", "汇总数量单位", "含税金额",
    ]
    ws3.append(cols)
    if not detail.empty:
        show = detail[cols].copy()
        show["汇总数量"] = show["汇总数量"].round(0)
        show["数量"] = show["数量"].round(0)
        for row in show.itertuples(index=False, name=None):
            ws3.append(list(row))
    _style_sheet(
        ws3,
        qty_headers={"数量", "汇总数量"},
        amount_headers={"含税金额"},
    )

    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="各项目采购分类汇总")
    parser.add_argument("--input", default="", help="输入明细 Excel")
    parser.add_argument("--output", default="", help="输出路径")
    parser.add_argument("--sum-dir", default="", help="汇总目录")
    args = parser.parse_args()

    sum_dir = Path(args.sum_dir).expanduser() if args.sum_dir else SUM_DIR
    src = find_input_file(sum_dir, args.input)
    out = (
        Path(args.output).expanduser().resolve()
        if args.output
        else sum_dir / DEFAULT_OUTPUT
    )

    print(f"输入: {src}")
    print(f"输出: {out}")

    df = load_detail(src)
    print(f"明细行数: {len(df)}")

    unit_conv = _load_unit_conv()
    print("单位转换:", "已加载" if unit_conv else "未加载")

    wide, class_sum, detail = build_summary(df, unit_conv)
    if wide.empty:
        print("无有效数据，结束")
        return

    write_output(out, wide, class_sum, detail, src)

    print("\n—— 分类金额概览 ——")
    if not class_sum.empty:
        unit_map = {c[0]: c[3] for c in CLASS_SPECS}
        for _, r in class_sum.iterrows():
            u = unit_map.get(r["分类"], "")
            print(
                f"  {r['分类']}: 数量={r['汇总数量']:,.0f} {u}  "
                f"金额={r['含税金额']:,.2f}  行数={int(r['行数'])}"
            )
    print(f"\n项目数: {len(wide)}")
    print(f"已输出: {out}")


if __name__ == "__main__":
    main()
