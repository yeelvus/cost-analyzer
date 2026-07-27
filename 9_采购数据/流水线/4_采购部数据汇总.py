#!/usr/bin/env python3
"""
步骤4：汇总 STD → 纯合并到 C_raw_sum，标准化到 D_sum。

默认输出：
  1) C_采购部采购台账_raw_sum/采购部采购台账_merged.xlsx — 纯合并
  2) D_采购部采购台账_sum/采购部采购台账_std_merged.xlsx — 标准化
  3) D_…/对比版（可选）
  4) 将 std「采购总金额」同步到 D_…/项目采购分类汇总.xlsx

不改动 STD 目录下的单项目 *_processed.xlsx。
可用 --plain-only / --std-only / --no-compare 控制输出。
"""
from __future__ import annotations

import argparse
import importlib.util
import shutil
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

_PKG_ROOT = Path(__file__).resolve().parents[1]
_CODE_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(_PKG_ROOT), str(_CODE_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from common.excel import apply_common_excel_layout

from _采购路径 import (
    CATEGORY_SUMMARY_NAME,
    CONTENT_STD_UTILS,
    MERGED_PLAIN_NAME,
    MERGED_STD_NAME,
    RAW_SUM_DIR,
    STD_DIR,
    SUM_DIR,
    UNIT_PIPELINE,
    is_excel,
)

# 正式输出列（merged / std_merged 不变）
FINAL_OUTPUT_COLUMNS = [
    "项目编码", "项目名称", "材料归类", "物品名称", "规格", "单位", "数量",
    "单价", "不含税金额", "7%增值税", "含税金额", "货币种类", "汇率",
    "标准日期", "供货商名称", "备注", "年月", "签订日期",
]

# 对比版：原始列 + 转换列 + 标准化后列
COMPARE_OUTPUT_COLUMNS = [
    "项目编码", "项目名称",
    "原始物品名称", "原始规格", "原始单位",
    "转换单位", "计算数量",
    "材料归类", "物品名称", "规格", "单位", "数量",
    "单价", "不含税金额", "7%增值税", "含税金额", "货币种类", "汇率",
    "标准日期", "供货商名称", "备注", "年月", "签订日期",
]

OUTPUT_PLAIN = MERGED_PLAIN_NAME
OUTPUT_WITH_STD = MERGED_STD_NAME
OUTPUT_COMPARE = "采购部采购台账_std_merged_对比.xlsx"
UNIT_CONV_PIPELINE = UNIT_PIPELINE

AMOUNT_FIELD_NAMES = {"单价", "不含税金额", "7%增值税", "含税金额", "数量", "计算数量", "汇率"}
# Excel 错误字面量 / 非数值，写入或汇总时当作空
_EXCEL_ERROR_TOKENS = {
    "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!",
    "#GETTING_DATA", "#SPILL!", "#CALC!", "#ERROR!",
}

_STD_UTILS = None
_UNIT_CONV = None


def _is_excel_error_token(value: object) -> bool:
    if value is None:
        return False
    s = str(value).strip().upper()
    if not s:
        return False
    if s in {t.upper() for t in _EXCEL_ERROR_TOKENS}:
        return True
    return s.startswith("#") and s.endswith("!")


def coerce_amount(value: object) -> float | None:
    """将金额/数量安全转为 float；#VALUE!/空/无法解析 → None。"""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if pd.isna(value):
            return None
        return float(value)
    if isinstance(value, str):
        s = value.strip().replace(",", "").replace("，", "")
        if not s or _is_excel_error_token(s):
            return None
        try:
            return float(s)
        except ValueError:
            return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _load_procurement_standardization_utils():
    """加载采购内容标准化工具（采购内容标准化_utils.py）。"""
    global _STD_UTILS
    if _STD_UTILS is not None:
        return _STD_UTILS if _STD_UTILS is not False else None

    if not CONTENT_STD_UTILS.exists():
        print(f"⚠️ 未找到内容标准化工具: {CONTENT_STD_UTILS}")
        _STD_UTILS = False
        return None

    spec = importlib.util.spec_from_file_location(
        "procurement_content_std_utils", str(CONTENT_STD_UTILS)
    )
    if spec is None or spec.loader is None:
        print(f"⚠️ 无法加载内容标准化工具: {CONTENT_STD_UTILS}")
        _STD_UTILS = False
        return None

    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    if not hasattr(mod, "apply_standardization_pipeline"):
        print("⚠️ 内容标准化工具缺少 apply_standardization_pipeline")
        _STD_UTILS = False
        return None

    _STD_UTILS = mod
    return mod


def _load_unit_conversion_pipeline():
    """加载单位数量转换流水线（钢筋网片 m2→kg、砖块→m3 等）。"""
    global _UNIT_CONV
    if _UNIT_CONV is not None:
        return _UNIT_CONV if _UNIT_CONV is not False else None

    if not UNIT_CONV_PIPELINE.exists():
        print(f"⚠️ 未找到单位转换流水线: {UNIT_CONV_PIPELINE}")
        _UNIT_CONV = False
        return None

    spec = importlib.util.spec_from_file_location(
        "unit_qty_conversion_pipeline", str(UNIT_CONV_PIPELINE)
    )
    if spec is None or spec.loader is None:
        _UNIT_CONV = False
        return None
    mod = importlib.util.module_from_spec(spec)
    try:
        spec.loader.exec_module(mod)
    except Exception as exc:
        print(f"⚠️ 单位转换流水线加载失败: {exc}")
        _UNIT_CONV = False
        return None
    _UNIT_CONV = mod
    return mod


def apply_unified_workbook_format(output_wb) -> None:
    amount_keywords = ("金额", "单价", "税")
    for sheet_name in ("合并数据", "项目汇总", "采购总金额"):
        if sheet_name not in output_wb.sheetnames:
            continue
        ws = output_wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        amount_headers = [
            h for h in headers if isinstance(h, str) and any(k in h for k in amount_keywords)
        ]
        if sheet_name == "采购总金额":
            apply_common_excel_layout(
                ws,
                numeric_start_col=3,
                amount_headers=[h for h in headers[2:] if isinstance(h, str)],
                header_row=1,
                header_height=40,
                freeze_panes="A2",
            )
        else:
            apply_common_excel_layout(
                ws,
                numeric_start_col=ws.max_column + 1,
                amount_headers=amount_headers,
                header_row=1,
                header_height=40,
                freeze_panes="A2",
            )


def _period_key(period_text: str) -> tuple[int, int]:
    y, m = str(period_text).split("-")
    return int(y), int(m)


def build_procurement_total_sheet(merged_df: pd.DataFrame) -> pd.DataFrame:
    required = {"项目编码", "项目名称", "年月", "含税金额"}
    if not required.issubset(set(merged_df.columns)):
        return pd.DataFrame(columns=["项目编码", "项目名称"])

    work = merged_df[["项目编码", "项目名称", "年月", "含税金额"]].copy()
    work["项目编码"] = work["项目编码"].fillna("").astype(str).str.strip()
    work["项目名称"] = work["项目名称"].fillna("").astype(str).str.strip()
    work["年月"] = work["年月"].fillna("").astype(str).str.strip()
    work["含税金额"] = work["含税金额"].map(coerce_amount)
    work = work[(work["项目编码"] != "") | (work["项目名称"] != "")]
    work = work[work["年月"].str.fullmatch(r"(?:19|20)\d{2}-(?:0[1-9]|1[0-2])", na=False)]
    work = work[work["含税金额"].notna()]
    if work.empty:
        return pd.DataFrame(columns=["项目编码", "项目名称"])

    grouped = work.groupby(["项目编码", "项目名称", "年月"], as_index=False).agg(
        含税金额=("含税金额", "sum")
    )
    periods = sorted(grouped["年月"].unique().tolist(), key=_period_key)
    label_map = {p: f"{int(p.split('-')[0])}年{int(p.split('-')[1])}月" for p in periods}
    pivot = (
        grouped.pivot_table(
            index=["项目编码", "项目名称"],
            columns="年月",
            values="含税金额",
            aggfunc="sum",
            fill_value=0.0,
        )
        .reindex(columns=periods, fill_value=0.0)
    )
    amount_wide = pivot.rename(columns=label_map).reset_index()
    sort_cols = [label_map[p] for p in reversed(periods)]
    amount_wide["__编码为空"] = amount_wide["项目编码"].astype(str).str.strip().eq("")
    amount_wide = amount_wide.sort_values(
        ["__编码为空", *sort_cols],
        ascending=[True] + [False] * len(sort_cols),
    ).drop(columns=["__编码为空"]).reset_index(drop=True)
    return amount_wide


def list_std_excels(folder: Path) -> list[Path]:
    files = []
    for p in folder.rglob("*"):
        if not is_excel(p):
            continue
        name = p.name
        if name.startswith("merged") or name.endswith("_std_merged.xlsx"):
            continue
        if "对比" in name:
            continue
        if name.startswith("采购部采购台账_"):
            continue
        files.append(p)
    return sorted(files)


def _new_merge_workbook(columns: list[str]) -> tuple[openpyxl.Workbook, openpyxl.worksheet.worksheet.Worksheet]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "合并数据"
    ws.append(columns)
    return wb, ws


def _append_data_row(ws, row_idx: int, columns: list[str], row_data: dict) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        value = row_data.get(col_name, "")
        try:
            if pd.isna(value):
                value = ""
        except Exception:
            pass

        # 金额/数量列：清洗 #VALUE! 等，只写真正的数字或空
        if col_name in AMOUNT_FIELD_NAMES or col_name.endswith("金额"):
            num = coerce_amount(value)
            cell = ws.cell(row_idx, col_idx, num if num is not None else None)
            if num is not None and col_name in {
                "单价", "不含税金额", "7%增值税", "含税金额", "计算数量",
            }:
                cell.number_format = "#,##0.00"
            continue

        if _is_excel_error_token(value):
            value = ""
        cell = ws.cell(row_idx, col_idx, value)


def _finalize_workbook(
    output_wb: openpyxl.Workbook,
    columns: list[str],
    project_totals: dict[tuple[str, str], float],
    data_last_row: int,
) -> openpyxl.Workbook | None:
    if data_last_row < 2:
        return None

    merged_ws = output_wb["合并数据"]
    merged_ws.row_dimensions[1].height = 40
    for cell in merged_ws[1]:
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    merged_ws.freeze_panes = "A2"
    last_col = get_column_letter(len(columns))
    merged_ws.auto_filter.ref = f"A1:{last_col}{data_last_row}"

    summary_ws = output_wb.create_sheet("项目汇总")
    summary_ws.append(["项目编码", "项目名称", "总采购金额"])
    grand_total = 0.0
    for row_idx, ((code, name), total) in enumerate(sorted(project_totals.items()), start=2):
        amt = float(total) if total is not None else 0.0
        summary_ws.append([code, name, amt])
        summary_ws.cell(row_idx, 3).number_format = "#,##0.00"
        grand_total += amt
    # 合计行（数值，不用公式，避免 #VALUE!）
    last = summary_ws.max_row + 1
    summary_ws.cell(last, 1, "")
    summary_ws.cell(last, 2, "【合计】")
    summary_ws.cell(last, 3, round(grand_total, 2)).number_format = "#,##0.00"

    merged_df = pd.DataFrame(merged_ws.values)
    merged_df.columns = merged_df.iloc[0]
    merged_df = merged_df.iloc[1:].reset_index(drop=True)
    # 金额列再清洗一遍，供宽表聚合
    if "含税金额" in merged_df.columns:
        merged_df["含税金额"] = merged_df["含税金额"].map(coerce_amount)

    total_df = build_procurement_total_sheet(merged_df)
    total_ws = output_wb.create_sheet("采购总金额")
    total_ws.append(total_df.columns.tolist())
    for row in total_df.itertuples(index=False, name=None):
        cleaned = []
        for v in row:
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cleaned.append(float(v))
            else:
                # 非金额文本列（项目编码/名称）原样
                cleaned.append("" if _is_excel_error_token(v) else v)
        total_ws.append(cleaned)

    # 只保留业务 sheet，去掉 openpyxl/Excel 可能残留的空 Sheet
    keep = {"合并数据", "项目汇总", "采购总金额"}
    for sn in list(output_wb.sheetnames):
        if sn not in keep:
            del output_wb[sn]

    return output_wb


def merge_standardized_excels(
    folder: Path,
    *,
    use_content_standardization: bool = False,
    include_compare: bool = False,
) -> tuple[openpyxl.Workbook | None, openpyxl.Workbook | None]:
    """
    返回 (正式汇总工作簿, 对比版工作簿)。
    对比版仅在 use_content_standardization 且 include_compare 时生成。
    """
    excel_paths = list_std_excels(folder)
    std_utils = None
    build_compare = bool(use_content_standardization and include_compare)

    if use_content_standardization:
        std_utils = _load_procurement_standardization_utils()
        if std_utils is None:
            print("⚠️ 内容标准化不可用，本轮回退为纯汇总")
            use_content_standardization = False
            build_compare = False
        else:
            try:
                unit_ref = std_utils.build_global_unit_ref(folder)
                std_utils.set_global_unit_ref(unit_ref)
                print(f"全局单位参考库: {len(unit_ref)} 个物品")
            except Exception as exc:
                print(f"⚠️ 全局单位库构建失败（将仅用文件内补全）: {exc}")

    mode = "启用采购内容标准化" if use_content_standardization else "纯汇总"
    if build_compare:
        mode += " + 对比版"
    print("\n" + "=" * 72)
    print(f"【{mode}】目录: {folder}")
    print(f"【{mode}】文件数: {len(excel_paths)}")
    print("=" * 72 + "\n")

    if not excel_paths:
        print("⚠️ 未找到任何 Excel 文件")
        return None, None

    formal_wb, formal_ws = _new_merge_workbook(FINAL_OUTPUT_COLUMNS)
    compare_wb = compare_ws = None
    unit_conv = None
    if build_compare:
        compare_wb, compare_ws = _new_merge_workbook(COMPARE_OUTPUT_COLUMNS)
        unit_conv = _load_unit_conversion_pipeline()
        if unit_conv is None:
            print("⚠️ 单位转换不可用：对比版「转换单位/计算数量」将为空")

    formal_row = 2
    compare_row = 2
    project_totals: dict[tuple[str, str], float] = defaultdict(float)
    compare_totals: dict[tuple[str, str], float] = defaultdict(float)

    for idx, file_path in enumerate(excel_paths, start=1):
        print(f"[{idx}/{len(excel_paths)}] {file_path.name}")
        read_path = file_path
        temp_dir = None
        orig_snapshot = None

        if use_content_standardization and std_utils is not None:
            try:
                # 标准化前快照原始 物品名称/规格/单位（按行序，供对比版）
                if build_compare:
                    try:
                        orig_df = pd.read_excel(file_path)
                        for c in ("物品名称", "规格", "单位"):
                            if c not in orig_df.columns:
                                orig_df[c] = ""
                        orig_snapshot = orig_df[["物品名称", "规格", "单位"]].copy()
                        orig_snapshot.columns = ["原始物品名称", "原始规格", "原始单位"]
                    except Exception as exc:
                        print(f"  ⚠️ 读取原始列失败，对比版该文件将无原始值: {exc}")
                        orig_snapshot = None

                temp_dir = tempfile.mkdtemp(prefix="procurement_std_")
                temp_file = Path(temp_dir) / file_path.name
                shutil.copy2(file_path, temp_file)
                std_utils.apply_standardization_pipeline(
                    str(temp_file),
                    auto_prompt_units=False,
                    resolve_pending=False,
                )
                read_path = temp_file
            except Exception as exc:
                print(f"  ⚠️ 标准化失败，用原文件: {exc}")
                read_path = file_path

        try:
            df = pd.read_excel(read_path)
        except Exception as exc:
            print(f"  ✗ 读取失败: {exc}")
            continue
        finally:
            if temp_dir and Path(temp_dir).is_dir():
                shutil.rmtree(temp_dir, ignore_errors=True)

        if df.empty:
            continue

        for col in FINAL_OUTPUT_COLUMNS:
            if col not in df.columns:
                df[col] = ""

        # 清洗源文件里残留的 #VALUE! 等错误字面量
        for col in list(df.columns):
            if col in AMOUNT_FIELD_NAMES or str(col).endswith("金额"):
                df[col] = df[col].map(coerce_amount)

        # 对齐原始快照行数
        if build_compare:
            if orig_snapshot is None or len(orig_snapshot) != len(df):
                orig_snapshot = pd.DataFrame(
                    {
                        "原始物品名称": df["物品名称"].values,
                        "原始规格": df["规格"].values,
                        "原始单位": df["单位"].values,
                    }
                )
            else:
                orig_snapshot = orig_snapshot.reset_index(drop=True)
            df = df.reset_index(drop=True)
            df = pd.concat([orig_snapshot, df], axis=1)

            # 单位数量转换：用标准化后的 名称/规格/单位/数量 计算
            conv_units: list = []
            conv_qtys: list = []
            for _, r in df.iterrows():
                if unit_conv is not None and hasattr(unit_conv, "convert_row"):
                    cr = unit_conv.convert_row(
                        item_name=r.get("物品名称", ""),
                        spec=r.get("规格", ""),
                        unit=r.get("单位", ""),
                        qty=r.get("数量", ""),
                    )
                    conv_units.append(cr.get("转换单位", ""))
                    conv_qtys.append(cr.get("计算数量", ""))
                else:
                    conv_units.append("")
                    conv_qtys.append("")
            df["转换单位"] = conv_units
            df["计算数量"] = conv_qtys

        for i, row in df.iterrows():
            formal_data = {c: row[c] for c in FINAL_OUTPUT_COLUMNS}
            _append_data_row(formal_ws, formal_row, FINAL_OUTPUT_COLUMNS, formal_data)

            code = "" if pd.isna(row["项目编码"]) else str(row["项目编码"]).strip()
            name = "" if pd.isna(row["项目名称"]) else str(row["项目名称"]).strip()
            amt = coerce_amount(row["含税金额"])
            if amt is not None:
                project_totals[(code, name)] += float(amt)
            formal_row += 1

            if build_compare and compare_ws is not None:
                compare_data = {c: row[c] if c in row.index else "" for c in COMPARE_OUTPUT_COLUMNS}
                _append_data_row(compare_ws, compare_row, COMPARE_OUTPUT_COLUMNS, compare_data)
                if amt is not None:
                    compare_totals[(code, name)] += float(amt)
                compare_row += 1

    formal_out = _finalize_workbook(formal_wb, FINAL_OUTPUT_COLUMNS, project_totals, formal_row - 1)
    compare_out = None
    if build_compare and compare_wb is not None:
        compare_out = _finalize_workbook(
            compare_wb, COMPARE_OUTPUT_COLUMNS, compare_totals, compare_row - 1
        )

    if formal_out is None:
        print("⚠️ 没有有效数据")
    return formal_out, compare_out


def sync_monthly_total_to_category_summary(
    std_merged_path: Path,
    category_summary_path: Path,
) -> bool:
    """
    把 std_merged 的「采购总金额」sheet（项目×月份宽表）写入
    项目采购分类汇总.xlsx（存在则替换该 sheet，保留其它 sheet）。
    """
    std_merged_path = Path(std_merged_path)
    category_summary_path = Path(category_summary_path)
    if not std_merged_path.is_file():
        print(f"  ⚠️ 无 std 汇总，跳过同步采购总金额: {std_merged_path}")
        return False

    try:
        total_df = pd.read_excel(std_merged_path, sheet_name="采购总金额")
    except Exception as exc:
        print(f"  ⚠️ 读取「采购总金额」失败: {exc}")
        return False

    if total_df is None or total_df.empty:
        print("  ⚠️ 「采购总金额」为空，跳过同步")
        return False

    category_summary_path.parent.mkdir(parents=True, exist_ok=True)
    sheet_name = "采购总金额"

    if category_summary_path.is_file():
        try:
            book = openpyxl.load_workbook(category_summary_path)
        except Exception as exc:
            print(f"  ⚠️ 打开分类汇总失败，将新建: {exc}")
            book = openpyxl.Workbook()
            # 去掉默认空表，后面统一建
            if book.active and book.active.max_row == 1 and book.active.max_column == 1:
                pass
    else:
        book = openpyxl.Workbook()

    if sheet_name in book.sheetnames:
        del book[sheet_name]
    # 若仅有一个空 Sheet，可先删再建
    if len(book.sheetnames) == 1:
        only = book.sheetnames[0]
        ws0 = book[only]
        if ws0.max_row <= 1 and (ws0.max_column or 1) <= 1 and not ws0["A1"].value:
            del book[only]

    ws = book.create_sheet(sheet_name)
    headers = [str(c) for c in total_df.columns.tolist()]
    ws.append(headers)
    for row in total_df.itertuples(index=False, name=None):
        cleaned = []
        for v in row:
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                try:
                    if pd.isna(v):
                        cleaned.append(None)
                    else:
                        cleaned.append(float(v))
                except Exception:
                    cleaned.append(v)
            else:
                cleaned.append("" if _is_excel_error_token(v) else v)
        ws.append(cleaned)

    # 金额列格式：从第 3 列起多为月份金额
    for r in range(2, (ws.max_row or 1) + 1):
        for c in range(3, (ws.max_column or 2) + 1):
            cell = ws.cell(r, c)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"

    ws.freeze_panes = "C2"
    if ws.max_row and ws.max_column:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.row_dimensions[1].height = 36
    for cell in ws[1]:
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    try:
        apply_unified_workbook_format(book)
    except Exception:
        pass

    book.save(category_summary_path)
    print(f"已同步「采购总金额」→ {category_summary_path}")
    return True


def main() -> None:
    parser = argparse.ArgumentParser(description="汇总采购部 std 台账（融入内容标准化）")
    parser.add_argument(
        "--plain-only",
        action="store_true",
        help="只输出纯合并文件到 C_raw_sum（不跑内容标准化）",
    )
    parser.add_argument(
        "--std-only",
        action="store_true",
        help="只输出内容标准化后的汇总到 D_sum（含对比版，除非 --no-compare）",
    )
    parser.add_argument(
        "--no-compare",
        action="store_true",
        help="不生成对比版 Excel",
    )
    parser.add_argument("--input", default="", help="覆盖输入目录，默认 STD_DIR")
    parser.add_argument(
        "--output-dir",
        default="",
        help="覆盖 std 输出目录，默认 SUM_DIR（D_sum）；纯合并仍写 RAW_SUM_DIR",
    )
    parser.add_argument(
        "--raw-sum-dir",
        default="",
        help="纯合并输出目录，默认 C_采购部采购台账_raw_sum",
    )
    args = parser.parse_args()

    in_dir = Path(args.input).expanduser() if args.input else STD_DIR
    out_dir = Path(args.output_dir).expanduser() if args.output_dir else SUM_DIR
    raw_sum_dir = (
        Path(args.raw_sum_dir).expanduser() if args.raw_sum_dir else RAW_SUM_DIR
    )
    if not in_dir.is_dir():
        raise FileNotFoundError(f"输入目录不存在: {in_dir}")
    out_dir.mkdir(parents=True, exist_ok=True)
    raw_sum_dir.mkdir(parents=True, exist_ok=True)

    print("汇总模式")
    print(f"- 输入 STD:     {in_dir}")
    print(f"- 纯合并输出:   {raw_sum_dir / OUTPUT_PLAIN}")
    print(f"- 标准化输出:   {out_dir / OUTPUT_WITH_STD}")
    print(f"- 内容标准化工具: {CONTENT_STD_UTILS}")

    if args.plain_only and args.std_only:
        raise SystemExit("不能同时指定 --plain-only 与 --std-only")

    # jobs: (use_std, formal_filename, include_compare, dest_dir)
    if args.plain_only:
        jobs = [(False, OUTPUT_PLAIN, False, raw_sum_dir)]
    elif args.std_only:
        jobs = [(True, OUTPUT_WITH_STD, not args.no_compare, out_dir)]
    else:
        jobs = [
            (False, OUTPUT_PLAIN, False, raw_sum_dir),
            (True, OUTPUT_WITH_STD, not args.no_compare, out_dir),
        ]

    std_merged_path: Path | None = None
    for use_std, filename, include_compare, dest in jobs:
        formal_wb, compare_wb = merge_standardized_excels(
            in_dir,
            use_content_standardization=use_std,
            include_compare=include_compare,
        )
        if formal_wb is None:
            print(f"未产生有效数据: {filename}")
            continue

        dest.mkdir(parents=True, exist_ok=True)
        out_path = dest / filename
        apply_unified_workbook_format(formal_wb)
        formal_wb.save(out_path)
        print(f"已输出: {out_path}")

        if use_std and filename == OUTPUT_WITH_STD:
            std_merged_path = out_path

        if compare_wb is not None:
            compare_path = out_dir / OUTPUT_COMPARE
            apply_unified_workbook_format(compare_wb)
            compare_wb.save(compare_path)
            print(f"已输出对比版: {compare_path}")

    # std 汇总的「采购总金额」→ 项目采购分类汇总.xlsx
    if std_merged_path and std_merged_path.is_file():
        cat_path = out_dir / CATEGORY_SUMMARY_NAME
        sync_monthly_total_to_category_summary(std_merged_path, cat_path)


if __name__ == "__main__":
    main()
