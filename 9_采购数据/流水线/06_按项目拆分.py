#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
按项目拆分 C_raw_sum/采购部采购台账_raw_sum.xlsx
→ E_采购部采购台账_std_split/{项目编码}_{项目名称}_采购数据.xlsx

每个文件仅含该项目明细（默认取「合并数据」sheet）。

用法：
  python3 流水线/06_按项目拆分.py
  python3 流水线/06_按项目拆分.py --input ... --output-dir ...
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

_PKG_ROOT = Path(__file__).resolve().parents[1]
_CODE_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(_PKG_ROOT), str(_CODE_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from common.excel import apply_common_excel_layout

from _采购路径 import (
    RAW_SUM_DIR,
    RAW_SUM_FILE_NAME,
    STD_SPLIT_DIR,
)


def _safe_filename_part(text: str, *, max_len: int = 40) -> str:
    s = str(text or "").strip()
    s = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", s)
    s = re.sub(r"\s+", "", s)
    s = s.strip("._ ")
    if not s:
        s = "未命名"
    return s[:max_len]


def _s(v) -> str:
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except Exception:
        pass
    t = str(v).strip()
    return "" if t.lower() in ("nan", "none") else t


def _load_raw_sum(path: Path) -> pd.DataFrame:
    try:
        df = pd.read_excel(path, sheet_name="合并数据")
    except Exception:
        df = pd.read_excel(path, sheet_name=0)
    return df


def _write_project_excel(df: pd.DataFrame, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    # 首列序号
    out = df.copy().reset_index(drop=True)
    if "序号" in out.columns:
        out = out.drop(columns=["序号"])
    out.insert(0, "序号", range(1, len(out) + 1))

    wb = Workbook()
    ws = wb.active
    ws.title = "采购数据"
    headers = [str(c) for c in out.columns.tolist()]
    ws.append(headers)
    for row in out.itertuples(index=False, name=None):
        cleaned = []
        for v in row:
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                try:
                    cleaned.append(None if pd.isna(v) else v)
                except Exception:
                    cleaned.append(v)
            else:
                cleaned.append(_s(v))
        ws.append(cleaned)

    # 金额列
    amount_keys = ("金额", "单价", "税", "汇率")
    amount_headers = [h for h in headers if any(k in h for k in amount_keys)]
    apply_common_excel_layout(
        ws,
        numeric_start_col=ws.max_column + 1,
        amount_headers=amount_headers,
        header_row=1,
        header_height=36,
        freeze_panes="A2",
    )
    # 左对齐文本
    left_keys = ("物品名称", "规格", "项目名称", "材料归类", "供货商", "备注")
    left = Alignment(horizontal="left", vertical="center")
    center = Alignment(horizontal="center", vertical="center")
    for col_idx, h in enumerate(headers, start=1):
        for r in range(2, (ws.max_row or 1) + 1):
            cell = ws.cell(r, col_idx)
            if h == "序号":
                cell.alignment = center
            elif any(k in h for k in left_keys):
                cell.alignment = left

    if ws.max_row and ws.max_column:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    wb.save(out_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="按项目拆分 raw_sum → E_std_split/{编码}_{名称}_采购数据.xlsx"
    )
    parser.add_argument(
        "--input",
        default="",
        help=f"输入 raw_sum，默认 {RAW_SUM_DIR / RAW_SUM_FILE_NAME}",
    )
    parser.add_argument(
        "--output-dir",
        default="",
        help=f"输出目录，默认 {STD_SPLIT_DIR}",
    )
    parser.add_argument(
        "--clear",
        action="store_true",
        help="写出前清空输出目录中已有 *_采购数据.xlsx",
    )
    args = parser.parse_args(argv)

    src = Path(args.input).expanduser() if args.input else (RAW_SUM_DIR / RAW_SUM_FILE_NAME)
    out_dir = Path(args.output_dir).expanduser() if args.output_dir else STD_SPLIT_DIR
    if not src.is_file():
        raise FileNotFoundError(
            f"未找到 raw_sum: {src}\n请先运行 流水线/04a_raw_sum纯合并.py"
        )
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.clear:
        for p in out_dir.glob("*_采购数据.xlsx"):
            try:
                p.unlink()
            except Exception:
                pass

    print("按项目拆分 raw_sum")
    print(f"- 输入: {src}")
    print(f"- 输出: {out_dir}")

    df = _load_raw_sum(src)
    if df.empty:
        print("无数据")
        return 1

    for c in ("项目编码", "项目名称"):
        if c not in df.columns:
            df[c] = ""

    df["_code"] = df["项目编码"].map(_s)
    df["_name"] = df["项目名称"].map(_s)
    # 无编码时用名称兜底分组
    df["_key"] = df.apply(
        lambda r: r["_code"] or r["_name"] or "未知项目",
        axis=1,
    )

    groups = df.groupby("_key", dropna=False)
    n_files = 0
    for key, g in groups:
        code = _s(g["_code"].iloc[0]) if len(g) else ""
        name = _s(g["_name"].iloc[0]) if len(g) else ""
        if not code and not name:
            code, name = "未知", str(key)
        if not code:
            code = "无编码"
        if not name:
            name = "未命名项目"

        fname = (
            f"{_safe_filename_part(code)}_"
            f"{_safe_filename_part(name)}_采购数据.xlsx"
        )
        part = g.drop(columns=["_code", "_name", "_key"], errors="ignore")
        out_path = out_dir / fname
        _write_project_excel(part, out_path)
        n_files += 1
        print(f"  ✓ {fname}  ({len(part)} 行)")

    print(f"\n完成: 共 {n_files} 个项目文件 → {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
