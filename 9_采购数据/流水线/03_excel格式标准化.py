#!/usr/bin/env python3
"""
步骤3：读取 EXTRACT_DIR 去重后的台账 → 列映射/项目匹配/日期
→ 写出 STD_DIR/*_processed.xlsx。

不做：
- 单位纠错/补全
- 物品规格等内容规则标准化（桩/钢筋/网片/砂浆等由其他脚本处理）
"""
from __future__ import annotations

import re
import sys
from datetime import date as date_type
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

_PKG_ROOT = Path(__file__).resolve().parents[1]
_CODE_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(_PKG_ROOT), str(_CODE_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from settings import MATCH_DB_PATH

from _采购路径 import (
    EXTRACT_DIR,
    PROJECT_INFO_FILE,
    STD_DIR,
    extract_month_iso,
    extract_project_name,
    list_excels,
)

# ── 项目匹配 ──────────────────────────────────────────────
_MATCHER_PATH = _CODE_ROOT / "0_项目基本信息" / "project_name_matcher.py"


def _load_matcher():
    import importlib.util

    spec = importlib.util.spec_from_file_location("_proj_matcher", str(_MATCHER_PATH))
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载: {_MATCHER_PATH}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


_matcher = _load_matcher()

# ── 列配置 ────────────────────────────────────────────────
COLUMN_MAPPINGS = {
    "材料归类": ["材料归类"],
    "物品名称": ["物品名称", "采购物资名称", "物资名称"],
    "规格": ["规格", "规格型号", "型号", "规格型号/参数"],
    "单位": ["标准单位", "单位"],
    "数量": ["数量"],
    "单价": ["单价", "不含税单价", "含税单价"],
    "不含税金额": ["不含税金额", "金额"],
    "7%增值税": ["7%增值税", "税额", "增值税", "税金"],
    "含税金额": ["含税金额", "总金额", "价税合计"],
    "货币种类": ["货币种类", "币种"],
    "汇率": ["汇率"],
    "签订日期": ["签订日期", "日期"],
    "供货商名称": ["供货商名称", "供货单位", "供应商", "供应商名称"],
    "备注": ["备注"],
}
STANDARD_COLUMNS = list(COLUMN_MAPPINGS.keys())
CORE_REQUIRED = ["物品名称", "单位", "数量"]
FINAL_COLUMNS = [
    "项目编码", "项目名称", "材料归类", "物品名称", "规格", "单位", "数量",
    "单价", "不含税金额", "7%增值税", "含税金额", "货币种类", "汇率",
    "标准日期", "供货商名称", "备注", "年月", "签订日期",
]
AMOUNT_COLS = {"单价", "不含税金额", "7%增值税", "含税金额"}
CENTER_COLS = {"单位", "数量", "货币种类", "汇率"}
# 正文左对齐（表头仍居中）
LEFT_COLS = {
    "物品名称",
    "规格",
    "项目名称",
    "材料归类",
    "供货商名称",
    "备注",
    "签订日期",
}
OUTPUT_SUFFIX = "_processed"


# ── 项目配置 / 匹配 ───────────────────────────────────────
def load_project_config() -> list[dict]:
    try:
        df = pd.read_excel(PROJECT_INFO_FILE, sheet_name=0)
    except Exception as exc:
        print(f"警告：无法加载项目配置: {exc}")
        return []
    if not {"项目编码", "项目名称"}.issubset(df.columns):
        return []
    projects = []
    for _, row in df.iterrows():
        code = str(row["项目编码"]).strip() if pd.notna(row["项目编码"]) else ""
        name = str(row["项目名称"]).strip() if pd.notna(row["项目名称"]) else ""
        if code or name:
            projects.append({"project_code": code, "project_name": name})
    return projects


def resolve_project(raw_name: str, projects: list[dict], match_db) -> dict | None:
    """优先匹配库（含模糊）；再回落项目表精确/包含。"""
    hit = _matcher.match_project(raw_name, match_db, fuzzy=True, skip_marked_unmatched=True)
    if hit:
        return {
            "project_code": hit.get("project_code", "") or "",
            "project_name": hit.get("project_name", raw_name),
            "source": hit.get("source", ""),
        }
    for p in projects:
        if p["project_name"] == raw_name:
            return {**p, "source": "config_exact"}
    # 唯一包含再接受
    contain = [
        p
        for p in projects
        if raw_name and (raw_name in p["project_name"] or p["project_name"] in raw_name)
    ]
    if len(contain) == 1:
        return {**contain[0], "source": "config_contain"}
    return None


def preload_project_matches(
    files: list[Path],
    projects: list[dict],
    match_db,
    *,
    interactive: bool = True,
) -> None:
    """
    批量处理未匹配项目。
    - 已匹配 / 模糊可判 / 曾选「保留原名」的 → 不再弹窗
    - interactive=False → 静默 mark_unmatched，适合无人值守
    """
    names = [extract_project_name(f.name) for f in files]
    # 先让匹配库吃一遍模糊，写入别名
    still = []
    for name in names:
        if not name:
            continue
        if resolve_project(name, projects, match_db):
            continue
        still.append(name)

    need_prompt = _matcher.collect_names_needing_prompt(still, match_db, fuzzy=True)
    if not need_prompt:
        print("项目匹配: 全部可自动处理（精确/模糊/历史保留），无需弹窗")
        return

    print(f"\n发现 {len(need_prompt)} 个新未识别项目" + ("，打开批量匹配…" if interactive else "（非交互，保留原名）"))
    results = _matcher.prompt_unmatched_project_mappings(
        need_prompt,
        match_db,
        window_title="采购-批量项目匹配",
        db_path=str(MATCH_DB_PATH),
        interactive=interactive,
    )
    lookup = _matcher.get_project_lookup(match_db)
    for raw, std in results.items():
        if std:
            code, _ = lookup.get(std, ("", ""))
            print(f"  ✓ {raw} -> [{code}] {std}")
        else:
            print(f"  - 保留原名(下次不再弹): {raw}")


# ── 表头 / 行提取 ─────────────────────────────────────────
def normalize_headers(headers) -> list[str]:
    out = []
    for col in headers:
        s = str(col).strip() if col else ""
        for std, aliases in COLUMN_MAPPINGS.items():
            if s in aliases:
                out.append(std)
                break
        else:
            out.append(s)
    return out


def find_header_row(sheet) -> int | None:
    need = ["物品名称", "单位", "数量"]
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if not row or not row[0]:
            continue
        if str(row[0]).strip() not in ("序号", "No.", "编号"):
            continue
        vals = [str(v).strip() for v in row if v is not None]
        if all(any(a in vals for a in COLUMN_MAPPINGS.get(k, [k])) for k in need):
            return idx
        matched = sum(
            1
            for col in STANDARD_COLUMNS
            if any(a in vals for a in COLUMN_MAPPINGS.get(col, [col]))
        )
        if matched >= len(STANDARD_COLUMNS) // 2:
            return idx
    return None


def is_summary_row(ws, row_idx: int, amount_col: int) -> bool:
    cell = ws.cell(row_idx, amount_col)
    if cell.data_type != "f":
        return False
    f = str(cell.value).upper()
    if "SUM" not in f:
        return False
    letter = get_column_letter(amount_col)
    if f"{letter}:{letter}" in f.replace(" ", ""):
        return True
    return bool(re.search(rf"\$?{letter}\$?\d+:\$?{letter}\$?\d+", f.replace(" ", "")))


def _is_blank_cell(v) -> bool:
    """判断读取后是否视为空（含 None / 空白串 / NaN）。"""
    if v is None:
        return True
    if isinstance(v, float) and pd.isna(v):
        return True
    if isinstance(v, str) and v.strip() == "":
        return True
    return False


def vertical_merge_slave_rows(ws, col_1based: int, header_row: int) -> set[int]:
    """
    某列上「纵向合并」的从属行号集合（非合并区左上角所在行）。

    Excel 合并单元格只在左上角存值，从属格 data_only 读出来是空。
    对金额列：拆散后从属格应填 0，以便与「真实空缺」区分，方便后续核对。
    仅处理跨多行的合并；纯横向单行合并不产生从属行。
    """
    slaves: set[int] = set()
    merged = getattr(ws, "merged_cells", None)
    if merged is None:
        return slaves
    for mr in merged.ranges:
        if not (mr.min_col <= col_1based <= mr.max_col):
            continue
        if mr.max_row <= mr.min_row:
            continue
        for r in range(mr.min_row + 1, mr.max_row + 1):
            if r > header_row:
                slaves.add(r)
    return slaves


def vertical_merge_fill_map(ws, col_1based: int, header_row: int) -> dict[int, object]:
    """
    纵向合并：从属行 → 主格（左上角）的值。

    用于文本/枚举列（如货币种类）：拆散后按合并主值填写从属行，
    而非填 0。主格本身为空时不生成映射（从属行保持空）。
    """
    fill: dict[int, object] = {}
    merged = getattr(ws, "merged_cells", None)
    if merged is None:
        return fill
    for mr in merged.ranges:
        if not (mr.min_col <= col_1based <= mr.max_col):
            continue
        if mr.max_row <= mr.min_row:
            continue
        master_val = ws.cell(mr.min_row, col_1based).value
        if _is_blank_cell(master_val):
            continue
        # 文本列规范化为去首尾空白的字符串；数值原样保留
        if isinstance(master_val, str):
            master_val = master_val.strip()
        for r in range(mr.min_row + 1, mr.max_row + 1):
            if r > header_row:
                fill[r] = master_val
    return fill


# 金额类：合并从属格空 → 填 0（区分真·空）
_MERGE_FILL_ZERO_COLS = ("含税金额", "不含税金额", "7%增值税", "单价")
# 文本/枚举类：合并从属格空 → 填主格原值
_MERGE_FILL_MASTER_COLS = ("货币种类",)


def extract_rows(file_path: Path, project_code: str, project_name: str):
    """
    两遍读取：先公式模式找汇总行与合并单元格，再 data_only 取数值。
    第二遍使用 read_only 加速大表。

    合并单元格拆散规则：
    - 金额列（含税金额等）：从属行空 → 填 0（真·空保持空，便于核对）
    - 货币种类等：从属行空 → 填合并主格原值
    """
    fwb = openpyxl.load_workbook(file_path, data_only=False, read_only=False)
    fws = fwb.active
    if fws is None:
        fwb.close()
        return [], "未找到活动工作表"

    header_row = find_header_row(fws)
    if header_row is None:
        fwb.close()
        return [], "未识别到表头"

    headers = normalize_headers([c.value for c in fws[header_row]])
    hmap = {n: i for i, n in enumerate(headers)}
    missing = [c for c in CORE_REQUIRED if c not in hmap]
    if missing:
        fwb.close()
        return [], f"缺少核心列: {', '.join(missing)}"
    if "含税金额" not in hmap:
        fwb.close()
        return [], "缺少含税金额列"

    amount_col = hmap["含税金额"] + 1
    max_row = fws.max_row or header_row
    summary = {
        r for r in range(header_row + 1, max_row + 1) if is_summary_row(fws, r, amount_col)
    }
    # 金额列纵向合并从属行：拆散后空值填 0
    amount_slave_rows: dict[str, set[int]] = {}
    for col_name in _MERGE_FILL_ZERO_COLS:
        if col_name not in hmap:
            continue
        slaves = vertical_merge_slave_rows(fws, hmap[col_name] + 1, header_row)
        if slaves:
            amount_slave_rows[col_name] = slaves
    # 文本列纵向合并：从属行 → 主格值（货币种类等）
    text_merge_fill: dict[str, dict[int, object]] = {}
    for col_name in _MERGE_FILL_MASTER_COLS:
        if col_name not in hmap:
            continue
        fmap = vertical_merge_fill_map(fws, hmap[col_name] + 1, header_row)
        if fmap:
            text_merge_fill[col_name] = fmap
    fwb.close()

    # data_only + read_only：显著加快大文件数值读取
    vwb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    rows: list[dict] = []
    merge_zero_count = 0
    merge_master_count = 0
    try:
        vws = vwb.active
        if vws is None:
            return [], "未找到活动工作表"

        key_cols = [
            hmap[c]
            for c in ("物品名称", "规格", "单位", "数量", "单价", "不含税金额")
            if c in hmap
        ]
        for ridx, vals in enumerate(
            vws.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            if ridx in summary:
                continue
            if not vals:
                continue
            if all(v is None or str(v).strip() == "" for v in vals):
                continue
            if not any(
                (i < len(vals) and vals[i] is not None and str(vals[i]).strip() != "")
                for i in key_cols
            ):
                continue

            obj = {"项目编码": project_code, "项目名称": project_name}
            for col in STANDARD_COLUMNS:
                if col in hmap and hmap[col] < len(vals):
                    v = vals[hmap[col]]
                    if col == "单位":
                        obj[col] = "" if v is None else str(v).strip()
                    else:
                        obj[col] = "" if v is None else v
                else:
                    obj[col] = ""

            # 金额：合并从属格空 → 0（非合并真·空保持空）
            for col_name, slaves in amount_slave_rows.items():
                if ridx in slaves and _is_blank_cell(obj.get(col_name, "")):
                    obj[col_name] = 0
                    merge_zero_count += 1
            # 货币种类等：合并从属格空 → 主格原值
            for col_name, fmap in text_merge_fill.items():
                if ridx in fmap and _is_blank_cell(obj.get(col_name, "")):
                    obj[col_name] = fmap[ridx]
                    merge_master_count += 1
            rows.append(obj)
    finally:
        vwb.close()

    parts = []
    if merge_zero_count:
        parts.append(f"金额填0 {merge_zero_count} 格")
    if merge_master_count:
        cols = "、".join(text_merge_fill.keys())
        parts.append(f"{cols}填主值 {merge_master_count} 格")
    if parts:
        print(f"  · 合并单元格拆散: {'；'.join(parts)}")
    return rows, ""


# ── 日期 ──────────────────────────────────────────────────
def parse_date(value) -> tuple[int | None, int | None, int | None]:
    if value is None:
        return None, None, None
    if isinstance(value, datetime):
        return value.year, value.month, value.day
    if isinstance(value, date_type):
        return value.year, value.month, value.day
    s = str(value).strip()
    if not s:
        return None, None, None
    for pat in (
        r"^(\d{4})/(\d{1,2})/(\d{1,2})",
        r"^(\d{4})-(\d{1,2})-(\d{1,2})",
        r"^(\d{4})\.(\d{1,2})(?:\.(\d{1,2}))?",
    ):
        m = re.match(pat, s)
        if not m:
            continue
        y, mo = int(m.group(1)), int(m.group(2))
        d = int(m.group(3)) if m.lastindex and m.lastindex >= 3 and m.group(3) else None
        if 1 <= mo <= 12 and (d is None or 1 <= d <= 31):
            return y, mo, d
    return None, None, None


def enrich_dates(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        df["标准日期"] = ""
        df["年月"] = ""
        return df

    parsed = [parse_date(v) for v in df.get("签订日期", pd.Series(dtype=object))]
    # 前向/后向填充缺失日期
    filled = list(parsed)
    for i in range(1, len(filled)):
        if filled[i][0] is None:
            for j in range(i - 1, -1, -1):
                if filled[j][0] is not None:
                    filled[i] = filled[j]
                    break
    for i in range(len(filled) - 2, -1, -1):
        if filled[i][0] is None:
            for j in range(i + 1, len(filled)):
                if filled[j][0] is not None:
                    filled[i] = filled[j]
                    break

    std_dates, yms = [], []
    for y, m, d in filled:
        if y is None:
            std_dates.append("")
            yms.append("")
        else:
            std_dates.append(f"{y}-{m:02d}-{d:02d}" if d else f"{y}-{m:02d}")
            yms.append(f"{y}-{m:02d}" if m else "")
    df["标准日期"] = std_dates
    df["年月"] = yms
    return df


# ── 写出 ──────────────────────────────────────────────────
def write_output(df: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    out = df.copy()
    for col in FINAL_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    out = out[FINAL_COLUMNS]
    out.to_excel(output_path, sheet_name="合并数据", index=False)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["合并数据"]
    ws.row_dimensions[1].height = 40
    ws.freeze_panes = "A2"
    for col_idx, header in enumerate(FINAL_COLUMNS, start=1):
        ws.cell(1, col_idx).alignment = Alignment(
            wrap_text=True, horizontal="center", vertical="center"
        )
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, col_idx)
            if header in AMOUNT_COLS and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
            if header in CENTER_COLS:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif header in LEFT_COLS:
                cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(output_path)


def process_one(file_path: Path, projects: list[dict], match_db) -> tuple[bool, str]:
    print(f"\n处理: {file_path.name}")
    raw_name = extract_project_name(file_path.name)
    matched = resolve_project(raw_name, projects, match_db)
    if matched:
        code, name = matched.get("project_code", ""), matched.get("project_name", raw_name)
        src = matched.get("source", "")
        # 已是标准映射时不必反复 add；新模糊命中才写入
        if src not in ("unmatched_keep",) and name and name != raw_name:
            _matcher.add_mapping(match_db, raw_name, name)
        if src and src not in ("alias", "config_exact"):
            print(f"  项目匹配[{src}]: {raw_name} → {name}")
    else:
        code, name = "", raw_name
        _matcher.mark_unmatched(match_db, raw_name)

    try:
        rows, reason = extract_rows(file_path, code, name)
    except Exception as exc:
        print(f"  ✗ {exc}")
        return False, f"处理异常: {exc}"

    if not rows:
        print(f"  ⚠️ {reason or '无有效数据'}")
        return False, reason or "无有效数据"

    df = pd.DataFrame(rows)
    # 金额过滤策略（配合合并单元格拆散填 0）：
    # - 明确 0：合并从属格，保留，便于汇总时不把「无独立金额」当缺失
    # - 真·空（NaN/空串）：保留有物品名称的行，方便后续人工核对
    # - 仅丢弃：无物品名称 且 全部金额为空 的无效行
    amt_cols = [c for c in AMOUNT_COLS if c in df.columns]
    amt = df[amt_cols].apply(pd.to_numeric, errors="coerce") if amt_cols else pd.DataFrame(index=df.index)
    name_s = df["物品名称"] if "物品名称" in df.columns else pd.Series("", index=df.index)
    has_item = name_s.map(
        lambda x: bool(str(x).strip()) and str(x).strip().lower() not in ("nan", "none")
        if pd.notna(x)
        else False
    )
    has_amt = amt.notna().any(axis=1) if not amt.empty else pd.Series(False, index=df.index)
    df = df[has_item | has_amt]
    if df.empty:
        return False, "无有效明细行"

    df = enrich_dates(df)
    file_ym = extract_month_iso(file_path.name) or ""
    if file_ym and "年月" in df.columns:
        empty = df["年月"].fillna("").astype(str).str.strip().eq("")
        if empty.any():
            df.loc[empty, "年月"] = file_ym
            print(f"  ✓ 年月兜底 {int(empty.sum())} 行 -> {file_ym}")

    out_path = STD_DIR / f"{file_path.stem}{OUTPUT_SUFFIX}.xlsx"
    write_output(df, out_path)
    print(f"  ✓ {out_path.name}  ({len(df)} 行)")
    return True, ""


def main(argv: list[str] | None = None) -> None:
    import argparse

    parser = argparse.ArgumentParser(description="采购台账 Excel 格式标准化 → std")
    parser.add_argument(
        "--no-prompt",
        action="store_true",
        help="不弹项目匹配窗：新未匹配名保留原名并记入库，下次也不再弹",
    )
    parser.add_argument(
        "--force-prompt",
        action="store_true",
        help="强制对所有当前无法精确/模糊匹配的名弹窗（忽略历史「保留原名」）",
    )
    args = parser.parse_args(argv)

    if not EXTRACT_DIR.is_dir():
        raise FileNotFoundError(f"提取目录不存在: {EXTRACT_DIR}")
    STD_DIR.mkdir(parents=True, exist_ok=True)

    print(f"输入(提取): {EXTRACT_DIR}")
    print(f"输出(标准): {STD_DIR}")
    print(f"项目匹配: {_MATCHER_PATH}")
    print(f"弹窗: {'关闭' if args.no_prompt else ('强制' if args.force_prompt else '仅新未匹配')}")

    projects = load_project_config()
    match_db = _matcher.load_or_init_db(str(MATCH_DB_PATH), projects)

    files = list_excels(EXTRACT_DIR, recursive=False)
    if not files:
        print("未发现 Excel，结束")
        return

    if args.force_prompt:
        # 清空与当前文件相关的「保留原名」记录，允许重新弹窗
        names = {extract_project_name(f.name) for f in files}
        names = {n for n in names if n}
        before = list(match_db.get("unmatched_aliases", []))
        match_db["unmatched_aliases"] = [
            n
            for n in before
            if n not in names
            and _matcher._normalize_key(n)
            not in {_matcher._normalize_key(x) for x in names}
        ]
        print(
            f"强制弹窗: 已清除 {len(before) - len(match_db['unmatched_aliases'])} 条历史保留原名"
        )

    preload_project_matches(
        files,
        projects,
        match_db,
        interactive=not args.no_prompt,
    )

    ok = skip = 0
    skipped = []
    print(f"\n扫描到 {len(files)} 个文件，开始格式标准化…")
    for f in files:
        success, reason = process_one(f, projects, match_db)
        if success:
            ok += 1
        else:
            skip += 1
            skipped.append((f, reason))

    _matcher.save_db(str(MATCH_DB_PATH), match_db)
    print(f"\n完成: 成功 {ok}，跳过/失败 {skip}")
    for path, reason in skipped:
        print(f"  - {path.name}: {reason}")
    print(f"输出目录: {STD_DIR}")


if __name__ == "__main__":
    main()
