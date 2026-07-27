"""
采购内容标准化工具。

职责：
1. 单位处理（本模块内完成，不再放在 3_excel 格式标准化）
   - fix_unit_qty：单位/数量错位纠错
   - fill_missing_units：空单位按同物品+单价相似行补全
   - build_global_unit_ref：跨文件单位参考库
   - standardize_unit：调用 unit_map_utils，别名 → 标准单位
2. 物品名称/规格规则流水线（规格标准化/*）
"""
from __future__ import annotations

import importlib.util
import re
from collections import Counter
from pathlib import Path

import openpyxl
import pandas as pd

CURRENT_DIR = Path(__file__).resolve().parent
PACKAGE_ROOT = CURRENT_DIR.parent  # 9_采购数据
CODE_ROOT = PACKAGE_ROOT.parent  # Desktop/code
UNIT_MAP_UTILS_PATH = CODE_ROOT / "0_项目基本信息" / "unit_map_utils.py"
STANDARDIZATION_PIPELINE_DIR = PACKAGE_ROOT / "规格标准化"

_PIPELINE_CONTEXT = None
# 跨文件单位兜底：{物品名称: 最常见单位}
_GLOBAL_UNIT_REF: dict[str, str] = {}


def _load_module_from_path(module_name: str, file_path: Path):
    spec = importlib.util.spec_from_file_location(module_name, str(file_path))
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载模块: {file_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


_unit_map_utils = _load_module_from_path("unit_map_utils", UNIT_MAP_UTILS_PATH)


# ── unit_map 封装 ─────────────────────────────────────────
def standardize_unit(value, auto_prompt: bool = True) -> str:
    """别名映射到标准单位（见 0_项目基本信息/unit_map.json）。"""
    if value is None:
        return ""
    return str(_unit_map_utils.standardize_unit(str(value).strip(), auto_prompt=auto_prompt) or "")


def resolve_pending_unit_mappings() -> int:
    """批量弹窗处理本轮暂存的未知单位，写回 unit_map.json。"""
    return int(_unit_map_utils.resolve_pending_unit_mappings())


def set_global_unit_ref(ref: dict[str, str] | None) -> None:
    global _GLOBAL_UNIT_REF
    _GLOBAL_UNIT_REF = dict(ref or {})


def get_global_unit_ref() -> dict[str, str]:
    return dict(_GLOBAL_UNIT_REF)


# ── 单位/数量纠错与补全 ───────────────────────────────────
def is_numeric_value(value) -> bool:
    if value is None:
        return False
    if isinstance(value, (int, float)) and pd.notna(value):
        return True
    s = str(value).strip().replace(",", "")
    if not s:
        return False
    return re.fullmatch(r"[+-]?\d+(?:\.\d+)?", s) is not None


def fix_unit_qty(unit_val, qty_val):
    """
    纠正单位/数量填反或错位。
    1) 单位为空且数量是非数值文本 → 挪到单位列
    2) 单位是数值、数量是非数值文本 → 交换
    """
    unit_text = "" if unit_val is None else str(unit_val).strip()
    qty_text = "" if qty_val is None else str(qty_val).strip()

    if unit_text == "" and qty_text and not is_numeric_value(qty_val):
        return qty_text, ""
    if unit_text and qty_text and is_numeric_value(unit_val) and not is_numeric_value(qty_val):
        return qty_text, unit_val
    return unit_val, qty_val


def build_global_unit_ref(std_root: str | Path) -> dict[str, str]:
    """
    扫描目录下 Excel，统计每个物品名称出现最多的单位。
    返回 {物品名称: 单位}，供空单位兜底。
    """
    root = Path(std_root)
    ref: dict[str, Counter] = {}
    if not root.is_dir():
        return {}

    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if path.name.startswith("~$") or path.name.startswith("."):
            continue
        if path.suffix.lower() not in {".xlsx", ".xls", ".xlsm"}:
            continue
        try:
            df = pd.read_excel(path, sheet_name=0, usecols=lambda c: str(c).strip() in {"物品名称", "单位"}, dtype=str)
        except Exception:
            continue
        if "物品名称" not in df.columns or "单位" not in df.columns:
            continue
        df = df.dropna(subset=["物品名称"])
        df["物品名称"] = df["物品名称"].astype(str).str.strip()
        df["单位"] = df["单位"].fillna("").astype(str).str.strip()
        for item_name, grp in df[df["单位"].ne("")].groupby("物品名称"):
            ref.setdefault(str(item_name), Counter()).update(grp["单位"].tolist())

    return {k: v.most_common(1)[0][0] for k, v in ref.items() if v}


def _header_map(ws) -> dict[str, int]:
    return {
        str(cell.value).strip(): idx
        for idx, cell in enumerate(ws[1], start=1)
        if cell.value is not None and str(cell.value).strip()
    }


def _within_price_tolerance(base_price, candidate_price, tolerance: float = 0.50) -> bool:
    if base_price is None or candidate_price is None:
        return False
    try:
        base = float(base_price)
        cand = float(candidate_price)
    except (TypeError, ValueError):
        return False
    if base == 0:
        return abs(cand - base) < 1e-9
    return abs(cand - base) / abs(base) <= tolerance


def process_units_on_worksheet(
    worksheet,
    *,
    global_unit_ref: dict[str, str] | None = None,
    auto_prompt: bool = False,
) -> dict:
    """
    对单表执行：错位纠错 → 空单位补全 → unit_map 标准化。
    需要列：单位、数量；补全还需要 物品名称/单价（可选 项目名称/标准日期）。
    """
    hmap = _header_map(worksheet)
    unit_col = hmap.get("单位") or hmap.get("标准单位")
    qty_col = hmap.get("数量")
    name_col = hmap.get("物品名称")
    price_col = hmap.get("单价")
    project_col = hmap.get("项目名称")
    date_col = hmap.get("标准日期") or hmap.get("签订日期")

    stats = {
        "sheet_name": worksheet.title,
        "fixed_swap": 0,
        "filled_units": 0,
        "mapped_units": 0,
        "skipped": False,
    }

    if unit_col is None or qty_col is None:
        stats["skipped"] = True
        return stats

    max_row = worksheet.max_row
    # 先读出行数据，便于补全时做相似度匹配
    rows_cache: list[dict] = []
    for r in range(2, max_row + 1):
        unit_v = worksheet.cell(r, unit_col).value
        qty_v = worksheet.cell(r, qty_col).value
        new_unit, new_qty = fix_unit_qty(unit_v, qty_v)
        if (new_unit, new_qty) != (unit_v, qty_v):
            worksheet.cell(r, unit_col).value = new_unit
            worksheet.cell(r, qty_col).value = new_qty
            stats["fixed_swap"] += 1
            unit_v, qty_v = new_unit, new_qty

        name = ""
        if name_col:
            nv = worksheet.cell(r, name_col).value
            name = "" if nv is None else str(nv).strip()
        price = None
        if price_col:
            price = pd.to_numeric(pd.Series([worksheet.cell(r, price_col).value]), errors="coerce").iloc[0]
            if pd.isna(price):
                price = None
            else:
                price = float(price)
        project = ""
        if project_col:
            pv = worksheet.cell(r, project_col).value
            project = "" if pv is None else str(pv).strip()
        date_ts = None
        if date_col:
            date_ts = pd.to_datetime(worksheet.cell(r, date_col).value, errors="coerce")
            if pd.isna(date_ts):
                date_ts = None

        unit_text = "" if unit_v is None else str(unit_v).strip()
        rows_cache.append(
            {
                "row": r,
                "name": name,
                "unit": unit_text,
                "price": price,
                "project": project,
                "date_ts": date_ts,
            }
        )

    # 空单位补全
    if name_col and price_col:
        # 各物品平均单价
        avg_price: dict[str, float] = {}
        for name, grp in _groupby_name(rows_cache):
            prices = [x["price"] for x in grp if x["price"] is not None and x["price"] != 0]
            if prices:
                avg_price[name] = sum(prices) / len(prices)

        candidates = [x for x in rows_cache if x["unit"] and x["name"] and x["price"] is not None]
        for item in rows_cache:
            if item["unit"] or not item["name"] or item["price"] is None:
                continue
            name = item["name"]
            base_avg = avg_price.get(name)
            pool = [
                c
                for c in candidates
                if c["name"] == name and _within_price_tolerance(base_avg, c["price"], 0.50)
            ]
            best_unit = ""
            if pool:
                target_proj = item["project"]
                target_ts = item["date_ts"]
                ranked = []
                for c in pool:
                    same = 0 if c["project"] == target_proj else 1
                    if target_ts is not None and c["date_ts"] is not None:
                        gap = abs((c["date_ts"] - target_ts).days)
                    else:
                        gap = 10**9
                    ranked.append((same, gap, c["unit"]))
                ranked.sort(key=lambda x: (x[0], x[1]))
                best_unit = ranked[0][2] if ranked else ""
            if not best_unit and global_unit_ref:
                best_unit = str(global_unit_ref.get(name, "") or "").strip()
            if best_unit:
                worksheet.cell(item["row"], unit_col).value = best_unit
                item["unit"] = best_unit
                stats["filled_units"] += 1

    # unit_map 标准化（auto_prompt=False 时未知单位先原样，可稍后 resolve_pending）
    for item in rows_cache:
        r = item["row"]
        raw = worksheet.cell(r, unit_col).value
        if raw is None or str(raw).strip() == "":
            continue
        mapped = standardize_unit(raw, auto_prompt=auto_prompt)
        if str(mapped) != str(raw).strip():
            worksheet.cell(r, unit_col).value = mapped
            stats["mapped_units"] += 1
        else:
            # 即使相同也可能是 canonicalize；仍写回标准写法
            if mapped != raw:
                worksheet.cell(r, unit_col).value = mapped

    return stats


def _groupby_name(rows: list[dict]):
    buckets: dict[str, list] = {}
    for r in rows:
        if r["name"]:
            buckets.setdefault(r["name"], []).append(r)
    return buckets.items()


def process_units_on_workbook(
    workbook,
    *,
    global_unit_ref: dict[str, str] | None = None,
    auto_prompt: bool = False,
) -> dict:
    totals = {"fixed_swap": 0, "filled_units": 0, "mapped_units": 0, "sheets": 0}
    ref = global_unit_ref if global_unit_ref is not None else _GLOBAL_UNIT_REF
    for ws in workbook.worksheets:
        st = process_units_on_worksheet(ws, global_unit_ref=ref, auto_prompt=auto_prompt)
        if st.get("skipped"):
            continue
        totals["sheets"] += 1
        totals["fixed_swap"] += st["fixed_swap"]
        totals["filled_units"] += st["filled_units"]
        totals["mapped_units"] += st["mapped_units"]
    return totals


# ── 规格规则流水线 ────────────────────────────────────────
def _get_pipeline_context():
    global _PIPELINE_CONTEXT
    if _PIPELINE_CONTEXT is not None:
        return _PIPELINE_CONTEXT

    pipeline_script = STANDARDIZATION_PIPELINE_DIR / "0_统一标准化流水线.py"
    if not pipeline_script.exists():
        _PIPELINE_CONTEXT = {"ok": False, "reason": f"未找到标准化流水线脚本: {pipeline_script}"}
        return _PIPELINE_CONTEXT

    spec = importlib.util.spec_from_file_location("_standardization_pipeline", str(pipeline_script))
    if spec is None or spec.loader is None:
        _PIPELINE_CONTEXT = {"ok": False, "reason": f"无法加载标准化流水线: {pipeline_script}"}
        return _PIPELINE_CONTEXT

    pipeline_module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(pipeline_module)

    rule_names = getattr(pipeline_module, "DEFAULT_RULES", [])
    if not rule_names:
        _PIPELINE_CONTEXT = {"ok": False, "reason": "标准化流水线 DEFAULT_RULES 为空"}
        return _PIPELINE_CONTEXT

    rule_modules = []
    for rule_name in rule_names:
        rule_file = STANDARDIZATION_PIPELINE_DIR / rule_name
        if not rule_file.exists():
            continue
        try:
            rule_module = pipeline_module.load_module_from_file(rule_file)
            rule_modules.append((rule_file.name, rule_module))
        except Exception:
            continue

    _PIPELINE_CONTEXT = {
        "ok": True,
        "pipeline_module": pipeline_module,
        "rule_modules": rule_modules,
    }
    return _PIPELINE_CONTEXT


def apply_standardization_pipeline(
    excel_path: str,
    *,
    global_unit_ref: dict[str, str] | None = None,
    auto_prompt_units: bool = False,
    resolve_pending: bool = False,
) -> None:
    """
    对单个 Excel：
      1) 单位错位纠错 / 空单位补全 / unit_map 标准化
      2) 规格标准化规则（桩/钢筋/网片/砂浆…）
    """
    workbook = openpyxl.load_workbook(excel_path)

    unit_stats = process_units_on_workbook(
        workbook,
        global_unit_ref=global_unit_ref,
        auto_prompt=auto_prompt_units,
    )
    print(
        f"  ✓ 单位处理: 错位 {unit_stats['fixed_swap']} 行, "
        f"补全 {unit_stats['filled_units']} 行, "
        f"映射 {unit_stats['mapped_units']} 行 "
        f"(有效表 {unit_stats['sheets']})"
    )

    context = _get_pipeline_context()
    if not context.get("ok"):
        print(f"  ⚠️ {context.get('reason', '规格流水线不可用，已跳过')}")
    else:
        pipeline_module = context["pipeline_module"]
        for rule_name, rule_module in context["rule_modules"]:
            try:
                stats = pipeline_module.run_rule_on_workbook(rule_module, workbook, rule_name)
                print(
                    f"  ✓ {rule_name}: "
                    f"命中 {stats['matched']} 行, 更新 {stats['updated']} 行, "
                    f"未识别 {stats['unresolved']} 行"
                )
            except Exception as exc:
                print(f"  ✗ 规则执行失败 {rule_name}: {exc}")

    workbook.save(excel_path)

    if resolve_pending:
        n = resolve_pending_unit_mappings()
        if n:
            print(f"  ✓ 已批量确认未知单位 {n} 个")
            # 二次应用映射（用户刚写入的）
            workbook = openpyxl.load_workbook(excel_path)
            process_units_on_workbook(
                workbook,
                global_unit_ref=global_unit_ref,
                auto_prompt=False,
            )
            workbook.save(excel_path)
