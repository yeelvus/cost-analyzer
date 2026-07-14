import os
import shutil
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

# 允许脚本从 code/ 下直接运行时导入 common 包
_CODE_ROOT = Path(__file__).resolve().parents[1]
if str(_CODE_ROOT) not in sys.path:
    sys.path.insert(0, str(_CODE_ROOT))

from common.excel import apply_common_excel_layout

# 仅保留当前在用的直连汇总能力：
# 输入目录（递归汇总全部 Excel）-> 输出目录（单个汇总文件）
DIRECT_STD_INPUT_ROOT = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/9_采购台账📦/B_采购部采购台账_std"
DIRECT_OUTPUT_ROOT = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/9_采购台账📦/D_采购部采购台账_sum"
DIRECT_OUTPUT_FILENAME_WITH_STD = "采购部采购台账_std_merged.xlsx"
DIRECT_OUTPUT_FILENAME_WITHOUT_STD = "采购部采购台账_merged.xlsx"

FINAL_OUTPUT_COLUMNS = [
    "项目编码",
    "项目名称",
    "材料归类",
    "物品名称",
    "规格",
    "单位",
    "数量",
    "单价",
    "不含税金额",
    "7%增值税",
    "含税金额",
    "货币种类",
    "汇率",
    "年月",
    "标准日期",
    "供货商名称",
    "备注",
    "签订日期",
]


def _load_procurement_standardization_utils():
    """动态加载采购内容标准化工具。"""
    utils_path = Path(__file__).resolve().parent / "采购内容标准化_utils.py"
    if not utils_path.exists():
        raise FileNotFoundError(f"未找到标准化工具: {utils_path}")

    import importlib.util

    spec = importlib.util.spec_from_file_location("procurement_content_std_utils", str(utils_path))
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载标准化工具: {utils_path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def apply_unified_workbook_format(output_wb):
    """统一工作簿样式。"""
    amount_keywords = ("金额", "单价", "税")

    for sheet_name in ["合并数据", "项目汇总", "采购总金额"]:
        if sheet_name not in output_wb.sheetnames:
            continue

        ws = output_wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        amount_headers = [
            h for h in headers if isinstance(h, str) and any(k in h for k in amount_keywords)
        ]

        if sheet_name == "采购总金额":
            apply_common_excel_layout(
                ws,
                numeric_start_col=3,
                amount_headers=[h for h in headers[2:] if isinstance(h, str)],
                header_row=1,
                header_height=40,
                freeze_panes="A2",
            )
        else:
            apply_common_excel_layout(
                ws,
                numeric_start_col=ws.max_column + 1,
                amount_headers=amount_headers,
                header_row=1,
                header_height=40,
                freeze_panes="A2",
            )


def _period_key(period_text):
    y, m = str(period_text).split("-")
    return int(y), int(m)


def build_procurement_total_sheet(merged_df):
    """由合并数据构建采购总金额宽表（项目 x 年月）。"""
    required = {"项目编码", "项目名称", "年月", "含税金额"}
    if not required.issubset(set(merged_df.columns)):
        return pd.DataFrame(columns=["项目编码", "项目名称"])

    work = merged_df[["项目编码", "项目名称", "年月", "含税金额"]].copy()
    work["项目编码"] = work["项目编码"].fillna("").astype(str).str.strip()
    work["项目名称"] = work["项目名称"].fillna("").astype(str).str.strip()
    work["年月"] = work["年月"].fillna("").astype(str).str.strip()
    work["含税金额"] = pd.to_numeric(work["含税金额"], errors="coerce")

    work = work[(work["项目编码"] != "") | (work["项目名称"] != "")]
    work = work[work["年月"].str.fullmatch(r"(?:19|20)\d{2}-(?:0[1-9]|1[0-2])", na=False)]
    work = work[work["含税金额"].notna()]
    if work.empty:
        return pd.DataFrame(columns=["项目编码", "项目名称"])

    grouped = (
        work.groupby(["项目编码", "项目名称", "年月"], as_index=False)
        .agg(含税金额=("含税金额", "sum"))
    )

    periods = sorted(grouped["年月"].unique().tolist(), key=_period_key)
    label_map = {p: f"{int(p.split('-')[0])}年{int(p.split('-')[1])}月" for p in periods}

    pivot = grouped.pivot_table(
        index=["项目编码", "项目名称"],
        columns="年月",
        values="含税金额",
        aggfunc="sum",
        fill_value=0.0,
    ).reindex(columns=periods, fill_value=0.0)

    amount_wide = pivot.rename(columns=label_map).reset_index()
    sort_cols = [label_map[p] for p in reversed(periods)]
    amount_wide["__编码为空"] = amount_wide["项目编码"].astype(str).str.strip().eq("")
    amount_wide = amount_wide.sort_values(
        ["__编码为空", *sort_cols],
        ascending=[True] + [False] * len(sort_cols),
    )
    amount_wide = amount_wide.drop(columns=["__编码为空"]).reset_index(drop=True)
    return amount_wide


def list_source_excel_files_recursive(folder_path):
    """递归列出待合并的源 Excel 文件（绝对路径）。"""
    excel_files = []
    for root, _, files in os.walk(folder_path):
        for f in files:
            lower = f.lower()
            if not (lower.endswith(".xlsx") or lower.endswith(".xls") or lower.endswith(".xlsm")):
                continue
            if f == "merged_excel.xlsx":
                continue
            if f.startswith("~$") or f.startswith("merged_excel") or f.endswith("_std_merged.xlsx"):
                continue
            excel_files.append(os.path.join(root, f))
    return sorted(excel_files)


def merge_standardized_excels(folder_path, use_content_standardization=False):
    """汇总目录下全部 Excel，输出工作簿对象。"""
    excel_paths = list_source_excel_files_recursive(folder_path)
    std_utils = None
    if use_content_standardization:
        std_utils = _load_procurement_standardization_utils()

    print("\n" + "=" * 80)
    mode_name = "启用采购内容标准化" if use_content_standardization else "不启用采购内容标准化"
    print(f"【纯汇总模式 - {mode_name}】开始处理文件夹: {folder_path}")
    print(f"【纯汇总模式 - {mode_name}】发现 {len(excel_paths)} 个Excel文件")
    print("=" * 80 + "\n")

    if not excel_paths:
        print("⚠️  警告：未找到任何Excel文件！")
        return None

    output_wb = openpyxl.Workbook()
    merged_ws = output_wb.active
    if merged_ws is None:
        raise RuntimeError("无法创建工作表: 合并数据")
    merged_ws.title = "合并数据"
    merged_ws.append(FINAL_OUTPUT_COLUMNS)

    current_row = 2
    project_totals = defaultdict(float)

    for idx, file_path in enumerate(excel_paths, start=1):
        file_name = os.path.basename(file_path)
        print(f"[{idx}/{len(excel_paths)}] 汇总文件: {file_name}")

        read_path = file_path
        temp_dir = None
        if use_content_standardization and std_utils is not None:
            try:
                temp_dir = tempfile.mkdtemp(prefix="procurement_std_")
                temp_file = os.path.join(temp_dir, file_name)
                shutil.copy2(file_path, temp_file)
                std_utils.apply_standardization_pipeline(temp_file)
                read_path = temp_file
            except Exception as exc:
                print(f"  ⚠️ 标准化失败，回退原文件: {exc}")
                read_path = file_path

        try:
            df = pd.read_excel(read_path)
        except Exception as e:
            print(f"  ✗ 读取失败: {e}")
            if temp_dir and os.path.isdir(temp_dir):
                shutil.rmtree(temp_dir, ignore_errors=True)
            continue
        finally:
            if temp_dir and os.path.isdir(temp_dir):
                shutil.rmtree(temp_dir, ignore_errors=True)

        if df.empty:
            print("  ⚠️ 空文件，跳过")
            continue

        for col in FINAL_OUTPUT_COLUMNS:
            if col not in df.columns:
                df[col] = ""
        df = df[FINAL_OUTPUT_COLUMNS]

        for _, row in df.iterrows():
            for col_idx, col_name in enumerate(FINAL_OUTPUT_COLUMNS, start=1):
                value = row[col_name]
                merged_ws.cell(current_row, col_idx).value = value
                if col_name in {"单价", "不含税金额", "7%增值税", "含税金额"} and isinstance(value, (int, float)):
                    merged_ws.cell(current_row, col_idx).number_format = "0.00"

            project_code = "" if pd.isna(row["项目编码"]) else str(row["项目编码"]).strip()
            project_name = "" if pd.isna(row["项目名称"]) else str(row["项目名称"]).strip()
            amount_value = pd.to_numeric(pd.Series([row["含税金额"]]), errors="coerce").iloc[0]
            if pd.notna(amount_value):
                project_totals[(project_code, project_name)] += float(amount_value)

            current_row += 1

    if current_row == 2:
        print("⚠️  没有找到任何有效数据！")
        return None

    merged_ws.row_dimensions[1].height = 40
    for cell in merged_ws[1]:
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    merged_ws.freeze_panes = "A2"
    last_col = get_column_letter(len(FINAL_OUTPUT_COLUMNS))
    merged_ws.auto_filter.ref = f"A1:{last_col}{current_row - 1}"

    summary_ws = output_wb.create_sheet("项目汇总")
    summary_ws.append(["项目编码", "项目名称", "总采购金额"])
    for row_idx, ((project_code, project_name), total) in enumerate(sorted(project_totals.items()), start=2):
        summary_ws.append([project_code, project_name, total])
        summary_ws.cell(row_idx, 3).number_format = "0.00"

    merged_df = pd.DataFrame(merged_ws.values)
    merged_df.columns = merged_df.iloc[0]
    merged_df = merged_df.iloc[1:].reset_index(drop=True)

    procurement_total_df = build_procurement_total_sheet(merged_df)
    total_ws = output_wb.create_sheet("采购总金额")
    total_ws.append(procurement_total_df.columns.tolist())
    for row in procurement_total_df.itertuples(index=False, name=None):
        total_ws.append(list(row))

    return output_wb


def main():
    standardized_root = DIRECT_STD_INPUT_ROOT
    out_root = DIRECT_OUTPUT_ROOT

    if not os.path.isdir(standardized_root):
        raise FileNotFoundError(f"输入目录不存在: {standardized_root}")

    os.makedirs(out_root, exist_ok=True)
    print("\n直连模式：")
    print(f"- 输入目录: {standardized_root}")
    print(f"- 输出目录: {out_root}")

    version_configs = [
        (False, DIRECT_OUTPUT_FILENAME_WITHOUT_STD),
        (True, DIRECT_OUTPUT_FILENAME_WITH_STD),
    ]

    for use_std, output_filename in version_configs:
        output_wb = merge_standardized_excels(
            standardized_root,
            use_content_standardization=use_std,
        )
        if output_wb is None:
            print(f"未产生有效汇总数据: {output_filename}")
            continue

        output_path = os.path.join(out_root, output_filename)
        apply_unified_workbook_format(output_wb)
        output_wb.save(output_path)
        print(f"已输出（直连汇总）: {output_path}")


if __name__ == "__main__":
    main()
