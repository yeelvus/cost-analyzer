#!/usr/bin/env python3
"""
步骤3：读取 EXTRACT_DIR 去重后的台账 → 列映射/项目匹配/日期
→ 写出 STD_DIR/*_processed.xlsx。

不做：
- 单位纠错/补全
- 物品规格等内容规则标准化（桩/钢筋/网片/砂浆等由其他脚本处理）
"""
from __future__ import annotations

import re
import sys
from datetime import date as date_type
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

_CODE_ROOT = Path(__file__).resolve().parents[1]
if str(_CODE_ROOT) not in sys.path:
    sys.path.insert(0, str(_CODE_ROOT))

from settings import MATCH_DB_PATH

from _采购路径 import (
    EXTRACT_DIR,
    PROJECT_INFO_FILE,
    STD_DIR,
    extract_month_iso,
    extract_project_name,
    list_excels,
)

# ── 项目匹配 ──────────────────────────────────────────────
_MATCHER_PATH = _CODE_ROOT / "0_项目基本信息" / "project_name_matcher.py"


def _load_matcher():
    import importlib.util

    spec = importlib.util.spec_from_file_location("_proj_matcher", str(_MATCHER_PATH))
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载: {_MATCHER_PATH}")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


_matcher = _load_matcher()

# ── 列配置 ────────────────────────────────────────────────
COLUMN_MAPPINGS = {
    "材料归类": ["材料归类"],
    "物品名称": ["物品名称", "采购物资名称", "物资名称"],
    "规格": ["规格", "规格型号", "型号", "规格型号/参数"],
    "单位": ["标准单位", "单位"],
    "数量": ["数量"],
    "单价": ["单价", "不含税单价", "含税单价"],
    "不含税金额": ["不含税金额", "金额"],
    "7%增值税": ["7%增值税", "税额", "增值税", "税金"],
    "含税金额": ["含税金额", "总金额", "价税合计"],
    "货币种类": ["货币种类", "币种"],
    "汇率": ["汇率"],
    "签订日期": ["签订日期", "日期"],
    "供货商名称": ["供货商名称", "供货单位", "供应商", "供应商名称"],
    "备注": ["备注"],
}
STANDARD_COLUMNS = list(COLUMN_MAPPINGS.keys())
CORE_REQUIRED = ["物品名称", "单位", "数量"]
FINAL_COLUMNS = [
    "项目编码", "项目名称", "材料归类", "物品名称", "规格", "单位", "数量",
    "单价", "不含税金额", "7%增值税", "含税金额", "货币种类", "汇率",
    "标准日期", "供货商名称", "备注", "年月", "签订日期",
]
AMOUNT_COLS = {"单价", "不含税金额", "7%增值税", "含税金额"}
CENTER_COLS = {"单位", "数量", "货币种类", "汇率"}
OUTPUT_SUFFIX = "_processed"


# ── 项目配置 / 匹配 ───────────────────────────────────────
def load_project_config() -> list[dict]:
    try:
        df = pd.read_excel(PROJECT_INFO_FILE, sheet_name=0)
    except Exception as exc:
        print(f"警告：无法加载项目配置: {exc}")
        return []
    if not {"项目编码", "项目名称"}.issubset(df.columns):
        return []
    projects = []
    for _, row in df.iterrows():
        code = str(row["项目编码"]).strip() if pd.notna(row["项目编码"]) else ""
        name = str(row["项目名称"]).strip() if pd.notna(row["项目名称"]) else ""
        if code or name:
            projects.append({"project_code": code, "project_name": name})
    return projects


def resolve_project(raw_name: str, projects: list[dict], match_db) -> dict | None:
    hit = _matcher.match_project(raw_name, match_db)
    if hit:
        return {
            "project_code": hit.get("project_code", ""),
            "project_name": hit.get("project_name", raw_name),
        }
    for p in projects:
        if p["project_name"] == raw_name:
            return p
    for p in projects:
        if raw_name in p["project_name"] or p["project_name"] in raw_name:
            return p
    return None


def preload_project_matches(files: list[Path], projects: list[dict], match_db) -> None:
    unmatched = []
    for f in files:
        name = extract_project_name(f.name)
        if resolve_project(name, projects, match_db):
            continue
        if name not in unmatched:
            unmatched.append(name)
    if not unmatched:
        return

    print(f"\n发现 {len(unmatched)} 个未识别项目，打开批量匹配…")
    results = _matcher.prompt_unmatched_project_mappings(
        unmatched, match_db, window_title="采购-批量项目匹配"
    )
    lookup = _matcher.get_project_lookup(match_db)
    for raw, std in results.items():
        if std:
            _matcher.add_mapping(match_db, raw, std)
            code, _ = lookup.get(std, ("", ""))
            print(f"  ✓ {raw} -> [{code}] {std}")
        else:
            _matcher.mark_unmatched(match_db, raw)
            print(f"  - 留空: {raw}")


# ── 表头 / 行提取 ─────────────────────────────────────────
def normalize_headers(headers) -> list[str]:
    out = []
    for col in headers:
        s = str(col).strip() if col else ""
        for std, aliases in COLUMN_MAPPINGS.items():
            if s in aliases:
                out.append(std)
                break
        else:
            out.append(s)
    return out


def find_header_row(sheet) -> int | None:
    need = ["物品名称", "单位", "数量"]
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if not row or not row[0]:
            continue
        if str(row[0]).strip() not in ("序号", "No.", "编号"):
            continue
        vals = [str(v).strip() for v in row if v is not None]
        if all(any(a in vals for a in COLUMN_MAPPINGS.get(k, [k])) for k in need):
            return idx
        matched = sum(
            1
            for col in STANDARD_COLUMNS
            if any(a in vals for a in COLUMN_MAPPINGS.get(col, [col]))
        )
        if matched >= len(STANDARD_COLUMNS) // 2:
            return idx
    return None


def is_summary_row(ws, row_idx: int, amount_col: int) -> bool:
    cell = ws.cell(row_idx, amount_col)
    if cell.data_type != "f":
        return False
    f = str(cell.value).upper()
    if "SUM" not in f:
        return False
    letter = get_column_letter(amount_col)
    if f"{letter}:{letter}" in f.replace(" ", ""):
        return True
    return bool(re.search(rf"\$?{letter}\$?\d+:\$?{letter}\$?\d+", f.replace(" ", "")))


def extract_rows(file_path: Path, project_code: str, project_name: str):
    fwb = openpyxl.load_workbook(file_path, data_only=False)
    fws = fwb.active
    if fws is None:
        fwb.close()
        return [], "未找到活动工作表"

    header_row = find_header_row(fws)
    if header_row is None:
        fwb.close()
        return [], "未识别到表头"

    headers = normalize_headers([c.value for c in fws[header_row]])
    hmap = {n: i for i, n in enumerate(headers)}
    missing = [c for c in CORE_REQUIRED if c not in hmap]
    if missing:
        fwb.close()
        return [], f"缺少核心列: {', '.join(missing)}"
    if "含税金额" not in hmap:
        fwb.close()
        return [], "缺少含税金额列"

    amount_col = hmap["含税金额"] + 1
    summary = {
        r for r in range(header_row + 1, fws.max_row + 1) if is_summary_row(fws, r, amount_col)
    }
    fwb.close()

    vwb = openpyxl.load_workbook(file_path, data_only=True)
    vws = vwb.active
    if vws is None:
        vwb.close()
        return [], "未找到活动工作表"

    key_cols = [hmap[c] for c in ("物品名称", "规格", "单位", "数量", "单价", "不含税金额") if c in hmap]
    rows = []
    for ridx, vals in enumerate(
        vws.iter_rows(min_row=header_row + 1, max_row=vws.max_row, values_only=True),
        start=header_row + 1,
    ):
        if ridx in summary:
            continue
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        if not any(
            (vals[i] is not None and str(vals[i]).strip() != "")
            for i in key_cols
            if i < len(vals)
        ):
            continue

        obj = {"项目编码": project_code, "项目名称": project_name}
        for col in STANDARD_COLUMNS:
            if col in hmap and hmap[col] < len(vals):
                v = vals[hmap[col]]
                # 单位原样保留（仅 strip 字符串空白），不做纠错/补全
                if col == "单位":
                    obj[col] = "" if v is None else str(v).strip()
                else:
                    obj[col] = "" if v is None else v
            else:
                obj[col] = ""
        rows.append(obj)

    vwb.close()
    return rows, ""


# ── 日期 ──────────────────────────────────────────────────
def parse_date(value) -> tuple[int | None, int | None, int | None]:
    if value is None:
        return None, None, None
    if isinstance(value, datetime):
        return value.year, value.month, value.day
    if isinstance(value, date_type):
        return value.year, value.month, value.day
    s = str(value).strip()
    if not s:
        return None, None, None
    for pat in (
        r"^(\d{4})/(\d{1,2})/(\d{1,2})",
        r"^(\d{4})-(\d{1,2})-(\d{1,2})",
        r"^(\d{4})\.(\d{1,2})(?:\.(\d{1,2}))?",
    ):
        m = re.match(pat, s)
        if not m:
            continue
        y, mo = int(m.group(1)), int(m.group(2))
        d = int(m.group(3)) if m.lastindex and m.lastindex >= 3 and m.group(3) else None
        if 1 <= mo <= 12 and (d is None or 1 <= d <= 31):
            return y, mo, d
    return None, None, None


def enrich_dates(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        df["标准日期"] = ""
        df["年月"] = ""
        return df

    parsed = [parse_date(v) for v in df.get("签订日期", pd.Series(dtype=object))]
    # 前向/后向填充缺失日期
    filled = list(parsed)
    for i in range(1, len(filled)):
        if filled[i][0] is None:
            for j in range(i - 1, -1, -1):
                if filled[j][0] is not None:
                    filled[i] = filled[j]
                    break
    for i in range(len(filled) - 2, -1, -1):
        if filled[i][0] is None:
            for j in range(i + 1, len(filled)):
                if filled[j][0] is not None:
                    filled[i] = filled[j]
                    break

    std_dates, yms = [], []
    for y, m, d in filled:
        if y is None:
            std_dates.append("")
            yms.append("")
        else:
            std_dates.append(f"{y}-{m:02d}-{d:02d}" if d else f"{y}-{m:02d}")
            yms.append(f"{y}-{m:02d}" if m else "")
    df["标准日期"] = std_dates
    df["年月"] = yms
    return df


# ── 写出 ──────────────────────────────────────────────────
def write_output(df: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    out = df.copy()
    for col in FINAL_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    out = out[FINAL_COLUMNS]
    out.to_excel(output_path, sheet_name="合并数据", index=False)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["合并数据"]
    ws.row_dimensions[1].height = 40
    ws.freeze_panes = "A2"
    for col_idx, header in enumerate(FINAL_COLUMNS, start=1):
        ws.cell(1, col_idx).alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, col_idx)
            if header in AMOUNT_COLS and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
            if header in CENTER_COLS:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif header == "签订日期":
                cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    wb.save(output_path)


def process_one(file_path: Path, projects: list[dict], match_db) -> tuple[bool, str]:
    print(f"\n处理: {file_path.name}")
    raw_name = extract_project_name(file_path.name)
    matched = resolve_project(raw_name, projects, match_db)
    if matched:
        code, name = matched.get("project_code", ""), matched.get("project_name", raw_name)
        _matcher.add_mapping(match_db, raw_name, name)
    else:
        code, name = "", raw_name
        _matcher.mark_unmatched(match_db, raw_name)

    try:
        rows, reason = extract_rows(file_path, code, name)
    except Exception as exc:
        print(f"  ✗ {exc}")
        return False, f"处理异常: {exc}"

    if not rows:
        print(f"  ⚠️ {reason or '无有效数据'}")
        return False, reason or "无有效数据"

    df = pd.DataFrame(rows)
    amt = df[[c for c in AMOUNT_COLS if c in df.columns]].apply(pd.to_numeric, errors="coerce")
    df = df[amt.fillna(0).sum(axis=1).ne(0)]
    if df.empty:
        return False, "金额合计均为0"

    df = enrich_dates(df)
    file_ym = extract_month_iso(file_path.name) or ""
    if file_ym and "年月" in df.columns:
        empty = df["年月"].fillna("").astype(str).str.strip().eq("")
        if empty.any():
            df.loc[empty, "年月"] = file_ym
            print(f"  ✓ 年月兜底 {int(empty.sum())} 行 -> {file_ym}")

    out_path = STD_DIR / f"{file_path.stem}{OUTPUT_SUFFIX}.xlsx"
    write_output(df, out_path)
    print(f"  ✓ {out_path.name}")
    return True, ""


def main() -> None:
    if not EXTRACT_DIR.is_dir():
        raise FileNotFoundError(f"提取目录不存在: {EXTRACT_DIR}")
    STD_DIR.mkdir(parents=True, exist_ok=True)

    print(f"输入(提取): {EXTRACT_DIR}")
    print(f"输出(标准): {STD_DIR}")
    print(f"项目匹配: {_MATCHER_PATH}")

    projects = load_project_config()
    match_db = _matcher.load_or_init_db(str(MATCH_DB_PATH), projects)

    files = list_excels(EXTRACT_DIR, recursive=False)
    if not files:
        print("未发现 Excel，结束")
        return

    preload_project_matches(files, projects, match_db)

    ok = skip = 0
    skipped = []
    print(f"\n扫描到 {len(files)} 个文件，开始格式标准化…")
    for f in files:
        success, reason = process_one(f, projects, match_db)
        if success:
            ok += 1
        else:
            skip += 1
            skipped.append((f, reason))

    _matcher.save_db(str(MATCH_DB_PATH), match_db)
    print(f"\n完成: 成功 {ok}，跳过/失败 {skip}")
    for path, reason in skipped:
        print(f"  - {path.name}: {reason}")
    print(f"输出目录: {STD_DIR}")


if __name__ == "__main__":
    main()
