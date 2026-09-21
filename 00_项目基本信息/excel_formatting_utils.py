from __future__ import annotations

from typing import Iterable
from openpyxl.styles import Alignment, Font, PatternFill


def _display_len(text: str) -> int:
    """估算显示宽度：中文按2个字符宽，英文按1个字符宽。"""
    total = 0
    for ch in text:
        total += 2 if ord(ch) > 127 else 1
    return total


def _safe_text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def apply_common_excel_layout(
    worksheet,
    *,
    numeric_start_col: int = 3,
    amount_headers: Iterable[str] | None = None,
    percent_headers: Iterable[str] | None = None,
    header_row: int = 1,
    header_height: float = 40,
    freeze_panes: str = "A2",
    thousand_number_format: str = "#,##0.00",
    percent_number_format: str = "0.00%",
    min_col_width: float = 10,
    max_col_width: float = 60,
    extra_padding: float = 2,
    header_vertical: str = "center",
    header_horizontal: str = "center",
    header_wrap_text: bool = True,
) -> None:
    """应用通用布局：首行高度/冻结/筛选/数字格式/列宽自适应。"""
    worksheet.row_dimensions[header_row].height = header_height
    worksheet.freeze_panes = freeze_panes

    worksheet.auto_filter.ref = worksheet.dimensions

    # 首行表头统一对齐。
    for cell in worksheet[header_row]:
        cell.alignment = Alignment(
            vertical=header_vertical,
            horizontal=header_horizontal,
            wrap_text=header_wrap_text,
        )

    for row in worksheet.iter_rows(min_row=header_row + 1, max_row=worksheet.max_row):
        for cell in row[numeric_start_col - 1 :]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = thousand_number_format

    # 若指定了金额列表头名称，则按表头精确应用金额格式。
    if amount_headers:
        header_values = [cell.value for cell in worksheet[header_row]]
        amount_indices = [
            idx + 1 for idx, value in enumerate(header_values) if value in set(amount_headers)
        ]
        for col_idx in amount_indices:
            for row_idx in range(header_row + 1, worksheet.max_row + 1):
                worksheet.cell(row=row_idx, column=col_idx).number_format = thousand_number_format

    # 若指定了百分比列头名称，则按表头精确应用百分比格式。
    if percent_headers:
        header_values = [cell.value for cell in worksheet[header_row]]
        percent_indices = [
            idx + 1 for idx, value in enumerate(header_values) if value in set(percent_headers)
        ]
        for col_idx in percent_indices:
            for row_idx in range(header_row + 1, worksheet.max_row + 1):
                worksheet.cell(row=row_idx, column=col_idx).number_format = percent_number_format

    for column_cells in worksheet.iter_cols(min_col=1, max_col=worksheet.max_column):
        col_letter = column_cells[0].column_letter
        max_len = 0
        for cell in column_cells:
            cur_len = _display_len(_safe_text(cell.value))
            if cur_len > max_len:
                max_len = cur_len
        width = max(min_col_width, min(max_col_width, max_len + extra_padding))
        worksheet.column_dimensions[col_letter].width = width


def set_header_row_style(
    worksheet,
    *,
    header_row: int = 1,
    header_height: float = 40,
    wrap_text: bool = True,
    vertical: str = "center",
    horizontal: str = "center",
    bold: bool = False,
) -> None:
    """设置表头行高、对齐及加粗。"""
    worksheet.row_dimensions[header_row].height = header_height
    for cell in worksheet[header_row]:
        cell.alignment = Alignment(
            vertical=vertical,
            horizontal=horizontal,
            wrap_text=wrap_text,
        )
        if bold:
            cell.font = Font(bold=True)


def set_columns_number_format(
    worksheet,
    *,
    col_indices: Iterable[int],
    number_format: str,
    start_row: int = 2,
    end_row: int | None = None,
) -> None:
    """按列索引批量设置数字格式。"""
    actual_end_row = worksheet.max_row if end_row is None else end_row
    for col_idx in col_indices:
        for row_idx in range(start_row, actual_end_row + 1):
            worksheet.cell(row=row_idx, column=col_idx).number_format = number_format


def set_auto_filter_range(
    worksheet,
    *,
    start_col: int = 1,
    end_col: int | None = None,
    start_row: int = 1,
    end_row: int | None = None,
) -> None:
    """设置自动筛选范围。"""
    if end_col is None:
        end_col = worksheet.max_column
    if end_row is None:
        end_row = worksheet.max_row
    start_col_letter = worksheet.cell(row=start_row, column=start_col).column_letter
    end_col_letter = worksheet.cell(row=start_row, column=end_col).column_letter
    worksheet.auto_filter.ref = f"{start_col_letter}{start_row}:{end_col_letter}{end_row}"


def style_full_row(
    worksheet,
    *,
    row_idx: int,
    max_col: int,
    fill_color: str | None = None,
    bold: bool = False,
) -> None:
    """整行应用填充色和加粗。"""
    fill = PatternFill("solid", fgColor=fill_color) if fill_color else None
    for col_idx in range(1, max_col + 1):
        cell = worksheet.cell(row=row_idx, column=col_idx)
        if fill is not None:
            cell.fill = fill
        if bold:
            cell.font = Font(bold=True)


def auto_adjust_column_widths(
    worksheet,
    *,
    start_col: int = 1,
    end_col: int | None = None,
    start_row: int = 1,
    end_row: int | None = None,
    min_col_width: float = 10,
    max_col_width: float = 60,
    extra_padding: float = 2,
    ignore_formula: bool = True,
) -> None:
    """按显示长度自动调整列宽。"""
    actual_end_col = worksheet.max_column if end_col is None else end_col
    actual_end_row = worksheet.max_row if end_row is None else end_row

    for col_idx in range(start_col, actual_end_col + 1):
        col_letter = worksheet.cell(row=start_row, column=col_idx).column_letter
        max_len = 0
        for row_idx in range(start_row, actual_end_row + 1):
            value = worksheet.cell(row=row_idx, column=col_idx).value
            if ignore_formula and isinstance(value, str) and value.startswith("="):
                continue
            cur_len = _display_len(_safe_text(value))
            if cur_len > max_len:
                max_len = cur_len
        width = max(min_col_width, min(max_col_width, max_len + extra_padding))
        worksheet.column_dimensions[col_letter].width = width
