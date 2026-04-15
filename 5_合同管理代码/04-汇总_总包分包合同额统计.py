"""
合同管理 - 总包分包合同额汇总工具
功能:
1. 扫描所有项目文件夹
2. 提取总包合同和分包合同的金额
3. 按项目汇总,转换为万泰铢(使用JSON中的实际汇率)
4. 生成4个Sheet: 全部项目、人民币项目、泰铢项目、混合币种项目
"""

import os
import json
import re
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# =========================
# 配置区域
# =========================
CONTRACT_BASE_DIR = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/5_合同管理📑"
PROJECT_MATCH_DB_FILE = "/Users/wyattsun/Desktop/code/0_项目基本信息/project_name_match_db.json"
OUTPUT_DIR = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/5_合同管理📑"
GENERAL_CONTRACT_DIR = "总包合同"
SUBPACKAGE_CONTRACT_DIR = "分包合同"


def load_project_name_by_code(match_db_file):
    """从项目库加载 project_code -> 原始项目名称 的映射。"""
    mapping = {}

    if not os.path.exists(match_db_file):
        print(f"警告: 项目库文件不存在: {match_db_file}")
        return mapping

    try:
        with open(match_db_file, 'r', encoding='utf-8') as f:
            loaded_data = json.load(f)

        projects = loaded_data.get('projects', {}) if isinstance(loaded_data, dict) else {}

        for project_name, project_info in projects.items():
            if not isinstance(project_info, dict):
                continue
            code = str(project_info.get('project_code', '')).strip()
            name = str(project_name).strip()
            if not code or not name:
                continue

            # 若同一编码重复定义，优先保留先出现的名称
            mapping.setdefault(code, name)

        print(f"已加载项目编码映射: {len(mapping)} 个 project_code")
    except Exception as e:
        print(f"警告: 读取项目库失败: {e}")

    return mapping


def extract_contract_amount_from_json(json_file):
    """从JSON文件提取合同金额、币种和汇率"""
    try:
        with open(json_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        amount = float(data.get('合同金额', 0))
        currency = str(data.get('币种', '')).strip()
        exchange_rate = float(data.get('汇率', 1.0))
        
        if currency == '':
            currency = 'THB'
        
        return amount, currency, exchange_rate
    except Exception as e:
        return 0, 'THB', 1.0


def calculate_contract_total(contract_folder):
    """计算合同文件夹下所有合同的总金额(万泰铢),并统计币种"""
    total_thb = 0.0
    contract_count = 0
    currencies_used = set()
    currency_details = {'CNY': 0.0, 'THB': 0.0}
    
    if not os.path.exists(contract_folder):
        return total_thb, contract_count, currencies_used, currency_details
    
    for item in os.listdir(contract_folder):
        item_path = os.path.join(contract_folder, item)
        
        if not os.path.isdir(item_path) or item.startswith('.'):
            continue
        
        json_file = os.path.join(item_path, '合同信息.json')
        
        if os.path.exists(json_file):
            amount, currency, exchange_rate = extract_contract_amount_from_json(json_file)
            
            amount_in_wan_thb = (amount * exchange_rate) / 10000
            
            total_thb += amount_in_wan_thb
            contract_count += 1
            currencies_used.add(currency)
            
            if currency == 'CNY':
                currency_details['CNY'] += amount_in_wan_thb
            else:
                currency_details['THB'] += amount_in_wan_thb
    
    return total_thb, contract_count, currencies_used, currency_details


def extract_project_code(project_dir_name):
    """从项目目录名中提取项目编码"""
    match = re.search(r'(THCC-\d{4}-\d{3}|LSBM-\d+)', project_dir_name)
    return match.group(1) if match else None


def resolve_contract_dir_name(base_dir, logical_name):
    """
    自动解析合同类型目录名。
    优先级: 完全匹配 > 后缀匹配(支持 A_总包合同) > 包含匹配
    """
    exact_path = os.path.join(base_dir, logical_name)
    if os.path.isdir(exact_path):
        return logical_name

    candidates = []
    try:
        for name in os.listdir(base_dir):
            path = os.path.join(base_dir, name)
            if os.path.isdir(path) and not name.startswith('.'):
                candidates.append(name)
    except Exception:
        return logical_name

    suffix_matches = sorted([name for name in candidates if name.endswith(logical_name)])
    if suffix_matches:
        return suffix_matches[0]

    contains_matches = sorted([name for name in candidates if logical_name in name])
    if contains_matches:
        return contains_matches[0]

    return logical_name


def collect_contract_project_paths(base_dir, contract_dir_name):
    """收集合同类型目录下 项目状态/项目目录 的路径映射"""
    result = {}
    contract_root = os.path.join(base_dir, contract_dir_name)

    if not os.path.exists(contract_root):
        return result

    for status_dir in os.listdir(contract_root):
        status_path = os.path.join(contract_root, status_dir)
        if not os.path.isdir(status_path) or status_dir.startswith('.'):
            continue

        for project_dir in os.listdir(status_path):
            project_path = os.path.join(status_path, project_dir)
            if not os.path.isdir(project_path) or project_dir.startswith('.'):
                continue

            project_code = extract_project_code(project_dir)
            if not project_code:
                continue

            result[(status_dir, project_dir, project_code)] = project_path

    return result


def scan_all_projects(base_dir, project_name_by_code, general_contract_dir, subpackage_contract_dir):
    """扫描所有项目并汇总总包分包合同额"""
    
    print("\n开始扫描项目文件夹...")
    
    projects_data = []

    general_paths = collect_contract_project_paths(base_dir, general_contract_dir)
    subpackage_paths = collect_contract_project_paths(base_dir, subpackage_contract_dir)
    all_project_keys = sorted(set(general_paths.keys()) | set(subpackage_paths.keys()))

    for project_status, project_dir, project_code in all_project_keys:
        standard_name = project_name_by_code.get(project_code, project_dir)
        name_source = 'project_code_db' if project_code in project_name_by_code else 'folder_name_fallback'

        general_total = 0.0
        general_count = 0
        general_currencies = set()
        general_currency_details = {'CNY': 0.0, 'THB': 0.0}
        general_path = general_paths.get((project_status, project_dir, project_code))
        if general_path:
            amount, count, currencies, details = calculate_contract_total(general_path)
            general_total += amount
            general_count += count
            general_currencies.update(currencies)
            general_currency_details['CNY'] += details['CNY']
            general_currency_details['THB'] += details['THB']

        subpackage_total = 0.0
        subpackage_count = 0
        subpackage_currencies = set()
        subpackage_currency_details = {'CNY': 0.0, 'THB': 0.0}
        subpackage_path = subpackage_paths.get((project_status, project_dir, project_code))
        if subpackage_path:
            amount, count, currencies, details = calculate_contract_total(subpackage_path)
            subpackage_total += amount
            subpackage_count += count
            subpackage_currencies.update(currencies)
            subpackage_currency_details['CNY'] += details['CNY']
            subpackage_currency_details['THB'] += details['THB']

        if general_count > 0 or subpackage_count > 0:
            all_currencies = general_currencies | subpackage_currencies

            if 'CNY' in all_currencies and 'THB' in all_currencies:
                currency_type = 'MIXED'
            elif 'CNY' in all_currencies:
                currency_type = 'CNY'
            else:
                currency_type = 'THB'

            projects_data.append({
                '项目状态': project_status,
                '项目编码': project_code,
                '项目名称': standard_name,
                '总包合同额(万泰铢)': general_total,
                '分包合同额(万泰铢)': subpackage_total,
                '_currency_type': currency_type
            })

            currency_label = {'CNY': '[人民币]', 'THB': '[泰铢]', 'MIXED': '[混合币种]'}.get(currency_type, '')
            print(f"  {project_code} {currency_label}: 总包 {general_count}个({general_total:,.2f}万) + 分包 {subpackage_count}个({subpackage_total:,.2f}万) [名称匹配:{name_source}]")
    
    print(f"\n共找到 {len(projects_data)} 个项目")
    
    return projects_data


def generate_excel_report(projects_data, output_dir):
    """生成Excel报告,含4个Sheet"""
    
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_file = os.path.join(output_dir, f'合同管理_总包分包汇总_{timestamp}.xlsx')
    
    print(f"\n生成Excel报告: {output_file}")
    
    df_all = pd.DataFrame(projects_data)
    display_columns = ['项目状态', '项目编码', '项目名称', '总包合同额(万泰铢)', '分包合同额(万泰铢)']
    
    # Sheet 1: 全部项目
    df_sheet1 = df_all[display_columns].copy().sort_values(by='项目编码', ascending=False)
    total_row_1 = {'项目状态': '', '项目编码': '', '项目名称': '总计',
                   '总包合同额(万泰铢)': df_sheet1['总包合同额(万泰铢)'].sum(),
                   '分包合同额(万泰铢)': df_sheet1['分包合同额(万泰铢)'].sum()}
    df_sheet1 = pd.concat([df_sheet1, pd.DataFrame([total_row_1])], ignore_index=True)
    
    # Sheet 2: 人民币项目
    df_cny = df_all[df_all['_currency_type'] == 'CNY'].copy()
    df_sheet2 = df_cny[display_columns].copy().sort_values(by='项目编码', ascending=False)
    if not df_sheet2.empty:
        total_row_2 = {'项目状态': '', '项目编码': '', '项目名称': '总计',
                       '总包合同额(万泰铢)': df_sheet2['总包合同额(万泰铢)'].sum(),
                       '分包合同额(万泰铢)': df_sheet2['分包合同额(万泰铢)'].sum()}
        df_sheet2 = pd.concat([df_sheet2, pd.DataFrame([total_row_2])], ignore_index=True)
    
    # Sheet 3: 泰铢项目
    df_thb = df_all[df_all['_currency_type'] == 'THB'].copy()
    df_sheet3 = df_thb[display_columns].copy().sort_values(by='项目编码', ascending=False)
    if not df_sheet3.empty:
        total_row_3 = {'项目状态': '', '项目编码': '', '项目名称': '总计',
                       '总包合同额(万泰铢)': df_sheet3['总包合同额(万泰铢)'].sum(),
                       '分包合同额(万泰铢)': df_sheet3['分包合同额(万泰铢)'].sum()}
        df_sheet3 = pd.concat([df_sheet3, pd.DataFrame([total_row_3])], ignore_index=True)
    
    # Sheet 4: 混合币种项目
    df_mixed = df_all[df_all['_currency_type'] == 'MIXED'].copy()
    df_sheet4 = df_mixed[display_columns].copy().sort_values(by='项目编码', ascending=False)
    if not df_sheet4.empty:
        total_row_4 = {'项目状态': '', '项目编码': '', '项目名称': '总计',
                       '总包合同额(万泰铢)': df_sheet4['总包合同额(万泰铢)'].sum(),
                       '分包合同额(万泰铢)': df_sheet4['分包合同额(万泰铢)'].sum()}
        df_sheet4 = pd.concat([df_sheet4, pd.DataFrame([total_row_4])], ignore_index=True)
    
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df_sheet1.to_excel(writer, index=False, sheet_name='1_总包分包汇总(全部)')
        if not df_sheet2.empty:
            df_sheet2.to_excel(writer, index=False, sheet_name='2_人民币项目(CNY)')
        if not df_sheet3.empty:
            df_sheet3.to_excel(writer, index=False, sheet_name='3_泰铢项目(THB)')
        if not df_sheet4.empty:
            df_sheet4.to_excel(writer, index=False, sheet_name='4_混合币种项目')
    
    format_excel(output_file)
    
    stats = {
        'all': len(df_all), 'cny': len(df_cny), 'thb': len(df_thb), 'mixed': len(df_mixed),
        'total_general': df_sheet1['总包合同额(万泰铢)'].iloc[-1] if not df_sheet1.empty else 0,
        'total_subpackage': df_sheet1['分包合同额(万泰铢)'].iloc[-1] if not df_sheet1.empty else 0
    }
    
    return output_file, stats


def format_excel(file_path):
    """格式化Excel文件,所有Sheet"""
    
    wb = load_workbook(file_path)
    
    for ws in wb.worksheets:
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=12)
        total_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
        total_font = Font(bold=True, size=11)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                             top=Side(style='thin'), bottom=Side(style='thin'))
        
        ws.row_dimensions[1].height = 40
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = thin_border
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(vertical='center')
        
        last_row = ws.max_row
        for col_idx, cell in enumerate(ws[last_row], start=1):
            cell.fill = total_fill
            cell.font = total_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center' if col_idx <= 3 else 'right', vertical='center')
        
        for col_idx in [4, 5]:
            col_letter = get_column_letter(col_idx)
            for cell in ws[col_letter][1:]:
                if cell.value and isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'
        
        column_widths = {'A': 20, 'B': 18, 'C': 50, 'D': 18, 'E': 18}
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    
    wb.save(file_path)


def main():
    """主函数"""
    
    print("=" * 80)
    print("合同管理 - 总包分包合同额汇总工具")
    print("=" * 80)
    print(f"\n合同目录: {CONTRACT_BASE_DIR}")
    print(f"总包目录: {GENERAL_CONTRACT_DIR}")
    print(f"分包目录: {SUBPACKAGE_CONTRACT_DIR}")
    print(f"映射文件: {PROJECT_MATCH_DB_FILE}")
    print(f"输出目录: {OUTPUT_DIR}")
    print(f"\n汇率设置: 使用JSON中的实际汇率值")
    
    if not os.path.exists(CONTRACT_BASE_DIR):
        print(f"\n错误: 合同目录不存在: {CONTRACT_BASE_DIR}")
        return

    resolved_general_dir = resolve_contract_dir_name(CONTRACT_BASE_DIR, GENERAL_CONTRACT_DIR)
    resolved_subpackage_dir = resolve_contract_dir_name(CONTRACT_BASE_DIR, SUBPACKAGE_CONTRACT_DIR)
    print(f"自动识别后-总包目录: {resolved_general_dir}")
    print(f"自动识别后-分包目录: {resolved_subpackage_dir}")
    
    project_name_by_code = load_project_name_by_code(PROJECT_MATCH_DB_FILE)
    projects_data = scan_all_projects(
        CONTRACT_BASE_DIR,
        project_name_by_code,
        resolved_general_dir,
        resolved_subpackage_dir,
    )
    
    if not projects_data:
        print("\n未找到任何合同数据!")
        return
    
    report_file, stats = generate_excel_report(projects_data, OUTPUT_DIR)
    
    print("\n" + "=" * 80)
    print("✅ 汇总完成!")
    print("=" * 80)
    print(f"\n报告文件: {report_file}")
    print(f"\n统计信息:")
    print(f"  项目总数: {stats['all']}")
    print(f"    - 人民币项目: {stats['cny']}")
    print(f"    - 泰铢项目: {stats['thb']}")
    print(f"    - 混合币种项目: {stats['mixed']}")
    print(f"\n  总包合同额: {stats['total_general']:,.2f} 万泰铢")
    print(f"  分包合同额: {stats['total_subpackage']:,.2f} 万泰铢")
    print(f"\nExcel包含Sheet:")
    print(f"  1. 总包分包汇总(全部) - {stats['all']} 个项目")
    print(f"  2. 人民币项目(CNY) - {stats['cny']} 个项目")
    print(f"  3. 泰铢项目(THB) - {stats['thb']} 个项目")
    print(f"  4. 混合币种项目 - {stats['mixed']} 个项目")
    print("\nExcel格式:")
    print("  ✓ 首行高度: 40")
    print("  ✓ 冻结首行: A2")
    print("  ✓ 自动筛选: 已开启")
    print("  ✓ 排序方式: 项目编码倒序")
    print("  ✓ 总计行: 每个Sheet都已添加")
    print("  ✓ 汇率计算: 使用JSON中的实际汇率")
    print("\n" + "=" * 80)


if __name__ == '__main__':
    main()
