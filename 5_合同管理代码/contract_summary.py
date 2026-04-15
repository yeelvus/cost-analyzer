import os
import json
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 合同数据自动汇总统计”脚本，主要用来把分散在各项目目录里的合同信息做成一份年度统计表。


DEFAULT_OUTPUT_DIR = r'/Users/wyattsun/Desktop/邯郸中材/01成本统计/Z_report'
DEFAULT_OUTPUT_FILE = '合同汇总统计.xlsx'


def _format_date_get_year(v):
    if v is None:
        return None
    s = str(v).strip()
    if s == "":
        return None
    # 先处理类似 20241008
    digits = ''.join(ch for ch in s if ch.isdigit())
    if len(digits) >= 4:
        # 年一般为前4位
        y = digits[0:4]
        if y.isdigit():
            return y
    # 处理像 2024-10-08 或 2024/10/08
    for sep in ['-', '/']:
        if sep in s:
            parts = [p for p in s.replace('/', '-').split('-') if p]
            if len(parts) >= 1 and parts[0].isdigit() and len(parts[0]) == 4:
                return parts[0]
    return None


def _determine_category_from_path(path):
    p = path.replace('\\', '/').lower()
    if '总包' in p or 'zongbao' in p:
        return '总包'
    if '分包' in p or 'fenbao' in p:
        return '分包'
    # 如果上级文件夹含有关键词
    parts = p.split('/')
    for part in parts[::-1]:
        if '总包' in part:
            return '总包'
        if '分包' in part:
            return '分包'
    return '未知'


def safe_float(x, default=0.0):
    try:
        return float(x)
    except Exception:
        return default


def collect_contracts(root_dir):
    records = []
    for root, dirs, files in os.walk(root_dir):
        for fn in files:
            if fn != '合同信息.json':
                continue
            fp = os.path.join(root, fn)
            try:
                with open(fp, 'r', encoding='utf-8') as f:
                    data = json.load(f)
            except Exception:
                continue
            contract_no = data.get('合同编号', '')
            date_raw = data.get('合同签订日期', '')
            year = _format_date_get_year(date_raw) or ''
            amount = safe_float(data.get('合同金额', 0))
            currency = data.get('币种', '')
            rate = safe_float(data.get('汇率', 1.0), 1.0)
            category = _determine_category_from_path(root)
            converted = amount * rate
            records.append({
                '年份': year,
                '合同编号': contract_no,
                '合同签订日期': date_raw,
                '合同金额': amount,
                '币种': currency,
                '汇率': rate,
                '折算金额': converted,
                '合同类别': category
            })
    return records


def make_summary(df):
    # 汇总：按 年份 和 合同类别 统计合同数量与金额总和
    grp = df.groupby(['年份', '合同类别']).agg(
        合同数量=('合同编号', 'count'),
        折算金额合计=('折算金额', 'sum'),
        合同金额合计=('合同金额', 'sum')
    ).reset_index()

    # 透视表：合同数量
    pivot_count = grp.pivot(index='年份', columns='合同类别', values='合同数量').fillna(0).astype(int).reset_index()

    # 透视表：折算金额
    pivot_amount = grp.pivot(index='年份', columns='合同类别', values='折算金额合计').fillna(0.0).reset_index()

    return grp, pivot_count, pivot_amount


def _style_worksheet(ws):
    header_fill = PatternFill(fill_type='solid', fgColor='1F4E78')
    header_font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
    body_font = Font(name='Microsoft YaHei', bold=False, color='000000')
    center_alignment = Alignment(horizontal='center', vertical='center')
    border_side = Side(style='thin', color='D9D9D9')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    ws.freeze_panes = 'A2'

    max_col = ws.max_column
    max_row = ws.max_row
    if max_col == 0 or max_row == 0:
        return

    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    for col_idx in range(1, max_col + 1):
        header_cell = ws.cell(row=1, column=col_idx)
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.alignment = center_alignment
        header_cell.border = cell_border

        max_len = len(str(header_cell.value or ''))
        for row_idx in range(2, max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = body_font
            cell.alignment = center_alignment
            cell.border = cell_border

            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))

        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 36)

    ws.row_dimensions[1].height = 24


def _apply_number_formats(ws):
    money_headers = {'合同金额', '折算金额', '折算金额合计', '合同金额合计', '总包', '分包', '未知'}
    rate_headers = {'汇率'}

    max_col = ws.max_column
    max_row = ws.max_row

    for col_idx in range(1, max_col + 1):
        header = str(ws.cell(row=1, column=col_idx).value or '').strip()
        if header in money_headers:
            for row_idx in range(2, max_row + 1):
                ws.cell(row=row_idx, column=col_idx).number_format = '#,##0.00'
        elif header in rate_headers:
            for row_idx in range(2, max_row + 1):
                ws.cell(row=row_idx, column=col_idx).number_format = '0.0000'


def export_to_excel(records, out_path):
    df = pd.DataFrame(records)
    if df.empty:
        print('未找到合同数据。')
        return

    out_dir = os.path.dirname(out_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    grp, pivot_count, pivot_amount = make_summary(df)

    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='原始明细', index=False)
        grp.to_excel(writer, sheet_name='年度类别汇总', index=False)
        pivot_count.to_excel(writer, sheet_name='合同数量透视', index=False)
        pivot_amount.to_excel(writer, sheet_name='折算金额透视', index=False)

        for ws in writer.book.worksheets:
            _style_worksheet(ws)
            _apply_number_formats(ws)

    print('已导出：', out_path)


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='合同汇总：按年份与总包/分包统计（金额*汇率）')
    parser.add_argument('--root', '-r', default='.', help='合同根目录（递归查找 合同信息.json）')
    parser.add_argument(
        '--out',
        '-o',
        default=os.path.join(DEFAULT_OUTPUT_DIR, DEFAULT_OUTPUT_FILE),
        help='输出 Excel 路径'
    )
    args = parser.parse_args()

    records = collect_contracts(args.root)
    export_to_excel(records, args.out)
