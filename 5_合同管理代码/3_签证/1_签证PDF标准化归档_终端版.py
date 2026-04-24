import json
import re
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook


TXT_ENCODING = "utf-8-sig"
PDF_PATTERN = re.compile(
    r"^\s*(?P<index>\d+)\s*[_ ]?\s*(?P<name>.+?)\s*_\s*(?P<amount>-?\d+(?:\.\d+)?)\s*$"
)
SUPPORTED_EXTENSIONS = {".pdf", ".jpg", ".jpeg"}
INVALID_PATH_CHARS = str.maketrans({
    "/": "／",
    "\\": "＼",
    ":": "：",
    "*": "＊",
    "?": "？",
    '"': "＂",
    "<": "＜",
    ">": "＞",
    "|": "｜",
})
SUMMARY_HEADERS = [
    "标准序号",
    "签证名称",
    "签证金额",
    "签证日期",
    "签证完成日期",
    "签证是否确认",
]
GENERATED_TXT_FILES = [
    "签证名称.txt",
    "签证金额.txt",
    "签证日期.txt",
    "签证完成日期.txt",
    "是否签证确认.txt",
]


@dataclass
class ParsedPdf:
    source_path: Path
    index: int
    visa_name: str
    amount: str
    extension: str


@dataclass
class OrganizeTask:
    parsed: ParsedPdf
    target_stem: str
    target_dir: Path
    target_file: Path


def clean_text_segment(text: str) -> str:
    cleaned = re.sub(r"\s+", " ", text.strip(" _"))
    return cleaned.translate(INVALID_PATH_CHARS)


def normalize_input_path(raw_path: str) -> Path:
    normalized = raw_path.strip().strip('"').strip("'")
    return Path(normalized).expanduser().resolve()


def parse_pdf_file(pdf_path: Path) -> ParsedPdf | None:
    extension = pdf_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        return None

    match = PDF_PATTERN.match(pdf_path.stem)
    if not match:
        return None

    visa_name = clean_text_segment(match.group("name"))
    amount = match.group("amount").strip()
    if not visa_name:
        return None

    return ParsedPdf(
        source_path=pdf_path,
        index=int(match.group("index")),
        visa_name=visa_name,
        amount=amount,
        extension=extension,
    )


def build_tasks(folder_path: Path) -> tuple[list[OrganizeTask], list[Path]]:
    pdf_files = sorted(
        [
            file
            for file in folder_path.iterdir()
            if file.is_file() and file.suffix.lower() in SUPPORTED_EXTENSIONS
        ],
        key=lambda file: file.name,
    )

    parsed_items: list[ParsedPdf] = []
    unmatched_files: list[Path] = []
    for pdf_file in pdf_files:
        parsed = parse_pdf_file(pdf_file)
        if parsed is None:
            unmatched_files.append(pdf_file)
            continue
        parsed_items.append(parsed)

    if not parsed_items:
        return [], unmatched_files

    index_width = max(2, len(str(max(item.index for item in parsed_items))))
    tasks: list[OrganizeTask] = []
    for parsed in sorted(parsed_items, key=lambda item: (item.index, item.visa_name, item.amount)):
        target_stem = f"{str(parsed.index).zfill(index_width)}_{parsed.visa_name}_{parsed.amount}"
        target_dir = folder_path / target_stem
        target_file = target_dir / f"{target_stem}{parsed.extension}"
        tasks.append(
            OrganizeTask(
                parsed=parsed,
                target_stem=target_stem,
                target_dir=target_dir,
                target_file=target_file,
            )
        )

    return tasks, unmatched_files


def load_json_if_exists(file_path: Path) -> dict:
    if not file_path.exists():
        return {}
    try:
        content = json.loads(file_path.read_text(encoding="utf-8"))
        return content if isinstance(content, dict) else {}
    except Exception:
        return {}


def build_core_metadata(
    serial_no: str,
    visa_name: str,
    amount: str,
    existing_data: dict,
) -> dict[str, str]:
    serial_text = str(serial_no).strip()
    if serial_text.isdigit():
        serial_text = serial_text.zfill(max(2, len(serial_text)))

    return {
        "标准序号": serial_text,
        "签证名称": str(existing_data.get("签证名称") or visa_name or "").strip(),
        "签证金额": str(existing_data.get("签证金额") or amount or "").strip(),
        "签证日期": str(existing_data.get("签证日期") or "").strip(),
        "签证完成日期": str(existing_data.get("签证完成日期") or "").strip(),
        "签证是否确认": str(existing_data.get("签证是否确认") or "").strip(),
    }


def write_metadata_files(task: OrganizeTask) -> dict[str, str]:
    metadata_path = task.target_dir / "签证信息.json"
    existing_data = load_json_if_exists(metadata_path)
    metadata = build_core_metadata(
        serial_no=task.target_stem.split("_", 1)[0],
        visa_name=task.parsed.visa_name,
        amount=task.parsed.amount,
        existing_data=existing_data,
    )

    metadata_path.write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return metadata


def move_pdf_file(task: OrganizeTask) -> None:
    task.target_dir.mkdir(parents=True, exist_ok=True)

    if task.target_file.exists():
        raise FileExistsError(f"目标文件已存在: {task.target_file}")

    shutil.move(str(task.parsed.source_path), str(task.target_file))


def delete_generated_txt_files(folder_path: Path) -> int:
    deleted_count = 0
    for subfolder in folder_path.iterdir():
        if not subfolder.is_dir():
            continue
        for txt_name in GENERATED_TXT_FILES:
            txt_path = subfolder / txt_name
            if txt_path.exists():
                txt_path.unlink()
                deleted_count += 1
    return deleted_count


def collect_and_normalize_metadata(folder_path: Path) -> list[dict[str, str]]:
    summary_rows: list[dict[str, str]] = []

    for subfolder in sorted([item for item in folder_path.iterdir() if item.is_dir()], key=lambda p: p.name):
        metadata_path = subfolder / "签证信息.json"
        existing_data = load_json_if_exists(metadata_path)

        match = PDF_PATTERN.match(subfolder.name)
        serial_no = ""
        visa_name = ""
        amount = ""
        if match:
            serial_no = match.group("index")
            visa_name = clean_text_segment(match.group("name"))
            amount = match.group("amount").strip()

        metadata = build_core_metadata(
            serial_no=existing_data.get("标准序号") or serial_no,
            visa_name=existing_data.get("签证名称") or visa_name,
            amount=existing_data.get("签证金额") or amount,
            existing_data=existing_data,
        )

        # 仅处理看起来像签证目录的数据
        if not (metadata["标准序号"] or metadata["签证名称"] or metadata["签证金额"]):
            continue

        metadata_path.write_text(
            json.dumps(metadata, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        summary_rows.append(metadata)

    summary_rows.sort(
        key=lambda row: (0, int(row["标准序号"])) if str(row["标准序号"]).isdigit() else (1, str(row["标准序号"]))
    )
    return summary_rows


def write_summary_excel(folder_path: Path, summary_rows: list[dict[str, str]]) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = folder_path / f"签证整理汇总_{timestamp}.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        raise RuntimeError("Excel工作表初始化失败")
    worksheet.title = "签证汇总"

    for col_index, header in enumerate(SUMMARY_HEADERS, start=1):
        worksheet.cell(row=1, column=col_index, value=header)

    for row_index, row_data in enumerate(summary_rows, start=2):
        for col_index, header in enumerate(SUMMARY_HEADERS, start=1):
            worksheet.cell(row=row_index, column=col_index, value=row_data.get(header, ""))

    workbook.save(output_path)
    return output_path


def print_preview(tasks: list[OrganizeTask], unmatched_files: list[Path]) -> None:
    print("\n扫描完成。")
    print(f"可处理文件数量: {len(tasks)}")
    print(f"跳过的文件数量: {len(unmatched_files)}")

    if tasks:
        print("\n将执行以下标准化操作:")
        for task in tasks:
            print(f"- {task.parsed.source_path.name}")
            print(f"  -> {task.target_dir.name}/{task.target_file.name}")

    if unmatched_files:
        print("\n以下文件未匹配到“序号_名称_金额.扩展名”规则，已跳过:")
        for file_path in unmatched_files:
            print(f"- {file_path.name}")


def process_tasks(tasks: list[OrganizeTask]) -> tuple[int, list[str]]:
    success_count = 0
    errors: list[str] = []

    for task in tasks:
        try:
            move_pdf_file(task)
            write_metadata_files(task)
            success_count += 1
            print(f"✓ 已处理: {task.parsed.source_path.name} -> {task.target_file}")
        except Exception as exc:
            errors.append(f"{task.parsed.source_path.name}: {exc}")
            print(f"✗ 处理失败: {task.parsed.source_path.name}: {exc}")

    return success_count, errors


def main() -> None:
    print("=" * 70)
    print("签证PDF标准化归档工具")
    print("功能: 序号补零、创建文件夹、移动PDF/JPG、仅输出JSON、删除历史txt、输出Excel汇总")
    print("=" * 70)
    print("说明: 只处理当前目录下一级的PDF/JPG/JPEG文件，不递归处理子文件夹。")

    folder_input = input("\n请输入需要处理的文件夹路径: ").strip()
    if not folder_input:
        print("未输入路径，程序结束。")
        return

    folder_path = normalize_input_path(folder_input)
    if not folder_path.exists() or not folder_path.is_dir():
        print(f"路径无效: {folder_path}")
        return

    tasks, unmatched_files = build_tasks(folder_path)
    print_preview(tasks, unmatched_files)

    if not tasks:
        print("\n未找到可处理的文件，将仅执行：删除旧txt、规范json、生成Excel汇总。")

    dry_run_choice = input("\n是否先仅预演、不实际移动文件？(y/N): ").strip().lower()
    if dry_run_choice in {"y", "yes", "是"}:
        print("\n已完成预演，未修改任何文件。")
        return

    confirm_choice = input("\n确认开始执行以上操作？(y/N): ").strip().lower()
    if confirm_choice not in {"y", "yes", "是"}:
        print("已取消执行。")
        return

    deleted_txt_count = delete_generated_txt_files(folder_path)
    if deleted_txt_count:
        print(f"\n已删除历史txt文件: {deleted_txt_count} 个")

    success_count, errors = process_tasks(tasks)
    summary_rows = collect_and_normalize_metadata(folder_path)

    if summary_rows:
        summary_path = write_summary_excel(folder_path, summary_rows)
        print(f"\n汇总文件已生成: {summary_path}")
    else:
        print("\n未收集到可汇总的签证数据，未生成Excel汇总。")

    print("\n处理完成。")
    print(f"成功处理文件: {success_count} 个")
    print(f"汇总签证目录: {len(summary_rows)} 个")
    print(f"失败: {len(errors)} 个")

    if errors:
        print("\n失败明细:")
        for error in errors:
            print(f"- {error}")


if __name__ == "__main__":
    main()