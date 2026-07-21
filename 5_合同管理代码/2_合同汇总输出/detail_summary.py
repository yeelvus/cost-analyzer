# -*- coding: utf-8 -*-
"""
合同明细汇总（原 03_合同管理-汇总）
扫描 总包合同/分包合同 下的 合同信息.json，生成明细与分类汇总 Sheet。
"""

from __future__ import annotations

import json
import os
import re
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook

from summary_config import (
    COLUMN_ORDER_FILE,
    OUTPUT_EXCLUDED_COLUMNS,
    SETTLEMENT_CHECKS,
    TARGET_KEYWORDS,
    get_contract_root,
    get_project_db,
)


def read_contract_data(folder_path: str) -> dict[str, Any]:
    default_data = {
        "对方单位公司名称": "",
        "合同编号": "",
        "合同签订日期": "",
        "币种": "",
        "合同金额": 0.0,
    }
    json_path = os.path.join(folder_path, "合同信息.json")
    if not os.path.exists(json_path):
        return default_data
    try:
        with open(json_path, "r", encoding="utf-8") as f:
            json_data = json.load(f)
        if "合同金额" in json_data:
            try:
                json_data["合同金额"] = float(json_data["合同金额"])
            except (TypeError, ValueError):
                json_data["合同金额"] = 0.0
        # 字段别名归一
        if not json_data.get("对方单位公司名称") and json_data.get("对方公司名称"):
            json_data["对方单位公司名称"] = json_data.get("对方公司名称", "")
        return {**default_data, **json_data}
    except Exception as e:
        print(f" 警告：JSON 读取失败 {json_path}: {e}")
        return default_data


def is_folder_non_empty(folder_path: str) -> bool:
    if not os.path.isdir(folder_path):
        return False
    return len([f for f in os.listdir(folder_path) if not f.startswith(".")]) > 0


def natural_sort_key(text: str):
    def convert(s):
        return int(s) if s.isdigit() else s.lower()

    return [convert(c) for c in re.split(r"([0-9]+)", text)]


def detect_attachment_files(folder_path: str) -> dict[str, str]:
    if not os.path.isdir(folder_path):
        return {"有PDF": "", "有Word": "", "有Excel": ""}
    has_pdf = has_word = has_excel = False
    try:
        for item in os.listdir(folder_path):
            if item.startswith("."):
                continue
            item_path = os.path.join(folder_path, item)
            if not os.path.isfile(item_path):
                continue
            lower_name = item.lower()
            if lower_name.endswith(".pdf"):
                has_pdf = True
            elif lower_name.endswith((".doc", ".docx")):
                has_word = True
            elif lower_name.endswith((".xls", ".xlsx")):
                has_excel = True
    except OSError:
        pass
    return {
        "有PDF": "✓" if has_pdf else "",
        "有Word": "✓" if has_word else "",
        "有Excel": "✓" if has_excel else "",
    }


def load_project_code_mapping(db_path: str) -> dict[str, tuple[str, str]]:
    code_to_meta: dict[str, tuple[str, str]] = {}
    if not os.path.exists(db_path):
        print(f"警告：项目编码库不存在：{db_path}")
        return code_to_meta
    try:
        with open(db_path, "r", encoding="utf-8") as f:
            db = json.load(f)
    except Exception as e:
        print(f"警告：读取项目编码库失败：{e}")
        return code_to_meta

    projects = db.get("projects", {}) if isinstance(db, dict) else {}
    for std_name, meta in projects.items():
        if not isinstance(meta, dict):
            continue
        project_code = str(meta.get("project_code", "")).strip()
        if project_code:
            code_to_meta[project_code] = (std_name, project_code)
    print(f"已加载项目编码库：{db_path}，可匹配项目编码 {len(code_to_meta)} 个")
    return code_to_meta


def extract_project_code_from_folder_name(folder_name: str) -> str:
    text = str(folder_name).strip()
    if not text:
        return ""
    return text.split("_", 1)[0].strip()


def resolve_project_by_code(raw_project_name: str, code_to_meta: dict) -> tuple[str, str]:
    raw = str(raw_project_name).strip()
    if not raw:
        return "", ""
    project_code = extract_project_code_from_folder_name(raw)
    if project_code and project_code in code_to_meta:
        return code_to_meta[project_code]
    return raw, project_code


def infer_project_context(
    contract_folder_path: str, sub_abs_path: str, code_to_meta: dict
) -> tuple[str, str, str, str]:
    rel_path = os.path.relpath(sub_abs_path, contract_folder_path)
    parts = [] if rel_path == "." else rel_path.split(os.sep)

    project_status = parts[0] if parts else os.path.basename(contract_folder_path)
    raw_project_name = ""
    project_dir = ""

    for idx, part in enumerate(parts):
        code = extract_project_code_from_folder_name(part)
        if code and (code in code_to_meta or re.match(r"^[A-Za-z]{2,}-\d{4}-\d+$", code)):
            raw_project_name = part
            project_dir = os.path.join(contract_folder_path, *parts[: idx + 1])
            break

    if not raw_project_name:
        if len(parts) >= 2:
            raw_project_name = parts[1]
            project_dir = os.path.join(contract_folder_path, parts[0], parts[1])
        elif len(parts) == 1:
            raw_project_name = parts[0]
            project_dir = os.path.join(contract_folder_path, parts[0])
        else:
            raw_project_name = os.path.basename(os.path.dirname(contract_folder_path))
            project_dir = os.path.dirname(contract_folder_path)

    project_name, project_code = resolve_project_by_code(raw_project_name, code_to_meta)
    return project_status, project_name, project_code, project_dir


def load_column_order(columns: list[str], default_order: list[str]) -> list[str]:
    if os.path.exists(COLUMN_ORDER_FILE):
        try:
            with open(COLUMN_ORDER_FILE, "r", encoding="utf-8") as f:
                saved = [line.strip() for line in f.readlines() if line.strip() in columns]
            return saved + [c for c in columns if c not in saved]
        except Exception:
            return default_order
    return default_order


def scan_contract_details(
    root_dir: str | None = None,
    project_db_file: str | None = None,
    target_keywords: list[str] | None = None,
) -> tuple[dict[str, list[dict]], dict[str, dict], dict[str, int]]:
    """扫描合同明细。返回 (all_data, settlement_results, project_counts)。"""
    root_dir = root_dir or get_contract_root()
    project_db_file = project_db_file or get_project_db()
    target_keywords = target_keywords or TARGET_KEYWORDS

    all_data: dict[str, list[dict]] = {}
    project_counts: dict[str, int] = {}
    settlement_results: dict[str, dict] = {}

    print(f"开始扫描根目录：{root_dir}")
    print(f"正在查找文件夹名包含 {target_keywords} 的目录...\n")

    project_code_map = load_project_code_mapping(project_db_file)

    for current_root, dirs, _files in os.walk(root_dir):
        folder_name = os.path.basename(current_root)
        if not any(kw in folder_name for kw in target_keywords):
            continue
        parent_name = os.path.basename(os.path.dirname(current_root))
        if any(kw in parent_name for kw in target_keywords):
            continue

        dirs[:] = []
        hit_kw = next(kw for kw in target_keywords if kw in folder_name)

        if hit_kw not in all_data:
            all_data[hit_kw] = []
            project_counts[hit_kw] = 0
        project_counts[hit_kw] += 1

        contract_folder_path = current_root
        print(f"[{project_counts[hit_kw]}] 发现 {hit_kw} 文件夹：{contract_folder_path}")

        subfolder_data: dict[str, dict] = {}
        for sub_root, _, _ in os.walk(contract_folder_path):
            rel_path = os.path.relpath(sub_root, contract_folder_path)
            sub_key = rel_path if rel_path != "." else folder_name
            subfolder_data[sub_key] = read_contract_data(sub_root)

        for sub, fields in subfolder_data.items():
            if all(str(v).strip() == "" or v == 0.0 for k, v in fields.items() if k != "币种"):
                continue

            sub_abs_path = (
                os.path.join(contract_folder_path, sub)
                if sub != folder_name
                else contract_folder_path
            )
            sub_abs_path = os.path.normpath(sub_abs_path)

            project_status, project_name, project_code, project_dir = infer_project_context(
                contract_folder_path, sub_abs_path, project_code_map
            )
            settlement_key = project_code if project_code else project_name

            if settlement_key and settlement_key not in settlement_results:
                settlement_results[settlement_key] = {check["name"]: "" for check in SETTLEMENT_CHECKS}

            if settlement_key:
                for check in SETTLEMENT_CHECKS:
                    if settlement_results[settlement_key].get(check["name"]) == "有资料":
                        continue
                    if not os.path.isdir(project_dir):
                        continue
                    for item in os.listdir(project_dir):
                        item_path = os.path.join(project_dir, item)
                        if not os.path.isdir(item_path):
                            continue
                        if re.match(check["pattern"], item):
                            if is_folder_non_empty(item_path):
                                settlement_results[settlement_key][check["name"]] = "有资料"
                            break

            attachments = detect_attachment_files(sub_abs_path)
            rec = {
                "项目状态": project_status,
                "项目编码": project_code,
                "项目名称": project_name,
                "合同名称": os.path.basename(sub_abs_path),
                "合同编号": fields.get("合同编号", ""),
                "合同签订日期": fields.get("合同签订日期", ""),
                "币种": fields.get("币种", ""),
                "合同金额": fields.get("合同金额", 0.0),
                "有PDF": attachments["有PDF"],
                "有Word": attachments["有Word"],
                "有Excel": attachments["有Excel"],
                "_folder_path": sub_abs_path,
            }
            for k, v in fields.items():
                if k not in rec:
                    rec[k] = v
            all_data[hit_kw].append(rec)

    print("\n扫描完成：")
    for kw in target_keywords:
        cnt = project_counts.get(kw, 0)
        rec = len(all_data.get(kw, []))
        print(f" {kw}：{cnt} 个文件夹，{rec} 条记录")

    unmatched_projects = sorted(
        {
            rec.get("项目名称", "")
            for rows in all_data.values()
            for rec in rows
            if not str(rec.get("项目编码", "")).strip()
        }
    )
    if unmatched_projects:
        print("\n以下项目未在项目编码库匹配到编码：")
        for name in unmatched_projects:
            print(f" - {name}")

    return all_data, settlement_results, project_counts


def _choose_rate(row: dict) -> float:
    try:
        v = row.get("汇率")
        if v is not None and str(v).strip() != "":
            fv = float(v)
            if fv != 0 and fv != 1:
                return fv
    except (TypeError, ValueError):
        pass
    return 5 if row.get("币种") == "CNY" else 1


def build_detail_dataframes(
    all_data: dict[str, list[dict]],
    settlement_results: dict[str, dict],
) -> dict[str, pd.DataFrame]:
    """构建各明细 Sheet 的 DataFrame（含分类汇总与不分类两套）。"""
    fixed_cols = ["项目状态", "项目编码", "项目名称", "合同名称", "币种", "合同金额", "有PDF", "有Word", "有Excel"]
    settlement_cols = [
        check["name"] for check in SETTLEMENT_CHECKS if check["name"] not in OUTPUT_EXCLUDED_COLUMNS
    ]
    default_cols_order = fixed_cols + ["汇率", "金额（万泰铢）"] + settlement_cols

    sheets: dict[str, pd.DataFrame] = {}
    has_any = any(len(rows) > 0 for rows in all_data.values())
    if not has_any:
        sheets["说明"] = pd.DataFrame([{"提示": "未找到符合条件的 合同信息.json 文件"}])
        return sheets

    sorted_keywords = sorted(all_data.keys(), key=lambda x: (x != "总包合同", x))
    for kw in sorted_keywords:
        rows = all_data[kw]
        if not rows:
            continue

        df = pd.DataFrame(rows).fillna("")
        dynamic_cols = sorted(
            [
                c
                for c in df.columns
                if c not in fixed_cols and c != "_folder_path" and c not in OUTPUT_EXCLUDED_COLUMNS
            ]
        )
        for col in ["汇率", "金额（万泰铢）"] + settlement_cols:
            if col not in df.columns:
                df[col] = ""

        all_columns = list(
            set(fixed_cols + ["汇率", "金额（万泰铢）"] + settlement_cols + dynamic_cols)
        )
        default_order = [c for c in default_cols_order if c in all_columns] + [
            c for c in all_columns if c not in default_cols_order
        ]
        final_cols_order = load_column_order(all_columns, default_order)
        output_cols = [
            c for c in final_cols_order if c in df.columns and c not in OUTPUT_EXCLUDED_COLUMNS
        ]
        if not output_cols:
            continue

        df = df[output_cols]
        if "合同签订日期" in df.columns:
            df["_date_sort"] = df["合同签订日期"].apply(
                lambda x: str(x).strip() if str(x).strip() else "0000-00-00"
            )
            df = df.sort_values(by="_date_sort", ascending=False).drop(columns=["_date_sort"])

        detail_final = df.copy()

        grouped = []
        for (code, name), g in df.groupby(["项目编码", "项目名称"], sort=False):
            dates = (
                [d for d in g["合同签订日期"] if str(d).strip()]
                if "合同签订日期" in g.columns
                else []
            )
            min_date = min(dates) if dates else ""
            cur = "CNY" if "CNY" in g["币种"].values else (g["币种"].iloc[0] if not g.empty else "")
            settlement_key = str(code).strip() if str(code).strip() else name

            sum_row = {col: "" for col in g.columns}
            if "项目状态" in sum_row:
                sum_row["项目状态"] = g["项目状态"].iloc[0]
            if "项目编码" in sum_row:
                sum_row["项目编码"] = "" if code is None else str(code)
            if "项目名称" in sum_row:
                sum_row["项目名称"] = "" if name is None else str(name)
            if "合同名称" in sum_row:
                sum_row["合同名称"] = "（分类汇总）"
            if "币种" in sum_row:
                sum_row["币种"] = cur
            if "合同签订日期" in sum_row:
                sum_row["合同签订日期"] = min_date

            for check_name in settlement_cols:
                if check_name in sum_row:
                    sum_row[check_name] = settlement_results.get(str(settlement_key).strip(), {}).get(
                        check_name, ""
                    )

            for c in dynamic_cols:
                if c not in sum_row:
                    continue
                if any(k in c for k in ["金额", "数量", "面积"]):
                    sum_row[c] = pd.to_numeric(g[c], errors="coerce").sum()
                elif "日期" in c:
                    v = [x for x in g[c] if str(x).strip()]
                    sum_row[c] = min(v) if v else ""
                else:
                    sum_row[c] = ""

            grouped.append(pd.concat([g, pd.DataFrame([sum_row])], ignore_index=True))

        final = pd.concat(grouped, ignore_index=True) if grouped else df.copy()
        final["汇率"] = final.apply(lambda r: _choose_rate(r.to_dict()), axis=1)
        final["金额（万泰铢）"] = ""
        detail_final["汇率"] = detail_final.apply(lambda r: _choose_rate(r.to_dict()), axis=1)
        detail_final["金额（万泰铢）"] = ""

        sheets[kw[:31]] = final
        sheets[f"{kw}_不分类汇总"[:31]] = detail_final

    return sheets


def write_detail_sheets(writer: pd.ExcelWriter, sheets: dict[str, pd.DataFrame]) -> list[str]:
    """将明细 Sheet 写入已有 ExcelWriter，返回写入的 sheet 名列表。"""
    written = []
    for name, df in sheets.items():
        safe_name = name[:31]
        df.to_excel(writer, index=False, sheet_name=safe_name)
        written.append(safe_name)
    return written


def format_detail_sheets(workbook: Workbook, sheet_names: list[str] | None = None) -> None:
    """对明细类 Sheet 做格式化（高亮分类汇总、公式、分组）。"""
    targets = sheet_names if sheet_names is not None else list(workbook.sheetnames)
    for sheet_name in targets:
        if sheet_name not in workbook.sheetnames:
            continue
        ws = workbook[sheet_name]
        if ws.max_row == 1 and ws["A1"].value == "提示":
            continue

        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        highlight = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
        bold = Font(bold=True)

        col_map = {}
        for idx, col_name in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
            cell = col_name[0]
            col_map[cell.value] = get_column_letter(idx)

        rate_col = col_map.get("汇率")
        money_col = col_map.get("合同金额")
        target_col = col_map.get("金额（万泰铢）")
        subfolder_col = col_map.get("合同名称")
        proj_col = col_map.get("项目名称")

        for r_idx in range(2, ws.max_row + 1):
            sub_val = ws[f"{subfolder_col}{r_idx}"].value if subfolder_col else ""
            proj_val = ws[f"{proj_col}{r_idx}"].value if proj_col else ""

            if sub_val == "（分类汇总）":
                for c in ws[r_idx]:
                    c.fill = highlight
                    c.font = bold
                if target_col and proj_col:
                    start = r_idx
                    while start > 2 and ws[f"{proj_col}{start - 1}"].value == proj_val:
                        start -= 1
                    if r_idx - 1 >= start:
                        ws[f"{target_col}{r_idx}"] = f"=SUM({target_col}{start}:{target_col}{r_idx - 1})"
            else:
                if target_col and rate_col and money_col:
                    ws[f"{target_col}{r_idx}"] = f"=({money_col}{r_idx}*{rate_col}{r_idx}/10000)"

        if proj_col and subfolder_col:
            block_start = 2
            for r_idx in range(2, ws.max_row + 1):
                sub_val = ws[f"{subfolder_col}{r_idx}"].value
                if sub_val == "（分类汇总）":
                    detail_start = block_start
                    detail_end = r_idx - 1
                    if detail_end >= detail_start:
                        ws.row_dimensions.group(
                            detail_start, detail_end, outline_level=1, hidden=False
                        )
                    block_start = r_idx + 1
            ws.sheet_properties.outlinePr.summaryBelow = True
            ws.sheet_properties.outlinePr.applyStyles = True

        if money_col:
            for cell in ws[money_col][1:]:
                cell.number_format = "#,##0.00"
        if target_col:
            for cell in ws[target_col][1:]:
                cell.number_format = "#,##0.00"

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            row_idx = row[0].row
            if row_idx is not None:
                ws.row_dimensions[row_idx].height = 35

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions


def export_detail_only(output_path: str, root_dir: str | None = None) -> str:
    """单独导出明细汇总 Excel（兼容原脚本行为）。"""
    all_data, settlement_results, _ = scan_contract_details(root_dir=root_dir)
    sheets = build_detail_dataframes(all_data, settlement_results)
    if os.path.exists(output_path):
        try:
            os.remove(output_path)
        except PermissionError as e:
            raise PermissionError(f"文件被占用，请关闭后重试：{output_path}") from e

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        written = write_detail_sheets(writer, sheets)

    wb = load_workbook(output_path)
    format_detail_sheets(wb, written)
    wb.save(output_path)
    wb.close()
    print(f"明细汇总已保存：{output_path}")
    return output_path


def main() -> None:
    root = get_contract_root()
    if not os.path.isdir(root):
        raise FileNotFoundError(f"合同根目录无效：{root}")
    from summary_config import DETAIL_OUTPUT_NAME

    output_path = os.path.join(root, DETAIL_OUTPUT_NAME)
    export_detail_only(output_path, root_dir=root)
    print(f"成功生成：{output_path}")


if __name__ == "__main__":
    main()
