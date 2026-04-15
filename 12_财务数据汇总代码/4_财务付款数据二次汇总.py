import os
import re
import sys
import importlib.util

import pandas as pd
from openpyxl import load_workbook


CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_INFO_DIR = os.path.join(os.path.dirname(CURRENT_DIR), "0_项目基本信息")
if PROJECT_INFO_DIR not in sys.path:
    sys.path.append(PROJECT_INFO_DIR)


def load_apply_common_excel_layout():
    module_path = os.path.join(PROJECT_INFO_DIR, "excel_formatting_utils.py")
    if not os.path.exists(module_path):
        raise FileNotFoundError(f"未找到格式模块: {module_path}")

    spec = importlib.util.spec_from_file_location("excel_formatting_utils", module_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载格式模块: {module_path}")

    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.apply_common_excel_layout


APPLY_COMMON_EXCEL_LAYOUT = load_apply_common_excel_layout()


INPUT_ROOT = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/12-财务数据💰/A_财务付款数据/B_财务付款_标准化数据"
OUTPUT_ROOT = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/12-财务数据💰/A_财务付款数据/C_财务付款_汇总数据"
AMOUNT_COLUMNS = ["应付金额", "预扣税额", "实付金额", "增值税额", "不含税额"]


def parse_folder_tag(folder_name):
    # 从目录名中提取统计日期键（列头），如 20260406_财务付款数据 -> 20260406。
    match = re.search(r"(20\d{6})", folder_name)
    if match:
        return match.group(1)
    digits = "".join(ch for ch in str(folder_name) if ch.isdigit())
    return digits[:8] if len(digits) >= 8 else str(folder_name)


def iter_project_files(folder_path):
    for name in sorted(os.listdir(folder_path)):
        if name.startswith("~$"):
            continue
        if os.path.splitext(name)[1].lower() in {".xlsx", ".xls", ".xlsm"}:
            yield os.path.join(folder_path, name)


def safe_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"", "nan", "none"}:
        return ""
    return text


def extract_project_summary(file_path, ymd):
    try:
        df = pd.read_excel(file_path)
    except Exception as exc:
        print(f"  ⚠️ 读取失败: {os.path.basename(file_path)} | {exc}")
        return None

    if df.empty:
        return None

    if "实付金额" not in df.columns:
        print(f"  ⚠️ 缺少实付金额列，已跳过: {os.path.basename(file_path)}")
        return None

    # 金额字段合计为0时，不参与二次汇总。
    existing_amount_cols = [col for col in AMOUNT_COLUMNS if col in df.columns]
    if existing_amount_cols:
        amount_df = df[existing_amount_cols].apply(pd.to_numeric, errors="coerce")
        row_amount_sum = amount_df.fillna(0).sum(axis=1)
        valid_mask = row_amount_sum.ne(0)
        df = df[valid_mask]
        if df.empty:
            return None

    project_code = ""
    if "项目编码" in df.columns:
        code_series = df["项目编码"].dropna().astype(str).str.strip()
        code_series = code_series[~code_series.str.lower().isin(["", "nan", "none"])]
        if not code_series.empty:
            project_code = code_series.iloc[0]

    project_name = ""
    if "项目名称" in df.columns:
        name_series = df["项目名称"].dropna().astype(str).str.strip()
        name_series = name_series[~name_series.str.lower().isin(["", "nan", "none"])]
        if not name_series.empty:
            project_name = name_series.iloc[0]

    if not project_name:
        project_name = os.path.splitext(os.path.basename(file_path))[0]

    paid_amount = pd.to_numeric(df["实付金额"], errors="coerce").fillna(0)
    vat_amount = pd.Series([0] * len(df))
    if "增值税额" in df.columns:
        vat_amount = pd.to_numeric(df["增值税额"], errors="coerce").fillna(0)

    amount = (paid_amount + vat_amount).sum()

    return {
        "项目编码": safe_text(project_code),
        "项目名称": safe_text(project_name),
        "年月日": ymd,
        "付款总金额": float(amount),
    }


def build_summary_dataframe():
    if not os.path.isdir(INPUT_ROOT):
        raise FileNotFoundError(f"输入目录不存在: {INPUT_ROOT}")

    rows = []
    folder_names = [
        name
        for name in sorted(os.listdir(INPUT_ROOT))
        if os.path.isdir(os.path.join(INPUT_ROOT, name))
    ]

    for folder_name in folder_names:
        folder_path = os.path.join(INPUT_ROOT, folder_name)
        ymd = parse_folder_tag(folder_name)
        print(f"\n处理目录: {folder_name}")

        for file_path in iter_project_files(folder_path):
            row = extract_project_summary(file_path, ymd)
            if row is not None:
                rows.append(row)

    if not rows:
        return pd.DataFrame(columns=["项目编码", "项目名称"])

    detail_df = pd.DataFrame(rows)
    detail_df = (
        detail_df.groupby(["项目名称", "年月日"], dropna=False, as_index=False)
        .agg(
            项目编码=(
                "项目编码",
                lambda s: next((v for v in s if safe_text(v) != ""), ""),
            ),
            付款总金额=("付款总金额", "sum"),
        )
    )

    # 项目编码只做展示：同一项目名称下优先取非空编码。
    project_code_df = (
        detail_df.groupby("项目名称", dropna=False, as_index=False)
        .agg(
            项目编码=(
                "项目编码",
                lambda s: next((v for v in s if safe_text(v) != ""), ""),
            )
        )
    )

    pivot_df = detail_df.pivot_table(
        index="项目名称",
        columns="年月日",
        values="付款总金额",
        aggfunc="sum",
    )
    pivot_df = pivot_df.reset_index()

    date_cols = [col for col in pivot_df.columns if col != "项目名称"]
    date_cols = sorted(date_cols)
    pivot_df = pivot_df[["项目名称"] + date_cols]

    out_df = project_code_df.merge(pivot_df, on="项目名称", how="left")
    out_df = out_df[["项目编码", "项目名称"] + date_cols]
    return out_df


def build_increase_sheet(df):
    if df.empty:
        return df.copy()

    inc_df = df.copy()
    date_cols = [col for col in inc_df.columns if col not in ["项目编码", "项目名称"]]

    if not date_cols:
        return inc_df

    # 使用原始金额快照计算差值，避免前面写入"-"后污染后续计算。
    base_numeric = {
        col: pd.to_numeric(df[col], errors="coerce") for col in date_cols
    }

    for idx, col in enumerate(date_cols):
        if idx == 0:
            inc_df[col] = "-"
            continue

        prev_col = date_cols[idx - 1]
        cur_vals = base_numeric[col]
        prev_vals = base_numeric[prev_col]
        diff_vals = cur_vals - prev_vals

        # 仅当当前与上期都存在金额时计算差值，否则显示“-”。
        has_both = cur_vals.notna() & prev_vals.notna()
        inc_df[col] = diff_vals.where(has_both, "-")

    return inc_df


def apply_output_format(output_path):
    wb = load_workbook(output_path)

    ws_total = wb["付款总金额汇总"]
    total_headers = [cell.value for cell in ws_total[1]]
    total_amount_headers = total_headers[2:]
    APPLY_COMMON_EXCEL_LAYOUT(
        ws_total,
        numeric_start_col=ws_total.max_column + 1,
        amount_headers=total_amount_headers,
        header_row=1,
        header_height=40,
        freeze_panes="A2",
    )

    ws_inc = wb["较上期增加汇总"]
    inc_headers = [cell.value for cell in ws_inc[1]]
    inc_amount_headers = inc_headers[2:]
    APPLY_COMMON_EXCEL_LAYOUT(
        ws_inc,
        numeric_start_col=ws_inc.max_column + 1,
        amount_headers=inc_amount_headers,
        header_row=1,
        header_height=40,
        freeze_panes="A2",
    )

    wb.save(output_path)


def main():
    summary_df = build_summary_dataframe()
    increase_df = build_increase_sheet(summary_df)

    os.makedirs(OUTPUT_ROOT, exist_ok=True)
    output_path = os.path.join(OUTPUT_ROOT, "财务付款数据汇总.xlsx")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="付款总金额汇总", index=False)
        increase_df.to_excel(writer, sheet_name="较上期增加汇总", index=False)

    apply_output_format(output_path)

    print("\n处理完成")
    print(f"输出文件: {output_path}")
    print(f"付款总金额汇总行数: {len(summary_df)}")
    print(f"较上期增加汇总行数: {len(increase_df)}")


if __name__ == "__main__":
    main()
