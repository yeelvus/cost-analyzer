import os
import re
import sys
import importlib.util

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_INFO_DIR = os.path.join(os.path.dirname(CURRENT_DIR), "0_项目基本信息")
if PROJECT_INFO_DIR not in sys.path:
    sys.path.append(PROJECT_INFO_DIR)


def load_apply_common_excel_layout():
    module_path = os.path.join(PROJECT_INFO_DIR, "excel_formatting_utils.py")
    if not os.path.exists(module_path):
        raise FileNotFoundError(f"未找到格式模块: {module_path}")

    spec = importlib.util.spec_from_file_location("excel_formatting_utils", module_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载格式模块: {module_path}")

    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.apply_common_excel_layout


def load_project_name_matcher_module():
    module_path = os.path.join(PROJECT_INFO_DIR, "project_name_matcher.py")
    if not os.path.exists(module_path):
        raise FileNotFoundError(f"未找到项目匹配模块: {module_path}")

    spec = importlib.util.spec_from_file_location("project_name_matcher", module_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载项目匹配模块: {module_path}")

    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def load_subject_mapper_class():
    module_path = os.path.join(CURRENT_DIR, "subject_mapper.py")
    if not os.path.exists(module_path):
        raise FileNotFoundError(f"未找到摘要匹配模块: {module_path}")

    spec = importlib.util.spec_from_file_location("subject_mapper", module_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载摘要匹配模块: {module_path}")

    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.SubjectMapper


APPLY_COMMON_EXCEL_LAYOUT = load_apply_common_excel_layout()
PROJECT_NAME_MATCHER = load_project_name_matcher_module()
SUBJECT_MAPPER_CLASS = load_subject_mapper_class()


INPUT_ROOT = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/12-财务部💰/A_财务付款数据/A_原始数据"
OUTPUT_ROOT = "/Users/wyattsun/Desktop/邯郸中材/01成本统计/12-财务部💰/A_财务付款数据/B_标准化数据"
PROJECT_MATCH_DB_PATH = os.path.join(PROJECT_INFO_DIR, "project_name_match_db.json")
SUBJECT_RULES_PATH = os.path.join(CURRENT_DIR, "rules.json")


TARGET_COLUMNS = [
    "项目名称",
    "科目",
    "摘要",
    "供应商名称",
    "应付金额",
    "预扣税额",
    "实付金额",
    "增值税额",
    "不含税额",
    "日期",
    "月",
    "年",
]


MIN_REQUIRED_COLUMNS = ["项目名称", "摘要", "实付金额"]


AMOUNT_COLUMNS = ["应付金额", "预扣税额", "实付金额", "增值税额", "不含税额"]

DEFAULT_CATEGORY = "6_管理费"
LEVEL1_ORDER = ["0_分包", "2_材料采购", "1_人工费", "3_机具费", "6_管理费"]


COLUMN_ALIASES = {
    "项目名称": ["项目名称", "项目", "项目名", "工程名称", "工程项目"],
    "科目": ["科目", "会计科目", "财务科目"],
    "摘要": ["摘要", "业务摘要", "凭证摘要", "说明"],
    "供应商名称": ["供应商名称", "供应商", "收款单位", "单位名称", "往来单位"],
    "应付金额": ["应付金额", "应付", "应付款", "应支付金额"],
    "预扣税额": ["预扣税额", "预扣税", "扣税", "税额"],
    "实付金额": ["实付金额", "实付", "实际付款", "支付金额", "付款金额"],
    "增值税额": ["增值税额", "增值税", "进项税额", "销项税额"],
    "不含税额": ["不含税额", "不含税金额", "未税金额", "税前金额"],
    "日期": ["日期", "付款日期", "业务日期", "发生日期"],
    "月": ["月", "月份", "所属月份", "会计期间(月)"],
    "年": ["年", "年份", "所属年份", "会计年度"],
}


def normalize_text(text):
    if text is None:
        return ""
    s = str(text).strip()
    s = s.replace("\n", "").replace("\r", "")
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"[()（）\[\]【】:：_\-]", "", s)
    return s.lower()


def extract_time_tag(folder_name):
    match = re.search(r"(20\d{6})", folder_name)
    if match:
        return match.group(1)
    match = re.search(r"(\d{6,8})", folder_name)
    if match:
        return match.group(1)
    return folder_name


def extract_month_tag(time_tag):
    digits = "".join(ch for ch in str(time_tag) if ch.isdigit())
    if len(digits) >= 6:
        return digits[:6]
    return str(time_tag)


def safe_filename(name):
    if name is None:
        return "未命名项目"
    text = str(name).strip()
    if not text:
        return "未命名项目"
    text = re.sub(r'[\\/:*?"<>|]', "_", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:80]


def normalize_group_key(key):
    if isinstance(key, tuple):
        if len(key) == 0:
            return ""
        return key[0]
    return key


def excel_sheet_ref(sheet_name):
    return "'" + sheet_name.replace("'", "''") + "'"


def _level1_sort_key(level1):
    text = str(level1).strip()
    m = re.match(r"^(\d+)_", text)
    if m:
        return (0, int(m.group(1)), text)
    return (1, 9999, text)


def init_subject_matcher():
    mapper = SUBJECT_MAPPER_CLASS(SUBJECT_RULES_PATH)
    level1_values = sorted(
        {
            str(rule.get("一级科目", "")).strip()
            for rule in mapper.rules
            if str(rule.get("一级科目", "")).strip()
        },
        key=_level1_sort_key,
    )

    default_level1 = next((x for x in level1_values if "管理费" in x), "6_管理费")
    ordered = [x for x in level1_values if x != default_level1] + [default_level1]
    return mapper, ordered if ordered else LEVEL1_ORDER, default_level1


SUBJECT_MAPPER, LEVEL1_ORDER, DEFAULT_CATEGORY = init_subject_matcher()


def classify_subject_levels(subject_value, summary_value):
    raw_subject = "" if subject_value is None else str(subject_value).strip()

    try:
        level1, level2 = SUBJECT_MAPPER.classify(subject_value, summary_value)
    except Exception:
        level1, level2 = "", ""

    level1 = str(level1).strip() if level1 is not None else ""
    level2 = str(level2).strip() if level2 is not None else ""

    # 兼容旧版本 mapper：未命中时可能返回 (原始科目, 原始科目)。
    if raw_subject and level1 == raw_subject and level2 == raw_subject:
        level1 = ""

    if not level1:
        level1 = DEFAULT_CATEGORY

    if not level2:
        level2 = raw_subject if raw_subject else "(空白)"

    return level1, level2


def append_subject_levels(df):
    enriched = df.copy()
    if "科目" not in enriched.columns and "摘要" not in enriched.columns:
        enriched["一级科目"] = DEFAULT_CATEGORY
        enriched["二级科目"] = "(空白)"
        return enriched

    level_pairs = enriched.apply(
        lambda row: classify_subject_levels(row.get("科目"), row.get("摘要")),
        axis=1,
    )
    enriched["一级科目"] = level_pairs.apply(lambda x: x[0])
    enriched["二级科目"] = level_pairs.apply(lambda x: x[1])

    # 将科目分级列放到“摘要”之前，便于查看分类依据。
    cols = [c for c in enriched.columns if c not in {"一级科目", "二级科目"}]
    if "摘要" in cols:
        idx = cols.index("摘要")
        cols = cols[:idx] + ["一级科目", "二级科目"] + cols[idx:]
    else:
        cols = cols + ["一级科目", "二级科目"]
    enriched = enriched[cols]

    return enriched


def build_detail_wan_sheet(detail_df):
    wan_df = detail_df.copy()
    for col in AMOUNT_COLUMNS:
        if col in wan_df.columns:
            wan_df[col] = pd.to_numeric(wan_df[col], errors="coerce") / 10000.0
    return wan_df


def find_column_mapping(df):
    mapping = {}
    normalized_to_original = {normalize_text(col): col for col in df.columns}

    for target in TARGET_COLUMNS:
        candidates = COLUMN_ALIASES.get(target, [target])
        found_col = None

        for alias in candidates:
            norm_alias = normalize_text(alias)
            if norm_alias in normalized_to_original:
                found_col = normalized_to_original[norm_alias]
                break

        if found_col is None:
            for alias in candidates:
                norm_alias = normalize_text(alias)
                for norm_col, origin_col in normalized_to_original.items():
                    if norm_alias and (norm_alias in norm_col or norm_col in norm_alias):
                        found_col = origin_col
                        break
                if found_col is not None:
                    break

        mapping[target] = found_col

    return mapping


def to_numeric(series):
    cleaned = (
        series.astype(str)
        .str.replace(",", "", regex=False)
        .str.replace(" ", "", regex=False)
        .replace({"": pd.NA, "None": pd.NA, "nan": pd.NA, "NaN": pd.NA})
    )
    return pd.to_numeric(cleaned, errors="coerce")


def normalize_date_month_year(df):
    if "日期" in df.columns:
        parsed_date = pd.to_datetime(df["日期"], errors="coerce")
        df["日期"] = parsed_date.dt.strftime("%Y-%m-%d").where(parsed_date.notna(), df["日期"])

        if "年" in df.columns:
            year_numeric = pd.to_numeric(df["年"], errors="coerce")
            df["年"] = year_numeric.where(year_numeric.notna(), parsed_date.dt.year)
        if "月" in df.columns:
            month_numeric = pd.to_numeric(df["月"], errors="coerce")
            df["月"] = month_numeric.where(month_numeric.notna(), parsed_date.dt.month)

    if "年" in df.columns:
        df["年"] = pd.to_numeric(df["年"], errors="coerce").astype("Int64")
    if "月" in df.columns:
        df["月"] = pd.to_numeric(df["月"], errors="coerce").astype("Int64")

    return df


def load_project_matcher_context():
    db = PROJECT_NAME_MATCHER.load_or_init_db(PROJECT_MATCH_DB_PATH, [])
    return {"db": db}


def enrich_project_fields(df, matcher_ctx):
    db = matcher_ctx["db"]
    raw_series = df["项目名称"].fillna("").astype(str).str.strip()
    unique_names = sorted(set(name for name in raw_series.tolist() if name and name.lower() != "nan"))

    unmatched_names = []
    code_map = {}
    name_map = {}
    for raw_name in unique_names:
        match_result = PROJECT_NAME_MATCHER.match_project(raw_name, db)
        if match_result:
            code_map[raw_name] = match_result.get("project_code", "")
            name_map[raw_name] = match_result.get("project_name", raw_name)
        else:
            PROJECT_NAME_MATCHER.mark_unmatched(db, raw_name)
            code_map[raw_name] = ""
            name_map[raw_name] = raw_name
            unmatched_names.append(raw_name)

    # 对未匹配项目弹窗，手动选择后回写数据库并再次匹配。
    if unmatched_names:
        print(f"  ⚠️ 检测到未匹配项目 {len(unmatched_names)} 个，弹出手动匹配窗口...")
        PROJECT_NAME_MATCHER.prompt_unmatched_project_mappings(
            unmatched_names,
            db,
            window_title="财务付款数据项目匹配",
        )

        for raw_name in unmatched_names:
            retry_match = PROJECT_NAME_MATCHER.match_project(raw_name, db)
            if retry_match:
                code_map[raw_name] = retry_match.get("project_code", "")
                name_map[raw_name] = retry_match.get("project_name", raw_name)

    enriched = df.copy()
    enriched["项目编码"] = raw_series.map(code_map).fillna("")
    enriched["项目名称"] = raw_series.map(name_map).fillna(raw_series)
    enriched = enriched[["项目编码"] + TARGET_COLUMNS]

    PROJECT_NAME_MATCHER.save_db(PROJECT_MATCH_DB_PATH, db)
    return enriched


def process_single_sheet(file_path, sheet_name):
    try:
        raw_df = pd.read_excel(file_path, sheet_name=sheet_name)
    except Exception as exc:
        print(f"  ⚠️ 读取失败: {os.path.basename(file_path)} | {sheet_name} | {exc}")
        return None

    if raw_df.empty:
        return None

    mapping = find_column_mapping(raw_df)

    missing_required = [col for col in MIN_REQUIRED_COLUMNS if mapping.get(col) is None]
    if missing_required:
        print(f"    - 跳过sheet: {sheet_name}，缺少关键列: {', '.join(missing_required)}")
        return None

    out_df = pd.DataFrame()
    for col in TARGET_COLUMNS:
        src = mapping.get(col)
        out_df[col] = raw_df[src] if src is not None else pd.NA

    for amount_col in AMOUNT_COLUMNS:
        if amount_col in out_df.columns:
            out_df[amount_col] = to_numeric(out_df[amount_col])

    # 五个金额字段合计为0的记录不参与汇总。
    row_amount_sum = out_df[AMOUNT_COLUMNS].fillna(0).sum(axis=1)
    out_df = out_df[row_amount_sum.ne(0)]

    out_df = normalize_date_month_year(out_df)
    out_df = out_df.dropna(how="all", subset=TARGET_COLUMNS)
    if out_df.empty:
        return None

    return out_df


def iter_excel_files(folder_path):
    for name in sorted(os.listdir(folder_path)):
        if name.startswith("~$"):
            continue
        ext = os.path.splitext(name)[1].lower()
        if ext in {".xlsx", ".xls", ".xlsm"}:
            yield os.path.join(folder_path, name)


def process_time_folder(folder_path):
    folder_name = os.path.basename(os.path.normpath(folder_path))
    print(f"\n{'=' * 60}")
    print(f"开始处理时间目录: {folder_name}")

    result_parts = []
    excel_files = list(iter_excel_files(folder_path))
    if not excel_files:
        print("  ⚠️ 未发现Excel文件，跳过")
        return []

    for file_path in excel_files:
        file_name = os.path.basename(file_path)
        print(f"  处理文件: {file_name}")

        try:
            sheet_names = pd.ExcelFile(file_path).sheet_names
        except Exception as exc:
            print(f"    ⚠️ 获取sheet失败: {exc}")
            continue

        for sheet_name in sheet_names:
            part = process_single_sheet(file_path, sheet_name)
            if part is None:
                continue
            result_parts.append(part)

    if not result_parts:
        print("  ⚠️ 未提取到有效数据")
        return []

    final_df = pd.concat(result_parts, ignore_index=True)
    final_df = enrich_project_fields(final_df, load_project_matcher_context())
    final_df = append_subject_levels(final_df)

    output_dir = os.path.join(OUTPUT_ROOT, folder_name)
    os.makedirs(output_dir, exist_ok=True)

    output_paths = []

    grouped = final_df.groupby(["项目名称"], dropna=False, sort=True)
    for project_name_key, project_df in grouped:
        project_name = normalize_group_key(project_name_key)
        project_part = safe_filename(project_name)
        output_name = f"{project_part}_{folder_name}.xlsx"
        output_path = os.path.join(output_dir, output_name)

        detail_sheet_name = "明细数据"
        detail_wan_sheet_name = "明细数据_万泰铢"
        detail_wan_df = build_detail_wan_sheet(project_df)

        if os.path.exists(output_path):
            updated = SUBJECT_MAPPER.update_excel_subject_columns(
                output_path,
                sheet_names=[detail_wan_sheet_name, detail_sheet_name],
                subject_col="科目",
                summary_col="摘要",
                level1_col="一级科目",
                level2_col="二级科目",
            )
            print(f"  ♻️ 已存在文件，仅更新科目分类列：{os.path.basename(output_path)} | 更新行数: {updated}")
            output_paths.append(output_path)
            continue

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            detail_wan_df.to_excel(writer, sheet_name=detail_wan_sheet_name, index=False)
            project_df.to_excel(writer, sheet_name=detail_sheet_name, index=False)

        apply_output_format(output_path)
        output_paths.append(output_path)

    print(f"  ✅ 已输出目录: {output_dir}")
    print(f"  📊 项目文件数: {len(output_paths)}")
    print(f"  📊 数据行数: {len(final_df)}")
    return output_paths


def apply_output_format(output_path):
    wb = load_workbook(output_path)
    if not wb.worksheets:
        wb.save(output_path)
        return

    amount_headers = set(AMOUNT_COLUMNS)
    for ws in wb.worksheets:
        APPLY_COMMON_EXCEL_LAYOUT(
            ws,
            numeric_start_col=1,
            amount_headers=amount_headers,
            header_row=1,
            header_height=40,
            freeze_panes="A2",
        )

    wb.save(output_path)


def main():
    if not os.path.isdir(INPUT_ROOT):
        raise FileNotFoundError(f"输入目录不存在: {INPUT_ROOT}")

    time_folders = [
        os.path.join(INPUT_ROOT, name)
        for name in sorted(os.listdir(INPUT_ROOT))
        if os.path.isdir(os.path.join(INPUT_ROOT, name))
    ]

    if not time_folders:
        print("未发现时间子目录，程序结束")
        return

    outputs = []
    for folder_path in time_folders:
        outputs.extend(process_time_folder(folder_path))

    print(f"\n{'=' * 60}")
    print("处理完成")
    print(f"成功输出文件数: {len(outputs)}")
    for path in outputs:
        print(f"- {path}")


if __name__ == "__main__":
    main()
